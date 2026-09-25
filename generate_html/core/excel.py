from __future__ import annotations

import base64
import logging
import gzip
import hashlib
import json
import os
import tempfile
from itertools import groupby
import re
import weakref
from statistics import median
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles.numbers import is_date_format


GRAIN_MAP = {"时": "hour", "天": "day", "周": "week", "月": "month"}
GRAIN_LABELS = {value: key for key, value in GRAIN_MAP.items()}
LOGGER = logging.getLogger(__name__)


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\u2011", "-")).strip()


def compact_text(value: Any) -> str:
    return re.sub(r"\s+", "", clean_text(value))


def compact_identifier(value: Any) -> str:
    """Normalize separators used inconsistently in workbook and Sheet names."""
    return re.sub(r"[\s_\-—–:：/\\（）()【】\[\]]+", "", clean_text(value)).lower()


def display_period(value: Any, number_format: str = "") -> str:
    if isinstance(value, (datetime, date)):
        return format_excel_value(value, number_format or "yyyy-mm-dd")
    return clean_text(value)


def safe_number(value: Any, default: float = 0.0) -> float:
    if value in (None, ""):
        return default
    if isinstance(value, str):
        text = clean_text(value)
        negative = text.startswith("(") and text.endswith(")")
        percent = "%" in text or "％" in text
        text = text.replace(",", "").replace("，", "").replace("%", "").replace("％", "")
        text = re.sub(r"(?:人民币|RMB|CNY|元|万元|亿元|￥|¥|\$)", "", text, flags=re.IGNORECASE).strip(" ()")
        try:
            number = float(text)
            if negative:
                number = -number
            return number / 100 if percent else number
        except ValueError:
            return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def cell_number(cell: Any, rate: bool = False, default: float = 0.0) -> float:
    """Read a number while respecting percent-like strings and cell formats."""
    value = safe_number(cell.value, default)
    if not rate or cell.value in (None, ""):
        return value
    if isinstance(cell.value, str):
        text = clean_text(cell.value)
        return value / 100 if "%" not in text and "％" not in text and abs(value) > 1.5 else value
    number_format = clean_text(getattr(cell, "number_format", ""))
    if "%" not in number_format and abs(value) > 1.5:
        return value / 100
    return value


def _format_precision(number_format: str, marker: str = "") -> int:
    section = number_format.split(";", 1)[0]
    if marker and marker in section:
        section = section.split(marker, 1)[0]
    match = re.search(r"[0#](?:\.([0#]+))?[^0#]*$", section)
    return len(match.group(1) or "") if match else 0


def _format_date(value: date | datetime, number_format: str) -> str:
    fmt = number_format.lower().replace("\\", "")
    has_time = any(token in fmt for token in ("h", "s", "am/pm"))
    has_day = "d" in fmt
    has_month = "m" in fmt
    has_year = "y" in fmt
    separator = "/" if "/" in fmt else "." if "." in fmt else "-"
    if has_year and has_month and has_day:
        result = value.strftime(f"%Y{separator}%m{separator}%d")
    elif has_year and has_month:
        result = value.strftime(f"%Y{separator}%m")
    elif has_month and has_day:
        result = value.strftime(f"%m{separator}%d")
    else:
        result = value.isoformat()
    if has_time and isinstance(value, datetime):
        time_text = value.strftime("%H:%M:%S" if "s" in fmt else "%H:%M")
        result = f"{result} {time_text}" if has_year or has_month or has_day else time_text
    return result


def format_excel_value(value: Any, number_format: str = "General") -> str:
    """Render common Excel formats without assuming fixed precision or date syntax."""
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return _format_date(value, number_format)
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if not isinstance(value, (int, float)):
        return clean_text(value)
    fmt = clean_text(number_format or "General")
    if fmt.lower() == "general" or fmt == "@":
        return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
    if is_date_format(fmt):
        return str(value)
    if "%" in fmt:
        precision = _format_precision(fmt, "%")
        return f"{value * 100:,.{precision}f}%"
    precision = _format_precision(fmt)
    grouping = "," in fmt.split(";", 1)[0]
    rendered = f"{value:,.{precision}f}" if grouping else f"{value:.{precision}f}"
    currency = next((symbol for symbol in ("￥", "¥", "$", "€", "£") if symbol in fmt), "")
    return f"{currency}{rendered}"


def format_excel_cell(cell: Any) -> str:
    return format_excel_value(cell.value, getattr(cell, "number_format", "General"))


def safe_rate(numerator: Any, denominator: Any) -> float:
    den = safe_number(denominator)
    return safe_number(numerator) / den if den else 0.0


def sheet_subject(name: str) -> str:
    value = clean_text(name)
    patterns = [
        r"净(?:大定|小订)by[时天周月](?:第\s*\d+\s*期)?图表$",
        r"by[时天周月](?:第\s*\d+\s*期)?图表$",
        r"by[时天周月](?:第\s*\d+\s*期)?$",
        r"_(SKU|组合|日度退订|选配退订|小订分时退订|选配渗透率排名|付费选配占比|选配渗透率)$",
    ]
    for pattern in patterns:
        value = re.sub(pattern, "", value, flags=re.IGNORECASE)
    # Generic analysis sheets may append a free-form suffix. Keep the suffix as
    # the analysis name while resolving the sheet back to its known group,
    # brand, or generation subject.
    subject_prefixes = [
        # “总计/汇总”可以是代际名称本身的一部分，例如“问界 M9 2026款总计”。
        # 只清理它后面的分析后缀，不要把它误当作汇总周期截掉。
        r"^(.+?20\d{2}\s*款(?:\s*(?:总计|汇总))?)\s*[_\-—]\s*.+$",
        # 代号型代际（无款号），例如“问界 F2N”。
        r"^([\u4e00-\u9fff]{1,3}界\s*[A-Za-z0-9]+(?:\s*20\d{2}\s*款)?(?:\s*(?:总计|汇总))?)\s*[_\-—]\s*.+$",
        r"^(鸿蒙智行)\s*[_\-—]\s*.+$",
        r"^([\u4e00-\u9fff]{1,3}界)\s*[_\-—]\s*.+$",
    ]
    for pattern in subject_prefixes:
        match = re.match(pattern, value, flags=re.IGNORECASE)
        if match:
            value = match.group(1)
            break
    return clean_text(value.rstrip("_"))


def is_aggregate_generation(name: str) -> bool:
    """以“总计/汇总”结尾的代际是多个代际的聚合视图，品牌级统计时须排除以免重复。"""
    return clean_text(name).endswith(("总计", "汇总"))


def subject_type(name: str) -> str:
    if compact_text(name) == "鸿蒙智行":
        return "group"
    if re.search(r"20\d{2}\s*款", name):
        return "generation"
    if re.fullmatch(r"[\u4e00-\u9fff]{1,3}界\s*[A-Za-z0-9]+", compact_text(name)):
        return "generation"
    return "brand"


def subject_parent(name: str, kind: str) -> str | None:
    if kind == "brand":
        return "鸿蒙智行"
    if kind != "generation":
        return None
    match = re.match(r"^([\u4e00-\u9fff]{1,3}界)", compact_text(name))
    return match.group(1) if match else None


def grain_from_sheet(name: str) -> str | None:
    match = re.search(r"by([时天周月])", name)
    return GRAIN_MAP.get(match.group(1)) if match else None



_DATA_CACHE_VERSION = 3


def load_data_workbook(path: Path, cache_dir: Path | None = None):
    """Return cached values/formats without retaining drawings/decorative styles.

    Cache is derived JSON (not executable pickle), keyed by the exact source
    identity, nanosecond timestamps and size. A failed cache is always optional.
    """
    path = Path(path)
    stat = path.stat()
    signature = [_DATA_CACHE_VERSION, str(path.resolve()), stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns]
    cache = None
    if cache_dir is not None:
        cache = Path(cache_dir) / (hashlib.sha256(str(path.resolve()).encode()).hexdigest() + ".jsonl.gz")

    def decode(value, kind):
        if kind == "datetime":
            return datetime.fromisoformat(value)
        if kind == "date":
            return date.fromisoformat(value)
        if kind == "time":
            from datetime import time
            return time.fromisoformat(value)
        if kind == "timedelta":
            from datetime import timedelta
            return timedelta(seconds=value)
        return value

    def restore(stream):
        if json.loads(stream.readline()) != signature:
            raise ValueError("stale workbook cache")
        result = Workbook()
        result.remove(result.active)
        sheet = None
        complete = False
        for line in stream:
            row = json.loads(line)
            if row == ["end"]:
                complete = True
                break
            if row[0] == "sheet":
                sheet = result.create_sheet(row[1])
                sheet.sheet_state = row[2]
            elif row[0] == "row" and sheet is not None:
                for col, value, kind, data_type, number_format in row[2]:
                    cell = sheet.cell(row[1], col, decode(value, kind))
                    cell.data_type = data_type
                    cell.number_format = number_format
            else:
                raise ValueError("invalid workbook cache")
        if not result.worksheets or not complete:
            raise ValueError("incomplete workbook cache")
        return result

    if cache is not None and cache.exists():
        try:
            with gzip.open(cache, "rt", encoding="utf-8") as stream:
                result = restore(stream)
            LOGGER.debug("Excel解析缓存命中: %s", path.name)
            return result
        except (OSError, EOFError, ValueError, TypeError, IndexError, KeyError):
            LOGGER.debug("Excel解析缓存失效，重新读取: %s", path.name)

    # The normal reader applies Excel merged-cell semantics. Read-only mode
    # exposes hidden XML values inside merged ranges and can change real orders.
    # Only cache misses pay this cost; the returned view and cache stay compact.
    source = load_workbook(path, data_only=True, read_only=False, keep_links=False)
    result = Workbook()
    result.remove(result.active)
    pending = None
    writer = None
    try:
        if cache is not None:
            try:
                cache.parent.mkdir(parents=True, exist_ok=True)
                fd, name = tempfile.mkstemp(prefix=cache.stem + ".", suffix=".tmp", dir=cache.parent)
                os.close(fd)
                pending = Path(name)
                writer = gzip.open(pending, "wt", encoding="utf-8", compresslevel=1)
                writer.write(json.dumps(signature) + "\n")
            except OSError:
                LOGGER.debug("Excel解析缓存不可写，将直接读取: %s", path.name)

        def write(record):
            nonlocal writer
            if writer is not None:
                try:
                    writer.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n")
                except OSError:
                    writer.close()
                    writer = None

        from datetime import time, timedelta
        for raw in source.worksheets:
            sheet = result.create_sheet(raw.title)
            sheet.sheet_state = raw.sheet_state
            write(["sheet", sheet.title, sheet.sheet_state])
            # Iterate stored cells, never materialize a huge styled blank grid.
            for row_number, cells in groupby(raw._cells.values(), key=lambda cell: cell.row):
                entries = []
                for cell in cells:
                    column, value = cell.column, cell.value
                    if value is None:
                        continue
                    dest = sheet.cell(row_number, column, value)
                    dest.data_type = cell.data_type
                    dest.number_format = cell.number_format
                    kind = ""
                    if isinstance(value, datetime):
                        value, kind = value.isoformat(), "datetime"
                    elif isinstance(value, date):
                        value, kind = value.isoformat(), "date"
                    elif isinstance(value, time):
                        value, kind = value.isoformat(), "time"
                    elif isinstance(value, timedelta):
                        value, kind = value.total_seconds(), "timedelta"
                    entries.append([column, value, kind, cell.data_type, cell.number_format])
                if entries:
                    write(["row", row_number, entries])
        write(["end"])
        if writer is not None:
            writer.close()
            writer = None
            # Never commit a cache for a source edited while it was being read.
            after = path.stat()
            if (after.st_size, after.st_mtime_ns, after.st_ctime_ns) == (stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns):
                try:
                    os.replace(pending, cache)
                    pending = None
                except OSError:
                    LOGGER.debug("Excel解析缓存提交失败，当前读取结果仍有效: %s", path.name)
        return result
    finally:
        source.close()
        if writer is not None:
            writer.close()
        if pending is not None:
            pending.unlink(missing_ok=True)


def _source_key(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat().replace("T00:00:00", "")
    return clean_text(value)


def _period_sort(value, *, keep_aggregate_order=False):
    text = str(value)
    aggregate = any(word in text for word in ("总计", "汇总", "累计", "近"))
    if aggregate and keep_aggregate_order:
        return (True, ())  # Stable sort retains source ordering among summary columns.
    return (aggregate, tuple((0, int(part)) if part.isdigit() else (1, part)
                             for part in re.split(r"(\d+)", text)))


def _merge_matrix(matches, header_rows, key_columns, *, forward_rows=(), forward_headers=(), time_row=None,
                  column_group_row=None, sparse_headers=()):
    """Union an identified table by business labels; first nonblank value wins.

    Used only across files, never to add totals or combine different launch phases.
    Physical originals stay untouched and remain available in the raw-table centre.
    """
    rows, columns, values = {}, {}, {}
    conflicts = []
    conflict_count = 0
    for item, sheet in matches:
        row_keys, col_keys = {}, {}
        inherited = {}
        for row in range(header_rows + 1, sheet.max_row + 1):
            labels = []
            for col in range(1, key_columns + 1):
                cell = sheet.cell(row, col)
                value = cell.value
                if col in forward_rows:
                    if value not in (None, ""):
                        inherited[col] = value
                    value = inherited.get(col) if value in (None, "") else value
                labels.append(value)
            if not any(value not in (None, "") for value in labels):
                continue
            key = tuple(_source_key(value) for value in labels)
            row_keys[row] = key
            rows.setdefault(key, labels)
        inherited = {}
        for col in range(key_columns + 1, sheet.max_column + 1):
            labels = []
            for row in range(1, header_rows + 1):
                cell = sheet.cell(row, col)
                value = cell.value
                if row in forward_headers:
                    if value not in (None, ""):
                        inherited[row] = value
                    value = inherited.get(row) if value in (None, "") else value
                labels.append((value, cell.number_format))
            if not any(value not in (None, "") for value, _ in labels):
                continue
            if time_row and labels[time_row - 1][0] not in (None, ""):
                value = labels[time_row - 1][0]
                if isinstance(value, (int, float)) and 20000 < value < 100000:
                    from openpyxl.utils.datetime import from_excel
                    value = from_excel(value)
                key = (_source_key(value),)
            else:
                key = tuple(_source_key(value) for value, _ in labels)
            col_keys[col] = key
            columns.setdefault(key, labels)
        for row, row_key in row_keys.items():
            for col, col_key in col_keys.items():
                cell = sheet.cell(row, col)
                if cell.value in (None, ""):
                    continue
                key = row_key, col_key
                previous = values.get(key)
                if previous is None:
                    values[key] = (cell.value, cell.number_format, cell.data_type, item.path.name)
                elif previous[0] != cell.value:
                    conflict_count += 1
                    if len(conflicts) < 3:
                        conflicts.append(f"{'/'.join(row_key)} @ {'/'.join(col_key)}: {previous[3]}={previous[0]} / {item.path.name}={cell.value}")
    book = Workbook()
    target = book.active
    target.title = matches[0][1].title
    first = matches[0][1]
    for row in range(1, header_rows + 1):
        for col in range(1, key_columns + 1):
            cell = first.cell(row, col)
            target.cell(row, col, cell.value).number_format = cell.number_format
    if column_group_row is None:
        column_order = sorted(columns, key=lambda key: _period_sort("|".join(key), keep_aggregate_order=True))
    else:
        # Keep a month's bands and subtotal together, in source band order.
        column_order = sorted(columns, key=lambda key: (
            _period_sort(key[column_group_row - 1]),
            any(word in "|".join(key) for word in ("总计", "汇总", "合计")),
        ))
    previous_headers = {}
    for index, key in enumerate(column_order, key_columns + 1):
        for row, (value, fmt) in enumerate(columns[key], 1):
            repeated = row in sparse_headers and previous_headers.get(row) == value
            target.cell(row, index, None if repeated else value).number_format = fmt
            previous_headers[row] = value
    for row, (row_key, labels) in enumerate(rows.items(), header_rows + 1):
        for col, value in enumerate(labels, 1):
            target.cell(row, col, value)
        for col, col_key in enumerate(column_order, key_columns + 1):
            value = values.get((row_key, col_key))
            if value is not None:
                cell = target.cell(row, col, value[0])
                cell.number_format, cell.data_type = value[1:3]
    target._source_matches = matches
    if conflict_count:
        LOGGER.warning("[多文件数值冲突] Sheet=%s | 共%d项 | 保留原文件顺序中首个非空值，其他文件仍保留为来源 | 示例=%s",
                       target.title, conflict_count, "；".join(conflicts))
    return target


def _merge_records(matches, key_aliases, *, header_depth=1):
    """Union records by named date/hour/period fields, not physical positions."""
    aliases = [{compact_identifier(value) for value in group} for group in key_aliases]
    columns, records = {}, {}
    first_header = None
    conflict_count, conflicts = 0, []
    for item, sheet in matches:
        header_row = None
        for row in range(1, min(sheet.max_row, 15) + 1):
            found = {}
            for col in range(1, sheet.max_column + 1):
                label = compact_identifier(sheet.cell(row, col).value)
                for field, names in enumerate(aliases):
                    if label in names:
                        found.setdefault(field, col)
            if len(found) == len(aliases):
                header_row, key_cols = row, [found[i] for i in range(len(aliases))]
                break
        if header_row is None:
            raise ValueError(f"多文件合并无法识别业务键表头: {item.path.name} / {sheet.title}")
        if first_header is None:
            first_header = header_row
        by_column, inherited = {}, ""
        for col in range(1, sheet.max_column + 1):
            labels = []
            for offset in range(header_depth):
                cell = sheet.cell(header_row + offset, col)
                value = cell.value
                if header_depth > 1 and offset == 0:
                    if value not in (None, ""):
                        inherited = value
                    value = inherited
                labels.append((value, cell.number_format))
            if not any(value not in (None, "") for value, _ in labels):
                continue
            # Key aliases describe the same field even if the header wording differs.
            key = (f"__key_{key_cols.index(col)}",) if col in key_cols else tuple(
                compact_identifier(value) for value, _ in labels)
            by_column[col] = key
            columns.setdefault(key, labels)
        for row in range(header_row + header_depth, sheet.max_row + 1):
            keys = [sheet.cell(row, col).value for col in key_cols]
            if any(value in (None, "") for value in keys):
                continue
            row_key = tuple(_source_key(value) for value in keys)
            record = records.setdefault(row_key, {})
            for col, field in by_column.items():
                cell = sheet.cell(row, col)
                if cell.value in (None, ""):
                    continue
                previous = record.get(field)
                if previous is None:
                    record[field] = (cell.value, cell.number_format, cell.data_type, item.path.name)
                elif previous[0] != cell.value:
                    conflict_count += 1
                    if len(conflicts) < 3:
                        conflicts.append(f"{'/'.join(row_key)} @ {'/'.join(field)}: "
                                         f"{previous[3]}={previous[0]} / {item.path.name}={cell.value}")
    target = Workbook().active
    target.title = matches[0][1].title
    if first_header > 1:
        _copy_cells(matches[0][1], target, row_end=first_header - 1)
    for col, labels in enumerate(columns.values(), 1):
        for offset, (value, fmt) in enumerate(labels):
            target.cell(first_header + offset, col, value).number_format = fmt
    for row, key in enumerate(sorted(records, key=lambda key: _period_sort("|".join(key))),
                              first_header + header_depth):
        for col, field in enumerate(columns, 1):
            value = records[key].get(field)
            if value is not None:
                cell = target.cell(row, col, value[0])
                cell.number_format, cell.data_type = value[1:3]
    target._source_matches = matches
    if conflict_count:
        LOGGER.warning("[多文件数值冲突] Sheet=%s | 共%d项 | 保留首个非空值 | 示例=%s",
                       target.title, conflict_count, "；".join(conflicts))
    return target


def _merge_option_fee(matches):
    """Align dynamic period/band headers, retaining sparse grouped headers."""
    parts, first_period_row = [], None
    for item, sheet in matches:
        # An amount-band header is not a month such as 26-03.
        def is_band(value):
            text = compact_text(value).replace("元", "")
            if "免费" in text:
                return True
            numbers = re.findall(r"\d+", text)
            return bool(numbers and max(map(int, numbers)) >= 100
                        and re.search(r"[-—–~～至到]|以上|以下", text))
        scored = [(sum(is_band(sheet.cell(row, col).value)
                       for col in range(3, sheet.max_column + 1)), row)
                  for row in range(2, min(sheet.max_row, 15) + 1)]
        score, band_row = max(scored, default=(0, 0))
        if not score:
            raise ValueError(f"多文件合并无法识别金额段表头: {item.path.name} / {sheet.title}")
        period_row = band_row - 1
        if first_period_row is None:
            first_period_row = period_row
        part = Workbook().active
        part.title = sheet.title
        _copy_cells(sheet, part, row_start=period_row)
        parts.append((item, part))
    merged = _merge_matrix(parts, 2, 2, forward_rows=(1,), forward_headers=(1,),
                           column_group_row=1, sparse_headers=(1,))
    target = Workbook().active
    target.title = matches[0][1].title
    if first_period_row > 1:
        _copy_cells(matches[0][1], target, row_end=first_period_row - 1)
    _copy_cells(merged, target, row_offset=first_period_row - 1)
    target._source_matches = matches
    return target


def _copy_cells(source, target, row_start=1, row_end=None, col_start=1, col_end=None, row_offset=0, col_offset=0):
    for cells in source.iter_rows(min_row=row_start, max_row=row_end or source.max_row,
                                  min_col=col_start, max_col=col_end or source.max_column):
        for cell in cells:
            if cell.value is not None:
                dest = target.cell(cell.row - row_start + 1 + row_offset,
                                   cell.column - col_start + 1 + col_offset, cell.value)
                dest.number_format, dest.data_type = cell.number_format, cell.data_type


def _merge_partitioned(matches, *, horizontal):
    groups = {}
    for item, sheet in matches:
        if horizontal:
            starts = [col for col in range(1, sheet.max_column + 1) if sheet.cell(1, col).value not in (None, "")]
            limit = sheet.max_column + 1
        else:
            starts = [row for row in range(1, sheet.max_row + 1)
                      if clean_text(sheet.cell(row, 1).value).endswith("退订比例")
                      and all(sheet.cell(row, col).value in (None, "") for col in range(2, sheet.max_column + 1))]
            limit = sheet.max_row + 1
        for start, end in zip(starts, starts[1:] + [limit]):
            title = clean_text(sheet.cell(1, start).value if horizontal else sheet.cell(start, 1).value)
            book = Workbook()
            part = book.active
            part.title = sheet.title
            if horizontal:
                _copy_cells(sheet, part, col_start=start, col_end=end - 1)
            else:
                _copy_cells(sheet, part, row_start=start, row_end=end - 1)
            groups.setdefault(title, []).append((item, part))
    if not groups:
        return None
    book = Workbook()
    target = book.active
    target.title = matches[0][1].title
    offset = 0
    for title, parts in groups.items():
        if horizontal:
            first = parts[0][1]
            period_column = next((col for col in range(2, first.max_column + 1)
                                  if re.match(r"^(?:20)?\d{2}(?:WK|[-/.])", clean_text(first.cell(2, col).value), re.I)), None)
            if period_column is None:
                return None
            merged = _merge_matrix(parts, 2, period_column - 1)
            _copy_cells(merged, target, col_offset=offset)
            offset += merged.max_column + 1
        else:
            merged = _merge_matrix(parts, 2, 1)
            _copy_cells(merged, target, row_offset=offset)
            offset += merged.max_row + 1
    target._source_matches = matches
    return target

def merge_source_sheets(matches):
    """Known source layouts share one merged read view, not a copied source file."""
    if len(matches) == 1:
        return matches[0]
    first, sheet = matches[0]
    if "图表" not in sheet.title and all("选配比例" in item.path.name for item, _ in matches):
        return first, _merge_matrix(matches, 1, 3, forward_rows=(1, 2))
    if all("首销期订单节奏" in item.path.name for item, _ in matches):
        return first, _merge_matrix(matches, 2, 1, time_row=2)
    if "小订分时退订" in sheet.title:
        return first, _merge_records(matches, ({"小订日期", "下单日期", "日期"}, {"小订小时", "小时", "时段"}))
    if "日度退订" in sheet.title:
        return first, _merge_records(matches, ({"取消日期", "退订日期", "日期"},))
    if all("订单7级转化" in item.path.name for item, _ in matches):
        return first, _merge_records(matches, ({"周", "月", "周期", "统计周期"},), header_depth=2)
    if all("选配金统计" in item.path.name for item, _ in matches):
        if "排名" in sheet.title:
            return first, _merge_matrix(matches, 2, 2, forward_rows=(1,))
        return first, _merge_option_fee(matches)
    if "选配退订" in sheet.title:
        merged = _merge_partitioned(matches, horizontal=False)
        return (first, merged) if merged is not None else None
    if all("SKU" in item.path.name for item, _ in matches):
        merged = _merge_partitioned(matches, horizontal=True)
        return (first, merged) if merged is not None else None
    return None


class WorkbookView:
    def __init__(self, sheets):
        self.worksheets = list(sheets)
        self.sheetnames = [sheet.title for sheet in self.worksheets]

    def __getitem__(self, name):
        return next(sheet for sheet in self.worksheets if sheet.title == name)

    def close(self):
        pass  # The enclosing store/context owns the physical workbooks.


@dataclass
class WorkbookItem:
    path: Path
    workbook: Any


class WorkbookStore:
    def __init__(self, input_dir: Path):
        self.input_dir = input_dir
        self.items: list[WorkbookItem] = []
        self.chart_assets: dict[str, str] = {}
        self.load_errors: list[str] = []
        self._find_all_cache: dict[str, list[WorkbookItem]] = {}
        self._combined_cache: dict[str, WorkbookItem] = {}
        self._source_aliases: dict[tuple[str, str], list[tuple[WorkbookItem, Any]]] = {}
        self._subject_sheets_cache: dict[tuple[Any, ...], list[tuple[WorkbookItem, Any]]] = {}

    def load(self, exclude_names: Iterable[str] = (), *, data_only_view: bool = False) -> None:
        if not self.input_dir.exists():
            raise FileNotFoundError(f"输入目录不存在: {self.input_dir}")
        self._find_all_cache.clear()
        self._combined_cache.clear()
        self._source_aliases.clear()
        self._subject_sheets_cache.clear()
        excluded = set(exclude_names)
        paths = sorted(path for path in self.input_dir.glob("*.xlsx") if not path.name.startswith("~$") and path.name not in excluded)
        if not paths:
            LOGGER.warning("输入目录中没有找到 .xlsx 文件: %s", self.input_dir)
        for path in paths:
            try:
                workbook = (load_data_workbook(path, self.input_dir / ".cache" / "excel")
                            if data_only_view else load_workbook(path, data_only=True, read_only=False))
                self.items.append(WorkbookItem(path=path, workbook=workbook))
                LOGGER.debug("已读取 Excel: %s（%d 个 Sheet）", path.name, len(workbook.sheetnames))
            except Exception as exc:
                message = f"Excel读取失败: {path.name}: {exc}"
                self.load_errors.append(message)
                LOGGER.exception(message)

    def close(self) -> None:
        for item in self.items:
            item.workbook.close()

    def find(self, keyword: str) -> WorkbookItem | None:
        items = self.find_all(keyword)
        if not items:
            return None
        if len(items) == 1:
            return items[0]
        if keyword in self._combined_cache:
            return self._combined_cache[keyword]
        groups = {}
        for item in items:
            for sheet in item.workbook.worksheets:
                groups.setdefault(compact_text(sheet.title), []).append((item, sheet))
        sheets = []
        for matches in groups.values():
            title = matches[0][1].title
            combined = merge_source_sheets(matches)
            if combined is not None:
                sheets.append(combined[1])
            else:
                # Unknown layouts remain separate; consumers using all matching
                # sheets or the raw centre still see every original.
                sheets.extend(sheet for _, sheet in matches)
            self._source_aliases[(items[0].path.name, title)] = matches
            for source_item, source_sheet in matches:
                self._source_aliases[(source_item.path.name, source_sheet.title)] = matches
                self._source_aliases[(items[0].path.name, source_sheet.title)] = matches
        result = WorkbookItem(items[0].path, WorkbookView(sheets))
        self._combined_cache[keyword] = result
        LOGGER.info("多文件来源: %s | 共%d个文件全部纳入: %s", keyword, len(items), "、".join(item.path.name for item in items))
        return result

    def expand_sources(self, sources):
        from core.models import SourceRef
        result, seen = [], set()
        for source in sources:
            matches = self._source_aliases.get((source.file, source.sheet))
            candidates = [SourceRef(item.path.name, sheet.title, source.note) for item, sheet in matches] if matches else [source]
            for candidate in candidates:
                key = candidate.file, candidate.sheet, candidate.note
                if key not in seen:
                    seen.add(key)
                    result.append(candidate)
        return result

    def resolve_dashboard_sources(self, dashboard):
        """Keep physical source links valid for merged/secondary-file sheets."""
        def visit(value):
            if isinstance(value, dict):
                matches = self._source_aliases.get((value.get("file"), value.get("sheet")))
                if matches:
                    value["file"], value["sheet"] = matches[0][0].path.name, matches[0][1].title
                    if len(matches) > 1:
                        value["source_files"] = list(dict.fromkeys(item.path.name for item, _ in matches))
                matches = self._source_aliases.get((value.get("source_file"), value.get("source_sheet")))
                if matches:
                    value["source_file"] = matches[0][0].path.name
                    if len(matches) > 1:
                        value["source_files"] = list(dict.fromkeys(item.path.name for item, _ in matches))
                for child in value.values():
                    visit(child)
            elif isinstance(value, list):
                for child in value:
                    visit(child)
        dashboard.sources = self.expand_sources(dashboard.sources)
        visit(dashboard.views)
        return dashboard

    def find_all(self, keyword: str) -> list[WorkbookItem]:
        cached = self._find_all_cache.get(keyword)
        if cached is None:
            cached = [item for item in self.items if keyword in item.path.name]
            self._find_all_cache[keyword] = cached
        return cached

    def find_subject_sheet(
        self,
        keyword: str,
        subject: str,
        grain: str | None = None,
        suffix: str | tuple[str, ...] | list[str] | None = None,
        include_charts: bool = False,
    ) -> tuple[WorkbookItem, Any] | None:
        matches = self.find_subject_sheets(keyword, subject, grain, suffix, include_charts)
        return matches[0] if matches else None

    def find_subject_sheets(
        self,
        keyword: str,
        subject: str,
        grain: str | None = None,
        suffix: str | tuple[str, ...] | list[str] | None = None,
        include_charts: bool = False,
    ) -> list[tuple[WorkbookItem, Any]]:
        """Return every matching sheet; callers can retain multiple launch phases."""
        suffixes = () if suffix is None else ((suffix,) if isinstance(suffix, str) else tuple(suffix))
        cache_key = (keyword, compact_text(subject), grain, suffixes, include_charts)
        cached = self._subject_sheets_cache.get(cache_key)
        if cached is not None:
            return cached
        combined = self.find(keyword)
        items = [combined] if combined else []
        if not items:
            return []
        result: list[tuple[WorkbookItem, Any]] = []
        for item in items:
            for sheet in item.workbook.worksheets:
                if not include_charts and "图表" in sheet.title:
                    continue
                if compact_text(sheet_subject(sheet.title)) != compact_text(subject):
                    continue
                if grain and grain_from_sheet(sheet.title) != grain:
                    continue
                if suffixes:
                    normalized_title = compact_identifier(sheet.title)
                    if not any(compact_identifier(candidate) in normalized_title for candidate in suffixes):
                        continue
                origins = self._source_aliases.get((item.path.name, sheet.title))
                owner = WorkbookItem(origins[0][0].path, item.workbook) if origins else item
                result.append((owner, sheet))
                LOGGER.debug("Sheet映射: 关键词=%s, 主体=%s -> %s / %s", keyword, subject, item.path.name, sheet.title)
        self._subject_sheets_cache[cache_key] = result
        return result

    def matching_sheet_names(self, keyword: str, subject: str) -> list[str]:
        """Return candidate Sheet names for diagnostics when a module mapping fails."""
        items = self.find_all(keyword)
        target = compact_text(subject)
        return [
            f"{item.path.name} / {sheet.title}"
            for item in items
            for sheet in item.workbook.worksheets
            if compact_text(sheet_subject(sheet.title)) == target
        ]

    def all_sheet_names(self) -> Iterable[tuple[str, str]]:
        for item in self.items:
            for name in item.workbook.sheetnames:
                yield item.path.name, name

    def find_chart_data(
        self, keyword: str, subject: str, grain: str
    ) -> tuple[WorkbookItem, Any, dict[str, Any]] | None:
        """Read the preferred subject block from EVERY matching workbook."""
        self.find(keyword)  # register source aliases and reuse merged read views
        target = compact_text(subject)
        kind = subject_type(subject)
        owner = subject_parent(subject, kind) if kind == "generation" else subject
        found = []
        for item in self.find_all(keyword):
            sheets = [sheet for sheet in item.workbook.worksheets
                      if "图表" in sheet.title and grain_from_sheet(sheet.title) == grain]
            sheets.sort(key=lambda sheet: compact_text(sheet_subject(sheet.title)) != compact_text(owner))
            match = None
            for sheet in sheets:
                for row in range(1, sheet.max_row + 1):
                    if compact_text(sheet.cell(row, 1).value) == target:
                        data = self._parse_chart_block(item, sheet, row, target)
                        if data:
                            match = (item, sheet, data)
                            break
                if match:
                    break
            if match:
                found.append(match)
        if not found:
            return None
        return found[0][0], found[0][1], self._merge_chart_data(found)

    def find_chart_blocks(
        self, keyword: str, subject: str, grain: str
    ) -> tuple[WorkbookItem, Any, list[dict[str, Any]]] | None:
        """Union owner/subject chart blocks and periods across source files."""
        self.find(keyword)
        kind = subject_type(subject)
        owner = subject_parent(subject, kind) if kind == "generation" else subject
        grouped = {}
        for item in self.find_all(keyword):
            for sheet in item.workbook.worksheets:
                if "图表" not in sheet.title or grain_from_sheet(sheet.title) != grain:
                    continue
                if compact_text(sheet_subject(sheet.title)) != compact_text(owner):
                    continue
                for row in range(1, sheet.max_row + 1):
                    name = clean_text(sheet.cell(row, 1).value)
                    if not name or (kind == "generation" and compact_text(name) != compact_text(subject)):
                        continue
                    data = self._parse_chart_block(item, sheet, row, compact_text(name))
                    if data:
                        grouped.setdefault(compact_text(name), []).append((item, sheet, data))
        if grouped:
            first = next(iter(grouped.values()))[0]
            return first[0], first[1], [self._merge_chart_data(matches) for matches in grouped.values()]
        exact = self.find_chart_data(keyword, subject, grain)
        return (exact[0], exact[1], [exact[2]]) if exact else None

    @staticmethod
    def _merge_chart_data(matches):
        if len(matches) == 1:
            data = matches[0][2]
            return {**data, "totals": [0 if value is None else value for value in data["totals"]],
                    "series": [{**entry, "values": [0 if value is None else value for value in entry["values"]]}
                               for entry in data["series"]]}
        periods, series, totals = set(), {}, {}
        conflicts = 0
        samples = []
        for item, sheet, data in matches:
            periods.update(data["periods"])
            for index, period in enumerate(data["periods"]):
                value = data["totals"][index]
                if value is not None and period not in totals:
                    totals[period] = value
                elif value is not None and totals[period] != value:
                    conflicts += 1
                    if len(samples) < 3:
                        samples.append(period)
                for entry in data["series"]:
                    values = series.setdefault(entry["name"], {})
                    value = entry["values"][index]
                    if value is not None and period not in values:
                        values[period] = value
                    elif value is not None and values[period] != value:
                        conflicts += 1
        order = sorted(periods, key=_period_sort)
        if conflicts:
            LOGGER.warning("[多文件图表冲突] 主体=%s | 共%d项 | 重叠周期保留首个值，不重复累加 | 周期示例=%s",
                           matches[0][2]["subject"], conflicts, "、".join(samples))
        return {**matches[0][2], "periods": order,
                "totals": [totals.get(period, 0) for period in order],
                "series": [{"name": name, "values": [values.get(period, 0) for period in order]}
                           for name, values in series.items()],
                "image_key": None}

    def _parse_chart_block(
        self, item: WorkbookItem, sheet: Any, row: int, target: str
    ) -> dict[str, Any] | None:
        header_row = row + 1
        columns: list[int] = []
        periods: list[str] = []
        for col in range(3, sheet.max_column + 1):
            cell = sheet.cell(header_row, col)
            value = display_period(cell.value, cell.number_format)
            if not value or value == "总计":
                continue
            columns.append(col)
            periods.append(value)
        if not columns:
            return None
        series = []
        # Missing cells must remain missing until cross-file fallback finishes.
        totals = [None] * len(columns)
        for data_row in range(header_row + 1, sheet.max_row + 1):
            label = clean_text(sheet.cell(data_row, 2).value)
            next_subject = clean_text(sheet.cell(data_row, 1).value)
            if next_subject:
                break
            if not label:
                if series:
                    break
                continue
            if label == "总计":
                totals = [cell_number(sheet.cell(data_row, col))
                          if sheet.cell(data_row, col).value not in (None, "") else None for col in columns]
                break
            values = [cell_number(sheet.cell(data_row, col), rate=True)
                      if sheet.cell(data_row, col).value not in (None, "") else None for col in columns]
            series.append({"name": label, "values": values})
        if not series:
            return None
        block_subject = clean_text(sheet.cell(row, 1).value)
        image_key = self._extract_chart_image(item, sheet, row, compact_text(target))
        return {
            "subject": block_subject,
            "headline": clean_text(sheet.cell(header_row, 2).value),
            "periods": periods,
            "series": series,
            "totals": totals,
            "image_key": image_key,
        }

    def _extract_chart_image(self, item: WorkbookItem, sheet: Any, subject_row: int, target: str) -> str | None:
        """Extract the image embedded in a subject block once and keep it as a shared HTML asset."""
        asset_key = f"{item.path.name}|{sheet.title}|{target}"
        if asset_key in self.chart_assets:
            return asset_key
        next_subject_row = None
        for row in range(subject_row + 1, sheet.max_row + 1):
            if clean_text(sheet.cell(row, 1).value):
                next_subject_row = row
                break

        if next_subject_row is None:
            # The final subject has no real lower boundary. Infer one from the
            # preceding block spacing instead of relying on sheet.max_row: Excel
            # images may be anchored one or more rows below the last data cell.
            subject_rows = [
                row for row in range(1, subject_row + 1)
                if clean_text(sheet.cell(row, 1).value)
            ]
            spacings = [
                later - earlier
                for earlier, later in zip(subject_rows, subject_rows[1:])
                if later > earlier
            ]
            if spacings:
                block_span = max(1, int(round(median(spacings))))
            else:
                # A single-block sheet still gets a finite search window. Cover
                # its existing cells plus the immediately following anchor row.
                block_span = max(30, sheet.max_row - subject_row + 2)
            next_subject_row = subject_row + block_span

        images = sorted(
            getattr(sheet, "_images", []),
            key=lambda image: getattr(getattr(getattr(image, "anchor", None), "_from", None), "row", -1),
        )
        for image in images:
            anchor = getattr(image, "anchor", None)
            marker = getattr(anchor, "_from", None)
            image_row = getattr(marker, "row", -1) + 1
            if not subject_row <= image_row < next_subject_row:
                continue
            try:
                image_format = clean_text(getattr(image, "format", "png")).lower() or "png"
                mime = "jpeg" if image_format in {"jpg", "jpeg"} else image_format
                encoded = base64.b64encode(image._data()).decode("ascii")
                self.chart_assets[asset_key] = f"data:image/{mime};base64,{encoded}"
                return asset_key
            except (OSError, ValueError, AttributeError):
                return None
        return None


_METRIC_SHEET_CACHE: "weakref.WeakKeyDictionary[Any, dict[str, dict[str, Any]]]" = weakref.WeakKeyDictionary()


def parse_metric_sheet(sheet: Any) -> dict[str, dict[str, Any]]:
    """Parse one metric sheet once per build.

    Dashboard modules only read the returned structure.  Caching by Worksheet
    identity avoids reparsing the same order/lock sheet for every subject and
    every period while allowing closed workbooks to be garbage-collected.
    """
    cached = _METRIC_SHEET_CACHE.get(sheet)
    if cached is not None:
        return cached
    headers = [display_period(sheet.cell(1, col).value, sheet.cell(1, col).number_format) for col in range(4, sheet.max_column + 1)]
    # 原表周期列完整保留，包括“总计/汇总”；KPI Card 直接使用原表对应列的值。
    periods = [header for header in headers if header]
    result = {period: {"metrics": {}, "structures": {}} for period in periods}
    current_metric = ""
    current_stat = ""
    for row in range(2, sheet.max_row + 1):
        metric = clean_text(sheet.cell(row, 1).value)
        stat = clean_text(sheet.cell(row, 2).value)
        category = clean_text(sheet.cell(row, 3).value)
        if metric:
            current_metric = metric
        if stat:
            current_stat = stat
        if not current_metric or not category:
            continue
        for offset, period in enumerate(headers, start=4):
            if period not in result:
                continue
            cell = sheet.cell(row, offset)
            if current_stat == "数量" and category == "数量":
                result[period]["metrics"][current_metric] = cell_number(cell)
            elif current_stat == "占比":
                result[period]["structures"].setdefault(current_metric, []).append(
                    {"label": category, "share": cell_number(cell, rate=True), "count": None}
                )
            elif current_stat == "数量":
                rows = result[period]["structures"].setdefault(current_metric, [])
                target = next((entry for entry in rows if entry["label"] == category), None)
                if target:
                    target["count"] = cell_number(cell)
    _METRIC_SHEET_CACHE[sheet] = result
    return result
