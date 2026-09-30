from __future__ import annotations

import re
from functools import lru_cache
from pathlib import Path
from typing import Any


_REMOVABLE_WORDS = re.compile(r"鸿蒙智行|车型|款|总计|汇总|合计", re.I)
_SEPARATORS = re.compile(r"[\s()（）/\\_\-&]+")


def model_key(value: Any) -> str:
    """Return a brand-preserving identity key for propagation and generation names.

    Brand names are deliberately retained.  Removing them makes similarly named
    vehicles from different brands indistinguishable and can bind the wrong data.
    """

    text = str(value or "").strip().lower().replace("（", "(").replace("）", ")")
    text = _REMOVABLE_WORDS.sub("", text)
    return _SEPARATORS.sub("", text)


def usable_attribute(value: Any, default: str = "未维护") -> str:
    """Normalize blank and placeholder master-data attributes for the UI."""

    text = str(value or "").strip()
    return default if text in {"", "待维护", "未维护"} else text


def stage_name(record: dict) -> str:
    """Identity of a reservation/launch campaign; the parent is never its alias."""
    return str(record.get("二级代际名") or record.get("代际名") or record.get("车型") or record.get("历史传播名") or "").strip()


def has_reservation(record: dict, *, modern=True) -> bool:
    """The integrated attributes sheet also contains direct-launch vehicles."""
    marker = record.get("有小订")
    if marker not in (None, ""):
        return marker not in (False, 0, "否", "无", "false", "False")
    if not modern:
        return True  # Legacy small-order-only roster.
    return any(record.get(k) not in (None, "") for k in
               ("小订开始日期", "小订结束日期", "小订开始", "小订结束")) or any(
        isinstance(record.get(k), (int, float)) and record[k] > 0
        for k in ("总小订", "小订天数"))


def stage_records(path) -> list[dict]:
    source = Path(path)
    if not source.exists():
        return []
    stat = source.stat()
    return _stage_records(str(source.resolve()), stat.st_mtime_ns, stat.st_size)


@lru_cache(maxsize=8)
def _stage_records(path, stamp, size):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(v or "").strip() for v in next(rows, ())]
        if "代际名" not in headers:
            return []  # Legacy mappings have a separate compatibility reader.
        records = []
        for values in rows:
            record = dict(zip(headers, values))
            primary = str(record.get("代际名") or "").strip()
            if not primary:
                continue
            secondary = str(record.get("二级代际名") or "").strip()
            record.update({"代际名": primary, "二级代际名": secondary,
                           "历史传播名": secondary or primary,
                           "订单分析代际名": secondary or primary,
                           "primary_generation": primary, "secondary_generation": secondary})
            record.setdefault("品牌", primary.split(" ", 1)[0])
            records.append(record)
        return records
    finally:
        book.close()


def parent_generation(name, records) -> str:
    key = model_key(name)
    found = next((r for r in records if r.get("二级代际名") and model_key(r["二级代际名"]) == key), None)
    return str(found["代际名"]) if found else str(name or "")


def stage_label(name, parent) -> str:
    text = str(name or "").strip()
    if text.startswith(str(parent)):
        return text[len(str(parent)):].strip() or ""
    return text


def generation_records() -> list[dict]:
    code_root = Path(__file__).resolve().parents[2]
    project = code_root.parent if code_root.name.lower() == "scripts" else code_root
    return stage_records(project / "input_file" / "销量预测输入文件" / "小订及首销数据整理.xlsx")


def resolve_stage_identity(name, records, launch=None):
    """Compatibility migration: exact identity or one confirmed family/window."""
    key = model_key(name)
    exact = [r for r in records if model_key(stage_name(r)) == key]
    if len(exact) == 1:
        return stage_name(exact[0])
    def family(value):
        return re.findall(r"[a-z]+\d+", model_key(re.sub(r"(?:20)?\d{2}款", "", str(value), flags=re.I)))
    families = family(name)
    if not families:
        return None
    year = re.search(r"(?:20)?(\d{2})款", str(name))
    launch_day = str(launch or "").split(" ")[0]
    brand = re.match(r"^([\u4e00-\u9fff]{1,3}界)", str(name).strip())
    candidates = []
    for r in records:
        if brand and not str(r.get("代际名") or stage_name(r)).startswith(brand[1]):
            continue
        candidate_key = model_key(stage_name(r))
        if family(stage_name(r)) != families:
            continue
        if year and ("20" + year[1]) not in candidate_key:
            continue
        if ("ultimate" in key) != ("ultimate" in candidate_key):
            continue
        versions = ("ultra", "pro", "纯电", "增程", "典藏大观")
        candidate_attributes = candidate_key + model_key(r.get("能源类型"))
        if any(v not in candidate_attributes for v in versions if v in key):
            continue
        if launch_day and str(r.get("开始大定日期") or r.get("首销开始") or "").split(" ")[0] != launch_day:
            continue
        candidates.append(r)
    return stage_name(candidates[0]) if len(candidates) == 1 else None
