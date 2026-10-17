"""Read forecast inputs from one workbook while retaining original source identities."""
from __future__ import annotations

from contextlib import contextmanager
from contextvars import ContextVar
from datetime import date, datetime, timedelta
from pathlib import Path
import math
import logging
import json

from openpyxl import load_workbook

from core.excel import WorkbookItem, WorkbookStore

SUMMARY_NAME = "鸿蒙智行销量数据汇总.xlsx"
GUIDE_SHEET = "说明与来源"
MASTER_SHEET = "车型基本信息"
DAILY_SHEET = "小订首销平销by天"
DAILY_FIELD_LABELS = {"小订数量": "小订", "退订数量": "小订退订"}
GROSS_DAILY_SHEET = "大定当日数量"
GROSS_PROGRESS_SHEET = "大定累计完成度"
REFERENCE_SHEET_ALIASES = {"总大定当日数量": GROSS_DAILY_SHEET, "累计大定完成度": GROSS_PROGRESS_SHEET}
SMALL_HOURLY_SHEET = "小订by时"
LAUNCH_HOURLY_SHEET = "首销by时"
WEEKLY_SHEET = "首销平销by周"
INDEX_SHEET = "数据来源目录"
SNAPSHOT_SHEET = "_预测缓存"
ACTIVE_SUMMARY = ContextVar("forecast_summary", default=None)


G9_TRACE_MODEL = "享界 G9 2027款"


def forecast_trace_enabled(model):
    """Temporary, read-only diagnostics for the one intranet investigation."""
    from core.model_identity import model_key
    return model_key(model) == model_key(G9_TRACE_MODEL)


def forecast_trace(model, step, **details):
    if forecast_trace_enabled(model):
        logging.getLogger('sales_forecast_refresh').info(
            '[G9取数排查] 代际=%s | 环节=%s | %s', model, step,
            json.dumps(details, ensure_ascii=False, default=str, allow_nan=True))


def forecast_trace_series(values, labels=None):
    """Describe values without coercing, filling or changing any business cell."""
    from core.excel import compact_diagnostic_ranges
    labels = list(labels) if labels is not None else [f'D{i+1}' for i in range(len(values))]
    known = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) and v >= 0
    probes = sorted(set(range(min(2, len(values)))) | set(range(max(0, len(values)-3), len(values))))
    return {'数量': len(values), '有效数': sum(known(v) for v in values),
            '零值数': sum(known(v) and v == 0 for v in values),
            '空白或无效范围': compact_diagnostic_ranges(labels[i] for i, v in enumerate(values) if not known(v)),
            '首尾样例': {str(labels[i]): values[i] for i in probes}}


class SheetView:
    def __init__(self, sheet, title):
        self._sheet, self.title = sheet, title

    def __getattr__(self, name):
        return getattr(self._sheet, name)

    def __getitem__(self, key):
        return self._sheet[key]


class WorkbookView:
    def __init__(self, sheets):
        self.worksheets = list(sheets)
        self.sheetnames = [sheet.title for sheet in self.worksheets]

    def __getitem__(self, name):
        return next(sheet for sheet in self.worksheets if sheet.title == name)

    def close(self):
        pass  # The enclosing summary_scope owns the workbook.


def input_path(default):
    active = ACTIVE_SUMMARY.get()
    return active["path"] if active else default


def open_input(path, **kwargs):
    active = ACTIVE_SUMMARY.get()
    if active and Path(path) == active["path"]:
        return active["inputs"]
    return load_workbook(path, **kwargs)


def table_records(workbook, name):
    """Read typed public cells, never a serialized dashboard or hidden payload."""
    canonical = REFERENCE_SHEET_ALIASES.get(name, name)
    if canonical in workbook.sheetnames:
        name = canonical
    if name not in workbook.sheetnames:
        return []
    rows = workbook[name].iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows, ())]
    result = [dict(zip(headers, values)) for values in rows if any(value is not None for value in values)]
    for row in result:
        if name == DAILY_SHEET:
            # A single read adapter keeps existing consumers and old workbooks
            # compatible; these aliases are never extra exported columns.
            row.update({old: row[new] for old, new in DAILY_FIELD_LABELS.items() if new in row})
        if "二级代际名" in headers and row.get("代际名"):
            event = str(row.get("二级代际名") or row["代际名"])
            row.update({"primary_generation": row["代际名"], "secondary_generation": row.get("二级代际名") or "",
                        "订单分析代际名": event, "历史传播名": event, "传播名": event})
    return result


def visible_target_names(workbook):
    return list(dict.fromkeys(
        str(row.get("primary_generation") or row["订单分析代际名"]) for row in table_records(workbook, MASTER_SHEET)
        if row.get("订单来源文件") and row.get("订单分析代际名")
    ))


ORDER_FIELDS = {"gross": "大定", "net": "留存大定", "small_to_big": "小转大", "direct": "直接大定", "lock": "交车锁单"}


def iso_day(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    return value.isoformat() if isinstance(value, date) else str(value or "")


def calendar_day(value):
    """Only valid calendar dates may anchor relative forecast curves."""
    try:
        return date.fromisoformat(iso_day(value)).isoformat()
    except (TypeError, ValueError):
        return ""


def date_cell(value):
    try:
        return datetime.fromisoformat(value) if value else None
    except (TypeError, ValueError):
        return value


def source_text(value):
    if isinstance(value, dict):
        return "｜".join(str(value.get(key) or "") for key in ("file", "sheet"))
    return str(value or "")


SMALL_CURVE_SHEET = "小订累计完成度"
SMALL_DAILY_SHEET = "小订当日数量"
D12_SHEET = "首销D1D2预测指标"
LEGACY_D12_SHEET = "D1_D2预测指标"
LEGACY_SMALL_CURVE_SHEET = "小订参考曲线"
SMALL_TOTAL_SHEET = "小订来源总量"
SMALL_REFERENCE_FIELDS = ("小订总量来源", "小订参考来源", "小订参考总量有效")
D12_HEADERS = [
    "传播名", "代际名", "映射状态", "产品档位", "首销天数", "总小转大", "总直接大定", "总大定", "总小订",
    "D1小转大", "D1小转大/总小转大", "D2小转大", "D2小转大/总小转大",
    "D1+D2小转大", "D1+D2小转大/总小转大", "D2小转大/D1小转大", "D1小转大/D1+D2小转大",
    "D1直接大", "D1直接大/总直接大", "D2直接大", "D2直接大/总直接大",
    "D1+D2直接大", "D1+D2直接大/总直接大", "D2直接大/D1直接大", "D1直接大/D1+D2直接大",
    "D1大定", "D1小转大/D1大定", "D1直接大/D1大定",
    "D2大定", "D2小转大/D2大定", "D2直接大/D2大定",
    "D1+D2大定", "D1+D2小转大/D1+D2大定", "D1+D2直接大/D1+D2大定",
    "D1退订", "D1退订率", "D2退订", "D2退订率", "D1+D2退订", "D1+D2退订率",
    "D2退订/D1退订", "D1退订/D1+D2退订",
    "D1口径状态", "D2口径状态", "D1+D2口径状态", "结构异常说明",
]


def _source_parts(value):
    if isinstance(value, dict):
        return value.get("file") or "", value.get("sheet") or ""
    parts = str(value or "").split("｜", 1)
    return parts[0], parts[1] if len(parts) > 1 else ""


def summary_quality(
    total_small: float,
    small_to_big: float,
    conversion: float,
    gross: float,
    direct: float,
    direct_share: float,
    cancel: float,
    cancel_rate: float,
    net: float,
    net_rate: float,
    lock: float,
    lock_rate: float,
    provided_fields: tuple[bool, ...] | None = None,
) -> tuple[float, float, list[str]]:
    fields = (total_small, small_to_big, conversion, gross, direct, direct_share, cancel, cancel_rate, net, net_rate, lock, lock_rate)
    presence = provided_fields or tuple(value not in (None, "") for value in fields)
    field_completeness = sum(bool(value) for value in presence) / len(fields)
    checks: list[tuple[bool, str]] = []
    if gross > 0 and (small_to_big > 0 or direct > 0):
        checks.append((abs(small_to_big + direct - gross) <= max(5.0, max(abs(small_to_big + direct), abs(gross)) * .002), "总小转大+总直接大定与总大定不一致"))
    for label, value in (("小订转化率", conversion), ("直接大定占比", direct_share), ("退订率", cancel_rate), ("留存大定率", net_rate), ("大定到锁单率", lock_rate)):
        if value not in (None, 0):
            checks.append((0 <= value <= 1, f"{label}超出0%～100%"))
    if gross > 0 and direct > 0 and direct_share > 0:
        checks.append((abs(direct_share - direct / gross) <= 0.02, "直接大定占比与数量不一致"))
    if gross > 0 and net > 0:
        checks.append((net <= gross * 1.02, "首销期留存大定高于总大定"))
    if gross > 0 and lock > 0:
        checks.append((lock <= gross * 1.02, "首销期锁单高于总大定"))
    issues = [message for passed, message in checks if not passed]
    consistency = sum(passed for passed, _ in checks) / len(checks) if checks else 0.0
    return field_completeness, consistency, issues


def public_forecast_tables(book, data, emit, weekly_rows=(), as_of_date=None, raw_profiles=None, weekly_builder=None):
    """Resolve each business cell once; every downstream table uses these cells."""
    references = data.setdefault("reference_daily", {})
    if GROSS_DAILY_SHEET not in references and "总大定当日数量" in references:
        references[GROSS_DAILY_SHEET] = references["总大定当日数量"]
    master = table_records(book, MASTER_SHEET)
    baseline = table_records(book, "预测基准总表")
    aliases = {"传播名": "历史传播名", "代际名": "订单分析代际名", "发布日": "首销开始", "首销截止": "首销结束"}
    for record in baseline:
        row = next((r for r in master if r.get("历史传播名") == record.get("传播名")), None)
        if row is None:
            row = {"历史传播名": record.get("传播名")}
            master.append(row)
        row.update({aliases.get(k, k): v for k, v in record.items() if v is not None})
    sources = table_records(book, INDEX_SHEET)
    history_file = next((r.get("原始文件") for r in sources if "历史" in str(r.get("来源类型"))), None)
    history_file = history_file or "小订及首销数据整理.xlsx"
    raw_profiles = list(raw_profiles if raw_profiles is not None else data.get("actuals", []))
    targets = {r["name"]: r for r in data.get("targets", [])}
    mix_sources = {}
    for table in (data.get("mix_daily", {}), data.get("mix_weekly", {})):
        for (model, _, _), (_, source) in table.items():
            filename = _source_parts(source)[0]
            if filename:
                mix_sources.setdefault(model, {})[filename] = None
    for target in targets.values():
        row = next((r for r in master if r.get("历史传播名") == target.get("history_model")
                    and r.get("订单分析代际名") == target["name"]), None)
        row = row if row is not None else next((r for r in master if r.get("订单分析代际名") == target["name"]), None)
        dates = {"首销开始": target.get("launch_date"), "首销结束": target.get("end_date"),
                 "小订开始": target.get("small_start_date"), "小订结束": target.get("small_end_date")}
        if row is None:
            row = {"历史传播名": target.get("history_model"), "订单分析代际名": target["name"]}
            master.append(row)
        elif any(row.get(k) is not None and iso_day(row[k]) != iso_day(v) for k, v in dates.items()):
            row = {k: row.get(k) for k in ("订单分析代际名", "品牌", "产品档位", "能源类型", "发布类型", "小订发布时段", "首销发布时段")}
            master.append(row)
        row.update({k: date_cell(v) for k, v in dates.items()})
        if not row.get("历史传播名"):
            row["历史传播名"] = target.get("history_model") or target["name"]
        row["首销天数"] = target.get("days")
        row["有小订"] = target.get("has_small", True)
        row["小订发布时段"] = target.get("small_period") or row.get("小订发布时段") or "未维护"
        row["首销发布时段"] = target.get("launch_period") or row.get("首销发布时段") or "未维护"
        raw = next((r for r in raw_profiles if r["model"] == target["name"]), {})
        files = [_source_parts(raw.get(k))[0] for k in ("day_source", "hour_source", "small_hour_source", "cancel_source")]
        files += [_source_parts(v)[0] for v in raw.get("small_daily_sources", {}).values()]
        files += [r.get("source_file") for r in data.get("steady_history", []) if r["model"] == target["name"]]
        files += list(mix_sources.get(target["name"], {}))
        files = [history_file if filename and filename.endswith(".refreshing.xlsx") else filename for filename in files]
        row["订单来源文件"] = "、".join(dict.fromkeys(f for f in files if f)) or history_file
    master_headers = ["历史传播名", "订单分析代际名", "原始表简称/别名", "品牌", "产品档位", "能源类型",
        "发布类型", "小订发布时段", "首销发布时段", "有小订", "小订开始", "小订结束", "首销开始", "首销结束", "小订天数", "首销天数",
        "总小订", "总小转大", "小订转化率", "总大定", "总直接大定", "直接大定占比",
        "总退订", "退订率", "首销期留存大定", "留存大定率", "首销期锁单", "留存大定到锁单率", "大定到锁单率", "字段完整度", "订单来源文件"]
    emit(MASTER_SHEET, master_headers, [[r.get(k) for k in master_headers] for r in master],
         {k for k in master_headers if k.endswith("率") or "占比" in k or k == "字段完整度"})

    rows = []
    def add(model, historical, day, stage, kind, source, values):
        filename, sheet = _source_parts(source)
        rows.append({"订单分析代际名": model, "历史传播名": historical, "日期": date_cell(day.get("date")),
                     "订单阶段": stage, "生命周期": day.get("day"), **values,
                     "来源类型": kind, "来源文件": filename, "来源Sheet": sheet})
    for raw in raw_profiles:
        model, target = raw["model"], targets.get(raw["model"], {})
        for day in raw.get("days", []):
            stage = "平销" if day.get("date") and target.get("end_date") and day["date"] > target["end_date"] else "首销"
            add(model, target.get("history_model"), day, stage, "首销订单", raw.get("day_source"),
                {c: day.get(k) for k, c in ORDER_FIELDS.items()})
            if day.get("_hourly_backfill"):
                rows[-1]["_origins"] = {ORDER_FIELDS[field]: source_text(raw.get("hour_source")) + (
                                           "（完整24小时汇总，分时补缺）" if field in day.get("_complete_hour_fields", [])
                                           else "（日/时大定一致，分时补缺）")
                                       for field in day["_hourly_backfill"]}
        for day in raw.get("small_daily_days", []):
            add(model, target.get("history_model"), day, "小订", "逐日小订", raw.get("small_daily_sources", {}).get(day.get("date")),
                {"小订数量": day.get("orders")})
    for item in data.get("history", []):
        if item["model"] in references.get(GROSS_DAILY_SHEET, {}):
            continue  # The explicit quantity table below owns this history field.
        for index, value in enumerate(item.get("daily_orders", [])):
            start = item.get("launch_date")
            day = {"day": f"D{index + 1}", "date": (date.fromisoformat(start) + timedelta(days=index)).isoformat() if start else None}
            add(item.get("generation") or item["model"], item["model"], day, "首销", "历史首销",
                {"file": history_file, "sheet": GROSS_DAILY_SHEET}, {"大定": value})
    small_records = []
    max_curve = max((len(r.get("standard_progress", [])) for r in data.get("small_order_history", [])), default=0)
    for item in data.get("small_order_history", []):
        if item.get("daily_actual"):
            target = targets.get(item.get("generation") or item["model"], {})
            for i, value in enumerate(item.get("daily_orders", [])):
                dates = item.get("dates", [])
                observed = dates[i] if i < len(dates) else None
                if (not target.get("small_start_date") or item.get("small_start_date") != target["small_start_date"]
                        or not observed or not target["small_start_date"] <= observed <= (target.get("small_end_date") or "")):
                    continue
                add(item.get("generation") or item["model"], item["model"],
                    {"day": f"D{i+1}", "date": observed}, "小订", "历史小订",
                    {"file": history_file, "sheet": item.get("source_sheet")}, {"小订数量": value})
        curve = item.get("standard_progress", [])
        small_records.append([item["model"], item.get("generation"), date_cell(item.get("small_start_date")),
            date_cell(item.get("small_end_date")), item.get("days"), item.get("total") if item.get("total_complete") or item.get("daily_actual") is False else None,
            item.get("total_source"), item.get("source_model") or item["model"],
            history_file, item.get("source_sheet"), date_cell(next(iter(item.get("dates") or []), None)),
            *curve, *([None] * (max_curve-len(curve)))])
    for item in data.get("steady_history", []):
        for day in item.get("daily", []):
            add(item["model"], None, day, "平销", "平销锁单",
                {"file": item.get("source_file"), "sheet": "、".join(item.get("daily_source_sheets", []))}, {"交车锁单": day.get("lock")})
    rows = resolved_daily_rows(rows, data, master, history_file, as_of_date or date.today(), raw_profiles)
    headers = ["订单分析代际名", "历史传播名", "日期", "订单阶段", "生命周期", "日期状态", "小订数量", "退订数量",
               "小转大", "直接大定", "大定", "留存大定", "交车锁单", "字段来源", "首销分时核对"]
    for sheet, key, field, label, source_key in (
        (SMALL_HOURLY_SHEET, "small_hourly_days", "orders", "小时小订", "small_hour_source"),
        (LAUNCH_HOURLY_SHEET, "hourly_days", "gross", "小时大定", "hour_source")):
        records = {}
        for raw in raw_profiles:
            for day in raw.get(key, []):
                hours = list(day.get("hours", []))
                for hour in hours:
                    record = [raw["model"], date_cell(day.get("date")), hour.get("hour"), hour.get(field),
                              _source_parts(raw.get(source_key))[0]]
                    if key == "hourly_days":
                        record.extend(hour.get(metric) for metric in ("small_to_big", "direct", "lock"))
                    identity = (raw["model"], iso_day(day.get("date")), hour.get("hour"))
                    if identity not in records:
                        records[identity] = record
                    else:
                        records[identity] = [a if a is not None else b for a, b in zip(records[identity], record)]
        component_headers = ["小时小转大", "小时直接大定", "小时交车锁单"] if key == "hourly_days" else []
        emit(sheet, ["订单分析代际名", "日期", "小时", label, "来源文件", *component_headers], [records[k] for k in sorted(records)])
        if key == "hourly_days":
            reconcile_intraday_fields(rows, data.get("actuals", []), records)
    emit(DAILY_SHEET, [DAILY_FIELD_LABELS.get(k, k) for k in headers], [[r.get(k) for k in headers] for r in rows])
    if weekly_builder:
        weekly_rows = weekly_builder(rows)
    weekly = [list(r[:11]) for r in weekly_rows]
    for item in data.get("steady_history", []):
        for week in item.get("weeks", []):
            match = next((r for r in weekly if r[0] == item["model"] and r[1] == week["period"] and r[2] == "平销"), None)
            if match is None:
                weekly.append([item["model"], week["period"], "平销", date_cell(week["start_date"]), date_cell(week["end_date"]),
                    None, None, week.get("lock"), None, None, source_text({"file": item.get("source_file"), "sheet": item.get("source_sheet")})])
            else:
                match[7] = week.get("lock")
                match[10] = source_text({"file": item.get("source_file"), "sheet": item.get("source_sheet")})
    emit(WEEKLY_SHEET, ["订单分析代际名", "周期", "订单阶段", "统计开始", "统计结束", "大定", "留存大定", "交车锁单",
         "大定来源", "留存大定来源", "锁单来源", *[stage + field for stage in ("首销", "平销") for field in ("大定", "留存大定", "交车锁单")]],
         resolved_weekly_rows(weekly, data.get("mix_weekly")))
    emit(SMALL_CURVE_SHEET, ["历史传播名", "订单分析代际名", "小订开始", "小订结束", "小订天数", "总小订", "总量来源",
         "原始名称", "来源文件", "来源Sheet", "首条数据日期", *[f"D{i+1}" for i in range(max_curve)]],
         small_records, {f"D{i+1}" for i in range(max_curve)})
    guide = [
        ["口径", "数量与空值", "单位：单；0是已知无销量，空白是未知。未来日期不视为真实完成日。", None, None],
        ["口径", "订单by天", "同代际同日期一行，逐字段按优先级选值；选配比例表覆盖范围内省略日期按0。未更新日期及显式空白不补0。by时保存小时数量及已知小时分项，不存累计列；by天的首销分时核对列记录可用于完整日核对的分项。", None, None],
        ["口径", "参考曲线", "小订累计完成度、小订当日数量按小订窗口排列D1、D2；已有日量和可计算比例均保留，分母未知才留空。进行中、当日或未来数据不作为完整终局参考；参考来源与终值资格见车型基本信息，D1/D2另标口径状态。", None, None],
        ["口径", "退订", "退订当日数量的D1为截至首销D1的累计小订退订（含首销前），D2起为当日新增；by天和D1/D2指标仍保存真实当日量。日量按整理表→小订退订分析的当日小订退补缺，真实0有效；累计只用兼容的真实累计，不混入大定退订。累计退订率仅截首销日期，分母为总小订。", None, None],
        ["口径", "小订及首销前", "整理表真实小订by天 → 小订选配比例by天 → 小订退订分析分时按日汇总；整理表须同一小订开始日期及维护窗口。D1分时只取小订分时退订。", None, None],
        ["口径", "首销窗口", "整理表 → 首销期订单节奏 → 对应大定/锁单选配比例表补缺；退订独立取数。进入平销不改变历史首销日的优先级。", None, None],
        ["口径", "首销结束后的日期", "大定、留存大定取大定选配比例，交车锁单取锁单选配比例；整理表不延伸到首销后，小订/小转大/直接大定不延伸各自维护窗口。", None, None],
        ["口径", "阶段未知", "保留实际数量：整理表 → 小订退订分析/首销期订单节奏 → 对应选配比例；阶段日期不明不参与正常预测。周总量已知但阶段拆分不足时，保留周总量、拆分留空。", None, None],
        ["口径", "平销", "预测交车锁单。完整周按周起止和阶段日期判断；不足完整周不当作完整周参考。", None, None],
        ["口径", "预测指标", "D1/D2和后续曲线使用同一份多来源取值结果；保存源表明确填写的值并标记日期/阶段/口径状态。网页读取时排除未结束日和不合格终值，不因不能参与预测而清空Excel。不用预测补真实值，不补未知分母。", None, None],
        ["口径", "字段补齐", "车型基本信息优先保留原维护时段，开始小时优先按上午9/下午14/晚上19；留白才用同代际（含二级事件）本阶段D1首个正数订单小时，并回填时段（0–11时上午、12–18时下午、19–23时晚上），另列来源并保留精确小时。日期可确定的天数、星期和平销开始在汇总生成时补齐。无依据留空，不补预测终值。", None, None],
        ["口径", "字段完整度", "按汇总最终字段是否有值计算；口径一致性和质量状态由同一组数量计算。", None, None],
        ["生成", "数据判定日期", (as_of_date or date.today()).isoformat(), None, None],
    ]
    for record in sources:
        guide.append(["来源", record.get("原始文件"), record.get("包含内容"),
                      record.get("文件数据开始日期"), record.get("文件数据结束日期")])
    if "汇总说明" in book.sheetnames:
        info = dict(book["汇总说明"].iter_rows(min_row=2, values_only=True))
        book.properties.identifier = info.get("刷新签名")
    emit(GUIDE_SHEET, ["类型", "项目", "内容", "数据开始", "数据结束"], guide)
    refresh_reference_tables(book, data, emit, rows, as_of_date or date.today())
    for name in ("汇总说明", "字段说明", INDEX_SHEET, "预测基准总表", "小订及退订逐日", "当前小订分时", "当前订单逐日", "当前首销分时", SNAPSHOT_SHEET, LEGACY_SMALL_CURVE_SHEET, SMALL_TOTAL_SHEET):
        if name in book.sheetnames:
            del book[name]
    if LEGACY_D12_SHEET in book.sheetnames:
        del book[LEGACY_D12_SHEET]
    for legacy, canonical in REFERENCE_SHEET_ALIASES.items():
        if legacy in book.sheetnames and canonical in book.sheetnames:
            del book[legacy]
    first = [GUIDE_SHEET, MASTER_SHEET, SMALL_HOURLY_SHEET, LAUNCH_HOURLY_SHEET, DAILY_SHEET, WEEKLY_SHEET,
             SMALL_DAILY_SHEET, SMALL_CURVE_SHEET, D12_SHEET]
    book._sheets = [book[name] for name in first] + [s for s in book.worksheets if s.title not in first]


def resolved_daily_rows(candidates, data, masters, history_file, today, raw_profiles=None):
    """One model/date row, field-wise precedence, never sum overlapping sources."""
    targets = {r["name"]: r for r in [*data.get("targets", []), *data.get("steady_targets", [])]}
    references = data.get("reference_daily", {})
    cancel_cumulative = data.setdefault("resolved_cancel_cumulative", {})
    history_cancel_days = data["resolved_cancel_history_days"] = set()
    for profile in data.get("actuals", []):
        previous_cancel = {iso_day(r.get("date")): r.get("cancel") for r in profile.get("cancel_days", [])}
        for day in profile.get("days", []):
            filename, sheet = _source_parts(profile.get("day_source"))
            provenance, kinds = {}, {}
            for key, field in ORDER_FIELDS.items():
                owner = day.get("_field_sources", {}).get(key)
                kinds[field] = "历史首销" if owner == "history" else "首销订单"
                provenance[field] = source_text({"file": history_file, "sheet": {"gross": GROSS_DAILY_SHEET, "small_to_big": "小转大当日数量", "direct": "直接大定当日数量"}.get(key, "车型汇总")}) if owner == "history" else source_text({"file": filename, "sheet": sheet})
            candidates.append({"订单分析代际名": profile["model"], "日期": date_cell(day.get("date")),
                "订单阶段": "首销", "生命周期": day.get("day"), **{c: day.get(k) for k, c in ORDER_FIELDS.items()},
                "来源类型": "已解析", "来源文件": filename, "来源Sheet": sheet, "_origins": provenance, "_kinds": kinds})
        # Read cancellation quantities independently of launch stage and its date spine.
        # An explicit daily value is authoritative; cumulative differencing requires
        # the immediately preceding calendar day, never just the previous record.
        for entry in profile.get("cancel_days", []):
            day_key = iso_day(entry.get("date"))
            value = entry.get("cancel")
            valid = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) and v >= 0
            if not calendar_day(day_key):
                continue
            if valid(value):
                cancel_cumulative[(profile["model"], day_key)] = value
            daily_cancel = entry.get("daily_cancel")
            if not valid(daily_cancel) and valid(value):
                prior_day = (date.fromisoformat(day_key)-timedelta(days=1)).isoformat()
                prior = previous_cancel.get(prior_day)
                if prior is None and day_key == targets.get(profile["model"], {}).get("small_start_date"):
                    prior = 0
                daily_cancel = value-prior if valid(prior) and value >= prior else None
            if valid(daily_cancel):
                filename, sheet = _source_parts(profile.get("cancel_source"))
                candidates.append({"订单分析代际名": profile["model"], "日期": date_cell(day_key),
                    "退订数量": daily_cancel, "来源类型": "小订退订",
                    "来源文件": filename, "来源Sheet": sheet})
    for item in data.get("history", []):
        model = item.get("generation") or item["model"]
        start = item.get("launch_date")
        if not start:
            continue
        for field, sheet in (("小转大", "小转大当日数量"), ("直接大定", "直接大定当日数量"), ("大定", GROSS_DAILY_SHEET), ("退订数量", "退订当日数量")):
            values = references.get(sheet, {}).get(item["model"], [])
            for index, value in enumerate(values[:max(int(item.get("days") or len(values)), 0)]):
                candidates.append({"订单分析代际名": model, "历史传播名": item["model"],
                    "日期": date_cell((date.fromisoformat(start) + timedelta(days=index)).isoformat()),
                    "订单阶段": "首销", "生命周期": f"D{index+1}", field: value,
                    "来源类型": "历史首销", "来源文件": history_file, "来源Sheet": sheet})
    for (model, day, field), (value, source) in data.get("mix_daily", {}).items():
        filename, sheet = _source_parts(source)
        candidates.append({"订单分析代际名": model, "日期": date_cell(day),
            ORDER_FIELDS[field]: value, "来源类型": "选配实际", "来源文件": filename, "来源Sheet": sheet})
    rows, ranks, origins = {}, {}, {}
    def rank_for(source, field, target, day):
        kind = source.get("来源类型")
        if field == "退订数量":
            if kind == "历史首销":
                start, end = target.get("launch_date"), target.get("end_date")
                return 0 if not start or (day >= start and (not end or day <= end)) else None
            return 1 if kind == "小订退订" else None
        if field == "小订数量":
            start, end = target.get("small_start_date"), target.get("small_end_date")
            if start and end and not start <= day <= end:
                return None
            if kind == "历史小订":
                return 0
            origin = source.get("来源文件") or ""
            return 2 if "小订退订" in origin else 1
        if kind == "已解析":
            kind = source.get("_kinds", {}).get(field, "首销订单")
        if "选配比例" in str(source.get("来源文件")):
            kind = "选配实际"
        start, end = target.get("launch_date"), target.get("end_date")
        known = bool(start and end)
        in_launch = bool(known and start <= day <= end)
        mix = kind in {"选配实际", "平销锁单", "无小订大定补缺"}
        if field in {"小转大", "直接大定"} and known and not in_launch:
            return None
        if known and not in_launch:
            return 0 if mix else None
        return {"历史首销": 0, "首销订单": 1 if known else 2,
                "选配实际": 3, "平销锁单": 3, "无小订大定补缺": 3}.get(kind, 4)
    fields = ["小订数量", *ORDER_FIELDS.values(), "退订数量"]
    for source in candidates:
        if source.get("来源类型") == "首销分时累计":
            continue
        day = iso_day(source.get("日期"))
        model = source.get("订单分析代际名")
        if not model or not calendar_day(day):
            continue
        day = day[:10]
        key = (model, day)
        row = rows.setdefault(key, {"订单分析代际名": model, "日期": date_cell(day)})
        for label in ("历史传播名", "生命周期"):
            if source.get(label) and not row.get(label):
                row[label] = source[label]
        for field in fields:
            if field in source:
                row.setdefault(field, None)
            value = source.get(field)
            if not isinstance(value, (int, float)) or isinstance(value, bool) or not math.isfinite(value) or value < 0:
                continue
            rank = rank_for(source, field, targets.get(model, {}), day)
            if rank is None:
                continue
            identity = (*key, field)
            if identity in ranks and rank >= ranks[identity]:
                continue
            ranks[identity] = rank
            row[field] = value
            if field == "退订数量" and source.get("来源类型") == "历史首销":
                history_cancel_days.add(key)
            origins.setdefault(key, {})[field] = source.get("_origins", {}).get(field) or source_text({"file": source.get("来源文件"), "sheet": source.get("来源Sheet")})
    backfill_launch_cumulative(rows, ranks, origins, data, targets,
                               raw_profiles if raw_profiles is not None else data.get("actuals", []), today)
    result = []
    for key, row in sorted(rows.items()):
        model, day = key
        target = targets.get(model, {})
        in_launch = bool(target.get("launch_date") and target.get("end_date") and target["launch_date"] <= day <= target["end_date"])
        if in_launch and row.get("大定") is None and all(row.get(field) is not None for field in ("小转大", "直接大定")):
            row["大定"] = row["小转大"] + row["直接大定"]
            origins.setdefault(key, {})["大定"] = "已选小转大+直接大定（日量相加）"
        if in_launch and target.get("has_small") is False and isinstance(row.get("大定"), (int, float)):
            # Business identity, not a historical reference-ratio estimate.
            row["小转大"] = 0 if row.get("小转大") is None else row["小转大"]
            row["直接大定"] = row["大定"] if row.get("直接大定") is None else row["直接大定"]
        stages = []
        small_start, small_end = target.get("small_start_date"), target.get("small_end_date")
        if "小订数量" in row and (not small_start or not small_end or small_start <= day <= small_end):
            stages.append("小订")
        if any(row.get(c) is not None for c in ORDER_FIELDS.values()) or "退订数量" in row or (in_launch and any(c in row for c in ORDER_FIELDS.values())):
            stages.append("阶段未知" if not target.get("launch_date") or not target.get("end_date") else
                          "首销前" if day < target["launch_date"] else "首销" if in_launch else "平销")
        row["订单阶段"] = "/".join(stages)
        row["日期状态"] = "已结束日" if day < today.isoformat() else "当日未结束" if day == today.isoformat() else "未来日期"
        if not stages:
            continue  # Do not emit empty padding outside every maintained window.
        if target.get("launch_date") and day >= target["launch_date"]:
            row["生命周期"] = f"D{(date.fromisoformat(day)-date.fromisoformat(target['launch_date'])).days+1}"
        by_origin = {}
        for field, origin in origins.get(key, {}).items():
            by_origin.setdefault(origin, []).append(field)
        row["字段来源"] = "；".join(f"{'/'.join(fields)}={origin}" for origin, fields in by_origin.items())
        result.append(row)
    for model, target in targets.items():
        if not forecast_trace_enabled(model):
            continue
        start, finish = calendar_day(target.get('launch_date')), calendar_day(target.get('end_date'))
        dates = [(date.fromisoformat(start)+timedelta(days=i)).isoformat()
                 for i in range((date.fromisoformat(finish)-date.fromisoformat(start)).days+1)] if start and finish and finish >= start else []
        selected = {iso_day(r['日期']): r for r in result if r['订单分析代际名'] == model}
        forecast_trace(model, '逐日选值概况', 首销窗口=[start, finish], 当前阶段=target.get('stage'),
                       候选模型键=sorted({str(r.get('订单分析代际名')) for r in candidates if forecast_trace_enabled(r.get('订单分析代际名'))}),
                       已选行数=len(selected), 各字段={field: forecast_trace_series(
                           [selected.get(day, {}).get(field) for day in dates]) for field in fields})
        focus = set(dates[:2]+dates[-3:])
        for field in fields:
            focus.update([day for day in dates if selected.get(day, {}).get(field) is None][:2])
        observed = {}
        for candidate in candidates:
            day = iso_day(candidate.get('日期'))[:10]
            if candidate.get('订单分析代际名') == model and day in focus:
                observed.setdefault(day, []).append({
                    '来源': source_text({'file': candidate.get('来源文件'), 'sheet': candidate.get('来源Sheet')}),
                    '类型': candidate.get('来源类型'), '值': {f: candidate[f] for f in fields if f in candidate},
                    '优先级': {f: rank_for(candidate, f, target, day) for f in fields if f in candidate}})
        for day in sorted(focus):
            forecast_trace(model, '逐日选值明细', 日期=day, 候选=observed.get(day, []),
                           最终值={f: selected.get(day, {}).get(f) for f in fields},
                           字段来源=selected.get(day, {}).get('字段来源'),
                           说明='优先级数字越小越优先；空白、非法数及不适用窗口不采用；真实0有效')
    return result


def backfill_launch_cumulative(rows, ranks, origins, data, targets, profiles, today):
    """Fill gaps from actual cumulative counts, never from reference percentages.

    Daily history and launch facts win. Adjacent cumulative differences precede
    generic mix fallback; non-adjacent snapshots can fill only cumulative fields.
    A contradictory higher-priority interval invalidates that source curve.
    """
    resolved, conflicts = {}, []
    for profile in profiles:
        model = profile["model"]
        target = targets.get(model, {})
        start, finish = calendar_day(target.get("launch_date")), calendar_day(target.get("end_date"))
        if not start or not finish or calendar_day(profile.get("cumulative_launch_date")) != start:
            forecast_trace(model, '累计源不参与补缺', 原因='首销窗口缺失或源D1未锚定维护起点',
                           维护开始=start, 维护结束=finish, 源D1=profile.get('cumulative_launch_date'),
                           源累计行数=len(profile.get('cumulative_days', [])))
            continue
        source = source_text(profile.get("cumulative_source") or profile.get("day_source"))
        for key, field in ORDER_FIELDS.items():
            anchors = {calendar_day(r.get("date")): r[key] for r in profile.get("cumulative_days", [])
                       if start <= calendar_day(r.get("date")) <= finish
                       and isinstance(r.get(key), (int, float)) and not isinstance(r[key], bool)
                       and math.isfinite(r[key]) and r[key] >= 0}
            if not anchors:
                forecast_trace(model, '累计源不参与补缺', 字段=field, 来源=source, 原因='维护窗口内没有有效累计数量')
                continue
            previous_day, previous_value = (date.fromisoformat(start)-timedelta(days=1)).isoformat(), 0
            incompatible = False
            for day, value in sorted(anchors.items()):
                interval = (date.fromisoformat(day)-date.fromisoformat(previous_day)).days
                known = [rows.get((model, stamp), {}).get(field)
                         if ranks.get((model, stamp, field), 99) < 2 else None
                         for stamp in ((date.fromisoformat(previous_day)+timedelta(days=i+1)).isoformat()
                                       for i in range(interval))]
                known_sum = sum(v for v in known if v is not None)
                delta = value-previous_value
                if delta < 0 or known_sum > delta or (all(v is not None for v in known) and known_sum != delta):
                    incompatible = True
                    break
                previous_day, previous_value = day, value
            if incompatible:
                forecast_trace(model, '累计源拒绝', 字段=field, 来源=source, 首冲突日期=day,
                               前一锚点日期=previous_day, 前一累计=previous_value, 当前累计=value,
                               累计增量=delta, 高优先级日量合计=known_sum, 高优先级区间日量=known,
                               影响='本字段整条累计来源不用于补缺，不改变已选日量')
                conflicts.append((model, source, field, day, delta, known_sum))
                continue
            forecast_trace(model, '累计源接受', 字段=field, 来源=source, 锚点数=len(anchors),
                           首末锚点=[min(anchors), max(anchors)], 结束日=finish, 结束日累计=anchors.get(finish))
            resolved[(model, field)] = anchors
            for day, value in sorted(anchors.items()):
                previous = (date.fromisoformat(day)-timedelta(days=1)).isoformat()
                base = 0 if day == start else anchors.get(previous)
                if base is None or ranks.get((model, day, field), 99) < 2:
                    continue
                rows.setdefault((model, day), {"订单分析代际名": model, "日期": date_cell(day)})[field] = value-base
                ranks[(model, day, field)] = 2
                origins.setdefault((model, day), {})[field] = source + "（相邻日累计差分）"
    data["resolved_launch_cumulative"] = resolved
    if conflicts:
        grouped = {}
        for model, source, field, day, delta, known_sum in conflicts:
            grouped.setdefault((model, source), []).append(f"{field}@{day}：累计增量{delta:g}/高优先级日量{known_sum:g}")
        logging.getLogger(__name__).warning(
            "[首销累计口径冲突] %d组 | 各字段首处冲突=%s | 保留高优先级真实日量，冲突累计量不用于补缺",
            len(conflicts), "；".join(f"{model}（{source}）：{'、'.join(items)}" for (model, source), items in grouped.items()))


def anchored_cumulative(values, dates, anchors=None, initial=0, prefer_anchor=False):
    """Continue a known cumulative baseline; gaps stay blank until a real anchor."""
    anchors = anchors or {}
    running, result = initial, []
    for day, value in zip(dates, values):
        running = running+value if running is not None and value is not None else None
        if day in anchors and (running is None or prefer_anchor):
            running = anchors[day]
        result.append(running)
    return result


def selected_cancel_cumulative(values, dates, anchors, initial, history_days, model):
    """Keep primary day facts; accept compatible cumulative evidence per point."""
    by_date = dict(zip(dates, values))
    compatible, rejected = {}, []
    previous_day = (date.fromisoformat(dates[0])-timedelta(days=1)).isoformat() if dates and dates[0] else ""
    previous_value = initial
    for day, value in sorted(anchors.items()):
        interval = [stamp for stamp in dates if previous_day < stamp <= day]
        known = [by_date[stamp] for stamp in interval if by_date[stamp] is not None]
        if previous_value is not None and any(stamp in history_days for stamp in interval):
            delta = value-previous_value
            if sum(known) > delta or (len(known) == len(interval) and sum(known) != delta):
                rejected.append(day)
                continue  # Reject this point, not another day's known quantity.
        compatible[day] = value
        previous_day, previous_value = day, value
    if rejected:
        forecast_trace(model, '退订累计逐日取舍', 未采用累计日期=rejected,
                       原因='日度累计与更高优先级整理表日量不一致', 处理='保留逐日选定数量，不整行舍弃')
    return anchored_cumulative(values, dates, compatible, initial, prefer_anchor=True)


def resolved_weekly_rows(candidates, raw_weeks=None):
    grouped = {}
    for row in candidates:
        grouped.setdefault((row[0], row[1]), {}).setdefault(row[2], row)
    result = []
    for _, stages in sorted(grouped.items()):
        parts = list(stages.values())
        row = list(parts[0][:11])
        row[2] = "/".join(stages)
        row[3] = min((p[3] for p in parts), key=iso_day)
        row[4] = max((p[4] for p in parts), key=iso_day)
        for i in range(5, 8):
            values = [p[i] for p in parts]
            row[i] = sum(values) if all(isinstance(v, (int, float)) for v in values) else None
            row[i+3] = "、".join(dict.fromkeys(str(p[i+3]) for p in parts if p[i+3]))
        # Only transition weeks need separate stage quantities. Ordinary weeks
        # already have an unambiguous stage and must not duplicate all measures.
        row += [stages[s][i] if len(parts) > 1 and s in stages else None for s in ("首销", "平销") for i in range(5, 8)]
        for i, field in enumerate(("gross", "net", "lock"), 5):
            raw = (raw_weeks or {}).get((row[0], row[1], field))
            if row[i] is None and raw and raw[0] is not None:
                row[i], row[i+3] = raw[0], str(raw[1]) + "（阶段拆分不足，保留原周总量）"
        result.append(row)
    return result


def reconcile_intraday_fields(daily, profiles, hourly_records):
    """Verify day fields against the selected hour snapshot once, during refresh.

    Equal dates or gross alone do not prove component consistency. Store only
    the verified field names, not another copy of their quantities.
    """
    valid = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) and v >= 0
    totals = {}
    for (model, day, _), record in hourly_records.items():
        key = (model, day)
        previous = totals.get(key, 0)
        totals[key] = previous + record[3] if previous is not None and valid(record[3]) else None
    snapshots = {(p["model"], d["date"]): d for p in profiles for d in p.get("hourly_days", [])}
    for row in daily:
        key = (row["订单分析代际名"], iso_day(row["日期"]))
        if key not in totals:
            continue
        snapshot, total = snapshots.get(key, {}), totals[key]
        if not valid(total) or not valid(row.get("大定")) or not valid(snapshot.get("gross")):
            row["首销分时核对"] = "大定快照缺失"
            continue
        if abs(row["大定"] - total) > .5 or abs(snapshot["gross"] - total) > .5:
            row["首销分时核对"] = "日/时大定数量不一致"
            continue
        fields = [label for metric, label in (("small_to_big", "小转大"), ("direct", "直接大定"), ("lock", "交车锁单"))
                  if valid(row.get(label)) and valid(snapshot.get(metric)) and abs(row[label] - snapshot[metric]) <= .5]
        row["首销分时核对"] = "一致：" + "/".join(fields) if fields else "分项缺失或不一致"


def complete_master_fields(book, masters, headers):
    """Persist deterministic defaults, scoped to the exact event and its D1."""
    fields = ["小订天数", "小订发布星期", "首销发布星期", "平销开始",
              "小订开始小时", "小订开始小时来源", "首销开始小时", "首销开始小时来源"]
    headers.extend(k for k in fields if k not in headers)
    anchors = {"上午": 9, "下午": 14, "晚上": 19}
    for stage, sheet, quantity in (("小订", SMALL_HOURLY_SHEET, "小时小订"),
                                    ("首销", LAUNCH_HOURLY_SHEET, "小时大定")):
        first = {}
        for row in table_records(book, sheet):
            hour, value = row.get("小时"), row.get(quantity)
            if (isinstance(hour, bool) or not isinstance(hour, (int, float)) or
                    not math.isfinite(hour) or int(hour) != hour or not 0 <= hour < 24 or
                    isinstance(value, bool) or not isinstance(value, (int, float)) or
                    not math.isfinite(value) or value <= 0):
                continue
            key = (row.get("订单分析代际名"), calendar_day(row.get("日期")))
            if key not in first or hour < first[key][0]:
                first[key] = (int(hour), row.get("来源文件") or sheet)
        for row in masters:
            start, end = calendar_day(row.get(stage + "开始")), calendar_day(row.get(stage + "结束"))
            if start:
                row[stage + "发布星期"] = "周" + "一二三四五六日"[date.fromisoformat(start).weekday()]
            if start and end and end >= start:
                row[stage + "天数"] = (date.fromisoformat(end) - date.fromisoformat(start)).days + 1
            if stage == "首销" and end:
                row["平销开始"] = date_cell((date.fromisoformat(end) + timedelta(days=1)).isoformat())
            maintained = anchors.get(str(row.get(stage + "发布时段") or "").strip())
            inferred_before = str(row.get(stage + "开始小时来源") or "").startswith("D1首个正数订单小时：")
            if inferred_before:
                maintained = None  # A previously inferred label is not a maintained anchor.
            evidence = first.get((row.get("订单分析代际名"), start)) if start else None
            if maintained is None and evidence:
                row[stage + "发布时段"] = "上午" if evidence[0] < 12 else "下午" if evidence[0] < 19 else "晚上"
            elif inferred_before:
                row[stage + "发布时段"] = "未维护"
            row[stage + "开始小时"] = maintained if maintained is not None else evidence[0] if evidence else None
            row[stage + "开始小时来源"] = ("原表维护时段" if maintained is not None else
                f"D1首个正数订单小时：{evidence[1]}" if evidence else "缺少本阶段D1正数分时")


def refresh_small_reference_tables(book, masters, headers, lookup, emit, today):
    """Finalize reservation quantities once, before launch features consume them."""
    from tools.refresh_sales_forecast_data import cumulative
    curves = table_records(book, SMALL_CURVE_SHEET)
    identities = {(r.get("订单分析代际名"), r.get("历史传播名"), iso_day(r.get("小订开始"))) for r in curves}
    for master in masters:
        identity = (master.get("订单分析代际名"), master.get("历史传播名"), iso_day(master.get("小订开始")))
        if identity not in identities and identity[2] and iso_day(master.get("小订结束")):
            curves.append(dict(master))
            identities.add(identity)
    for field in SMALL_REFERENCE_FIELDS:
        if field not in headers:
            headers.append(field)
    small_daily, small_curves = [], []
    small_span = 0
    for row in curves:
        model, name = row.get("订单分析代际名") or row.get("历史传播名"), row.get("历史传播名")
        start, end = iso_day(row.get("小订开始")), iso_day(row.get("小订结束"))
        master = next((r for r in masters if r.get("历史传播名") == name and
                       r.get("订单分析代际名") == model and iso_day(r.get("小订开始")) == start), None)
        if master is None:
            master = next((r for r in masters if r.get("历史传播名") == name and r.get("订单分析代际名") == model), None)
        if master is None:
            continue
        curve = [v for k, v in row.items() if k.startswith("D") and k[1:].isdigit()]
        while curve and curve[-1] is None:
            curve.pop()
        count = max((date.fromisoformat(end)-date.fromisoformat(start)).days+1, 0) if start and end else int(row.get("小订天数") or len(curve))
        values = [lookup.get((model, (date.fromisoformat(start)+timedelta(days=i)).isoformat()), {}).get("小订数量")
                  for i in range(count)] if start else [None]*count
        total = master.get("总小订") if master.get("总小订") is not None else row.get("总小订")
        ended = bool(end and end < today.isoformat())
        if values and all(v is not None for v in values):
            total = master.get("总小订") if master.get("总小订") is not None else sum(values)
            row["总量来源"] = "车型基本信息" if master.get("总小订") is not None else "汇总实际小订合计"
            row["来源文件"], row["来源Sheet"] = SUMMARY_NAME, DAILY_SHEET
            curve = cumulative(values, total or 0)
        if isinstance(total, (int, float)) and total > 0:
            derived = cumulative(values, total)
            curve = [derived[i] if i < len(derived) and derived[i] is not None
                     else curve[i] if i < len(curve) else None for i in range(max(len(derived), len(curve)))]
        # Availability for forecasting is separate from preserving source facts.
        if master.get("总小订") is None and isinstance(total, (int, float)) and total > 0:
            master["总小订"] = total
        master.update({"小订总量来源": row.get("总量来源"),
            "小订参考来源": source_text({"file": row.get("来源文件"), "sheet": row.get("来源Sheet")}),
            "小订参考总量有效": ended and isinstance(total, (int, float)) and total > 0})
        status = "未开始" if start and start > today.isoformat() else "进行中" if end and end >= today.isoformat() else "已结束" if end else "时间未维护"
        small_daily.append([name, status, *values])
        small_curves.append([name, status, *curve])
        if forecast_trace_enabled(model):
            forecast_trace(model, '小订参考计算', 历史键=name, 小订窗口=[start, end], 天数=count,
                           总小订=total, 总量来源=row.get('总量来源'), 参考终值有效=master.get('小订参考总量有效'),
                           当日数量=forecast_trace_series(values), 完成度=forecast_trace_series(curve))
        small_span = max(small_span, count, len(curve))
    small_headers = ["传播名", "小订阶段", *[f"D{i+1}" for i in range(small_span)]]
    for name, records in ((SMALL_CURVE_SHEET, small_curves), (SMALL_DAILY_SHEET, small_daily)):
        emit(name, small_headers, [r+[None]*(len(small_headers)-len(r)) for r in records],
             set(small_headers[2:]) if name == SMALL_CURVE_SHEET else set())


def refresh_reference_tables(book, data, emit, daily, today):
    """Rebuild reference features from the same resolved quantities as actuals."""
    from tools.refresh_sales_forecast_data import safe_div, day_structure_quality
    valid_quantity = lambda value: (isinstance(value, (int, float)) and not isinstance(value, bool)
                                    and math.isfinite(value) and value >= 0)
    masters = table_records(book, MASTER_SHEET)
    headers = [c.value for c in book[MASTER_SHEET][1]]
    headers.extend(k for k in ("首销阶段", "首销参考终值有效") if k not in headers)
    lookup = {(r["订单分析代际名"], iso_day(r["日期"])): r for r in daily}
    actuals = {p["model"]: p for p in data.get("actuals", [])}
    mapping_status = {p["model"]: p.get("mapping_status") for p in data.get("history", [])}
    targets = {p["name"]: p for p in data.get("targets", [])}
    # Resolve the denominator before any rates, D1/D2 features or curves.
    for r in masters:
        model = r.get("订单分析代际名")
        profile = actuals.get(model, {})
        if r.get("总小订") in (None, "") and iso_day(r.get("首销开始")) == targets.get(model, {}).get("launch_date") and profile.get("total_small", 0) > 0:
            r["总小订"] = profile["total_small"]
    refresh_small_reference_tables(book, masters, headers, lookup, emit, today)
    derived = {name: [] for name in ("小转大当日数量", "直接大定当日数量", GROSS_DAILY_SHEET, "退订当日数量", "小转大累计完成度",
        "直接大定累计完成度", GROSS_PROGRESS_SHEET, "累计小订转化率", "累计退订率", "累计直接大定占比", "每日直接大定占比")}
    metric_rows = []
    max_days = max((int(r.get("首销天数") or 0) for r in masters), default=0)
    for r in masters:
        model, name = r.get("订单分析代际名"), r.get("历史传播名")
        start, finish = iso_day(r.get("首销开始")), iso_day(r.get("首销结束"))
        forecast_trace(model, '参考基准进入计算', 历史键=name, 首销开始=start, 首销结束=finish,
                       维护天数=r.get('首销天数'), 已有总量={k:r.get(k) for k in
                           ('总小订','总小转大','总直接大定','总大定','总退订','首销期锁单','首销期留存大定')})
        if not model or not name:
            continue
        count = int(r.get("首销天数") or ((date.fromisoformat(finish)-date.fromisoformat(start)).days+1 if start and finish else 0))
        source_daily = {field: data.get("reference_daily", {}).get(sheet, {}).get(name, [])
                        for field, sheet in (("小转大", "小转大当日数量"), ("直接大定", "直接大定当日数量"), ("大定", GROSS_DAILY_SHEET), ("退订数量", "退订当日数量"))}
        count = max(count, max((i+1 for values in source_daily.values() for i, value in enumerate(values)
                                if isinstance(value, (int, float)) and not isinstance(value, bool)), default=0))
        if start and not finish:
            observed_end = max((day for (owner, day), row in lookup.items() if owner == model and day >= start
                                and any(row.get(field) is not None for field in ORDER_FIELDS.values())), default="")
            if observed_end:
                count = max(count, (date.fromisoformat(observed_end)-date.fromisoformat(start)).days+1)
        max_days = max(max_days, count)
        dates = [(date.fromisoformat(start)+timedelta(days=i)).isoformat() for i in range(max(count, 0))] if start else [""]*count
        selected = [dict(lookup.get((model, day), {})) for day in dates]
        if not start:
            for i, row in enumerate(selected):
                row.update({field: values[i] for field, values in source_daily.items() if i < len(values)})
        s, d, g, c = ([day.get(field) for day in selected] for field in ("小转大", "直接大定", "大定", "退订数量"))
        cumulative_counts = {field: anchored_cumulative([day.get(field) for day in selected], dates,
            data.get("resolved_launch_cumulative", {}).get((model, field))) for field in ORDER_FIELDS.values()}
        cancel_anchors = {day: value for (owner, day), value in data.get("resolved_cancel_cumulative", {}).items()
                          if owner == model and start and start <= day <= (finish or day)}
        prior_day = (date.fromisoformat(start)-timedelta(days=1)).isoformat() if start else ""
        cancel_base = data.get("resolved_cancel_cumulative", {}).get((model, prior_day), 0)
        if (model, prior_day) not in data.get("resolved_cancel_cumulative", {}) and start in cancel_anchors:
            raw_start = next((day for day in actuals.get(model, {}).get("cancel_days", []) if iso_day(day.get("date")) == start), {})
            raw_daily = raw_start.get("daily_cancel")
            cancel_base = (cancel_anchors[start]-raw_daily if valid_quantity(raw_daily) and cancel_anchors[start] >= raw_daily else None)
        history_days = {day for owner, day in data.get("resolved_cancel_history_days", ()) if owner == model}
        cumulative_counts["退订数量"] = selected_cancel_cumulative(c, dates, cancel_anchors, cancel_base, history_days, model)
        if forecast_trace_enabled(model):
            forecast_trace(model, 'Dn映射及累计结果', 代际查找键=model, 历史查找键=name, 展开天数=count,
                           首末日期=[dates[0], dates[-1]] if dates else [],
                           命中日记录数=sum(bool(row) for row in selected),
                           日量={f:forecast_trace_series([row.get(f) for row in selected]) for f in (*ORDER_FIELDS.values(), '退订数量')},
                           累计={f:forecast_trace_series(v) for f,v in cumulative_counts.items()}, 退订期初基数=cancel_base)
        ended = bool(start and finish and finish < today.isoformat())
        r["首销阶段"] = "已结束" if ended else "未开始" if start > today.isoformat() else "进行中" if start and finish else "时间未维护"
        r["首销参考终值有效"] = ended
        terminal_index = dates.index(finish) if finish in dates else -1
        for field, total in (("小转大", "总小转大"), ("直接大定", "总直接大定"), ("大定", "总大定"),
                             ("留存大定", "首销期留存大定"), ("交车锁单", "首销期锁单"), ("退订数量", "总退订")):
            # A missing terminal does not erase a maintained source value.
            # Completed cumulative facts fill gaps; an ongoing sum is not a final.
            terminal = cumulative_counts[field][terminal_index] if terminal_index >= 0 else None
            forecast_trace(model, '终值判定', 字段=total, 结束日期=finish, 结束日索引=terminal_index,
                           已有维护值=r.get(total), 结束日累计=terminal,
                           动作=('首销结束日未落入Dn序列' if terminal_index < 0 else
                                 '累计到结束日仍未知，保留已有值或留空' if terminal is None else
                                 '写入末日累计' if r.get(total) in (None, '') or field == '退订数量' else '保留已有维护值'))
            if terminal is not None and (r.get(total) in (None, "") or field == "退订数量"):
                r[total] = terminal
        for output, numerator, denominator in (("小订转化率", "总小转大", "总小订"), ("直接大定占比", "总直接大定", "总大定"),
            ("退订率", "总退订", "总小订"), ("留存大定率", "首销期留存大定", "总大定"),
            ("留存大定到锁单率", "首销期锁单", "首销期留存大定"), ("大定到锁单率", "首销期锁单", "总大定")):
            top, bottom, previous = r.get(numerator), r.get(denominator), r.get(output)
            calculated = safe_div(top, bottom) if valid_quantity(top) and valid_quantity(bottom) and bottom > 0 else None
            maintained_valid = valid_quantity(previous) and previous <= 1
            forecast_trace(model, '比例判定', 字段=output, 分子字段=numerator, 分子=r.get(numerator),
                           分母字段=denominator, 分母=r.get(denominator), 已有值=r.get(output), 计算值=calculated,
                           动作=('按最终选定数量重算' if calculated is not None else
                                 '数量不足，保留有效维护比例' if maintained_valid else '数量不足且维护比例无效，留空'))
            # Derive rates after quantity resolution, even when an old rate is nonempty.
            # Without quantities keep only a valid independently maintained rate;
            # never replace unknown inputs with zero or swap lock-rate denominators.
            r[output] = calculated if calculated is not None else previous if maintained_valid else None
        fields = ("总小订", "总小转大", "小订转化率", "总大定", "总直接大定", "直接大定占比", "总退订", "退订率", "首销期留存大定", "留存大定率", "首销期锁单", "大定到锁单率")
        r["字段完整度"] = sum(r.get(k) is not None for k in fields)/len(fields)
        derived["小转大当日数量"].append([name, *s])
        derived["直接大定当日数量"].append([name, *d])
        derived[GROSS_DAILY_SHEET].append([name, *g])
        # Seed the exported roll-forward at launch D1, not at the latest date.
        # Keep c unchanged: by-day facts and D1/D2 features are actual daily counts.
        cancel_export = list(c)
        if cancel_export:
            cancel_export[0] = cumulative_counts["退订数量"][0]
        derived["退订当日数量"].append([name, *cancel_export])
        for title, field, denominator in (("小转大累计完成度", "小转大", "总小转大"), ("直接大定累计完成度", "直接大定", "总直接大定"),
            (GROSS_PROGRESS_SHEET, "大定", "总大定"), ("累计小订转化率", "小转大", "总小订"), ("累计退订率", "退订数量", "总小订")):
            progress = [safe_div(value, r.get(denominator)) for value in cumulative_counts[field]]
            if title == "累计退订率":
                progress = [value if start and start <= day and (not finish or day <= finish) else None
                            for day, value in zip(dates, progress)]
            derived[title].append([name, *progress])
        cg, cd = cumulative_counts["大定"], cumulative_counts["直接大定"]
        derived["累计直接大定占比"].append([name, *[safe_div(a, b) for a, b in zip(cd, cg)]])
        derived["每日直接大定占比"].append([name, *[safe_div(a, b) for a, b in zip(d, g)]])
        s1, s2 = (s+[None, None])[:2]
        d1, d2 = (d+[None, None])[:2]
        g1, g2 = (g+[None, None])[:2]
        c1, c2 = (c+[None, None])[:2]
        def pair(a, b):
            return a+b if a is not None and b is not None else None
        s12, d12, g12, c12 = pair(s1, s2), pair(d1, d2), pair(g1, g2), pair(c1, c2)
        triples = ((s1, d1, g1), (s2, d2, g2), (s12, d12, g12))
        checks = []
        for (a, b, v), label in zip(triples, ("D1", "D2", "D1+D2")):
            if all(x is not None for x in (a, b, v)):
                checks.append(day_structure_quality(a, b, v, label))
            elif v is not None and any(x is not None and x > v*1.02 for x in (a, b)):
                checks.append((False, label+"单项订单来源高于总大定"))
            else:
                checks.append((False, label+"数量未完整"))
        statuses = ["通过" if ok else "部分字段有效" if issue.endswith("数量未完整") and any(x is not None for x in triple)
                    else "数据不足或异常" for (ok, issue), triple in zip(checks, triples)]
        # An unrelated missing component must not erase an independently valid
        # D1/terminal or D2/D1 ratio. A proven structure conflict still rejects it.
        # Keep computable facts even when their structure/time is unsuitable
        # for a reference. The reader applies these explicit statuses.
        (vs1, vd1, vg1), (vs2, vd2, vg2), (vs12, vd12, vg12) = triples
        for i, day_index in enumerate((0, 1, 1)):
            if not start:
                statuses[i] = "时间未维护"
            elif day_index < len(dates) and dates[day_index] >= today.isoformat():
                statuses[i] = "当日未结束" if dates[day_index] == today.isoformat() else "未来日期"
        status = mapping_status.get(name) or ("已映射" if model != name else "未映射·暂按传播名")
        forecast_trace(model, 'D1D2计算', D1={'小转大':s1,'直接大定':d1,'大定':g1,'退订':c1},
                       D2={'小转大':s2,'直接大定':d2,'大定':g2,'退订':c2},
                       分母={k:r.get(k) for k in ('总小订','总小转大','总直接大定','总大定')},
                       口径状态=statuses, 结构说明=checks)
        metric_rows.append([name, model, status, r.get("产品档位"), count, r.get("总小转大"), r.get("总直接大定"), r.get("总大定"), r.get("总小订"),
            s1, safe_div(vs1, r.get("总小转大")), s2, safe_div(vs2, r.get("总小转大")), s12, safe_div(vs12, r.get("总小转大")), safe_div(vs2, vs1), safe_div(vs1, vs12),
            d1, safe_div(vd1, r.get("总直接大定")), d2, safe_div(vd2, r.get("总直接大定")), d12, safe_div(vd12, r.get("总直接大定")), safe_div(vd2, vd1), safe_div(vd1, vd12),
            g1, safe_div(vs1, vg1), safe_div(vd1, vg1), g2, safe_div(vs2, vg2), safe_div(vd2, vg2), g12, safe_div(vs12, vg12), safe_div(vd12, vg12),
            c1, safe_div(c1, r.get("总小订")), c2, safe_div(c2, r.get("总小订")), c12, safe_div(c12, r.get("总小订")), safe_div(c2, c1), safe_div(c1, c12),
            *statuses, "；".join(issue for _, issue in checks if issue)])
    emit(D12_SHEET, D12_HEADERS, metric_rows, {k for k in D12_HEADERS if "/" in k or k.endswith("率")})
    day_headers = [f"D{i+1}" for i in range(max_days)]
    for title, records in derived.items():
        padded = [row+[None]*(max_days+1-len(row)) for row in records]
        emit(title, ["传播名", *day_headers], padded, set(day_headers) if "当日数量" not in title else set())
    complete_master_fields(book, masters, headers)
    emit(MASTER_SHEET, headers, [[r.get(k) for k in headers] for r in masters],
         {k for k in headers if k.endswith("率") or "占比" in k or k == "字段完整度"})
    if SMALL_TOTAL_SHEET in book.sheetnames:
        del book[SMALL_TOTAL_SHEET]


def restrict_launch_reference(item, master, today):
    """Apply reference eligibility to web inputs, never to saved workbook facts."""
    start, end = calendar_day(master.get("首销开始")), calendar_day(master.get("首销结束"))
    completed = max((today-date.fromisoformat(start)).days, 0) if start else 0
    ended = bool(start and end and end < today.isoformat() and master.get("首销参考终值有效", True))
    item["terminal_available"] = ended
    for field in ("daily_orders", "daily_small", "daily_direct", "small_progress", "direct_progress",
                  "gross_progress", "direct_share_progress", "cancel_progress"):
        item[field] = list(item.get(field) or [])[:completed]
    if not ended:
        for field in ("gross", "net", "lock", "small_to_big", "direct", "cancel", "conversion",
                      "direct_share", "cancel_rate", "net_rate", "lock_rate",
                      "d1_small_completion", "d2_small_completion", "d12_small_completion",
                      "d1_direct_completion", "d2_direct_completion", "d12_direct_completion"):
            item[field] = None
        for field in ("small_progress", "direct_progress", "gross_progress"):
            item[field] = []
    if master.get("小订参考总量有效") is False:
        item["small"] = 0
        for field in ("conversion", "cancel_rate", "d1_cancel_rate", "d2_cancel_rate", "d12_cancel_rate"):
            item[field] = None
    for prefix, required in (("d1", 1), ("d2", 2), ("d12", 2)):
        if completed < required:
            for field in list(item):
                if field.startswith(prefix + "_"):
                    item[field] = False if field.endswith("_valid") else None
    if completed < 2:
        for field in ("small_d2_d1", "small_d1_d12", "direct_d2_d1", "direct_d1_d12",
                      "cancel_d2_d1", "cancel_d1_d12"):
            item[field] = None


def read_public_forecast(book, history, today=None):
    """Reconstruct raw candidates from public business cells; resolve stages later."""
    from core.model_identity import model_key
    from modules.sales_forecast import _read_model_mapping, _read_model_master, _master_record, _brand
    current = today or date.today()

    masters = table_records(book, MASTER_SHEET)
    days = table_records(book, DAILY_SHEET)
    canonical = "字段来源" in [c.value for c in book[DAILY_SHEET][1]]
    mapping, model_master = _read_model_mapping(), _read_model_master()
    grouped = {}
    for row in days:
        model = row.get("订单分析代际名")
        if canonical:
            if str(row.get("订单阶段")) in {"首销", "小订/首销", "阶段未知", "小订/阶段未知"}:
                grouped.setdefault((model, "首销订单"), []).append(row)
            if "小订" in str(row.get("订单阶段")):
                grouped.setdefault((model, "逐日小订"), []).append(row)
                grouped.setdefault((model, "历史小订"), []).append(row)
            if "平销" in str(row.get("订单阶段")) and row.get("交车锁单") is not None:
                grouped.setdefault((model, "平销锁单"), []).append(row)
        else:
            grouped.setdefault((model, row.get("来源类型")), []).append(row)
    def lifecycle(row):
        return int(str(row.get("生命周期") or "D0")[1:])
    for item in history:
        if canonical:
            master = next((r for r in masters if r.get("历史传播名") == item["model"]), {})
            model = item.get("generation") or master.get("订单分析代际名")
            start = calendar_day(item.get("launch_date") or master.get("首销开始"))
            end = calendar_day(item.get("end_date") or master.get("首销结束"))
            selected = sorted([r for r in days if r.get("订单分析代际名") == model and start and
                start <= calendar_day(r.get("日期")) and (not end or calendar_day(r.get("日期")) <= end)], key=lambda r: calendar_day(r["日期"]))
            by_date = {calendar_day(r["日期"]): r for r in selected}
            last = max((d for d, r in by_date.items() if r.get("大定") is not None), default="")
            selected = [by_date.get((date.fromisoformat(start)+timedelta(days=i)).isoformat(), {})
                        for i in range((date.fromisoformat(last)-date.fromisoformat(start)).days+1)] if last else []
        else:
            selected = sorted([r for r in days if r.get("来源类型") == "历史首销" and r.get("历史传播名") == item["model"]], key=lifecycle)
        item["daily_orders"] = [r.get("大定") for r in selected]
        item["hourly_curve"] = []

    profiles, windows = [], {}
    by_model = {}
    cancel_curves = {r.get("传播名"): r for r in table_records(book, "累计退订率")}
    for row in masters:
        if not row.get("订单来源文件"):
            continue
        model = row["订单分析代际名"]
        if model in by_model:
            continue
        small_total = row.get("总小订") if row.get("小订参考总量有效", True) else None
        windows[model_key(model)] = {"generation": model, "history_model": row.get("历史传播名"),
            "primary_generation": row.get("primary_generation") or model,
            "secondary_generation": row.get("secondary_generation") or "",
            "has_small": row.get("有小订") not in (False, 0, "否"),
            "launch_date": calendar_day(row.get("首销开始")), "end_date": calendar_day(row.get("首销结束")),
            "small_start_date": calendar_day(row.get("小订开始")), "small_end_date": calendar_day(row.get("小订结束")),
            "days": row.get("首销天数"), "small": small_total, "source_sheet": MASTER_SHEET}
        profile = {"model": model, "launch_date": windows[model_key(model)]["launch_date"],
                   "days": [], "cancel_days": [], "hourly_days": [], "small_hourly_days": [],
                   "small_daily_days": [], "small_daily_sources": {}}
        if canonical:
            profile.update(_summary_resolved=True, total_small=small_total or 0,
                day_source={"file": SUMMARY_NAME, "sheet": DAILY_SHEET})
        launch_start = profile["launch_date"]
        launch_end = windows[model_key(model)]["end_date"]
        for r in grouped.get((model, "首销订单"), []):
            order_date = calendar_day(r.get("日期"))
            if canonical and (not order_date or order_date > current.isoformat() or (launch_start and order_date < launch_start) or
                              (launch_end and order_date > launch_end)):
                continue
            profile["days"].append({"date": order_date, "day": r.get("生命周期"),
                                    **{k: r.get(c) for k, c in ORDER_FIELDS.items()}})
            if canonical:
                # D1/D2 reference features use daily cancellations, while the
                # cumulative curve can include cancellations before launch.
                profile["days"][-1]["daily_cancel"] = r.get("退订数量")
            # Keep dated actuals when the launch window is missing, but do not
            # invent D1 or apply a relative cancellation curve to those rows.
            if canonical and launch_start:
                index = (date.fromisoformat(order_date)-date.fromisoformat(launch_start)).days+1
                rate = cancel_curves.get(row.get("历史传播名"), {}).get(f"D{index}")
                if isinstance(rate, (int, float)) and profile["total_small"] > 0:
                    profile["days"][-1]["cancel"] = rate * profile["total_small"]
            if not canonical:
                profile["day_source"] = {"file": r.get("来源文件"), "sheet": r.get("来源Sheet")}
        profile["days"].sort(key=lambda r: (r.get("date") or "", int(str(r.get("day") or "D0")[1:])))
        for r in grouped.get((model, "逐日小订"), []):
            d = iso_day(r.get("日期"))
            if canonical and d > current.isoformat():
                continue
            profile["small_daily_days"].append({"date": d, "orders": r.get("小订数量")})
            profile["small_daily_sources"][d] = {"file": SUMMARY_NAME, "sheet": DAILY_SHEET} if canonical else {"file": r.get("来源文件"), "sheet": r.get("来源Sheet")}
        profiles.append(profile)
        by_model[model] = profile
    for row in ([] if canonical else table_records(book, SMALL_TOTAL_SHEET)):
        profile = by_model.get(row.get("订单分析代际名"))
        if profile is None:
            continue
        prefix = "cancel" if row.get("来源类型") == "小订退订" else "launch"
        value = row.get("总小订")
        profile[prefix + "_total_small"] = value or 0
        profile[prefix + "_total_small_valid"] = isinstance(value, (int, float)) and value > 0
        profile[prefix + "_total_small_reason"] = row.get("总量口径")
        profile[prefix + "_total_small_estimated"] = row.get("总量口径") == "首销累计进度反推"
        profile["cancel_source" if prefix == "cancel" else "day_source"] = {"file": row.get("来源文件"), "sheet": row.get("来源Sheet")}
        if prefix == "cancel":
            start, end = calendar_day(row.get("数据开始")), calendar_day(row.get("数据结束"))
            profile["cancel_latest_date"] = end
            if start and end:
                span = (date.fromisoformat(end) - date.fromisoformat(start)).days + 1
                profile["cancel_days"] = [{"date": (date.fromisoformat(start) + timedelta(days=i)).isoformat()} for i in range(max(span, 0))]
    for sheet, key, field, label, source_key in (
        (SMALL_HOURLY_SHEET, "small_hourly_days", "orders", "小时小订", "small_hour_source"),
        (LAUNCH_HOURLY_SHEET, "hourly_days", "gross", "小时大定", "hour_source")):
        buckets = {}
        for row in table_records(book, sheet):
            model, d = row.get("订单分析代际名"), iso_day(row.get("日期"))
            profile = by_model.get(model)
            if profile is None or d > current.isoformat():
                continue
            value = row.get(label)
            if not isinstance(value, (int, float)) or isinstance(value, bool) or not math.isfinite(value) or value < 0:
                continue  # Keep partial source components in Excel; an unknown gross is not zero.
            profile[source_key] = {"file": row.get("来源文件"), "sheet": sheet}
            bucket = buckets.setdefault((model, d), {"date": d, "last_hour": 0, field: 0, "hours": []})
            hour = int(row["小时"])
            bucket["last_hour"] = max(bucket["last_hour"], hour)
            bucket["hours"].append({"hour": hour, field: value})
            if isinstance(value, (int, float)):
                bucket[field] += value
        daily_index = {(model, iso_day(r.get("日期"))): r for (model, kind), entries in grouped.items()
                       if kind == "首销订单" for r in entries} if key == "hourly_days" else {}
        for (model, d), bucket in sorted(buckets.items()):
            bucket["hours"].sort(key=lambda r: r["hour"])
            if key == "hourly_days":
                daily = daily_index.get((model, d), {})
                status = str(daily.get("首销分时核对") or "未核对，日/时快照不可混用")
                verified = status.removeprefix("一致：").split("/") if status.startswith("一致：") else []
                bucket["component_status"] = status
                for metric, column in (("small_to_big", "小转大"), ("direct", "直接大定"), ("lock", "交车锁单")):
                    bucket[metric] = daily.get(column) if column in verified else None
                for row in grouped.get((model, "首销分时累计"), []):
                    if iso_day(row.get("日期")) == d:
                        bucket.update({k: row[c] for k, c in ORDER_FIELDS.items() if row.get(c) is not None})
            by_model[model][key].append(bucket)
    small = []
    curve_sheet = SMALL_CURVE_SHEET if SMALL_CURVE_SHEET in book.sheetnames else LEGACY_SMALL_CURVE_SHEET
    for row in table_records(book, curve_sheet):
        if "总小订" not in row:
            identity = row.get("历史传播名") or row.get("传播名")
            metadata = next((m for m in masters if m.get("历史传播名") == identity), {})
            filename, source_sheet = _source_parts(metadata.get("小订参考来源"))
            row = {**metadata, **row, "总小订": metadata.get("总小订") if metadata.get("小订参考总量有效", True) else None,
                   "总量来源": metadata.get("小订总量来源"), "来源文件": filename, "来源Sheet": source_sheet}
        name, generation = row.get("历史传播名"), row.get("订单分析代际名") or ""
        start, end = calendar_day(row.get("小订开始")), calendar_day(row.get("小订结束"))
        first_date = calendar_day(row.get("首条数据日期"))
        selected = [r for r in grouped.get((generation or name, "历史小订"), [])
                    if (canonical and start and start <= iso_day(r.get("日期")) <= (end or "9999")) or (not canonical and r.get("历史传播名") == name and
                    (not first_date or iso_day(r.get("日期")) ==
                     (date.fromisoformat(first_date) + timedelta(days=lifecycle(r)-1)).isoformat()))]
        selected.sort(key=lambda r: iso_day(r.get("日期")) if canonical else lifecycle(r))
        if canonical:
            # A historical reference uses completed calendar days, not row order.
            # Keep live snapshots in profiles above; missing/duplicate days here
            # stay unknown. Confirmed sparse-source zeros were resolved upstream.
            by_date = {}
            for observation in selected:
                day = calendar_day(observation.get("日期"))
                if day and day < current.isoformat():
                    by_date[day] = None if day in by_date else observation.get("小订数量")
            last = max(by_date, default="")
            dates = [(date.fromisoformat(start) + timedelta(days=i)).isoformat()
                     for i in range((date.fromisoformat(last) - date.fromisoformat(start)).days + 1)] if start and last else []
            daily = [by_date.get(day) for day in dates]
        else:
            daily = [r.get("小订数量") for r in selected]
            dates = [iso_day(r.get("日期")) for r in selected]
        curve = [v for k, v in row.items() if k.startswith("D") and k[1:].isdigit()]
        while curve and curve[-1] is None:
            curve.pop()
        final_total = row.get("总小订")
        complete = isinstance(final_total, (int, float)) and final_total > 0 and bool(end and end < (today or date.today()).isoformat())
        total = final_total if complete else sum(v for v in daily if v is not None)
        actual = bool(canonical and selected) or row.get("来源Sheet") != "小订进度"
        cumulative, running, prefix_complete = [], 0, True
        for value in daily:
            prefix_complete = prefix_complete and value is not None
            if prefix_complete:
                running += value
            cumulative.append(min(running / max(total, 1), 1) if prefix_complete and complete else None)
        if not actual and curve and total > 0:
            terminal = curve[-1] or 0
            if terminal > 0:
                daily, assigned, previous = [], 0, 0
                for index, value in enumerate(curve):
                    progress = min(max(value / terminal, 0), 1)
                    count = max(int(round((progress-previous)*total)), 0) if index < len(curve)-1 else max(total-assigned, 0)
                    daily.append(count)
                    assigned += count
                    previous = progress
                dates = [(date.fromisoformat(start) + timedelta(days=i)).isoformat() for i in range(len(daily))] if start else []
        attrs = _master_record(model_master, name, generation) or {}
        item = {"model": name, "generation": generation, "mapped": model_key(row.get("原始名称") or name) in mapping,
            "brand": attrs.get("品牌") or _brand(name), "tier": attrs.get("产品档位") or "未维护",
            "energy": attrs.get("能源类型") or "未维护", "node": attrs.get("发布类型") or attrs.get("发布节点") or "未维护",
            "release_start_hour": row.get("小订开始小时", attrs.get("小订开始小时")),
            "launch_period": row.get("小订发布时段") or attrs.get("小订发布时段") or "未维护", "small_start_date": start, "small_end_date": end,
            "days": row.get("小订天数") or ((date.fromisoformat(end)-date.fromisoformat(start)).days+1 if start and end else max(len(daily), 1)), "total": total, "total_complete": complete,
            "total_source": row.get("总量来源"),
            "daily_orders": daily, "dates": dates, "small_progress": cumulative, "standard_progress": curve,
            "d1_share": (daily[0] or 0) / max(total, 1) if daily and complete else 0,
            "small_hourly_curve": [], "daily_actual": actual, "source_sheet": row.get("来源Sheet"),
            "source_model": row.get("原始名称"), "event_id": "|".join(("small", model_key(name),
                model_key(generation or name), start or "no-date", str(row.get("来源Sheet") or "no-sheet")))}
        if not actual and complete:
            item.pop("total_complete", None)
            item["d1_share"] = min(curve[0] / curve[-1], 1) if curve and curve[-1] > 0 else 0
        small.append(item)
    # Flat sales belongs to the primary generation, even when launch inputs
    # are split into independent editions. Derive its window without duplicating
    # master rows or assigning primary-only quantities to each edition.
    parents = {}
    for window in list(windows.values()):
        if window.get("secondary_generation"):
            parents.setdefault(window["primary_generation"], []).append(window)
    for parent, children in parents.items():
        if not all(w.get("launch_date") and w.get("end_date") for w in children):
            continue
        windows[model_key(parent)] = {**children[-1], "generation": parent, "history_model": parent,
            "primary_generation": parent, "secondary_generation": "",
            "launch_date": min(w["launch_date"] for w in children),
            "end_date": max(w["end_date"] for w in children)}
        if parent not in by_model:
            by_model[parent] = {"model": parent, "days": [], "hourly_days": [], "small_hourly_days": [],
                               "small_daily_days": [], "small_daily_sources": {}, "_summary_resolved": canonical}
            profiles.append(by_model[parent])
    steady = []
    week_records = table_records(book, WEEKLY_SHEET)
    for model, profile in by_model.items():
        window = windows[model_key(model)]
        end = window.get("end_date")
        start = (date.fromisoformat(end) + timedelta(days=1)).isoformat() if end else ""
        daily_rows = sorted([r for r in grouped.get((model, "平销锁单"), []) if iso_day(r.get("日期")) <= current.isoformat()],
                            key=lambda r: iso_day(r.get("日期")))
        weeks = []
        for row in week_records:
            if row.get("订单分析代际名") != model or row.get("订单阶段") != "平销":
                continue
            # Order-mix lock totals remain visible, but are not lock-mix references.
            if "锁单选配比例" not in _source_parts(row.get("锁单来源"))[0]:
                continue
            a, b = calendar_day(row.get("统计开始")), calendar_day(row.get("统计结束"))
            if not a or not b or not start:
                continue
            first, last = date.fromisoformat(a), date.fromisoformat(b)
            if first.weekday() != 0 or (last-first).days != 6 or last >= current or a < start:
                continue
            if not isinstance(row.get("交车锁单"), (int, float)) or row["交车锁单"] < 0:
                continue
            weeks.append({"period": row["周期"], "start_date": a, "end_date": b, "lock": row["交车锁单"], "gross": row.get("大定") if isinstance(row.get("大定"), (int, float)) and row["大定"] >= 0 else None})
        weeks.sort(key=lambda r: r["start_date"])
        if not daily_rows and not weeks:
            continue
        # Unknown gaps do not invalidate independent, completed weekly facts.
        # Preserve calendar positions; consumers check adjacency before ratios.
        for week in weeks:
            week["week"] = (date.fromisoformat(week["start_date"])-date.fromisoformat(weeks[0]["start_date"])).days // 7 + 1
        attrs = _master_record(model_master, model) or {}
        filename = (daily_rows[0].get("来源文件") if daily_rows else None) or SUMMARY_NAME
        sheets = list(dict.fromkeys(r.get("来源Sheet") for r in daily_rows if r.get("来源Sheet")))
        steady.append({"model": model, "generation": model, "mapped": True,
            "brand": attrs.get("品牌") or _brand(model), "tier": attrs.get("产品档位") or "未维护",
            "energy": attrs.get("能源类型") or "未维护", "node": attrs.get("发布类型") or "未维护",
            "steady_start_date": start, "weeks": weeks,
            "daily": [{"date": iso_day(r["日期"]), "lock": r.get("交车锁单"), "gross": r.get("大定"), "complete": iso_day(r["日期"]) < current.isoformat()} for r in daily_rows],
            "source_file": filename, "source_sheet": WEEKLY_SHEET, "daily_source_sheets": sheets,
            "event_id": "|".join(("steady", model_key(model), str(filename))), "lock_rate": None, "launch_days": None})
    if canonical:
        for item in history:
            master = next((r for r in masters if r.get("历史传播名") == item["model"]), {})
            restrict_launch_reference(item, master, current)
    return profiles, windows, small, steady


@contextmanager
def summary_scope(path, workbook=None):
    path = Path(path)
    owns_workbook = workbook is None
    if owns_workbook:
        workbook = load_workbook(path, read_only=False, data_only=True)
    token = None
    try:
        if DAILY_SHEET in workbook.sheetnames and GUIDE_SHEET in workbook.sheetnames:
            active = {"path": path, "inputs": WorkbookView(workbook.worksheets),
                      "workbook": workbook, "visible": True}
            token = ACTIVE_SUMMARY.set(active)
            yield active
            return
        if SNAPSHOT_SHEET in workbook.sheetnames:
            raise ValueError("销量汇总仍为旧版缓存结构，请运行main重新生成可见数据表")
        if INDEX_SHEET not in workbook.sheetnames:
            raise ValueError(f"{path.name}缺少数据来源目录，请运行main重新生成")
        groups, inputs, source_ranges = {}, list(workbook.worksheets), {}
        headers = [cell.value for cell in workbook[INDEX_SHEET][1]]
        range_columns = ([headers.index(label) for label in ("文件数据开始日期", "文件数据结束日期")]
                         if all(label in headers for label in ("文件数据开始日期", "文件数据结束日期")) else [])
        seen = set()
        for record in workbook[INDEX_SHEET].iter_rows(min_row=2, values_only=True):
            kind, filename, original, stored, *_ = record
            if not stored:
                continue
            if stored not in workbook.sheetnames or (filename, original) in seen:
                raise ValueError(f"汇总目录重复或缺少Sheet：{filename}/{original}")
            seen.add((filename, original))
            workbook[stored]._log_origin = (original, 0, 0)
            view = SheetView(workbook[stored], original)
            if kind == "订单":
                groups.setdefault(filename, []).append(view)
                if range_columns:
                    try:
                        start, end = (date.fromisoformat(str(record[column])) for column in range_columns)
                        source_ranges[filename] = (start, end) if start <= end else None
                    except (TypeError, ValueError):
                        source_ranges[filename] = None
            elif kind in {"历史", "映射"}:
                inputs = [sheet for sheet in inputs if sheet.title != original]
                inputs.append(view)
        store = WorkbookStore(path.parent)
        store.items = [WorkbookItem(Path(name), WorkbookView(sheets)) for name, sheets in groups.items()]
        for item in store.items:
            if item.path.name in source_ranges:
                item._boundary_date_range = source_ranges[item.path.name]
        active = {"path": path, "inputs": WorkbookView(inputs), "store": store}
        token = ACTIVE_SUMMARY.set(active)
        yield active
    finally:
        if token is not None:
            ACTIVE_SUMMARY.reset(token)
        if owns_workbook:
            workbook.close()
