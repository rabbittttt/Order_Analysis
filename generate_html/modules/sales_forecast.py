from __future__ import annotations

import logging
import math
import re
from copy import deepcopy
from datetime import date, datetime, timedelta
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from core.china_calendar import calendar_payload
from core.model_identity import model_key, usable_attribute
from core.excel import display_period, grain_from_sheet, is_aggregate_generation, parse_metric_sheet, sheet_subject, subject_type
from core.models import Dashboard, SourceRef, Subject
from core.forecast_summary import ACTIVE_SUMMARY, SUMMARY_NAME, input_path, open_input, summary_scope


LOGGER = logging.getLogger(__name__)


CODE_ROOT = Path(__file__).resolve().parents[2]
PROJECT_ROOT = CODE_ROOT.parent if CODE_ROOT.name.lower() == "scripts" else CODE_ROOT
FORECAST_INPUT_ROOT = PROJECT_ROOT / "input_file" / "销量预测输入文件"
RAW_FORECAST_DATA = FORECAST_INPUT_ROOT / "小订及首销数据整理.xlsx"
MODEL_MASTER_PATH = FORECAST_INPUT_ROOT / "车型基本信息.xlsx"
HISTORY_CANDIDATES = (
    PROJECT_ROOT / "output_file" / SUMMARY_NAME,
    FORECAST_INPUT_ROOT / SUMMARY_NAME,
    FORECAST_INPUT_ROOT / "销量预测数据汇总.xlsx",
    FORECAST_INPUT_ROOT / "小订及首销预测二次处理.xlsx",
    FORECAST_INPUT_ROOT / "小订及首销数据整理.xlsx",
    FORECAST_INPUT_ROOT / "小订及首销期数据整理.xlsx",
    FORECAST_INPUT_ROOT / "小订及首销期历史基准.xlsx",
    FORECAST_INPUT_ROOT / "小订及首销期数据整理_测试数据.xlsx",
)
MODEL_MAPPING_CANDIDATES = (
    MODEL_MASTER_PATH,
    FORECAST_INPUT_ROOT / "传播代际名映射.xlsx",
    FORECAST_INPUT_ROOT / "车型代际映射.xlsx",
)

TASKS = [
    ("hourly", "D1 全天与分时参考", "按发布节点、星期及时段预测或滚动反推首销第一天终值"),
    ("small_progress", "方法一 · 小转大累计进度", "先按当前已发生的小转大结构与斜率选参考，再用历史参考同期完成率反推终局"),
    ("direct_progress", "方法一 · 直接大定累计进度", "先按当前D1/D2订单来源结构选参考，再用历史参考同期完成率反推终局"),
    ("conversion", "方法二 · 小订转化率", "根据产品、可观测小转大结构、D2斜率与早期退订质量选择最终转化率"),
    ("direct_share", "方法二 · 直接大定占比", "根据发布节点及D1/D2可观测订单来源结构选择最终占比"),
    ("lock", "大定到锁单率", "由历史总大定、首销期锁单结果及当前早期退订质量确定大定到锁单率"),
    ("daily_slope", "每日斜率（仅展示）", "独立评分与主辅参考只用于比较已结束真实日环比，不参与销量预测终局或余量分配"),
    ("daily", "到天基础曲线", "历史每日总大定先适配到当前首销周期并剔除原日期类型系数，再除以剔除日历影响后的D1形成相对日强度；两种方法均用主辅加权后的基础曲线衔接前一真实日，再应用当前日历系数与封顶渐进权重分配各自剩余量"),
]

SMALL_ORDER_TASKS = [
    ("small_total", "小订最终总量", "按产品属性、小订窗口、线索量和互联网热度选择小订总量参考"),
    ("small_hourly", "小订D1分时", "按发布类型、星期、时段和历史分时完成率反推D1终值"),
    ("small_progress", "小订累计进度", "按当前已发生小订累计量与历史同期完成率反推最终总量"),
]

STEADY_TASKS = [
    ("steady_level", "平销锁单周量级", "优先本车型平销真实锁单，其次首销直接大定折算锁单，最后历史车型"),
    ("steady_curve", "平销锁单周曲线", "趋势与量级采用同级来源，真实日冻结后预测滚动锁单量"),
]

def _number(value: Any, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _optional_number(value: Any) -> float | None:
    """Preserve the difference between a missing cell and an explicit zero."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _stable_implied_total(
    numerators: list[Any],
    rates: list[Any],
    *,
    minimum_points: int = 2,
    relative_tolerance: float = .05,
) -> tuple[int, bool, str]:
    """Infer a fixed denominator only when repeated cumulative ratios agree."""
    candidates = [
        _number(numerator) / _number(rate)
        for numerator, rate in zip(numerators, rates)
        if _number(numerator) > 0 and _number(rate) > 0
    ]
    if len(candidates) < minimum_points:
        return (int(round(candidates[-1])) if candidates else 0, False, "有效反推点不足")
    ordered = sorted(candidates)
    median = ordered[len(ordered) // 2]
    spread = max(abs(value - median) / max(abs(median), 1) for value in candidates)
    if spread > relative_tolerance:
        return int(round(candidates[-1])), False, f"多期反推分母波动{spread:.1%}"
    return int(round(median)), True, f"{len(candidates)}期反推一致"


def _optional_count(values: list[Any], index: int) -> int | None:
    value = values[index] if index < len(values) else None
    if not isinstance(value, (int, float)) or not math.isfinite(float(value)) or float(value) < 0:
        return None
    return int(round(float(value)))


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError, OverflowError):
            return None
    text = str(value or "").strip()
    normalized = (
        text.replace("年", "-")
        .replace("月", "-")
        .replace("日", "")
        .replace("/", "-")
        .replace(".", "-")
        .replace("\\", "-")
    )
    normalized = re.sub(
        r"(?<!\d)(\d{2})-(\d{1,2})-(\d{1,2})(?!\d)",
        lambda match: f"20{match.group(1)}-{match.group(2)}-{match.group(3)}",
        normalized,
        count=1,
    )
    try:
        return datetime.fromisoformat(normalized).date()
    except (TypeError, ValueError):
        match = re.search(r"(?<!\d)(20\d{2})-(\d{1,2})-(\d{1,2})(?!\d)", normalized)
        if not match:
            return None
        try:
            return date(*(int(part) for part in match.groups()))
        except ValueError:
            return None


def _as_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
            if isinstance(converted, datetime):
                return converted
            if isinstance(converted, date):
                return datetime.combine(converted, datetime.min.time())
        except (TypeError, ValueError, OverflowError):
            return None
    text = str(value or "").strip()
    if not text:
        return None
    normalized = (
        text.replace("年", "-")
        .replace("月", "-")
        .replace("日", " ")
        .replace("时", ":")
        .replace("分", ":")
        .replace("秒", "")
        .replace("/", "-")
        .replace(".", "-")
        .replace("\\", "-")
        .strip()
        .rstrip(":")
    )
    normalized = re.sub(
        r"(?<!\d)(\d{2})-(\d{1,2})-(\d{1,2})(?!\d)",
        lambda match: f"20{match.group(1)}-{match.group(2)}-{match.group(3)}",
        normalized,
        count=1,
    )
    try:
        return datetime.fromisoformat(normalized)
    except ValueError:
        parsed_date = _as_date(value)
        if parsed_date is None:
            return None
        match = re.search(
            r"20\d{2}-\d{1,2}-\d{1,2}[ T]+(\d{1,2})(?::(\d{1,2}))?(?::(\d{1,2}))?",
            normalized,
        )
        if not match:
            return datetime.combine(parsed_date, datetime.min.time())
        hour, minute, second = (int(part or 0) for part in match.groups())
        try:
            return datetime.combine(parsed_date, datetime.min.time()).replace(
                hour=hour, minute=minute, second=second
            )
        except ValueError:
            return None


def _iso(value: Any) -> str:
    parsed = _as_date(value)
    return parsed.isoformat() if parsed else ""


def _configured_forecast_date(value: Any) -> date | None:
    """Parse the optional manual as-of date used for lifecycle classification."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(value).strip()):
        raise ValueError("config.json 的 forecast_as_of_date 必须是 YYYY-MM-DD 或留空")
    parsed = _as_date(value)
    if parsed is None:
        raise ValueError("config.json 的 forecast_as_of_date 必须是 YYYY-MM-DD 或留空")
    return parsed


def _forecast_stage(launch_date: Any, end_date: Any = None, days: Any = None, today: date | None = None) -> dict[str, Any]:
    """Classify the launch window by absolute dates; D1 is the launch date itself."""
    current = today or date.today()
    start = _as_date(launch_date)
    explicit_end = _as_date(end_date)
    maintained_span = max(int(_number(days, 0)), 0)
    if start is None:
        return {"key": "unknown", "label": "时间缺失", "day": 0, "days": 0, "launch_date": "", "end_date": ""}
    if explicit_end and explicit_end < start:
        return {
            "key": "unknown", "label": "首销窗口冲突", "day": 0, "days": 0,
            "launch_date": start.isoformat(), "end_date": explicit_end.isoformat(),
        }
    end = explicit_end or (start + timedelta(days=maintained_span - 1) if maintained_span > 0 else None)
    if end is None:
        return {
            "key": "unknown", "label": "首销截止日期与天数缺失", "day": 0, "days": 0,
            "launch_date": start.isoformat(), "end_date": "",
        }
    span = (end - start).days + 1
    if current < start:
        key, label, day_number = "before", "首销期未开始", 0
    elif end and current > end:
        key, label, day_number = "ended", "首销期已结束", span
    else:
        key, label = "active", "首销期进行中"
        day_number = min(max((current - start).days + 1, 1), span)
    return {
        "key": key,
        "label": label,
        "day": day_number,
        "days": span,
        "launch_date": start.isoformat(),
        "end_date": end.isoformat() if end else "",
    }


def _small_order_stage(start_date: Any, end_date: Any = None, today: date | None = None) -> dict[str, Any]:
    """Classify the small-order window; D1 is kept separate for hourly forecasting."""
    current = today or date.today()
    start = _as_date(start_date)
    end = _as_date(end_date)
    if start is None:
        return {"key": "unknown", "label": "小订时间缺失", "day": 0, "start_date": "", "end_date": "", "days": 0}
    if end is None:
        return {
            "key": "unknown", "label": "小订结束日期缺失", "day": 0,
            "start_date": start.isoformat(), "end_date": "", "days": 0,
        }
    if end < start:
        return {
            "key": "unknown", "label": "小订窗口冲突", "day": 0,
            "start_date": start.isoformat(), "end_date": end.isoformat(), "days": 0,
        }
    span = (end - start).days + 1
    if current < start:
        key, label, day_number = "before", "小订D1未到", 0
    elif current == start:
        key, label, day_number = "d1", "小订D1进行中", 1
    elif current > end:
        key, label, day_number = "ended", "小订期已结束", span
    else:
        key, label = "active", "小订D1已过"
        day_number = min((current - start).days + 1, span)
    return {
        "key": key,
        "label": label,
        "day": day_number,
        "start_date": start.isoformat(),
        "end_date": end.isoformat(),
        "days": span,
    }


def _weekday(value: Any) -> str:
    parsed = _as_date(value)
    return ("周一", "周二", "周三", "周四", "周五", "周六", "周日")[parsed.weekday()] if parsed else "未维护"


def _model_key(value: Any) -> str:
    return model_key(value)


def _maintained_alias_keys(value: Any) -> set[str]:
    """Return exact and year-equivalent keys for one explicitly maintained name."""
    text = str(value or "").strip()
    if not text:
        return set()
    variants = {text}
    variants.add(re.sub(r"(?<!\d)(\d{2})\s*款", lambda match: f"20{match.group(1)}款", text))
    variants.add(re.sub(r"(?<!\d)20(\d{2})\s*款", lambda match: f"{match.group(1)}款", text))
    return {key for item in variants if (key := _model_key(item))}


def _canonical_model(value: Any, mapping: dict[str, str] | None = None) -> str:
    """Resolve a propagation name, generation name, or maintained alias to one generation name."""
    name = str(value or "").strip()
    if not name:
        return ""
    resolved = (mapping if mapping is not None else _read_model_mapping()).get(_model_key(name))
    return str(resolved or name).strip()


def _same_model(left: Any, right: Any) -> bool:
    mapping = _read_model_mapping()
    left_key, right_key = _model_key(_canonical_model(left, mapping)), _model_key(_canonical_model(right, mapping))
    return bool(left_key and right_key and left_key == right_key)


def _read_model_master(path: Path | None = None) -> dict[str, dict[str, Any]]:
    source = path or input_path(MODEL_MASTER_PATH)
    if not source.exists():
        return {}
    workbook = open_input(source, read_only=True, data_only=True)
    try:
        if "车型基本信息" not in workbook.sheetnames:
            return {}
        sheet = workbook["车型基本信息"]
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        result: dict[str, dict[str, Any]] = {}
        for values in sheet.iter_rows(min_row=2, values_only=True):
            record = dict(zip(headers, values))
            model = str(record.get("历史传播名") or record.get("历史车型名") or "").strip()
            generation = str(record.get("订单分析代际名") or "").strip()
            keys = [model, generation, *str(record.get("原始表简称/别名") or "").split("|")]
            for value in keys:
                key = _model_key(value)
                if key:
                    result.setdefault(key, record)
        return result
    finally:
        workbook.close()


def _master_record(master: dict[str, dict[str, Any]], *names: Any) -> dict[str, Any] | None:
    return next((master[_model_key(name)] for name in names if _model_key(name) in master), None)


def _master_values(master: dict[str, dict[str, Any]], field: str) -> list[str]:
    placeholders = {"未维护", "待维护"}
    return sorted({
        str(record.get(field) or "").strip()
        for record in master.values()
        if str(record.get(field) or "").strip() not in placeholders | {""}
    })


def _read_model_mapping() -> dict[str, str]:
    active = ACTIVE_SUMMARY.get()
    if active is not None:
        if "mapping" not in active:
            active["mapping"] = _read_model_mapping_file.__wrapped__(active["path"])
        return active["mapping"]
    mapping_path = next((path for path in MODEL_MAPPING_CANDIDATES if path.exists()), None)
    if mapping_path is None:
        return {}
    return _read_model_mapping_file(mapping_path)


@lru_cache(maxsize=4)
def _read_model_mapping_file(mapping_path) -> dict[str, str]:
    workbook = open_input(mapping_path, read_only=True, data_only=True)
    try:
        mapping_sheet = next((name for name in ("车型基本信息", "传播名代际映射", "车型代际映射") if name in workbook.sheetnames), "")
        if not mapping_sheet:
            return {}
        sheet = workbook[mapping_sheet]
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        result: dict[str, str] = {}
        conflicts: dict[str, set[str]] = {}

        def register(value: Any, generation: str) -> None:
            for key in _maintained_alias_keys(value):
                existing = result.get(key)
                if existing and _model_key(existing) != _model_key(generation):
                    conflicts.setdefault(key, {existing}).add(generation)
                    result.pop(key, None)
                    continue
                if key not in conflicts:
                    result[key] = generation

        for values in sheet.iter_rows(min_row=2, values_only=True):
            record = dict(zip(headers, values))
            model = str(record.get("历史传播名") or record.get("历史车型名") or "").strip()
            generation = str(record.get("订单分析代际名") or "").strip()
            if model and generation:
                register(model, generation)
                register(generation, generation)
                aliases = str(record.get("原始表简称/别名") or "")
                for alias in aliases.split("|"):
                    if alias.strip():
                        register(alias, generation)
        for key, generations in conflicts.items():
            LOGGER.warning(
                "[映射诊断] 标准键冲突=%s | 对应代际=%s | 处理=该键不自动匹配，请修正车型基本信息",
                key,
                "、".join(sorted(generations)),
            )
        return result
    finally:
        workbook.close()


def _read_stage_windows() -> tuple[Path | None, dict[str, dict[str, Any]]]:
    """Read absolute small-order and launch windows, normalized to order-analysis generation names."""
    source = input_path(RAW_FORECAST_DATA)
    if not source.exists():
        return None, {}
    workbook = open_input(source, read_only=True, data_only=True)
    try:
        if "车型汇总" not in workbook.sheetnames:
            return source, {}
        sheet = workbook["车型汇总"]
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        model_mapping = _read_model_mapping()
        result: dict[str, dict[str, Any]] = {}
        for values in sheet.iter_rows(min_row=2, values_only=True):
            record = dict(zip(headers, values))
            model = str(record.get("车型") or record.get("传播名") or "").strip()
            if not model:
                continue
            generation = model_mapping.get(_model_key(model), model)
            launch_date = _iso(record.get("开始大定日期") or record.get("发布日"))
            end_date = _iso(record.get("小转大结束日期") or record.get("首销截止"))
            days = max(int(_number(record.get("首销期天数"), 0)), 0)
            item = {
                "generation": generation,
                "history_model": model,
                "small_start_date": _iso(record.get("小订开始日期")),
                "small_end_date": _iso(record.get("小订结束日期")),
                "launch_date": launch_date,
                "end_date": end_date,
                "days": days,
                "small": int(round(_number(record.get("总小订")))),
                "source_sheet": sheet.title,
            }
            key = _model_key(generation)
            existing = result.get(key)
            if not existing or item["launch_date"] > existing["launch_date"]:
                result[key] = item
        return source, result
    finally:
        workbook.close()


def _stage_window(windows: dict[str, dict[str, Any]], model: str) -> dict[str, Any] | None:
    exact = windows.get(_model_key(model))
    return exact or next((item for item in windows.values() if _same_model(item.get("generation"), model)), None)


def _first_record_value(record: dict[str, Any], *names: str) -> Any:
    return next((record.get(name) for name in names if record.get(name) not in (None, "")), None)


def _read_small_order_history(path: Path | None = None) -> tuple[Path | None, list[dict[str, Any]]]:
    """Read actual by-day small-order curves; standardized progress is fallback evidence only."""
    source = path or input_path(RAW_FORECAST_DATA)
    if not source.exists():
        return None, []
    workbook = open_input(source, read_only=False, data_only=True)
    try:
        if "小订by天" not in workbook.sheetnames:
            LOGGER.warning(
                "[销量预测数据诊断] 文件=%s | 未找到'小订by天'Sheet | "
                "处理=小订历史缺少真实逐日曲线，仅能按车型汇总回退合成，网页小订预测证据将受限",
                source.name,
            )
            return source, []
        model_mapping = _read_model_mapping()
        model_master = _read_model_master()
        # A generation may contain multiple propagation events. Only aliases
        # explicitly declared on the same master row identify the same event.
        event_aliases: dict[str, dict[str, str]] = {}
        for record in model_master.values():
            history_name = str(record.get("历史传播名") or record.get("历史车型名") or "").strip()
            if not history_name:
                continue
            for name in [history_name, *str(record.get("原始表简称/别名") or "").split("|")]:
                for alias_key in _maintained_alias_keys(name):
                    event_aliases.setdefault(alias_key, {})[_model_key(history_name)] = history_name
        event_names = {
            key: next(iter(names.values()))
            for key, names in event_aliases.items() if len(names) == 1
        }
        summaries: list[dict[str, Any]] = []
        if "车型汇总" in workbook.sheetnames:
            sheet = workbook["车型汇总"]
            headers = [str(cell.value or "").strip() for cell in sheet[1]]
            for values in sheet.iter_rows(min_row=2, values_only=True):
                record = dict(zip(headers, values))
                model = str(record.get("车型") or record.get("传播名") or "").strip()
                if model:
                    record["_model"] = model
                    record["_generation"] = _canonical_model(model, model_mapping)
                    summaries.append(record)

        standardized = _progress_rows(workbook, "小订进度")
        sheet = workbook["小订by天"]
        current_header: list[Any] = []
        parsed: dict[tuple[str, str], dict[str, Any]] = {}
        skip_labels = {"", "合计", "预测", "```", "小订"}
        for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            label = str(values[0] or "").strip() if values else ""
            if label == "小订":
                current_header = list(values)
                continue
            if label in skip_labels or not current_header:
                continue
            date_columns = [
                (index, parsed_date)
                for index, value in enumerate(current_header)
                if index > 0 and (parsed_date := _as_date(value)) is not None
            ]
            if not date_columns:
                LOGGER.warning(
                    "[销量预测字段校验] 传播名=%s | Sheet=%s | 行=%d | 小订by天表头没有可识别日期列 | "
                    "处理=整条真实逐日记录不参与预测",
                    label, sheet.title, row_number,
                )
                continue
            invalid_daily: list[str] = []
            daily: list[int] = []
            for day_index, (column_index, header_date) in enumerate(date_columns, start=1):
                raw_value = values[column_index] if column_index < len(values) else None
                number = _optional_number(raw_value)
                if number is None or not math.isfinite(number):
                    invalid_daily.append(f"D{day_index}({header_date.isoformat()})缺失或非数字")
                elif number < 0:
                    invalid_daily.append(f"D{day_index}({header_date.isoformat()})为负数{number:g}")
                else:
                    daily.append(int(round(number)))
            if invalid_daily:
                LOGGER.warning(
                    "[销量预测字段校验] 传播名=%s | Sheet=%s | 行=%d | 小订by天异常=%s | "
                    "处理=整条真实逐日记录不参与预测，不把缺失或负数改写为0",
                    label, sheet.title, row_number, "；".join(invalid_daily),
                )
                continue
            total_columns = [
                index for index, value in enumerate(current_header)
                if str(value or "").strip() == "合计"
            ]
            if total_columns:
                total_column = max(total_columns)
                raw_total = values[total_column] if total_column < len(values) else None
                total_value = _optional_number(raw_total)
                if total_value is None or not math.isfinite(total_value) or total_value < 0:
                    LOGGER.warning(
                        "[销量预测字段校验] 传播名=%s | Sheet=%s | 行=%d | 合计=%r无效 | "
                        "处理=整条真实逐日记录不参与预测",
                        label, sheet.title, row_number, raw_total,
                    )
                    continue
                total = int(round(total_value))
                daily_sum = sum(daily)
                if abs(total - daily_sum) > 1:
                    LOGGER.warning(
                        "[销量预测字段校验] 传播名=%s | Sheet=%s | 行=%d | 合计=%d与逐日求和=%d不一致 | "
                        "处理=整条真实逐日记录不参与预测",
                        label, sheet.title, row_number, total, daily_sum,
                    )
                    continue
            else:
                total = sum(daily)
                LOGGER.warning(
                    "[销量预测字段校验] 传播名=%s | Sheet=%s | 行=%d | 未提供合计列 | "
                    "处理=使用全部日期列求和%d，不把最后一个日期误作合计",
                    label, sheet.title, row_number, total,
                )
            if total <= 0:
                continue
            first_header_date = date_columns[0][1]
            generation = _canonical_model(label, model_mapping)
            summary_candidates = [
                record for record in summaries
                if _model_key(record.get("_generation")) == _model_key(generation)
            ]
            dated_candidates = [
                record for record in summary_candidates
                if first_header_date and _as_date(record.get("小订开始日期")) == first_header_date
            ]
            exact_candidates = [
                record for record in summary_candidates
                if _model_key(record.get("_model")) == _model_key(label)
            ]
            matching = dated_candidates or exact_candidates
            summary = matching[0] if len(matching) == 1 else {}
            if len(matching) > 1:
                LOGGER.warning(
                    "[映射诊断] 小订历史=%s | 开始日期=%s | 车型汇总命中%d条 | "
                    "处理=不绑定汇总值，仅使用该传播事件的by天真实合计",
                    label, first_header_date.isoformat() if first_header_date else "缺失", len(matching),
                )
            elif not matching:
                LOGGER.warning(
                    "[映射诊断] 小订历史=%s | 标准代际=%s | 开始日期=%s | "
                    "处理=未绑定车型汇总，仅使用by天真实合计，不跨传播事件猜测",
                    label, generation, first_header_date.isoformat() if first_header_date else "缺失",
                )
            header_dates = [parsed_date for _, parsed_date in date_columns]
            discontinuity = next((
                (index, current_date)
                for index, current_date in enumerate(header_dates)
                if (current_date - first_header_date).days != index
            ), None)
            if discontinuity:
                index, current_date = discontinuity
                LOGGER.warning(
                    "[销量预测字段校验] 代际=%s | Sheet=%s | 小订by天表头日期D%d=%s与首日间隔不符（D1=%s）| "
                    "处理=仅保留中断前%d个完整日参与日期对齐",
                    generation, sheet.title, index + 1, current_date.isoformat(), first_header_date.isoformat(),
                    index,
                )
                daily = daily[:index]
                header_dates = header_dates[:index]
            dates = [(first_header_date + timedelta(days=index)).isoformat() for index in range(len(daily))] if first_header_date else []
            master_record = _master_record(model_master, label, generation) or {}
            summary_start = _iso(summary.get("小订开始日期"))
            summary_end = _iso(summary.get("小订结束日期"))
            fallback_curve = next(
                (curve for name, curve in standardized.items() if _same_model(name, label) or _same_model(name, generation)),
                [],
            )
            summary_total = _optional_number(summary.get("总小订"))
            if summary_total is not None and summary_total > 0 and abs(summary_total - total) > 1:
                LOGGER.warning(
                    "[销量预测字段校验] 代际=%s | 车型汇总总小订%d与by天合计%d不一致 | "
                    "处理=按规则优先采用车型汇总总小订作为终值，请核对两处口径",
                    generation, int(round(summary_total)), total,
                )
            final_total = int(round(summary_total)) if summary_total is not None and summary_total > 0 else total
            running = 0
            cumulative = []
            for value in daily:
                running += value
                cumulative.append(min(running / max(final_total, 1), 1))
            item = {
                "model": label,
                "generation": generation,
                "mapped": _model_key(label) in model_mapping,
                "brand": str(master_record.get("品牌") or _brand(label)),
                "tier": usable_attribute(master_record.get("产品档位")),
                "energy": usable_attribute(master_record.get("能源类型")),
                "node": usable_attribute(master_record.get("发布类型") or master_record.get("发布节点")),
                "launch_period": usable_attribute(master_record.get("发布时段")),
                "small_start_date": summary_start or (dates[0] if dates else ""),
                "small_end_date": summary_end or (dates[-1] if dates else ""),
                "days": max(int(_number(summary.get("小订天数"), len(daily))), len(daily), 1),
                "total": final_total,
                "total_source": "车型汇总" if (_optional_number(summary.get("总小订")) or 0) > 0 else "小订by天合计",
                "leads": _number(_first_record_value(summary, "线索量", "累计线索量", "线索数")),
                "heat": _number(_first_record_value(summary, "互联网热度", "热度指数", "网络热度")),
                "daily_orders": daily,
                "dates": dates,
                "small_progress": cumulative,
                "standard_progress": [value for value in fallback_curve if value is not None],
                "d1_share": daily[0] / max(final_total, 1),
                "small_hourly_curve": [],
                "daily_actual": True,
                "source_sheet": sheet.title,
            }
            parsed[(_model_key(label), item["small_start_date"])] = item
        parsed_windows = {
            (_model_key(item.get("generation") or item.get("model")), str(item.get("small_start_date") or ""))
            for item in parsed.values()
        }
        for summary in summaries:
            generation_key = _model_key(summary.get("_generation"))
            summary_start_key = _iso(summary.get("小订开始日期"))
            if (generation_key, summary_start_key) in parsed_windows:
                continue
            model = str(summary.get("车型") or summary.get("传播名") or "").strip()
            generation = _canonical_model(model, model_mapping)
            total = int(round(_number(summary.get("总小订"))))
            fallback_curve = next(
                (curve for name, curve in standardized.items() if _same_model(name, model) or _same_model(name, generation)),
                [],
            )
            numeric_curve = [max(_number(value), 0) for value in fallback_curve if value is not None]
            terminal = numeric_curve[-1] if numeric_curve else 0
            if total <= 0:
                continue
            normalized = [min(value / terminal, 1) for value in numeric_curve] if terminal > 0 else []
            daily = []
            previous = 0.0
            assigned = 0
            for index, value in enumerate(normalized):
                count = total - assigned if index == len(normalized) - 1 else int(round(max(value - previous, 0) * total))
                count = max(count, 0)
                daily.append(count)
                assigned += count
                previous = value
            start_date = _as_date(summary.get("小订开始日期"))
            dates = [(start_date + timedelta(days=index)).isoformat() for index in range(len(daily))] if start_date else []
            master_record = _master_record(model_master, model, generation) or {}
            LOGGER.warning(
                "[销量预测估算] 代际=%s | 小订by天无真实逐日行 | "
                "处理=按'小订进度'标准化曲线合成逐日形状（daily_actual=False），仅作历史参考形状",
                model,
            )
            parsed[(_model_key(model), dates[0] if dates else "")] = {
                "model": model,
                "generation": generation,
                "mapped": _model_key(model) in model_mapping,
                "brand": str(master_record.get("品牌") or _brand(model)),
                "tier": usable_attribute(master_record.get("产品档位")),
                "energy": usable_attribute(master_record.get("能源类型")),
                "node": usable_attribute(master_record.get("发布类型") or master_record.get("发布节点")),
                "launch_period": usable_attribute(master_record.get("发布时段")),
                "small_start_date": _iso(summary.get("小订开始日期")),
                "small_end_date": _iso(summary.get("小订结束日期")),
                "days": max(int(_number(summary.get("小订天数"), len(daily))), len(daily), 1),
                "total": total,
                "total_source": "车型汇总",
                "leads": _number(_first_record_value(summary, "线索量", "累计线索量", "线索数")),
                "heat": _number(_first_record_value(summary, "互联网热度", "热度指数", "网络热度")),
                "daily_orders": daily,
                "dates": dates,
                "small_progress": [],
                "standard_progress": numeric_curve,
                "d1_share": normalized[0] if normalized else 0,
                "small_hourly_curve": [],
                "daily_actual": False,
                "source_sheet": "小订进度",
            }
        deduped: dict[tuple[str, str], dict[str, Any]] = {}
        for item in parsed.values():
            source_model = str(item.get("model") or "")
            item["source_model"] = source_model
            item["model"] = event_names.get(_model_key(source_model), source_model)
            key = (_model_key(item.get("model")), str(item.get("small_start_date") or ""))
            existing = deduped.get(key)
            if existing is None:
                deduped[key] = item
                continue
            def preference(candidate: dict[str, Any]) -> tuple[bool, bool, int]:
                return (
                    bool(candidate.get("daily_actual")),
                    _model_key(candidate.get("source_model")) == _model_key(candidate.get("model")),
                    len(candidate.get("dates") or []),
                )
            keep_new = preference(item) > preference(existing)
            deduped[key] = item if keep_new else existing
            LOGGER.warning(
                "[映射诊断] 小订历史=%s | 开始日期=%s | 出现重复记录 | "
                "处理=按明确别名关系归并，保留%s，避免同一传播事件重复参与匹配（不累加数量）",
                item.get("model"), item.get("small_start_date") or "缺失",
                deduped[key].get("source_model"),
            )
        for item in deduped.values():
            item["event_id"] = "small|{}|{}|{}|{}".format(
                _model_key(item.get("model")),
                _model_key(item.get("generation") or item.get("model")),
                item.get("small_start_date") or "no-date",
                item.get("source_sheet") or "no-sheet",
            )
        unmapped_small = [str(item.get("model")) for item in deduped.values() if not item.get("mapped")]
        if unmapped_small:
            LOGGER.warning(
                "[映射诊断] 模块=小订预测 | 未映射传播名=%d个 | 明细=%s | "
                "处理=仍作为历史参考并在网页标记未映射，不猜测性绑定订单分析代际",
                len(unmapped_small), "、".join(unmapped_small),
            )
        return source, list(deduped.values())
    finally:
        workbook.close()


def _steady_history_paths() -> list[Path]:
    candidates = [FORECAST_INPUT_ROOT / name for name in STEADY_HISTORY_FILENAMES]
    candidates.extend(path for path in FORECAST_INPUT_ROOT.glob("*.xlsx") if "平销" in path.stem)
    if RAW_FORECAST_DATA.exists():
        candidates.append(RAW_FORECAST_DATA)
    unique: list[Path] = []
    for path in candidates:
        if path.exists() and path not in unique:
            unique.append(path)
    return unique


def _read_legacy_steady_history() -> tuple[list[dict[str, Any]], list[SourceRef]]:
    """Read long-form weekly steady-sales history. Flat-period orders are direct orders only."""
    model_mapping = _read_model_mapping()
    model_master = _read_model_master()
    grouped: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    sources: list[SourceRef] = []
    model_headers = ("代际名", "车型/代际名", "传播名", "车型")
    week_headers = ("周序号", "平销周", "周次", "周期")
    direct_headers = ("平销直接大定", "周直接大定", "直接大定", "平销大定", "周大定", "大定量")
    start_headers = ("周开始日期", "平销周开始", "开始日期")
    end_headers = ("周结束日期", "平销周结束", "结束日期")
    steady_start_headers = ("平销开始日期", "平销期开始日期")
    paths = _steady_history_paths()
    saw_steady_sheet = False
    saw_valid_header = False
    for path in paths:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                if "平销" not in sheet.title:
                    continue
                saw_steady_sheet = True
                header_row = 0
                header_index: dict[str, int] = {}
                for row_index in range(1, min(sheet.max_row or 0, 12) + 1):
                    headers = [str(sheet.cell(row=row_index, column=column).value or "").strip() for column in range(1, (sheet.max_column or 0) + 1)]
                    if any(name in headers for name in model_headers) and any(name in headers for name in week_headers) and any(name in headers for name in direct_headers):
                        header_row = row_index
                        header_index = {header: index for index, header in enumerate(headers) if header}
                        break
                if not header_row:
                    LOGGER.warning(
                        "[销量预测字段校验] 文件=%s | Sheet=%s | 未识别车型、周序号或直接大定必需表头 | "
                        "处理=该Sheet全部平销历史拒绝",
                        path.name, sheet.title,
                    )
                    continue
                saw_valid_header = True
                column = lambda aliases: next((header_index[name] for name in aliases if name in header_index), None)
                model_column, week_column, direct_column = column(model_headers), column(week_headers), column(direct_headers)
                start_column, end_column, steady_start_column = column(start_headers), column(end_headers), column(steady_start_headers)
                cancel_column = column(("平销退订", "周退订", "退订"))
                net_column = column(("平销留存大定", "周留存大定", "留存大定", "平销净大定", "周净大定", "净大定"))
                rights_column = column(("权益/价格变化", "权益变化", "价格变化"))
                event_column = column(("活动标记", "促销标记", "活动"))
                supply_column = column(("供应限制", "供应标记"))
                found = False
                for row_number, values in enumerate(
                    sheet.iter_rows(min_row=header_row + 1, values_only=True),
                    start=header_row + 1,
                ):
                    if not any(value not in (None, "") for value in values):
                        continue
                    model = str(values[model_column] or "").strip() if model_column is not None and model_column < len(values) else ""
                    if not model:
                        LOGGER.warning(
                            "[销量预测字段校验] 文件=%s | Sheet=%s | 行=%d | 平销传播名缺失 | "
                            "处理=该行拒绝，无法绑定到任何车型历史",
                            path.name, sheet.title, row_number,
                        )
                        continue
                    generation = _canonical_model(model, model_mapping)
                    master_record = _master_record(model_master, model, generation) or {}
                    key = (str(path.resolve()).casefold(), sheet.title, _model_key(generation or model), _model_key(model))
                    item = grouped.setdefault(key, {
                        "model": model,
                        "generation": generation,
                        "mapped": _model_key(model) in model_mapping,
                        "brand": str(master_record.get("品牌") or _brand(model)),
                        "tier": usable_attribute(master_record.get("产品档位")),
                        "energy": usable_attribute(master_record.get("能源类型")),
                        "node": usable_attribute(master_record.get("发布类型") or master_record.get("发布节点")),
                        "steady_start_date": "",
                        "weeks": [],
                        "source_file": path.name,
                        "source_sheet": sheet.title,
                        "invalid_reasons": [],
                    })
                    raw_week = values[week_column] if week_column is not None and week_column < len(values) else None
                    week_text = str(raw_week or "").strip().upper()
                    week_match = re.fullmatch(r"(?:WK?|第)?\s*(\d+)\s*(?:周)?", week_text)
                    if not week_match:
                        LOGGER.warning(
                            "[销量预测字段校验] 代际=%s | 文件=%s | Sheet=%s | 行=%d | 周序号=%r无效 | "
                            "处理=整条平销历史拒绝",
                            generation, path.name, sheet.title, row_number, raw_week,
                        )
                        item["invalid_reasons"].append(f"行{row_number}周序号无效")
                        continue
                    week = int(week_match.group(1))
                    if week <= 0:
                        LOGGER.warning(
                            "[销量预测字段校验] 代际=%s | 文件=%s | Sheet=%s | 行=%d | 周序号=WK%d（应从WK1起且为正整数）| "
                            "处理=整条平销历史拒绝",
                            generation, path.name, sheet.title, row_number, week,
                        )
                        item["invalid_reasons"].append(f"行{row_number}周序号非正")
                        continue
                    raw_direct = values[direct_column] if direct_column is not None and direct_column < len(values) else None
                    direct = _optional_number(raw_direct)
                    if direct is None or not math.isfinite(direct) or direct < 0:
                        LOGGER.warning(
                            "[销量预测字段校验] 代际=%s | 文件=%s | Sheet=%s | 行=%d | WK%d直接大定=%r缺失、非数字或为负 | "
                            "处理=整条平销历史拒绝",
                            generation, path.name, sheet.title, row_number, week, raw_direct,
                        )
                        item["invalid_reasons"].append(f"WK{week}直接大定异常")
                        continue
                    raw_start = values[start_column] if start_column is not None and start_column < len(values) else None
                    raw_end = values[end_column] if end_column is not None and end_column < len(values) else None
                    raw_steady_start = values[steady_start_column] if steady_start_column is not None and steady_start_column < len(values) else None
                    start_date = _iso(raw_start) if start_column is not None else ""
                    end_date = _iso(raw_end) if end_column is not None else ""
                    explicit_steady_start = _iso(raw_steady_start) if steady_start_column is not None else ""
                    date_errors: list[str] = []
                    if start_column is not None and not start_date:
                        date_errors.append(f"WK{week}周开始日期={raw_start!r}无效")
                    if end_column is not None and not end_date:
                        date_errors.append(f"WK{week}周结束日期={raw_end!r}无效")
                    if steady_start_column is not None and raw_steady_start not in (None, "") and not explicit_steady_start:
                        date_errors.append(f"WK{week}平销开始日期={raw_steady_start!r}无效")
                    if date_errors:
                        LOGGER.warning(
                            "[销量预测字段校验] 代际=%s | 文件=%s | Sheet=%s | 行=%d | %s | "
                            "处理=整条平销历史拒绝",
                            generation, path.name, sheet.title, row_number, "；".join(date_errors),
                        )
                        item["invalid_reasons"].extend(date_errors)
                        continue
                    item["steady_start_date"] = item["steady_start_date"] or explicit_steady_start or (start_date if week == 1 else "")
                    optional_values: dict[str, float | None] = {}
                    optional_invalid = False
                    for field_name, field_column in (("退订", cancel_column), ("留存大定", net_column)):
                        raw_value = values[field_column] if field_column is not None and field_column < len(values) else None
                        number = _optional_number(raw_value)
                        if raw_value not in (None, "") and (number is None or not math.isfinite(number) or number < 0):
                            LOGGER.warning(
                                "[销量预测字段校验] 代际=%s | 文件=%s | Sheet=%s | 行=%d | WK%d%s=%r非数字或为负 | "
                                "处理=整条平销历史拒绝",
                                generation, path.name, sheet.title, row_number, week, field_name, raw_value,
                            )
                            item["invalid_reasons"].append(f"WK{week}{field_name}异常")
                            optional_invalid = True
                        optional_values[field_name] = number
                    if optional_invalid:
                        continue
                    text_value = lambda index: str(values[index] or "").strip() if index is not None and index < len(values) else ""
                    item["weeks"].append({
                        "week": week,
                        "start_date": start_date,
                        "end_date": end_date,
                        "direct": int(round(direct)),
                        "cancel": optional_values["退订"],
                        "net": optional_values["留存大定"],
                        "rights": text_value(rights_column),
                        "event": text_value(event_column),
                        "supply": text_value(supply_column),
                    })
                    found = True
                if found:
                    sources.append(SourceRef(path.name, sheet.title, "平销期按周直接大定历史"))
        finally:
            workbook.close()
    if not paths:
        LOGGER.info(
            "[销量预测数据诊断] 目录=%s | 未提供平销历史文件 | "
            "处理=平销新功能不可用，不影响已有模块生成",
            FORECAST_INPUT_ROOT.name,
        )
    elif saw_steady_sheet and not saw_valid_header:
        LOGGER.warning(
            "[销量预测数据诊断] 目录=%s | 已找到平销Sheet，但缺少车型/周序号/直接大定必需列 | "
            "处理=平销预测按规则停止，网页显示停止原因",
            FORECAST_INPUT_ROOT.name,
        )
    elif not sources:
        LOGGER.info(
            "[销量预测数据诊断] 目录=%s | 未提供可读取的平销周历史 | "
            "处理=平销新功能不可用，不影响已有模块生成",
            FORECAST_INPUT_ROOT.name,
        )
    result = []
    for item in grouped.values():
        item["weeks"].sort(key=lambda row: row["week"])
        model_name = item.get("generation") or item.get("model")
        week_numbers = [row["week"] for row in item["weeks"]]
        duplicates = sorted({week for week in week_numbers if week_numbers.count(week) > 1})
        invalid_reasons: list[str] = list(item.pop("invalid_reasons", []))
        if not week_numbers:
            LOGGER.warning(
                "[销量预测字段校验] 代际=%s | 文件=%s | Sheet=%s | 没有通过校验的平销周行 | "
                "处理=整条平销历史拒绝",
                model_name, item.get("source_file"), item.get("source_sheet"),
            )
            invalid_reasons.append("没有有效周数据")
        elif week_numbers[0] != 1:
            LOGGER.warning(
                "[销量预测字段校验] 代际=%s | 文件=%s | Sheet=%s | 平销历史从WK%d开始（必须从WK1开始）| "
                "处理=整条平销历史拒绝",
                model_name, item.get("source_file"), item.get("source_sheet"), week_numbers[0],
            )
            invalid_reasons.append("未从WK1开始")
        if duplicates:
            LOGGER.warning(
                "[销量预测字段校验] 代际=%s | 平销周序号重复=%s | "
                "处理=整条历史不参与预测，不任意选取重复值",
                model_name, "、".join(f"WK{week}" for week in duplicates),
            )
            invalid_reasons.append("周序号重复")
        for previous, current in zip(item["weeks"], item["weeks"][1:]):
            if current["week"] - previous["week"] > 1:
                LOGGER.warning(
                    "[销量预测字段校验] 代际=%s | 平销周序号从WK%d跳到WK%d | "
                    "处理=整条历史不参与预测，不用末两周比值伪造缺周",
                    model_name, previous["week"], current["week"],
                )
                invalid_reasons.append("周序号不连续")
            if previous["start_date"] and current["start_date"]:
                interval = (_as_date(current["start_date"]) - _as_date(previous["start_date"])).days
                if interval != 7:
                    LOGGER.warning(
                        "[销量预测字段校验] 代际=%s | 平销周WK%d与WK%d开始日期间隔%d天（应相隔7天）| "
                        "处理=整条平销历史拒绝",
                        model_name, previous["week"], current["week"], interval,
                    )
                    invalid_reasons.append("周开始日期不连续")
        for row in item["weeks"]:
            start_date, end_date = _as_date(row.get("start_date")), _as_date(row.get("end_date"))
            if start_date and end_date and (end_date - start_date).days != 6:
                LOGGER.warning(
                    "[销量预测字段校验] 代际=%s | WK%d日期范围=%s~%s（应为7个自然日）| "
                    "处理=整条历史不参与预测",
                    model_name, row["week"], row["start_date"], row["end_date"],
                )
                invalid_reasons.append("周日期范围错误")
            if row.get("cancel") is not None and row.get("net") is not None and abs(
                float(row["direct"]) - float(row["cancel"]) - float(row["net"])
            ) > 1:
                LOGGER.warning(
                    "[销量预测字段校验] 代际=%s | WK%d | 直接大定-退订与留存大定不一致 | "
                    "处理=整条历史不参与预测",
                    model_name, row["week"],
                )
                invalid_reasons.append("留存大定勾稽不一致")
        if invalid_reasons:
            continue
        item["event_id"] = "steady|{}|{}|{}|{}".format(
            _model_key(item.get("generation") or item.get("model")),
            item.get("steady_start_date") or "no-date",
            item.get("source_file") or "no-file",
            item.get("source_sheet") or "no-sheet",
        )
        result.append(item)
    unmapped_steady = [str(item.get("model")) for item in result if not item.get("mapped")]
    if unmapped_steady:
        LOGGER.warning(
            "[映射诊断] 模块=平销预测 | 未映射传播名=%d个 | 明细=%s | "
            "处理=仍作为历史参考并在网页标记未映射，不猜测性绑定订单分析代际",
            len(unmapped_steady), "、".join(unmapped_steady),
        )
    return result, sources


def _iso_week_bounds(value: Any) -> tuple[date, date] | None:
    match = re.fullmatch(r"(\d{2}|\d{4})WK(\d{1,2})", str(value or "").strip().upper())
    if not match:
        return None
    year = int(match.group(1))
    if year < 100:
        year += 2000
    try:
        start = date.fromisocalendar(year, int(match.group(2)), 1)
    except ValueError:
        return None
    return start, start + timedelta(days=6)


def _fill_sparse_dates(rows, fields):
    """Absent export columns inside observed coverage mean zero, not blank cells.

    Never fill before/after the observed interval or overwrite an explicit blank.
    Duplicate dates remain intact so downstream validation can reject them.
    """
    by_date = {row["date"]: row for row in rows if row.get("date")}
    if len(by_date) != len(rows):
        return rows
    observed = sorted(day for day, row in by_date.items()
                      if any(_optional_number(row.get(field)) is not None for field in fields))
    if len(observed) < 2:
        return rows
    current, end = date.fromisoformat(observed[0]), date.fromisoformat(observed[-1])
    while current <= end:
        key = current.isoformat()
        if key not in by_date:
            by_date[key] = {"date": key, **{field: 0 for field in fields}, "absent_as_zero": True}
        current += timedelta(days=1)
    return [by_date[key] for key in sorted(by_date)]


def _read_steady_daily(workbook, generation: str, start: date, today: date):
    result = []
    source_sheets = []
    for sheet in workbook.worksheets:
        if '图表' in sheet.title or grain_from_sheet(sheet.title) != 'day' or not _same_model(sheet_subject(sheet.title), generation):
            continue
        metric = ''
        for row in range(2, sheet.max_row + 1):
            metric = str(sheet.cell(row, 1).value or metric).strip()
            if metric != '交车锁单' or sheet.cell(row, 2).value != '数量' or sheet.cell(row, 3).value != '数量':
                continue
            for col in range(4, sheet.max_column + 1):
                day = _as_date(sheet.cell(1, col).value)
                if day is None or day < start or day > today:
                    continue
                value = _optional_number(sheet.cell(row, col).value)
                valid = value is not None and math.isfinite(value) and value >= 0
                result.append({'date':day.isoformat(), 'lock':int(round(value)) if valid else None, 'complete':day < today})
            source_sheets.append(sheet.title)
            break
    result = _fill_sparse_dates(sorted(result, key=lambda row:row['date']), ("lock",))
    for row in result:
        row["complete"] = row["date"] < today.isoformat()
    return result, source_sheets


def _read_steady_history(
    store,
    stage_windows: dict[str, dict[str, Any]],
    today: date | None = None,
) -> tuple[list[dict[str, Any]], list[SourceRef]]:
    """Read completed flat-sales lock weeks from generation-level lock-mix workbooks."""
    current = today or date.today()
    model_mapping = _read_model_mapping()
    model_master = _read_model_master()
    lock_items = store.find_all("锁单选配比例") if hasattr(store, "find_all") else []
    if not lock_items:
        item = store.find("锁单选配比例")
        lock_items = [item] if item else []
    result: list[dict[str, Any]] = []
    sources: list[SourceRef] = []
    for lock_item in lock_items:
        for sheet in lock_item.workbook.worksheets:
            grain = grain_from_sheet(sheet.title)
            if "图表" in sheet.title or grain not in {"week", "day"}:
                continue
            if grain == 'day' and any(grain_from_sheet(other.title) == 'week' and '图表' not in other.title and _same_model(sheet_subject(other.title), sheet_subject(sheet.title)) for other in lock_item.workbook.worksheets):
                continue
            generation = _canonical_model(sheet_subject(sheet.title), model_mapping)
            if not generation or subject_type(generation) != "generation" or is_aggregate_generation(generation):
                continue
            window = _stage_window(stage_windows, generation)
            launch_end = _as_date((window or {}).get("end_date"))
            if launch_end is None:
                LOGGER.warning(
                    "[销量预测数据诊断] 代际=%s | 文件=%s | Sheet=%s | 首销截止日期缺失 | "
                    "处理=无法切分平销期，该代际锁单周历史不参与平销预测",
                    generation, lock_item.path.name, sheet.title,
                )
                continue
            steady_start = launch_end + timedelta(days=1)
            daily, daily_sheets = _read_steady_daily(lock_item.workbook, generation, steady_start, current)
            parsed = parse_metric_sheet(sheet)
            raw_locks: dict[str, Any] = {}
            current_metric = ""
            for row_index in range(2, sheet.max_row + 1):
                metric = str(sheet.cell(row_index, 1).value or "").strip()
                if metric:
                    current_metric = metric
                stat = str(sheet.cell(row_index, 2).value or "").strip()
                category = str(sheet.cell(row_index, 3).value or "").strip()
                if current_metric != "交车锁单" or stat != "数量" or category != "数量":
                    continue
                for column_index in range(4, sheet.max_column + 1):
                    period = display_period(
                        sheet.cell(1, column_index).value,
                        sheet.cell(1, column_index).number_format,
                    )
                    if period:
                        raw_locks[period] = sheet.cell(row_index, column_index).value
                break
            weeks: list[dict[str, Any]] = []
            for period, payload in parsed.items():
                bounds = _iso_week_bounds(period)
                if not bounds:
                    continue
                week_start, week_end = bounds
                if week_start < steady_start or week_end >= current:
                    continue
                raw_lock = raw_locks.get(period)
                lock_value = _optional_number(raw_lock)
                if lock_value is None or not math.isfinite(lock_value) or lock_value < 0:
                    LOGGER.warning(
                        "[销量预测字段校验] 代际=%s | 文件=%s | Sheet=%s | 周期=%s | "
                        "交车锁单=%r缺失、非数字或为负 | 处理=该周不进入平销锁单历史",
                        generation, lock_item.path.name, sheet.title, period, raw_lock,
                    )
                    continue
                weeks.append({
                    "period": str(period),
                    "start_date": week_start.isoformat(),
                    "end_date": week_end.isoformat(),
                    "lock": int(round(lock_value)),
                })
            weeks.sort(key=lambda row: row["start_date"])
            if any(
                (_as_date(current_row["start_date"]) - _as_date(previous["start_date"])).days != 7
                for previous, current_row in zip(weeks, weeks[1:])
            ):
                LOGGER.warning(
                    "[销量预测字段校验] 代际=%s | 文件=%s | Sheet=%s | 平销锁单周不连续 | "
                    "处理=周历史不参与环比，保留可用的平销真实日",
                    generation, lock_item.path.name, sheet.title,
                )
                weeks = []
            if not weeks and not daily:
                continue
            for index, row in enumerate(weeks, start=1):
                row["week"] = index
            master_record = _master_record(model_master, generation) or {}
            result.append({
                "model": generation,
                "generation": generation,
                "mapped": True,
                "brand": str(master_record.get("品牌") or _brand(generation)),
                "tier": usable_attribute(master_record.get("产品档位")),
                "energy": usable_attribute(master_record.get("能源类型")),
                "node": usable_attribute(master_record.get("发布类型") or master_record.get("发布节点")),
                "steady_start_date": steady_start.isoformat(),
                "weeks": weeks,
                "daily": daily,
                "daily_source_sheets": daily_sheets,
                "source_file": lock_item.path.name,
                "source_sheet": sheet.title,
                "event_id": "steady-lock|{}|{}|{}".format(
                    _model_key(generation), lock_item.path.name, sheet.title,
                ),
            })
            sources.append(SourceRef(lock_item.path.name, sheet.title, "平销期按周交车锁单历史"))
            sources.extend(SourceRef(lock_item.path.name, title, '平销期按天交车锁单真实进度') for title in daily_sheets)
    if not lock_items:
        LOGGER.warning(
            "[销量预测数据诊断] 未找到《锁单选配比例分析》 | "
            "处理=平销锁单预测不可用，不再读取独立平销历史文件",
        )
    return result, sources


def _progress_rows(workbook, sheet_name: str) -> dict[str, list[float | None]]:
    if sheet_name not in workbook.sheetnames:
        return {}
    sheet = workbook[sheet_name]
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    day_columns = [(index, header) for index, header in enumerate(headers) if re.fullmatch(r"D\d+", header)]
    result: dict[str, list[float | None]] = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not values or not values[0]:
            continue
        result[str(values[0])] = [_optional_number(values[index]) for index, _ in day_columns]
    return result


def _sheet_records(workbook, sheet_name: str) -> dict[str, dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        return {}
    sheet = workbook[sheet_name]
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    records: dict[str, dict[str, Any]] = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, values))
        model = str(record.get("传播名") or record.get("车型") or "").strip()
        if model:
            records[_model_key(model)] = record
    return records


def _history_record(history: list[dict[str, Any]], model: str) -> dict[str, Any] | None:
    requested = _model_key(model)
    exact = [item for item in history if _model_key(item.get("model")) == requested]
    if len(exact) == 1:
        return exact[0]
    mapped = [
        item for item in history
        if _same_model(item.get("model"), model) or _same_model(item.get("generation"), model)
    ]
    if not mapped:
        return None
    if len(mapped) > 1:
        mapped.sort(key=lambda item: (_iso(item.get("launch_date")), str(item.get("model") or "")), reverse=True)
        LOGGER.warning(
            "[映射诊断] 预测对象=%s | 匹配到多个历史传播名=%s | 处理=统一采用发布日最新的%s，避免跨记录拼接",
            model,
            "、".join(str(item.get("model") or "") for item in mapped),
            mapped[0].get("model"),
        )
    return mapped[0]


def _brand(model: str) -> str:
    return next((brand for brand in ("问界", "智界", "享界", "尊界", "尚界") if brand in model), "")


def _history_item(record: dict[str, Any], processed: bool) -> dict[str, Any]:
    model = str(record.get("传播名") or record.get("车型") or "")
    gross_key = "总大定" if processed else "首销期大定"
    direct_key = "总直接大定" if processed else "直接大定量"
    cancel_key = "小订后退订"
    item = {
        "model": model,
        "generation": str(record.get("代际名") or ""),
        "mapping_status": str(record.get("映射状态") or ("已映射" if record.get("代际名") else "未映射·暂按传播名")),
        "mapped": not str(record.get("映射状态") or "").startswith("未映射") and bool(str(record.get("代际名") or "").strip()),
        "brand": str(record.get("品牌") or _brand(model)),
        "tier": usable_attribute(record.get("产品档位")),
        "energy": usable_attribute(record.get("能源类型")),
        "launch_date": _iso(record.get("发布日") or record.get("开始大定日期")),
        "node": usable_attribute(record.get("发布类型") or record.get("发布节点")),
        "launch_weekday": str(record.get("发布星期") or _weekday(record.get("发布日") or record.get("开始大定日期"))),
        "launch_period": usable_attribute(record.get("发布时段")),
        "end_date": _iso(record.get("首销截止") or record.get("小转大结束日期")),
        "days": int(_number(record.get("首销期天数") or record.get("首销天数"), 0)),
        "gross": int(_number(record.get(gross_key) or record.get("大定量"))),
        "net": int(_number(record.get("首销期留存大定") or record.get("首销期净大定"))),
        "net_rate": _number(record.get("留存大定率") or record.get("净大定率")),
        "lock": int(_number(record.get("首销期锁单"))),
        "lock_rate": _number(record.get("大定到锁单率") or record.get("锁单率")),
        "small": int(_number(record.get("总小订"))),
        "small_to_big": int(_number(record.get("小订转大") or record.get("小订转大定量") or record.get("小订转大定"))),
        "conversion": _number(record.get("小订转化率")),
        "direct": int(_number(record.get(direct_key) or record.get("直接大定量"))),
        "direct_share": _number(record.get("直接大定占比")),
        "cancel": int(_number(record.get(cancel_key) or record.get("小订后退定"))),
        "cancel_rate": _number(record.get("退订率") or record.get("小订后退定占比")),
        "d1_small": int(_number(record.get("D1小转大") or record.get("首日小转大"))),
        "d1_small_completion": _number(record.get("D1小转大/总小转大")),
        "d2_small": int(_number(record.get("D2小转大"))),
        "d2_small_completion": _number(record.get("D2小转大/总小转大")),
        "d12_small": int(_number(record.get("D1+D2小转大"))),
        "d12_small_completion": _number(record.get("D1+D2小转大/总小转大")),
        "small_d2_d1": _number(record.get("D2小转大/D1小转大")),
        "small_d1_d12": _number(record.get("D1小转大/D1+D2小转大")),
        "d1_direct": int(_number(record.get("D1直接大") or record.get("首日直接大定"))),
        "d1_direct_completion": _number(record.get("D1直接大/总直接大")),
        "d2_direct_completion": _number(record.get("D2直接大/总直接大")),
        "d1_gross": int(_number(record.get("D1大定"))),
        "d1_small_share": _number(record.get("D1小转大/D1大定")),
        "d1_direct_share": _number(record.get("D1直接大/D1大定") or record.get("首日直接大定占比")),
        "d12_direct": int(_number(record.get("D1+D2直接大"))),
        "d12_direct_completion": _number(record.get("D1+D2直接大/总直接大")),
        "direct_d2_d1": _number(record.get("D2直接大/D1直接大")),
        "direct_d1_d12": _number(record.get("D1直接大/D1+D2直接大")),
        "d2_gross": int(_number(record.get("D2大定"))),
        "d2_direct": int(_number(record.get("D2直接大"))),
        "d2_small_share": _number(record.get("D2小转大/D2大定")),
        "d2_direct_share": _number(record.get("D2直接大/D2大定")),
        "d12_gross": int(_number(record.get("D1+D2大定"))),
        "d12_small_share": _number(record.get("D1+D2小转大/D1+D2大定")),
        "d12_direct_share": _number(record.get("D1+D2直接大/D1+D2大定")),
        "d1_cancel": int(_number(record.get("D1退订"))),
        "d1_cancel_rate": _number(record.get("D1退订率")),
        "d2_cancel": int(_number(record.get("D2退订"))),
        "d2_cancel_rate": _number(record.get("D2退订率")),
        "d12_cancel": int(_number(record.get("D1+D2退订"))),
        "d12_cancel_rate": _number(record.get("D1+D2退订率")),
        "cancel_d2_d1": _number(record.get("D2退订/D1退订")),
        "cancel_d1_d12": _number(record.get("D1退订/D1+D2退订")),
        "field_completeness": _number(record.get("字段完整度")),
        "consistency": _number(record.get("口径一致性")),
        "quality_status": str(record.get("质量状态") or "未校验"),
        "quality_issues": str(record.get("质量问题") or record.get("结构异常说明") or ""),
        "d1_valid": str(record.get("D1口径状态") or "通过") == "通过",
        "d2_valid": str(record.get("D2口径状态") or "通过") == "通过",
        "d12_valid": str(record.get("D1+D2口径状态") or "通过") == "通过",
        "data_type": "二次处理基准" if processed else str(record.get("数据性质") or "历史原始数据"),
    }
    if not item["d1_valid"]:
        for field in ("d1_small_completion", "d1_direct_completion", "d1_small_share", "d1_direct_share"):
            item[field] = 0
    if not item["d2_valid"]:
        for field in ("d2_small_completion", "d2_direct_completion", "d2_small_share", "d2_direct_share", "small_d2_d1", "direct_d2_d1"):
            item[field] = 0
    if not item["d12_valid"]:
        for field in ("d12_small_completion", "d12_direct_completion", "d12_small_share", "d12_direct_share", "small_d1_d12", "direct_d1_d12"):
            item[field] = 0
    if item["d1_valid"]:
        item["d1_small_share"] = item["d1_small"] / item["d1_gross"] if item["d1_gross"] else 0
    item["quality_issues"] = "；".join(filter(None, (
        str(record.get("质量问题") or "").strip(), str(record.get("结构异常说明") or "").strip()
    )))
    return item


def _read_history() -> tuple[Path | None, list[dict[str, Any]]]:
    path = input_path(next((candidate for candidate in HISTORY_CANDIDATES if candidate.exists()), None))
    if path is None:
        return None, []
    workbook = open_input(path, read_only=False, data_only=True)
    try:
        processed = "预测基准总表" in workbook.sheetnames
        base_sheet = "预测基准总表" if processed else "首销预测基准" if "首销预测基准" in workbook.sheetnames else "传播名汇总" if "传播名汇总" in workbook.sheetnames else "车型汇总" if "车型汇总" in workbook.sheetnames else ""
        if not base_sheet:
            return path, []
        progress_maps = {
            "small_progress": _progress_rows(workbook, "小转大累计完成度" if processed else "小转大进度"),
            "direct_progress": _progress_rows(workbook, "直接大定累计完成度" if processed else "直接大定进度"),
            "direct_share_progress": _progress_rows(workbook, "累计直接大定占比" if processed else "大定进度"),
            "cancel_progress": _progress_rows(workbook, "累计退订率" if processed else "退订进度"),
            "gross_progress": _progress_rows(workbook, "累计大定完成度"),
            "daily_orders": _progress_rows(workbook, "总大定当日数量"),
            "daily_small": _progress_rows(workbook, "小转大当日数量"),
            "daily_direct": _progress_rows(workbook, "直接大定当日数量"),
        }
        d12_records = _sheet_records(workbook, "D1_D2预测指标") if processed else {}
        sheet = workbook[base_sheet]
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        result = []
        model_mapping = _read_model_mapping()
        model_master = _read_model_master()
        for row in sheet.iter_rows(min_row=2, values_only=True):
            record = dict(zip(headers, row))
            model = str(record.get("传播名") or record.get("车型") or "").strip()
            if not model:
                continue
            d12_record = d12_records.get(_model_key(model)) or next(
                (candidate for name, candidate in d12_records.items() if _same_model(name, model)), None
            )
            if d12_record:
                record.update(d12_record)
            item = _history_item(record, processed)
            if not item.get("generation"):
                item["generation"] = model_mapping.get(_model_key(model), "")
            master_record = _master_record(model_master, model, item.get("generation"))
            if master_record:
                item["brand"] = str(master_record.get("品牌") or item.get("brand") or "")
                item["tier"] = str(master_record.get("产品档位") or "未维护")
                item["energy"] = str(master_record.get("能源类型") or "未维护")
                item["node"] = str(master_record.get("发布类型") or master_record.get("发布节点") or "未维护")
                item["launch_period"] = str(master_record.get("发布时段") or "未维护")
            for field, mapping in progress_maps.items():
                # Different historical events/editions may map to one generation.
                # Preserve their own named curves before considering an alias.
                values = next((curve for name, curve in mapping.items() if _model_key(name) == _model_key(model)), None)
                if values is None:
                    matches = [curve for name, curve in mapping.items() if _same_model(name, model)]
                    values = matches[0] if len(matches) == 1 else []
                item[field] = values
            if not processed and item.get("small_progress"):
                terminal = _number(item.get("conversion"))
                if 0 < terminal <= 1:
                    item["small_progress"] = [
                        min(max(value / terminal, 0), 1) if value else 0
                        for value in item["small_progress"]
                    ]
                else:
                    LOGGER.warning(
                        "[销量预测字段校验] %s | 非二次处理历史文件缺少有效小订转化率，"
                        "小转大完成度曲线已标记为不可用，不再用0.1%%伪完成度强制归一化",
                        model,
                    )
                    item["small_progress"] = []
            result.append(item)
        return path, result
    finally:
        workbook.close()


def _sheet_rows(sheet) -> dict[str, list[Any]]:
    return {
        str(sheet.cell(row=row, column=1).value or "").strip(): [sheet.cell(row=row, column=column).value for column in range(2, sheet.max_column + 1)]
        for row in range(1, sheet.max_row + 1)
        if sheet.cell(row=row, column=1).value
    }


def _read_actual_profiles(store) -> tuple[list[dict[str, Any]], list[SourceRef]]:
    profiles: dict[str, dict[str, Any]] = {}
    sources: list[SourceRef] = []
    model_mapping = _read_model_mapping()
    launch_item = store.find("首销期订单节奏")
    if launch_item:
        for sheet in launch_item.workbook.worksheets:
            model = _canonical_model(sheet_subject(sheet.title), model_mapping)
            grain = grain_from_sheet(sheet.title)
            if not model or grain not in ("day", "hour"):
                continue
            rows = _sheet_rows(sheet)
            profile = profiles.setdefault(model, {"model": model, "days": [], "hourly_days": []})
            if grain == "day":
                dates = rows.get("时间", [])
                gross = rows.get("当日大定数量", [])
                net = rows.get("当日留存大定数量", rows.get("当日净大定数量", []))
                small = rows.get("当日小订转大定数量", [])
                direct = rows.get("当日直接大定数量", [])
                locks = rows.get("当日交车锁单数量", [])
                actual_days = []
                for index, value in enumerate(gross):
                    if index >= len(dates) or not dates[index] or not isinstance(value, (int, float)):
                        continue
                    if not math.isfinite(float(value)) or float(value) < 0:
                        LOGGER.warning(
                            "[销量预测字段校验] 代际=%s | Sheet=%s | 行=%d | "
                            "当日大定=%r | 处理=负数或非有限数已拒绝，该日不进入真实进度",
                            model,
                            sheet.title,
                            index + 1,
                            value,
                        )
                        continue
                    for label, values in (
                        ("当日留存大定数量", net),
                        ("当日小订转大定数量", small),
                        ("当日直接大定数量", direct),
                        ("当日交车锁单数量", locks),
                    ):
                        component = values[index] if index < len(values) else None
                        if isinstance(component, (int, float)) and (
                            not math.isfinite(float(component)) or float(component) < 0
                        ):
                            LOGGER.warning(
                                "[销量预测字段校验] 代际=%s | Sheet=%s | 行=%d | %s=%r | "
                                "处理=负数或非有限数已置为缺失，后续按当前阶段逐字段回退",
                                model,
                                sheet.title,
                                index + 1,
                                label,
                                component,
                            )
                    actual_days.append({
                        "day": f"D{len(actual_days) + 1}",
                        "date": _iso(dates[index]),
                        "gross": int(round(_number(value))),
                        "net": _optional_count(net, index),
                        "small_to_big": _optional_count(small, index),
                        "direct": _optional_count(direct, index),
                        "lock": _optional_count(locks, index),
                    })
                if not actual_days and (dates or gross):
                    d1_date = dates[0] if dates else None
                    d1_gross = gross[0] if gross else None
                    LOGGER.warning(
                        "[销量预测D1排查] by天未构建真实日：原Sheet=%s，标准代际=%s，"
                        "时间指标=%s，大定指标=%s，D1时间=%r(%s)，D1大定=%r(%s)，"
                        "判定要求=时间非空且大定为数字",
                        sheet.title,
                        model,
                        "有" if "时间" in rows else "无",
                        "有" if "当日大定数量" in rows else "无",
                        d1_date,
                        type(d1_date).__name__,
                        d1_gross,
                        type(d1_gross).__name__,
                    )
                profile["days"] = actual_days
                profile["launch_date"] = actual_days[0]["date"] if actual_days else ""
                cumulative_small = rows.get("累计小订转大定数量", [])
                cumulative_rate = rows.get("累计小订转化率", [])
                explicit_small = next((_number(value) for value in reversed(rows.get("总小订", [])) if _number(value) > 0), 0)
                inferred_small, inferred_valid, inferred_reason = _stable_implied_total(cumulative_small, cumulative_rate)
                cumulative_latest = max((_number(value) for value in cumulative_small), default=0)
                explicit_valid = explicit_small > 0 and explicit_small >= cumulative_latest
                if explicit_small > 0 and not explicit_valid:
                    LOGGER.warning(
                        "[销量预测估算] 代际=%s | Sheet=%s | 显式总小订%d小于累计小转大%d，判定不可用 | "
                        "处理=总小订按缺失处理，不用反推值强行兜底",
                        model, sheet.title, int(round(explicit_small)), int(round(cumulative_latest)),
                    )
                profile["launch_total_small"] = int(round(explicit_small or inferred_small))
                profile["launch_total_small_valid"] = bool(explicit_valid or (not explicit_small and inferred_valid))
                profile["launch_total_small_estimated"] = bool(not explicit_small and inferred_valid)
                profile["launch_total_small_reason"] = "首销节奏明确总小订" if explicit_valid else inferred_reason
                profile["total_small"] = profile["launch_total_small"] if profile["launch_total_small_valid"] else 0
                profile["day_source"] = SourceRef(launch_item.path.name, sheet.title, "已发生首销日真实数据").to_dict()
                sources.append(SourceRef(launch_item.path.name, sheet.title, "已发生首销日真实数据"))
            elif grain == "hour":
                timestamps = rows.get("时间", [])
                gross = rows.get("当日大定数量", [])
                small = rows.get("当日小订转大定数量", [])
                direct = rows.get("当日直接大定数量", [])
                locks = rows.get("当日交车锁单数量", [])
                grouped: dict[str, dict[str, Any]] = {}
                for index, timestamp in enumerate(timestamps):
                    parsed_timestamp = _as_datetime(timestamp)
                    if parsed_timestamp is None or index >= len(gross) or not isinstance(gross[index], (int, float)):
                        continue
                    if not math.isfinite(float(gross[index])) or float(gross[index]) < 0:
                        LOGGER.warning(
                            "[销量预测字段校验] 代际=%s | Sheet=%s | 时间=%s | "
                            "分时大定=%r | 处理=负数或非有限数已拒绝",
                            model,
                            sheet.title,
                            timestamp,
                            gross[index],
                        )
                        continue
                    day = parsed_timestamp.date().isoformat()
                    bucket = grouped.setdefault(day, {"date": day, "last_hour": -1, "gross": 0, "small_to_big": 0, "direct": 0, "lock": 0, "hours": []})
                    bucket["last_hour"] = max(bucket["last_hour"], parsed_timestamp.hour)
                    bucket["gross"] += int(round(_number(gross[index])))
                    for field, values in (("small_to_big", small), ("direct", direct), ("lock", locks)):
                        value = _optional_count(values, index)
                        bucket[field] = bucket[field] + value if bucket[field] is not None and value is not None else None
                    bucket["hours"].append({"hour": parsed_timestamp.hour, "gross": int(round(_number(gross[index])))})
                if not grouped and (timestamps or gross):
                    first_time = timestamps[0] if timestamps else None
                    first_gross = gross[0] if gross else None
                    LOGGER.warning(
                        "[销量预测D1排查] by时未构建分时进度：原Sheet=%s，标准代际=%s，"
                        "时间指标=%s，大定指标=%s，首列时间=%r(%s)，首列大定=%r(%s)，"
                        "判定要求=时间可解析且大定为数字",
                        sheet.title,
                        model,
                        "有" if "时间" in rows else "无",
                        "有" if "当日大定数量" in rows else "无",
                        first_time,
                        type(first_time).__name__,
                        first_gross,
                        type(first_gross).__name__,
                    )
                profile["hourly_days"] = list(grouped.values())
                profile["hour_source"] = SourceRef(launch_item.path.name, sheet.title, "首销分时真实进度").to_dict()

    small_mix_item = store.find("小订选配比例")
    if small_mix_item:
        for sheet in small_mix_item.workbook.worksheets:
            if "图表" in sheet.title or grain_from_sheet(sheet.title) != "day":
                continue
            model = _canonical_model(sheet_subject(sheet.title), model_mapping)
            if not model or subject_type(model) != "generation" or is_aggregate_generation(model):
                continue
            quantity_row = next((
                row for row in range(2, sheet.max_row + 1)
                if str(sheet.cell(row, 1).value or "").strip() == "小订"
                and str(sheet.cell(row, 2).value or "").strip() == "数量"
                and str(sheet.cell(row, 3).value or "").strip() == "数量"
            ), None)
            if quantity_row is None:
                continue
            daily_rows = []
            for column in range(4, sheet.max_column + 1):
                date_key = _iso(sheet.cell(1, column).value)
                orders = _optional_number(sheet.cell(quantity_row, column).value)
                if not date_key:
                    continue
                if orders is None:
                    daily_rows.append({"date": date_key, "orders": None})
                    continue
                if not math.isfinite(orders) or orders < 0:
                    LOGGER.warning(
                        "[销量预测字段校验] 代际=%s | 文件=%s | Sheet=%s | 日期=%s | "
                        "小订=%r | 处理=负数或非有限数已拒绝",
                        model, small_mix_item.path.name, sheet.title, date_key,
                        sheet.cell(quantity_row, column).value,
                    )
                    daily_rows.append({"date": date_key, "orders": None})
                    continue
                daily_rows.append({"date": date_key, "orders": int(round(orders))})
            if not daily_rows:
                continue
            daily_rows = _fill_sparse_dates(daily_rows, ("orders",))
            profile_name = next((name for name in profiles if _same_model(name, model)), model)
            profile = profiles.setdefault(profile_name, {"model": profile_name, "days": [], "hourly_days": [], "small_hourly_days": []})
            source = SourceRef(small_mix_item.path.name, sheet.title, "小订选配比例by天真实小订")
            by_date = {row["date"]: row for row in profile.get("small_daily_days", [])}
            by_date.update({row["date"]: row for row in daily_rows})
            profile["small_daily_days"] = [by_date[key] for key in sorted(by_date)]
            profile.setdefault("small_daily_sources", {}).update({row["date"]: source.to_dict() for row in daily_rows})
            sources.append(source)

    cancel_item = store.find("小订退订分析")
    if cancel_item:
        for sheet in cancel_item.workbook.worksheets:
            if "小订分时退订" in sheet.title:
                model = _canonical_model(sheet.title.split("_", 1)[0], model_mapping)
                profile_name = next((name for name in profiles if _same_model(name, model)), model)
                profile = profiles.setdefault(profile_name, {"model": profile_name, "days": [], "hourly_days": [], "small_hourly_days": []})
                aliases = {
                    "date": {"小订日期", "下单日期", "日期"},
                    "hour": {"小订小时", "小时", "时段"},
                    "orders": {"小订数", "小订数量"},
                }
                resolved: tuple[int, dict[str, int]] | None = None
                for header_row in range(1, min(sheet.max_row, 10) + 1):
                    headers = [str(sheet.cell(row=header_row, column=column).value or "").strip() for column in range(1, sheet.max_column + 1)]
                    columns = {
                        key: next((index for index, header in enumerate(headers) if header in names), -1)
                        for key, names in aliases.items()
                    }
                    if all(index >= 0 for index in columns.values()):
                        resolved = header_row, columns
                        break
                grouped: dict[str, dict[str, Any]] = {}
                invalid_small_dates: set[str] = set()
                seen_small_hours: set[tuple[str, int]] = set()
                if resolved:
                    header_row, columns = resolved
                    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                        date_key = _iso(values[columns["date"]] if columns["date"] < len(values) else None)
                        raw_hour = values[columns["hour"]] if columns["hour"] < len(values) else None
                        orders = _optional_number(values[columns["orders"]] if columns["orders"] < len(values) else None)
                        if isinstance(raw_hour, (datetime,)):
                            hour = raw_hour.hour
                        else:
                            hour_match = re.search(r"\d{1,2}", str(raw_hour if raw_hour is not None else ""))
                            hour = int(hour_match.group()) if hour_match else -1
                        if not date_key:
                            continue
                        if not 0 <= hour <= 23 or orders is None or not math.isfinite(orders) or orders < 0 or (date_key, hour) in seen_small_hours:
                            invalid_small_dates.add(date_key)
                            continue
                        seen_small_hours.add((date_key, hour))
                        bucket = grouped.setdefault(date_key, {"date": date_key, "last_hour": -1, "orders": 0, "hours": []})
                        bucket["last_hour"] = max(bucket["last_hour"], hour)
                        bucket["orders"] += int(round(orders))
                        bucket["hours"].append({"hour": hour, "orders": int(round(orders))})
                profile["small_hourly_days"] = list(grouped.values())
                cancel_daily_rows = [
                    {"date": day, "orders": bucket["orders"] if day not in invalid_small_dates else None}
                    for day, bucket in sorted(grouped.items())
                ]
                source = SourceRef(cancel_item.path.name, sheet.title, "小订D1分时真实进度")
                by_date = {row["date"]: row for row in profile.get("small_daily_days", [])}
                source_by_date = profile.setdefault("small_daily_sources", {})
                for row in cancel_daily_rows:
                    if row["date"] not in by_date or by_date[row["date"]].get("orders") is None:
                        by_date[row["date"]] = row
                        source_by_date[row["date"]] = source.to_dict()
                profile["small_daily_days"] = [by_date[key] for key in sorted(by_date)]
                profile["small_hour_source"] = source.to_dict()
                if grouped:
                    sources.append(source)
                continue
            if "日度退订" not in sheet.title:
                continue
            model = _canonical_model(sheet.title.split("_", 1)[0], model_mapping)
            profile_name = next((name for name in profiles if _same_model(name, model)), model)
            profile = profiles.setdefault(profile_name, {"model": profile_name, "days": [], "hourly_days": []})
            headers = [str(sheet.cell(row=2, column=column).value or "").strip() for column in range(1, sheet.max_column + 1)]
            header_index = {header: index for index, header in enumerate(headers)}
            cancel_by_date: dict[str, int] = {}
            cancel_rows = []
            cumulative_column = next((header_index[name] for name in ("累计小订退", "累计退订数", "累计总退订") if name in header_index), 2)
            rate_column = next((header_index[name] for name in ("整体 - 小订退 %", "小订后退订比例", "小订退订率") if name in header_index), None)
            for values in sheet.iter_rows(min_row=3, values_only=True):
                if not values or not values[0]:
                    continue
                date_key = _iso(values[0])
                if not date_key:
                    continue
                cumulative = _optional_count(values, cumulative_column)
                rate = _optional_number(values[rate_column]) if rate_column is not None and rate_column < len(values) else None
                cancel_by_date[date_key] = cumulative
                cancel_rows.append({"date": date_key, "cancel": cumulative, "cancel_rate": rate})
            latest_cancel = None
            for row in profile.get("days", []):
                if row["date"] in cancel_by_date:
                    latest_cancel = cancel_by_date[row["date"]]
                row["cancel"] = latest_cancel
            if cancel_rows:
                latest = cancel_rows[-1]
                inferred_small, inferred_valid, inferred_reason = _stable_implied_total(
                    [row["cancel"] for row in cancel_rows], [row["cancel_rate"] for row in cancel_rows]
                )
                profile["cancel_total_small"] = inferred_small
                profile["cancel_total_small_valid"] = inferred_valid
                profile["cancel_total_small_reason"] = inferred_reason
                profile["cancel_days"] = cancel_rows
                profile["cancel_latest_date"] = latest["date"]
            profile["cancel_source"] = SourceRef(cancel_item.path.name, sheet.title, "逐日真实退订进度").to_dict()
            sources.append(SourceRef(cancel_item.path.name, sheet.title, "逐日真实退订进度"))

    lock_item = store.find("锁单选配比例")
    if lock_item:
        for sheet in lock_item.workbook.worksheets:
            model = _canonical_model(sheet_subject(sheet.title), model_mapping)
            if not model or subject_type(model) != "generation" or is_aggregate_generation(model):
                continue
            profile_name = next((name for name in profiles if _same_model(name, model)), model)
            if profile_name in profiles:
                continue
            profiles[profile_name] = {"model": profile_name, "days": [], "hourly_days": []}
            sources.append(SourceRef(lock_item.path.name, sheet.title, "预测候选代际"))
    return list(profiles.values()), sources


def _history_has_data(item: dict[str, Any] | None) -> bool:
    if not item:
        return False
    return any(_number(item.get(field)) > 0 for field in ("small", "small_to_big", "gross", "direct", "lock")) or any(
        any(_number(value) != 0 for value in item.get(field, []))
        for field in ("daily_orders", "daily_small", "daily_direct", "gross_progress", "cancel_progress")
    )


def _profile_from_history(
    model: str,
    item: dict[str, Any],
    source: SourceRef,
    target_launch_date: Any = None,
) -> dict[str, Any]:
    # D1-Dn in the secondary workbook are lifecycle positions. Once the row is
    # mapped to a current generation, anchor those positions to that target's
    # confirmed launch date so the browser does not discard valid D1 values as
    # belonging to another or missing absolute date.
    launch = _as_date(target_launch_date) or _as_date(item.get("launch_date"))
    daily_orders = list(item.get("daily_orders") or [])
    daily_small = list(item.get("daily_small") or [])
    daily_direct = list(item.get("daily_direct") or [])
    cancel_progress = list(item.get("cancel_progress") or [])
    span = min(max(len(daily_orders), len(daily_small), len(daily_direct), 0), max(int(_number(item.get("days"), 35)), 1))
    rows = []
    for index in range(span):
        gross_value = _optional_number(daily_orders[index] if index < len(daily_orders) else None)
        small_value = _optional_number(daily_small[index] if index < len(daily_small) else None)
        direct_value = _optional_number(daily_direct[index] if index < len(daily_direct) else None)
        if direct_value is None and gross_value is not None and small_value is not None:
            direct_value = max(gross_value - small_value, 0)
        cancel_rate = _optional_number(cancel_progress[index] if index < len(cancel_progress) else None)
        rows.append({
            "day": f"D{index + 1}",
            "date": (launch + timedelta(days=index)).isoformat() if launch else "",
            "gross": int(round(gross_value)) if gross_value is not None else None,
            "net": None,
            "small_to_big": int(round(small_value)) if small_value is not None else None,
            "direct": int(round(direct_value)) if direct_value is not None else None,
            "lock": None,
            "cancel": int(round(_number(item.get("small")) * cancel_rate)) if cancel_rate is not None else None,
        })
    return {
        "model": model,
        "launch_date": _iso(item.get("launch_date")),
        "days": rows,
        "hourly_days": [],
        "total_small": int(round(_number(item.get("small")))),
        "day_source": source.to_dict(),
        "selected_source": "history",
        "selected_source_label": "小订及首销数据整理",
    }


SOURCE_LABELS = {
    "history": "小订及首销数据整理",
    "launch": "首销期订单节奏",
    "cancel": "小订退订分析",
    "small_mix": "小订选配比例分析",
    "missing": "数据缺失",
    "mixed": "多来源逐日回退",
}


def _field_value(candidates: dict[str, dict[str, Any]], priority: tuple[str, ...], field: str, validator) -> tuple[Any, str]:
    for source in priority:
        value = candidates.get(source, {}).get(field)
        if validator(value):
            return value, source
    return None, "missing"


def _valid_daily_value(field: str, value: Any, candidate: dict[str, Any]) -> bool:
    if not isinstance(value, (int, float)) or not math.isfinite(float(value)) or value < 0:
        return False
    gross = candidate.get("gross")
    if field in {"small_to_big", "direct", "net", "lock"} and isinstance(gross, (int, float)) and gross >= 0:
        return value <= gross * 1.02 + 5
    return True


def _candidate_for_lifecycle_row(
    base: dict[str, Any],
    index: int,
    by_date: dict[str, dict[str, Any]],
    by_day: dict[str, dict[str, Any]],
    rows: list[dict[str, Any]],
) -> dict[str, Any]:
    """Match absolute dates first; only use Dn/index when a side has no valid date."""

    base_date = _iso(base.get("date"))
    if base_date and base_date in by_date:
        return by_date[base_date]
    lifecycle = str(base.get("day") or f"D{index + 1}").upper()
    by_lifecycle = by_day.get(lifecycle)
    if by_lifecycle and (not base_date or not _iso(by_lifecycle.get("date"))):
        return by_lifecycle
    if not base_date and index < len(rows) and not _iso(rows[index].get("date")):
        return rows[index]
    return {}


def _merge_day_fields(
    candidates: dict[str, dict[str, Any]],
    priority: tuple[str, ...],
    *,
    log_issues: bool = False,
    diagnostic_context: str = "",
) -> tuple[list[dict[str, Any]], dict[str, str]]:
    """Use one date spine and fill fields without crossing different absolute dates."""
    base_rows, base_source = None, "missing"
    for source in priority:
        candidate_rows = list(candidates.get(source, {}).get("days") or [])
        if any(any(isinstance(row.get(field), (int, float)) for field in ("gross", "small_to_big", "direct", "lock")) for row in candidate_rows):
            base_rows, base_source = candidate_rows, source
            break
    if not base_rows:
        return [], {}
    lookups: dict[str, tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]], list[dict[str, Any]]]] = {}
    for source in priority:
        rows = list(candidates.get(source, {}).get("days") or [])
        by_date = {_iso(row.get("date")): row for row in rows if _iso(row.get("date"))}
        by_day = {str(row.get("day") or f"D{index + 1}").upper(): row for index, row in enumerate(rows)}
        lookups[source] = (by_date, by_day, rows)
    fields = ("gross", "net", "small_to_big", "direct", "lock", "cancel")
    source_sets: dict[str, set[str]] = {field: set() for field in fields}
    merged = []
    structure_issues: list[str] = []
    for index, base in enumerate(base_rows):
        row = {**base}
        # 基准行可能来自某个优先来源；先清空所有可回退数值，再逐字段验证取值。
        # 否则基准行中的负数/非法值在所有来源都校验失败时仍会残留。
        for field in fields:
            row[field] = None
        row_sources: dict[str, str] = {}
        for field in fields:
            for source in priority:
                by_date, by_day, rows = lookups[source]
                candidate = _candidate_for_lifecycle_row(base, index, by_date, by_day, rows)
                value = candidate.get(field)
                if _valid_daily_value(field, value, candidate):
                    row[field] = value
                    row_sources[field] = source
                    source_sets[field].add(source)
                    break
        gross, small, direct = row.get("gross"), row.get("small_to_big"), row.get("direct")
        if all(isinstance(value, (int, float)) for value in (gross, small, direct)):
            tolerance = max(5.0, abs(float(gross)) * 0.02)
            if abs(float(small) + float(direct) - float(gross)) > tolerance:
                structure_issues.append(
                    f"{row.get('date') or row.get('day') or f'D{index + 1}'}："
                    f"大定{gross}/小转大{small}/直接大定{direct}"
                )
                row["small_to_big"] = None
                row["direct"] = None
                row_sources.pop("small_to_big", None)
                row_sources.pop("direct", None)
        row["_field_sources"] = row_sources
        merged.append(row)
    if log_issues and structure_issues:
        preview = "、".join(structure_issues[:8])
        suffix = f"共{len(structure_issues)}天"
        LOGGER.warning(
            "[销量预测字段校验] %s | 小转大+直接大定与总大定不一致=%s | 明细=%s | 处理=对应日期分项标记缺失，不参与真实结构计算",
            diagnostic_context or "当前对象",
            suffix,
            preview,
        )
    field_sources = {
        field: next(iter(sources)) if len(sources) == 1 else "mixed"
        for field, sources in source_sets.items()
        if sources
    }
    field_sources["date_spine"] = base_source
    return merged, field_sources


def _d1_snapshot(candidate: dict[str, Any]) -> str:
    rows = list(candidate.get("days") or [])
    first = rows[0] if rows else {}
    return (
        f"行数={len(rows)},D1日期={first.get('date')!r},"
        f"大定={first.get('gross')!r},小转大={first.get('small_to_big')!r},"
        f"直接大定={first.get('direct')!r},来源Sheet={(candidate.get('day_source') or {}).get('sheet', '')!r}"
    )


def _anchor_lifecycle_rows(rows: list[dict[str, Any]], launch_date: Any) -> list[dict[str, Any]]:
    """Fill unparseable D1-Dn dates from the target's confirmed launch date."""
    anchor = _as_date(launch_date)
    anchored = []
    for index, original in enumerate(rows):
        row = {**original}
        if anchor and not _as_date(row.get("date")):
            row["date"] = (anchor + timedelta(days=index)).isoformat()
        anchored.append(row)
    return anchored


PRE_LAUNCH_SOURCE_PRIORITY = ("cancel", "history")
POST_LAUNCH_SOURCE_PRIORITY = ("history", "launch", "cancel")


STAGE_SOURCE_PRIORITIES: dict[str, tuple[str, ...]] = {
    "ended": POST_LAUNCH_SOURCE_PRIORITY,
    "before": PRE_LAUNCH_SOURCE_PRIORITY,
    "active": ("launch", "history"),
    "unknown": ("launch", "history", "cancel"),
}


def _resolve_stage_candidate(
    candidates: dict[str, dict[str, Any]],
    stage: str,
    *,
    log_issues: bool = False,
    diagnostic_context: str = "",
) -> dict[str, Any]:
    priority = STAGE_SOURCE_PRIORITIES.get(stage, STAGE_SOURCE_PRIORITIES["unknown"])
    availability = {
        source: bool(candidate.get("days") or candidate.get("hourly_days") or _number(candidate.get("total_small")) > 0)
        for source, candidate in candidates.items()
    }
    selected = next((source for source in priority if availability.get(source)), "missing")
    days, day_field_sources = _merge_day_fields(
        candidates,
        priority,
        log_issues=log_issues,
        diagnostic_context=diagnostic_context,
    )
    hourly_days, hourly_source = _field_value(candidates, priority, "hourly_days", lambda value: bool(value))
    total_small, total_small_source = _field_value(candidates, priority, "total_small", lambda value: _number(value) > 0)
    day_source, day_source_owner = _field_value(candidates, priority, "day_source", lambda value: bool(value))
    missing_fields = []
    if not days and stage != "before":
        missing_fields.append("首销日真实进度")
    if not _number(total_small):
        missing_fields.append("总小订")
    field_sources = {
        "首销日明细": SOURCE_LABELS.get(day_field_sources.get("date_spine", day_source_owner), "数据缺失"),
        "总小订": SOURCE_LABELS.get(total_small_source, "数据缺失"),
        "分时进度": SOURCE_LABELS.get(hourly_source, "数据缺失"),
        **{
            f"首销日·{field}": SOURCE_LABELS.get(source, "数据缺失")
            for field, source in day_field_sources.items()
            if field != "date_spine"
        },
    }
    latest_date = max(
        [row.get("date", "") for row in days if row.get("date")]
        + [row.get("date", "") for row in (hourly_days or []) if row.get("date")]
        + [str(candidates.get("cancel", {}).get("cancel_latest_date") or "")],
        default="",
    )
    return {
        "stage": stage,
        "priority": priority,
        "days": days,
        "hourly_days": hourly_days or [],
        "total_small": int(_number(total_small)),
        "total_small_source": total_small_source,
        "day_source": day_source,
        "selected_source": selected,
        "selected_source_label": SOURCE_LABELS[selected],
        "field_sources": field_sources,
        "missing_fields": missing_fields,
        "data_missing": selected == "missing",
        "source_latest_date": latest_date,
    }


def _resolve_actual_profiles(
    history: list[dict[str, Any]],
    profiles: list[dict[str, Any]],
    targets: list[dict[str, Any]],
    history_source: SourceRef,
    today: date | None = None,
) -> list[dict[str, Any]]:
    """Resolve every required field independently using the absolute-stage source order."""
    resolved = []
    for target in targets:
        name = target["name"]
        raw = next((profile for profile in profiles if _same_model(profile.get("model"), name)), None) or {}
        reference = _history_record(history, name)
        stage = target.get("stage") or "unknown"

        history_profile = _profile_from_history(
            name,
            reference,
            history_source,
            target.get("launch_date") or raw.get("launch_date"),
        ) if _history_has_data(reference) else {}
        if not history_profile and _number(target.get("history_small")) > 0:
            history_profile = {
                "model": name,
                "days": [],
                "hourly_days": [],
                "total_small": int(_number(target.get("history_small"))),
                "day_source": history_source.to_dict(),
            }
        has_explicit_history_small = _number(history_profile.get("total_small")) > 0
        launch_small = (
            raw.get("launch_total_small") if raw.get("launch_total_small_valid") else 0
        ) if "launch_total_small" in raw else raw.get("total_small")
        if has_explicit_history_small and raw.get("launch_total_small_estimated"):
            launch_small = 0
        cancel_small = (
            raw.get("cancel_total_small") if raw.get("cancel_total_small_valid") else 0
        ) if "cancel_total_small_valid" in raw else raw.get("cancel_total_small")
        launch_profile = {
            "days": _anchor_lifecycle_rows(
                list(raw.get("days") or []),
                target.get("launch_date") or raw.get("launch_date"),
            ),
            "hourly_days": list(raw.get("hourly_days") or []),
            "total_small": int(_number(launch_small)),
            "day_source": raw.get("day_source"),
            "hour_source": raw.get("hour_source"),
        }
        cancel_days = [
            {"day": f"D{index + 1}", "date": row.get("date", ""), "cancel": row.get("cancel"), "cancel_rate": row.get("cancel_rate")}
            for index, row in enumerate(raw.get("cancel_days") or [])
        ]
        cancel_profile = {
            "days": cancel_days,
            "hourly_days": [],
            "total_small": int(_number(cancel_small)),
            "day_source": raw.get("cancel_source"),
            "cancel_source": raw.get("cancel_source"),
            "cancel_latest_date": raw.get("cancel_latest_date", ""),
        }
        candidates = {"history": history_profile, "launch": launch_profile, "cancel": cancel_profile}
        stage_profiles = {
            stage_key: _resolve_stage_candidate(
                candidates,
                stage_key,
                log_issues=stage_key == stage,
                diagnostic_context=f"预测对象={name} | 阶段={stage_key}",
            )
            for stage_key in ("before", "active", "ended")
        }
        current_profile = stage_profiles.get(stage) or _resolve_stage_candidate(candidates, "unknown")
        total_small_source = current_profile.get("total_small_source")
        if total_small_source == "launch" and raw.get("launch_total_small_estimated"):
            LOGGER.warning(
                "[销量预测估算] 代际=%s | 总小订=%d为首销累计进度反推值（%s）| "
                "处理=当前无显式总小订可用，反推值作为预测分母",
                name, int(_number(current_profile.get("total_small"))), raw.get("launch_total_small_reason") or "反推口径",
            )
        elif total_small_source == "cancel" and raw.get("cancel_total_small_valid"):
            LOGGER.warning(
                "[销量预测估算] 代际=%s | 总小订=%d为累计退订/退订率反推值（%s）| "
                "处理=当前无显式总小订可用，反推值作为预测分母",
                name, int(_number(current_profile.get("total_small"))), raw.get("cancel_total_small_reason") or "反推口径",
            )
        priority = current_profile["priority"]
        selected = current_profile["selected_source"]
        days = current_profile["days"]
        hourly_days = current_profile["hourly_days"]
        total_small = current_profile["total_small"]
        day_source = current_profile["day_source"]
        day_field_sources = {
            key.removeprefix("首销日·"): next(
                (source for source, label in SOURCE_LABELS.items() if label == value),
                "missing",
            )
            for key, value in current_profile["field_sources"].items()
            if key.startswith("首销日·")
        }
        missing_fields = current_profile["missing_fields"]
        if missing_fields:
            LOGGER.warning(
                "[销量预测数据诊断] 预测对象=%s | 阶段=%s | 缺失字段=%s | 来源顺序=%s | 已选来源=%s | 处理=按字段继续回退，仍缺失的字段在网页标记为数据缺失",
                name,
                stage,
                "、".join(missing_fields),
                "→".join(SOURCE_LABELS[source] for source in priority),
                SOURCE_LABELS[selected],
            )
        elif selected != priority[0]:
            LOGGER.warning(
                "[销量预测来源回退] 预测对象=%s | 阶段=%s | 首选来源=%s无可用数据 | 实际来源=%s | 处理=已按阶段优先级回退",
                name,
                stage,
                SOURCE_LABELS[priority[0]],
                SOURCE_LABELS[selected],
            )

        if stage in ("active", "ended"):
            as_of_date = today or date.today()
            target_launch = _as_date(target.get("launch_date"))
            aligned_d1 = next(
                (row for row in days if target_launch and _as_date(row.get("date")) == target_launch),
                None,
            )
            hourly_d1 = next(
                (
                    row for row in (hourly_days or [])
                    if target_launch and _as_date(row.get("date")) == target_launch
                    and isinstance(row.get("gross"), (int, float))
                ),
                None,
            )
            browser_d1_ready = bool(
                target_launch
                and target_launch <= as_of_date
                and (
                    (aligned_d1 and isinstance(aligned_d1.get("gross"), (int, float)))
                    or (target_launch == as_of_date and hourly_d1)
                )
            )
            if not browser_d1_ready:
                if not target_launch:
                    browser_reason = "预测对象首销日期为空"
                elif target_launch > as_of_date:
                    browser_reason = f"D1日期{target_launch.isoformat()}晚于判定日期{as_of_date.isoformat()}"
                elif aligned_d1 is None:
                    browser_reason = f"合并结果没有日期={target_launch.isoformat()}的D1行"
                else:
                    browser_reason = f"对齐D1的大定不是数字：{aligned_d1.get('gross')!r}"
                LOGGER.warning(
                    "[销量预测D1排查] 网页将判定D1缺失：预测对象=%s，阶段=%s，首销日期=%r，"
                    "来源顺序=%s，命中首销对象=%r，命中历史传播名=%r，命中历史代际名=%r，原因=%s",
                    name,
                    stage,
                    target.get("launch_date"),
                    "→".join(SOURCE_LABELS[source] for source in priority),
                    raw.get("model"),
                    reference.get("model") if reference else None,
                    reference.get("generation") if reference else None,
                    browser_reason,
                )
                LOGGER.warning(
                    "[销量预测D1排查] 三来源D1：首销期订单节奏{%s}；小订及首销数据整理{%s}；"
                    "小订退订分析{%s}；合并后{%s}；字段来源=%s",
                    _d1_snapshot(launch_profile),
                    _d1_snapshot(history_profile),
                    _d1_snapshot(cancel_profile),
                    _d1_snapshot({"days": days, "day_source": day_source}),
                    day_field_sources,
                )
            else:
                LOGGER.debug(
                    "[销量预测D1排查] D1可供网页识别：预测对象=%s，阶段=%s，首销日期=%s，"
                    "大定=%s，小转大=%s，直接大定=%s，字段来源=%s",
                    name,
                    stage,
                    target_launch.isoformat(),
                    aligned_d1.get("gross") if aligned_d1 else hourly_d1.get("gross"),
                    aligned_d1.get("small_to_big") if aligned_d1 else hourly_d1.get("small_to_big"),
                    aligned_d1.get("direct") if aligned_d1 else hourly_d1.get("direct"),
                    day_field_sources,
                )

        profile = {
            **raw,
            "model": name,
            "days": days,
            "hourly_days": hourly_days or [],
            "total_small": int(_number(total_small)),
            "day_source": day_source,
            "selected_source": selected,
            "selected_source_label": SOURCE_LABELS[selected],
            "field_sources": current_profile["field_sources"],
            "missing_fields": missing_fields,
            "stage_profiles": stage_profiles,
        }

        target["small"] = profile["total_small"]
        target["selected_source"] = selected
        target["selected_source_label"] = SOURCE_LABELS[selected]
        target["field_sources"] = profile["field_sources"]
        target["missing_fields"] = missing_fields
        target["data_missing"] = current_profile["data_missing"]
        target["source_latest_date"] = current_profile["source_latest_date"] or profile.get("cancel_latest_date") or ""
        profile["stage"] = stage
        profile["data_missing"] = target["data_missing"]
        resolved.append(profile)
    return resolved


def _attach_actual_shapes(history: list[dict[str, Any]], profiles: list[dict[str, Any]], today: date | None = None) -> None:
    for item in history:
        profile = next((profile for profile in profiles if _same_model(profile["model"], item["model"]) or _same_model(profile["model"], item.get("generation"))), None)
        if not profile:
            item.setdefault("daily_orders", [])
            item["hourly_curve"] = []
            continue
        if not item.get("daily_orders"):
            item["daily_orders"] = [row["gross"] for row in profile.get("days", [])]
        first_hourly = next((row for row in profile.get("hourly_days", [])
                             if row.get("date") == item.get("launch_date")
                             and row.get("date", "") < (today or date.today()).isoformat()), {})
        hourly_values = first_hourly.get("hours", [])
        by_hour = [0.0] * 24
        for row in hourly_values:
            hour = row.get("hour")
            value = _optional_number(row.get("gross"))
            if isinstance(hour, int) and 0 <= hour <= 23 and value is not None and value >= 0:
                by_hour[hour] += value
        total = sum(by_hour)
        cumulative = 0.0
        item["hourly_curve"] = []
        if total > 0:
            for value in by_hour:
                cumulative += value
                item["hourly_curve"].append(cumulative / total)


def _attach_small_hourly_curves(small_history: list[dict[str, Any]], profiles: list[dict[str, Any]]) -> None:
    for item in small_history:
        profile = next(
            (
                candidate for candidate in profiles
                if _same_model(candidate.get("model"), item.get("model"))
                or _same_model(candidate.get("model"), item.get("generation"))
            ),
            None,
        )
        if not profile:
            continue
        buckets = profile.get("small_hourly_days") or []
        bucket = next((row for row in buckets if row.get("date") == item.get("small_start_date")), None)
        if bucket is None and buckets and not item.get("small_start_date"):
            bucket = min(buckets, key=lambda row: str(row.get("date") or ""))
        if not bucket:
            continue
        by_hour: dict[int, float] = {}
        for row in bucket.get("hours") or []:
            hour = int(_number(row.get("hour"), -1))
            if 0 <= hour <= 23:
                by_hour[hour] = by_hour.get(hour, 0) + max(_number(row.get("orders")), 0)
        terminal = sum(by_hour.values())
        if terminal <= 0:
            continue
        running = 0.0
        curve = []
        for hour in range(24):
            running += by_hour.get(hour, 0)
            curve.append(min(running / terminal, 1))
        item["small_hourly_curve"] = curve


def _attach_steady_launch_features(steady_history: list[dict[str, Any]], history: list[dict[str, Any]]) -> None:
    for item in steady_history:
        name = item.get("generation") or item.get("model")
        reference = _history_record(history, name)
        if reference:
            item["lock_rate"] = _number(reference.get("lock_rate"))
            item["launch_days"] = int(_number(reference.get("days")))


def _validate_stage_window(model: str, *, start: Any, end: Any, kind: str) -> None:
    """Window boundaries must be internally consistent; contradictions are data defects, not inputs."""
    start_date, end_date = _as_date(start), _as_date(end)
    if start_date and end_date and end_date < start_date:
        LOGGER.warning(
            "[销量预测字段校验] 代际=%s | %s结束日期%s早于开始日期%s | "
            "处理=停止该阶段预测，不用天数覆盖冲突日期，请修正车型汇总",
            model, kind, end_date.isoformat(), start_date.isoformat(),
        )


def _target_options(
    history: list[dict[str, Any]],
    profiles: list[dict[str, Any]],
    stage_windows: dict[str, dict[str, Any]],
    today: date | None = None,
    model_master: dict[str, dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    options = []
    for profile in profiles:
        reference = _history_record(history, profile["model"])
        window = _stage_window(stage_windows, profile["model"])
        master_record = _master_record(
            model_master or {},
            reference.get("model") if reference else "",
            profile["model"],
            reference.get("generation") if reference else "",
        )
        launch_date = (window or {}).get("launch_date") or profile.get("launch_date") or (reference["launch_date"] if reference else "")
        maintained_days = (window or {}).get("days") or (reference["days"] if reference else 0)
        end_date = (window or {}).get("end_date") or (reference.get("end_date") if reference else None)
        _validate_stage_window(profile["model"], start=launch_date, end=end_date, kind="首销窗口")
        _validate_stage_window(profile["model"], start=(window or {}).get("small_start_date"), end=(window or {}).get("small_end_date"), kind="小订窗口")
        parsed_start, parsed_end = _as_date(launch_date), _as_date(end_date)
        if parsed_start and parsed_end and parsed_end >= parsed_start:
            date_span = (parsed_end - parsed_start).days + 1
            if _number(maintained_days) > 0 and int(_number(maintained_days)) != date_span:
                LOGGER.warning(
                    "[销量预测字段校验] 代际=%s | 首销开始=%s | 截止=%s | "
                    "日期包含%d天但维护天数=%d | 处理=以明确开始/截止日期为准",
                    profile["model"], parsed_start.isoformat(), parsed_end.isoformat(),
                    date_span, int(_number(maintained_days)),
                )
        elif parsed_start and not parsed_end and _number(maintained_days) > 0:
            LOGGER.warning(
                "[销量预测估算] 代际=%s | 首销截止日期缺失 | "
                "处理=按明确维护的%d天推导截止日，请补齐截止日期",
                profile["model"], int(_number(maintained_days)),
            )
        if (window or {}).get("small_start_date") and not (window or {}).get("small_end_date"):
            LOGGER.warning(
                "[销量预测字段校验] 代际=%s | 小订结束日期缺失 | "
                "处理=停止小订阶段预测，不猜测为1天窗口",
                profile["model"],
            )
        stage = _forecast_stage(launch_date, end_date, maintained_days, today)
        small_stage = _small_order_stage(
            (window or {}).get("small_start_date"),
            (window or {}).get("small_end_date"),
            today,
        )
        steady_start = (_as_date(stage.get("end_date")) + timedelta(days=1)).isoformat() if _as_date(stage.get("end_date")) else ""
        options.append({
            "name": profile["model"],
            "history_model": reference["model"] if reference else "",
            "tier": usable_attribute(master_record.get("产品档位")) if master_record else usable_attribute(reference["tier"] if reference else None),
            "energy": usable_attribute(master_record.get("能源类型")) if master_record else usable_attribute(reference["energy"] if reference else None),
            "node": usable_attribute(master_record.get("发布类型") or master_record.get("发布节点")) if master_record else usable_attribute(reference["node"] if reference else None),
            "launch_date": stage["launch_date"],
            "end_date": stage["end_date"],
            "stage": stage["key"],
            "stage_label": stage["label"],
            "calendar_day": stage["day"],
            "date_source_label": "小订及首销数据整理 · 车型汇总" if window else "首销期订单节奏 · by天第二行" if profile.get("launch_date") else "日期数据缺失",
            "small_start_date": (window or {}).get("small_start_date", ""),
            "small_end_date": (window or {}).get("small_end_date", ""),
            "small_stage": small_stage["key"],
            "small_stage_label": small_stage["label"],
            "small_calendar_day": small_stage["day"],
            "small_days": small_stage["days"],
            "steady_start_date": steady_start,
            "small_date_source_label": "车型汇总维护" if (window or {}).get("small_start_date") else "车型汇总未维护",
            "steady_date_source_label": "由首销截止次日推导" if (window or {}).get("end_date") else ("由首销开始+天数推算" if _as_date(stage.get("end_date")) else "缺失"),
            "launch_weekday": reference["launch_weekday"] if reference else _weekday(launch_date),
            "launch_period": usable_attribute(master_record.get("发布时段")) if master_record else usable_attribute(reference["launch_period"] if reference else None),
            "days": stage["days"],
            "launch_days_maintained": bool(parsed_end or _number(maintained_days) > 0),
            "small": profile.get("total_small") or (window or {}).get("small") or (reference["small"] if reference else 0),
            "history_small": (window or {}).get("small") or (reference["small"] if reference else 0),
            "actual_model": profile["model"],
        })
    return options


def _merge_small_daily_history(
    profiles: list[dict[str, Any]],
    small_history: list[dict[str, Any]],
    targets: list[dict[str, Any]],
    history_source: SourceRef | None,
) -> None:
    """Resolve real reservation days once for both the webpage and summary detail."""
    for profile in profiles:
        target = next((row for row in targets if _same_model(row.get("name"), profile.get("model"))), None)
        if not target:
            continue
        start = str(target.get("small_start_date") or "")
        end = str(target.get("small_end_date") or "")
        if not (_as_date(start) and _as_date(end)):
            continue
        event = next(
            (
                row for row in small_history
                if row.get("daily_actual") is True
                and row.get("small_start_date") == start
                and _same_model(row.get("generation") or row.get("model"), profile.get("model"))
            ),
            None,
        )
        if not event:
            continue
        by_date = {row["date"]: row for row in profile.get("small_daily_days", []) if row.get("date")}
        sources = profile.setdefault("small_daily_sources", {})
        for day, orders in zip(event.get("dates") or [], event.get("daily_orders") or []):
            if not day or not start <= day <= end or (day in by_date and by_date[day].get("orders") is not None):
                continue
            if not isinstance(orders, (int, float)) or not math.isfinite(orders) or orders < 0:
                continue
            by_date[day] = {"date": day, "orders": int(round(orders))}
            if history_source:
                sources[day] = history_source.to_dict()
        profile["small_daily_days"] = [by_date[day] for day in sorted(by_date)]


def _profile_hard_errors(
    target: dict[str, Any],
    profile: dict[str, Any],
    today: date | None = None,
) -> list[str]:
    """Return blocking raw-data errors for one forecast target.

    Forecasting may fall back between the three approved sources field by field,
    but it must never silently estimate an already completed calendar day or a
    missing core denominator.  Those are source-data defects, not forecast inputs.
    """
    current = today or date.today()
    name = str(target.get("name") or profile.get("model") or "未命名代际")
    # Final small-order volume is an output while reservations are ongoing.
    # Completed reservation days are validated by the small-order workspace.
    small_end = _as_date(target.get("small_end_date"))
    if small_end is not None and current <= small_end:
        return []
    stage = str(target.get("stage") or "unknown")
    start = _as_date(target.get("launch_date"))
    span = max(int(_number(target.get("days"), 0)), 0)
    errors: list[str] = []

    if start is None:
        errors.append("首销开始日期缺失或无法解析")
    if span <= 0:
        errors.append("首销期天数缺失或不是正整数")
    elif not target.get("launch_days_maintained", True):
        errors.append("首销截止日期和首销期天数均未维护")
    if _number(profile.get("total_small")) <= 0:
        errors.append("总小订缺失或不大于0")

    profile_dates = [
        parsed for row in profile.get("days", [])
        if (parsed := _as_date(row.get("date"))) is not None
    ]
    if start and profile_dates and min(profile_dates) != start:
        errors.append(
            f"首销真实by天起始日{min(profile_dates).isoformat()}"
            f"与车型汇总开始日{start.isoformat()}不一致"
        )

    if stage == "unknown":
        errors.append("无法按绝对日期判断首销阶段")
    if start is None or span <= 0 or stage not in {"active", "ended"}:
        return list(dict.fromkeys(errors))

    completed_days = span if stage == "ended" else min(max((current - start).days, 0), span)
    rows_by_date = {
        parsed.isoformat(): row
        for row in profile.get("days", [])
        if (parsed := _as_date(row.get("date"))) is not None
    }
    missing_days: list[int] = []
    invalid_days: list[str] = []
    for index in range(completed_days):
        expected_date = (start + timedelta(days=index)).isoformat()
        row = rows_by_date.get(expected_date)
        if row is None:
            missing_days.append(index + 1)
            continue
        gross = _optional_number(row.get("gross"))
        small_to_big = _optional_number(row.get("small_to_big"))
        direct = _optional_number(row.get("direct"))
        missing_fields = [
            label
            for label, value in (("大定", gross), ("小转大", small_to_big), ("直接大定", direct))
            if value is None or not math.isfinite(float(value)) or value < 0
        ]
        if missing_fields:
            invalid_days.append(f"D{index + 1}缺{'+'.join(missing_fields)}")
            continue
        if abs((small_to_big + direct) - gross) > max(1.0, abs(gross) * .005):
            invalid_days.append(
                f"D{index + 1}分项不一致（小转大{small_to_big:g}+直接大定{direct:g}≠大定{gross:g}）"
            )
    if missing_days:
        labels = "、".join(f"D{day}" for day in missing_days[:12])
        suffix = f"等{len(missing_days)}天" if len(missing_days) > 12 else ""
        errors.append(f"已结束日期缺少真实数据：{labels}{suffix}")
    if invalid_days:
        labels = "；".join(invalid_days[:8])
        suffix = f"；另有{len(invalid_days) - 8}天" if len(invalid_days) > 8 else ""
        errors.append(f"已结束日期字段缺失或校验失败：{labels}{suffix}")
    return list(dict.fromkeys(errors))


def _default_target_option(targets: list[dict[str, Any]]) -> dict[str, Any] | None:
    """Prefer the currently active launch, then the nearest future launch, then the latest ended launch."""
    if not targets:
        return None
    active = [item for item in targets if item.get("stage") == "active"]
    if active:
        return max(active, key=lambda item: int(_number(item.get("calendar_day"))))
    before = [item for item in targets if item.get("stage") == "before"]
    if before:
        return min(before, key=lambda item: _as_date(item.get("launch_date")) or date.max)
    ended = [item for item in targets if item.get("stage") == "ended"]
    if ended:
        return max(ended, key=lambda item: _as_date(item.get("launch_date")) or date.min)
    return targets[0]


class SalesForecastModule:
    id = "sales_forecast"
    label = "销量预测"

    def __init__(self) -> None:
        self.as_of_date: date | None = None
        self.summary_path = PROJECT_ROOT / "output_file" / SUMMARY_NAME

    def set_config(self, config: dict[str, Any]) -> None:
        """Apply the optional manual date without changing source data."""
        self.as_of_date = _configured_forecast_date(config.get("forecast_as_of_date"))
        if self.as_of_date:
            LOGGER.info(
                "[销量预测日期] 使用 config.json 手工判定日期=%s；仅影响阶段、已发生/未来切分和预测窗口",
                self.as_of_date.isoformat(),
            )

    def build(self, store, subject: Subject) -> Dashboard | None:
        if subject.type != "group":
            return None
        if self.summary_path.exists():
            with summary_scope(self.summary_path) as summary:
                if summary.get("snapshot"):
                    snapshot = summary["snapshot"]
                    views = deepcopy(snapshot["views"])
                    page = views["week"]["pages"]["预测方案"]
                    page["workspace"]["source"] = SourceRef(
                        self.summary_path.name,
                        "当前订单逐日",
                        "销量预测汇总数据",
                    ).to_dict()
                    sources = [
                        SourceRef(self.summary_path.name, "车型基本信息", "车型信息与当前阶段"),
                        SourceRef(self.summary_path.name, "小订及退订逐日", "历史及当前小订与退订"),
                        SourceRef(self.summary_path.name, "当前订单逐日", "当前订单"),
                        SourceRef(self.summary_path.name, "预测基准总表", "首销历史参考"),
                    ]
                    return Dashboard(self.id, subject.id, views, sources)
                return self._build_from_sources(summary["store"], subject)
        return self._build_from_sources(store, subject)

    def _build_from_sources(self, store, subject: Subject) -> Dashboard | None:
        path, history = _read_history()
        if not history or path is None:
            return None
        raw_profiles, actual_sources = _read_actual_profiles(store)
        _attach_actual_shapes(history, raw_profiles, today=self.as_of_date)
        small_history_path, small_history = _read_small_order_history()
        _attach_small_hourly_curves(small_history, raw_profiles)
        stage_window_path, stage_windows = _read_stage_windows()
        steady_history, steady_sources = _read_steady_history(
            store, stage_windows, today=self.as_of_date
        )
        _attach_steady_launch_features(steady_history, history)
        model_master = _read_model_master()
        model_mapping = _read_model_mapping()
        if not model_master:
            LOGGER.warning(
                "[映射诊断] 文件=%s | 状态=未读取到车型基本信息 | 影响=传播名—代际名映射和车型属性可能缺失",
                MODEL_MAPPING_CANDIDATES[0],
            )
        unmapped_history = [item["model"] for item in history if not item.get("mapped")]
        if unmapped_history:
            LOGGER.warning(
                "[映射诊断] 模块=销量预测 | 未映射历史传播名=%d个 | 明细=%s | 处理=仅作为历史参考，不自动绑定订单分析代际",
                len(unmapped_history),
                "、".join(unmapped_history),
            )
        targets = _target_options(
            history,
            raw_profiles,
            stage_windows,
            today=self.as_of_date,
            model_master=model_master,
        )
        history_sheet = "预测基准总表" if ACTIVE_SUMMARY.get() or path.name in {SUMMARY_NAME, "销量预测数据汇总.xlsx", "小订及首销预测二次处理.xlsx"} else "首销预测基准" if path.name == "小订及首销期历史基准.xlsx" else "车型汇总"
        history_source = SourceRef(path.name, history_sheet, "历史参考传播名与预测二次处理基准")
        profiles = _resolve_actual_profiles(
            history, raw_profiles, targets, history_source, today=self.as_of_date
        )
        _merge_small_daily_history(
            profiles,
            small_history,
            targets,
            SourceRef(small_history_path.name, "小订by天", "历史真实逐日小订") if small_history_path else None,
        )
        for option in targets:
            profile = next((item for item in profiles if _same_model(item.get("model"), option["name"])), {})
            hard_errors = _profile_hard_errors(option, profile, today=self.as_of_date)
            option["hard_errors"] = hard_errors
            option["data_error"] = bool(hard_errors)
            if profile:
                profile["hard_errors"] = hard_errors
                profile["data_error"] = bool(hard_errors)
            if hard_errors:
                LOGGER.error(
                    "[销量预测原始数据错误] 预测对象=%s | 阶段=%s | %s | 处理=停止该对象两种预测，修正原始数据后重新生成",
                    option["name"],
                    option.get("stage_label", option.get("stage", "未知")),
                    "；".join(hard_errors),
                )
        default_option = _default_target_option(targets)
        if default_option is None:
            return None
        default_reference = _history_record(history, default_option["name"])
        target = {
            "name": default_option["name"],
            "history_model": default_option["history_model"],
            "tier": default_option["tier"] or "未维护",
            "energy": default_option["energy"] or "未维护",
            "launch_node": default_option["node"] or "未维护",
            "launch_date": default_option["launch_date"],
            "end_date": default_option.get("end_date", ""),
            "stage": default_option.get("stage", "unknown"),
            "stage_label": default_option.get("stage_label", "时间缺失"),
            "calendar_day": default_option.get("calendar_day", 0),
            "date_source_label": default_option.get("date_source_label", "日期数据缺失"),
            "small_start_date": default_option.get("small_start_date", ""),
            "small_end_date": default_option.get("small_end_date", ""),
            "small_stage": default_option.get("small_stage", "unknown"),
            "small_stage_label": default_option.get("small_stage_label", "小订时间缺失"),
            "small_calendar_day": default_option.get("small_calendar_day", 0),
            "small_days": default_option.get("small_days", 0),
            "steady_start_date": default_option.get("steady_start_date", ""),
            "small_date_source_label": default_option.get("small_date_source_label", ""),
            "steady_date_source_label": default_option.get("steady_date_source_label", ""),
            "selected_source": default_option.get("selected_source", "missing"),
            "selected_source_label": default_option.get("selected_source_label", "数据缺失"),
            "data_missing": default_option.get("data_missing", True),
            "field_sources": default_option.get("field_sources", {}),
            "missing_fields": default_option.get("missing_fields", []),
            "hard_errors": default_option.get("hard_errors", []),
            "data_error": default_option.get("data_error", False),
            "source_latest_date": default_option.get("source_latest_date", ""),
            "launch_weekday": default_option.get("launch_weekday") or _weekday(default_option["launch_date"]),
            "launch_period": default_option.get("launch_period") or "未维护",
            "launch_days": default_option["days"] or 35,
            "total_small": default_option["small"],
            "conversion": _number(default_reference and default_reference.get("conversion")),
            "direct_share": _number(default_reference and default_reference.get("direct_share")),
            "lock_rate": _number(default_reference and default_reference.get("lock_rate")),
        }
        default_profile = next((item for item in profiles if _same_model(item["model"], target["name"])), None)
        workspace_source = SourceRef(
            default_profile["day_source"]["file"], default_profile["day_source"]["sheet"], default_profile.get("selected_source_label", "当前阶段优先数据源")
        ) if default_profile and default_profile.get("day_source") else history_source
        data = {
            "target": target,
            "targets": targets,
            "history": history,
            "history_source": history_source.to_dict(),
            "model_aliases": model_mapping,
            "actuals": profiles,
            "small_order_history": small_history,
            "small_order_source": SourceRef(small_history_path.name, "小订by天", "小订历史真实逐日曲线").to_dict() if small_history_path else None,
            "small_order_tasks": [{"key": key, "label": label, "purpose": purpose} for key, label, purpose in SMALL_ORDER_TASKS],
            "steady_history": steady_history,
            "steady_tasks": [{"key": key, "label": label, "purpose": purpose} for key, label, purpose in STEADY_TASKS],
            "tiers": _master_values(model_master, "产品档位"),
            "energies": _master_values(model_master, "能源类型"),
            "nodes": _master_values(model_master, "发布类型"),
            "day_type_defaults": {"workday": 1.0, "weekend": 1.15, "holiday": 1.30},
            "china_calendar": calendar_payload(),
            "tasks": [{"key": key, "label": label, "purpose": purpose} for key, label, purpose in TASKS],
            "method": [
                "先按产品档位确定候选传播名；保留标准发布、年度换代、年度改款、品牌首发等发布类型，并按首销天数及首销窗口内法定节假日、周末和调休工作日结构确定强参考传播名",
                "D1 分时参考同时比较发布类型、发布星期与上午/下午/晚上时段；首日未结束时，用当前分时累计进度反推 D1 终值",
                "方法一：当前车型只提供已发生累计数量；先用分母已知的D1/D2订单结构与斜率选择参考传播名，再用历史参考传播名同期累计完成率反推两个终局分量，不计算未知的当前车型自身完成率",
                "方法一只采用历史参考传播名的可靠累计完成率曲线；每条曲线先按当前车型首销天数截取并重新定基，参考周期较短时按末段日增量衰减趋势外推，再计算当前Dn完成率；不使用D1或D1+D2汇总值自动回退",
                "方法二：小转大 = 总小订 × 小订转化率；直接大定按最终直接大定占比换算；总大定始终等于两个分量相加",
                "D1 未开始时显示专项预测，进行中按分时完成率滚动，D1 结束后只读取并冻结真实值，不再用预测值替代",
                "两种方法保留各自终局和剩余量；共同优先使用到天基础曲线安排未来走势，小转大、直接大定分别以前一完整真实日衔接，剔除前日及参考曲线的日历影响后，乘未来日期系数，再按人工可调的逐日权重渐进分配差额（默认0、0.5、1、1……，第3天起封顶）。真实日不变；缺少已结束日时停止预测，不补估真实数据。",
                "总体量级只做合理性校验：按可比传播名大定/总小订比率折算当前量级区间，不直接缩放或覆盖预测结果",
            ],
            "source_rules": [
                "可选预测代际取《首销期订单节奏》《小订退订分析》《小订选配比例分析》《锁单选配比例分析》中可识别代际的并集；锁单选配表不改变小订和首销的来源优先级，平销真实锁单按其by天及by周读取。",
                "所有关联先通过《车型基本信息》的传播名—代际名映射归一到订单分析代际名；网页当前对象与数据源匹配均使用代际名。",
                "传播名没有维护代际映射时，二次处理文件保留该历史记录并将代际名暂按传播名，同时标记“未映射”；它仍可作为历史参考，但不会猜测性绑定当前订单分析代际，需补充映射后刷新。别名冲突同样不自动匹配。",
                "当前阶段按绝对日期判断：小订开始/结束、首销开始/结束优先读取《小订及首销数据整理》的“车型汇总”；若该代际没有汇总日期，才用《首销期订单节奏》对应by天 Sheet第二行的首个绝对日期作为首销开始日，并按首销天数计算截止日。",
                "首销结束且真实数据完整时，两种方法均以真实累计收口，不依赖预测参考或分配参数，也不再分配剩余量；真实日期缺失仍停止预测。今天早于D1为首销期未开始；位于D1至首销截止之间（含首尾日期）为首销期进行中；晚于首销截止为首销期已结束。当前Dn按今天与D1的自然日差计算，不按已读取的数据行数计算。",
                "小订结束至首销开始前视为同一个阶段：优先使用《小订退订分析》，其次使用《小订及首销数据整理》；两处都没有则标记数据缺失。",
                "首销期已结束，以及平销期计算中需要引用的首销与小订数据，使用同一来源顺序：优先使用《小订及首销数据整理》，缺失时依次回退《首销期订单节奏》《小订退订分析》；三处都没有则标记数据缺失。平销期已发生周固定读取《锁单选配比例分析》的交车锁单。",
                "首销期进行中：优先使用《首销期订单节奏》，其次使用《小订及首销数据整理》；两处都没有则标记数据缺失。",
                "by天中早于今天、有日期且大定为数值的行视为已结束真实日期；今天的数据不直接当作完整日冻结。",
                "今天优先取匹配的by时 Sheet，并按最后一个已采集小时反推当日终值；没有by时时，今天的by天行只作为当日快照。",
                "各阶段先确定来源优先顺序，再对总小订、首销日明细、分时进度及各日小转大/直接大定/锁单/退订逐字段回退；主数据源字段为空或一致性校验不通过时自动读取下一顺位，三处均无可靠值才标记该字段缺失。",
                "日期类型按国务院办公厅年度节假日安排自动判断；周末调休上班日按工作日处理，不读取人工维护的节假日日期字段。",
            ],
            "small_order_rules": [
                "小订D1未到：优先使用线索量、互联网热度与产品/发布属性测算最终总小订；驱动字段未维护时仅输出可比车型量级并标低置信度。",
                "小订D1当天：当前分时累计小订除以主辅参考车型相同小时完成率得到D1终值，再除以历史D1占最终小订比例得到小订期总量。",
                "小订D1已过：已发生累计小订除以主辅参考车型同Dn累计完成率反推最终总量；已结束日期冻结真实值。",
                "未来小订以前一完整真实日为衔接基准，参考曲线先剔除日历影响，再应用当前日期系数并渐进分配与剩余量的差额；网页可分阶段调整工作日、周末、节假日系数与逐日差额权重。无前日真实值时明确标注参考分配。",
                "小订结束后即进入首销前衔接阶段，最终总小订和相关实际字段按《小订退订分析》→《小订及首销数据整理》的顺序逐字段取值。",
                "逐日小订按《小订选配比例分析》代际by天→《小订退订分析》分时汇总→《小订及首销数据整理》小订by天的顺序逐字段取值；前两者没有同车型同日期可用值时才使用历史小订by天。“小订进度”标准化曲线只作为最后回退。",
            ],
            "steady_rules": [
                "平销真实日只读取《锁单选配比例分析》的代际by天，完整周读取by周；指标固定为交车锁单，不再读取独立平销历史文件，也不再预测平销大定。",
                "平销从首销截止次日开始；当前周已结束日冻结真实锁单，今天及未来按天预测再汇总为周，混合周明确区分已实现和预测；首销重叠周只纳入平销日期。",
                "基线优先本车型已有平销真实日（最近最多7日）及完整周；其次本车型首销已结束日的直接大定×同期大定到交车锁单率，明确标为估算；尚无已结束首销日时才参考历史车型。自身首销已开始但数据不足时显示缺失原因。",
                "趋势与基线采用同级来源：自身日数据用最近两组完整7日均量之比，不足两组或前组为0时暂按持平；剔除日历影响后预测4周滚动量，已实现真实锁单始终冻结。",
                "滚动区间总量扣除已实现锁单，剩余量以前一天真实锁单衔接，乘未来日历系数并按网页逐日权重渐进分配；不改真实日，不使用大定兜底；昨天或本周已结束日缺失时停止预测。",
                "参考车型评分使用产品档位、能源类型、发布类型及首销期大定到锁单率；所有平销KPI、图表、周合计和来源说明均为交车锁单口径。",
            ],
            "description": "车型汇总提供完整绝对时间窗口；阶段内真实数据按阶段选择订单节奏、退订分析或整理表，二次处理历史库提供参考完成比例与未来斜率。",
        }
        page = {
            "layout": "workspace",
            "workspace": {
                "kind": "forecast_workspace",
                "title": "销量预测",
                "meta": "小订、首销与平销全生命周期 · 真实进度滚动更新 · 分阶段参考",
                "data": data,
                "source": workspace_source.to_dict(),
            },
            "kpis": [],
            "sections": [],
        }
        stage_window_source = SourceRef(stage_window_path.name, "车型汇总", "小订及首销绝对日期窗口") if stage_window_path else None
        small_history_source = SourceRef(small_history_path.name, "小订by天", "小订历史真实逐日曲线") if small_history_path else None
        sources = [history_source, *([stage_window_source] if stage_window_source else []), *([small_history_source] if small_history_source else []), *actual_sources, *steady_sources]
        views = {"week": {"periods": ["预测方案"], "default_period": "预测方案", "pages": {"预测方案": page}}}
        return Dashboard(self.id, subject.id, views, sources)
