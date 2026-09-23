from __future__ import annotations

import unittest
from datetime import date, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook

from core.models import SourceRef
from modules.sales_forecast import (
    _as_datetime,
    _attach_steady_launch_features,
    _configured_forecast_date,
    _default_target_option,
    _forecast_stage,
    _history_item,
    _iso,
    _master_values,
    _merge_small_daily_history,
    _merge_day_fields,
    _profile_hard_errors,
    _read_actual_profiles,
    _read_model_master,
    _read_small_order_history,
    _read_steady_history,
    _resolve_actual_profiles,
    _same_model,
    _small_order_stage,
    _stable_implied_total,
    _target_options,
)


class SalesForecastStageTests(unittest.TestCase):
    def test_configured_forecast_date_accepts_blank_or_iso_date(self):
        self.assertIsNone(_configured_forecast_date(""))
        self.assertEqual(_configured_forecast_date("2026-09-05"), date(2026, 9, 5))
        with self.assertRaises(ValueError):
            _configured_forecast_date("2026/09/05")

    def test_steady_daily_only_reads_lock_after_launch_and_preserves_missing(self):
        workbook=Workbook()
        sheet=workbook.active
        sheet.title='问界 M9 2026款by天'
        sheet.append(['指标','统计类型','分类',date(2026,7,1),date(2026,7,2),date(2026,7,3),date(2026,7,4)])
        sheet.append(['交车锁单','数量','数量',99,10,None,30])
        item=SimpleNamespace(path=Path('锁单选配比例分析.xlsx'),workbook=workbook)
        store=SimpleNamespace(find_all=lambda _: [item],find=lambda _:item)
        windows={'m9':{'generation':'问界 M9 2026款','end_date':'2026-07-01'}}
        with patch('modules.sales_forecast._read_model_mapping',return_value={}),patch('modules.sales_forecast._read_model_master',return_value={}):
            rows,_=_read_steady_history(store,windows,today=date(2026,7,3))
        self.assertEqual(rows[0]['weeks'],[])
        self.assertEqual(rows[0]['daily'],[{'date':'2026-07-02','lock':10,'complete':True},{'date':'2026-07-03','lock':None,'complete':False}])
        workbook.close()

    def test_small_period_does_not_require_final_total_or_launch_days(self):
        target = {'name':'测试车','small_end_date':'2026-09-10','stage':'before'}
        self.assertEqual(_profile_hard_errors(target, {'total_small':0}, today=date(2026,9,4)), [])

    def test_cancel_hourly_orders_are_aggregated_by_small_order_date(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '问界 M9 2026款_小订分时退订'
        sheet.append(['小订日期', '小订小时', '小订数'])
        sheet.append([date(2026,9,1), 0, 10])
        sheet.append([date(2026,9,1), 12, 20])
        sheet.append([date(2026,9,2), 10, 0])
        item = SimpleNamespace(path=Path('小订退订分析.xlsx'), workbook=workbook)
        store = SimpleNamespace(find=lambda keyword: item if keyword=='小订退订分析' else None)
        with patch('modules.sales_forecast._read_model_mapping', return_value={}):
            profiles, _ = _read_actual_profiles(store)
        self.assertEqual(profiles[0]['small_daily_days'], [{'date':'2026-09-01','orders':30},{'date':'2026-09-02','orders':0}])
        workbook.close()

    def test_small_daily_priority_is_small_mix_then_cancel(self):
        mix_book = Workbook()
        mix_sheet = mix_book.active
        mix_sheet.title = '问界 M9 2026款by天'
        mix_sheet.append(['指标', '统计类型', '分类', date(2026, 9, 1), date(2026, 9, 2), date(2026, 9, 3)])
        mix_sheet.append(['小订', '数量', '数量', 100, 0, None])
        cancel_book = Workbook()
        cancel_sheet = cancel_book.active
        cancel_sheet.title = '问界 M9 2026款_小订分时退订'
        cancel_sheet.append(['小订日期', '小订小时', '小订数'])
        cancel_sheet.append([date(2026, 9, 1), 12, 30])
        cancel_sheet.append([date(2026, 9, 2), 12, 25])
        cancel_sheet.append([date(2026, 9, 3), 12, 40])
        items = {
            '小订选配比例': SimpleNamespace(path=Path('鸿蒙智行小订选配比例分析.xlsx'), workbook=mix_book),
            '小订退订分析': SimpleNamespace(path=Path('鸿蒙智行小订退订分析.xlsx'), workbook=cancel_book),
        }
        store = SimpleNamespace(find=lambda keyword: items.get(keyword))
        with patch('modules.sales_forecast._read_model_mapping', return_value={}):
            profiles, _ = _read_actual_profiles(store)
        self.assertEqual(profiles[0]['small_daily_days'], [
            {'date': '2026-09-01', 'orders': 100},
            {'date': '2026-09-02', 'orders': 0},
            {'date': '2026-09-03', 'orders': 40},
        ])
        self.assertEqual(profiles[0]['small_daily_sources']['2026-09-01']['file'], '鸿蒙智行小订选配比例分析.xlsx')
        self.assertEqual(profiles[0]['small_daily_sources']['2026-09-02']['file'], '鸿蒙智行小订选配比例分析.xlsx')
        self.assertEqual(profiles[0]['small_daily_sources']['2026-09-03']['file'], '鸿蒙智行小订退订分析.xlsx')
        mix_book.close()
        cancel_book.close()

    def test_small_daily_history_fallback_matches_window_and_keeps_real_zero(self):
        profiles = [{
            "model": "问界 M9 2026款",
            "small_daily_days": [
                {"date": "2026-09-01", "orders": 100},
                {"date": "2026-09-02", "orders": 0},
            ],
            "small_daily_sources": {
                "2026-09-01": {"file": "鸿蒙智行小订选配比例分析.xlsx"},
                "2026-09-02": {"file": "鸿蒙智行小订选配比例分析.xlsx"},
            },
        }]
        history = [{
            "model": "M9 2026款", "generation": "问界 M9 2026款",
            "small_start_date": "2026-09-01", "daily_actual": True,
            "dates": ["2026-09-01", "2026-09-02", "2026-09-03"],
            "daily_orders": [20, 30, 40],
        }]
        targets = [{
            "name": "问界 M9 2026款",
            "small_start_date": "2026-09-01",
            "small_end_date": "2026-09-03",
        }]
        with patch("modules.sales_forecast._read_model_mapping", return_value={}):
            _merge_small_daily_history(
                profiles, history, targets, SourceRef("历史整理.xlsx", "小订by天", "真实逐日"),
            )
        self.assertEqual(profiles[0]["small_daily_days"], [
            {"date": "2026-09-01", "orders": 100},
            {"date": "2026-09-02", "orders": 0},
            {"date": "2026-09-03", "orders": 40},
        ])
        self.assertEqual(profiles[0]["small_daily_sources"]["2026-09-03"]["sheet"], "小订by天")
        self.assertEqual(profiles[0]["small_daily_sources"]["2026-09-02"]["file"], "鸿蒙智行小订选配比例分析.xlsx")
        history[0]["daily_actual"] = False
        profiles[0]["small_daily_days"] = []
        with patch("modules.sales_forecast._read_model_mapping", return_value={}):
            _merge_small_daily_history(profiles, history, targets, None)
        self.assertEqual(profiles[0]["small_daily_days"], [])

    def test_small_order_stage_keeps_d1_separate(self):
        self.assertEqual(_small_order_stage("2026-09-01", "2026-09-10", date(2026, 8, 31))["key"], "before")
        self.assertEqual(_small_order_stage("2026-09-01", "2026-09-10", date(2026, 9, 1))["key"], "d1")
        active = _small_order_stage("2026-09-01", "2026-09-10", date(2026, 9, 4))
        self.assertEqual((active["key"], active["day"]), ("active", 4))
        self.assertEqual(_small_order_stage("2026-09-01", "2026-09-10", date(2026, 9, 11))["key"], "ended")

    def test_lock_mix_generation_is_added_to_forecast_targets(self):
        workbook = Workbook()
        workbook.active.title = "问界 M9 2026款by周"
        workbook.create_sheet("问界by周")
        workbook.create_sheet("鸿蒙智行by周")
        workbook.create_sheet("问界 M9 2026款汇总by周")
        lock_item = SimpleNamespace(path=Path("鸿蒙智行锁单选配比例分析.xlsx"), workbook=workbook)
        store = SimpleNamespace(find=lambda keyword: lock_item if keyword == "锁单选配比例" else None)

        with patch("modules.sales_forecast._read_model_mapping", return_value={}):
            profiles, sources = _read_actual_profiles(store)

        self.assertEqual([profile["model"] for profile in profiles], ["问界 M9 2026款"])
        self.assertEqual([(source.file, source.sheet, source.note) for source in sources], [
            ("鸿蒙智行锁单选配比例分析.xlsx", "问界 M9 2026款by周", "预测候选代际"),
        ])

    def test_small_order_history_prefers_actual_by_day_curve(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "小订及首销数据整理.xlsx"
            workbook = Workbook()
            summary = workbook.active
            summary.title = "车型汇总"
            summary.append(["车型", "小订开始日期", "小订结束日期", "小订天数", "总小订"])
            summary.append(["测试车", datetime(2026, 9, 1), datetime(2026, 9, 3), 3, 100])
            daily = workbook.create_sheet("小订by天")
            daily.append(["小订", datetime(2026, 9, 1), datetime(2026, 9, 2), datetime(2026, 9, 3), "合计"])
            daily.append(["测试车", 50, 30, 20, 100])
            progress = workbook.create_sheet("小订进度")
            progress.append(["车型", "D1", "D2", "D3"])
            progress.append(["测试车", .5, .8, 1])
            workbook.save(path)
            workbook.close()
            with patch("modules.sales_forecast._read_model_mapping", return_value={}), patch("modules.sales_forecast._read_model_master", return_value={}):
                _, rows = _read_small_order_history(path)
        self.assertEqual(rows[0]["daily_orders"], [50, 30, 20])
        self.assertEqual(rows[0]["small_progress"], [.5, .8, 1])
        self.assertEqual(rows[0]["d1_share"], .5)

    def test_small_order_history_preserves_distinct_propagation_events_of_same_generation(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "小订及首销数据整理.xlsx"
            workbook = Workbook()
            summary = workbook.active
            summary.title = "车型汇总"
            summary.append(["车型", "小订开始日期", "小订结束日期", "小订天数", "总小订"])
            summary.append(["测试车", datetime(2026, 9, 1), datetime(2026, 9, 2), 2, 150])
            daily = workbook.create_sheet("小订by天")
            daily.append(["小订", datetime(2026, 9, 1), datetime(2026, 9, 2), "合计"])
            daily.append(["变体A", 60, 40, 100])
            daily.append(["变体B", 80, 70, 150])
            workbook.save(path)
            workbook.close()
            mapping = {"变体a": "测试车", "变体b": "测试车", "测试车": "测试车"}
            with patch("modules.sales_forecast._read_model_mapping", return_value=mapping), patch("modules.sales_forecast._read_model_master", return_value={}):
                with self.assertLogs("modules.sales_forecast", level="WARNING") as logs:
                    _, rows = _read_small_order_history(path)
        self.assertEqual(len(rows), 2)
        self.assertEqual({row["model"] for row in rows}, {"变体A", "变体B"})
        self.assertEqual(len({row["event_id"] for row in rows}), 2)
        self.assertEqual({tuple(row["daily_orders"]) for row in rows}, {(60, 40), (80, 70)})

    def test_small_order_aliases_share_one_reference_per_start_date(self):
        for alias_first in (True, False):
            with self.subTest(alias_first=alias_first), TemporaryDirectory() as directory:
                path = Path(directory) / "small.xlsx"
                workbook = Workbook()
                summary = workbook.active
                summary.title = "车型汇总"
                summary.append(["车型", "小订开始日期", "小订结束日期", "小订天数", "总小订"])
                daily = workbook.create_sheet("小订by天")
                for start_day in (1, 10):
                    start, end = datetime(2026, 8, start_day), datetime(2026, 8, start_day + 1)
                    summary.append(["历史车型", start, end, 2, 100])
                    daily.append(["小订", start, end, "合计"])
                    rows = [["简称", 70, 30, 100], ["历史车型", 60, 40, 100]]
                    for row in rows if alias_first else reversed(rows):
                        daily.append(row)
                workbook.save(path)
                workbook.close()
                record = {"历史传播名": "历史车型", "订单分析代际名": "标准代际", "原始表简称/别名": "简称|其他别名"}
                master = {"历史车型": record, "简称": record, "标准代际": record}
                mapping = {"历史车型": "标准代际", "简称": "标准代际", "标准代际": "标准代际"}
                with patch("modules.sales_forecast._read_model_mapping", return_value=mapping), patch("modules.sales_forecast._read_model_master", return_value=master):
                    _, rows = _read_small_order_history(path)
                self.assertEqual(len(rows), 2)
                self.assertEqual({row["model"] for row in rows}, {"历史车型"})
                self.assertEqual({row["small_start_date"] for row in rows}, {"2026-08-01", "2026-08-10"})
                self.assertEqual(len({row["event_id"] for row in rows}), 2)
                self.assertTrue(all(row["daily_orders"] == [60, 40] and row["total"] == 100 for row in rows))

    def test_small_order_alias_year_abbreviation_maps_to_history_name(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "small.xlsx"
            workbook = Workbook()
            daily = workbook.active
            daily.title = "小订by天"
            daily.append(["小订", datetime(2026, 8, 1), datetime(2026, 8, 2), "合计"])
            daily.append(["R7 2026款", 60, 40, 100])
            daily.append(["智界R7 2026款", 60, 40, 100])
            workbook.save(path)
            workbook.close()
            record = {
                "历史传播名": "智界R7 2026款",
                "订单分析代际名": "智界 R7 2026款",
                "原始表简称/别名": "R7 26款",
            }
            master = {"智界r72026": record, "r726": record}
            mapping = {name: record["订单分析代际名"] for name in master}
            with patch("modules.sales_forecast._read_model_mapping", return_value=mapping), patch("modules.sales_forecast._read_model_master", return_value=master):
                _, rows = _read_small_order_history(path)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["model"], "智界R7 2026款")
        self.assertEqual(rows[0]["generation"], "智界 R7 2026款")
        self.assertEqual(rows[0]["source_model"], "智界R7 2026款")
        self.assertEqual(rows[0]["daily_orders"], [60, 40])

    def test_small_order_shared_generation_is_not_an_alias_relationship(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "small.xlsx"
            workbook = Workbook()
            daily = workbook.active
            daily.title = "小订by天"
            daily.append(["小订", datetime(2026, 8, 1), datetime(2026, 8, 2), "合计"])
            daily.append(["简称A", 60, 40, 100])
            daily.append(["车型A", 60, 40, 100])
            daily.append(["简称B", 80, 70, 150])
            daily.append(["车型B", 80, 70, 150])
            workbook.save(path)
            workbook.close()
            a = {"历史传播名": "车型A", "订单分析代际名": "同代际", "原始表简称/别名": "简称A"}
            b = {"历史传播名": "车型B", "订单分析代际名": "同代际", "原始表简称/别名": "简称B"}
            master = {"车型a": a, "简称a": a, "车型b": b, "简称b": b}
            mapping = {name: "同代际" for name in master}
            with patch("modules.sales_forecast._read_model_mapping", return_value=mapping), patch("modules.sales_forecast._read_model_master", return_value=master):
                _, rows = _read_small_order_history(path)
            self.assertEqual({row["model"] for row in rows}, {"车型A", "车型B"})
            self.assertEqual({tuple(row["daily_orders"]) for row in rows}, {(60, 40), (80, 70)})
            self.assertEqual(len({row["event_id"] for row in rows}), 2)

    def test_small_order_by_day_header_dates_must_be_consecutive(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "小订及首销数据整理.xlsx"
            workbook = Workbook()
            summary = workbook.active
            summary.title = "车型汇总"
            summary.append(["车型", "小订开始日期", "小订结束日期", "小订天数", "总小订"])
            summary.append(["测试车", datetime(2026, 9, 1), datetime(2026, 9, 3), 3, 90])
            daily = workbook.create_sheet("小订by天")
            daily.append(["小订", datetime(2026, 9, 1), datetime(2026, 9, 2), datetime(2026, 9, 4), "合计"])
            daily.append(["测试车", 30, 30, 30, 90])
            workbook.save(path)
            workbook.close()
            with patch("modules.sales_forecast._read_model_mapping", return_value={}), patch("modules.sales_forecast._read_model_master", return_value={}):
                with self.assertLogs("modules.sales_forecast", level="WARNING") as logs:
                    _, rows = _read_small_order_history(path)
        self.assertEqual(rows[0]["daily_orders"], [30, 30])
        self.assertEqual(rows[0]["dates"], ["2026-09-01", "2026-09-02"])
        message = "\n".join(logs.output)
        self.assertIn("小订by天表头日期", message)
        self.assertIn("与首日间隔不符", message)

    def test_small_order_total_prefers_summary_and_flags_mismatch(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "小订及首销数据整理.xlsx"
            workbook = Workbook()
            summary = workbook.active
            summary.title = "车型汇总"
            summary.append(["车型", "小订开始日期", "小订结束日期", "小订天数", "总小订"])
            summary.append(["测试车", datetime(2026, 9, 1), datetime(2026, 9, 2), 2, 100])
            daily = workbook.create_sheet("小订by天")
            daily.append(["小订", datetime(2026, 9, 1), datetime(2026, 9, 2), "合计"])
            daily.append(["测试车", 40, 40, 80])
            workbook.save(path)
            workbook.close()
            with patch("modules.sales_forecast._read_model_mapping", return_value={}), patch("modules.sales_forecast._read_model_master", return_value={}):
                with self.assertLogs("modules.sales_forecast", level="WARNING") as logs:
                    _, rows = _read_small_order_history(path)
        self.assertEqual(rows[0]["total"], 100)
        self.assertEqual(rows[0]["total_source"], "车型汇总")
        self.assertEqual(rows[0]["small_progress"], [.4, .8])
        self.assertEqual(rows[0]["d1_share"], .4)
        self.assertIn("车型汇总总小订100与by天合计80不一致", "\n".join(logs.output))

        with TemporaryDirectory() as directory:
            path = Path(directory) / "小订及首销数据整理.xlsx"
            workbook = Workbook()
            summary = workbook.active
            summary.title = "车型汇总"
            summary.append(["车型", "小订开始日期", "小订结束日期", "小订天数"])
            summary.append(["测试车", datetime(2026, 9, 1), datetime(2026, 9, 2), 2])
            daily = workbook.create_sheet("小订by天")
            daily.append(["小订", datetime(2026, 9, 1), datetime(2026, 9, 2), "合计"])
            daily.append(["测试车", 30, 40, 70])
            workbook.save(path)
            workbook.close()
            with patch("modules.sales_forecast._read_model_mapping", return_value={}), patch("modules.sales_forecast._read_model_master", return_value={}):
                _, rows = _read_small_order_history(path)
        self.assertEqual(rows[0]["total"], 70)
        self.assertEqual(rows[0]["total_source"], "小订by天合计")

    def test_small_order_without_total_header_uses_every_date_column(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "小订及首销数据整理.xlsx"
            workbook = Workbook()
            summary = workbook.active
            summary.title = "车型汇总"
            summary.append(["车型", "小订开始日期", "小订结束日期", "小订天数"])
            summary.append(["测试车", datetime(2026, 9, 1), datetime(2026, 9, 3), 3])
            daily = workbook.create_sheet("小订by天")
            daily.append(["小订", datetime(2026, 9, 1), datetime(2026, 9, 2), datetime(2026, 9, 3)])
            daily.append(["测试车", 50, 30, 20])
            workbook.save(path)
            workbook.close()
            with patch("modules.sales_forecast._read_model_mapping", return_value={}), patch("modules.sales_forecast._read_model_master", return_value={}):
                with self.assertLogs("modules.sales_forecast", level="WARNING") as logs:
                    _, rows = _read_small_order_history(path)
        self.assertEqual(rows[0]["daily_orders"], [50, 30, 20])
        self.assertEqual(rows[0]["total"], 100)
        self.assertIn("未提供合计列", "\n".join(logs.output))

    def test_small_order_missing_or_negative_day_rejects_the_real_curve(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "小订及首销数据整理.xlsx"
            workbook = Workbook()
            summary = workbook.active
            summary.title = "车型汇总"
            summary.append(["车型", "小订开始日期", "小订结束日期", "小订天数", "总小订"])
            summary.append(["缺失车", datetime(2026, 9, 1), datetime(2026, 9, 2), 2, 100])
            summary.append(["负数车", datetime(2026, 9, 1), datetime(2026, 9, 2), 2, 100])
            daily = workbook.create_sheet("小订by天")
            daily.append(["小订", datetime(2026, 9, 1), datetime(2026, 9, 2), "合计"])
            daily.append(["缺失车", 60, None, 100])
            daily.append(["负数车", 110, -10, 100])
            workbook.save(path)
            workbook.close()
            with patch("modules.sales_forecast._read_model_mapping", return_value={}), patch("modules.sales_forecast._read_model_master", return_value={}):
                with self.assertLogs("modules.sales_forecast", level="WARNING") as logs:
                    _, rows = _read_small_order_history(path)
        self.assertTrue(all(row["daily_actual"] is False for row in rows))
        message = "\n".join(logs.output)
        self.assertIn("缺失或非数字", message)
        self.assertIn("为负数", message)

    def test_steady_history_reads_completed_lock_weeks_after_launch(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "问界 M9 2026款by周"
        sheet.append(["指标", "统计类型", "分类", "26WK27", "26WK28", "26WK29", "26WK30"])
        sheet.append(["交车锁单", "数量", "数量", 100, 200, 300, 400])
        item = SimpleNamespace(path=Path("问界锁单选配比例分析.xlsx"), workbook=workbook)
        store = SimpleNamespace(find_all=lambda _: [item], find=lambda _: None)
        windows = {"m9": {"generation": "问界 M9 2026款", "end_date": "2026-06-30"}}
        with patch("modules.sales_forecast._read_model_mapping", return_value={}), patch(
            "modules.sales_forecast._read_model_master", return_value={}
        ):
            rows, sources = _read_steady_history(store, windows, today=date(2026, 7, 27))
        workbook.close()
        self.assertEqual([week["period"] for week in rows[0]["weeks"]], ["26WK28", "26WK29", "26WK30"])
        self.assertEqual([week["lock"] for week in rows[0]["weeks"]], [200, 300, 400])
        self.assertNotIn("direct", rows[0]["weeks"][0])
        self.assertEqual(sources[0].note, "平销期按周交车锁单历史")

    def test_steady_launch_features_only_attach_lock_evidence(self):
        history = [{"model": "测试车", "lock_rate": .72, "days": 30}]
        steady = [{"model": "测试车", "generation": "测试车"}]
        with patch("modules.sales_forecast._history_record", return_value=history[0]):
            _attach_steady_launch_features(steady, history)
        self.assertEqual(steady[0]["lock_rate"], .72)
        self.assertEqual(steady[0]["launch_days"], 30)
        self.assertNotIn("launch_tail_direct", steady[0])

    def test_stage_window_contradiction_is_logged(self):
        profiles = [{"model": "问界 M6 2026款", "launch_date": "2026-08-05", "days": [{"gross": 1}], "total_small": 10}]
        windows = {"m62026": {"generation": "问界 M6 2026款", "launch_date": "2026-09-01", "end_date": "2026-08-01", "days": 40}}
        with self.assertLogs("modules.sales_forecast", level="WARNING") as logs:
            options = _target_options([], profiles, windows, today=date(2026, 8, 27))
        message = "\n".join(logs.output)
        self.assertIn("首销窗口", message)
        self.assertIn("早于开始日期", message)
        self.assertEqual(options[0]["stage"], "unknown")
        self.assertEqual(options[0]["end_date"], "2026-08-01")

    def test_steady_week_sequence_gaps_are_logged(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "问界 M9 2026款by周"
        sheet.append(["指标", "统计类型", "分类", "26WK28", "26WK30"])
        sheet.append(["交车锁单", "数量", "数量", 700, 630])
        item = SimpleNamespace(path=Path("问界锁单选配比例分析.xlsx"), workbook=workbook)
        store = SimpleNamespace(find_all=lambda _: [item], find=lambda _: None)
        windows = {"m9": {"generation": "问界 M9 2026款", "end_date": "2026-06-30"}}
        with patch("modules.sales_forecast._read_model_mapping", return_value={}), patch(
            "modules.sales_forecast._read_model_master", return_value={}
        ):
            with self.assertLogs("modules.sales_forecast", level="WARNING") as logs:
                rows, _ = _read_steady_history(store, windows, today=date(2026, 8, 1))
        workbook.close()
        message = "\n".join(logs.output)
        self.assertIn("平销锁单周不连续", message)
        self.assertEqual(rows, [])

    def test_steady_history_excludes_launch_overlap_and_current_week(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "问界 M9 2026款by周"
        sheet.append(["指标", "统计类型", "分类", "26WK27", "26WK28", "26WK29", "26WK30"])
        sheet.append(["交车锁单", "数量", "数量", 100, 200, 300, 400])
        item = SimpleNamespace(path=Path("问界锁单选配比例分析.xlsx"), workbook=workbook)
        store = SimpleNamespace(find_all=lambda _: [item], find=lambda _: None)
        windows = {"m9": {"generation": "问界 M9 2026款", "end_date": "2026-06-30"}}
        with patch("modules.sales_forecast._read_model_mapping", return_value={}), patch(
            "modules.sales_forecast._read_model_master", return_value={}
        ):
            rows, _ = _read_steady_history(store, windows, today=date(2026, 7, 22))
        workbook.close()
        self.assertEqual([week["period"] for week in rows[0]["weeks"]], ["26WK28", "26WK29"])

    def test_steady_invalid_lock_values_are_excluded(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "问界 M9 2026款by周"
        sheet.append(["指标", "统计类型", "分类", "26WK28", "26WK29", "26WK30"])
        sheet.append(["交车锁单", "数量", "数量", 700, None, -1])
        item = SimpleNamespace(path=Path("问界锁单选配比例分析.xlsx"), workbook=workbook)
        store = SimpleNamespace(find_all=lambda _: [item], find=lambda _: None)
        windows = {"m9": {"generation": "问界 M9 2026款", "end_date": "2026-06-30"}}
        with patch("modules.sales_forecast._read_model_mapping", return_value={}), patch(
            "modules.sales_forecast._read_model_master", return_value={}
        ):
            with self.assertLogs("modules.sales_forecast", level="WARNING") as logs:
                rows, _ = _read_steady_history(store, windows, today=date(2026, 8, 1))
        workbook.close()
        self.assertEqual([week["lock"] for week in rows[0]["weeks"]], [700])
        message = "\n".join(logs.output)
        self.assertIn("26WK29", message)
        self.assertIn("26WK30", message)
        self.assertIn("该周不进入平销锁单历史", message)

    def test_forecast_dropdown_attributes_come_from_vehicle_master(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "车型基本信息.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "车型基本信息"
            sheet.append(["历史传播名", "订单分析代际名", "产品档位", "发布类型"])
            sheet.append(["传播名A", "代际A", "旗舰", "年度换代"])
            sheet.append(["传播名B", "代际B", "主流", "品牌首发"])
            workbook.save(path)
            workbook.close()
            master = _read_model_master(path)

        self.assertEqual(_master_values(master, "产品档位"), ["主流", "旗舰"])
        self.assertEqual(_master_values(master, "发布类型"), ["品牌首发", "年度换代"])

    def test_model_matching_does_not_use_vehicle_specific_aliases(self):
        self.assertFalse(_same_model("问界 F2N 2026款", "V800/V680"))

    def test_model_matching_never_drops_brand_identity(self):
        self.assertFalse(_same_model("问界 S9 2026款", "享界 S9 2026款"))

    def test_model_matching_uses_maintained_propagation_generation_mapping_both_ways(self):
        mapping = {"尊界g9": "尊界 G9 2027款", "尊界g92027": "尊界 G9 2027款"}
        with patch("modules.sales_forecast._read_model_mapping", return_value=mapping):
            self.assertTrue(_same_model("尊界 G9", "尊界 G9 2027款"))
            self.assertTrue(_same_model("尊界 G9 2027款", "尊界 G9"))

    def test_actual_sources_are_merged_by_canonical_generation_name(self):
        workbook = Workbook()
        by_day = workbook.active
        by_day.title = "尊界 G9by天"
        by_day.append(["周期", "D1"])
        by_day.append(["时间", datetime(2026, 8, 1)])
        by_day.append(["当日大定数量", 300])
        by_day.append(["当日小订转大定数量", 180])
        by_day.append(["当日直接大定数量", 120])
        by_hour = workbook.create_sheet("尊界 G9 2027款by时")
        by_hour.append(["周期", "H1"])
        by_hour.append(["时间", "2026/8/1 10:00"])
        by_hour.append(["当日大定数量", 100])
        by_hour.append(["当日小订转大定数量", 60])
        by_hour.append(["当日直接大定数量", 40])
        item = SimpleNamespace(path=Path("首销期订单节奏.xlsx"), workbook=workbook)

        class Store:
            @staticmethod
            def find(keyword):
                return item if keyword == "首销期订单节奏" else None

        mapping = {"尊界g9": "尊界 G9 2027款", "尊界g92027": "尊界 G9 2027款"}
        with patch("modules.sales_forecast._read_model_mapping", return_value=mapping):
            profiles, _ = _read_actual_profiles(Store())

        self.assertEqual(len(profiles), 1)
        self.assertEqual(profiles[0]["model"], "尊界 G9 2027款")
        self.assertEqual(profiles[0]["days"][0]["gross"], 300)
        self.assertEqual(profiles[0]["hourly_days"][0]["gross"], 100)
        workbook.close()

    def test_invalid_by_day_d1_prints_cell_type_diagnostic(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "尊界 G9by天"
        sheet.append(["周期", "D1"])
        sheet.append(["时间", datetime(2026, 8, 1)])
        sheet.append(["当日大定数量", "未计算公式"])
        item = SimpleNamespace(path=Path("首销期订单节奏.xlsx"), workbook=workbook)

        class Store:
            @staticmethod
            def find(keyword):
                return item if keyword == "首销期订单节奏" else None

        mapping = {"尊界g9": "尊界 G9 2027款", "尊界g92027": "尊界 G9 2027款"}
        with patch("modules.sales_forecast._read_model_mapping", return_value=mapping):
            with self.assertLogs("modules.sales_forecast", level="WARNING") as logs:
                _read_actual_profiles(Store())

        message = "\n".join(logs.output)
        self.assertIn("by天未构建真实日", message)
        self.assertIn("尊界 G9by天", message)
        self.assertIn("D1大定='未计算公式'(str)", message)
        workbook.close()

    def test_invalid_by_hour_timestamp_prints_diagnostic(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "尊界 G9 2027款by时"
        sheet.append(["周期", "H1"])
        sheet.append(["时间", "无法解析的时间"])
        sheet.append(["当日大定数量", 100])
        item = SimpleNamespace(path=Path("首销期订单节奏.xlsx"), workbook=workbook)

        class Store:
            @staticmethod
            def find(keyword):
                return item if keyword == "首销期订单节奏" else None

        with self.assertLogs("modules.sales_forecast", level="WARNING") as logs:
            _read_actual_profiles(Store())

        message = "\n".join(logs.output)
        self.assertIn("by时未构建分时进度", message)
        self.assertIn("首列时间='无法解析的时间'(str)", message)
        workbook.close()

    def test_negative_by_day_gross_is_rejected_and_logged(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "尊界 G9 2027款by天"
        sheet.append(["周期", "D1"])
        sheet.append(["时间", datetime(2026, 8, 1)])
        sheet.append(["当日大定数量", -10])
        item = SimpleNamespace(path=Path("首销期订单节奏.xlsx"), workbook=workbook)

        class Store:
            @staticmethod
            def find(keyword):
                return item if keyword == "首销期订单节奏" else None

        with self.assertLogs("modules.sales_forecast", level="WARNING") as logs:
            profiles, _ = _read_actual_profiles(Store())

        self.assertEqual(profiles[0]["days"], [])
        self.assertIn("负数或非有限数已拒绝", "\n".join(logs.output))
        workbook.close()

    def test_missing_vehicle_attributes_are_not_inferred_from_model_name(self):
        item = _history_item({"传播名": "问界 M9 2026款纯电"}, processed=True)
        self.assertEqual(item["tier"], "未维护")
        self.assertEqual(item["energy"], "未维护")
        self.assertEqual(item["node"], "未维护")

    def test_processed_baseline_reads_its_own_launch_days_header(self):
        item = _history_item({"传播名": "车型A", "首销天数": 70}, processed=True)
        self.assertEqual(item["days"], 70)

    def test_excel_serial_in_by_day_row_two_is_an_absolute_date(self):
        self.assertEqual(_iso(46134), "2026-04-22")

    def test_common_text_date_formats_are_supported(self):
        self.assertEqual(_iso("2026/8/20"), "2026-08-20")
        self.assertEqual(_iso("2026.08.20 00:00:00"), "2026-08-20")
        self.assertEqual(_iso("2026年8月20日"), "2026-08-20")
        self.assertEqual(_iso("26-04-22"), "2026-04-22")
        self.assertEqual(_as_datetime("2026/8/20 09:30").isoformat(), "2026-08-20T09:30:00")
        self.assertEqual(_as_datetime("2026年8月20日 10时15分").isoformat(), "2026-08-20T10:15:00")
        self.assertEqual(_as_datetime("26-04-22 19:00").isoformat(), "2026-04-22T19:00:00")

    def test_stage_uses_absolute_launch_window(self):
        self.assertEqual(_forecast_stage("2026-09-01", days=10, today=date(2026, 8, 27))["key"], "before")
        active = _forecast_stage("2026-08-05", days=65, today=date(2026, 8, 27))
        self.assertEqual(active["key"], "active")
        self.assertEqual(active["day"], 23)
        self.assertEqual(active["end_date"], "2026-10-08")
        self.assertEqual(_forecast_stage("2026-07-01", days=10, today=date(2026, 8, 27))["key"], "ended")

    def test_default_target_prefers_currently_active_launch(self):
        targets = [
            {"name": "已结束", "stage": "ended", "launch_date": "2026-07-01"},
            {"name": "进行中", "stage": "active", "calendar_day": 8, "launch_date": "2026-08-20"},
            {"name": "未开始", "stage": "before", "launch_date": "2026-09-01"},
        ]
        self.assertEqual(_default_target_option(targets)["name"], "进行中")

    def test_completed_calendar_gap_is_a_blocking_raw_data_error(self):
        target = {
            "name": "车型A", "stage": "active", "launch_date": "2026-08-01",
            "days": 10, "launch_days_maintained": True,
        }
        profile = {
            "model": "车型A", "total_small": 1000,
            "days": [
                {"date": "2026-08-01", "gross": 100, "small_to_big": 80, "direct": 20},
                {"date": "2026-08-03", "gross": 90, "small_to_big": 70, "direct": 20},
            ],
        }
        errors = _profile_hard_errors(target, profile, today=date(2026, 8, 4))
        self.assertTrue(any("D2" in error for error in errors))

    def test_complete_raw_days_pass_blocking_validation(self):
        target = {
            "name": "车型A", "stage": "active", "launch_date": "2026-08-01",
            "days": 10, "launch_days_maintained": True,
        }
        profile = {
            "model": "车型A", "total_small": 1000,
            "days": [
                {"date": "2026-08-01", "gross": 100, "small_to_big": 80, "direct": 20},
                {"date": "2026-08-02", "gross": 90, "small_to_big": 70, "direct": 20},
                {"date": "2026-08-03", "gross": 80, "small_to_big": 60, "direct": 20},
            ],
        }
        self.assertEqual(_profile_hard_errors(target, profile, today=date(2026, 8, 4)), [])

    def test_source_priority_changes_with_absolute_stage(self):
        history = []
        profiles = []
        targets = []
        for name, stage in (("车型A", "ended"), ("车型B", "before"), ("车型C", "active")):
            history.append({
                "model": name, "generation": name, "launch_date": "2026-01-01", "days": 2,
                "small": 1000, "small_to_big": 300, "gross": 500, "direct": 200,
                "lock": 400, "daily_orders": [300, 200],
                "daily_small": [180, 120], "daily_direct": [120, 80], "cancel_progress": [0.01, 0.02],
            })
            profiles.append({
                "model": name, "launch_date": "2026-01-01", "total_small": 900,
                "days": [{"day": "D1", "date": "2026-01-01", "gross": 250, "small_to_big": 150, "direct": 100, "lock": 0}],
                "hourly_days": [], "day_source": {"file": "节奏.xlsx", "sheet": f"{name}by天"},
                "cancel_source": {"file": "退订.xlsx", "sheet": f"{name}_日度退订"},
                "cancel_days": [{"date": "2025-12-31", "cancel": 10, "cancel_rate": 0.01}],
                "cancel_total_small": 1100, "cancel_latest_date": "2025-12-31",
            })
            targets.append({"name": name, "stage": stage, "small": 0})

        resolved = _resolve_actual_profiles(history, profiles, targets, SourceRef("整理.xlsx", "预测基准总表", "历史"))
        self.assertEqual([item["selected_source"] for item in resolved], ["history", "cancel", "launch"])
        self.assertEqual([target["small"] for target in targets], [1000, 1100, 900])

    def test_active_and_before_stages_do_not_use_a_third_out_of_stage_source(self):
        history = [{
            "model": "车型A", "generation": "车型A", "small": 1000, "gross": 500,
            "daily_orders": [500], "daily_small": [300], "daily_direct": [200],
        }]
        profile = {
            "model": "车型A", "days": [], "hourly_days": [], "total_small": 0,
            "cancel_days": [], "cancel_total_small": 0,
        }
        before = {"name": "车型A", "stage": "before", "small": 0}
        active = {"name": "车型A", "stage": "active", "small": 0}

        resolved = _resolve_actual_profiles(
            history, [profile], [before, active], SourceRef("整理.xlsx", "预测基准总表", "历史")
        )

        self.assertEqual(resolved[0]["selected_source"], "history")
        self.assertEqual(resolved[1]["selected_source"], "history")
        self.assertNotIn("cancel", resolved[1]["stage_profiles"]["active"]["priority"])
        self.assertNotIn("launch", resolved[0]["stage_profiles"]["before"]["priority"])

    def test_day_merge_never_matches_different_absolute_dates_by_row_number(self):
        candidates = {
            "launch": {"days": [{"day": "D1", "date": "2026-08-20", "gross": 100}]},
            "cancel": {"days": [{"day": "D1", "date": "2026-07-01", "cancel": 999}]},
        }

        rows, _ = _merge_day_fields(candidates, ("launch", "cancel"))

        self.assertEqual(rows[0]["gross"], 100)
        self.assertIsNone(rows[0].get("cancel"))

    def test_day_merge_does_not_leave_invalid_base_values_behind(self):
        candidates = {
            "launch": {"days": [{"day": "D1", "date": "2026-08-20", "gross": -10}]},
            "history": {"days": [{"day": "D1", "date": "2026-08-20", "gross": -20}]},
        }

        rows, _ = _merge_day_fields(candidates, ("launch", "history"))

        self.assertIsNone(rows[0]["gross"])

    def test_negative_completed_values_are_blocking_errors(self):
        target = {
            "name": "车型A", "stage": "active", "launch_date": "2026-08-01",
            "days": 10, "launch_days_maintained": True,
        }
        profile = {
            "model": "车型A", "total_small": 1000,
            "days": [
                {"date": "2026-08-01", "gross": -100, "small_to_big": -80, "direct": -20},
            ],
        }

        errors = _profile_hard_errors(target, profile, today=date(2026, 8, 2))

        self.assertTrue(any("字段缺失或校验失败" in error for error in errors))

    def test_transition_and_post_launch_stages_keep_distinct_total_small_priorities(self):
        history = [{
            "model": "车型A", "generation": "车型A", "launch_date": "2026-01-01", "days": 2,
            "small": 1000, "small_to_big": 300, "gross": 500, "direct": 200,
            "daily_orders": [300, 200], "daily_small": [180, 120], "daily_direct": [120, 80],
        }]
        profile = {
            "model": "车型A", "launch_date": "2026-01-01",
            "launch_total_small": 887, "launch_total_small_valid": False,
            "days": [{"day": "D1", "date": "2026-01-01", "gross": 250, "small_to_big": 150, "direct": 100}],
            "hourly_days": [], "day_source": {"file": "节奏.xlsx", "sheet": "车型Aby天"},
            "cancel_source": {"file": "退订.xlsx", "sheet": "车型A_日度退订"},
            "cancel_days": [{"date": "2025-12-31", "cancel": 10, "cancel_rate": 0.01}],
            "cancel_total_small": 1100, "cancel_total_small_valid": True,
        }
        targets = [
            {"name": "车型A", "stage": "ended", "small": 0},
            {"name": "车型A", "stage": "before", "small": 0},
            {"name": "车型A", "stage": "active", "small": 0},
        ]
        resolved = _resolve_actual_profiles(history, [profile], targets, SourceRef("整理.xlsx", "预测基准总表", "历史"))
        self.assertEqual([target["small"] for target in targets], [1000, 1100, 1000])
        self.assertEqual(
            [target["field_sources"]["总小订"] for target in targets],
            ["小订及首销数据整理", "小订退订分析", "小订及首销数据整理"],
        )
        self.assertEqual(resolved[0]["stage_profiles"]["ended"]["priority"], ("history", "launch", "cancel"))
        self.assertEqual(resolved[1]["stage_profiles"]["before"]["priority"], ("cancel", "history"))
        self.assertEqual(resolved[2]["selected_source"], "launch")

    def test_blank_ended_history_day_fields_fall_back_to_launch_by_field(self):
        history = [{
            "model": "车型A", "generation": "车型A", "launch_date": "2026-01-01", "days": 1,
            "small": 1000, "gross": 500,
            "daily_orders": [None], "daily_small": [None], "daily_direct": [None],
        }]
        profile = {
            "model": "车型A", "launch_date": "2026-01-01", "total_small": 900,
            "days": [{
                "day": "D1", "date": "2026-01-01", "gross": 250,
                "small_to_big": 150, "direct": 100, "lock": 80,
            }],
            "hourly_days": [], "day_source": {"file": "节奏.xlsx", "sheet": "车型Aby天"},
        }
        target = {"name": "车型A", "stage": "ended", "small": 0}

        resolved = _resolve_actual_profiles(
            history, [profile], [target], SourceRef("整理.xlsx", "预测基准总表", "历史")
        )

        self.assertEqual(resolved[0]["days"][0]["gross"], 250)
        self.assertEqual(resolved[0]["days"][0]["small_to_big"], 150)
        self.assertEqual(resolved[0]["days"][0]["direct"], 100)
        self.assertEqual(target["field_sources"]["首销日·gross"], "首销期订单节奏")

    def test_active_generation_reanchors_mapped_secondary_d1_to_current_launch_date(self):
        history = [{
            "model": "尊界 G9", "generation": "尊界 G9 2027款",
            "launch_date": "", "days": 35, "small": 1000, "gross": 500,
            "daily_orders": [300], "daily_small": [180], "daily_direct": [120],
        }]
        profile = {
            "model": "尊界 G9", "launch_date": "2026-08-20", "total_small": 0,
            "days": [], "hourly_days": [],
        }
        target = {
            "name": "尊界 G9 2027款", "stage": "active",
            "launch_date": "2026-08-20", "small": 0,
        }
        mapping = {"尊界g9": "尊界 G9 2027款", "尊界g92027": "尊界 G9 2027款"}

        with patch("modules.sales_forecast._read_model_mapping", return_value=mapping):
            resolved = _resolve_actual_profiles(
                history, [profile], [target], SourceRef("二次处理.xlsx", "预测基准总表", "历史")
            )

        self.assertEqual(resolved[0]["days"][0]["date"], "2026-08-20")
        self.assertEqual(resolved[0]["days"][0]["gross"], 300)
        self.assertEqual(target["field_sources"]["首销日·gross"], "小订及首销数据整理")

    def test_active_launch_rows_with_blank_dates_are_anchored_to_target_d1(self):
        profile = {
            "model": "享界 G9 2027款", "launch_date": "", "total_small": 0,
            "days": [{
                "day": "D1", "date": "", "gross": 2552,
                "small_to_big": 2412, "direct": 140, "lock": 0,
            }],
            "hourly_days": [],
            "day_source": {"file": "首销节奏.xlsx", "sheet": "享界 G9 2027款by天"},
        }
        target = {
            "name": "享界 G9 2027款", "stage": "active",
            "launch_date": "2026-08-20", "small": 0,
        }

        resolved = _resolve_actual_profiles(
            [], [profile], [target], SourceRef("二次处理.xlsx", "预测基准总表", "历史")
        )

        self.assertEqual(resolved[0]["days"][0]["date"], "2026-08-20")
        self.assertEqual(resolved[0]["days"][0]["gross"], 2552)
        self.assertEqual(target["field_sources"]["首销日·gross"], "首销期订单节奏")

    def test_missing_active_d1_prints_mapping_and_three_source_diagnostics(self):
        target = {
            "name": "尊界 G9 2027款", "stage": "active",
            "launch_date": "2026-08-20", "small": 0,
        }
        mapping = {"尊界g9": "尊界 G9 2027款", "尊界g92027": "尊界 G9 2027款"}
        with patch("modules.sales_forecast._read_model_mapping", return_value=mapping):
            with self.assertLogs("modules.sales_forecast", level="WARNING") as logs:
                _resolve_actual_profiles([], [], [target], SourceRef("二次处理.xlsx", "预测基准总表", "历史"))

        message = "\n".join(logs.output)
        self.assertIn("网页将判定D1缺失", message)
        self.assertIn("预测对象=尊界 G9 2027款", message)
        self.assertIn("三来源D1", message)
        self.assertIn("字段来源={}", message)

    def test_implied_total_requires_a_stable_denominator(self):
        total, valid, _ = _stable_implied_total([10, 20, 30], [.01, .02, .03])
        self.assertEqual(total, 1000)
        self.assertTrue(valid)
        _, valid, reason = _stable_implied_total([10, 20, 185], [.18, .18, .208568])
        self.assertFalse(valid)
        self.assertIn("波动", reason)

    def test_generation_window_from_summary_overrides_incomplete_launch_sheet(self):
        profiles = [{"model": "问界 M6 2026款", "launch_date": "2026-08-05", "days": [{"gross": 1}], "total_small": 10}]
        windows = {"m62026": {"generation": "问界 M6 2026款", "launch_date": "2026-09-01", "end_date": "2026-10-10", "days": 40}}
        options = _target_options([], profiles, windows, today=date(2026, 8, 27))
        self.assertEqual(len(options), 1)
        self.assertEqual(options[0]["name"], "问界 M6 2026款")
        self.assertEqual(options[0]["launch_date"], "2026-09-01")
        self.assertEqual(options[0]["stage"], "before")
        self.assertEqual(options[0]["date_source_label"], "小订及首销数据整理 · 车型汇总")

    def test_inferred_total_small_is_flagged_as_estimate(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "尊界 G9by天"
        sheet.append(["周期", "D1", "D2", "D3"])
        sheet.append(["时间", datetime(2026, 8, 1), datetime(2026, 8, 2), datetime(2026, 8, 3)])
        sheet.append(["当日大定数量", 100, 100, 100])
        sheet.append(["累计小订转大定数量", 100, 200, 300])
        sheet.append(["累计小订转化率", .1, .2, .3])
        item = SimpleNamespace(path=Path("首销期订单节奏.xlsx"), workbook=workbook)

        class Store:
            @staticmethod
            def find(keyword):
                return item if keyword == "首销期订单节奏" else None

        mapping = {"尊界g9": "尊界 G9", "尊界g92027": "尊界 G9"}
        with patch("modules.sales_forecast._read_model_mapping", return_value=mapping):
            profiles, _ = _read_actual_profiles(Store())
            with self.assertLogs("modules.sales_forecast", level="WARNING") as logs:
                _resolve_actual_profiles(
                    [], profiles,
                    [{"name": "尊界 G9", "stage": "active", "launch_date": "2026-08-01", "history_small": 0}],
                    SourceRef("历史.xlsx", "历史", "历史"),
                )

        self.assertTrue(profiles[0]["launch_total_small_valid"])
        self.assertEqual(profiles[0]["total_small"], 1000)
        message = "\n".join(logs.output)
        self.assertIn("[销量预测估算]", message)
        self.assertIn("总小订=1000为首销累计进度反推值", message)
        self.assertIn("累计进度反推值", message)
        workbook.close()

    def test_cancel_side_total_small_estimate_is_logged(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "尊界 G9_日度退订"
        sheet.append(["占位", "占位", "占位"])
        sheet.append(["日期", "累计小订退", "整体 - 小订退 %"])
        sheet.append([datetime(2026, 8, 1), 100, .1])
        sheet.append([datetime(2026, 8, 2), 200, .2])
        sheet.append([datetime(2026, 8, 3), 300, .3])
        item = SimpleNamespace(path=Path("小订退订分析.xlsx"), workbook=workbook)

        class Store:
            @staticmethod
            def find(keyword):
                return item if keyword == "小订退订分析" else None

        mapping = {"尊界g9": "尊界 G9", "尊界g92027": "尊界 G9"}
        with patch("modules.sales_forecast._read_model_mapping", return_value=mapping):
            profiles, _ = _read_actual_profiles(Store())
            with self.assertLogs("modules.sales_forecast", level="WARNING") as logs:
                _resolve_actual_profiles(
                    [], profiles,
                    [{"name": "尊界 G9", "stage": "before", "launch_date": "2026-09-01", "history_small": 0}],
                    SourceRef("历史.xlsx", "历史", "历史"),
                )

        self.assertTrue(profiles[0]["cancel_total_small_valid"])
        self.assertEqual(profiles[0]["cancel_total_small"], 1000)
        message = "\n".join(logs.output)
        self.assertIn("[销量预测估算]", message)
        self.assertIn("总小订=1000为累计退订/退订率反推值", message)
        workbook.close()

    def test_conflicting_explicit_total_small_is_logged_not_silently_dropped(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "尊界 G9by天"
        sheet.append(["周期", "D1", "D2"])
        sheet.append(["时间", datetime(2026, 8, 1), datetime(2026, 8, 2)])
        sheet.append(["当日大定数量", 100, 100])
        sheet.append(["累计小订转大定数量", 150, 300])
        sheet.append(["总小订", 200])
        item = SimpleNamespace(path=Path("首销期订单节奏.xlsx"), workbook=workbook)

        class Store:
            @staticmethod
            def find(keyword):
                return item if keyword == "首销期订单节奏" else None

        mapping = {"尊界g9": "尊界 G9", "尊界g92027": "尊界 G9"}
        with patch("modules.sales_forecast._read_model_mapping", return_value=mapping):
            with self.assertLogs("modules.sales_forecast", level="WARNING") as logs:
                profiles, _ = _read_actual_profiles(Store())

        self.assertFalse(profiles[0]["launch_total_small_valid"])
        message = "\n".join(logs.output)
        self.assertIn("[销量预测估算]", message)
        self.assertIn("显式总小订200小于累计小转大300", message)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
