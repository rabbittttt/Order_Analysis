from __future__ import annotations

import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.utils.datetime import to_excel

from core.excel import WorkbookItem, WorkbookStore, _quantity_field
from core.models import Subject
from modules import channel_analysis as channel


MODEL = "问界 M9 2026款"
A, B = "华为授权体验店", "问界用户中心"


class ChannelAnalysisTests(unittest.TestCase):
    def book(self, periods, rows, title=MODEL + "by天"):
        book = Workbook()
        book.active.title = title
        book.active.append(["指标", "统计类型", "分类", *periods])
        for row in rows:
            book.active.append(row)
        return book

    def store(self, *books):
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path(f"大定选配比例分析_{index}.xlsx"), book)
                       for index, book in enumerate(books)]
        self.addCleanup(store.close)
        return store

    def build(self, store, name=MODEL, kind="generation", parent="问界"):
        return channel.ChannelAnalysisModule().build(store, Subject(name, name, kind, parent))

    def row(self, dashboard, grain="day", period=None, index=0):
        view = dashboard.views[grain]
        return view["pages"][period or view["default_period"]]["sections"][0]["data"]["rows"][index]

    def test_original_share_formulas_and_absolute_changes(self):
        book = self.book(["2026-01-02", "2026-01-01"], [
            ["门店类型", "数量", A, 40, 20],
            [None, None, B, 10, 0],
            [None, "占比", A, .123, .10],
            [None, None, B, "87.7%", "90%"],
            ["贡献订单门店数", "数量", A, 4, 2],
            [None, None, B, 1, 0],
        ])
        with patch.dict(channel.CHANNEL_STORE_COUNTS, {"问界": {A: 20, B: 10}}, clear=True):
            dashboard = self.build(self.store(book))
        row = self.row(dashboard)
        self.assertEqual(row["count"], 40)
        self.assertEqual(row["share"], .123)  # Never recompute the source share as 40 / 50.
        self.assertEqual(row["store_count"], 20)
        self.assertEqual(row["active_rate"], .2)
        self.assertEqual(row["productivity"], 2)
        self.assertEqual(row["count_change"], 20)
        self.assertAlmostEqual(row["share_change"], .023)
        self.assertEqual(row["active_rate_change"], .1)
        self.assertEqual(row["productivity_change"], 1)
        self.assertEqual(self.row(dashboard, period="2026-01-01")["previous_count"], None)
        self.assertEqual(self.row(dashboard, index=1)["share"], .877)
        page = dashboard.views["day"]["pages"]["2026-01-02"]
        self.assertEqual([item["value"] for item in page["kpis"][:2]], [50, 5])
        self.assertAlmostEqual(page["kpis"][2]["value"], 5 / 30 * 100)
        self.assertAlmostEqual(page["kpis"][3]["value"], 50 / 30)
        totals = page["sections"][0]["data"]["totals"]
        self.assertEqual(totals["store_count"], 30)
        self.assertEqual(totals["count"], 50)
        self.assertEqual(totals["active_count"], 5)
        self.assertAlmostEqual(totals["active_rate"], 5 / 30)
        self.assertAlmostEqual(totals["productivity"], 50 / 30)
        self.assertIs(totals, page["sections"][0]["data"]["history"][-1]["totals"])
        self.assertNotIn("active_share", row)
        for block in page["sections"]:
            self.assertEqual(block["kind"], "channel_panel")
            self.assertEqual(block["source"]["file"], "大定选配比例分析_0.xlsx")
            self.assertEqual(block["source"]["sheet"], MODEL + "by天")
            self.assertIs(block["data"]["rows"], page["sections"][0]["data"]["rows"])
        previous_page = dashboard.views["day"]["pages"]["2026-01-01"]
        self.assertIs(page["sections"][0]["data"]["history"], previous_page["sections"][1]["data"]["history"])

    def test_missing_invalid_and_genuine_zero_are_distinct(self):
        periods = ["2026-01-01", "2026-01-02", "2026-01-03", "2026-01-04", "2026-01-05"]
        book = self.book(periods, [
            ["门店类型", "数量", A, 0, None, "错误", -2, 10],
            [None, "占比", A, 0, "—", -0.1, "NaN", .5],
            ["贡献订单门店数", "数量", A, 0, None, "错误", -2, 3],
        ])
        with patch.dict(channel.CHANNEL_STORE_COUNTS, {"问界": {A: 20}}, clear=True):
            dashboard = self.build(self.store(book))
        zero = self.row(dashboard, period=periods[0])
        self.assertEqual([zero[key] for key in ("count", "share", "active_count", "active_rate", "productivity")], [0] * 5)
        for period in periods[1:4]:
            row = self.row(dashboard, period=period)
            self.assertTrue(all(row[key] is None for key in ("count", "share", "active_count", "active_rate", "productivity")))
        self.assertIsNone(self.row(dashboard)["count_change"])

    def test_source_channels_are_not_classified_and_unknown_config_is_missing(self):
        strange = "新 AITO 直营门店"
        book = self.book(["2026-01-01"], [
            ["门店类型", "占比", strange, .4],
            [None, None, A, .6],
            [None, "数量", strange, 40],
            [None, None, A, 60],
            ["贡献订单门店数", "数量", strange, 4],
            [None, None, A, 6],
        ])
        dashboard = self.build(self.store(book))
        row = self.row(dashboard)
        self.assertEqual(row["channel"], strange)
        self.assertEqual(row["active_count"], 4)
        self.assertIsNone(row["store_count"])
        self.assertIsNone(row["active_rate"])
        self.assertIsNone(row["productivity"])
        page = dashboard.views["day"]["pages"]["2026-01-01"]
        self.assertIsNone(page["kpis"][2]["value"])
        self.assertIsNone(page["kpis"][3]["value"])

    def test_configuration_priority_is_subject_then_parent_then_brand(self):
        book = self.book(["2026-01-01"], [["门店类型", "数量", A, 40], [None, None, B, 20]])
        with patch.dict(channel.CHANNEL_STORE_COUNTS, {"问界": {A: 100, B: 50}, "智界": {A: 129}}, clear=True), patch.dict(
            channel.CHANNEL_STORE_COUNTS_BY_SUBJECT, {"问界": {A: 80, B: 20}, MODEL: {A: 40}}, clear=True
        ):
            dashboard = self.build(self.store(book))
            self.assertEqual(self.row(dashboard)["store_count"], 40)
            self.assertEqual(self.row(dashboard, index=1)["store_count"], 20)
            self.assertEqual(channel._store_counts(Subject("other", "智界 S7 2026款", "generation", "智界"))[A], 129)
            self.assertNotIn(A, channel._store_counts(Subject("z", "尊界 S800 2026款", "generation", "尊界")))

    def test_categories_only_follow_store_type_block_and_keep_source_order(self):
        source_name, legacy_name = "AITO授权用户中心", "旧版问界中心"
        book = self.book(["2026-01-01"], [
            ["贡献订单门店数", "数量", legacy_name, 99],
            [None, None, source_name, 3],
            [None, None, A, 2],
            ["门店类型", "占比", source_name, .37],
            [None, None, A, .63],
            [None, "数量", source_name, 30],
            [None, None, A, 70],
        ])
        with patch.dict(channel.CHANNEL_STORE_COUNTS, {
            "问界": {A: 20, source_name: 10, "仅配置未出现在原表": 5}
        }, clear=True):
            data = self.build(self.store(book)).views["day"]["pages"]["2026-01-01"]["sections"][0]["data"]
        self.assertEqual([row["channel"] for row in data["rows"]], [source_name, A])
        self.assertEqual([row["active_count"] for row in data["rows"]], [3, 2])
        self.assertEqual([row["share"] for row in data["rows"]], [.37, .63])
        self.assertEqual(data["totals"]["count"], 100)
        self.assertEqual(data["totals"]["active_count"], 5)
        self.assertIsNone(self.build(self.store(self.book(
            ["2026-01-01"], [["贡献订单门店数", "数量", legacy_name, 99]]
        ))))

    def test_zero_denominator_and_impossible_active_counts_keep_original_data(self):
        for total, count, active in [(0, 10, 2), (2, 10, 3)]:
            with self.subTest(total=total, count=count, active=active):
                book = self.book(["2026-01-01"], [["门店类型", "数量", A, count], ["贡献订单门店数", "数量", A, active]])
                with patch.dict(channel.CHANNEL_STORE_COUNTS, {"问界": {A: total}}, clear=True), self.assertLogs(channel.LOGGER, level="WARNING"):
                    dashboard = self.build(self.store(book))
                row = self.row(dashboard)
                self.assertEqual(row["count"], count)
                self.assertEqual(row["active_count"], active)
                self.assertIsNone(row["active_rate"])
                self.assertEqual(row["productivity"], count / total if total else None)

    def test_active_store_denominator_is_configured_stores_not_retained_orders(self):
        book = self.book(["2026-01-01"], [["门店类型", "数量", A, 2], ["贡献订单门店数", "数量", A, 3]])
        with patch.dict(channel.CHANNEL_STORE_COUNTS, {"问界": {A: 20}}, clear=True):
            row = self.row(self.build(self.store(book)))
        self.assertEqual(row["active_rate"], 3 / 20)

    def test_brand_store_counts_match_business_configuration(self):
        expected = {
            "问界": {"问界用户中心": 374, "鸿蒙智行用户中心": 28, "华为旗舰店": 11,
                   "华为授权体验店": 131, "华为智能生活馆": 425},
            "智界": {"智界用户中心": 48, "鸿蒙智行用户中心": 334, "华为旗舰店": 11,
                   "华为授权体验店": 129, "华为智能生活馆": 388},
            "尊界": {"鸿蒙智行用户中心": 157, "华为旗舰店": 7, "华为授权体验店": 3,
                   "华为智能生活馆": 35, "尊界用户中心": 9, "智界用户中心": 1, "问界用户中心": 2},
            "尚界": {"尚界用户中心": 221, "鸿蒙智行用户中心": 259, "华为旗舰店": 13,
                   "华为授权体验店": 153, "华为智能生活馆": 450},
            "享界": {"享界用户中心": 14, "鸿蒙智行用户中心": 344, "华为旗舰店": 11,
                   "华为授权体验店": 100, "华为智能生活馆": 387},
        }
        for brand, counts in expected.items():
            with self.subTest(brand=brand):
                self.assertEqual(channel._store_counts(Subject(brand, brand, "brand", "鸿蒙智行")), counts)
                self.assertEqual(channel._store_counts(Subject(brand, brand + "车型", "generation", brand)), counts)
        self.assertNotIn("综合店", channel._store_counts(Subject("z", "尊界", "brand")))
        self.assertNotIn("尊界体验中心", channel._store_counts(Subject("z", "尊界", "brand")))

    def test_group_sums_brand_channel_counts(self):
        counts = channel._store_counts(Subject("g", "鸿蒙智行", "group"))
        self.assertEqual(counts, {
            "问界用户中心": 376, "智界用户中心": 49, "尚界用户中心": 221, "享界用户中心": 14,
            "鸿蒙智行用户中心": 1122, "华为旗舰店": 53, "华为授权体验店": 516,
            "华为智能生活馆": 1685, "尊界用户中心": 9,
        })
        book = self.book(["2026-01-01"], [["门店类型", "数量", A, 10], [None, "贡献门店数", A, 2]], "鸿蒙智行by天")
        dashboard = self.build(self.store(book), "鸿蒙智行", "group", None)
        row = self.row(dashboard)
        self.assertEqual(row["store_count"], 516)
        self.assertEqual(row["active_count"], 2)
        self.assertEqual(row["active_rate"], 2 / 516)
        self.assertEqual(row["productivity"], 10 / 516)
        totals = dashboard.views["day"]["pages"]["2026-01-01"]["sections"][0]["data"]["totals"]
        self.assertEqual(totals["store_count"], 4045)
        self.assertEqual(totals["active_rate"], 2 / 4045)
        self.assertEqual(totals["productivity"], 10 / 4045)

    def test_group_uses_brand_overrides_without_adding_models_or_group_overrides(self):
        with patch.dict(channel.CHANNEL_STORE_COUNTS, {"问界": {A: 100, B: 50}, "智界": {A: 129}}, clear=True), patch.dict(
            channel.CHANNEL_STORE_COUNTS_BY_SUBJECT,
            {"问界": {A: 80}, MODEL: {A: 40}, "鸿蒙智行": {A: 999}}, clear=True
        ):
            self.assertEqual(channel._store_counts(Subject("g", "鸿蒙智行", "group")), {A: 209, B: 50})

    def test_group_preserves_unknown_and_zero_configurations(self):
        for missing in (None, "错误", -1):
            with self.subTest(missing=missing), patch.dict(channel.CHANNEL_STORE_COUNTS, {
                "问界": {A: 100, B: 0}, "智界": {A: missing, B: 0, "智界用户中心": 48}
            }, clear=True):
                counts = channel._store_counts(Subject("g", "鸿蒙智行", "group"))
                self.assertIsNone(counts[A])
                self.assertEqual(counts[B], 0)
                self.assertEqual(counts["智界用户中心"], 48)

    def test_totals_use_all_configured_channels_even_when_source_has_only_one(self):
        book = self.book(["2026-01-01"], [["门店类型", "数量", A, 10], [None, "贡献门店数", A, 2]])
        with patch.dict(channel.CHANNEL_STORE_COUNTS, {"问界": {A: 20, B: 80}}, clear=True):
            dashboard = self.build(self.store(book))
        data = dashboard.views["day"]["pages"]["2026-01-01"]["sections"][0]["data"]
        self.assertEqual(data["rows"][0]["active_rate"], 2 / 20)
        self.assertEqual(data["totals"]["store_count"], 100)
        self.assertEqual(data["totals"]["active_rate"], 2 / 100)
        self.assertEqual(data["totals"]["productivity"], 10 / 100)

    def test_new_store_format_takes_priority_including_blank_and_zero(self):
        periods = ["2026-01-01", "2026-01-02", "2026-01-03"]
        for new_first in (False, True):
            with self.subTest(new_first=new_first):
                store_rows = [["门店类型", "贡献门店数", A, 3, None, 0],
                              ["贡献订单门店数", "数量", A, 8, 9, 10]]
                if not new_first:
                    store_rows.reverse()
                book = self.book(periods, [["门店类型", "数量", A, 20, 20, 20], *store_rows])
                dashboard = self.build(self.store(book))
                self.assertEqual([self.row(dashboard, period=period)["active_count"] for period in periods], [3, None, 0])

    def test_mixed_format_files_preserve_old_history_and_new_explicit_blank(self):
        old = self.book(["2026-01-01", "2026-01-02"], [
            ["门店类型", "数量", A, 20, 20], ["贡献订单门店数", "数量", A, 4, 5]])
        new = self.book(["2026-01-02", "2026-01-03"], [
            ["门店类型", "数量", A, 20, 20], [None, "贡献门店数", A, None, 6]])
        for books in ((old, new), (new, old)):
            with self.subTest(new_first=books[0] is new):
                dashboard = self.build(self.store(*books))
                self.assertEqual([self.row(dashboard, period=period)["active_count"] for period in dashboard.views["day"]["periods"]], [4, None, 6])

    def test_new_format_preserves_missing_and_zero_store_counts(self):
        book = self.book(["2026-01-01", "2026-01-02", "2026-01-03"], [
            ["门店类型", "数量", A, 0, 10, 20], [None, "贡献门店数", A, 0, "—", -3]])
        dashboard = self.build(self.store(book))
        self.assertEqual(self.row(dashboard, period="2026-01-01")["active_count"], 0)
        self.assertEqual(self.row(dashboard, period="2026-01-01")["active_rate"], 0)
        self.assertIsNone(self.row(dashboard, period="2026-01-02")["active_count"])
        self.assertIsNone(self.row(dashboard, period="2026-01-03")["active_count"])

    def test_day_week_month_sorting_and_aggregate_exclusion(self):
        book = self.book([date(2026, 1, 3), "2026/1/1", to_excel(date(2026, 1, 2)), "总计", "近7天"],
                         [["门店类型", "数量", A, 30, 10, 20, 60, 60]])
        week = book.create_sheet(MODEL + "by周")
        week.append(["指标", "统计类型", "分类", "2026WK03", "26WK01", "汇总"])
        week.append(["门店类型", "数量", A, 30, 10, 40])
        month = book.create_sheet(MODEL + "by月")
        month.append(["指标", "统计类型", "分类", "2026-10", "2026-02", "累计"])
        month.append(["门店类型", "数量", A, 30, 10, 40])
        dashboard = self.build(self.store(book))
        self.assertEqual(dashboard.views["day"]["periods"], ["2026-01-01", "2026-01-02", "2026-01-03"])
        self.assertEqual(dashboard.views["week"]["periods"], ["26WK01", "26WK03"])
        self.assertEqual(dashboard.views["month"]["periods"], ["2026-02", "2026-10"])
        for view in dashboard.views.values():
            self.assertEqual(view["default_period"], view["periods"][-1])

    def test_subjects_read_only_their_own_sheet(self):
        book = self.book(["2026-01-01"], [["门店类型", "数量", A, 10]])
        for name, value in [("问界", 200), ("智界", 300), ("鸿蒙智行", 500)]:
            sheet = book.create_sheet(name + "by天")
            sheet.append(["指标", "统计类型", "分类", "2026-01-01"])
            sheet.append(["门店类型", "数量", A, value])
        store = self.store(book)
        for name, kind, expected in [(MODEL, "generation", 10), ("问界", "brand", 200), ("智界", "brand", 300), ("鸿蒙智行", "group", 500)]:
            self.assertEqual(self.row(self.build(store, name, kind))["count"], expected)
        self.assertIsNone(self.build(store, "问界 M5 2026款"))

    def test_no_channel_data_returns_none_for_disabled_navigation(self):
        for rows in [[], [["大定", "数量", "数量", 10]], [["门店类型", "数量", A, None]]]:
            self.assertIsNone(self.build(self.store(self.book(["2026-01-01"], rows))))

    def test_multi_source_missing_channels_never_become_zero(self):
        first = self.book(["2026-01-01"], [["门店类型", "数量", A, 10]])
        second = self.book(["2026-01-02"], [["门店类型", "数量", B, 20]])
        store = self.store(first, second)
        dashboard = self.build(store)
        self.assertIsNone(self.row(dashboard)["count"])
        self.assertEqual(self.row(dashboard, index=1)["count"], 20)
        self.assertIsNone(self.row(dashboard, period="2026-01-01", index=1)["count"])
        store.resolve_dashboard_sources(dashboard)
        self.assertEqual(len(dashboard.sources), 2)
        self.assertEqual(len(dashboard.views["day"]["pages"]["2026-01-02"]["sections"][0]["source"]["source_files"]), 2)

    def test_boundary_week_distinct_stores_cannot_sum_or_keep_partial_values(self):
        self.assertFalse(_quantity_field(("贡献订单门店数", "数量", A), "0"))
        self.assertFalse(_quantity_field(("门店类型", "贡献门店数", A), "0"))
        self.assertTrue(_quantity_field(("门店类型", "数量", A), "0"))
        for period, ranges in [("26WK53", [(date(2026, 12, 28), date(2026, 12, 31)), (date(2027, 1, 1), date(2027, 1, 3))]),
                               ("26WK27", [(date(2026, 6, 29), date(2026, 6, 30)), (date(2026, 7, 1), date(2026, 7, 5))])]:
            for active_values in [(4, 3), (4, 4), (4, None), (4, 3, 7)]:
                with self.subTest(period=period, active_values=active_values):
                    books = [self.book([period], [["门店类型", "占比", A, 1], [None, "数量", A, 10 + index],
                                                  ["贡献订单门店数", "数量", A, active]], MODEL + "by周")
                             for index, active in enumerate(active_values)]
                    store = self.store(*books)
                    for index, item in enumerate(store.items):
                        item._boundary_date_range = ranges[min(index, 1)]
                    with self.assertLogs("core.excel", level="WARNING") as logs:
                        dashboard = self.build(store)
                    row = self.row(dashboard, grain="week")
                    self.assertEqual(row["count"], 21)
                    self.assertEqual(row["share"], 1)
                    self.assertIsNone(row["active_count"])
                    self.assertIsNone(row["active_rate"])
                    self.assertIn("无法去重", "\n".join(logs.output))

    def test_boundary_week_missing_entire_store_row_cannot_keep_other_shard(self):
        for missing_side in (0, 1):
            for has_other_channel in (False, True):
                with self.subTest(missing_side=missing_side, has_other_channel=has_other_channel):
                    books = []
                    for index in range(2):
                        rows = [["门店类型", "数量", A, 10 + index]]
                        if index != missing_side:
                            rows.append(["贡献订单门店数", "数量", A, 4])
                        elif has_other_channel:
                            rows.append(["贡献订单门店数", "数量", B, 2])
                        books.append(self.book(["26WK27"], rows, MODEL + "by周"))
                    store = self.store(*books)
                    for item, span in zip(store.items, [(date(2026, 6, 29), date(2026, 6, 30)),
                                                       (date(2026, 7, 1), date(2026, 7, 5))]):
                        item._boundary_date_range = span
                    with self.assertLogs("core.excel", level="WARNING") as logs:
                        dashboard = self.build(store)
                    row = self.row(dashboard, grain="week")
                    self.assertEqual(row["count"], 21)
                    self.assertIsNone(row["active_count"])
                    self.assertIsNone(row["active_rate"])
                    self.assertIn("无法去重", "\n".join(logs.output))

    def test_boundary_week_new_and_mixed_format_store_counts_are_nonadditive(self):
        for labels in [(("门店类型", "贡献门店数"), ("门店类型", "贡献门店数")),
                       (("贡献订单门店数", "数量"), ("门店类型", "贡献门店数")),
                       (("门店类型", "贡献门店数"), ("贡献订单门店数", "数量")),
                       (("门店类型", "贡献门店数"), None)]:
            with self.subTest(labels=labels):
                books = []
                for index, label in enumerate(labels):
                    rows = [["门店类型", "数量", A, 10 + index]]
                    if label:
                        rows.append([*label, A, 4])
                    books.append(self.book(["26WK27"], rows, MODEL + "by周"))
                store = self.store(*books)
                for item, span in zip(store.items, [(date(2026, 6, 29), date(2026, 6, 30)),
                                                   (date(2026, 7, 1), date(2026, 7, 5))]):
                    item._boundary_date_range = span
                with self.assertLogs("core.excel", level="WARNING"):
                    dashboard = self.build(store)
                row = self.row(dashboard, grain="week")
                self.assertEqual(row["count"], 21)
                self.assertIsNone(row["active_count"])
                self.assertIsNone(row["active_rate"])


if __name__ == "__main__":
    unittest.main()
