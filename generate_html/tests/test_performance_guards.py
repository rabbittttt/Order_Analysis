from __future__ import annotations

import os
import unittest
from copy import copy
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill, Protection
from openpyxl.utils.indexed_list import IndexedList
from tools import refresh_sales_forecast_data as refresh

from core.excel import parse_metric_sheet
from modules.option_fee import _option_layout
from tools.refresh_sales_forecast_data import (
    REFRESH_CODE_DEPENDENCIES,
    find_header_row,
    output_is_current,
    scan_header_rows,
)


class PerformanceGuardTests(unittest.TestCase):
    def test_summary_style_registrations_do_not_grow_with_repeated_rows(self):
        original_add = IndexedList.add
        counts = []
        for size in (20, 2000):
            book = Workbook()
            sheet = book.active
            sheet.append(['名称', '数量', '比例'])
            for _ in range(size):
                sheet.append(['测试', 12, .25])
            calls = []
            def tracked_add(registry, value):
                calls.append(1)
                return original_add(registry, value)
            with patch.object(IndexedList, 'add', tracked_add):
                refresh.style_sheet(sheet, {'比例'})
            counts.append(len(calls))
            book.close()
        self.assertEqual(counts[0], counts[1])
        self.assertLess(counts[1], 50)

    def test_summary_style_reuse_preserves_original_formatting_and_values(self):
        book = Workbook()
        sheet = book.active
        sheet.append(['文本', '数量', '比例', '日期', '空白'])
        for _ in range(4):
            sheet.append(['说明', 0, .125, datetime(2026, 1, 2, 15, 30), None])
        sheet['B3'] = 2.25
        sheet['D3'] = datetime(2026, 1, 3)
        sheet['A3'].fill = PatternFill('solid', fgColor='FF123456')
        sheet['A3'].protection = Protection(locked=False)
        sheet['E4'].number_format = '0.000'
        expected = book.copy_worksheet(sheet)
        before = list(sheet.values)
        # The pre-optimization contract, including non-overwritten attributes.
        for cell in expected[1]:
            cell.fill, cell.font = refresh._HEADER_FILL, refresh._HEADER_FONT
            cell.alignment, cell.border = refresh._HEADER_ALIGNMENT, refresh._HEADER_BORDER
        for row in expected.iter_rows(min_row=2):
            for cell in row:
                cell.font, cell.alignment, cell.border = refresh._BODY_FONT, refresh._BODY_ALIGNMENT, refresh._BODY_BORDER
                if cell.row % 2 == 0:
                    cell.fill = refresh._EVEN_FILL
                if cell.column == 3:
                    cell.number_format = '0.0%'
                elif isinstance(cell.value, datetime):
                    cell.number_format = 'yyyy-mm-dd hh:mm' if any((cell.value.hour, cell.value.minute, cell.value.second)) else 'yyyy-mm-dd'
                elif isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00' if not float(cell.value).is_integer() else '#,##0'
        refresh.style_sheet(sheet, {'比例'})
        self.assertEqual(list(sheet.values), before)
        for row in sheet:
            for cell in row:
                self.assertEqual(cell._style, expected[cell.coordinate]._style, cell.coordinate)
        preserved = copy(sheet['C4']._style)
        sheet['C2'].alignment = Alignment(wrap_text=True)
        sheet['C2'].number_format = '0.000%'
        self.assertEqual(sheet['C4']._style, preserved)
        # Reusing a sheet-local cache cannot mix style indexes between workbooks.
        other = Workbook()
        other.active['A1'].fill = PatternFill('solid', fgColor='FFABCDEF')
        other.active.append([.25])
        refresh.style_sheet(other.active)
        self.assertEqual(other.active['A2'].font.name, '微软雅黑')
        self.assertEqual(other.active['A2'].number_format, '#,##0.00')
        book.close()
        other.close()

    def test_forecast_refresh_skips_only_when_every_dependency_is_older(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            mapping = root / "mapping.xlsx"
            output = root / "output.xlsx"
            source.write_bytes(b"source")
            mapping.write_bytes(b"mapping")
            output.write_bytes(b"PK\x03\x04generated")

            newest_code = max(path.stat().st_mtime_ns for path in REFRESH_CODE_DEPENDENCIES)
            dependency_time = newest_code - 2_000_000_000
            os.utime(source, ns=(dependency_time, dependency_time))
            os.utime(mapping, ns=(dependency_time, dependency_time))
            output_time = newest_code + 2_000_000_000
            os.utime(output, ns=(output_time, output_time))
            self.assertTrue(output_is_current(source, output, mapping))

            source_time = output_time + 2_000_000_000
            os.utime(source, ns=(source_time, source_time))
            self.assertFalse(output_is_current(source, output, mapping))

    def test_header_scan_can_be_reused_without_changing_detection(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["说明"])
        sheet.append(["车型", "D1", "D2"])
        scanned = scan_header_rows(sheet)

        row, headers, days = find_header_row(sheet, require_days=True, scanned_rows=scanned)

        self.assertEqual(row, 2)
        self.assertEqual(headers[:3], ["车型", "D1", "D2"])
        self.assertEqual(days, [(1, 1), (2, 2)])

    def test_metric_and_option_layout_results_are_cached_per_sheet(self):
        workbook = Workbook()
        metric = workbook.active
        metric.cell(1, 4, "26WK35")
        metric.cell(2, 1, "大定")
        metric.cell(2, 2, "数量")
        metric.cell(2, 3, "数量")
        metric.cell(2, 4, 100)
        self.assertIs(parse_metric_sheet(metric), parse_metric_sheet(metric))

        option = workbook.create_sheet("选配金")
        option.cell(3, 3, "26-08")
        option.cell(4, 3, "免费")
        option.cell(4, 4, "0-5000")
        option.cell(5, 2, "订单数量")
        option.cell(5, 3, 10)
        option.cell(5, 4, 20)
        first = _option_layout(option)
        self.assertEqual(first, (3, 4, 5))
        self.assertIs(first, _option_layout(option))


if __name__ == "__main__":
    unittest.main()
