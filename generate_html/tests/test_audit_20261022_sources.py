"""Reservation aliases obey the same physical-source precedence as launch inputs."""
from datetime import date
from pathlib import Path
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from core.excel import WorkbookItem, WorkbookStore
from core.forecast_summary import DAILY_SHEET, table_records
from core.model_identity import model_key
from modules.sales_forecast import _forecast_event_sheets, _read_actual_profiles
from tests import test_forecast_source_backfill_e2e as fixtures


class SmallMixAliasSourceTests(TestCase):
    model = "问界 M9 2026款"
    days = [date(2026, 9, i) for i in (1, 2, 3)]

    def setUp(self):
        diagnostics = patch("core.excel._CONFLICT_CHECKS", set())
        diagnostics.start()
        self.addCleanup(diagnostics.stop)

    def item(self, filename, sheets, coverage=None):
        book = Workbook()
        book.remove(book.active)
        self.addCleanup(book.close)
        for name, days, values in sheets:
            sheet = book.create_sheet(name + "by天")
            sheet.append(["指标", "统计类型", "分类", *days])
            sheet.append(["小订", "数量", "数量", *values])
        item = WorkbookItem(Path(filename), book)
        if coverage:
            item._boundary_date_range = coverage
        return item

    def read(self, items, mapping=None):
        store = WorkbookStore(Path("."))
        store.items = items
        with patch("modules.sales_forecast._read_model_mapping", return_value=mapping or {model_key(self.model): self.model}), \
             patch("modules.sales_forecast.generation_records", return_value=[]):
            profiles, _ = _read_actual_profiles(store, today=date(2026, 9, 4))
        return {profile["model"]: profile for profile in profiles}

    def values(self, profile):
        return [row["orders"] for row in profile["small_daily_days"]]

    def test_first_zero_and_known_values_survive_conflicting_alias(self):
        item = self.item("小订选配比例分析.xlsx", [
            (self.model, self.days, [0, 20, 30]),
            (self.model.replace(" ", ""), self.days, [99, 25, 30]),
        ])
        with self.assertLogs("core.excel", level="WARNING") as logs:
            result = self.read([item])
        self.assertEqual(self.values(result[self.model]), [0, 20, 30])
        self.assertTrue(any("冲突" in line for line in logs.output))

    def test_explicit_blanks_fill_without_erasing_prior_observations(self):
        item = self.item("小订选配比例分析.xlsx", [
            (self.model, self.days, [0, 20, None]),
            (self.model.replace(" ", ""), self.days, [None, None, 30]),
        ])
        self.assertEqual(self.values(self.read([item])[self.model]), [0, 20, 30])

    def test_physical_omitted_day_stays_zero_before_alias_fallback(self):
        first = self.item("小订选配比例分析_A.xlsx", [
            (self.model, [self.days[0], self.days[2]], [1, None]),
        ], coverage=(self.days[0], self.days[2]))
        second = self.item("小订选配比例分析_B.xlsx", [
            (self.model.replace(" ", ""), [self.days[1], self.days[2]], [9, 7]),
        ], coverage=(self.days[0], self.days[2]))
        self.assertEqual(self.values(self.read([first, second])[self.model]), [1, 0, 7])

    def test_maintained_aliases_merge_but_distinct_models_stay_separate(self):
        alias, other = "M9旗舰2026款", "问界 M8 2026款"
        item = self.item("小订选配比例分析.xlsx", [
            (self.model, self.days, [0, 20, None]),
            (alias, self.days, [None, None, 30]),
            (other, self.days, [7, 8, 9]),
        ])
        mapping = {model_key(self.model): self.model, model_key(alias): self.model, model_key(other): other}
        result = self.read([item], mapping)
        self.assertEqual(set(result), {self.model, other})
        self.assertEqual(self.values(result[self.model]), [0, 20, 30])
        self.assertEqual(self.values(result[other]), [7, 8, 9])

    def test_first_physical_file_alias_precedes_later_standard_name(self):
        alias = "M9旗舰2026款"
        first = self.item("小订选配比例分析_A.xlsx", [
            (self.model, self.days, [None, None, None]),
            (alias, self.days, [10, 20, 30]),
        ])
        second = self.item("小订选配比例分析_B.xlsx", [
            (self.model, self.days, [99, 88, 77]),
        ])
        mapping = {model_key(self.model): self.model, model_key(alias): self.model}
        self.assertEqual(self.values(self.read([first, second], mapping)[self.model]), [10, 20, 30])

    def test_file_order_and_original_sheet_order_are_both_preserved(self):
        alias = "M9旗舰2026款"
        first = self.item("小订选配比例分析_Z.xlsx", [
            (alias, self.days, [0, None, 30]),
            (self.model, self.days, [5, 20, None]),
        ])
        second = self.item("小订选配比例分析_A.xlsx", [
            (self.model, self.days, [99, 88, 77]),
        ])
        mapping = {model_key(self.model): self.model, model_key(alias): self.model}
        self.assertEqual(self.values(self.read([first, second], mapping)[self.model]), [0, 20, 30])

    def test_later_file_only_alias_keeps_its_actual_source_identity(self):
        alias = "M9旗舰2026款"
        first = self.item("小订选配比例分析_A.xlsx", [(self.model, self.days, [0, 20, None])])
        second = self.item("小订选配比例分析_B.xlsx", [(alias, self.days, [99, 88, 30])])
        store = WorkbookStore(Path("."))
        store.items = [first, second]
        mapping = {model_key(self.model): self.model, model_key(alias): self.model}
        merged = list(_forecast_event_sheets(store.find("小订选配比例"), mapping))[0]
        self.assertEqual([origin.path for origin, _ in merged._source_matches], [first.path, second.path])
        self.assertEqual([cell.value for cell in merged[2][3:]], [0, 20, 30])

    def test_missing_physical_metadata_keeps_flattened_relative_order(self):
        alias = "M9旗舰2026款"
        known = self.item("小订选配比例分析.xlsx", [
            (self.model, self.days, [None, None, None]), (alias, self.days, [None, None, None]),
        ])
        external = self.item("小订选配比例分析_外部.xlsx", [
            (self.model, self.days, [10, 20, 30]), (alias, self.days, [99, 88, 77]),
        ])
        # One unknown origin and one known origin with a missing original-sheet
        # identity must retain their existing order when no rank can be proved.
        known.workbook.worksheets[0]._source_matches = [(external, external.workbook.worksheets[0])]
        known.workbook.worksheets[1]._source_matches = [(known, external.workbook.worksheets[1])]
        mapping = {model_key(self.model): self.model, model_key(alias): self.model}
        merged = list(_forecast_event_sheets(known, mapping))[0]
        self.assertEqual([origin.path for origin, _ in merged._source_matches], [external.path, known.path])
        self.assertEqual([cell.value for cell in merged[2][3:]], [10, 20, 30])

    def test_launch_and_cancel_aliases_also_keep_first_file_priority(self):
        alias = "M9旗舰2026款"
        mapping = {model_key(self.model): self.model, model_key(alias): self.model}
        for kind in ("首销期订单节奏", "小订退订分析"):
            with self.subTest(kind=kind):
                items = []
                for suffix, entries in (("A", ((self.model, None), (alias, 10))),
                                        ("B", ((self.model, 99),))):
                    book = Workbook()
                    book.remove(book.active)
                    self.addCleanup(book.close)
                    for name, value in entries:
                        if kind == "首销期订单节奏":
                            sheet = book.create_sheet(name + "by天")
                            sheet.append(["周期", "D1"])
                            sheet.append(["时间", self.days[0]])
                            sheet.append(["当日大定数量", value])
                        else:
                            sheet = book.create_sheet(name + "_日度退订")
                            sheet.append(["取消日期", "当日小订退", "累计小订退"])
                            sheet.append([self.days[0], value, value])
                    items.append(WorkbookItem(Path(kind + "_" + suffix + ".xlsx"), book))
                profile = self.read(items, mapping)[self.model]
                value = profile["days"][0]["gross"] if kind == "首销期订单节奏" else profile["cancel_days"][0]["daily_cancel"]
                self.assertEqual(value, 10)

    def test_saved_summary_and_web_inputs_keep_alias_precedence(self):
        fixture = fixtures.SourceBackfillEndToEndTests()
        self.addCleanup(fixture.doCleanups)

        def source_edit(path):
            book = load_workbook(path)
            for column in (2, 3, 4):
                book["小订by天"].cell(2, column).value = None
            book.save(path)
            book.close()

        def add_sources(folder):
            book = Workbook()
            book.remove(book.active)
            for name, values in ((fixture.model, [0, 20]), (fixture.model.replace(" ", ""), [99, None])):
                sheet = book.create_sheet(name + "by天")
                sheet.append(["指标", "统计类型", "分类", date(2026, 8, 30), date(2026, 8, 31)])
                sheet.append(["小订", "数量", "数量", *values])
            book.save(folder / "鸿蒙智行小订选配比例分析.xlsx")
            book.close()

        book = fixture.build(source_edit=source_edit, add_sources=add_sources)
        rows = {str(row["日期"])[:10]: row for row in table_records(book, DAILY_SHEET)}
        self.assertEqual([rows[day]["小订数量"] for day in ("2026-08-30", "2026-08-31")], [0, 20])
        _, profiles, _, _, _ = fixture.web_inputs(book, date(2026, 9, 5))
        profile = next(row for row in profiles if row["model"] == fixture.model)
        self.assertEqual(self.values(profile), [0, 20])
