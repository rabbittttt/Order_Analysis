from datetime import date, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook

from core.discovery import discover_subjects
from core.excel import WorkbookItem, WorkbookStore
from core.forecast_summary import table_records, visible_target_names, MASTER_SHEET, summary_scope
from core.model_identity import stage_records, stage_name, parent_generation, resolve_stage_identity, has_reservation, FORECAST_SOURCE_PATH, generation_records
from core.models import Subject
from modules.launch_rhythm import LaunchRhythmModule
from modules.sales_forecast import _combined_steady_inputs, _read_stage_windows, _read_model_mapping_file, _progress_rows, _small_campaign_rows, _read_model_master, _read_model_mapping, _primary_attributes
from tools.refresh_sales_forecast_data import load_mapping, normalize_public_names, write_rows, append_forecast_inputs, DEFAULT_SOURCE, DEFAULT_MAPPING


PARENT = '问界 M7 2024款'
ULTRA, PRO = PARENT + ' Ultra版', PARENT + ' Pro版'


def records():
    return [dict(代际名=PARENT, 二级代际名=name, 小订开始日期=datetime(2024, month, 1),
                 小订结束日期=datetime(2024, month, 2), 开始大定日期=datetime(2024, month, 3),
                 小转大结束日期=datetime(2024, month, 4), 首销期天数=2)
            for name, month in ((ULTRA, 5), (PRO, 8))]


class GenerationCampaignTests(TestCase):
    def test_default_source_paths_share_scripts_config(self):
        from modules.sales_forecast import RAW_FORECAST_DATA, HISTORY_CANDIDATES
        expected = Path(__file__).resolve().parents[2] / 'config' / '小订及首销数据整理.xlsx'
        self.assertEqual(FORECAST_SOURCE_PATH, expected)
        self.assertEqual(DEFAULT_SOURCE, expected)
        self.assertEqual(DEFAULT_MAPPING, expected)
        self.assertEqual(RAW_FORECAST_DATA, expected)
        self.assertIn(expected, HISTORY_CANDIDATES)
        with patch('core.model_identity.stage_records', return_value=[]) as reader:
            generation_records()
        reader.assert_called_once_with(expected)

    def make_source(self, root):
        source = root / '小订及首销数据整理.xlsx'
        book = Workbook(); book.active.title = '任意首Sheet名称'
        rows = records(); headers = list(rows[0])
        book.active.append(headers)
        for record in rows: book.active.append([record[k] for k in headers])
        book.save(source); book.close()
        return source

    def test_master_keeps_two_events_under_one_primary_without_legacy_mapping(self):
        with TemporaryDirectory() as folder:
            source = self.make_source(Path(folder))
            self.assertEqual([stage_name(r) for r in stage_records(source)], [ULTRA, PRO])
            mapping, aliases = load_mapping(source)
            self.assertEqual(len(mapping), 2)
            self.assertNotIn(PARENT, aliases)
            parsed = _read_model_mapping_file(source)
            self.assertEqual(set(parsed.values()), {ULTRA, PRO})
            with patch('modules.sales_forecast.RAW_FORECAST_DATA', source):
                _, windows = _read_stage_windows()
            self.assertEqual({w['generation'] for w in windows.values()}, {ULTRA, PRO})

    def test_legacy_name_requires_one_confirmed_version_or_window(self):
        rs = records()
        self.assertIsNone(resolve_stage_identity('问界 新M7', rs))
        self.assertEqual(resolve_stage_identity('问界 新M7 2024款 Ultra', rs), ULTRA)
        self.assertEqual(resolve_stage_identity('问界 新M7', rs, date(2024, 8, 3)), PRO)
        self.assertEqual(parent_generation(ULTRA, rs), PARENT)
        self.assertEqual(parent_generation(PARENT + '总计', rs), PARENT + '总计')

    def test_legacy_family_matching_does_not_consume_year(self):
        rs = [dict(代际名='智界 S7 2024款', 开始大定日期=date(2023, 11, 28)),
              dict(代际名='智界 S7 2025款', 开始大定日期=date(2024, 11, 26))]
        self.assertEqual(resolve_stage_identity('智界 S7', rs, date(2023, 11, 28)), '智界 S7 2024款')
        self.assertIsNone(resolve_stage_identity('智界 S7', rs))
        self.assertIsNone(resolve_stage_identity('问界 S7', rs, date(2023, 11, 28)))

    def test_same_integrated_file_imports_once_and_scope_reads_its_attributes(self):
        with TemporaryDirectory() as folder:
            root = Path(folder); source = self.make_source(root); orders = root / 'orders'; orders.mkdir()
            book = Workbook(); append_forecast_inputs(book, source, source, orders, '2024-08-03', preserve_layout=False)
            cached = root / 'temporary.xlsx'; book.save(cached)
            with summary_scope(cached, workbook=book), patch('modules.sales_forecast.stage_records', side_effect=AssertionError('汇总作用域不得补读原始工作簿')):
                self.assertEqual(set(_read_model_mapping().values()), {ULTRA, PRO})
                self.assertIn('问界m72024ultra版', _read_model_master())

    def test_primary_attributes_combine_energy_without_reusing_campaign_totals(self):
        master = {str(i): dict(代际名='问界 M8 2025款', 二级代际名='版'+str(i), 能源类型=e, 总小订=100)
                  for i,e in enumerate(('增程', '纯电'))}
        parent = _primary_attributes(master)['问界m82025']
        self.assertEqual(parent['能源类型'], '增程/纯电')
        self.assertNotIn('总小订', parent)

    def test_attribute_only_modern_rows_are_not_small_order_vehicles(self):
        self.assertFalse(has_reservation({'代际名': '无小订车', '总小订': 0}))
        self.assertTrue(has_reservation({'代际名': '小订车', '小订开始日期': date(2026, 1, 1)}))
        self.assertTrue(has_reservation({'代际名': '小订车', '有小订': True}))
        self.assertFalse(has_reservation({'有小订': False, '总小订': 100}))
        self.assertTrue(has_reservation({}, modern=False))

    def test_visible_names_and_progress_reader_use_campaign_identity(self):
        book = Workbook(); book.active.title = MASTER_SHEET
        book.active.append(['代际名', '二级代际名', '订单来源文件'])
        for name in (ULTRA, PRO): book.active.append([PARENT, name, '来源.xlsx'])
        self.assertEqual(visible_target_names(book), [PARENT])
        sheet = book.create_sheet('小转大累计完成度')
        sheet.append(['代际名', 'D1', 'D2', '二级代际名'])
        sheet.append([PARENT, .3, 1, ULTRA]); sheet.append([PARENT, .4, 1, PRO])
        self.assertEqual(_progress_rows(book, sheet.title), {ULTRA: [.3, 1], PRO: [.4, 1]})
        self.assertEqual([r['订单分析代际名'] for r in table_records(book, MASTER_SHEET)], [ULTRA, PRO])

    def test_normalization_preserves_quantities_and_never_exports_blank_identity(self):
        book = Workbook(); sheet = book.active; sheet.title = '小订参考曲线'
        sheet.append(['历史传播名', '订单分析代际名', 'D1'])
        sheet.append([ULTRA, ULTRA, .5]); sheet.append([None, None, None])
        def emit(name, headers, rows, percent_headers=None):
            del book[name]; write_rows(book, name, headers, rows, 'Campaign', percent_headers)
        data = {'targets': [{'name': ULTRA, 'primary_generation': PARENT, 'secondary_generation': ULTRA}]}
        with patch('tools.refresh_sales_forecast_data.stage_records', return_value=records()):
            normalize_public_names(book, data, emit)
        self.assertEqual(list(book.active.values), [('代际名', '二级代际名', 'D1'), (PARENT, ULTRA, .5)])

    def test_primary_selector_and_exact_secondary_sheet_pairing(self):
        book = Workbook(); book.remove(book.active)
        # Deliberately reverse file order: hourly sheets must pair by edition.
        for name, month, quantity in ((PRO, 8, 20), (ULTRA, 5, 10)):
            day = book.create_sheet(name + 'by天')
            day.append(['周期', 'D1']); day.append(['时间', datetime(2024, month, 3)])
            day.append(['当日大定数量', quantity])
            hourly = book.create_sheet(name + 'by时')
            hourly.append(['日期', datetime(2024, month, 3)])
            hourly.append(['时刻', 20]); hourly.append(['当日大定数量', quantity])
        store = WorkbookStore(Path('.')); store.items = [WorkbookItem(Path('首销期订单节奏.xlsx'), book)]
        with patch('core.discovery.generation_records', return_value=records()), patch('core.model_identity.generation_records', return_value=records()):
            subjects = discover_subjects(store)
            self.assertEqual([s.name for s in subjects if s.type == 'generation'], [PARENT])
            self.assertEqual(len(store.find_subject_sheets('首销期订单节奏', PARENT, 'day')), 2)
            self.assertEqual(len(store.find_subject_sheets('首销期订单节奏', ULTRA, 'day')), 1)
            board = LaunchRhythmModule().build(store, Subject('m7', PARENT, 'generation', '问界'))
        self.assertTrue(board.views['day']['default_period'].startswith('Pro版 · '))
        self.assertTrue(any(label.startswith('Ultra版 · ') for label in board.views['day']['periods']))

    def test_flat_sales_combines_actual_quantities_not_ratios_or_gap_zeroes(self):
        targets = [dict(name=ULTRA, primary_generation=PARENT, secondary_generation=ULTRA,
                        launch_date='2024-05-01', end_date='2024-05-01'),
                   dict(name=PRO, primary_generation=PARENT, secondary_generation=PRO,
                        launch_date='2024-05-03', end_date='2024-05-03')]
        actuals = [dict(model=t['name'], days=[dict(date=t['launch_date'], gross=amount,
                    direct=amount, lock=locks)]) for t, amount, locks in zip(targets, (100, 200), (10, 100))]
        merged, profiles, history = _combined_steady_inputs(targets, actuals, [])
        self.assertEqual(merged[0]['name'], PARENT)
        self.assertEqual(merged[0]['steady_start_date'], '2024-05-04')
        self.assertEqual([r['gross'] for r in profiles[0]['days']], [100, 200])
        self.assertEqual(history[0]['daily_direct'], [100, 200])
        self.assertEqual(history[0]['gross'], 300)
        self.assertEqual(history[0]['direct'], 300)
        self.assertAlmostEqual(history[0]['lock_rate'], 110/300)

    def test_flat_sales_unknown_component_is_not_zero(self):
        targets = [dict(name=n, primary_generation=PARENT, secondary_generation=n,
                        launch_date='2024-05-01', end_date='2024-05-01') for n in (ULTRA, PRO)]
        actuals = [dict(model=ULTRA, days=[dict(date='2024-05-01', gross=20, direct=10)])]
        _, profiles, history = _combined_steady_inputs(targets, actuals, [])
        self.assertIsNone(profiles[0]['days'][0]['gross'])
        self.assertIsNone(history[0]['lock_rate'])

    def test_unknown_campaign_window_never_falls_back_to_one_edition_flat_sales(self):
        targets = [dict(name=n, primary_generation=PARENT, secondary_generation=n,
                        launch_date='2024-05-01', end_date='2024-05-02' if n == ULTRA else '') for n in (ULTRA, PRO)]
        merged, profiles, history = _combined_steady_inputs(targets, [], [])
        self.assertEqual(merged[0]['name'], PARENT)
        self.assertTrue(merged[0]['data_error'])
        self.assertIn(PRO, merged[0]['date_source_label'])
        self.assertEqual(merged[0]['steady_start_date'], '')
        self.assertEqual((profiles, history), ([], []))

    def test_primary_small_daily_is_assigned_once_and_overlap_is_not_duplicated(self):
        windows = {name: dict(generation=name, primary_generation=PARENT, secondary_generation=name,
                   small_start_date=start, small_end_date=end)
                   for name, start, end in ((ULTRA, '2024-05-01', '2024-05-02'), (PRO, '2024-05-02', '2024-05-03'))}
        rows = [dict(date=f'2024-05-0{i}', orders=i) for i in (1, 2, 3)]
        owned, overlap = _small_campaign_rows(PARENT, rows, windows)
        self.assertEqual(owned[ULTRA], [rows[0]])
        self.assertEqual(owned[PRO], [rows[2]])
        self.assertEqual(overlap, ['2024-05-02'])
