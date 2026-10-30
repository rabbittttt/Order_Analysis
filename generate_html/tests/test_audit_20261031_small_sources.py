"""One invalid maintained small-order cell must not erase other known daily facts."""
from datetime import date, datetime
from hashlib import sha256
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from core.forecast_summary import DAILY_SHEET, SMALL_DAILY_SHEET, table_records
from core.model_identity import model_key
from modules.sales_forecast import _read_small_order_history
from tests import test_forecast_source_backfill_e2e as fixtures


class SmallMaintainedCellTests(TestCase):
    model = '问界 M9 2026款'

    def read(self, values, total=100, summary_total=100):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        path = Path(temporary.name) / '小订及首销数据整理.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.title = '车型汇总'
        sheet.append(['代际名', '小订开始日期', '小订结束日期', '小订天数', '总小订'])
        sheet.append([self.model, date(2026, 9, 1), date(2026, 9, 3), 3, summary_total])
        sheet = book.create_sheet('小订by天')
        sheet.append(['代际名', date(2026, 9, 1), date(2026, 9, 2), date(2026, 9, 3), '合计'])
        sheet.append([self.model, *values, total])
        book.save(path)
        book.close()
        before = sha256(path.read_bytes()).hexdigest()
        with patch('modules.sales_forecast._read_model_mapping', return_value={model_key(self.model): self.model}), \
             patch('modules.sales_forecast._read_model_master', return_value={}):
            _, items = _read_small_order_history(path)
        self.assertEqual(sha256(path.read_bytes()).hexdigest(), before)
        return items[0]

    def test_invalid_interior_cell_preserves_dates_and_other_valid_values(self):
        for value in ('待更新', '#DIV/0!', '—', True, False, -1, 'NaN', 'Infinity', '-Infinity'):
            with self.subTest(value=value):
                item = self.read([40, value, 20])
                self.assertTrue(item['daily_actual'])
                self.assertEqual(item['daily_orders'], [40, None, 20])
                self.assertEqual(item['dates'], ['2026-09-01', '2026-09-02', '2026-09-03'])
                self.assertEqual(item['small_progress'], [.4, None, None])

    def test_invalid_first_cell_does_not_shift_later_observations_into_d1(self):
        item = self.read(['待更新', 30, 20])
        self.assertTrue(item['daily_actual'])
        self.assertEqual(item['daily_orders'], [None, 30, 20])
        self.assertEqual(item['d1_share'], 0)
        self.assertEqual(item['small_progress'], [None, None, None])

    def test_real_zero_and_supported_count_strings_keep_priority(self):
        item = self.read(['0', '待更新', ' 20 '])
        self.assertTrue(item['daily_actual'])
        self.assertEqual(item['daily_orders'], [0, None, 20])
        self.assertEqual(item['small_progress'], [0, None, None])

    def test_missing_terminal_does_not_make_partial_daily_sum_a_final_total(self):
        item = self.read([40, '待更新', 20], total=None, summary_total=None)
        self.assertTrue(item['daily_actual'])
        self.assertEqual(item['daily_orders'], [40, None, 20])
        self.assertFalse(item['total_complete'])
        self.assertEqual(item['small_progress'], [None, None, None])

    def test_complete_valid_record_remains_unchanged(self):
        item = self.read([40, 40, 20])
        self.assertEqual(item['daily_orders'], [40, 40, 20])
        self.assertTrue(item['total_complete'])
        self.assertEqual(item['small_progress'], [.4, .8, 1])

    def test_contradictory_total_and_invalid_total_keep_existing_rejection(self):
        for values, total in (([110, -10, 0], 100), ([40, None, 20], '待更新')):
            with self.subTest(values=values, total=total):
                item = self.read(values, total)
                self.assertFalse(item['daily_actual'])


class SmallMaintainedCellEndToEndTests(TestCase):
    def build(self, first=40, fallback=None):
        fixture = fixtures.SourceBackfillEndToEndTests()
        self.addCleanup(fixture.doCleanups)
        source_digest = {}

        def edit(path):
            book = load_workbook(path)
            book['小订by天'].cell(2, 2).value = first
            book['小订by天'].cell(2, 3).value = '待更新'
            book.save(path)
            book.close()
            source_digest['value'] = sha256(path.read_bytes()).hexdigest()

        def add(folder):
            if fallback is None:
                return
            book = Workbook()
            sheet = book.active
            sheet.title = fixture.model + 'by天'
            sheet.append(['指标', '统计类型', '分类', date(2026, 8, 30), date(2026, 8, 31)])
            sheet.append(['小订', '数量', '数量', *fallback])
            book.save(folder / '鸿蒙智行小订选配比例分析.xlsx')
            book.close()

        book = fixture.build(source_edit=edit, add_sources=add)
        source = book._test_path.parent / '小订及首销数据整理.xlsx'
        self.assertEqual(sha256(source.read_bytes()).hexdigest(), source_digest['value'])
        daily = {str(row['日期'])[:10]: row for row in table_records(book, DAILY_SHEET)}
        _, profiles, _, history, _ = fixture.web_inputs(book, date(2026, 9, 5))
        profile = next(item for item in profiles if item['model'] == fixture.model)
        event = next(item for item in history if item['generation'] == fixture.model)
        return book, daily, profile, event

    def test_only_invalid_day_falls_back_and_saved_sources_follow_provider(self):
        book, daily, profile, event = self.build(fallback=[99, 60])
        self.assertEqual([daily[d]['小订数量'] for d in ('2026-08-30', '2026-08-31')], [40, 60])
        self.assertIn('小订及首销数据整理.xlsx', daily['2026-08-30']['字段来源'])
        self.assertNotIn('小订选配比例分析', daily['2026-08-30']['字段来源'])
        self.assertIn('小订选配比例分析.xlsx', daily['2026-08-31']['字段来源'])
        self.assertEqual([row['orders'] for row in profile['small_daily_days']], [40, 60])
        self.assertEqual(event['daily_orders'], [40, 60])
        row = table_records(book, SMALL_DAILY_SHEET)[0]
        self.assertEqual([row['D1'], row['D2']], [40, 60])

    def test_known_zero_wins_over_fallback_with_positive_value(self):
        _, daily, profile, _ = self.build(first=0, fallback=[99, 100])
        self.assertEqual([daily[d]['小订数量'] for d in ('2026-08-30', '2026-08-31')], [0, 100])
        self.assertEqual([row['orders'] for row in profile['small_daily_days']], [0, 100])

    def test_missing_fallback_stays_unknown_without_erasing_the_known_day(self):
        _, daily, _, event = self.build()
        self.assertEqual(daily['2026-08-30']['小订数量'], 40)
        self.assertIsNone(daily['2026-08-31']['小订数量'])
        self.assertEqual(event['daily_orders'], [40, None])
        self.assertEqual(event['small_progress'], [.4, None])
