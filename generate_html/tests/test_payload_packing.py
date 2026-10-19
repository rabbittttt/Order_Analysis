import base64
import gzip
import json
import shutil
import subprocess
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from core.renderer import pack_dashboard, render_dashboard
from core.validation import validate_manifest


ROOT = Path(__file__).resolve().parents[1]


def unpack(value):
    if isinstance(value, list):
        return [unpack(v) for v in value]
    if not isinstance(value, dict):
        return value
    if value.get("format") == "shared-v1":
        pool = value["pool"]
        def decode(item):
            if isinstance(item, list):
                return [decode(v) for v in item]
            if not isinstance(item, dict):
                return item
            if "$shared" in item:
                return decode(pool[item["$shared"]])
            entries = item["$literal"] if "$literal" in item else item.items()
            return {k: decode(v) for k, v in entries}
        return decode(value["root"])
    return value


class PayloadPackingTests(unittest.TestCase):
    def fixture(self):
        curve = {"dates": list(range(1000)), "values": [0, None, 12.5] * 334}
        return {"views": {"day": {"pages": {
            str(i): {"kpis": [{"value": i}], "sections": [
                {"data": curve}, {"data": {"history": curve}},
            ]} for i in range(100)
        }}}}

    def test_all_periods_values_and_missing_values_round_trip(self):
        board = self.fixture()
        packed = json.loads(json.dumps(pack_dashboard(board)))
        self.assertEqual(unpack(packed), board)
        self.assertLess(len(json.dumps(packed)), len(json.dumps(board)) / 10)

    def test_reserved_keys_distinct_equal_objects_and_no_mutation(self):
        source = {"$shared": 0, "$literal": [["x", None]]}
        board = {"a": source, "b": source, "c": dict(source), "zero": 0}
        before = json.dumps(board)
        packed = pack_dashboard(board)
        self.assertEqual(unpack(packed), board)
        self.assertEqual(json.dumps(board), before)
        self.assertNotEqual(packed["root"]["a"], packed["root"]["c"])

    def test_cycles_fail_clearly(self):
        value = {}
        value["cycle"] = value
        with self.assertRaises(ValueError):
            pack_dashboard(value)

    def test_nonfinite_business_values_are_null_in_strict_browser_json(self):
        shared = {'values': [0, 12.5, float('nan'), float('inf'), -float('inf'), None]}
        board = {'a': shared, 'b': shared, 'text': 'NaN is a diagnostic, not a number'}
        payload = json.dumps(pack_dashboard(board), allow_nan=False)
        decoded = unpack(json.loads(payload))
        self.assertEqual(decoded['a']['values'], [0, 12.5, None, None, None, None])
        self.assertEqual(decoded['a'], decoded['b'])
        self.assertEqual(decoded['text'], board['text'])
        self.assertNotEqual(shared['values'][2], shared['values'][2])

    def test_validation_visits_shared_values_and_reports_invalid_numbers(self):
        shared = {"value": float("nan")}
        board = {"subject_id": "a", "views": {}, "x": shared, "y": shared}
        result = validate_manifest({"subjects": [{"id": "a"}], "dashboards": {"a": board}})
        self.assertTrue(any("无效值" in text for text in result))
        shared["value"] = 0
        self.assertEqual(validate_manifest({"subjects": [{"id": "a"}], "dashboards": {"a": board}}), [])

    def test_real_javascript_decoder(self):
        node = shutil.which("node")
        if not node:
            self.skipTest("Node unavailable")
        script = (ROOT / "templates/dashboard.js").read_text(encoding="utf-8")
        decoder = script[script.index("  function unpackDashboard"):script.index("  const DATA =")]
        board = self.fixture()
        result = subprocess.run([node, "-e", decoder + "\nprocess.stdout.write(JSON.stringify(unpackDashboard(JSON.parse(require('fs').readFileSync(0,'utf8')))));"],
                                input=json.dumps(pack_dashboard(board)), text=True, capture_output=True, check=True)
        self.assertEqual(json.loads(result.stdout), board)

    def test_output_has_separate_blocks_and_small_bootstrap(self):
        import re
        board = self.fixture()
        with TemporaryDirectory() as folder:
            output = Path(folder) / "output.html"
            render_dashboard({"dashboards": {"a": board}, "raw_blocks": {}, "subjects": []}, ROOT / "templates", output)
            html = output.read_text(encoding="utf-8")
        bootstrap = re.search(r'window.DASHBOARD_DATA_B64="([^"]+)"', html)[1]
        main = json.loads(gzip.decompress(base64.b64decode(bootstrap)))
        self.assertNotIn("dashboard_blocks", main)
        blocks = json.loads(re.search(r'window.DASHBOARD_BLOCKS_B64=(.*?);</script>', html)[1])
        self.assertEqual(unpack(json.loads(gzip.decompress(base64.b64decode(blocks["a"])))), board)

    def test_rendered_nonfinite_board_keeps_other_boards_and_original_diagnostics(self):
        import re
        bad = {'subject_id': 'a', 'views': {}, 'value': float('nan')}
        good = {'subject_id': 'b', 'views': {}, 'value': 0}
        manifest = {'subjects': [{'id': 'a'}, {'id': 'b'}],
                    'dashboards': {'a': bad, 'b': good}, 'raw_blocks': {}}
        with TemporaryDirectory() as folder:
            output = Path(folder) / 'output.html'
            render_dashboard(manifest, ROOT / 'templates', output)
            blocks = json.loads(re.search(r'window.DASHBOARD_BLOCKS_B64=(.*?);</script>',
                output.read_text(encoding='utf-8'))[1])
        def reject_constant(value):
            raise AssertionError('Nonstandard JSON token: ' + value)
        decoded = {key: unpack(json.loads(gzip.decompress(base64.b64decode(value)),
                    parse_constant=reject_constant)) for key, value in blocks.items()}
        self.assertIsNone(decoded['a']['value'])
        self.assertEqual(decoded['b'], good)
        self.assertTrue(any('无效值' in text for text in validate_manifest(manifest)))


    def test_sparse_raw_table_keeps_formats_zero_and_order(self):
        from types import SimpleNamespace
        from openpyxl import Workbook
        from main import raw_tables
        book = Workbook()
        sheet = book.active
        sheet.title = "测试by天"
        sheet.append(["日期", None, "比例", "数量"])
        sheet.append(["2026-01-01", None, .125, 0])
        sheet["C2"].number_format = "0.0%"
        sheet["D2"].number_format = "#,##0"
        sheet.cell(100000, 1000).number_format = "0.00"
        files, blocks = raw_tables(SimpleNamespace(items=[SimpleNamespace(path=Path("原表.xlsx"), workbook=book)]))
        rows = json.loads(gzip.decompress(base64.b64decode(blocks["0|0"])))["rows"]
        self.assertEqual(rows, [["日期", "比例", "数量"], ["2026-01-01", "12.5%", "0"]])
        self.assertEqual(files[0]["sheets"][0]["total_rows"], 2)
        book.close()

    def test_option_percentage_format_repair_reaches_raw_tables_and_refreshes_cache(self):
        from types import SimpleNamespace
        from openpyxl import Workbook
        from core.excel import load_data_workbook
        from main import raw_tables

        with TemporaryDirectory() as folder:
            path = Path(folder) / "鸿蒙智行大定选配比例分析.xlsx"
            cache = Path(folder) / "cache"
            book = Workbook()
            book.remove(book.active)
            for name in ("智界by周", "问界by天", "尊界by月"):
                sheet = book.create_sheet(name)
                for row in (
                    ["指标", "统计类型", "分类", "当期", "上期"],
                    ["大定", "数量", "数量", 1234, 0],
                    ["动力", "占比", "增程", .3111, 0],
                    [None, None, "纯电", .6889, 1],
                    ["门店类型", "占比", "华为授权体验店", .375, .5],
                    [None, "数量", "华为授权体验店", 1234, 0],
                    [None, "贡献门店数", "华为授权体验店", 131, 0],
                    ["普通小数", "数值", "产能", 12.34, None],
                ):
                    sheet.append(row)
                sheet.merge_cells("B3:B4")
                for row in (2, 6, 7):
                    for cell in sheet[row][3:]:
                        cell.number_format = "#,##0"
                for row in (3, 4, 5):
                    for cell in sheet[row][3:]:
                        cell.number_format = "0.0%" if row == 5 else "0.00%"
                sheet["D3"].number_format = "#,##0.00"
                sheet["D8"].number_format = "#,##0.00"
            original_values = [[list(row) for row in sheet.values] for sheet in book]
            book.save(path)

            def read_display_rows():
                loaded = load_data_workbook(path, cache)
                self.assertEqual([[list(row) for row in sheet.values] for sheet in loaded], original_values)
                _, blocks = raw_tables(SimpleNamespace(items=[SimpleNamespace(path=path, workbook=loaded)]))
                loaded.close()
                return [json.loads(gzip.decompress(base64.b64decode(blocks[f"0|{index}"])))["rows"]
                        for index in range(3)]

            before = read_display_rows()
            for rows in before:
                self.assertEqual(rows[2][3], "0.31")
            # A source-only format repair must invalidate the cached display;
            # merged labels must not affect percent precision or numeric values.
            for sheet in book:
                sheet["D3"].number_format = "0.00%"
            book.save(path)
            book.close()
            after = read_display_rows()
            self.assertEqual(read_display_rows(), after)
            for rows in after:
                self.assertEqual(rows[1][3:], ["1,234", "0"])
                self.assertEqual(rows[2][3:], ["31.11%", "0.00%"])
                self.assertEqual(rows[3][3:], ["68.89%", "100.00%"])
                self.assertEqual(rows[4][3:], ["37.5%", "50.0%"])
                self.assertEqual(rows[5][3:], ["1,234", "0"])
                self.assertEqual(rows[6][3:], ["131", "0"])
                self.assertEqual(rows[7][3:], ["12.34", None])
            for old, repaired in zip(before, after):
                old[2][3] = "31.11%"
                self.assertEqual(old, repaired)
