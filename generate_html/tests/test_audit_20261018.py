"""Preserve usable weeks, intentional D1 zero and failed-build exit status."""
from argparse import Namespace
from datetime import date
import json
import logging
from pathlib import Path
import shutil
import subprocess
from unittest import TestCase, skipUnless
from unittest.mock import patch

from openpyxl import Workbook

import main as generator
from core.forecast_summary import DAILY_SHEET, MASTER_SHEET, WEEKLY_SHEET, read_public_forecast

ROOT = Path(__file__).resolve().parents[1]


class VisibleSteadyWeeksTests(TestCase):
    def read_weeks(self, entries, today=date(2026, 2, 1)):
        book = Workbook()
        self.addCleanup(book.close)
        master = book.active
        master.title = MASTER_SHEET
        master.append(['历史传播名', '订单分析代际名', '首销开始', '首销结束', '订单来源文件', '有小订'])
        master.append(['测试车', '测试车', '2025-12-20', '2026-01-04', '来源.xlsx', False])
        daily = book.create_sheet(DAILY_SHEET)
        daily.append(['订单分析代际名', '日期', '订单阶段', '交车锁单', '字段来源'])
        weeks = book.create_sheet(WEEKLY_SHEET)
        weeks.append(['订单分析代际名', '周期', '订单阶段', '统计开始', '统计结束', '交车锁单', '锁单来源'])
        for period, first, last, value in entries:
            weeks.append(['测试车', period, '平销', first, last, value,
                          '鸿蒙智行锁单选配比例分析.xlsx｜测试车by周'])
        with patch('modules.sales_forecast._read_model_mapping', return_value={}), \
             patch('modules.sales_forecast._read_model_master', return_value={}):
            return read_public_forecast(book, [], today=today)[3]

    def test_weekly_only_history_survives_an_unknown_gap(self):
        result = self.read_weeks([('26WK02', '2026-01-05', '2026-01-11', 70),
                                  ('26WK04', '2026-01-19', '2026-01-25', 140)])
        self.assertEqual(len(result), 1)
        self.assertEqual([(r['week'], r['lock']) for r in result[0]['weeks']], [(1, 70), (3, 140)])

    def test_zero_week_is_known_and_missing_week_is_not_invented(self):
        result = self.read_weeks([('26WK02', '2026-01-05', '2026-01-11', 0),
                                  ('26WK03', '2026-01-12', '2026-01-18', None),
                                  ('26WK04', '2026-01-19', '2026-01-25', 14)])
        self.assertEqual(len(result), 1)
        self.assertEqual([(r['week'], r['lock']) for r in result[0]['weeks']], [(1, 0), (3, 14)])

    def test_consecutive_weeks_and_unfinished_boundary_keep_existing_rules(self):
        entries = [('26WK02', '2026-01-05', '2026-01-11', 70),
                   ('26WK03', '2026-01-12', '2026-01-18', 140)]
        self.assertEqual([r['week'] for r in self.read_weeks(entries)[0]['weeks']], [1, 2])
        self.assertEqual([r['period'] for r in self.read_weeks(entries, date(2026, 1, 18))[0]['weeks']], ['26WK02'])


@skipUnless(shutil.which('node'), 'Node.js required')
class LaunchD1InputTests(TestCase):
    def run_js(self, body):
        result = subprocess.run(['node', '-e', "const m=require('./templates/forecast-math.js');" + body],
                                cwd=ROOT, capture_output=True, encoding='utf-8')
        self.assertEqual(result.returncode, 0, result.stderr)
        return json.loads(result.stdout)

    def test_page_honors_explicit_zero_instead_of_system_d1(self):
        result = self.run_js(r'''
const source=require('fs').readFileSync('templates/dashboard.js','utf8');
const declaration=source.split('\n').find(line=>line.includes('const systemD1='));
const d1Input=m.launchD1Input?m.launchD1Input('0'):null;
const run=new Function('get','parameterGross','d1Ratio','d1Input',declaration+';return plannedD1;');
console.log(JSON.stringify(run(()=>0,1000,.2,d1Input)));
''')
        self.assertEqual(result, 0)

    def test_blank_is_auto_but_zero_and_positive_integers_are_manual(self):
        result = self.run_js("console.log(JSON.stringify([null,undefined,'','  ','0','125','1e3'].map(m.launchD1Input)));")
        for item in result[:4]:
            self.assertFalse(item['supplied'])
            self.assertEqual(item['error'], '')
        for item, quantity in zip(result[4:], [0, 125, 1000]):
            self.assertTrue(item['supplied'])
            self.assertEqual(item['value'], quantity)
            self.assertEqual(item['error'], '')

    def test_invalid_manual_quantities_never_fall_back_to_auto(self):
        result = self.run_js("console.log(JSON.stringify([-1,'3.5','NaN','Infinity','9007199254740992','bad',true].map(m.launchD1Input)));")
        for item in result:
            self.assertTrue(item['supplied'])
            self.assertIsNone(item['value'])
            self.assertIn('D1预测大定', item['error'])


class GenerationOutcomeTests(TestCase):
    def run_main(self, severity=None, *, strict=False, auto_open=False, failure=None):
        args = Namespace(input=Path('input'), output=Path('output/report.html'), debug=False,
                         check=strict, no_open=not auto_open, disable_sales_forecast=False)
        def configure(*_):
            generator.DIAGNOSTIC_COUNTS.update(WARNING=0, ERROR=0)
            return Path('test.log')
        def build(*_, **__):
            if failure:
                raise failure
            if severity is not None:
                record = logging.LogRecord('audit', severity, __file__, 0, '生成阶段问题', (), None)
                generator.DiagnosticCollector().handle(record)
            return {'subjects': [], 'dashboards': {}}, ['生成阶段问题'] if severity else []
        with patch.object(generator, 'parse_args', return_value=args), \
             patch.object(generator, 'configure_runtime', side_effect=configure), \
             patch.object(generator, 'build', side_effect=build), \
             patch.object(generator, 'cleanup_input_cache') as cleanup, \
             patch.object(generator, 'AUTO_OPEN_HTML', True), \
             patch.object(generator.os, 'startfile', create=True) as startfile, \
             patch.object(generator.webbrowser, 'open', return_value=True) as browser, \
             self.assertLogs(generator.LOGGER, level='INFO') as captured:
            code = generator.main()
        cleanup.assert_called_once_with(args.input)
        return code, captured.output, startfile, browser

    def test_non_strict_run_with_module_or_excel_error_is_not_success(self):
        code, output, _, _ = self.run_main(logging.ERROR)
        self.assertEqual(code, 1)
        self.assertFalse(any('生成成功:' in message for message in output))

    def test_incomplete_build_does_not_auto_open_as_success(self):
        code, _, startfile, browser = self.run_main(logging.ERROR, auto_open=True)
        self.assertEqual(code, 1)
        startfile.assert_not_called()
        browser.assert_not_called()

    def test_warning_only_status_preserves_normal_and_strict_modes(self):
        self.assertEqual(self.run_main(logging.WARNING)[0], 0)
        self.assertEqual(self.run_main(logging.WARNING, strict=True)[0], 1)
        self.assertEqual(self.run_main()[0], 0)

    def test_fatal_exception_still_fails_and_cleans_cache(self):
        code, _, startfile, browser = self.run_main(failure=RuntimeError('failed'), auto_open=True)
        self.assertEqual(code, 1)
        startfile.assert_not_called()
        browser.assert_not_called()
