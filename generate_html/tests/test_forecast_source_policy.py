from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook
from core.excel import WorkbookItem
from core.forecast_summary import resolved_daily_rows, resolved_weekly_rows, table_records, DAILY_SHEET, MASTER_SHEET
from modules.sales_forecast import _merge_small_daily_history, _resolve_stage_candidate, _read_actual_profiles
from tools.refresh_sales_forecast_data import forecast_weekly_orders
from tests import test_public_forecast_tables as public_tests


class ForecastSourcePolicyTests(TestCase):
    def target(self):
        return dict(name='测试车', small_start_date='2026-01-01', small_end_date='2026-01-02',
                    launch_date='2026-01-03', end_date='2026-01-05')

    def row(self, day, kind, **values):
        return {'订单分析代际名': '测试车', '日期': day, '来源类型': kind, **values}

    def resolve(self, candidates, **extra):
        data = {'targets': [self.target()], **extra}
        return {str(r['日期'])[:10]: r for r in resolved_daily_rows(candidates, data, [], '整理表.xlsx', date(2026, 1, 10))}

    def test_launch_priority_is_per_field_and_preserves_zero(self):
        rows = self.resolve([
            self.row('2026-01-03', '首销订单', 大定=30, 小转大=10, 直接大定=20),
            self.row('2026-01-03', '历史首销', 大定=0, 小转大=None),
        ], mix_daily={('测试车', '2026-01-03', 'lock'): (7, '锁单选配比例.xlsx｜by天'),
                     ('测试车', '2026-01-03', 'gross'): (40, '大定选配比例.xlsx｜by天')})
        row = rows['2026-01-03']
        self.assertEqual((row['大定'], row['小转大'], row['直接大定'], row['交车锁单']), (0, 10, 20, 7))
        # A structure discrepancy does not erase the known quantities.
        self.assertIn('锁单选配比例', row['字段来源'])

    def test_post_launch_uses_mix_without_launch_components(self):
        rows = self.resolve([self.row('2026-01-06', '历史首销', 大定=99, 小转大=5, 直接大定=94),
                             self.row('2026-01-06', '逐日小订', 小订数量=8)],
                            mix_daily={('测试车', '2026-01-06', 'gross'): (40, '大定选配比例.xlsx'),
                                       ('测试车', '2026-01-06', 'lock'): (80, '锁单选配比例.xlsx')})
        row = rows['2026-01-06']
        self.assertEqual((row['大定'], row['交车锁单'], row['订单阶段']), (40, 80, '平销'))
        for field in ('小订数量', '小转大', '直接大定'):
            self.assertIsNone(row.get(field))

    def test_unknown_window_keeps_actuals(self):
        rows = self.resolve([], targets=[{'name': '测试车'}],
            mix_daily={('测试车', '2026-01-03', 'gross'): (40, '大定选配比例.xlsx')})
        self.assertEqual((rows['2026-01-03']['大定'], rows['2026-01-03']['订单阶段']), (40, '阶段未知'))

    def test_reservation_history_beats_mix_and_hourly_sum(self):
        def customize(data):
            data['actuals'][0]['small_daily_days'] = [{'date': '2026-01-01', 'orders': 99}, {'date': '2026-01-02', 'orders': 88}]
        book = public_tests.PublicForecastTablesTests().make_public(customize=customize)
        small = [r['小订数量'] for r in table_records(book, DAILY_SHEET) if r['订单阶段'] == '小订']
        self.assertEqual(small, [0, 10])
        book.close()

    def test_wrong_small_campaign_does_not_override(self):
        for start in ('2025-01-01', '2026-01-01'):
            profiles = [{'model': '测试车', 'small_daily_days': [{'date': '2026-01-01', 'orders': 99}]}]
            history = [{'model': '测试车', 'generation': '测试车', 'small_start_date': start,
                        'daily_actual': True, 'dates': ['2026-01-01', '2026-01-03'], 'daily_orders': [0, 88]}]
            _merge_small_daily_history(profiles, history, [self.target()], None)
            self.assertEqual(profiles[0]['small_daily_days'], [{'date': '2026-01-01', 'orders': 0 if start == '2026-01-01' else 99}])

    def test_cancel_does_not_fall_back_to_history_or_launch(self):
        candidates = {'history': {'days': [{'date': '2026-01-03', 'gross': 20, 'cancel': 3}]},
                      'launch': {'days': [{'date': '2026-01-03', 'gross': 30, 'cancel': 5}],
                                 'hourly_days': [{'date': '2026-01-03', 'gross': 9}]}}
        for stage in ('before', 'active', 'ended', 'unknown'):
            result = _resolve_stage_candidate(candidates, stage)
            self.assertEqual(result['days'][0]['gross'], 20)
            self.assertIsNone(result['days'][0]['cancel'])
            self.assertEqual(result['hourly_days'][0]['gross'], 9)

    def test_primary_and_secondary_quantities_are_not_duplicated(self):
        # Different identities must remain different rows even on the same day.
        result = resolved_daily_rows([
            self.row('2026-01-03', '历史首销', 大定=10)],
            {'targets': [self.target()], 'mix_daily': {('测试车 Pro', '2026-01-03', 'gross'): (20, '大定选配比例.xlsx')}},
            [], '整理表.xlsx', date(2026, 1, 10))
        self.assertEqual({r['订单分析代际名']: r['大定'] for r in result}, {'测试车': 10, '测试车 Pro': 20})

    def test_mix_only_target_lists_its_actual_source_not_history_file(self):
        def customize(data):
            data['targets'].append({'name': '新车2026款', 'has_small': False})
            data['mix_daily'] = {('新车2026款', '2026-01-03', 'gross'): (12, '大定选配比例.xlsx｜by天')}
            data['mix_weekly'] = {('新车2026款', '26WK01', 'lock'): (10, '锁单选配比例.xlsx｜by周')}
        book = public_tests.PublicForecastTablesTests().make_public(customize=customize)
        master = next(r for r in table_records(book, MASTER_SHEET) if r['订单分析代际名'] == '新车2026款')
        self.assertEqual(master['订单来源文件'], '大定选配比例.xlsx、锁单选配比例.xlsx')
        self.assertFalse(master['有小订'])
        book.close()

    def test_raw_week_total_survives_missing_stage_split(self):
        parts = [['测试车', '26WK01', stage, date(2025, 12, 29), date(2026, 1, 4),
                  None, None, None, None, None, None] for stage in ('首销', '平销')]
        rows = resolved_weekly_rows(parts, {('测试车', '26WK01', 'gross'): (100, '大定.xlsx')})
        self.assertEqual(rows[0][5], 100)
        self.assertIsNone(rows[0][11])
        self.assertIsNone(rows[0][14])

    def test_unknown_and_prelaunch_weeks_preserved_and_not_forecast_weeks(self):
        for window, stage in (({}, '阶段未知'), ({'launch_date': '2026-02-01', 'end_date': '2026-02-20'}, '首销前')):
            book = Workbook()
            sheet = book.active
            sheet.title = '测试车2026款by周'
            sheet.append(['指标', '统计类型', '分类', '26WK01'])
            sheet.append(['大定', '数量', '数量', 100])
            item = WorkbookItem(Path('大定选配比例.xlsx'), book)
            store = SimpleNamespace(find=lambda key: item if key == '大定选配比例' else None)
            data = {'targets': [{'name': '测试车2026款', **window}], 'actuals': []}
            dashboard = SimpleNamespace(views={'week': {'pages': {'预测方案': {'workspace': {'data': data}}}}})
            with patch('modules.sales_forecast._read_model_mapping', return_value={}):
                rows = forecast_weekly_orders(store, dashboard, date(2026, 1, 10))
            self.assertEqual((rows[0][2], rows[0][5]), (stage, 100))
            book.close()

    def test_order_mix_only_model_is_discovered_without_forecast_metadata(self):
        book = Workbook()
        sheet = book.active
        sheet.title = '测试车2026款by天'
        sheet.append(['指标', '统计类型', '分类', datetime(2026, 1, 1)])
        sheet.append(['大定', '数量', '数量', 12])
        item = WorkbookItem(Path('大定选配比例.xlsx'), book)
        store = SimpleNamespace(find=lambda key: item if key == '大定选配比例' else None)
        with patch('modules.sales_forecast._read_model_mapping', return_value={}):
            profiles, _ = _read_actual_profiles(store)
        self.assertIn('测试车2026款', [p['model'] for p in profiles])
        data = {'targets': [], 'actuals': []}
        dashboard = SimpleNamespace(views={'week': {'pages': {'预测方案': {'workspace': {'data': data}}}}})
        with patch('modules.sales_forecast._read_model_mapping', return_value={}):
            forecast_weekly_orders(store, dashboard, date(2026, 1, 10), collect_only=True)
        self.assertFalse(data['targets'][0]['has_small'])
        book.close()

    def test_mix_fields_choose_their_own_workbook_and_week_reuses_final_days(self):
        books = []
        items = {}
        for keyword, gross, lock in [('大定选配比例', 10, 2), ('锁单选配比例', 999, 4)]:
            book = Workbook()
            books.append(book)
            sheet = book.active
            sheet.title = '测试车2026款by天'
            dates = [date(2026, 1, day) for day in range(5, 12)]
            sheet.append(['指标', '统计类型', '分类', *dates])
            sheet.append(['大定', '数量', '数量', *([gross]*7)])
            sheet.append(['交车锁单', '数量', '数量', *([lock]*7)])
            items[keyword] = WorkbookItem(Path(keyword+'.xlsx'), book)
        seen = []
        def find(key):
            seen.append(key)
            return items.get(key)
        data = {'targets': [{'name': '测试车2026款', 'launch_date': '2026-01-05', 'end_date': '2026-01-11'}]}
        dashboard = SimpleNamespace(views={'week': {'pages': {'预测方案': {'workspace': {'data': data}}}}})
        with patch('modules.sales_forecast._read_model_mapping', return_value={}):
            forecast_weekly_orders(SimpleNamespace(find=find), dashboard, date(2026, 1, 12), collect_only=True)
            self.assertEqual(data['mix_daily'][('测试车2026款', '2026-01-05', 'gross')][0], 10)
            self.assertEqual(data['mix_daily'][('测试车2026款', '2026-01-05', 'lock')][0], 4)
            final = [{'订单分析代际名': '测试车2026款', '日期': date(2026, 1, day), '大定': 0, '交车锁单': 4}
                     for day in range(5, 12)]
            rows = forecast_weekly_orders(None, dashboard, date(2026, 1, 12), resolved_daily=final)
            self.assertEqual(rows[0][5:8], [0, None, 28])
            self.assertEqual(seen, ['大定选配比例', '锁单选配比例'])
        for book in books:
            book.close()
