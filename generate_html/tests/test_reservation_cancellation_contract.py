"""The daily export names and reservation-only cancellation source contract."""
from datetime import date, datetime, timedelta
from unittest import TestCase

from openpyxl import Workbook

from core.forecast_summary import (DAILY_SHEET, MASTER_SHEET, table_records,
                                   read_public_forecast, refresh_reference_tables)
from tools.refresh_sales_forecast_data import normalize_public_names, write_rows
from tests import test_public_forecast_tables as public_fixtures
from tests import test_forecast_source_backfill_e2e as source_fixtures


class ReservationCancellationContractTests(TestCase):
    model = '享界 G9 2027款'

    def public_book(self):
        def customize(data):
            data['actuals'][0].update(
                cancel_days=[{'date': '2026-01-02', 'cancel': 16, 'daily_cancel': 2},
                             {'date': '2026-01-03', 'cancel': 20, 'daily_cancel': 4}],
                cancel_source={'file': '小订退订分析.xlsx', 'sheet': '测试车_日度退订'})
        book = public_fixtures.PublicForecastTablesTests().make_public(customize=customize)
        self.addCleanup(book.close)
        def emit(name, headers, rows, percent_headers=None):
            del book[name]
            write_rows(book, name, headers, rows, 'DailyContract', percent_headers)
        normalize_public_names(book, {'targets': [{'name': '历史车', 'primary_generation': '测试车'}]}, emit)
        return book

    def test_daily_export_has_small_and_small_cancellation_adjacent(self):
        book = self.public_book()
        headers = [cell.value for cell in book[DAILY_SHEET][1]]
        self.assertEqual(headers[6:13], ['小订', '小订退订', '大定', '留存大定', '小转大', '直接大定', '交车锁单'])
        self.assertNotIn('小订数量', headers)
        self.assertNotIn('退订数量', headers)
        self.assertEqual(headers[:2], ['代际名', '二级代际名'])
        raw = [dict(zip(headers, values)) for values in book[DAILY_SHEET].iter_rows(min_row=2, values_only=True)]
        first_small = next(row for row in raw if row['日期'] == datetime(2026, 1, 1))
        launch = next(row for row in raw if row['日期'] == datetime(2026, 1, 3))
        self.assertEqual(first_small['小订'], 0)
        self.assertEqual(launch['小订退订'], 4)
        self.assertEqual(launch['大定'], 30)

    def test_web_reads_new_and_legacy_headers_identically(self):
        book = self.public_book()
        current = read_public_forecast(book, [], today=date(2026, 1, 4))
        self.assertEqual(current[0][0]['days'][0]['daily_cancel'], 4)
        self.assertEqual(current[0][0]['days'][0]['cancel'], 20)
        self.assertEqual(current[2][0]['daily_orders'], [0, 10])
        for cell in book[DAILY_SHEET][1]:
            cell.value = {'小订': '小订数量', '小订退订': '退订数量'}.get(cell.value, cell.value)
        legacy = read_public_forecast(book, [], today=date(2026, 1, 4))
        self.assertEqual(legacy, current)

    def test_new_header_zero_and_blank_are_not_replaced_by_legacy_cells(self):
        book = Workbook()
        self.addCleanup(book.close)
        book.active.title = DAILY_SHEET
        book.active.append(['小订', '小订退订', '小订数量', '退订数量'])
        book.active.append([0, None, 99, 88])
        row = table_records(book, DAILY_SHEET)[0]
        self.assertEqual(row['小订数量'], 0)
        self.assertIsNone(row['退订数量'])

    def source_book(self, headers, rows):
        def add_sources(folder):
            book = Workbook()
            sheet = book.active
            sheet.title = self.model+'_日度退订'
            sheet.append(['日度退订'])
            sheet.append(headers)
            for row in rows:
                sheet.append(row)
            book.save(folder/'鸿蒙智行小订退订分析.xlsx')
            book.close()
        return source_fixtures.SourceBackfillEndToEndTests.build(self, add_sources=add_sources)

    def test_source_small_cumulative_column_wins_and_is_sliced_to_launch_dates(self):
        book = self.source_book(
            ['累计总退订', '累计小订退', '当日退订数', '取消日期', '当日小订退'],
            [[999, 20, 99, datetime(2026, 8, 31), 2],
             [1999, 25, 199, datetime(2026, 9, 1), 5],
             [2999, 28, 299, datetime(2026, 9, 2), 3],
             [3999, 40, 399, datetime(2026, 9, 3), 12],
             [4999, 55, 499, datetime(2026, 9, 4), 15]])
        master = table_records(book, MASTER_SHEET)[0]
        self.assertEqual(master['总退订'], 40)
        curve = table_records(book, '累计退订率')[0]
        self.assertEqual([curve[f'D{i}'] for i in (1, 2, 3)], [.25, .28, .4])
        headers = [cell.value for cell in book[DAILY_SHEET][1]]
        rows = {str(row[headers.index('日期')])[:10]: dict(zip(headers, row))
                for row in book[DAILY_SHEET].iter_rows(min_row=2, values_only=True)}
        self.assertEqual([rows[f'2026-09-0{i}']['小订退订'] for i in (1, 2, 3, 4)], [5, 3, 12, 15])

    def test_generic_total_and_large_order_cancellation_do_not_become_small_cancellation(self):
        book = self.source_book(['取消日期', '累计总退订', '累计退订数', '当日退订数'],
                                [[datetime(2026, 9, i), i*100, i*100, 100] for i in (1, 2, 3)])
        master = table_records(book, MASTER_SHEET)[0]
        self.assertIsNone(master['总退订'])
        self.assertEqual(master['总大定'], 130)
        curve = table_records(book, '累计退订率')[0]
        self.assertTrue(all(curve.get(f'D{i}') is None for i in (1, 2, 3)))
        headers = [cell.value for cell in book[DAILY_SHEET][1]]
        self.assertTrue(all(row[headers.index('小订退订')] is None
                            for row in book[DAILY_SHEET].iter_rows(min_row=2, values_only=True)))

    def test_longer_other_reference_does_not_extend_cancel_rate_past_launch_end(self):
        book = Workbook()
        self.addCleanup(book.close)
        book.active.title = MASTER_SHEET
        fields = ['历史传播名', '订单分析代际名', '首销开始', '首销结束', '首销天数', '总小订', '总退订']
        book.active.append(fields)
        book.active.append(['测试车', '测试车', '2026-01-03', '2026-01-04', 2, 100, None])
        stamps = [(date(2026, 1, 2)+timedelta(days=i)).isoformat() for i in range(5)]
        values = [20, 25, 28, 40, 55]
        data = {'resolved_cancel_cumulative': {('测试车', day): value for day, value in zip(stamps, values)},
                'reference_daily': {'小转大当日数量': {'测试车': [1, 1, 1, 1]}}}
        daily = [{'订单分析代际名': '测试车', '日期': day, '退订数量': value}
                 for day, value in zip(stamps[1:], [5, 3, 12, 15])]
        def emit(name, headers, rows, percent_headers=None):
            if name in book.sheetnames:
                del book[name]
            write_rows(book, name, headers, rows, 'CancelStage', percent_headers)
        refresh_reference_tables(book, data, emit, daily, date(2026, 1, 9))
        curve = table_records(book, '累计退订率')[0]
        self.assertEqual([curve[f'D{i}'] for i in (1, 2, 3, 4)], [.25, .28, None, None])
