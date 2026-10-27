"""Keep reservation and launch cancellation sources on their own date axes."""
from datetime import date, datetime
from pathlib import Path
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from core.forecast_summary import (
    DAILY_SHEET, D12_SHEET, LAUNCH_CANCEL_SHEET, SMALL_CANCEL_SHEET,
    SMALL_CANCEL_DAILY_SHEET, LAUNCH_CANCEL_DAILY_SHEET, summary_scope, table_records, read_public_forecast,
)
from modules.sales_forecast import _read_history
from tools.refresh_sales_forecast_data import extract_quantity_table, workbook_header_index
from tests import test_forecast_source_backfill_e2e as source_fixtures
from tests import test_forecast_reference_backfill as reference_fixtures


class CancelStageCurvesTests(TestCase):
    model = '享界 G9 2027款'

    def build(self, small=(2, 0), launch=(7, None, 0), *, include_launch=True,
              fallback=True, maintained=None, small_headers=None, **kwargs):
        def edit(path):
            book = load_workbook(path)
            if include_launch:
                book['退订进度'].title = '首销期退订进度'
            else:
                del book['退订进度']
            if small is not None:
                sheet = book.create_sheet('小订期退订进度', 1)
                sheet.append(['代际名', *(small_headers or [f'D{i+1}' for i in range(len(small))])])
                sheet.append([self.model, *small])
                # Ratio blocks are not a second set of daily quantities.
                sheet.append(['代际名', 'D1', 'D2'])
                sheet.append([self.model, .88, .99])
            book.save(path)
            book.close()

        def sources(folder):
            path = folder/'鸿蒙智行小订退订分析.xlsx'
            book = load_workbook(path)
            sheet = book.active
            # Deliberately contradictory fallback tests source priority and zero.
            sheet.append([datetime(2026, 8, 30), 5, 5])
            sheet.append([datetime(2026, 8, 31), 6, 11])
            sheet.cell(2, 4, '当日退订数')
            sheet.cell(2, 5, '累计退订数')
            for row in range(3, sheet.max_row+1):
                sheet.cell(row, 4, 999)
                sheet.cell(row, 5, 9999)
            book.save(path)
            book.close()
            if not fallback:
                path.unlink()

        return source_fixtures.SourceBackfillEndToEndTests.build(self, source_quantities={'退订进度': launch},
            maintained=maintained, source_edit=edit, add_sources=sources, **kwargs)

    def values(self, book, name, count=2):
        row = table_records(book, name)[0]
        return [row.get(f'D{i+1}') for i in range(count)]

    def test_new_names_and_independent_stage_d1_preserve_primary_zero(self):
        book = self.build()
        self.assertEqual(self.values(book, SMALL_CANCEL_SHEET), [.02, .02])
        self.assertEqual(self.values(book, LAUNCH_CANCEL_SHEET, 3), [.18, .20, .20])
        self.assertNotIn('累计退订率', book.sheetnames)
        self.assertIn(SMALL_CANCEL_DAILY_SHEET, book.sheetnames)
        self.assertEqual(book.sheetnames[book.sheetnames.index('小订当日数量')+1], SMALL_CANCEL_DAILY_SHEET)
        self.assertEqual(self.values(book, SMALL_CANCEL_DAILY_SHEET), [2, 0])
        self.assertEqual(self.values(book, LAUNCH_CANCEL_DAILY_SHEET, 3), [18, 2, 0])
        self.assertNotIn('退订当日数量', book.sheetnames)
        for name in (SMALL_CANCEL_DAILY_SHEET, LAUNCH_CANCEL_DAILY_SHEET):
            headers = [c.value for c in book[name][1]]
            self.assertNotIn('%', book[name].cell(2, headers.index('D1')+1).number_format)
        self.assertEqual(book.sheetnames[-1], LAUNCH_CANCEL_SHEET)
        rows = {r['日期'].date(): r for r in table_records(book, DAILY_SHEET)}
        self.assertEqual(rows[date(2026, 8, 30)]['退订数量'], 2)
        self.assertEqual(rows[date(2026, 8, 31)]['退订数量'], 0)
        self.assertIn('小订期退订进度', rows[date(2026, 8, 31)]['字段来源'])
        self.assertEqual(table_records(book, D12_SHEET)[0]['D1退订'], 7)
        for name in (SMALL_CANCEL_SHEET, LAUNCH_CANCEL_SHEET):
            headers = [c.value for c in book[name][1]]
            self.assertIn('%', book[name].cell(2, headers.index('D1')+1).number_format)

    def test_missing_small_cell_uses_only_small_cancellations(self):
        book = self.build(small=(None, 0))
        self.assertEqual(self.values(book, SMALL_CANCEL_SHEET), [.05, .05])
        self.assertEqual(self.values(book, SMALL_CANCEL_DAILY_SHEET), [5, 0])

    def test_missing_small_sheet_falls_back_on_small_calendar_window(self):
        book = self.build(small=None)
        self.assertEqual(self.values(book, SMALL_CANCEL_SHEET), [.05, .11])
        self.assertEqual(self.values(book, SMALL_CANCEL_DAILY_SHEET), [5, 6])
        self.assertEqual(self.values(book, LAUNCH_CANCEL_SHEET, 3), [.18, .20, .20])

    def test_missing_launch_sheet_does_not_select_small_sheet(self):
        book = self.build(include_launch=False)
        self.assertEqual(self.values(book, LAUNCH_CANCEL_SHEET, 3), [.01, .03, .06])
        self.assertEqual(self.values(book, SMALL_CANCEL_SHEET), [.02, .02])

    def test_missing_both_sheets_uses_analysis_without_name_error(self):
        book = self.build(small=None, include_launch=False)
        self.assertEqual(self.values(book, SMALL_CANCEL_SHEET), [.05, .11])
        self.assertEqual(self.values(book, LAUNCH_CANCEL_SHEET, 3), [.01, .03, .06])

    def test_missing_all_cancellation_sources_stays_unknown(self):
        book = self.build(small=None, include_launch=False, fallback=False)
        self.assertEqual(self.values(book, SMALL_CANCEL_SHEET), [None, None])
        self.assertEqual(self.values(book, LAUNCH_CANCEL_SHEET, 3), [None]*3)
        self.assertEqual(self.values(book, SMALL_CANCEL_DAILY_SHEET), [None, None])

    def test_small_window_stops_before_later_source_days(self):
        book = self.build(small=(2, 3, 800))
        self.assertEqual(self.values(book, SMALL_CANCEL_SHEET, 3), [.02, .05, None])
        self.assertNotIn('D3', [c.value for c in book[SMALL_CANCEL_SHEET][1]])
        self.assertEqual(self.values(book, SMALL_CANCEL_DAILY_SHEET), [2, 3])
        self.assertNotIn('D3', [c.value for c in book[SMALL_CANCEL_DAILY_SHEET][1]])

    def test_sparse_small_days_keep_gap_not_shift_or_fill_zero(self):
        book = self.build(small=(2, 3), small_headers=['D1', 'D3'], fallback=False)
        self.assertEqual(self.values(book, SMALL_CANCEL_SHEET), [.02, None])
        self.assertEqual(self.values(book, SMALL_CANCEL_DAILY_SHEET), [2, None])

    def test_overlap_date_does_not_reuse_launch_quantity_for_small_curve(self):
        book = self.build(small=(2, 3, 0), maintained={'小订结束日期': datetime(2026, 9, 1)})
        self.assertEqual(self.values(book, SMALL_CANCEL_SHEET, 3), [.02, .05, .05])
        self.assertEqual(self.values(book, SMALL_CANCEL_DAILY_SHEET, 3), [2, 3, 0])
        self.assertEqual(table_records(book, D12_SHEET)[0]['D1退订'], 7)

    def test_missing_launch_value_on_overlap_does_not_take_small_stage_value(self):
        book = self.build(small=(2, 3, 50), include_launch=False,
                          maintained={'小订结束日期': datetime(2026, 9, 1)})
        self.assertEqual(self.values(book, SMALL_CANCEL_SHEET, 3), [.02, .05, .55])
        self.assertEqual(table_records(book, D12_SHEET)[0]['D1退订'], 1)

    def test_legacy_source_and_summary_readers_remain_compatible(self):
        book = source_fixtures.SourceBackfillEndToEndTests.build(self, source_quantities={'退订进度': [7, None, 0]})
        self.assertEqual(self.values(book, LAUNCH_CANCEL_SHEET, 3), [.07, .09, .09])
        with patch('modules.sales_forecast._read_model_mapping', return_value={}), \
             patch('modules.sales_forecast._read_model_master', return_value={}):
            with summary_scope(book._test_path, workbook=book):
                current = _read_history()[1]
            book[LAUNCH_CANCEL_SHEET].title = '累计退订率'
            with summary_scope(book._test_path, workbook=book):
                legacy = _read_history()[1]
        self.assertEqual(current, legacy)
        self.assertEqual(current[0]['cancel_progress'], [.07, .09, .09])

    def test_optional_missing_denominator_does_not_fabricate_rate(self):
        book = reference_fixtures.ForecastReferenceBackfillTests.build(self, small_total=None,
            raw_fields={'small_daily_days': [], 'small_daily_sources': {}},
            history={SMALL_CANCEL_DAILY_SHEET: {'测试车': [2, 0]}})
        self.assertEqual(self.values(book, SMALL_CANCEL_SHEET), [None, None])
        self.assertEqual(self.values(book, SMALL_CANCEL_DAILY_SHEET), [2, 0])

    def test_matching_stage_name_wins_even_when_legacy_and_small_come_first(self):
        book = Workbook()
        self.addCleanup(book.close)
        book.remove(book.active)
        for name, value in [('小订期退订进度', 2), ('退订进度', 99), ('首销期退订进度', 7)]:
            sheet = book.create_sheet(name)
            sheet.append(['代际名', 'D1'])
            sheet.append([self.model, value])
        index = workbook_header_index(book)
        self.assertEqual(extract_quantity_table(book, index, 'cancel')[0][self.model], [7])
        self.assertEqual(extract_quantity_table(book, index, 'small_cancel')[0][self.model], [2])

    def test_small_forecast_reader_uses_its_curve_and_excludes_unfinished_day(self):
        book = self.build()
        with patch('modules.sales_forecast._read_model_mapping', return_value={}), \
             patch('modules.sales_forecast._read_model_master', return_value={}):
            _, _, small, _ = read_public_forecast(book, [], date(2026, 8, 31))
        self.assertEqual(small[0]['cancel_progress'], [.02])
