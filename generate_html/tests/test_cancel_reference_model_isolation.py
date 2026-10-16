from datetime import date
from unittest import TestCase

from openpyxl import Workbook

from core.forecast_summary import (
    D12_SHEET, MASTER_SHEET, refresh_reference_tables, table_records,
)


class CancelReferenceModelIsolationTests(TestCase):
    def build(self, *, master_order=('甲车', '乙车'), profile_order=('甲车', '乙车'),
              first_daily=10, previous=None):
        book = Workbook()
        self.addCleanup(book.close)
        book.active.title = MASTER_SHEET
        headers = ['历史传播名', '订单分析代际名', '首销开始', '首销结束',
                   '首销天数', '总小订', '总退订', '退订率']
        book.active.append(headers)
        totals = {'甲车': 1000, '乙车': 2000}
        daily_counts = {'甲车': [first_daily, 2, 0], '乙车': [4, 3, 0]}
        cumulative = {'甲车': 40, '乙车': 100}
        for model in master_order:
            book.active.append([model, model, '2026-01-03', '2026-01-05', 3,
                                totals[model], None, None])
        days = ['2026-01-03', '2026-01-04', '2026-01-05']
        anchors = {(model, days[0]): value for model, value in cumulative.items()}
        if previous is not None:
            anchors.update({(model, '2026-01-02'): value
                            for model, value in previous.items()})
        data = {
            'actuals': [
                {'model': model, 'cancel_days': [
                    {'date': days[0], 'cancel': cumulative[model],
                     'daily_cancel': daily_counts[model][0]},
                ]}
                for model in profile_order
            ],
            'resolved_cancel_cumulative': anchors,
            'resolved_cancel_history_days': {(model, day)
                                             for model in totals for day in days},
        }
        daily = [{'订单分析代际名': model, '日期': day, '退订数量': count}
                 for model, counts in daily_counts.items()
                 for day, count in zip(days, counts)]

        def emit(name, fields, rows, percent_headers=None):
            if name in book.sheetnames:
                del book[name]
            sheet = book.create_sheet(name)
            sheet.append(fields)
            for row in rows:
                sheet.append(row)

        refresh_reference_tables(book, data, emit, daily, date(2026, 1, 6))
        return book

    def assert_cancellation_results(self, book):
        masters = {row['历史传播名']: row for row in table_records(book, MASTER_SHEET)}
        curves = {row['传播名']: row for row in table_records(book, '累计退订率')}
        metrics = {row['传播名']: row for row in table_records(book, D12_SHEET)}
        for model, counts, denominator in (
                ('甲车', [40, 42, 42], 1000), ('乙车', [100, 103, 103], 2000)):
            with self.subTest(model=model):
                for i, value in enumerate(counts, 1):
                    self.assertAlmostEqual(curves[model][f'D{i}'], value / denominator)
                self.assertEqual(masters[model]['总退订'], counts[-1])
                self.assertAlmostEqual(masters[model]['退订率'], counts[-1] / denominator)
                self.assertEqual(metrics[model]['D2退订'], 2 if model == '甲车' else 3)

    def test_each_model_uses_its_own_d1_daily_count_regardless_of_row_order(self):
        for masters in (('甲车', '乙车'), ('乙车', '甲车')):
            for profiles in (('甲车', '乙车'), ('乙车', '甲车')):
                with self.subTest(master_order=masters, profile_order=profiles):
                    self.assert_cancellation_results(
                        self.build(master_order=masters, profile_order=profiles))

    def test_missing_model_profile_keeps_its_anchor_without_borrowing_another_model(self):
        self.assert_cancellation_results(self.build(profile_order=('乙车',)))

    def test_zero_d1_daily_count_remains_valid_for_its_own_model(self):
        book = self.build(first_daily=0)
        self.assert_cancellation_results(book)
        metrics = {row['传播名']: row for row in table_records(book, D12_SHEET)}
        self.assertEqual(metrics['甲车']['D1退订'], 0)

    def test_known_previous_day_cumulative_takes_precedence(self):
        for daily, previous in ((10, 30), (40, 0)):
            with self.subTest(previous=previous):
                self.assert_cancellation_results(
                    self.build(first_daily=daily, previous={'甲车': previous, '乙车': 96}))
