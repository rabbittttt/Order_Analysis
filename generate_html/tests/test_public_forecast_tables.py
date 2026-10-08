from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch
from tempfile import TemporaryDirectory

from openpyxl import Workbook

from core.excel import WorkbookItem
from core.forecast_summary import (
    GUIDE_SHEET, MASTER_SHEET, DAILY_SHEET, SMALL_HOURLY_SHEET, LAUNCH_HOURLY_SHEET,
    WEEKLY_SHEET, SMALL_CURVE_SHEET, SMALL_DAILY_SHEET, SMALL_REFERENCE_FIELDS, D12_SHEET,
    public_forecast_tables, read_public_forecast, visible_target_names, table_records,
)
from tools.refresh_sales_forecast_data import forecast_weekly_orders, write_rows
from modules.sales_forecast import _history_item, _target_options


class PublicForecastTablesTests(TestCase):
    def test_separate_publication_periods_survive_public_workbook(self):
        book = self.make_public(customize=lambda data:data['targets'][0].update(
            small_period='下午',launch_period='晚上'))
        record = table_records(book, MASTER_SHEET)[0]
        self.assertEqual(record['小订发布时段'], '下午')
        self.assertEqual(record['首销发布时段'], '晚上')
        _, _, small, _ = read_public_forecast(book, [])
        self.assertEqual(small[0]['launch_period'], '下午')
        book.close()

    def test_legacy_period_does_not_fill_either_stage(self):
        book = self.make_public(master_periods={'发布时段': '晚上'})
        record = table_records(book, MASTER_SHEET)[0]
        self.assertEqual(record['小订发布时段'], '未维护')
        self.assertEqual(record['首销发布时段'], '未维护')
        self.assertNotIn('发布时段', record)
        # Also exercise reading a legacy field from an existing visible summary.
        sheet = book[MASTER_SHEET]
        col = sheet.max_column + 1
        sheet.cell(1, col, '发布时段')
        sheet.cell(2, col, '晚上')
        for i, cell in enumerate(sheet[1], 1):
            if cell.value == '小订发布时段':
                sheet.cell(2, i).value = None
        _, _, small, _ = read_public_forecast(book, [])
        self.assertEqual(small[0]['launch_period'], '未维护')
        book.close()

    def test_stage_periods_never_borrow_from_each_other(self):
        for small_period, launch_period in ((None, '晚上'), ('下午', None), ('下午', '晚上')):
            with self.subTest(small=small_period, launch=launch_period):
                attrs = {'发布时段': '上午', '小订发布时段': small_period, '首销发布时段': launch_period}
                book = self.make_public(master_periods=attrs)
                record = table_records(book, MASTER_SHEET)[0]
                self.assertEqual(record['小订发布时段'], small_period or '未维护')
                self.assertEqual(record['首销发布时段'], launch_period or '未维护')
                _, _, small, _ = read_public_forecast(book, [])
                self.assertEqual(small[0]['launch_period'], small_period or '未维护')
                target = _target_options([], [{'model': '测试车', 'launch_date': '2026-01-03'}], {},
                    today=date(2026, 1, 4), model_master={'测试车': {**attrs, '首销天数': 7}})[0]
                self.assertEqual(target['small_period'], small_period or '未维护')
                self.assertEqual(target['launch_period'], launch_period or '未维护')
                self.assertEqual(_history_item({'传播名': '测试车', **attrs}, True)['launch_period'],
                                 launch_period or '未维护')
                book.close()

    def make_public(self, target_start='2026-01-03', customize=None, as_of=date(2026, 1, 4), master_periods=None):
        book = Workbook()
        book.active.title = MASTER_SHEET
        book.active.append(['历史传播名', '订单分析代际名', '产品档位'])
        book.active.append(['历史车', '测试车', '中型SUV'])
        for col, (field, value) in enumerate((master_periods or {}).items(), 4):
            book.active.cell(1, col, field)
            book.active.cell(2, col, value)
        base = book.create_sheet('预测基准总表')
        base.append(['传播名', '代际名', '总大定', '发布日', '首销截止', '总小订'])
        base.append(['历史车', '测试车', 20, datetime(2026, 1, 3), datetime(2026, 1, 9), 100])
        metrics = book.create_sheet('D1_D2预测指标')
        metrics.append(['传播名', 'D1大定'])
        metrics.append(['历史车', 20])
        profile = {'model': '测试车', 'total_small': 100, 'stage_profiles': {},
                   'launch_total_small': 100, 'launch_total_small_valid': True,
                   'day_source': {'file': '首销期订单节奏.xlsx', 'sheet': '测试车by天'},
                   'days': [{'date': '2026-01-03', 'day': 'D1', 'gross': 30, 'small_to_big': 10, 'direct': 20}],
                   'small_daily_days': [{'date': '2026-01-01', 'orders': 0}, {'date': '2026-01-02', 'orders': 10}],
                   'small_hourly_days': [{'date': '2026-01-02', 'orders': 10, 'hours': [{'hour': 15, 'orders': 4}, {'hour': 16, 'orders': 6}]}]}
        for stage, gross, owner in [('before', 20, 'history'), ('active', 30, 'launch'), ('ended', 20, 'history')]:
            profile['stage_profiles'][stage] = {'total_small': 100, 'total_small_source': 'history', 'selected_source': owner,
                'days': [{'date': '2026-01-03', 'day': 'D1', 'gross': gross, 'small_to_big': 10, 'direct': gross - 10,
                          '_field_sources': {'gross': owner, 'small_to_big': owner, 'direct': owner}}]}
        data = {'targets': [{'name': '测试车', 'history_model': '历史车', 'launch_date': target_start, 'end_date': '2026-01-09',
                            'small_start_date': '2026-01-01', 'small_end_date': '2026-01-02', 'days': 7}],
                'actuals': [profile], 'history': [{'model': '历史车', 'generation': '测试车', 'launch_date': '2026-01-03', 'daily_orders': [20]}], 'steady_history': [],
                'small_order_history': [{'model': '历史车', 'generation': '测试车', 'event_id': 'event1',
                    'small_start_date': '2026-01-01', 'small_end_date': '2026-01-02', 'days': 2, 'total': 100,
                    'daily_actual': True, 'source_sheet': '小订by天', 'total_complete': True,
                    'dates': ['2026-01-01', '2026-01-02'], 'daily_orders': [0, 10],
                    'small_progress': [0, .1], 'standard_progress': [0, .1, .8, 1], 'small_hourly_curve': []}]}
        if customize:
            customize(data)
        def emit(name, headers, rows, percent_headers=None):
            if name in book.sheetnames:
                del book[name]
            write_rows(book, name, headers, rows, 'PublicTest', percent_headers)
        public_forecast_tables(book, data, emit, as_of_date=as_of)
        return book

    def test_no_cache_merged_tables_and_typed_values(self):
        book = self.make_public()
        self.assertEqual(book.sheetnames[:6], [GUIDE_SHEET, MASTER_SHEET, SMALL_HOURLY_SHEET, LAUNCH_HOURLY_SHEET, DAILY_SHEET, WEEKLY_SHEET])
        self.assertNotIn('预测基准总表', book.sheetnames)
        self.assertTrue(all(s.sheet_state == 'visible' for s in book))
        self.assertEqual(visible_target_names(book), ['测试车'])
        headers = [c.value for c in book[DAILY_SHEET][1]]
        self.assertNotIn('累计退订', headers)
        self.assertNotIn('累计退订率', headers)
        self.assertIsInstance(book[DAILY_SHEET].cell(2, headers.index('日期') + 1).value, datetime)
        book.close()

    def test_public_cells_are_authoritative_and_partial_history_is_not_padded(self):
        book = self.make_public()
        profiles, windows, small, _ = read_public_forecast(book, [])
        self.assertEqual(small[0]['daily_orders'], [0, 10])
        self.assertEqual(small[0]['standard_progress'], [0, .1])
        self.assertEqual(small[0]['small_hourly_curve'], [])
        self.assertEqual(profiles[0]['days'][0]['gross'], 30)
        self.assertEqual(profiles[0]['small_hourly_days'][0]['hours'][0]['hour'], 15)
        sheet = book[DAILY_SHEET]
        headers = [c.value for c in sheet[1]]
        for row in sheet.iter_rows(min_row=2):
            if '首销' in row[headers.index('订单阶段')].value:
                row[headers.index('大定')].value = 99
        updated = read_public_forecast(book, [])[0]
        self.assertEqual(updated[0]['days'][0]['gross'], 99)
        history = [{'model': '历史车'}]
        read_public_forecast(book, history)
        self.assertEqual(history[0]['daily_orders'], [99])
        book.close()

    def test_current_window_does_not_overwrite_historical_reference_dates(self):
        book = self.make_public(target_start='2026-01-05')
        rows = list(book[MASTER_SHEET].values)
        record = dict(zip(rows[0], rows[1]))
        self.assertEqual(record['首销开始'], datetime(2026, 1, 3))
        current = next(dict(zip(rows[0], r)) for r in rows[1:] if dict(zip(rows[0], r)).get('订单来源文件'))
        self.assertEqual(current['首销开始'], datetime(2026, 1, 5))
        windows = read_public_forecast(book, [])[1]
        self.assertEqual(next(iter(windows.values()))['launch_date'], '2026-01-05')
        book.close()

    def test_dashboard_reads_public_workbook_without_original_files(self):
        from modules.sales_forecast import SalesForecastModule
        from core.models import Subject
        book = self.make_public()
        with TemporaryDirectory() as folder:
            path = Path(folder) / 'summary.xlsx'
            book.save(path)
            module = SalesForecastModule()
            module.summary_path = path
            module.as_of_date = date(2026, 1, 4)
            with patch('modules.sales_forecast._read_actual_profiles', side_effect=AssertionError('raw orders must not be loaded')), \
                 patch('modules.sales_forecast._read_small_order_history', side_effect=AssertionError('raw history must not be loaded')), \
                 patch('modules.sales_forecast._read_steady_history', side_effect=AssertionError('raw steady must not be loaded')):
                dashboard = module.build(None, Subject('summary', '鸿蒙智行', 'group'))
            data = dashboard.views['week']['pages']['预测方案']['workspace']['data']
            self.assertEqual(data['actuals'][0]['days'][0]['gross'], 30)
            self.assertEqual(data['small_order_history'][0]['daily_orders'], [0, 10])
        book.close()

    def test_missing_launch_date_keeps_actuals_without_relative_cancel_curve(self):
        from core.model_identity import model_key
        for invalid in (None, '', 'not-a-date'):
            with self.subTest(start=invalid):
                book = self.make_public()
                sheet = book[MASTER_SHEET]
                headers = [c.value for c in sheet[1]]
                sheet.cell(2, headers.index('首销开始') + 1).value = invalid
                book['累计退订率'].cell(2, 2).value = .1
                history = [{'model': '历史车', 'generation': '测试车'}]
                profiles, windows, _, _ = read_public_forecast(book, history)
                self.assertEqual(windows[model_key('测试车')]['launch_date'], '')
                self.assertEqual(profiles[0]['days'][0]['date'], '2026-01-03')
                self.assertEqual(profiles[0]['days'][0]['gross'], 30)
                self.assertIsNone(profiles[0]['days'][0].get('cancel'))
                self.assertEqual(history[0]['daily_orders'], [])
                book.close()

    def test_missing_launch_date_does_not_disable_shared_forecast_dashboard(self):
        from modules.sales_forecast import SalesForecastModule
        from core.models import Subject
        book = self.make_public()
        for name in (MASTER_SHEET, DAILY_SHEET):
            sheet = book[name]
            headers = [c.value for c in sheet[1]]
            for values in list(sheet.values)[1:]:
                clone = list(values)
                clone[headers.index('订单分析代际名')] = '缺少日期车'
                clone[headers.index('历史传播名')] = '缺少日期车'
                if name == MASTER_SHEET:
                    clone[headers.index('首销开始')] = None
                sheet.append(clone)
        with TemporaryDirectory() as folder:
            path = Path(folder) / 'summary.xlsx'
            book.save(path)
            module = SalesForecastModule()
            module.summary_path = path
            module.as_of_date = date(2026, 1, 4)
            dashboard = module.build(None, Subject('summary', '鸿蒙智行', 'group'))
            data = dashboard.views['week']['pages']['预测方案']['workspace']['data']
            targets = {t['name']: t for t in data['targets']}
            self.assertFalse(targets['测试车']['data_error'])
            self.assertIn('首销开始日期缺失或无法解析', targets['缺少日期车']['hard_errors'])
            actuals = {p['model']: p for p in data['actuals']}
            self.assertEqual(actuals['缺少日期车']['days'][0]['gross'], 30)
        book.close()

    def test_valid_launch_date_still_applies_relative_cancel_curve(self):
        book = self.make_public()
        book['累计退订率'].cell(2, 2).value = .1
        history = [{'model': '历史车', 'generation': '测试车'}]
        profiles, _, _, _ = read_public_forecast(book, history)
        self.assertEqual(profiles[0]['days'][0]['cancel'], 10)
        self.assertEqual(history[0]['daily_orders'], [30])
        book.close()

    def test_invalid_window_dates_are_not_calendar_anchors(self):
        from core.model_identity import model_key
        for column, field in (('首销结束', 'end_date'), ('小订开始', 'small_start_date'),
                              ('小订结束', 'small_end_date')):
            with self.subTest(column=column):
                book = self.make_public()
                sheet = book[MASTER_SHEET]
                headers = [c.value for c in sheet[1]]
                sheet.cell(2, headers.index(column) + 1).value = 'not-a-date'
                _, windows, _, _ = read_public_forecast(book, [])
                self.assertEqual(windows[model_key('测试车')][field], '')
                book.close()

    def test_small_tables_match_launch_curve_format_and_order(self):
        book = self.make_public()
        self.assertEqual(book.sheetnames[5:9], [WEEKLY_SHEET, SMALL_DAILY_SHEET, SMALL_CURVE_SHEET, D12_SHEET])
        self.assertNotIn('D1_D2预测指标', book.sheetnames)
        self.assertNotIn('小订参考曲线', book.sheetnames)
        self.assertNotIn('小订来源总量', book.sheetnames)
        for name in (SMALL_CURVE_SHEET, SMALL_DAILY_SHEET):
            headers = [c.value for c in book[name][1]]
            self.assertEqual(headers[0], '传播名')
            self.assertTrue(all(h.startswith('D') and h[1:].isdigit() for h in headers[1:]))
        counts = table_records(book, SMALL_DAILY_SHEET)[0]
        curve = table_records(book, SMALL_CURVE_SHEET)[0]
        self.assertEqual((counts['D1'], counts['D2']), (0, 10))
        self.assertEqual((curve['D1'], curve['D2']), (0, .1))
        self.assertEqual(book[SMALL_CURVE_SHEET].cell(2, 2).number_format,
                         book['直接大定累计完成度'].cell(2, 2).number_format)
        self.assertEqual(book[SMALL_DAILY_SHEET].cell(2, 2).number_format,
                         book['直接大定当日数量'].cell(2, 2).number_format)
        book.close()

    def test_small_reference_metadata_survives_narrow_curve_table(self):
        def customize(data):
            data['small_order_history'][0].update(leads=500, heat=25)
        book = self.make_public(customize=customize)
        master = table_records(book, MASTER_SHEET)[0]
        self.assertEqual((master['线索数'], master['热度']), (500, 25))
        item = read_public_forecast(book, [])[2][0]
        self.assertEqual((item['leads'], item['heat'], item['total']), (500, 25, 100))
        self.assertEqual(item['daily_orders'], [0, 10])
        book.close()

    def test_removed_columns_are_not_exported(self):
        book = self.make_public()
        prohibited = {'适用取数阶段', '小订参考累计完成度', '小订标准进度', '真实逐日',
                      '已结束日', '参考事件', '数据状态', '用途', '截至末小时小转大',
                      '截至末小时直接大定', '截至末小时交车锁单', '平销完整周参考'}
        for sheet in book:
            self.assertFalse(prohibited.intersection(c.value for c in sheet[1]), sheet.title)
        self.assertEqual(book[MASTER_SHEET].max_column, 29 + len(SMALL_REFERENCE_FIELDS))
        self.assertIn('小订发布时段', [c.value for c in book[MASTER_SHEET][1]])
        self.assertIn('首销发布时段', [c.value for c in book[MASTER_SHEET][1]])
        self.assertIn("有小订", [c.value for c in book[MASTER_SHEET][1]])
        for name in (SMALL_HOURLY_SHEET, LAUNCH_HOURLY_SHEET):
            self.assertEqual(book[name].max_column, 5 if name == SMALL_HOURLY_SHEET else 8)
            self.assertEqual(book[name].cell(1, 5).value, '来源文件')
        book.close()

    def test_partial_small_data_preserves_blanks_without_inventing_terminal(self):
        def customize(data):
            item = data['small_order_history'][0]
            item.update(daily_orders=[0, None], total=0, total_complete=False,
                        total_source='小订by天已填累计（非终值）')
            data['actuals'][0]['small_daily_days'][1]['orders'] = None
        book = self.make_public(customize=customize)
        item = read_public_forecast(book, [])[2][0]
        self.assertEqual(item['daily_orders'], [0, None])
        self.assertEqual(item['small_progress'], [None, None])
        self.assertFalse(item['total_complete'])
        self.assertEqual(len(item['standard_progress']), 4)
        book.close()

    def test_synthetic_curve_is_not_exported_as_real_daily_orders(self):
        def customize(data):
            item = data['small_order_history'][0]
            item.update(daily_actual=False, source_sheet='小订进度', standard_progress=[.333, 1])
            item.pop('total_complete', None)
            data['actuals'][0]['small_daily_days'] = []
        book = self.make_public(customize=customize)
        item = read_public_forecast(book, [])[2][0]
        self.assertEqual(item['total'], 100)
        self.assertEqual(item['daily_orders'], [33, 67])
        self.assertEqual(item['d1_share'], .333)
        self.assertFalse(item['daily_actual'])
        self.assertEqual(item['small_progress'], [])
        from core.forecast_summary import table_records
        self.assertFalse(any(r['小订数量'] is not None for r in table_records(book, DAILY_SHEET)))
        book.close()

    def test_hourly_components_live_in_hourly_table_not_duplicate_daily_rows(self):
        def customize(data):
            data['actuals'][0]['hourly_days'] = [{'date': '2026-01-03', 'last_hour': 15,
                'gross': 9, 'small_to_big': 0, 'direct': 9, 'lock': None,
                'hours': [{'hour': 15, 'gross': 9}]}]
            data['actuals'][0]['hour_source'] = {'file': '首销期订单节奏.xlsx', 'sheet': '测试车by时'}
        book = self.make_public(customize=customize)
        bucket = read_public_forecast(book, [])[0][0]['hourly_days'][0]
        self.assertEqual((bucket['gross'], bucket['small_to_big'], bucket['direct'], bucket['lock']), (9, 0, 9, None))
        self.assertEqual(book[LAUNCH_HOURLY_SHEET].cell(2, 5).value, '首销期订单节奏.xlsx')
        book.close()

    def test_week_reference_eligibility_is_derived_from_dates(self):
        book = self.make_public()
        sheet = book[WEEKLY_SHEET]
        # Launch ends Fri Jan 9; the Jan 5 week is not a complete steady week.
        sheet.append(['测试车', '26WK02', '平销', datetime(2026, 1, 10), datetime(2026, 1, 11), None, None, 3])
        sheet.append(['测试车', '26WK03', '平销', datetime(2026, 1, 12), datetime(2026, 1, 18), 20, 12, 7])
        sheet.append(['测试车', '26WK04', '平销', datetime(2026, 1, 19), datetime(2026, 1, 25), None, None, 9])
        for row in range(2, 5):
            sheet.cell(row, 11, '锁单选配比例分析.xlsx｜测试车by周')
        steady = read_public_forecast(book, [], today=date(2026, 1, 25))[3]
        self.assertEqual([w['period'] for w in steady[0]['weeks']], ['26WK03'])
        self.assertEqual(steady[0]['weeks'][0]['gross'], 20)  # Not retained orders (12).
        sheet.cell(3, 6, 0)
        self.assertEqual(read_public_forecast(book, [], today=date(2026, 1, 25))[3][0]['weeks'][0]['gross'], 0)
        sheet.cell(3, 11, '大定选配比例分析.xlsx｜测试车by周')
        self.assertEqual(read_public_forecast(book, [], today=date(2026, 1, 25))[3], [])
        book.close()

    def test_summary_result_is_authoritative_across_stages(self):
        from modules.sales_forecast import _resolve_actual_profiles
        from core.models import SourceRef
        book = self.make_public()
        history = [{'model': '历史车', 'generation': '测试车', 'small': 100, 'gross': 20,
                    'days': 1, 'daily_small': [10], 'daily_direct': [10],
                    'launch_date': '2026-01-03'}]
        raw = read_public_forecast(book, history)[0]
        targets = [{'name': '测试车', 'stage': 'ended', 'launch_date': '2026-01-03'}]
        result = _resolve_actual_profiles(history, raw, targets, SourceRef('summary.xlsx', MASTER_SHEET))[0]
        self.assertEqual(result['stage_profiles']['active']['days'][0]['gross'], 30)
        self.assertEqual(result['stage_profiles']['ended']['days'][0]['gross'], 30)
        book.close()

    def weekly_fixture(self, *, complete_days=True, cutoff='2026-01-01'):
        model = '问界 M7 2026款'
        book = Workbook()
        week = book.active
        week.title = model + 'by周'
        week.append(['指标', '统计类型', '分类', '26WK01'])
        week.append(['大定', '数量', '数量', 100])
        week.append(['净大定', '数量', '数量', 70])
        week.append(['交车锁单', '数量', '数量', 50])
        if complete_days:
            day = book.create_sheet(model + 'by天')
            dates = [date.fromisocalendar(2026, 1, n) for n in range(1, 8)]
            day.append(['指标', '统计类型', '分类', *dates])
            day.append(['大定', '数量', '数量', 1, 2, 3, 4, 5, 6, 7])
            day.append(['留存大定', '数量', '数量', 1, 1, 2, 2, 3, 3, 4])
            day.append(['交车锁单', '数量', '数量', 0, 1, 1, 1, 2, 2, 2])
        item = WorkbookItem(Path('大定选配比例分析.xlsx'), book)
        store = SimpleNamespace(find=lambda key: item if key == '大定选配比例' else None)
        data = {'targets': [{'name': model, 'launch_date': '2025-12-20', 'end_date': cutoff}], 'actuals': []}
        dashboard = SimpleNamespace(views={'week': {'pages': {'预测方案': {'workspace': {'data': data}}}}})
        with patch('modules.sales_forecast._read_model_mapping', return_value={}):
            rows = forecast_weekly_orders(store, dashboard, date(2026, 1, 10))
        book.close()
        return rows

    def test_cross_year_week_is_split_at_stage_date_for_all_three_metrics(self):
        rows = self.weekly_fixture()
        self.assertEqual([(r[2], *r[5:8]) for r in rows], [('首销', 10, 6, 3), ('平销', 18, 10, 6)])
        self.assertEqual(sum(r[5] for r in rows), 28)  # Never repeat the original whole-week 100.

    def test_cannot_split_whole_week_without_daily_evidence(self):
        with self.assertLogs('sales_forecast_refresh', level='WARNING') as logs:
            rows = self.weekly_fixture(complete_days=False)
        self.assertEqual(len(logs.output), 6)
        for stage in ('首销', '平销'):
            for field in ('大定', '留存大定', '交车锁单'):
                self.assertTrue(any(f'阶段={stage} | 字段={field} |' in message for message in logs.output))
        self.assertTrue(all(r[5:8] == [None, None, None] for r in rows))

    def test_whole_steady_week_prefers_source_week_over_daily_sum(self):
        rows = self.weekly_fixture(cutoff='2025-12-28')
        self.assertEqual(rows[0][5:8], [100, 70, 50])
