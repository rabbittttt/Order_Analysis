"""Missing physical files versus lifecycle dates, checked against independent facts."""
from datetime import date, datetime, timedelta
from itertools import product
from unittest import TestCase
from openpyxl import Workbook, load_workbook

from core.forecast_summary import (MASTER_SHEET, DAILY_SHEET, D12_SHEET, SMALL_DAILY_SHEET,
    SMALL_CURVE_SHEET, SMALL_HOURLY_SHEET, LAUNCH_HOURLY_SHEET, WEEKLY_SHEET, table_records, iso_day)
from tests import test_forecast_source_backfill_e2e as fixtures


FAMILIES = ('launch', 'cancel', 'small_mix', 'order_mix', 'lock_mix')
STAGES = ('2026-08-29', '2026-08-30', '2026-08-31', '2026-09-01',
          '2026-09-02', '2026-09-03', '2026-09-04', '2026-09-20')
ALL = frozenset(FAMILIES)


class MissingSourceMatrixTests(TestCase):
    model = '享界 G9 2027款'
    web_inputs = fixtures.SourceBackfillEndToEndTests.web_inputs

    def build_case(self, present, stamp, *, maintained='empty', updated=None, hourly_only=False,
                   reservation=True, weekly_only=False, launch_quantities=None, extra_sources=None):
        cutoff = date.fromisoformat(updated) if updated else None
        def edit_source(path):
            book = load_workbook(path)
            small = book['小订by天']
            small['B2'], small['C2'] = ((40, 60) if maintained == 'complete' else (None, None))
            if not reservation:
                headers = [c.value for c in book['车型汇总'][1]]
                for field in ('小订开始日期', '小订结束日期', '总小订'):
                    book['车型汇总'].cell(2, headers.index(field)+1).value = None
                del book['小订by天']
            book.save(path)
            book.close()
        def sources(folder):
            launch = folder/'鸿蒙智行首销期订单节奏.xlsx'
            cancel = folder/'鸿蒙智行小订退订分析.xlsx'
            if 'launch' not in present:
                launch.unlink()
            elif hourly_only or cutoff:
                book = load_workbook(launch)
                if hourly_only:
                    del book[self.model+'by天']
                for sheet in book:
                    for col in range(2, sheet.max_column+1):
                        value = sheet.cell(2, col).value
                        if cutoff and value.date() > cutoff:
                            for row in range(3, sheet.max_row+1):
                                sheet.cell(row, col).value = None
                book.save(launch)
                book.close()
            if 'cancel' not in present:
                cancel.unlink()
            else:
                book = load_workbook(cancel)
                sheet = book.create_sheet(self.model+'_小订分时退订')
                sheet.append(['小订日期', '小订小时', '小订数'])
                for day, count in ((date(2026, 8, 30), 45), (date(2026, 8, 31), 55)):
                    if not cutoff or day <= cutoff:
                        sheet.append([day, 20, count])
                if cutoff:
                    sheet = book[self.model+'_日度退订']
                    for row in range(3, sheet.max_row+1):
                        if sheet.cell(row, 1).value.date() > cutoff:
                            for col in (2, 3):
                                sheet.cell(row, col).value = None
                book.save(cancel)
                book.close()
            for family, title, entries, days in (
                ('small_mix', '小订', [('小订', [40, 60])], [date(2026, 8, 30), date(2026, 8, 31)]),
                ('order_mix', '大定', [('大定', [55, 60, 35]+[13]*10), ('留存大定', [40, 45, 25]+[10]*10)],
                 [date(2026, 9, 1)+timedelta(days=i) for i in range(13)]),
                ('lock_mix', '锁单', [('交车锁单', [20, 20, 20]+[7]*10)],
                 [date(2026, 9, 1)+timedelta(days=i) for i in range(13)])):
                if family not in present:
                    continue
                book = Workbook()
                sheet = book.active
                sheet.title = self.model+'by天'
                included = [i for i, day in enumerate(days) if not cutoff or day <= cutoff]
                sheet.append(['指标', '统计类型', '分类', *[days[i] for i in included]])
                for label, values in entries:
                    sheet.append([label, '数量', '数量', *[values[i] for i in included]])
                if family in ('order_mix', 'lock_mix') and not cutoff:
                    sheet = book.create_sheet(self.model+'by周')
                    sheet.append(['指标', '统计类型', '分类', '26WK37', '26WK38'])
                    for label, values in entries:
                        quantity = {'大定':91, '留存大定':70, '交车锁单':49}[label]
                        sheet.append([label, '数量', '数量', quantity, quantity+1])
                    if weekly_only:
                        del book[self.model+'by天']
                book.save(folder/f'鸿蒙智行{title}选配比例分析.xlsx')
                book.close()
            if extra_sources:
                extra_sources(folder)
        quantities = None
        if maintained == 'complete':
            quantities = {'小转大进度':[10, 20, 10], '直接大定进度':[40, 30, 20], '大定进度':[50, 50, 30]}
        elif maintained == 'partial':
            quantities = {'小转大进度':[0, None, None], '直接大定进度':[50, None, None], '大定进度':[50, None, None]}
        if not reservation:
            launch_quantities = {name: [None]*3 for name in ('当日小订转大定数量', '累计小订转大定数量', '当日直接大定数量', '累计直接大定数量')}
        return fixtures.SourceBackfillEndToEndTests.build(self, as_of=stamp, launch_quantities=launch_quantities,
            source_quantities=quantities, source_edit=edit_source, add_sources=sources)

    def assert_case(self, book, present, stamp, *, maintained='empty', updated=None):
        cutoff = updated or '9999-12-31'
        rows = table_records(book, DAILY_SHEET)
        lookup = {iso_day(r['日期']): r for r in rows}
        self.assertEqual(len(lookup), len(rows), 'one model/date row')
        self.assertTrue(all(s.sheet_state == 'visible' for s in book))
        for day, row in lookup.items():
            self.assertEqual(row['日期状态'], '已结束日' if day < stamp else '当日未结束' if day == stamp else '未来日期')
        def expected(day, field, index):
            available = day <= cutoff
            if field == '小订数量':
                return ([40, 60][index] if maintained == 'complete' else
                        [40, 60][index] if 'small_mix' in present and available else
                        [45, 55][index] if 'cancel' in present and available else None)
            if field == '退订数量':
                return index+1 if 'cancel' in present and available else None
            source = {'大定':[50, 50, 30], '小转大':[10, 20, 10],
                      '直接大定':[40, 30, 20], '交车锁单':[25, 25, 15]}
            if field in ('大定', '小转大', '直接大定'):
                if maintained == 'complete':
                    return source[field][index]
                if maintained == 'partial' and index == 0:
                    return {'大定':50, '小转大':0, '直接大定':50}[field]
            if field in source and 'launch' in present and available:
                return source[field][index]
            if field == '大定' and 'order_mix' in present and available:
                return [55, 60, 35][index]
            if field == '留存大定' and 'order_mix' in present and available:
                return [40, 45, 25][index]
            if field == '交车锁单' and 'lock_mix' in present and available:
                return 20
            return None
        for i in range(2):
            day = f'2026-08-{30+i:02d}'
            value = expected(day, '小订数量', i)
            self.assertEqual(lookup.get(day, {}).get('小订数量'), value, ('小订', day, present))
            small = table_records(book, SMALL_DAILY_SHEET)[0]
            self.assertEqual(small.get(f'D{i+1}'), value)
        small_curve = table_records(book, SMALL_CURVE_SHEET)[0]
        small_running = 0
        for i in range(2):
            value = expected(f'2026-08-{30+i:02d}', '小订数量', i)
            small_running = small_running+value if small_running is not None and value is not None else None
            self.assertEqual(small_curve.get(f'D{i+1}'), small_running/100 if small_running is not None else None)
        fields = ('大定', '小转大', '直接大定', '留存大定', '交车锁单', '退订数量')
        expected_days = [{f: expected(f'2026-09-{i+1:02d}', f, i) for f in fields} for i in range(3)]
        for i, values in enumerate(expected_days, 1):
            for field, value in values.items():
                self.assertEqual(lookup.get(f'2026-09-{i:02d}', {}).get(field), value, (field, i, present))
        master = table_records(book, MASTER_SHEET)[0]
        for field, total in (('大定', '总大定'), ('小转大', '总小转大'), ('直接大定', '总直接大定'),
                             ('留存大定', '首销期留存大定'), ('交车锁单', '首销期锁单'), ('退订数量', '总退订')):
            values = [r[field] for r in expected_days]
            value = sum(values) if all(v is not None for v in values) else None
            self.assertEqual(master.get(total), value, (total, present, updated))
        for field, numerator, denominator in (('小订转化率', '总小转大', '总小订'),
                ('直接大定占比', '总直接大定', '总大定'), ('退订率', '总退订', '总小订')):
            a, b = master.get(numerator), master.get(denominator)
            wanted = a/b if a is not None and b else None
            self.assertEqual(master.get(field), wanted, (field, present, stamp))
        for field, denominator, curve in (('小转大', '总小转大', '小转大累计完成度'),
                ('直接大定', '总直接大定', '直接大定累计完成度'), ('大定', '总大定', '累计大定完成度'),
                ('小转大', '总小订', '累计小订转化率'), ('退订数量', '总小订', '累计退订率')):
            row = table_records(book, curve)[0]
            running = 0
            for i, values in enumerate(expected_days, 1):
                running = running+values[field] if running is not None and values[field] is not None else None
                den = master.get(denominator)
                wanted = running/den if running is not None and den else None
                if wanted is None:
                    self.assertIsNone(row.get(f'D{i}'), (curve, i, present))
                else:
                    self.assertAlmostEqual(row.get(f'D{i}'), wanted, msg=str((curve, i, present)))
        metrics = table_records(book, D12_SHEET)[0]
        for i in (1, 2):
            for field, suffix in (('大定', '大定'), ('小转大', '小转大'), ('直接大定', '直接大'), ('退订数量', '退订')):
                self.assertEqual(metrics.get(f'D{i}{suffix}'), expected_days[i-1][field])
        for day in range(4, 14):
            key = f'2026-09-{day:02d}'
            for field, family, count in (('大定', 'order_mix', 13), ('留存大定', 'order_mix', 10), ('交车锁单', 'lock_mix', 7)):
                self.assertEqual(lookup.get(key, {}).get(field), count if family in present and key <= cutoff else None)
            for field in ('小订数量', '小转大', '直接大定'):
                self.assertIsNone(lookup.get(key, {}).get(field), (field, key))
        for sheet in (SMALL_HOURLY_SHEET, LAUNCH_HOURLY_SHEET):
            hourly = table_records(book, sheet)
            self.assertEqual(len(hourly), len({(iso_day(r['日期']), r['小时']) for r in hourly}))
            dates = ('2026-08-30', '2026-08-31') if sheet == SMALL_HOURLY_SHEET else ('2026-09-01', '2026-09-02', '2026-09-03')
            family = 'cancel' if sheet == SMALL_HOURLY_SHEET else 'launch'
            self.assertEqual(len(hourly), sum(day <= cutoff for day in dates) if family in present else 0)
        history, profiles, _, small, _ = self.web_inputs(book, date.fromisoformat(stamp))
        self.assertTrue(all(r['date'] <= stamp for p in profiles for r in p['days']))
        if stamp <= '2026-09-03':
            self.assertIsNone(history[0]['conversion'])
        if stamp <= '2026-08-31':
            self.assertTrue(all(not r['total_complete'] for r in small))
        if not updated:
            weeks = {r['周期']: r for r in table_records(book, WEEKLY_SHEET)}
            for period, extra in (('26WK37', 0), ('26WK38', 1)):
                for field, family, value in (('大定', 'order_mix', 91), ('留存大定', 'order_mix', 70), ('交车锁单', 'lock_mix', 49)):
                    if family in present:
                        self.assertEqual(weeks.get(period, {}).get(field), value+extra, (period, field, stamp))

    def test_smoke_missing_each_family_and_maintained_quantities(self):
        for present, maintained in product((ALL, *(ALL-{family} for family in FAMILIES), frozenset()), ('empty', 'partial', 'complete')):
            with self.subTest(present=sorted(present), maintained=maintained):
                try:
                    book = self.build_case(present, STAGES[-1], maintained=maintained)
                    self.assert_case(book, present, STAGES[-1], maintained=maintained)
                finally:
                    self.doCleanups()

    def test_every_file_subset_across_eight_stage_dates(self):
        for bits, stamp in product(product((False, True), repeat=len(FAMILIES)), STAGES):
            present = frozenset(name for name, enabled in zip(FAMILIES, bits) if enabled)
            with self.subTest(present=sorted(present), stamp=stamp):
                try:
                    book = self.build_case(present, stamp)
                    self.assert_case(book, present, stamp)
                finally:
                    self.doCleanups()

    def test_files_updated_only_through_current_or_previous_day(self):
        for stamp, lag in product(STAGES, (0, 1)):
            updated = (date.fromisoformat(stamp)-timedelta(days=lag)).isoformat()
            with self.subTest(stamp=stamp, updated=updated):
                try:
                    book = self.build_case(ALL, stamp, updated=updated)
                    self.assert_case(book, ALL, stamp, updated=updated)
                finally:
                    self.doCleanups()

    def test_hourly_only_source_preserves_components_without_inventing_full_days(self):
        for stamp in STAGES:
            with self.subTest(stamp=stamp):
                try:
                    book = self.build_case(frozenset({'launch'}), stamp, hourly_only=True)
                    hours = table_records(book, LAUNCH_HOURLY_SHEET)
                    self.assertEqual(len(hours), 3)
                    for i, row in enumerate(hours):
                        for field, expected in (('小时大定', [50, 50, 30]), ('小时小转大', [10, 20, 10]),
                                                ('小时直接大定', [40, 30, 20]), ('小时交车锁单', [25, 25, 15])):
                            self.assertEqual(row.get(field), expected[i], (field, stamp, i))
                    master = table_records(book, MASTER_SHEET)[0]
                    self.assertIsNone(master['总小转大'])
                    self.assertIsNone(master['总直接大定'])
                finally:
                    self.doCleanups()

    def test_partial_hour_components_survive_unknown_and_invalid_hour_gross(self):
        book = self.build_case(frozenset({'launch'}), STAGES[-1], hourly_only=True, launch_quantities={
            '当日大定数量':[None, None, -1], '当日小订转大定数量':[10, None, 0],
            '当日直接大定数量':[None, 30, 0], '累计小订转大定数量':[None]*3,
            '累计直接大定数量':[None]*3})
        rows = table_records(book, LAUNCH_HOURLY_SHEET)
        self.assertEqual(len(rows), 3)
        self.assertEqual([r['小时大定'] for r in rows], [None]*3)
        self.assertEqual([r['小时小转大'] for r in rows], [10, None, 0])
        self.assertEqual([r['小时直接大定'] for r in rows], [None, 30, 0])
        self.assertEqual([r['小时交车锁单'] for r in rows], [25, 25, 15])
        _, profiles, *_ = self.web_inputs(book, date.fromisoformat(STAGES[-1]))
        self.assertTrue(all(not p['hourly_days'] for p in profiles))

    def test_week_only_files_preserve_week_facts_without_daily_estimates(self):
        for stamp in STAGES:
            with self.subTest(stamp=stamp):
                try:
                    book = self.build_case(frozenset({'order_mix', 'lock_mix'}), stamp, weekly_only=True)
                    rows = {r['周期']: r for r in table_records(book, WEEKLY_SHEET)}
                    for period, extra in (('26WK37', 0), ('26WK38', 1)):
                        for field, value in (('大定', 91), ('留存大定', 70), ('交车锁单', 49)):
                            self.assertEqual(rows[period][field], value+extra)
                    self.assertFalse(any(r.get('大定') is not None or r.get('交车锁单') is not None
                                         for r in table_records(book, DAILY_SHEET)))
                finally:
                    self.doCleanups()

    def test_direct_launch_models_with_each_order_file_subset(self):
        for bits, stamp in product(product((False, True), repeat=3), STAGES):
            present = frozenset(name for name, enabled in zip(('launch', 'order_mix', 'lock_mix'), bits) if enabled)
            with self.subTest(present=sorted(present), stamp=stamp):
                try:
                    book = self.build_case(present, stamp, reservation=False)
                    rows = {iso_day(r['日期']): r for r in table_records(book, DAILY_SHEET)}
                    self.assertFalse(table_records(book, MASTER_SHEET)[0]['有小订'])
                    for i in range(3):
                        day = f'2026-09-{i+1:02d}'
                        gross = [50, 50, 30][i] if 'launch' in present else [55, 60, 35][i] if 'order_mix' in present else None
                        self.assertEqual(rows.get(day, {}).get('大定'), gross)
                        self.assertEqual(rows.get(day, {}).get('直接大定'), gross)
                        self.assertEqual(rows.get(day, {}).get('小转大'), 0 if gross is not None else None)
                    master = table_records(book, MASTER_SHEET)[0]
                    self.assertIsNone(master['小订转化率'])
                    self.assertEqual(master['总直接大定'], 130 if 'launch' in present else 150 if 'order_mix' in present else None)
                finally:
                    self.doCleanups()

    def test_missing_required_maintenance_file_preserves_previous_summary(self):
        saved = []
        def seed(path):
            output = path.parent/'汇总.xlsx'
            book = Workbook()
            book.active['A1'] = '上次汇总'
            book.save(output)
            book.close()
            saved.append(output)
        def remove_source(folder):
            (folder.parent/'小订及首销数据整理.xlsx').unlink()
        with self.assertRaises(FileNotFoundError):
            fixtures.SourceBackfillEndToEndTests.build(self, source_edit=seed, add_sources=remove_source)
        book = load_workbook(saved[0])
        try:
            self.assertEqual(book.active['A1'].value, '上次汇总')
        finally:
            book.close()

    def full_hour_source(self, folder, defect=None):
        path = folder/'鸿蒙智行首销期订单节奏.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.title = self.model+'by时'
        times = [datetime(2026, 9, day, hour) for day in (1, 2, 3) for hour in range(24)]
        if defect == 'missing_hour':
            times.remove(datetime(2026, 9, 1, 10))
        if defect == 'duplicate_hour':
            times.append(datetime(2026, 9, 1, 10))
        sheet.append(['周期', *[f'H{i+1}' for i in range(len(times))]])
        sheet.append(['时间', *times])
        for label, quantities in (('当日大定数量',[50,50,30]), ('当日小订转大定数量',[10,20,10]),
                                  ('当日直接大定数量',[40,30,20]), ('当日交车锁单数量',[25,25,15])):
            values = [quantities[t.day-1] if t.hour == 23 else 0 for t in times]
            if defect == 'unknown_component' and label in ('当日大定数量', '当日小订转大定数量'):
                values[times.index(datetime(2026, 9, 1, 10))] = None
            sheet.append([label, *values])
        book.save(path)
        book.close()

    def test_complete_24_hour_sources_fill_daily_facts_across_stages(self):
        for stamp in STAGES:
            with self.subTest(stamp=stamp):
                try:
                    book = self.build_case(frozenset({'launch'}), stamp, hourly_only=True,
                                           extra_sources=self.full_hour_source)
                    master = table_records(book, MASTER_SHEET)[0]
                    self.assertEqual((master['总小转大'], master['总直接大定'], master['总大定']), (40,90,130))
                    rows = {iso_day(r['日期']):r for r in table_records(book,DAILY_SHEET)}
                    self.assertEqual([rows[f'2026-09-{i:02d}']['大定'] for i in (1,2,3)], [50,50,30])
                    self.assertEqual(len(table_records(book,LAUNCH_HOURLY_SHEET)), 72)
                    curve = table_records(book,'小转大累计完成度')[0]
                    self.assertEqual([curve[f'D{i}'] for i in (1,2,3)], [.25,.75,1])
                finally:
                    self.doCleanups()

    def test_incomplete_hour_grid_does_not_make_complete_daily_totals(self):
        for defect in ('missing_hour', 'duplicate_hour', 'unknown_component'):
            with self.subTest(defect=defect):
                try:
                    book = self.build_case(frozenset({'launch'}), STAGES[-1], hourly_only=True,
                        extra_sources=lambda folder:self.full_hour_source(folder, defect))
                    master = table_records(book,MASTER_SHEET)[0]
                    self.assertIsNone(master['总小转大'])
                    self.assertIsNone(master['总大定'])
                    self.assertEqual(master['总直接大定'], 90 if defect == 'unknown_component' else None)
                finally:
                    self.doCleanups()
