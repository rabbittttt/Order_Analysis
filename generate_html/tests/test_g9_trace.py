"""Removing temporary G9 diagnostics must preserve forecast cells and warnings."""
from unittest import TestCase

from openpyxl import load_workbook

import tests.test_forecast_source_backfill_e2e as fixtures
from core.forecast_summary import MASTER_SHEET, D12_SHEET, table_records


class RetiredG9TraceTests(TestCase):
    def fixture(self, **kwargs):
        helper = fixtures.SourceBackfillEndToEndTests()
        self.addCleanup(helper.doCleanups)
        return helper, helper.build(**kwargs)

    def assert_no_trace(self, logs):
        self.assertFalse(any("[G9取数排查]" in record.getMessage() for record in logs.records))

    def test_full_pipeline_keeps_quantities_without_special_traces_even_at_debug(self):
        with self.assertLogs(level="DEBUG") as logs:
            helper, book = self.fixture()
        helper.assert_complete(book)
        self.assert_no_trace(logs)
        self.assertTrue(all(sheet.sheet_state == "visible" for sheet in book))
        self.assertTrue(any("刷新完成" in record.getMessage() for record in logs.records))

    def test_missing_terminal_keeps_unknown_values_and_business_warning(self):
        def edit(folder):
            path = folder / "鸿蒙智行首销期订单节奏.xlsx"
            book = load_workbook(path)
            for sheet in book:
                for row in (4, 5, 6, 7):
                    sheet.cell(row, 4).value = None
            book.save(path)
            book.close()

        with self.assertLogs(level="INFO") as logs:
            _, book = self.fixture(add_sources=edit)
        self.assert_no_trace(logs)
        self.assertIsNone(table_records(book, MASTER_SHEET)[0]["总小转大"])
        self.assertEqual(table_records(book, D12_SHEET)[0]["D1小转大"], 10)
        self.assertTrue(any("[预测历史参考缺项]" in record.getMessage() for record in logs.records))

    def test_cumulative_conflict_keeps_original_selection_and_warning(self):
        with self.assertLogs(level="INFO") as logs:
            _, book = self.fixture(source_quantities={"小转大进度": [0, None, None]})
        self.assert_no_trace(logs)
        self.assertEqual(table_records(book, MASTER_SHEET)[0]["总小转大"], 30)
        self.assertTrue(any("[首销累计口径冲突]" in record.getMessage() for record in logs.records))
