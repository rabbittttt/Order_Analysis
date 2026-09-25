from __future__ import annotations

import unittest
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from core.components import add_kpi_comparisons
from core.discovery import discover_subjects, is_valid_subject_name, set_capabilities
from core.excel import WorkbookItem, WorkbookStore, compact_identifier, compact_text, format_excel_value, grain_from_sheet, parse_metric_sheet, safe_number, sheet_subject, subject_parent, subject_type
from core.models import Subject
from modules.generic import GenericModule
from modules.launch_rhythm import LaunchRhythmModule
from modules.option_fee import OptionFeeModule, _option_layout, _period_groups, _version_rows_by_period
from modules.order_mix import LockMixModule, OrderMixModule, SmallOrderMixModule
from modules.cancellation import CancellationModule, parse_select_retreat
from modules.conversion import ConversionModule
from modules.overview import OverviewModule, _drop_rankings, _group_model_series, _lock_share_rows, _log_lock_share_issues, _period_range_text, _rise_rankings
from modules.sku import SkuModule, find_header_row
from tools.refresh_sales_forecast_data import build_mapping_workbook, write_rows


class SubjectParsingTests(unittest.TestCase):
    def test_generated_mapping_workbook_does_not_embed_vehicle_facts(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "车型基本信息.xlsx"
            build_mapping_workbook(path, ["问界 M9 2026款纯电"])
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                row = list(workbook["车型基本信息"].iter_rows(min_row=2, max_row=2, values_only=True))[0]
            finally:
                workbook.close()

        self.assertEqual(row[0], "问界 M9 2026款纯电")
        self.assertIsNone(row[1])
        self.assertEqual(row[2], "问界")
        self.assertEqual(row[4:8], ("待维护", "待维护", "待维护", "待维护"))
        self.assertEqual(row[8], "待人工确认")
        self.assertIsNone(row[10])

    def test_multi_phase_launch_sheet_names_map_to_same_generation(self):
        self.assertEqual(sheet_subject("问界 M8 2025款by天第1期"), "问界 M8 2025款")
        self.assertEqual(sheet_subject("问界 M8 2025款by天第2期"), "问界 M8 2025款")

    def test_option_fee_reads_actual_amount_from_dynamic_period_total_column(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.cell(3, 3, "26-07")
        sheet.cell(4, 3, "免费")
        sheet.cell(4, 4, "0-10000")
        sheet.cell(4, 5, "总计")
        sheet.cell(3, 6, "26-08")
        sheet.cell(4, 6, "免费")
        sheet.cell(4, 7, "0-5000")
        sheet.cell(4, 8, "5000-10000")
        sheet.cell(4, 9, "总计")
        sheet.cell(5, 2, "订单数量")
        sheet.cell(5, 3, 100)
        sheet.cell(5, 4, 200)
        sheet.cell(5, 6, 100)
        sheet.cell(5, 7, 200)
        sheet.cell(5, 8, 300)
        sheet.cell(6, 1, "Max")
        sheet.cell(6, 2, "选配金额")
        sheet.cell(6, 5, 10000)
        sheet.cell(6, 9, 12000)
        sheet.cell(7, 2, "实际金额")
        sheet.cell(7, 5, 4500)
        sheet.cell(7, 9, 6500)

        periods, _ = _period_groups(sheet)
        rows = _version_rows_by_period(sheet, periods)

        self.assertEqual(rows["26-07"][0][1:3], [10000.0, 4500.0])
        self.assertEqual(rows["26-08"][0][1:3], [12000.0, 6500.0])

    def test_option_fee_layout_follows_inserted_instruction_rows(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["说明一"])
        sheet.append(["说明二"])
        sheet.cell(5, 3, "26-08")
        sheet.cell(6, 3, "免费")
        sheet.cell(6, 4, "0-5000")
        sheet.cell(7, 2, "订单数量")
        sheet.cell(7, 3, 100)
        sheet.cell(7, 4, 200)

        self.assertEqual(_option_layout(sheet), (5, 6, 7))
        periods, groups = _period_groups(sheet)
        self.assertEqual(periods, ["26-08"])
        self.assertEqual([band for _, band in groups["26-08"]], ["免费", "0-5000"])

    def test_sku_header_row_is_not_fixed_to_second_row(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["说明"])
        sheet.append([])
        sheet.append(["SKU", "版本", "外观", "内饰", "轮毂", "26WK35", "总计占比"])
        self.assertEqual(find_header_row(sheet, ("总计占比",)), 3)

    def test_sku_latest_page_compares_with_previous_sheet_period(self):
        workbook = Workbook()
        sku_sheet = workbook.active
        sku_sheet.title = "问界 M9 2026款_SKU"
        headers = ["SKU", "版本", "外观", "内饰", "轮毂", "26WK34", "26WK35", "总计占比", "累计锁单", "26WK34", "26WK35", "总计"]
        sku_sheet.append(headers)
        rows = [
            ["SKU1", "Max", "黑", "棕", "20", .6, .45, .45, "TOP1", .6, .45, .45],
            ["SKU2", "Max", "白", "灰", "20", .4, .3, .3, "TOP2", .8, .7, .7],
            ["SKU3", "Ultra", "蓝", "棕", "21", 0, .25, .25, "TOP3", .8, .85, .85],
            ["SKU4", "Ultra", "黑", "灰", "21", 0, 0, 0, "TOP5", .8, .95, .95],
        ]
        for row in rows:
            sku_sheet.append(row)
        combo_sheet = workbook.create_sheet("问界 M9 2026款_组合")
        combo_sheet.append(["累计锁单", "配置", "26WK34", "26WK35", "总计"])
        combo_sheet.append(["TOP1", "Max", .6, .45, .45])
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("鸿蒙智行SKU收敛度报告.xlsx"), workbook)]

        dashboard = SkuModule().build(store, Subject("m9", "问界 M9 2026款", "generation", "问界"))
        payload = dashboard.to_dict()
        add_kpi_comparisons({"m9|sku": payload})
        kpis = {item["label"]: item for item in payload["views"]["week"]["pages"]["26WK35"]["kpis"]}

        self.assertEqual(kpis["有效SKU"]["value"], 3)
        self.assertEqual(kpis["有效SKU"]["comparison"]["previous"], 2)
        self.assertEqual(kpis["TOP3累计"]["comparison"]["previous"], 80)
        self.assertEqual(kpis["覆盖80%的SKU"]["comparison"]["previous"], 2)
        self.assertTrue(all(item["comparison"]["status"] == "available" for item in kpis.values()))

    def test_capability_order_comes_from_configured_order(self):
        subject = Subject("group_test", "鸿蒙智行", "group", None)
        set_capabilities(
            [subject],
            {"鸿蒙智行": {"overview", "sales_forecast"}},
            ["sales_forecast", "overview", "raw"],
        )
        self.assertEqual(subject.modules, ["sales_forecast", "overview", "raw"])

    def test_launch_rhythm_keeps_multiple_phases_and_defaults_to_latest(self):
        workbook = Workbook()
        first = workbook.active
        first.title = "问界 M8 2025款by天第1期"
        second = workbook.create_sheet("问界 M8 2025款by天第2期")
        for sheet, start, orders in (
            (first, datetime(2025, 4, 10), (100, 80)),
            (second, datetime(2025, 5, 1), (120, 90)),
        ):
            sheet.cell(1, 2, "D1")
            sheet.cell(1, 3, "D2")
            sheet.cell(2, 2, start)
            sheet.cell(2, 3, start.replace(day=start.day + 1))
            sheet.cell(3, 1, "当日大定数量")
            sheet.cell(3, 2, orders[0])
            sheet.cell(3, 3, orders[1])
            sheet.cell(4, 1, "当日净大定数量")
            sheet.cell(4, 2, orders[0] - 10)
            sheet.cell(4, 3, orders[1] - 10)
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("鸿蒙智行首销期订单节奏.xlsx"), workbook)]

        dashboard = LaunchRhythmModule().build(store, Subject("m8", "问界 M8 2025款", "generation", "问界"))

        view = dashboard.views["day"]
        self.assertEqual(view["periods"][0], "全部阶段")
        self.assertTrue(view["default_period"].startswith("第2期"))
        self.assertEqual(len(view["periods"]), 3)
        summary = view["pages"]["全部阶段"]["sections"][0]["data"]["rows"]
        self.assertEqual(summary[1][4], 19)

    def test_launch_small_conversion_kpi_uses_source_total_small_rate(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "问界 F2N 2026款by天"
        for column, (period, day) in enumerate((("D1", 1), ("D2", 2)), start=2):
            sheet.cell(1, column, period)
            sheet.cell(2, column, datetime(2026, 8, day))
        for row, (label, values) in enumerate((
            ("当日大定数量", (80, 120)),
            ("当日净大定数量", (70, 100)),
            ("累计小订转大定数量", (20, 40)),
            ("累计小订转化率", (0.1, 0.2)),
            ("累计直接大定数量", (60, 160)),
        ), start=3):
            sheet.cell(row, 1, label)
            for column, value in enumerate(values, start=2):
                sheet.cell(row, column, value)
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("鸿蒙智行首销期订单节奏.xlsx"), workbook)]

        dashboard = LaunchRhythmModule().build(
            store,
            Subject("f2n", "问界 F2N 2026款", "generation", "问界"),
        )

        page = dashboard.views["day"]["pages"]["2026-08-02"]
        kpis = {item["label"]: item for item in page["kpis"]}
        matrix = next(item for item in page["sections"] if item["kind"] == "matrix")
        rate_row = next(row for row in matrix["data"]["rows"] if row[0] == "累计小订转化率")
        self.assertEqual(kpis["累计小转大率"]["value"], 20)
        self.assertEqual(kpis["累计小转大率"]["note"], "累计小转大 ÷ 总小订")
        self.assertEqual(kpis["累计留存大定"]["value"], 170)
        self.assertTrue(any(row[0] == "当日留存大定数量" for row in matrix["data"]["rows"]))
        self.assertEqual(rate_row[-1], "20.0%")

    def test_generated_forecast_range_uses_worksheet_filter_without_table(self):
        workbook = Workbook()
        workbook.remove(workbook.active)

        write_rows(workbook, "测试", ["车型", "D1"], [["问界 F2N 2026款", 100]], "ForecastTest")

        sheet = workbook["测试"]
        self.assertEqual(sheet.auto_filter.ref, "A1:B2")
        self.assertEqual(len(sheet.tables), 0)

    def test_metric_parser_keeps_source_total_column(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "问界by周"
        sheet.cell(1, 4, "26WK34")
        sheet.cell(1, 5, "总计")
        sheet.cell(2, 1, "大定")
        sheet.cell(2, 2, "数量")
        sheet.cell(2, 3, "数量")
        sheet.cell(2, 4, 443)
        sheet.cell(2, 5, 3239)

        parsed = parse_metric_sheet(sheet)

        self.assertEqual(list(parsed), ["26WK34", "总计"])
        self.assertEqual(parsed["总计"]["metrics"]["大定"], 3239)

    def test_brand_overview_excludes_total_from_kpis_and_history(self):
        order_book = Workbook()
        order_sheet = order_book.active
        order_sheet.title = "问界by周"
        lock_book = Workbook()
        lock_sheet = lock_book.active
        lock_sheet.title = "问界by周"
        for sheet, metrics in (
            (order_sheet, (("大定", 443, 3239), ("净大定", 412, 2952))),
            (lock_sheet, (("交车锁单", 320, 2400), ("已交付", 280, 2100))),
        ):
            sheet.cell(1, 4, "26WK34")
            sheet.cell(1, 5, "总计")
            for row, (label, latest, total) in enumerate(metrics, start=2):
                sheet.cell(row, 1, label)
                sheet.cell(row, 2, "数量")
                sheet.cell(row, 3, "数量")
                sheet.cell(row, 4, latest)
                sheet.cell(row, 5, total)
        store = WorkbookStore(Path("."))
        store.items = [
            WorkbookItem(Path("鸿蒙智行大定选配比例分析.xlsx"), order_book),
            WorkbookItem(Path("鸿蒙智行锁单选配比例分析.xlsx"), lock_book),
        ]

        dashboard = OverviewModule().build(store, Subject("brand_test", "问界", "brand", "鸿蒙智行"))

        self.assertEqual(dashboard.views["week"]["default_period"], "26WK34")
        self.assertEqual(dashboard.views["week"]["periods"], ["26WK34"])
        self.assertNotIn("总计", dashboard.views["week"]["pages"])
        history = dashboard.views["week"]["pages"]["26WK34"]["sections"][-1]["data"]["granularities"]["week"]["rows"]
        self.assertEqual(history, [["26WK34", 443, "—", 412, "—", 320, "—", 280, "—"]])

    def test_group_overview_ranks_model_net_order_and_delivery_lock_drops(self):
        model_series = {
            "问界 M5 2025款": {
                "留存大定": {
                    "26WK34": {"metrics": {"净大定": 100}},
                    "26WK35": {"metrics": {"净大定": 60}},
                },
                "交车锁单": {
                    "26WK34": {"metrics": {"交车锁单": 80}},
                    "26WK35": {"metrics": {"交车锁单": 40}},
                },
            },
            "智界 R7 2026款": {
                "留存大定": {
                    "26WK34": {"metrics": {"净大定": 50}},
                    "26WK35": {"metrics": {"净大定": 45}},
                },
                "交车锁单": {
                    "26WK34": {"metrics": {"交车锁单": 40}},
                    "26WK35": {"metrics": {"交车锁单": 60}},
                },
            },
        }

        rankings = _drop_rankings(model_series, "26WK35", "26WK34")

        self.assertEqual([row["label"] for row in rankings["留存大定"]], ["问界 M5 2025款", "智界 R7 2026款"])
        self.assertEqual([row["rank"] for row in rankings["留存大定"]], [1, 2])
        self.assertAlmostEqual(rankings["留存大定"][0]["rate"], 0.4)
        self.assertEqual([row["label"] for row in rankings["交车锁单"]], ["问界 M5 2025款"])
        self.assertAlmostEqual(rankings["交车锁单"][0]["rate"], 0.5)

    def test_group_rankings_keep_only_latest_available_model_year(self):
        order_book = Workbook()
        lock_book = Workbook()
        models = (
            ("问界 M9 2025款", 100, 50),
            ("问界 M9 2026款", 120, 60),
            ("问界 M9 2027款", 140, None),
            ("问界 M9 Ultimate 2025款", 80, 40),
            ("问界 M9 Ultimate 2026款", 90, 45),
            ("智界 R7 2025款", 70, 35),
        )
        for book, index, metric in ((order_book, 1, "净大定"), (lock_book, 2, "交车锁单")):
            for model_index, row in enumerate(models):
                value = row[index]
                if value is None:
                    continue
                sheet = book.active if model_index == 0 else book.create_sheet()
                sheet.title = f"{row[0]}by周"
                sheet.cell(1, 4, "26WK34")
                sheet.cell(2, 1, metric)
                sheet.cell(2, 2, "数量")
                sheet.cell(2, 3, "数量")
                sheet.cell(2, 4, value)
        store = WorkbookStore(Path("."))
        store.items = [
            WorkbookItem(Path("鸿蒙智行大定选配比例分析.xlsx"), order_book),
            WorkbookItem(Path("鸿蒙智行锁单选配比例分析.xlsx"), lock_book),
        ]

        series = _group_model_series(store, "week")

        self.assertEqual(
            list(series),
            ["问界 M9 2026款", "问界 M9 Ultimate 2026款", "智界 R7 2025款"],
        )
        self.assertEqual(series["问界 M9 2026款"]["留存大定"]["26WK34"]["metrics"]["净大定"], 120)

    def test_drop_rankings_require_a_previous_positive_period_and_actual_decline(self):
        model_series = {
            "问界 M5 2025款": {
                "留存大定": {"本期": {"metrics": {"净大定": 10}}},
                "交车锁单": {"上期": {"metrics": {"交车锁单": 0}}, "本期": {"metrics": {"交车锁单": 0}}},
            },
        }

        self.assertEqual(_drop_rankings(model_series, "本期", None), {})
        self.assertEqual(_drop_rankings(model_series, "本期", "上期"), {})

    def test_rise_rankings_sort_by_actual_increase_and_ignore_non_risers(self):
        model_series = {
            "问界 M5 2025款": {
                "留存大定": {
                    "上期": {"metrics": {"净大定": 50}},
                    "本期": {"metrics": {"净大定": 75}},
                },
                "交车锁单": {
                    "上期": {"metrics": {"交车锁单": 40}},
                    "本期": {"metrics": {"交车锁单": 60}},
                },
            },
            "智界 R7 2026款": {
                "留存大定": {
                    "上期": {"metrics": {"净大定": 100}},
                    "本期": {"metrics": {"净大定": 120}},
                },
                "交车锁单": {
                    "上期": {"metrics": {"交车锁单": 50}},
                    "本期": {"metrics": {"交车锁单": 50}},
                },
            },
        }

        rankings = _rise_rankings(model_series, "本期", "上期")

        self.assertEqual([row["label"] for row in rankings["留存大定"]], ["问界 M5 2025款", "智界 R7 2026款"])
        self.assertAlmostEqual(rankings["留存大定"][0]["rate"], 0.5)
        self.assertEqual([row["label"] for row in rankings["交车锁单"]], ["问界 M5 2025款"])
        self.assertAlmostEqual(rankings["交车锁单"][0]["rate"], 0.5)

    def test_overview_validation_logs_compact_period_ranges(self):
        periods = ["26WK01", "26WK02", "26WK03", "26WK04"]
        self.assertEqual(
            _period_range_text(["26WK01", "26WK02", "26WK04"], periods),
            "26WK01–26WK02（2期）、26WK04",
        )
        subject = Subject("generation_test", "问界 M9 2026款", "generation", "问界")
        issues = [("version", period, ()) for period in periods]
        with self.assertLogs("modules.overview", level="WARNING") as captured:
            _log_lock_share_issues(subject, "week", periods, issues)
        self.assertEqual(len(captured.output), 1)
        self.assertIn("缺少同期版本结构=26WK01–26WK04（4期）", captured.output[0])

    def test_overview_sparse_child_period_is_zero_not_missing(self):
        order_book = Workbook()
        order_sheet = order_book.active
        order_sheet.title = "问界by天"
        unrelated_generation = order_book.create_sheet("问界 M7 2026款by天")
        lock_book = Workbook()
        parent = lock_book.active
        parent.title = "问界by天"
        child_a = lock_book.create_sheet("问界 M5 2026款by天")
        child_b = lock_book.create_sheet("问界 M9 2026款by天")

        def metric_row(sheet, periods, values, metric="交车锁单"):
            for column, period in enumerate(periods, start=4):
                sheet.cell(1, column, period)
            sheet.cell(2, 1, metric)
            sheet.cell(2, 2, "数量")
            sheet.cell(2, 3, "数量")
            for column, value in enumerate(values, start=4):
                sheet.cell(2, column, value)

        dates = [f"2026-09-0{day}" for day in (1, 2, 3)]
        metric_row(order_sheet, dates, [5, 4, 7], "大定")
        metric_row(unrelated_generation, dates, [1, 1, 1], "大定")
        metric_row(parent, dates, [5, 4, 7])
        metric_row(child_a, [dates[0], dates[2]], [5, 7])
        metric_row(child_b, [dates[1]], [4])
        store = WorkbookStore(Path("."))
        store.items = [
            WorkbookItem(Path("鸿蒙智行大定选配比例分析.xlsx"), order_book),
            WorkbookItem(Path("鸿蒙智行锁单选配比例分析.xlsx"), lock_book),
        ]

        with self.assertNoLogs("modules.overview", level="WARNING"):
            dashboard = OverviewModule().build(store, Subject("brand_test", "问界", "brand", "鸿蒙智行"))

        pages = dashboard.views["day"]["pages"]
        first_share = next(section for section in pages["2026-09-01"]["sections"] if section["title"] == "各代际交车锁单占比")
        second_share = next(section for section in pages["2026-09-02"]["sections"] if section["title"] == "各代际交车锁单占比")
        first_titles = [section["title"] for section in pages["2026-09-01"]["sections"]]
        self.assertEqual(first_titles.index("各代际交车锁单占比") + 1, first_titles.index("门店渠道订单量"))
        self.assertEqual(first_share["kind"], "donut")
        self.assertEqual(second_share["kind"], "donut")
        self.assertEqual([row["count"] for row in first_share["data"]["rows"]], [5, 0])
        self.assertEqual([row["count"] for row in second_share["data"]["rows"]], [0, 4])

    def test_brand_overview_uses_chart_totals_for_net_order_and_delivery_lock(self):
        order_book = Workbook()
        order_sheet = order_book.active
        order_sheet.title = "问界by周"
        lock_book = Workbook()
        lock_sheet = lock_book.active
        lock_sheet.title = "问界by周"

        def metric_row(sheet, row, label, values):
            for column, period in enumerate(("26WK01", "26WK02", "总计"), start=4):
                sheet.cell(1, column, period)
            sheet.cell(row, 1, label)
            sheet.cell(row, 2, "数量")
            sheet.cell(row, 3, "数量")
            for column, value in enumerate(values, start=4):
                sheet.cell(row, column, value)

        metric_row(order_sheet, 2, "大定", (110, 130, 240))
        metric_row(order_sheet, 3, "净大定", (80, 90, 170))
        metric_row(lock_sheet, 2, "交车锁单", (60, 80, 140))
        metric_row(lock_sheet, 3, "已交付", (30, 40, 70))
        aggregate_order = order_book.create_sheet("问界 M9 2026款汇总by周")
        aggregate_lock = lock_book.create_sheet("问界 M9 2026款汇总by周")
        metric_row(aggregate_order, 2, "大定", (9999, 9999, 19998))
        metric_row(aggregate_order, 3, "净大定", (9999, 9999, 19998))
        metric_row(aggregate_lock, 2, "交车锁单", (9999, 9999, 19998))
        metric_row(aggregate_lock, 3, "已交付", (9999, 9999, 19998))

        def chart_sheet(book, title, headline, totals):
            sheet = book.create_sheet(title)
            sheet.cell(1, 1, "问界")
            sheet.cell(2, 2, headline)
            for column, period in enumerate(("26WK01", "26WK02", "总计"), start=3):
                sheet.cell(2, column, period)
            sheet.cell(3, 2, "测试结构")
            sheet.cell(3, 3, 1)
            sheet.cell(3, 4, 1)
            sheet.cell(4, 2, "总计")
            for column, value in enumerate((*totals, sum(totals)), start=3):
                sheet.cell(4, column, value)

        chart_sheet(order_book, "问界by周图表", "问界周净大定", (100, 120))
        chart_sheet(lock_book, "问界by周图表", "问界周交车锁单", (70, 90))
        store = WorkbookStore(Path("."))
        store.items = [
            WorkbookItem(Path("鸿蒙智行大定选配比例分析.xlsx"), order_book),
            WorkbookItem(Path("鸿蒙智行锁单选配比例分析.xlsx"), lock_book),
        ]

        dashboard = OverviewModule().build(store, Subject("brand_test", "问界", "brand", "鸿蒙智行"))

        view = dashboard.views["week"]
        self.assertEqual(view["periods"], ["26WK01", "26WK02"])
        self.assertNotIn("总计", view["pages"])
        page = view["pages"]["26WK02"]
        kpis = {item["label"]: item["value"] for item in page["kpis"]}
        self.assertEqual(kpis, {"大定": 130, "留存大定": 120, "交车锁单": 90, "已交付": 40})
        history = next(item for item in page["sections"] if item["kind"] == "history_table")
        rows = history["data"]["granularities"]["week"]["rows"]
        self.assertEqual([row[0] for row in rows], ["26WK01", "26WK02"])
        self.assertEqual([(row[3], row[5], row[7]) for row in rows], [(100, 70, 30), (120, 90, 40)])
        lock_trend = next(item for item in page["sections"] if item["kind"] == "stacked_trend" and "锁单" in item["title"])
        self.assertEqual(lock_trend["title"], "问界周度交车锁单趋势")
        self.assertEqual(lock_trend["data"]["headline"], "问界周交车锁单")
        self.assertEqual(lock_trend["data"]["totals"], [70, 90])

    def test_group_lock_share_display_counts_match_channel_total(self):
        lock_book = Workbook()
        group = lock_book.active
        group.title = "鸿蒙智行by天"
        brand_a = lock_book.create_sheet("问界by天")
        brand_b = lock_book.create_sheet("智界by天")

        for sheet, value in ((group, 7), (brand_a, 20), (brand_b, 10)):
            sheet.cell(1, 4, "2026-09-01")
            sheet.cell(2, 1, "交车锁单")
            sheet.cell(2, 2, "数量")
            sheet.cell(2, 3, "数量")
            sheet.cell(2, 4, value)
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("鸿蒙智行锁单选配比例分析.xlsx"), lock_book)]

        _, rows = _lock_share_rows(
            store,
            Subject("group", "鸿蒙智行", "group"),
            "day",
            "2026-09-01",
            display_total=7,
        )

        self.assertEqual([row["count"] for row in rows], [5, 2])
        self.assertEqual(sum(row["count"] for row in rows), 7)
        self.assertAlmostEqual(rows[0]["share"], 2 / 3)
        self.assertAlmostEqual(rows[1]["share"], 1 / 3)

    def test_small_order_rhythm_is_before_launch_rhythm(self):
        from core.discovery import MODULE_ORDER

        self.assertLess(MODULE_ORDER.index("cancellation"), MODULE_ORDER.index("launch_rhythm"))

    def test_metric_label_is_not_treated_as_a_subject(self):
        self.assertTrue(is_valid_subject_name("尊界"))
        self.assertTrue(is_valid_subject_name("尊界 S800 2025款"))
        self.assertFalse(is_valid_subject_name("尊界交车锁单"))

    def test_generation_is_discovered_from_sheet_prefix(self):
        name = sheet_subject("问界 M6 2026款by周")
        self.assertEqual(name, "问界 M6 2026款")
        self.assertEqual(subject_type(name), "generation")
        self.assertEqual(subject_parent(name, "generation"), "问界")

    def test_hourly_launch_sheet_is_recognized_and_rendered(self):
        workbook = Workbook()
        day_sheet = workbook.active
        day_sheet.title = "问界 M6 2026款by天"
        day_sheet.cell(1, 2, "D1")
        day_sheet.cell(2, 2, 46134)
        for row, (label, value) in enumerate([
            ("当日大定数量", 10),
            ("当日净大定数量", 9),
            ("当日小订转大定数量", 6),
            ("累计小订转大定数量", 6),
            ("累计小订转化率", 0.6),
            ("当日直接大定数量", 4),
            ("累计直接大定数量", 4),
            ("累计直接大定进度", 0.4),
            ("当日交车锁单数量", 8),
            ("累计交车锁单数量", 8),
        ], start=3):
            day_sheet.cell(row, 1, label)
            day_sheet.cell(row, 2, value)
        hourly_sheet = workbook.create_sheet("问界 M6 2026款by时")
        hourly_sheet.cell(1, 2, "H1")
        hourly_sheet.cell(2, 2, datetime(2026, 4, 22, 0, 0))
        for row, (label, value) in enumerate([
            ("当日大定数量", 10),
            ("当日净大定数量", 9),
            ("当日小订转大定数量", 6),
            ("累计小订转大定数量", 6),
            ("累计小订转化率", 0.6),
            ("当日直接大定数量", 4),
            ("累计直接大定数量", 4),
            ("累计直接大定进度", 0.4),
            ("当日交车锁单数量", 8),
            ("累计交车锁单数量", 8),
        ], start=3):
            hourly_sheet.cell(row, 1, label)
            hourly_sheet.cell(row, 2, value)
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("鸿蒙智行首销期订单节奏.xlsx"), workbook)]
        subject = Subject("generation_test", "问界 M6 2026款", "generation", "问界")
        dashboard = LaunchRhythmModule().build(store, subject)
        self.assertEqual(grain_from_sheet(hourly_sheet.title), "hour")
        self.assertEqual(dashboard.views["day"]["default_period"], "2026-04-22")
        page = dashboard.views["day"]["pages"][dashboard.views["day"]["default_period"]]
        hourly_section = next(item for item in page["sections"] if item["kind"] == "launch_hourly")
        self.assertEqual(hourly_section["title"], "分时首销节奏")
        self.assertEqual(hourly_section["meta"], "首销开启 2026-04-22 00:00 · 识别高峰时段与大定波动")
        self.assertEqual(hourly_section["data"][0]["period"], "2026-04-22")
        self.assertEqual(hourly_section["data"][0]["small"], 6)
        self.assertEqual(hourly_section["data"][0]["direct"], 4)
        week_sheet = workbook.create_sheet("问界 M6 2026款by周")
        for row in range(1, day_sheet.max_row + 1):
            for column in range(1, day_sheet.max_column + 1):
                week_sheet.cell(row, column, day_sheet.cell(row, column).value)
        week_sheet.cell(2, 2, "26WK24")
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("鸿蒙智行首销期订单节奏.xlsx"), workbook)]
        dashboard = LaunchRhythmModule().build(store, subject)
        week_page = dashboard.views["week"]["pages"][dashboard.views["week"]["default_period"]]
        self.assertEqual(dashboard.views["week"]["default_period"], "26WK24")
        self.assertIn("launch_hourly", [item["kind"] for item in week_page["sections"]])

    def test_generation_name_ending_in_total_is_preserved(self):
        name = sheet_subject("问界 M9 2026款总计by周")
        self.assertEqual(name, "问界 M9 2026款总计")
        self.assertTrue(is_valid_subject_name(name))
        self.assertEqual(subject_type(name), "generation")
        self.assertEqual(subject_parent(name, "generation"), "问界")
        self.assertEqual(
            sheet_subject("问界 M9 2026款总计_区域分析"),
            "问界 M9 2026款总计",
        )

    def test_total_generation_is_not_merged_with_base_generation(self):
        workbook = Workbook()
        workbook.active.title = "问界 M9 2026款by周"
        workbook.create_sheet("问界 M9 2026款总计by周")
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("鸿蒙智行大定选配比例分析.xlsx"), workbook)]

        subjects = discover_subjects(store)
        generation_names = {item.name for item in subjects if item.type == "generation"}
        self.assertIn("问界 M9 2026款", generation_names)
        self.assertIn("问界 M9 2026款总计", generation_names)
        self.assertNotEqual(
            next(item.id for item in subjects if item.name == "问界 M9 2026款"),
            next(item.id for item in subjects if item.name == "问界 M9 2026款总计"),
        )
        found = store.find_subject_sheet("大定选配比例", "问界 M9 2026款总计", "week")
        self.assertIsNotNone(found)
        self.assertEqual(found[1].title, "问界 M9 2026款总计by周")

    def test_brand_and_group_types(self):
        self.assertEqual(subject_type("鸿蒙智行"), "group")
        self.assertEqual(subject_type("智界"), "brand")
        self.assertEqual(subject_parent("智界", "brand"), "鸿蒙智行")

    def test_sheet_suffixes_are_removed(self):
        self.assertEqual(sheet_subject("问界 M6 2026款_选配渗透率"), "问界 M6 2026款")
        self.assertEqual(sheet_subject("问界 M6 2026款_区域分析"), "问界 M6 2026款")
        self.assertEqual(sheet_subject("问界_渠道分析"), "问界")
        self.assertEqual(sheet_subject("鸿蒙智行_区域分析"), "鸿蒙智行")
        self.assertEqual(grain_from_sheet("问界 M6 2026款by月"), "month")
        self.assertEqual(compact_text("问界 M6 2026款"), "问界M62026款")

    def test_sheet_mapping_tolerates_spaces_and_separator_variants(self):
        workbook = Workbook()
        workbook.active.title = "问界 - 选配渗透率排名"
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("鸿蒙智行选配金统计.xlsx"), workbook)]
        found = store.find_subject_sheet("选配金统计", "问界", suffix=("选配渗透率排名", "选配渗透率"))
        self.assertIsNotNone(found)
        self.assertEqual(sheet_subject(workbook.active.title), "问界")
        self.assertEqual(compact_identifier("问界_选配渗透率排名"), compact_identifier("问界 - 选配渗透率排名"))

    def test_brand_option_fee_does_not_require_paid_sheet(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "问界_选配渗透率排名"
        sheet.cell(3, 1, "问界 M5 2025款")
        sheet.cell(3, 2, 1)
        sheet.cell(2, 8, "总计")
        sheet.cell(3, 8, "电动遮阳帘 77.5%\n(3000)")
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("鸿蒙智行选配金统计.xlsx"), workbook)]
        subject = Subject("brand_test", "问界", "brand", "鸿蒙智行")
        dashboard = OptionFeeModule().build(store, subject)
        self.assertIsNotNone(dashboard)
        page = dashboard.views["month"]["pages"]["汇总"]
        self.assertEqual(page["kpis"][0]["value"], 1)
        rank_item = page["sections"][0]["data"]["问界 M5 2025款"][0]
        self.assertEqual(rank_item["label"], "电动遮阳帘")
        self.assertAlmostEqual(rank_item["rate"], 0.775)
        self.assertEqual(rank_item["amount"], 3000)

    def test_subject_sheet_searches_all_matching_workbooks(self):
        first = Workbook()
        first.active.title = "鸿蒙智行_无关汇总"
        second = Workbook()
        second.active.title = "问界_选配渗透率排名"
        store = WorkbookStore(Path("."))
        store.items = [
            WorkbookItem(Path("01_鸿蒙智行选配金统计_旧版.xlsx"), first),
            WorkbookItem(Path("02_鸿蒙智行选配金统计_新版.xlsx"), second),
        ]
        found = store.find_subject_sheet("选配金统计", "问界", suffix="选配渗透率排名")
        self.assertIsNotNone(found)
        self.assertEqual(found[0].path.name, "02_鸿蒙智行选配金统计_新版.xlsx")

    def test_generic_analysis_sheet_enables_existing_subject(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "问界 M6 2026款_区域分析"
        sheet.append(["区域", "订单量"])
        sheet.append(["华东", 120])
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("区域销售分析.xlsx"), workbook)]
        subject = Subject("generation_test", "问界 M6 2026款", "generation", "问界")
        self.assertIsNotNone(GenericModule().build(store, subject))

    def test_generic_analysis_separates_routine_and_unmatched_sheets(self):
        routine_book = Workbook()
        routine_book.active.title = "问界 M6 2026款by周"
        unmatched_book = Workbook()
        unmatched_book.active.title = "问界 M6 2026款_区域分析"
        unmatched_book.active.append(["区域", "订单量"])
        unmatched_book.active.append(["华东", 120])
        store = WorkbookStore(Path("."))
        store.items = [
            WorkbookItem(Path("鸿蒙智行大定选配比例分析.xlsx"), routine_book),
            WorkbookItem(Path("区域销售分析.xlsx"), unmatched_book),
        ]
        subject = Subject("generation_test", "问界 M6 2026款", "generation", "问界")
        with self.assertLogs("modules.generic", level="WARNING") as captured:
            dashboard = GenericModule().build(store, subject)
        page = dashboard.views["week"]["pages"]["汇总"]
        self.assertEqual(page["sections"][0]["title"], "待归类数据")
        self.assertEqual(page["sections"][1]["title"], "已识别数据")
        self.assertEqual([item["label"] for item in page["kpis"][:3]], ["待归类Sheet", "已识别Sheet", "需核查Sheet"])
        self.assertEqual(page["kpis"][0]["value"], 1)
        self.assertEqual(page["kpis"][1]["value"], 1)
        log_text = "\n".join(captured.output)
        self.assertIn("[数据诊断]", log_text)
        self.assertIn("[映射诊断]", log_text)
        self.assertIn("回退=导入分析通用预览", log_text)

    def test_small_order_chart_title_does_not_create_extra_subject(self):
        self.assertEqual(sheet_subject("鸿蒙智行净小订by天图表"), "鸿蒙智行")

    def test_brands_are_ordered_by_business_priority(self):
        workbook = Workbook()
        for name in ("尚界by周", "尊界by周", "享界by周", "智界by周", "问界by周"):
            workbook.create_sheet(name)
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("测试.xlsx"), workbook)]
        subjects = discover_subjects(store)
        brand_names = [subject.name for subject in subjects if subject.type == "brand"]
        self.assertEqual(brand_names, ["问界", "智界", "享界", "尊界", "尚界"])
        self.assertEqual(subjects[0].name, "鸿蒙智行")

    def test_forecast_target_without_raw_sheet_is_in_generation_selector(self):
        workbook = Workbook()
        workbook.active.title = "尊界by周"
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("测试.xlsx"), workbook)]
        subjects = discover_subjects(store, ["尊界 V680&V800", "享界 G9", "数据来源目录"])
        generations = {item.name: item for item in subjects if item.type == "generation"}
        self.assertEqual(generations["尊界 V680&V800"].parent, "尊界")
        self.assertEqual(generations["享界 G9"].parent, "享界")
        self.assertNotIn("数据来源目录", generations)

    def test_combined_code_and_edition_are_generations_not_metric_labels(self):
        for name in ("尊界 V680&V800", "尊界 S800典藏大观"):
            self.assertTrue(is_valid_subject_name(name))
            self.assertEqual(subject_type(name), "generation")
        self.assertFalse(is_valid_subject_name("尊界交车锁单"))
        self.assertEqual(subject_type("尊界交车锁单"), "brand")

    def test_select_retreat_groups_are_parsed(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "问界 M6 2026款_选配退订"
        sheet.append(["地理区域退订比例"])
        sheet.append(["地理区域", "小订数", "总退订数", "小订后退订比例", "大定后退订比例", "总退订比例"])
        sheet.append(["东北", 1601, 93, 0.043, 0.014, 0.058])
        sheet.append(["华东", 1672, 14, 0.006, 0.002, 0.008])
        sheet.append([])
        sheet.append(["版本退订比例"])
        sheet.append(["版本", "小订数", "总退订数", "小订后退订比例", "大定后退订比例", "总退订比例"])
        sheet.append(["四驱 Ultra", 1796, 28, 0.011, 0.003, 0.015])
        groups = parse_select_retreat(sheet)
        self.assertEqual([group["name"] for group in groups], ["地理区域退订比例", "版本退订比例"])
        self.assertEqual(groups[0]["rows"][0]["label"], "东北")
        self.assertAlmostEqual(groups[0]["rows"][0]["share"], 0.058)
        self.assertEqual(groups[0]["rows"][0]["count"], 93)
        self.assertEqual(groups[1]["rows"][0]["label"], "四驱 Ultra")

    def test_select_retreat_groups_follow_reordered_headers(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "问界 M6 2026款_选配退订"
        sheet.append(["版本退订比例"])
        sheet.append(["总退订比例", "版本", "大定后退订比例", "总退订数", "小订数", "小订后退订比例"])
        sheet.append([0.015, "四驱 Ultra", 0.003, 28, 1796, 0.011])

        groups = parse_select_retreat(sheet)

        self.assertEqual(groups[0]["rows"][0]["label"], "四驱 Ultra")
        self.assertEqual(groups[0]["rows"][0]["orders"], 1796)
        self.assertEqual(groups[0]["rows"][0]["count"], 28)
        self.assertAlmostEqual(groups[0]["rows"][0]["share"], 0.015)

    def test_conversion_follows_two_level_headers_after_column_reorder(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "问界 M6 2026款by周"
        sheet.append(["已交付", "周", "大定到交车锁单（天）", "大定", "销售锁单", "已发运", "交车锁单到交付（天）", "留存大定", "已到店", "交车锁单"])
        sheet.append(["数量", None, None, None, "数量", "数量", None, "数量", "数量", "数量"])
        sheet.append([70, "26WK35", 3.5, 100, 85, 78, 18, 95, 74, 80])
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("鸿蒙智行订单7级转化报表.xlsx"), workbook)]
        subject = Subject("generation_test", "问界 M6 2026款", "generation", "问界")

        dashboard = ConversionModule().build(store, subject)

        page = dashboard.views["week"]["pages"]["26WK35"]
        funnel = next(item for item in page["sections"] if item["kind"] == "funnel")
        self.assertEqual([row["value"] for row in funnel["data"]], [100, 95, 85, 80, 78, 74, 70])
        matrix = next(item for item in page["sections"] if item["kind"] == "cohort_matrix")
        self.assertEqual(matrix["data"][0]["order_to_lock_days"], 3.5)
        self.assertEqual(matrix["data"][0]["lock_to_delivery_days"], 18)

    def test_excel_display_formats_are_dynamic(self):
        self.assertEqual(format_excel_value(0.1234, "0.00%"), "12.34%")
        self.assertEqual(format_excel_value(1234.5, "#,##0.00"), "1,234.50")
        self.assertEqual(format_excel_value(datetime(2026, 8, 19), "yyyy/mm/dd"), "2026/08/19")
        self.assertAlmostEqual(safe_number("12.5%"), 0.125)

    def test_kpis_are_compared_with_the_previous_period(self):
        dashboards = {"subject|overview": {"views": {"week": {
            "periods": ["26WK32", "26WK33"],
            "pages": {
                "26WK32": {"kpis": [{"label": "大定", "value": 100}]},
                "26WK33": {"kpis": [{"label": "大定", "value": 120}]},
            },
        }}}}
        add_kpi_comparisons(dashboards)
        comparison = dashboards["subject|overview"]["views"]["week"]["pages"]["26WK33"]["kpis"][0]["comparison"]
        self.assertEqual(comparison["direction"], "up")
        self.assertEqual(comparison["delta"], 20)
        self.assertAlmostEqual(comparison["rate"], 0.2)

    def test_aggregate_period_does_not_compare_against_a_single_day(self):
        dashboards = {"subject|overview": {"views": {"day": {
            "periods": ["2026-08-18", "近28天"],
            "pages": {
                "2026-08-18": {"kpis": [{"label": "大定", "value": 100}]},
                "近28天": {"kpis": [{"label": "大定", "value": 2800}]},
            },
        }}}}
        add_kpi_comparisons(dashboards)
        comparison = dashboards["subject|overview"]["views"]["day"]["pages"]["近28天"]["kpis"][0]["comparison"]
        self.assertEqual(comparison["status"], "unavailable")

    def test_selection_dashboard_excludes_order_and_conversion_kpis_and_chart_sheet(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "问界 M6 2026款by天"
        sheet.cell(1, 4, "2026-08-05")
        sheet.cell(2, 1, "小订")
        sheet.cell(2, 2, "数量")
        sheet.cell(2, 3, "数量")
        sheet.cell(2, 4, 100)
        sheet.cell(3, 1, "外观")
        sheet.cell(3, 2, "占比")
        sheet.cell(3, 3, "黑色")
        sheet.cell(3, 4, 0.6)
        sheet.cell(4, 2, "数量")
        sheet.cell(4, 3, "黑色")
        sheet.cell(4, 4, 60)
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("鸿蒙智行小订选配比例分析.xlsx"), workbook)]
        subject = Subject("generation_test", "问界 M6 2026款", "generation", "问界")
        dashboard = SmallOrderMixModule().build(store, subject)
        page = dashboard.views["day"]["pages"]["2026-08-05"]
        self.assertNotIn("stacked_trend", [item["kind"] for item in page["sections"]])
        labels = [item["label"] for item in page["kpis"]]
        self.assertNotIn("小订", labels)
        self.assertFalse(any("转化率" in label for label in labels))

    def test_selection_dashboard_does_not_render_chart_sheet_blocks(self):
        workbook = Workbook()
        data_sheet = workbook.active
        data_sheet.title = "问界by周"
        data_sheet.cell(1, 4, "26WK34")
        data_sheet.cell(2, 1, "版本")
        data_sheet.cell(2, 2, "占比")
        data_sheet.cell(2, 3, "Max")
        data_sheet.cell(2, 4, 0.6)
        data_sheet.cell(3, 2, "数量")
        data_sheet.cell(3, 3, "Max")
        data_sheet.cell(3, 4, 60)
        chart_sheet = workbook.create_sheet("问界by周图表")
        for start, generation, share, total in (
            (1, "问界 M6 2026款", 0.6, 100),
            (25, "问界 M9 2026款", 0.7, 120),
        ):
            chart_sheet.cell(start, 1, generation)
            chart_sheet.cell(start + 1, 2, f"{generation}周净大定")
            chart_sheet.cell(start + 1, 3, "26WK34")
            chart_sheet.cell(start + 1, 4, "总计")
            chart_sheet.cell(start + 2, 2, "Max")
            chart_sheet.cell(start + 2, 3, share)
            chart_sheet.cell(start + 2, 4, share)
            chart_sheet.cell(start + 3, 2, "总计")
            chart_sheet.cell(start + 3, 3, total)
            chart_sheet.cell(start + 3, 4, total)
        for module, file_name in (
            (OrderMixModule(), "鸿蒙智行大定选配比例分析.xlsx"),
            (LockMixModule(), "鸿蒙智行锁单选配比例分析.xlsx"),
            (SmallOrderMixModule(), "鸿蒙智行小订选配比例分析.xlsx"),
        ):
            with self.subTest(module=module.id):
                store = WorkbookStore(Path("."))
                store.items = [WorkbookItem(Path(file_name), workbook)]
                dashboard = module.build(store, Subject("brand_test", "问界", "brand", "鸿蒙智行"))
                page = dashboard.views["week"]["pages"]["26WK34"]
                chart_sections = [item for item in page["sections"] if item["kind"] == "stacked_trend"]
                self.assertEqual(chart_sections, [])
                self.assertFalse(any(source.sheet == "问界by周图表" for source in dashboard.sources))

    def test_final_chart_block_uses_inferred_finite_image_boundary(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "问界by周图表"
        for row, generation in (
            (1, "问界 M5 2025款"),
            (25, "问界 M6 2026款"),
            (49, "问界 M9 2026款"),
            (73, "问界 M9 Ultimate 2026款"),
            (97, "问界 F2N 2026款"),
        ):
            sheet.cell(row, 1, generation)
        sheet.cell(102, 2, "最后一行数据")

        image = SimpleNamespace(
            anchor=SimpleNamespace(_from=SimpleNamespace(row=102, col=1)),
            format="png",
            _data=lambda: b"image-bytes",
        )
        sheet._images.append(image)
        store = WorkbookStore(Path("."))
        item = WorkbookItem(Path("鸿蒙智行大定选配比例分析.xlsx"), workbook)

        image_key = store._extract_chart_image(item, sheet, 97, compact_text("问界 F2N 2026款"))

        self.assertIsNotNone(image_key)
        self.assertIn(image_key, store.chart_assets)

    def test_final_chart_block_does_not_claim_image_beyond_inferred_boundary(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "问界by周图表"
        for row, generation in ((1, "问界 M5 2025款"), (25, "问界 F2N 2026款")):
            sheet.cell(row, 1, generation)
        image = SimpleNamespace(
            anchor=SimpleNamespace(_from=SimpleNamespace(row=60, col=1)),
            format="png",
            _data=lambda: b"unrelated-image",
        )
        sheet._images.append(image)
        store = WorkbookStore(Path("."))
        item = WorkbookItem(Path("鸿蒙智行大定选配比例分析.xlsx"), workbook)

        image_key = store._extract_chart_image(item, sheet, 25, compact_text("问界 F2N 2026款"))

        self.assertIsNone(image_key)

    def test_small_order_rhythm_combines_mix_and_cancellation_sources(self):
        mix_book = Workbook()
        mix_sheet = mix_book.active
        mix_sheet.title = "问界 M6 2026款by天"
        mix_sheet.cell(1, 4, "2026-08-05")
        mix_sheet.cell(1, 5, "近28天")
        mix_sheet.cell(2, 1, "小订")
        mix_sheet.cell(2, 2, "数量")
        mix_sheet.cell(2, 3, "数量")
        mix_sheet.cell(2, 4, 100)
        mix_sheet.cell(2, 5, 2800)
        mix_sheet.cell(3, 1, "留存小订")
        mix_sheet.cell(3, 2, "数量")
        mix_sheet.cell(3, 3, "数量")
        mix_sheet.cell(3, 4, 90)
        mix_sheet.cell(3, 5, 2400)

        cancel_book = Workbook()
        day_sheet = cancel_book.active
        day_sheet.title = "问界 M6 2026款_日度退订"
        day_sheet.append(["日度退订"])
        day_sheet.append(["取消日期", "当日小订退", "累计小订退", "当日退订数", "整体 - 小订转大定 %", "整体 - 小订退 %", "整体 - 大定退 %"])
        day_sheet.cell(3, 1, datetime(2026, 8, 5))
        day_sheet.cell(3, 2, 10)
        day_sheet.cell(3, 3, 10)
        day_sheet.cell(3, 4, 3)
        day_sheet.cell(3, 5, 0.2)
        day_sheet.cell(3, 6, 0.1)
        day_sheet.cell(3, 7, 0.03)
        hourly_sheet = cancel_book.create_sheet("问界 M6 2026款_小订分时退订")
        hourly_sheet.append(["小订分时退订"])
        hourly_sheet.append(["小订日期", "小订小时", "小订数", "总退订数"])
        hourly_sheet.cell(3, 1, datetime(2026, 8, 5, 0, 0))
        hourly_sheet.cell(3, 2, 10)
        hourly_sheet.cell(3, 3, 20)
        hourly_sheet.cell(3, 4, 2)

        store = WorkbookStore(Path("."))
        store.items = [
            WorkbookItem(Path("鸿蒙智行小订选配比例分析.xlsx"), mix_book),
            WorkbookItem(Path("鸿蒙智行小订退订分析.xlsx"), cancel_book),
        ]
        subject = Subject("generation_test", "问界 M6 2026款", "generation", "问界")
        dashboard = CancellationModule().build(store, subject)
        view = dashboard.views["day"]
        self.assertEqual(view["periods"][0], "2026-08-05")
        page = view["pages"][view["default_period"]]
        rhythm = next(item for item in page["sections"] if item["kind"] == "small_order_rhythm")
        self.assertEqual(len(rhythm["data"]), 1)
        # 小订节奏按分时退订Sheet聚合：小订20、退订2、留存18
        self.assertEqual(rhythm["data"][0]["orders"], 20)
        self.assertEqual(rhythm["data"][0]["retained"], 18)
        self.assertEqual(rhythm["data"][0]["cancel"], 2)
        hourly_section = next(item for item in page["sections"] if item["kind"] == "hourly")
        self.assertEqual(hourly_section["data"][0]["period"], "2026-08-05")
        self.assertNotIn("structure_cards", [item["kind"] for item in page["sections"]])
        self.assertFalse(any("选配" in item["title"] for item in page["sections"]))

    def test_cancellation_follows_reordered_day_and_hour_headers(self):
        mix_book = Workbook()
        mix_sheet = mix_book.active
        mix_sheet.title = "问界 M6 2026款by天"
        mix_sheet.cell(1, 4, "2026-08-05")
        mix_sheet.cell(2, 1, "小订")
        mix_sheet.cell(2, 2, "数量")
        mix_sheet.cell(2, 3, "数量")
        mix_sheet.cell(2, 4, 100)

        cancel_book = Workbook()
        day_sheet = cancel_book.active
        day_sheet.title = "问界 M6 2026款_日度退订"
        day_sheet.append(["日度退订"])
        day_sheet.append(["整体 - 大定退 %", "累计小订退", "取消日期", "整体 - 小订退 %", "当日退订数", "整体 - 小订转大定 %", "当日小订退"])
        day_sheet.append([0.03, 10, datetime(2026, 8, 5), 0.1, 3, 0.2, 10])
        hourly_sheet = cancel_book.create_sheet("问界 M6 2026款_小订分时退订")
        hourly_sheet.append(["小订分时退订"])
        hourly_sheet.append(["总退订数", "小订数", "小订日期", "小订小时"])
        hourly_sheet.append([2, 20, datetime(2026, 8, 5), 10])

        store = WorkbookStore(Path("."))
        store.items = [
            WorkbookItem(Path("鸿蒙智行小订选配比例分析.xlsx"), mix_book),
            WorkbookItem(Path("鸿蒙智行小订退订分析.xlsx"), cancel_book),
        ]
        subject = Subject("generation_test", "问界 M6 2026款", "generation", "问界")
        dashboard = CancellationModule().build(store, subject)
        page = dashboard.views["day"]["pages"]["2026-08-05"]

        rhythm = next(item for item in page["sections"] if item["kind"] == "small_order_rhythm")["data"][0]
        self.assertEqual((rhythm["orders"], rhythm["cancel"], rhythm["retained"]), (20, 2, 18))
        daily = next(item for item in page["sections"] if item["kind"] == "daily_cancel")["data"][0]
        self.assertEqual((daily["daily"], daily["big_daily"]), (10, 3))
        self.assertEqual((daily["convert_rate"], daily["small_rate"], daily["big_rate"]), (0.2, 0.1, 0.03))


if __name__ == "__main__":
    unittest.main()
