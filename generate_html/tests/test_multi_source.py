from __future__ import annotations

import os
import unittest
from datetime import date, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from core.excel import WorkbookItem, WorkbookStore, load_data_workbook, parse_metric_sheet
from core.models import Dashboard, SourceRef, Subject
from modules.sales_forecast import _read_actual_profiles, _read_steady_history
from modules.conversion import ConversionModule, STAGES
from modules.order_mix import OrderMixModule
from modules.option_fee import _period_groups, _option_layout, _version_rows_by_period
from tools.refresh_sales_forecast_data import append_forecast_inputs
from core.forecast_summary import summary_scope


MODEL = "问界 M9 2026款"


class MultiSourceTests(unittest.TestCase):
    def store(self, books, keyword="大定选配比例分析"):
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path(f"{keyword}{2021 + index}.xlsx"), book)
                       for index, book in enumerate(books)]
        self.addCleanup(store.close)
        return store

    def metric(self, periods, rows, title=MODEL + "by天"):
        book = Workbook()
        sheet = book.active
        sheet.title = title
        sheet.append(["指标", "统计类型", "分类", *periods])
        for row in rows:
            sheet.append(row)
        return book

    def test_metric_union_aligns_labels_dates_and_preserves_zero_and_blank(self):
        first = self.metric(["2021-01-01", "2022-01-01", "总计"], [
            ["大定", "数量", "数量", 0, None, 15],
            ["留存大定", "数量", "数量", 8, None, 8],
            ["动力", "占比", "增程", .4, None, .4],
            [None, None, "纯电", .6, None, .6],
        ])
        second = self.metric(["2022-01-01", "2021-01-01", "2023-01-01", "总计"], [
            ["留存大定", "数量", "数量", 18, 8, 28, 54],
            ["大定", "数量", "数量", 20, 99, 30, 50],
            ["动力", "占比", "纯电", .7, .6, .8, .7],
            [None, None, "增程", .3, .4, .2, .3],
        ])
        second.active["D4"].number_format = "0.00%"
        store = self.store([first, second])
        with self.assertLogs("core.excel", level="WARNING") as logs:
            item, merged = store.find_subject_sheet("大定选配比例", MODEL, "day")
        parsed = parse_metric_sheet(merged)
        self.assertEqual(list(parsed), ["2021-01-01", "2022-01-01", "2023-01-01", "总计"])
        self.assertEqual([parsed[day]["metrics"]["大定"] for day in parsed], [0, 20, 30, 15])
        self.assertEqual(parsed["2022-01-01"]["structures"]["动力"][0]["share"], .3)
        self.assertEqual(len(logs.output), 1)
        self.assertIn("多文件数值冲突", logs.output[0])
        self.assertIs(store.find_subject_sheet("大定选配比例", MODEL, "day")[1], merged)
        self.assertEqual(first.active["D2"].value, 0)
        self.assertIsNone(first.active["E2"].value)
        dashboard = Dashboard("test", "test", {"source": SourceRef(item.path.name, merged.title).to_dict()},
                              [SourceRef(item.path.name, merged.title)])
        store.resolve_dashboard_sources(dashboard)
        self.assertEqual(len(dashboard.sources), 2)
        self.assertEqual(len(dashboard.views["source"]["source_files"]), 2)

    def test_many_conflicting_sheets_have_one_keyword_summary(self):
        first = self.metric(["2026-01-01"], [["大定", "数量", "数量", 10]])
        second = self.metric(["2026-01-01"], [["大定", "数量", "数量", 20]])
        first.create_sheet("问界 M6 2026款by天")
        second.create_sheet("问界 M6 2026款by天")
        for sheet, value in ((first.worksheets[1], 5), (second.worksheets[1], 6)):
            sheet.append(["大定", "数量", "数量", "2026-01-01"])
            sheet.append(["大定", "数量", "数量", value])
        store = self.store([first, second])
        with self.assertLogs("core.excel", level="WARNING") as logs:
            store.find("大定选配比例")
        self.assertEqual(len(logs.output), 1)
        self.assertIn("2个Sheet共2项", logs.output[0])

    def test_subject_in_second_file_keeps_real_file_link(self):
        first = self.metric(["2021-01-01"], [["大定", "数量", "数量", 10]], "问界 M5 2025款by天")
        second = self.metric(["2022-01-01"], [["大定", "数量", "数量", 20]])
        store = self.store([first, second])
        item, _ = store.find_subject_sheet("大定选配比例", MODEL, "day")
        self.assertEqual(item.path.name, "大定选配比例分析2022.xlsx")
        self.assertEqual(len(store.items), 2)

    def test_duplicate_mix_keeps_summary_column_order_and_comparisons(self):
        book = self.metric(["2026-01-01", "2026-01-02", "近28天", "总计"], [
            ["大定", "数量", "数量", 10, 20, 30, 30],
            ["动力", "占比", "增程", .4, .6, .5, .6],
            [None, None, "纯电", .6, .4, .5, .4],
        ])
        subject = Subject("m9", MODEL, "generation", "问界")
        original = OrderMixModule().build(self.store([book]), subject)
        duplicate = OrderMixModule().build(self.store([book, book]), subject)
        self.assertEqual(duplicate.views, original.views)
        self.assertEqual(duplicate.views["day"]["periods"][-2:], ["近28天", "总计"])

    def test_chart_reads_periods_and_subjects_from_every_matching_file(self):
        books = []
        for period, count in [("2021-01-01", 10), ("2022-01-01", 20)]:
            book = Workbook()
            sheet = book.active
            sheet.title = "问界净大定by天图表"
            sheet.append([MODEL])
            sheet.append([None, "留存大定", period])
            sheet.append([None, "版本A", 1])
            sheet.append([None, "总计", count])
            books.append(book)
        store = self.store(books)
        data = store.find_chart_data("大定选配比例", MODEL, "day")[2]
        self.assertEqual(data["periods"], ["2021-01-01", "2022-01-01"])
        self.assertEqual(data["totals"], [10, 20])
        self.assertEqual(store.find_chart_blocks("大定选配比例", MODEL, "day")[2][0]["totals"], [10, 20])

    def test_launch_and_small_hourly_union_do_not_double_count(self):
        books = []
        for dates, counts in [
            ([datetime(2026, 1, 1), datetime(2026, 1, 2)], [0, 20]),
            ([datetime(2026, 1, 2), datetime(2026, 1, 3)], [20, 30]),
        ]:
            book = Workbook()
            sheet = book.active
            sheet.title = MODEL + "by天"
            sheet.append(["周期", "D1", "D2"])
            sheet.append(["时间", *dates])
            sheet.append(["当日大定数量", *counts])
            books.append(book)
        store = self.store(books, "首销期订单节奏")
        with patch("modules.sales_forecast._read_model_mapping", return_value={}):
            profiles, _ = _read_actual_profiles(store)
        self.assertEqual([row["gross"] for row in profiles[0]["days"]], [0, 20, 30])
        self.assertEqual([row["date"] for row in profiles[0]["days"]],
                         ["2026-01-01", "2026-01-02", "2026-01-03"])

        books = []
        for hour in [0, 1]:
            book = Workbook()
            sheet = book.active
            sheet.title = MODEL + "_小订分时退订"
            sheet.append(["小订分时退订统计"])
            sheet.append(["小订日期", "小订小时", "小订数"])
            sheet.append([datetime(2026, 1, 1), hour, 10])
            books.append(book)
        store = self.store(books, "小订退订分析")
        with patch("modules.sales_forecast._read_model_mapping", return_value={}):
            profiles, _ = _read_actual_profiles(store)
        self.assertEqual(profiles[0]["small_hourly_days"][0]["orders"], 20)
        self.assertEqual(len(profiles[0]["small_hourly_days"][0]["hours"]), 2)

    def test_steady_history_merges_year_shards_into_one_reference(self):
        first = self.metric(["26WK28"], [["交车锁单", "数量", "数量", 100]], MODEL + "by周")
        second = self.metric(["26WK29"], [["交车锁单", "数量", "数量", 200]], MODEL + "by周")
        for book, day, count in [(first, datetime(2026, 7, 7), 10), (second, datetime(2026, 7, 8), 20)]:
            sheet = book.create_sheet(MODEL + "by天")
            sheet.append(["指标", "统计类型", "分类", day])
            sheet.append(["交车锁单", "数量", "数量", count])
        store = self.store([first, second], "锁单选配比例分析")
        with patch("modules.sales_forecast._read_model_mapping", return_value={}), patch(
            "modules.sales_forecast._read_model_master", return_value={}
        ):
            rows, _ = _read_steady_history(store, {"m9": {"generation": MODEL, "end_date": "2026-06-30"}}, date(2026, 7, 27))
        self.assertEqual(len(rows), 1)
        self.assertEqual([week["lock"] for week in rows[0]["weeks"]], [100, 200])
        self.assertEqual([day["lock"] for day in rows[0]["daily"]], [10, 20])

    def test_duplicate_values_are_not_logged_or_added(self):
        books = [self.metric(["2026-01-01"], [["大定", "数量", "数量", 25]]) for _ in range(2)]
        store = self.store(books)
        with patch("core.excel.LOGGER.warning") as warning:
            _, sheet = store.find_subject_sheet("大定选配比例", MODEL, "day")
        warning.assert_not_called()
        self.assertEqual(parse_metric_sheet(sheet)["2026-01-01"]["metrics"]["大定"], 25)


    def test_conversion_uses_period_not_repeated_name_and_aligns_reordered_columns(self):
        books = []
        for index, periods in enumerate((("26WK34", "26WK35"), ("26WK35", "26WK36"))):
            book = Workbook()
            sheet = book.active
            sheet.title = MODEL + "by周"
            headers = ["名称", "周", *STAGES, "大定到交车锁单（天）", "交车锁单到交付（天）"]
            if index:
                headers = list(reversed(headers))
                sheet.append(["新增说明行"])
            sheet.append(headers)
            sheet.append(["数量" if label in STAGES else None for label in headers])
            for period in periods:
                record = dict(zip(STAGES, (100, 95, 90, 80, 75, 72, 70)))
                record.update({"名称": "总计", "周": period,
                               "大定到交车锁单（天）": 3.5, "交车锁单到交付（天）": 18})
                sheet.append([record[label] for label in headers])
            books.append(book)
        store = self.store(books, "订单7级转化")
        with patch("core.excel.LOGGER.warning") as warning:
            dashboard = ConversionModule().build(store, Subject("m9", MODEL, "generation", "问界"))
        warning.assert_not_called()
        view = dashboard.views["week"]
        self.assertEqual(view["periods"], ["26WK34", "26WK35", "26WK36"])
        for page in view["pages"].values():
            funnel = next(section["data"] for section in page["sections"] if section["kind"] == "funnel")
            self.assertEqual([row["value"] for row in funnel], [100, 95, 90, 80, 75, 72, 70])
        self.assertEqual(books[0].active.max_row, 4)

    def fee_book(self, periods, instruction_rows=0):
        book = Workbook()
        sheet = book.active
        sheet.title = MODEL + "_选配渗透率"
        sheet.append(["权益金额", "6000元"])
        sheet.append(["权益内容", "测试权益"])
        for _ in range(instruction_rows):
            sheet.append(["说明"])
        period_row = 3 + instruction_rows
        for i, period in enumerate(periods):
            col = 3 + i * 4
            sheet.cell(period_row, col, period)
            for offset, band in enumerate(("免费", "0-5000", "5000-10000", "汇总")):
                sheet.cell(period_row + 1, col + offset, band)
            sheet.cell(period_row + 2, 1, "订单量")
            sheet.cell(period_row + 2, 2, "总计")
            for offset, value in enumerate((100, 200, 300, 600)):
                sheet.cell(period_row + 2, col + offset, value)
            sheet.cell(period_row + 3, 1, "Max")
            sheet.cell(period_row + 3, 2, "选配金额")
            sheet.cell(period_row + 3, col + 3, 12000)
            sheet.cell(period_row + 4, 2, "实际金额")
            sheet.cell(period_row + 4, col + 3, 6500)
        return book

    def test_option_fee_keeps_grouped_headers_subtotals_and_dynamic_header_rows(self):
        first = self.fee_book(["26-07", "26-08"])
        second = self.fee_book(["26-08", "26-09"], instruction_rows=2)
        store = self.store([first, second], "选配金统计")
        with patch("core.excel.LOGGER.warning") as warning:
            sheet = store.find("选配金统计").workbook.worksheets[0]
        warning.assert_not_called()
        self.assertEqual(_option_layout(sheet), (3, 4, 5))
        periods, groups = _period_groups(sheet)
        self.assertEqual(periods, ["26-07", "26-08", "26-09"])
        for period in periods:
            self.assertEqual([band for _, band in groups[period]], ["免费", "0-5000", "5000-10000"])
            self.assertEqual([sheet.cell(5, col).value for col, _ in groups[period]], [100, 200, 300])
            self.assertEqual(_version_rows_by_period(sheet, periods)[period][0][1:3], [12000, 6500])
        self.assertEqual(sheet.cell(1, 2).value, "6000元")
        self.assertEqual(sheet.cell(2, 2).value, "测试权益")
        self.assertEqual([sheet.cell(3, col).value for col in range(3, 7)], ["26-07", None, None, None])

    def test_option_fee_identical_duplicate_preserves_metrics(self):
        first = self.fee_book(["26-07", "26-08"])
        store = self.store([first, first], "选配金统计")
        merged = store.find("选配金统计").workbook.worksheets[0]
        periods = _period_groups(first.active)[0]
        self.assertEqual(_period_groups(merged)[0], periods)
        self.assertEqual(_version_rows_by_period(merged, periods), _version_rows_by_period(first.active, periods))

    def test_spaced_sheet_names_share_priority_and_keep_physical_source_links(self):
        first = self.metric([datetime(2026, 1, 1)], [["小订", "数量", "数量", 10]])
        second = self.metric([datetime(2026, 1, 1)], [["小订", "数量", "数量", 20]],
                             MODEL.replace(" ", "") + "by天")
        store = self.store([first, second], "小订选配比例分析")
        with self.assertLogs("core.excel", level="WARNING") as logs, patch(
            "modules.sales_forecast._read_model_mapping", return_value={}
        ):
            profiles, sources = _read_actual_profiles(store)
        self.assertEqual(len(logs.output), 1)
        self.assertEqual(profiles[0]["small_daily_days"][0]["orders"], 10)
        item, sheet = store.find_subject_sheet("小订选配比例", MODEL, "day")
        self.assertEqual(next(iter(parse_metric_sheet(sheet).values()))["metrics"]["小订"], 10)
        self.assertEqual([(s.file, s.sheet) for s in store.expand_sources(sources)], [
            ("小订选配比例分析2021.xlsx", first.active.title),
            ("小订选配比例分析2022.xlsx", second.active.title),
        ])
        self.assertEqual(len(store.find("小订选配比例").workbook.worksheets), 1)
        self.assertEqual(second.active.cell(2, 4).value, 20)

    def test_spaced_sheet_merge_preserves_distinct_launch_phases(self):
        books = []
        for index in range(2):
            book = Workbook()
            book.remove(book.active)
            for phase in (1, 2):
                sheet = book.create_sheet((MODEL if index == 0 else MODEL.replace(" ", "")) + f"by天第{phase}期")
                sheet.append(["周期", "D1"])
                sheet.append(["时间", datetime(2026, phase, 1)])
                sheet.append(["当日大定数量", phase * 10])
            books.append(book)
        store = self.store(books, "首销期订单节奏")
        self.assertEqual(len(store.find_subject_sheets("首销期订单节奏", MODEL, "day")), 2)

    def chart_book(self, total, share):
        book = Workbook()
        sheet = book.active
        sheet.title = "问界净大定by天图表"
        sheet.append([MODEL])
        sheet.append([None, "留存大定", "2026-01-01"])
        sheet.append([None, "版本A", share])
        sheet.append([None, "总计", total])
        return book

    def test_chart_blank_falls_back_without_conflict(self):
        store = self.store([self.chart_book(None, None), self.chart_book(20, 1)])
        with patch("core.excel.LOGGER.warning") as warning:
            chart = store.find_chart_data("大定选配比例", MODEL, "day")[2]
        warning.assert_not_called()
        self.assertEqual(chart["totals"], [20])
        self.assertEqual(chart["series"][0]["values"], [1])
        self.assertEqual(store.find_chart_blocks("大定选配比例", MODEL, "day")[2][0]["totals"], [20])

    def test_chart_real_zero_is_not_overwritten_and_conflict_is_reported(self):
        store = self.store([self.chart_book(0, 0), self.chart_book(20, 1)])
        with self.assertLogs("core.excel", level="WARNING") as logs:
            chart = store.find_chart_data("大定选配比例", MODEL, "day")[2]
        self.assertEqual(chart["totals"], [0])
        self.assertEqual(chart["series"][0]["values"], [0])
        self.assertEqual(len(logs.output), 1)

    def test_chart_all_blank_keeps_existing_display_contract(self):
        for books in ([self.chart_book(None, None)], [self.chart_book(None, None), self.chart_book(None, None)]):
            with self.subTest(file_count=len(books)):
                chart = self.store(books).find_chart_data("大定选配比例", MODEL, "day")[2]
                self.assertEqual(chart["totals"], [0])
                self.assertEqual(chart["series"][0]["values"], [0])

    def test_hourly_record_keys_follow_reordered_headers(self):
        books = []
        for index, days in enumerate(((1, 2), (2, 3))):
            book = Workbook()
            sheet = book.active
            sheet.title = MODEL + "_小订分时退订"
            sheet.append(["小订分时退订"])
            headers = ["总退订数", "小订数", "小订日期", "小订小时"]
            if index:
                headers.reverse()
                sheet.append(["说明"])
            sheet.append(headers)
            for day in days:
                values = {"总退订数": 2, "小订数": 20, "小订日期": datetime(2026, 1, day), "小订小时": 0}
                sheet.append([values[label] for label in headers])
            books.append(book)
        store = self.store(books, "小订退订分析")
        with patch("modules.sales_forecast._read_model_mapping", return_value={}), patch("core.excel.LOGGER.warning") as warning:
            profiles, _ = _read_actual_profiles(store)
        warning.assert_not_called()
        days = profiles[0]["small_hourly_days"]
        self.assertEqual([day["date"] for day in days], ["2026-01-01", "2026-01-02", "2026-01-03"])
        self.assertEqual([day["orders"] for day in days], [20, 20, 20])

    def test_daily_record_keys_follow_date_column(self):
        books = []
        for day in (1, 2):
            book = Workbook()
            sheet = book.active
            sheet.title = MODEL + "_日度退订"
            sheet.append(["日度退订"])
            sheet.append(["整体 - 大定退 %", "累计小订退", "取消日期"])
            sheet.append([.03, 10, datetime(2026, 1, day)])
            books.append(book)
        store = self.store(books, "小订退订分析")
        sheet = store.find("小订退订").workbook.worksheets[0]
        self.assertEqual([sheet.cell(row, 3).value for row in (3, 4)],
                         [datetime(2026, 1, 1), datetime(2026, 1, 2)])


class WorkbookCacheTests(unittest.TestCase):
    def test_cache_keeps_typed_values_formats_and_only_reloads_changed_files(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            path, cache = root / "source.xlsx", root / "cache"
            book = Workbook()
            sheet = book.active
            sheet["A1"], sheet["B1"], sheet["C1"] = 0, .125, datetime(2026, 1, 1)
            sheet["B1"].number_format = "0.00%"
            sheet["C1"].number_format = "yyyy/mm/dd"
            sheet["D1"] = "=literal"
            sheet["D1"].data_type = "s"
            sheet["Z500"].font = Font(bold=True)
            book.save(path)
            book.close()
            first = load_data_workbook(path, cache)
            with patch("core.excel.load_workbook", side_effect=AssertionError("cache missed")):
                second = load_data_workbook(path, cache)
            for cell in ("A1", "B1", "C1", "D1"):
                self.assertEqual(first.active[cell].value, second.active[cell].value)
                self.assertEqual(first.active[cell].number_format, second.active[cell].number_format)
                self.assertEqual(first.active[cell].data_type, second.active[cell].data_type)
            self.assertEqual(first.active.max_row, 1)
            self.assertEqual(second.active["D1"].value, "=literal")
            book = load_workbook(path)
            book.active["A1"] = 9
            book.save(path)
            book.close()
            changed = load_data_workbook(path, cache)
            self.assertEqual(changed.active["A1"].value, 9)
            for item in (first, second, changed):
                item.close()

    def test_cache_keeps_merged_cells_blank_even_if_xml_contains_hidden_values(self):
        from openpyxl.cell.cell import Cell
        with TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "merged.xlsx"
            book = Workbook()
            sheet = book.active
            sheet["A1"] = "时间"
            sheet["B1"] = datetime(2026, 8, 11, 23)
            sheet["B1"].number_format = "yyyy-mm-dd hh:mm"
            sheet.merge_cells("B1:C1")
            sheet["B2"] = 0
            sheet.merge_cells("B2:C3")
            # Reproduce stale XML values inside horizontal and rectangular merges.
            sheet._cells[(1, 3)] = Cell(sheet, row=1, column=3, value=datetime(2026, 8, 12))
            sheet._cells[(3, 3)] = Cell(sheet, row=3, column=3, value=999)
            book.save(path)
            book.close()
            streaming = load_workbook(path, data_only=True, read_only=True)
            self.assertEqual(streaming.active["C3"].value, 999)
            streaming.close()
            original = load_workbook(path, data_only=True, read_only=False)
            self.addCleanup(original.close)
            cold = load_data_workbook(path, root / "cache")
            with patch("core.excel.load_workbook", side_effect=AssertionError("cache missed")):
                warm = load_data_workbook(path, root / "cache")
            for name, actual in (("cold", cold), ("warm", warm)):
                self.addCleanup(actual.close)
                with self.subTest(reader=name):
                    for row in original.active.iter_rows():
                        for expected in row:
                            observed = actual.active[expected.coordinate]
                            self.assertEqual(observed.value, expected.value, expected.coordinate)
                            if expected.value is not None:
                                self.assertEqual(observed.number_format, expected.number_format)
                                self.assertEqual(observed.data_type, expected.data_type)
                    self.assertEqual(actual.active["B2"].value, 0)
                    self.assertIsNone(actual.active["C1"].value)
                    self.assertIsNone(actual.active["C3"].value)

    def test_corrupt_cache_falls_back_to_source(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            path, cache = root / "source.xlsx", root / "cache"
            book = Workbook()
            book.active["A1"] = 123
            book.save(path)
            book.close()
            load_data_workbook(path, cache).close()
            next(cache.glob("*.jsonl.gz")).write_bytes(b"broken cache")
            reloaded = load_data_workbook(path, cache)
            self.assertEqual(reloaded.active["A1"].value, 123)
            reloaded.close()
            self.assertFalse(list(cache.glob("*.tmp")))


if __name__ == "__main__":
    unittest.main()
