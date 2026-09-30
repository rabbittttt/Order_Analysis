from datetime import date, datetime
from pathlib import Path
from unittest import TestCase
from unittest.mock import patch
from types import SimpleNamespace

from openpyxl import Workbook

from core.excel import WorkbookItem, _merge_matrix
from core.forecast_summary import (DAILY_SHEET, MASTER_SHEET, SMALL_HOURLY_SHEET,
    LAUNCH_HOURLY_SHEET, table_records, resolved_weekly_rows)
from tests import test_public_forecast_tables


class ResolvedSummaryTests(TestCase):
    def make(self, customize=None):
        return test_public_forecast_tables.PublicForecastTablesTests().make_public(customize=customize)

    def test_launch_partial_fields_survive_blank_gross(self):
        from modules.sales_forecast import _read_actual_profiles
        book = Workbook()
        sheet = book.active
        sheet.title = '测试车2026款by天'
        sheet.append(['周期', 'D1', 'D2'])
        sheet.append(['时间', datetime(2026, 1, 3), datetime(2026, 1, 4)])
        sheet.append(['当日大定数量', None, None])
        sheet.append(['当日小订转大定数量', 5, None])
        sheet.append(['当日直接大定数量', None, 9])
        item = WorkbookItem(Path('首销期订单节奏.xlsx'), book)
        store = SimpleNamespace(find=lambda keyword: item if keyword == '首销期订单节奏' else None)
        with patch('modules.sales_forecast._read_model_mapping', return_value={}):
            profiles, _ = _read_actual_profiles(store)
        self.assertEqual([r['gross'] for r in profiles[0]['days']], [None, None])
        self.assertEqual([r['small_to_big'] for r in profiles[0]['days']], [5, None])
        self.assertEqual([r['direct'] for r in profiles[0]['days']], [None, 9])
        book.close()

    def test_temporal_keys_are_unique_and_intraday_is_not_a_daily_row(self):
        def customize(data):
            raw = data['actuals'][0]
            raw['small_daily_days'].append({'date': '2026-01-03', 'orders': 5})
            raw['hourly_days'] = [{'date': '2026-01-03', 'last_hour': 15,
                'small_to_big': 2, 'direct': 3, 'hours': [{'hour': 15, 'gross': 5}, {'hour': 15, 'gross': 5}]}]
        book = self.make(customize)
        for sheet, keys in ((DAILY_SHEET, ('订单分析代际名', '日期')),
                            (SMALL_HOURLY_SHEET, ('订单分析代际名', '日期', '小时')),
                            (LAUNCH_HOURLY_SHEET, ('订单分析代际名', '日期', '小时'))):
            ids = [tuple(r[k] for k in keys) for r in table_records(book, sheet)]
            self.assertEqual(len(ids), len(set(ids)), sheet)
        day = next(r for r in table_records(book, DAILY_SHEET) if r['日期'] == datetime(2026, 1, 3))
        self.assertEqual((day['小订数量'], day['大定']), (5, 30))
        self.assertEqual(day['订单阶段'], '小订/首销')
        self.assertTrue(all(r['日期'].hour == 0 for r in table_records(book, DAILY_SHEET)))
        book.close()

    def test_partial_quantities_and_d12_use_resolved_cells_not_old_reference(self):
        def customize(data):
            data['reference_daily'] = {
                '小转大当日数量': {'历史车': [1, 8]},
                '直接大定当日数量': {'历史车': [19, 32]},
                '总大定当日数量': {'历史车': [20, 40]},
                '退订当日数量': {'历史车': [2, None]},
            }
        book = self.make(customize)
        metric = table_records(book, 'D1_D2预测指标')[0]
        self.assertEqual((metric['D1小转大'], metric['D2小转大'], metric['D1+D2小转大']), (10, 8, 18))
        self.assertEqual((metric['D1大定'], metric['D2大定']), (30, 40))
        self.assertIsNone(metric['D1+D2退订'])
        self.assertIsNone(metric['总大定'])  # The launch period has not ended.
        self.assertIsNone(metric['D1小转大/总小转大'])
        self.assertEqual(table_records(book, '小转大当日数量')[0]['D1'], 10)
        self.assertIsNone(table_records(book, '小转大累计完成度')[0]['D1'])
        self.assertAlmostEqual(table_records(book, '累计小订转化率')[0]['D1'], .1)
        self.assertIsNone(table_records(book, MASTER_SHEET)[0]['总小转大'])
        book.close()

    def test_cross_stage_week_has_one_row_and_keeps_stage_portions(self):
        rows = resolved_weekly_rows([
            ['车', '26WK01', '首销', date(2025, 12, 29), date(2026, 1, 1), 10, 8, 3, 'A', 'A', 'A'],
            ['车', '26WK01', '平销', date(2026, 1, 2), date(2026, 1, 4), 0, 0, 0, 'B', 'B', 'B'],
        ])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][2], '首销/平销')
        self.assertEqual(rows[0][5:8], [10, 8, 3])
        self.assertEqual(rows[0][11:], [10, 8, 3, 0, 0, 0])

    def test_cancel_source_refreshes_daily_and_cumulative_reference(self):
        def customize(data):
            raw = data['actuals'][0]
            raw['days'][0].update(cancel=20, _field_sources={'cancel': 'cancel'})
            raw['cancel_days'] = [{'date': '2026-01-02', 'cancel': 16}, {'date': '2026-01-03', 'cancel': 20}]
            raw['cancel_source'] = {'file': '小订退订分析.xlsx', 'sheet': '测试车日度退订'}
            data['reference_daily'] = {'退订当日数量': {'历史车': [1]}}
        book = self.make(customize)
        metric = table_records(book, 'D1_D2预测指标')[0]
        self.assertEqual(metric['D1退订'], 4)
        self.assertAlmostEqual(metric['D1退订率'], .04)
        # Cumulative includes the 16 cancellations already known before launch.
        self.assertAlmostEqual(table_records(book, '累计退订率')[0]['D1'], .2)
        book.close()

    def test_explicit_cancel_daily_wins_first_day_and_zero_falls_back_per_field(self):
        def customize(data):
            raw = data['actuals'][0]
            raw['cancel_days'] = [
                {'date': '2026-01-03', 'cancel': 38, 'daily_cancel': 38},
                {'date': '2026-01-04', 'cancel': 38, 'daily_cancel': 0},
                {'date': '2026-01-05', 'cancel': None, 'daily_cancel': None},
            ]
            raw['cancel_source'] = {'file': '小订退订分析.xlsx', 'sheet': '测试车日度退订'}
            data['reference_daily'] = {'退订当日数量': {'历史车': [32, 5, 7]}}
        book = test_public_forecast_tables.PublicForecastTablesTests().make_public(customize=customize, as_of=date(2026, 1, 6))
        days = {r['日期']: r for r in table_records(book, DAILY_SHEET)}
        self.assertEqual(days[datetime(2026, 1, 3)]['退订数量'], 38)
        self.assertEqual(days[datetime(2026, 1, 4)]['退订数量'], 0)
        self.assertEqual(days[datetime(2026, 1, 5)]['退订数量'], 7)
        self.assertIn('小订退订分析.xlsx', days[datetime(2026, 1, 3)]['字段来源'])
        self.assertAlmostEqual(table_records(book, '累计退订率')[0]['D1'], .38)
        self.assertEqual(table_records(book, 'D1_D2预测指标')[0]['D1退订'], 38)
        book.close()

    def test_cancel_missing_prior_day_does_not_invent_daily_quantity(self):
        def customize(data):
            raw = data['actuals'][0]
            raw['cancel_days'] = [{'date': '2026-01-03', 'cancel': 38}]
            raw['cancel_source'] = {'file': '小订退订分析.xlsx', 'sheet': '测试车日度退订'}
            data['reference_daily'] = {'退订当日数量': {'历史车': [32]}}
        book = self.make(customize)
        self.assertEqual(table_records(book, 'D1_D2预测指标')[0]['D1退订'], 32)
        self.assertAlmostEqual(table_records(book, '累计退订率')[0]['D1'], .38)
        book.close()

    def test_cancel_cumulative_priority_is_independent_of_launch_stage(self):
        from modules.sales_forecast import _resolve_stage_candidate
        candidates = {
            'launch': {'days': [{'date': '2026-01-03', 'gross': 30, 'cancel': 99}]},
            'history': {'days': [{'date': '2026-01-03', 'gross': 20, 'cancel': 32}]},
            'cancel': {'days': [{'date': '2026-01-03', 'cancel': 38}]},
        }
        for stage in ('before', 'active', 'ended', 'unknown'):
            with self.subTest(stage=stage):
                self.assertEqual(_resolve_stage_candidate(candidates, stage)['days'][0]['cancel'], 38)
                candidates['cancel']['days'][0]['cancel'] = 0
                self.assertEqual(_resolve_stage_candidate(candidates, stage)['days'][0]['cancel'], 0)
                candidates['cancel']['days'][0]['cancel'] = None
                self.assertEqual(_resolve_stage_candidate(candidates, stage)['days'][0]['cancel'], 32)
                candidates['cancel']['days'][0]['cancel'] = 38

    def test_raw_cancel_parser_keeps_explicit_daily_quantity(self):
        from modules.sales_forecast import _read_actual_profiles
        book = Workbook()
        sheet = book.active
        sheet.title = '测试车2026款_日度退订'
        sheet.append(['日度退订'])
        sheet.append(['取消日期', '当日小订退', '累计小订退', '整体 - 小订退 %'])
        sheet.append([datetime(2026, 1, 3), 38, 38, .38])
        sheet.append([datetime(2026, 1, 4), 0, 38, .38])
        item = WorkbookItem(Path('小订退订分析.xlsx'), book)
        store = SimpleNamespace(find=lambda keyword: item if keyword == '小订退订分析' else None)
        with patch('modules.sales_forecast._read_model_mapping', return_value={}):
            profiles, _ = _read_actual_profiles(store)
        self.assertEqual([r['daily_cancel'] for r in profiles[0]['cancel_days']], [38, 0])
        book.close()

    def test_omitted_day_in_physical_file_is_zero_before_merge(self):
        def part(name, dates, values):
            book = Workbook()
            sheet = book.active
            sheet.title = '车2026款by天'
            sheet.append(['指标', '类型', '类别', *dates])
            sheet.append(['小订', '数量', '数量', *values])
            item = WorkbookItem(Path(name), book)
            item._boundary_date_range = (date(2026, 1, 1), date(2026, 1, 3))
            return item, sheet
        a = part('小订选配比例分析_A.xlsx', [date(2026, 1, 1), date(2026, 1, 3)], [1, None])
        b = part('小订选配比例分析_B.xlsx', [date(2026, 1, 2), date(2026, 1, 3)], [9, 7])
        merged = _merge_matrix([a, b], 1, 3, forward_rows=(1, 2))
        # Omitted Jan 2 is a valid 0 in A and wins; explicit blank Jan 3 falls back.
        self.assertEqual([c.value for c in merged[2][3:]], [1, 0, 7])
        self.assertEqual(merged._merge_conflict_count, 2)
        a[0].workbook.close()
        b[0].workbook.close()
