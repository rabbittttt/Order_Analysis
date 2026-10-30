"""Forecast mix aliases retain physical file/sheet precedence and provenance."""
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from core.excel import WorkbookItem, WorkbookStore
from core.forecast_summary import DAILY_SHEET, WEEKLY_SHEET, table_records
from core.model_identity import model_key
from modules.sales_forecast import _read_steady_history
from tools.refresh_sales_forecast_data import forecast_weekly_orders
from tests import test_forecast_source_backfill_e2e as fixtures


class ForecastMixAliasTests(TestCase):
    model = '问界 M9 2026款'
    alias = 'M9旗舰2026款'
    day = date(2026, 9, 7)
    today = date(2026, 9, 15)

    def setUp(self):
        for target, value in (
            ('core.excel._CONFLICT_CHECKS', set()),
            ('modules.sales_forecast._read_model_mapping', {model_key(self.model): self.model, model_key(self.alias): self.model}),
            ('modules.sales_forecast._read_model_master', {}),
            ('modules.sales_forecast.generation_records', []),
        ):
            mocked = patch(target, return_value=value) if not target.endswith('_CONFLICT_CHECKS') else patch(target, value)
            mocked.start()
            self.addCleanup(mocked.stop)

    def item(self, filename, entries, field='交车锁单'):
        book = Workbook()
        book.remove(book.active)
        self.addCleanup(book.close)
        for name, values in entries:
            for grain, periods in [('天', [self.day, date(2026, 9, 8)]), ('周', ['26WK37', '26WK38'])]:
                sheet = book.create_sheet(name + 'by' + grain)
                sheet.append(['指标', '统计类型', '分类', *periods])
                sheet.append([field, '数量', '数量', *values])
        return WorkbookItem(Path(filename), book)

    def store(self, items):
        store = WorkbookStore(Path('.'))
        store.items = items
        return store

    def collect(self, items):
        data = {'targets': [{'name': self.model, 'launch_date': '2026-08-01', 'end_date': '2026-08-07'}]}
        dashboard = SimpleNamespace(views={'week': {'pages': {'预测方案': {'workspace': {'data': data}}}}})
        forecast_weekly_orders(self.store(items), dashboard, self.today, collect_only=True)
        return data

    def assert_counts(self, data, expected, field='lock'):
        for table, periods in [('mix_daily', ['2026-09-07', '2026-09-08']), ('mix_weekly', ['26WK37', '26WK38'])]:
            self.assertEqual([data[table][(self.model, p, field)][0] for p in periods], expected, table)

    def test_first_file_alias_precedes_later_standard_name_for_day_and_week(self):
        first = self.item('锁单选配比例分析_A.xlsx', [(self.model, [None, None]), (self.alias, [10, 20])])
        second = self.item('锁单选配比例分析_B.xlsx', [(self.model, [99, 88])])
        data = self.collect([first, second])
        self.assert_counts(data, [10, 20])
        for table in ('mix_daily', 'mix_weekly'):
            for _, source in data[table].values():
                self.assertIn(first.path.name, source)
                self.assertIn(self.alias, source)
                self.assertNotIn(second.path.name, source)

    def test_zero_is_authoritative_and_missing_field_follows_actual_provider(self):
        first = self.item('锁单选配比例分析_Z.xlsx', [(self.alias, [0, None]), (self.model, [5, None])])
        second = self.item('锁单选配比例分析_A.xlsx', [(self.model, [99, 30])])
        data = self.collect([first, second])
        self.assert_counts(data, [0, 30])
        self.assertIn(first.path.name, data['mix_daily'][(self.model, '2026-09-07', 'lock')][1])
        self.assertIn(second.path.name, data['mix_daily'][(self.model, '2026-09-08', 'lock')][1])

    def test_single_file_sheet_order_and_distinct_models_are_preserved(self):
        other = '问界 M8 2026款'
        item = self.item('锁单选配比例分析.xlsx', [(self.model, [0, 20]), (self.alias, [9, 88]), (other, [7, 8])])
        data = self.collect([item])
        self.assert_counts(data, [0, 20])
        self.assertEqual(data['mix_daily'][(other, '2026-09-07', 'lock')][0], 7)

    def test_field_specific_workbook_priority_is_unchanged(self):
        order = self.item('大定选配比例分析.xlsx', [(self.model, [90, 80])])
        locks = self.item('锁单选配比例分析.xlsx', [(self.model, [10, 20])])
        self.assert_counts(self.collect([order, locks]), [10, 20])

    def test_gross_and_net_aliases_use_the_same_precedence(self):
        for label, field in [('大定', 'gross'), ('留存大定', 'net')]:
            with self.subTest(field=field):
                first = self.item('大定选配比例分析_A.xlsx', [(self.model, [None, None]), (self.alias, [10, 20])], label)
                second = self.item('大定选配比例分析_B.xlsx', [(self.model, [99, 88])], label)
                self.assert_counts(self.collect([first, second]), [10, 20], field)

    def test_steady_history_has_one_reference_and_one_row_per_day(self):
        first = self.item('锁单选配比例分析_A.xlsx', [(self.model, [None, None]), (self.alias, [10, 20])])
        second = self.item('锁单选配比例分析_B.xlsx', [(self.model, [99, 88])])
        history, _ = _read_steady_history(self.store([first, second]), {model_key(self.model): {'generation': self.model, 'launch_date': '2026-08-01', 'end_date': '2026-08-07'}}, self.today)
        self.assertEqual(len(history), 1)
        self.assertEqual([r['lock'] for r in history[0]['daily']], [10, 20])
        self.assertEqual([r['lock'] for r in history[0]['weeks']], [10])
        for source in history[0]['daily_sources'].values():
            self.assertIn(first.path.name, source)
            self.assertNotIn(second.path.name, source)


class ForecastMixAliasEndToEndTests(TestCase):
    def test_saved_summary_and_web_steady_values_keep_first_file_alias(self):
        fixture = fixtures.SourceBackfillEndToEndTests()
        self.addCleanup(fixture.doCleanups)
        alias = 'G9旗舰2027款'

        def source_edit(path):
            book = load_workbook(path)
            sheet = book['车型汇总']
            column = sheet.max_column + 1
            sheet.cell(1, column, '原始表简称/别名')
            sheet.cell(2, column, alias)
            book.save(path)
            book.close()

        def add_sources(folder):
            for suffix, entries in [('A', [(fixture.model, None), (alias, 10)]), ('B', [(fixture.model, 99)])]:
                book = Workbook()
                book.remove(book.active)
                for name, value in entries:
                    for grain, period in [('天', date(2026, 9, 7)), ('周', '26WK37')]:
                        sheet = book.create_sheet(name + 'by' + grain)
                        sheet.append(['指标', '统计类型', '分类', period])
                        sheet.append(['交车锁单', '数量', '数量', value])
                book.save(folder / ('鸿蒙智行锁单选配比例分析_' + suffix + '.xlsx'))
                book.close()

        book = fixture.build(source_edit=source_edit, add_sources=add_sources, as_of='2026-09-15')
        day = next(row for row in table_records(book, DAILY_SHEET) if str(row['日期'])[:10] == '2026-09-07')
        week = next(row for row in table_records(book, WEEKLY_SHEET) if row['周期'] == '26WK37')
        self.assertEqual(day['交车锁单'], 10)
        self.assertEqual(week['交车锁单'], 10)
        self.assertIn('_A.xlsx', day['字段来源'])
        self.assertNotIn('_B.xlsx', day['字段来源'])
        _, _, _, _, steady = fixture.web_inputs(book, date(2026, 9, 15))
        reference = next(row for row in steady if row['model'] == fixture.model)
        self.assertEqual(next(row['lock'] for row in reference['weeks'] if row['period'] == '26WK37'), 10)
