"""Keep erroneous rates and real zero amounts from disabling valid dashboards."""
from __future__ import annotations

from pathlib import Path
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook

from core.excel import WorkbookItem, WorkbookStore, cell_number
from core.models import Subject
from modules.launch_rhythm import LaunchRhythmModule
from modules.option_fee import OptionFeeModule, _option_layout, _period_groups


MODEL = "问界 M9 2026款"


class LaunchUnknownRateTests(TestCase):
    def build(self, rate):
        book = Workbook()
        self.addCleanup(book.close)
        sheet = book.active
        sheet.title = MODEL + "by天"
        sheet.append(["日期", "D1", "D2"])
        sheet.append(["日期", "2026-10-01", "2026-10-02"])
        sheet.append(["当日大定数量", 100, 200])
        sheet.append(["当日净大定数量", 80, 160])
        sheet.append(["累计小订转化率", .1, rate])
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("首销期订单节奏.xlsx"), book)]
        with patch("core.model_identity.generation_records", return_value=[]):
            dashboard = LaunchRhythmModule().build(store, Subject("car", MODEL, "generation", "问界"))
        view = dashboard.views["day"]
        return view["pages"][view["default_period"]]

    def test_error_or_placeholder_rate_keeps_orders_and_unknown_cell(self):
        for rate in ("#DIV/0!", "#N/A", "—", "--", "待更新", None, ""):
            with self.subTest(rate=rate):
                page = self.build(rate)
                metrics = page["sections"][0]["data"]["metrics"]
                self.assertEqual(metrics["当日大定数量"], [100, 200])
                self.assertEqual(metrics["当日留存大定数量"], [80, 160])
                self.assertEqual(metrics["累计小订转化率"], [.1, None])
                self.assertEqual([item["value"] for item in page["kpis"][:2]], [300, 240])

    def test_real_zero_and_existing_percent_encodings_are_preserved(self):
        for rate, expected in ((0, 0), (.25, .25), (25, .25), ("25%", .25), ("25％", .25), ("25", .25)):
            with self.subTest(rate=rate):
                page = self.build(rate)
                self.assertEqual(page["sections"][0]["data"]["metrics"]["累计小订转化率"], [.1, expected])

    def test_nullable_rate_does_not_change_default_zero_or_numeric_formats(self):
        book = Workbook()
        self.addCleanup(book.close)
        cell = book.active.cell(1, 1, "#DIV/0!")
        self.assertIsNone(cell_number(cell, rate=True, default=None))
        self.assertEqual(cell_number(cell, rate=True), 0)
        cell.value = 2
        cell.number_format = "0%"
        self.assertEqual(cell_number(cell, rate=True, default=None), 2)
        cell.number_format = "General"
        self.assertEqual(cell_number(cell, rate=True, default=None), .02)


class OptionFeeZeroLayoutTests(TestCase):
    def book(self, *, instructions=0, free_header="免费", second_version=False, quantity_label="订单数量"):
        book = Workbook()
        self.addCleanup(book.close)
        sheet = book.active
        sheet.title = MODEL + "_选配渗透率"
        sheet.append(["权益金额", "6000元"])
        sheet.append(["权益内容", "测试权益"])
        for _ in range(instructions):
            sheet.append(["新增说明"])
        period_row = 3 + instructions
        sheet.cell(period_row, 3, "26-10")
        for col, value in enumerate((free_header, "0-5000", "总计"), 3):
            sheet.cell(period_row + 1, col, value)
        sheet.cell(period_row + 2, 2, quantity_label)
        sheet.cell(period_row + 2, 3, 0)
        sheet.cell(period_row + 2, 4, 100)
        for row, values in enumerate((
            ["Max", "选配金额", 0, 5000, 5000],
            [None, "实际金额", 0, 0, 0],
        ), period_row + 3):
            for col, value in enumerate(values, 1):
                sheet.cell(row, col, value)
        if second_version:
            for values in (
                ["Ultra", "选配金额", 0, 8000, 8000],
                [None, "实际金额", 0, 4000, 4000],
            ):
                sheet.append(values)
        return book

    def page(self, book):
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("选配金统计.xlsx"), book)]
        # The independent paid-sheet absence warning is unrelated to layout.
        with patch("modules.option_fee.LOGGER.warning"):
            dashboard = OptionFeeModule().build(store, Subject("car", MODEL, "generation", "问界"))
        self.assertIsNotNone(dashboard)
        return dashboard.views["month"]["pages"]["26-10"]

    def test_zero_amounts_do_not_replace_header_or_disable_dashboard(self):
        book = self.book()
        page = self.page(book)
        self.assertEqual(_option_layout(book.active), (3, 4, 5))
        rows = next(section["data"]["rows"] for section in page["sections"] if section["title"] == "版本选配金额与实际金额")
        self.assertEqual(rows, [["Max", 5000, 0, 1]])
        series = next(section["data"]["series"] for section in page["sections"] if section["kind"] == "fee_period_bands")
        self.assertEqual(series, [{"name": "免费", "values": [0]}, {"name": "0-5000", "values": [100]}])

    def test_later_positive_version_amounts_are_not_mistaken_for_order_counts(self):
        book = self.book(second_version=True)
        page = self.page(book)
        self.assertEqual(_option_layout(book.active), (3, 4, 5))
        rows = next(section["data"]["rows"] for section in page["sections"] if section["title"] == "版本选配金额与实际金额")
        self.assertEqual(rows, [["Max", 5000, 0, 1], ["Ultra", 8000, 4000, .5]])

    def test_instruction_rows_and_numeric_zero_free_header_remain_supported(self):
        for free_header in (0, "0", "免费"):
            with self.subTest(free_header=free_header):
                book = self.book(instructions=2, free_header=free_header, second_version=True)
                self.page(book)
                self.assertEqual(_option_layout(book.active), (5, 6, 7))
                periods, groups = _period_groups(book.active)
                self.assertEqual(periods, ["26-10"])
                self.assertEqual(groups["26-10"], [(3, "免费"), (4, "0-5000")])

    def test_summary_quantity_label_and_unlabelled_legacy_orders_still_work(self):
        for label in ("总计", "合计", ""):
            with self.subTest(quantity_label=label):
                book = self.book(quantity_label=label, second_version=True)
                self.assertEqual(_option_layout(book.active), (3, 4, 5))

    def test_real_zero_order_counts_remain_valid(self):
        book = self.book()
        book.active.cell(5, 4, 0)
        self.assertEqual(_option_layout(book.active), (3, 4, 5))
        self.page(book)

    def test_order_row_keeps_original_six_row_scan_window(self):
        book = self.book()
        book.active.insert_rows(5, 5)
        self.assertEqual(_option_layout(book.active), (3, 4, 10))
        self.page(book)

    def test_missing_order_row_is_still_rejected(self):
        book = self.book()
        for cell in book.active[5]:
            cell.value = None
        with self.assertRaisesRegex(ValueError, "订单量行"):
            _option_layout(book.active)
