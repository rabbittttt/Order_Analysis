"""Launch fallback shapes keep calendar gaps and actual-source priority."""
from datetime import date
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook
from core.forecast_summary import (MASTER_SHEET, DAILY_SHEET, GROSS_DAILY_SHEET, INDEX_SHEET,
                                   public_forecast_tables, read_public_forecast,
                                   table_records, iso_day)
from core.models import SourceRef
from modules import sales_forecast as forecast


class ActualLaunchShapeCalendarTests(TestCase):
    def history(self, **changes):
        return {"model": "test", "generation": "test",
                "launch_date": "2026-09-01", "end_date": "2026-09-03",
                "days": 3, "gross": 110, "small": 0,
                "daily_orders": [], "daily_small": [], "daily_direct": [],
                **changes}

    def profile(self, entries):
        return {"model": "test", "days": [
            {"date": day, "gross": value} for day, value in entries],
            "hourly_days": [],
            "day_source": {"file": "launch.xlsx", "sheet": "testbyday"}}

    def attach(self, entries, **history_changes):
        item = self.history(**history_changes)
        profile = self.profile(entries)
        with patch.object(forecast, "_read_model_mapping", return_value={}):
            forecast._attach_actual_shapes([item], [profile], today=date(2026, 9, 4))
        return item, profile

    def test_d2_start_keeps_unknown_d1(self):
        item, _ = self.attach([("2026-09-02", 50), ("2026-09-03", 60)])
        self.assertEqual(item["daily_orders"], [None, 50, 60])

    def test_middle_gap_keeps_its_calendar_position(self):
        item, _ = self.attach([("2026-09-01", 10), ("2026-09-03", 30)])
        self.assertEqual(item["daily_orders"], [10, None, 30])

    def test_unsorted_observations_follow_calendar_dates(self):
        item, _ = self.attach([("2026-09-03", 30), ("2026-09-01", 10),
                               ("2026-09-02", 20)])
        self.assertEqual(item["daily_orders"], [10, 20, 30])

    def test_real_zero_is_not_a_gap(self):
        item, _ = self.attach([("2026-09-01", 0), ("2026-09-02", 20),
                               ("2026-09-03", 30)])
        self.assertEqual(item["daily_orders"], [0, 20, 30])

    def test_invalid_dates_and_window_exterior_are_excluded(self):
        item, _ = self.attach([("2026-08-31", 900), ("invalid", 800),
                               ("", 700), ("2026-09-01", 10),
                               ("2026-09-03", 30), ("2026-09-04", 600)])
        self.assertEqual(item["daily_orders"], [10, None, 30])

    def test_no_launch_anchor_cannot_invent_d1(self):
        item, _ = self.attach([("2026-09-02", 50)], launch_date="")
        self.assertEqual(item["daily_orders"], [])

    def test_fallback_is_not_promoted_above_actual_source(self):
        item, profile = self.attach([("2026-09-02", 50), ("2026-09-03", 60)])
        history_profile = forecast._profile_from_history(
            "test", item, SourceRef("history.xlsx", "base", "history"),
            "2026-09-01")
        merged, _ = forecast._merge_day_fields(
            {"history": history_profile, "launch": profile}, ("history", "launch"))
        self.assertEqual([(r["date"], r["gross"], r["_field_sources"]["gross"])
                          for r in merged],
                         [("2026-09-02", 50, "launch"),
                          ("2026-09-03", 60, "launch")])

    def test_original_history_fields_keep_their_priority(self):
        item, profile = self.attach([("2026-09-01", 10), ("2026-09-02", 20),
                                    ("2026-09-03", 30)],
                                   daily_orders=[90, 80, 70],
                                   daily_small=[9, 8, 7], daily_direct=[81, 72, 63])
        self.assertEqual(item["daily_orders"], [90, 80, 70])
        history_profile = forecast._profile_from_history(
            "test", item, SourceRef("history.xlsx", "base", "history"),
            "2026-09-01")
        merged, _ = forecast._merge_day_fields(
            {"history": history_profile, "launch": profile}, ("history", "launch"))
        self.assertEqual([r["gross"] for r in merged], [90, 80, 70])
        self.assertTrue(all(r["_field_sources"]["gross"] == "history" for r in merged))

    def test_original_component_curves_survive_gross_fallback(self):
        item, profile = self.attach([("2026-09-02", 50), ("2026-09-03", 60)],
                                   daily_small=[1, 2, 3], daily_direct=[4, 5, 6])
        history_profile = forecast._profile_from_history(
            "test", item, SourceRef("history.xlsx", "base", "history"),
            "2026-09-01")
        self.assertEqual([r["small_to_big"] for r in history_profile["days"]], [1, 2, 3])
        self.assertEqual([r["direct"] for r in history_profile["days"]], [4, 5, 6])
        self.assertEqual([r["gross"] for r in history_profile["days"]], [None, None, None])

    def public_roundtrip(self, entries, **history_changes):
        item, raw = self.attach(entries, **history_changes)
        target = {"name": "test", "history_model": "test", "stage": "ended",
                  "launch_date": "2026-09-01", "end_date": "2026-09-03",
                  "days": 3, "has_small": False}
        source = SourceRef("history.xlsx", "base", "history")
        with patch.object(forecast, "_read_model_mapping", return_value={}):
            actuals = forecast._resolve_actual_profiles(
                [item], [raw], [target], source, today=date(2026, 9, 4))
        data = {"targets": [target], "history": [item], "actuals": actuals,
                "small_order_history": [], "steady_history": []}
        book = Workbook()
        book.active.title = MASTER_SHEET
        book.active.append(["\u5386\u53f2\u4f20\u64ad\u540d", "\u8ba2\u5355\u5206\u6790\u4ee3\u9645\u540d"])
        book.active.append(["test", "test"])
        index = book.create_sheet(INDEX_SHEET)
        index.append(["\u6765\u6e90\u7c7b\u578b", "\u539f\u59cb\u6587\u4ef6"])
        index.append(["\u5386\u53f2", "history.xlsx"])
        def emit(name, headers, rows, percent_headers=None):
            if name in book.sheetnames:
                del book[name]
            sheet = book.create_sheet(name)
            sheet.append(headers)
            for row in rows:
                sheet.append(row)
        public_forecast_tables(book, data, emit, as_of_date=date(2026, 9, 4),
                               raw_profiles=[raw])
        reader_history = [self.history()]
        with patch.object(forecast, "_read_model_mapping", return_value={}), \
             patch.object(forecast, "_read_model_master", return_value={}):
            profiles, _, _, _ = read_public_forecast(
                book, reader_history, today=date(2026, 9, 4))
        return book, reader_history, profiles

    def test_public_roundtrip_keeps_dates_and_actual_provenance(self):
        book, history, profiles = self.public_roundtrip(
            [("2026-09-02", 50), ("2026-09-03", 60)])
        self.addCleanup(book.close)
        rows = table_records(book, DAILY_SHEET)
        quantities = {iso_day(r["\u65e5\u671f"]): r.get("\u5927\u5b9a") for r in rows}
        self.assertEqual(quantities, {"2026-09-02": 50, "2026-09-03": 60})
        self.assertTrue(all("launch.xlsx" in r["\u5b57\u6bb5\u6765\u6e90"] for r in rows))
        self.assertTrue(all("history.xlsx" not in r["\u5b57\u6bb5\u6765\u6e90"] for r in rows))
        self.assertEqual(history[0]["daily_orders"], [None, 50, 60])
        self.assertEqual([(r["date"], r["gross"]) for r in profiles[0]["days"]],
                         [("2026-09-02", 50), ("2026-09-03", 60)])
        self.assertEqual([c.value for c in book[GROSS_DAILY_SHEET][2]][1:],
                         [None, 50, 60])

    def test_public_roundtrip_keeps_original_history_priority(self):
        book, history, _ = self.public_roundtrip(
            [("2026-09-01", 10), ("2026-09-02", 20), ("2026-09-03", 30)],
            daily_orders=[90, 80, 70])
        self.addCleanup(book.close)
        rows = table_records(book, DAILY_SHEET)
        self.assertEqual([r["\u5927\u5b9a"] for r in rows], [90, 80, 70])
        self.assertTrue(all("history.xlsx" in r["\u5b57\u6bb5\u6765\u6e90"] for r in rows))
        self.assertEqual(history[0]["daily_orders"], [90, 80, 70])
