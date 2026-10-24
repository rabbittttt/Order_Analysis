"""Preserve source month amounts and genuine zero counts in fee charts."""
from __future__ import annotations

import json
from pathlib import Path
import re
import shutil
import subprocess
from unittest import TestCase, skipUnless
from unittest.mock import patch

from openpyxl import Workbook

from core.excel import WorkbookItem, WorkbookStore
from core.models import Subject
from modules.option_fee import OptionFeeModule


MODEL = "\u95ee\u754c M9 2026\u6b3e"
ROOT = Path(__file__).resolve().parents[1]


class FeeMonthSourceTests(TestCase):
    def page(self, periods, actual_amounts):
        book = Workbook()
        self.addCleanup(book.close)
        sheet = book.active
        sheet.title = MODEL + "_\u9009\u914d\u6e17\u900f\u7387"
        for index, (period, actual) in enumerate(zip(periods, actual_amounts)):
            start = 3 + index * 3
            sheet.cell(3, start, period)
            for offset, label in enumerate(("\u514d\u8d39", "0-5000", "\u603b\u8ba1")):
                sheet.cell(4, start + offset, label)
            sheet.cell(5, 2, "\u8ba2\u5355\u6570\u91cf")
            for offset, count in enumerate((0, 100, 100)):
                sheet.cell(5, start + offset, count)
            for row, values in enumerate((
                ["Max", "\u9009\u914d\u91d1\u989d"],
                [None, "\u5b9e\u9645\u91d1\u989d"],
                ["\u5408\u8ba1", "\u9009\u914d\u91d1\u989d\u603b\u8ba1"],
                ["\u5408\u8ba1", "\u5b9e\u9645\u91d1\u989d\u603b\u8ba1"],
            ), 6):
                for column, value in enumerate(values, 1):
                    sheet.cell(row, column, value)
            for row in (6, 8):
                sheet.cell(row, start + 2, 10000 + index * 2000)
            for row in (7, 9):
                sheet.cell(row, start + 2, actual)
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("\u9009\u914d\u91d1\u7edf\u8ba1.xlsx"), book)]
        with patch("modules.option_fee.LOGGER.warning"):
            board = OptionFeeModule().build(store, Subject("car", MODEL, "generation", "\u95ee\u754c"))
        self.assertIsNotNone(board)
        return board.views["month"]["pages"][periods[-1]]

    def test_each_month_uses_its_source_amount_and_latest_kpi_stays_latest(self):
        page = self.page(["26-07", "26-08"], [4500, 6500])
        data = next(part["data"] for part in page["sections"] if part["kind"] == "fee_period_bands")
        self.assertEqual(data["periods"], ["26-07", "26-08"])
        self.assertEqual(data["actual_totals"], [4500, 6500])
        self.assertEqual(data["actual_total"], 6500)
        actual_kpi = next(card for card in page["kpis"] if card["label"] == "\u5e73\u5747\u5b9e\u9645\u91d1\u989d")
        self.assertEqual(actual_kpi["value"], 6500)
        self.assertEqual(actual_kpi["note"], "26-08 \u5408\u8ba1\u53e3\u5f84")

    def test_single_month_zero_amount_remains_zero_in_source_and_kpi(self):
        page = self.page(["26-08"], [0])
        data = next(part["data"] for part in page["sections"] if part["kind"] == "fee_period_bands")
        self.assertEqual(data["actual_totals"], [0])
        self.assertEqual(data["actual_total"], 0)
        actual_kpi = next(card for card in page["kpis"] if card["label"] == "\u5e73\u5747\u5b9e\u9645\u91d1\u989d")
        self.assertEqual(actual_kpi["value"], 0)


@skipUnless(shutil.which("node"), "Node.js is required for production fee renderers")
class FeeRendererValueTests(TestCase):
    def render(self, name, data):
        # Execute the production renderer; only generic number/text formatting
        # and the palette are supplied outside its normal browser closure.
        script = r"""
const fs = require('fs');
const source = fs.readFileSync('templates/dashboard.js', 'utf8');
const name = process.argv[1];
const start = source.indexOf('  function ' + name + '(');
const end = source.indexOf('\n  function ', start + 1);
if (start < 0 || end < 0) throw new Error('Production renderer not found: ' + name);
const esc = value => String(value), fmt = value => String(value);
const chartPalette = {yellow: 'yellow', green: 'green', blue: 'blue', violet: 'violet', rose: 'rose'};
eval(source.slice(start, end));
let input = '';
process.stdin.setEncoding('utf8');
process.stdin.on('data', value => input += value);
process.stdin.on('end', () => process.stdout.write(eval(name)(JSON.parse(input))));
"""
        result = subprocess.run(
            [shutil.which("node"), "-e", script, name], cwd=ROOT,
            input=json.dumps(data), capture_output=True, text=True,
            encoding="utf-8", timeout=20,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        return result.stdout

    def monthly_amounts(self, actual_totals, periods=None, legacy=9900):
        data = {"periods": periods or ["26-07", "26-08"], "series": [], "actual_total": legacy}
        if actual_totals is not ...:
            data["actual_totals"] = actual_totals
        html = self.render("renderFeePeriodBands", data)
        return re.findall(r"<small>(.*?)</small>", html)

    def test_historical_months_show_their_own_amounts(self):
        self.assertEqual(self.monthly_amounts([4500, 6500]), [
            "\u5b9e\u9645\u91d1\u989d\u603b\u8ba1 4500\u5143",
            "\u5b9e\u9645\u91d1\u989d\u603b\u8ba1 6500\u5143",
        ])

    def test_unknown_and_omitted_months_do_not_fall_back_to_latest_and_zero_is_known(self):
        self.assertEqual(self.monthly_amounts([None, 0], ["26-06", "26-07", "26-08"]), [
            "\u5b9e\u9645\u91d1\u989d\u603b\u8ba1 \u2014\u5143",
            "\u5b9e\u9645\u91d1\u989d\u603b\u8ba1 0\u5143",
            "\u5b9e\u9645\u91d1\u989d\u603b\u8ba1 \u2014\u5143",
        ])
        self.assertEqual(self.monthly_amounts([]), [
            "\u5b9e\u9645\u91d1\u989d\u603b\u8ba1 \u2014\u5143",
            "\u5b9e\u9645\u91d1\u989d\u603b\u8ba1 \u2014\u5143",
        ])

    def test_invalid_month_amount_does_not_display_nan_or_latest(self):
        self.assertEqual(self.monthly_amounts(["invalid", 6500]), [
            "\u5b9e\u9645\u91d1\u989d\u603b\u8ba1 \u2014\u5143",
            "\u5b9e\u9645\u91d1\u989d\u603b\u8ba1 6500\u5143",
        ])

    def test_old_payload_without_month_array_keeps_legacy_amount(self):
        self.assertEqual(self.monthly_amounts(..., legacy=6500), [
            "\u5b9e\u9645\u91d1\u989d\u603b\u8ba1 6500\u5143",
            "\u5b9e\u9645\u91d1\u989d\u603b\u8ba1 6500\u5143",
        ])

    def test_single_month_zero_is_not_replaced_by_legacy_amount(self):
        self.assertEqual(self.monthly_amounts([0], ["26-08"]), [
            "\u5b9e\u9645\u91d1\u989d\u603b\u8ba1 0\u5143",
        ])

    def test_zero_orders_remain_zero_while_positive_structure_keeps_real_total(self):
        for values, total in (([0, 0], 0), ([3, 7], 10)):
            with self.subTest(values=values):
                html = self.render("renderFeeBandStructure", {
                    "bands": ["free", "paid"], "rows": [{"label": "model", "values": values}],
                })
                self.assertIn(f"<b>{total}\u5355</b>", html)
                self.assertNotIn("NaN", html)
                self.assertNotIn("Infinity", html)
                if total:
                    self.assertIn("30.0%", html)
                    self.assertIn("70.0%", html)
