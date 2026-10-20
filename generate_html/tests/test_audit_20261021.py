"""Curve-only reservation references must not invent days or quantities."""
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook

from core.forecast_summary import read_public_forecast
from modules.sales_forecast import _read_small_order_history
from tests import test_public_forecast_tables


class SmallProgressFallbackTests(TestCase):
    def raw_reference(self, curve, total=100):
        with TemporaryDirectory() as folder:
            path = Path(folder) / 'source.xlsx'
            book = Workbook()
            book.active.title = '车型汇总'
            book.active.append(['代际名', '小订开始日期', '小订结束日期', '小订天数', '总小订'])
            book.active.append(['测试车', '2026-01-01', '2026-01-05', 5, total])
            book.create_sheet('小订by天')
            sheet = book.create_sheet('小订进度')
            sheet.append(['代际名', *[f'D{i+1}' for i in range(len(curve))]])
            sheet.append(['测试车', *curve])
            book.save(path)
            book.close()
            with patch('modules.sales_forecast._read_model_mapping', return_value={}), \
                 patch('modules.sales_forecast._read_model_master', return_value={}):
                return _read_small_order_history(path)[1][0]

    def public_reference(self, curve):
        def customize(data):
            data['small_order_history'][0].update(daily_actual=False, source_sheet='小订进度',
                                                   standard_progress=curve)
            data['actuals'][0]['small_daily_days'] = []
        book = test_public_forecast_tables.PublicForecastTablesTests().make_public(customize=customize)
        try:
            with patch('modules.sales_forecast._read_model_mapping', return_value={}), \
                 patch('modules.sales_forecast._read_model_master', return_value={}):
                return read_public_forecast(book, [], today=date(2026, 2, 1))[2][0]
        finally:
            book.close()

    def test_raw_internal_blank_keeps_its_day(self):
        item = self.raw_reference([.2, None, .6, .8, 1])
        self.assertEqual(item['standard_progress'], [.2, None, .6, .8, 1])
        self.assertFalse(item['daily_actual'])
        self.assertFalse(item['daily_orders'])

    def test_public_internal_blank_does_not_crash_forecast_module(self):
        item = self.public_reference([.2, None, .6, .8, 1])
        self.assertEqual(item['standard_progress'], [.2, None, .6, .8, 1])
        self.assertEqual(item['daily_orders'], [])
        self.assertEqual(item['d1_share'], 0)

    def test_first_day_missing_does_not_borrow_d2(self):
        for reader in (self.raw_reference, self.public_reference):
            with self.subTest(reader=reader.__name__):
                item = reader([None, .4, .6, .8, 1])
                self.assertIsNone(item['standard_progress'][0])
                self.assertEqual(item['d1_share'], 0)
                self.assertEqual(item['daily_orders'], [])

    def test_decreasing_curve_cannot_create_excess_orders(self):
        for reader in (self.raw_reference, self.public_reference):
            with self.subTest(reader=reader.__name__):
                item = reader([.2, .8, .3, .9, 1])
                self.assertEqual(item['daily_orders'], [])

    def test_rounding_respects_known_total(self):
        item = self.raw_reference([.2, .4, .6, .8, 1], total=3)
        self.assertEqual(sum(item['daily_orders']), 3)
        self.assertTrue(all(value >= 0 for value in item['daily_orders']))

    def test_complete_curve_and_real_zero_still_work(self):
        for reader in (self.raw_reference, self.public_reference):
            with self.subTest(reader=reader.__name__):
                item = reader([0, .2, .6, .8, 1])
                self.assertEqual(item['daily_orders'], [0, 20, 40, 20, 20])
                self.assertFalse(item['daily_actual'])

    def test_raw_and_public_readers_agree(self):
        for curve in ([.333, 1], [.2, None, 1], [0, .5, 1]):
            with self.subTest(curve=curve):
                raw, public = self.raw_reference(curve), self.public_reference(curve)
                for field in ('standard_progress', 'daily_orders', 'd1_share'):
                    self.assertEqual(raw[field], public[field])
