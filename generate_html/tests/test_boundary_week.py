from __future__ import annotations

import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from core.excel import WorkbookItem, WorkbookStore, _week_period, display_period, parse_metric_sheet
from modules.conversion import ConversionModule, STAGES
from modules.sales_forecast import _read_steady_history
from core.models import Subject


MODEL = "问界 M9 2026款"


class BoundaryWeekTests(unittest.TestCase):
    def store(self, books, keyword="锁单选配比例分析"):
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path(f"{keyword}_{index}.xlsx"), book) for index, book in enumerate(books)]
        self.addCleanup(store.close)
        return store

    def metric(self, period, quantity, share=None):
        book = Workbook()
        sheet = book.active
        sheet.title = MODEL + "by周"
        sheet.append(["指标", "统计类型", "分类", period])
        sheet.append(["交车锁单", "数量", "数量", quantity])
        if share is not None:
            sheet.append(["版本", "占比", "A", share])
            sheet.append([None, None, "B", 1 - share])
        return book

    def test_boundaries_use_calendar_not_hardcoded_years(self):
        for period, reason in [("26WK53", "跨年"), ("2026WK53", "跨年"), ("26WK27", "跨半年"),
                               ("25WK27", "跨半年"), ("24WK26", ""), ("23WK52", ""), ("26WK28", "")]:
            with self.subTest(period=period):
                self.assertEqual(_week_period(period)[3], reason)
        self.assertIsNone(_week_period("25WK53"))
        self.assertIsNone(_week_period("26-06"))
        self.assertEqual(display_period("2026WK53"), "26WK53")

    def test_boundary_quantities_sum_and_week_aliases_align(self):
        for period, alias, reason in [("26WK53", "2026WK53", "跨年"), ("26WK27", "26wk27", "跨半年")]:
            with self.subTest(period=period):
                books = [self.metric(period, 140), self.metric(alias, 90)]
                store = self.store(books)
                with self.assertLogs("core.excel", level="INFO") as logs:
                    _, sheet = store.find_subject_sheet("锁单选配比例", MODEL, "week")
                parsed = parse_metric_sheet(sheet)
                self.assertEqual(list(parsed), [period])
                self.assertEqual(parsed[period]["metrics"]["交车锁单"], 230)
                self.assertIn("边界周累加", "\n".join(logs.output))
                self.assertIn(reason, "\n".join(logs.output))
                self.assertIn("→ 230", "\n".join(logs.output))
                self.assertEqual(books[0].active["D2"].value, 140)
                self.assertEqual(len(sheet._source_matches), 2)

    def test_sunday_cutoffs_and_regular_weeks_keep_conflict_priority(self):
        for period in ["24WK26", "23WK52", "26WK28"]:
            with self.subTest(period=period):
                store = self.store([self.metric(period, 140), self.metric(period, 90)])
                with self.assertLogs("core.excel", level="WARNING"):
                    _, sheet = store.find_subject_sheet("锁单选配比例", MODEL, "week")
                self.assertEqual(parse_metric_sheet(sheet)[period]["metrics"]["交车锁单"], 140)
                self.assertEqual(sheet._boundary_week_merges, [])

    def test_identical_pieces_stay_deduplicated_after_sum(self):
        store = self.store([self.metric("26WK53", value) for value in [140, 90, 140, 230]])
        _, sheet = store.find_subject_sheet("锁单选配比例", MODEL, "week")
        self.assertEqual(parse_metric_sheet(sheet)["26WK53"]["metrics"]["交车锁单"], 460)
        same = self.store([self.metric("26WK53", 140), self.metric("26WK53", 140)])
        with patch("core.excel.LOGGER.info") as log:
            _, sheet = same.find_subject_sheet("锁单选配比例", MODEL, "week")
        log.assert_not_called()
        self.assertEqual(parse_metric_sheet(sheet)["26WK53"]["metrics"]["交车锁单"], 140)

    def test_boundary_zero_is_added_but_invalid_values_are_not(self):
        store = self.store([self.metric("26WK53", 0), self.metric("26WK53", 90)])
        _, sheet = store.find_subject_sheet("锁单选配比例", MODEL, "week")
        self.assertEqual(parse_metric_sheet(sheet)["26WK53"]["metrics"]["交车锁单"], 90)
        for invalid in [-1, "not a number", float("inf")]:
            with self.subTest(invalid=invalid):
                store = self.store([self.metric("26WK53", 140), self.metric("26WK53", invalid)])
                with patch("core.excel.LOGGER.warning"):
                    _, sheet = store.find_subject_sheet("锁单选配比例", MODEL, "week")
                self.assertEqual(parse_metric_sheet(sheet)["26WK53"]["metrics"]["交车锁单"], 140)

    def test_shares_are_weighted_not_added(self):
        store = self.store([self.metric("26WK53", 140, .4), self.metric("26WK53", 90, .8)])
        with patch("core.excel.LOGGER.warning") as warning:
            _, sheet = store.find_subject_sheet("锁单选配比例", MODEL, "week")
        warning.assert_not_called()
        rows = parse_metric_sheet(sheet)["26WK53"]["structures"]["版本"]
        self.assertAlmostEqual(rows[0]["share"], (140 * .4 + 90 * .8) / 230)
        self.assertAlmostEqual(sum(row["share"] for row in rows), 1)

    def test_shares_recalculate_from_summed_category_counts(self):
        books = [self.metric("26WK53", 140, .4), self.metric("26WK53", 90, .8)]
        for book, counts in zip(books, [(56, 84), (72, 18)]):
            book.active.append(["版本", "数量", "A", counts[0]])
            book.active.append([None, None, "B", counts[1]])
        store = self.store(books)
        _, sheet = store.find_subject_sheet("锁单选配比例", MODEL, "week")
        rows = parse_metric_sheet(sheet)["26WK53"]["structures"]["版本"]
        self.assertEqual([row["count"] for row in rows], [128, 102])
        self.assertAlmostEqual(rows[0]["share"], 128 / 230)

    def test_conversion_adds_counts_not_average_days_and_derives_rates(self):
        books = []
        for quantities, days in [([100, 95, 90, 80, 75, 72, 70], 3.5), ([50, 45, 40, 35, 30, 25, 20], 5.5)]:
            book = Workbook()
            sheet = book.active
            sheet.title = MODEL + "by周"
            sheet.append(["名称", "周", *STAGES, "大定到交车锁单（天）", "交车锁单到交付（天）"])
            sheet.append([None, None, *["数量"] * len(STAGES), None, None])
            sheet.append(["总计", "26WK53", *quantities, days, 18])
            books.append(book)
        store = self.store(books, "订单7级转化")
        with patch("core.excel.LOGGER.warning"):
            dashboard = ConversionModule().build(store, Subject("m9", MODEL, "generation", "问界"))
        page = dashboard.views["week"]["pages"]["26WK53"]
        funnel = next(section["data"] for section in page["sections"] if section["kind"] == "funnel")
        self.assertEqual(funnel[0]["value"], 150)
        self.assertEqual(funnel[-1]["value"], 90)
        self.assertAlmostEqual(funnel[-1]["total_rate"], .6)
        self.assertEqual(store.find("订单7级转化").workbook.worksheets[0].cell(3, 10).value, 3.5)

    def test_chart_total_and_shares_follow_the_same_rule_without_repeat_logs(self):
        books = []
        for period, total, rate in [("26WK53", 140, .4), ("2026WK53", 90, .8), ("26WK53", 140, .4)]:
            book = Workbook()
            sheet = book.active
            sheet.title = "问界净大定by周图表"
            sheet.append([MODEL])
            sheet.append([None, "留存大定", period])
            sheet.append([None, "A", rate])
            sheet.append([None, "B", 1 - rate])
            sheet.append([None, "总计", total])
            books.append(book)
        store = self.store(books, "大定选配比例分析")
        with self.assertLogs("core.excel", level="INFO") as logs:
            data = store.find_chart_data("大定选配比例", MODEL, "week")[2]
            store.find_chart_data("大定选配比例", MODEL, "week")
        self.assertEqual(data["periods"], ["26WK53"])
        self.assertEqual(data["totals"], [230])
        self.assertAlmostEqual(data["series"][0]["values"][0], 128 / 230)
        self.assertEqual(sum("边界周累加" in line for line in logs.output), 1)

    def test_steady_forecast_receives_the_complete_merged_week(self):
        store = self.store([self.metric("26WK53", 140), self.metric("26WK53", 90)])
        with patch("modules.sales_forecast._read_model_mapping", return_value={}), patch(
            "modules.sales_forecast._read_model_master", return_value={}
        ):
            rows, _ = _read_steady_history(store, {"m9": {"generation": MODEL, "end_date": "2026-11-30"}}, date(2027, 1, 11))
        self.assertEqual(rows[0]["weeks"][0]["lock"], 230)
        self.assertEqual(rows[0]["weeks"][0]["start_date"], "2026-12-28")
        self.assertEqual(rows[0]["weeks"][0]["end_date"], "2027-01-03")

    def test_sku_count_blocks_add_without_summing_percentage_blocks(self):
        books = []
        for quantity in [140, 90]:
            book = Workbook()
            sheet = book.active
            sheet.title = MODEL + "_SKU"
            share = .4 if quantity == 140 else .8
            sheet.append(["累计锁单", None, None, None, "总计占比", None, None])
            sheet.append(["SKU", "26WK53", "总计", None, "SKU", "26WK53", "总计"])
            sheet.append(["版本A", quantity, quantity, None, "版本A", share, share])
            books.append(book)
        store = self.store(books, "SKU收敛度报告")
        with self.assertLogs("core.excel", level="INFO") as logs:
            sheet = store.find("SKU").workbook.worksheets[0]
        self.assertEqual(sheet.cell(3, 2).value, 230)
        self.assertEqual(sheet.cell(3, 3).value, 140)  # Annual totals are not additive.
        self.assertEqual(sheet.cell(3, 6).value, .4)  # No percentage summation.
        self.assertTrue(any("边界周累加" in line for line in logs.output))


if __name__ == "__main__":
    unittest.main()
