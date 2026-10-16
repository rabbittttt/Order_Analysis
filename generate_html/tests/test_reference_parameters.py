from unittest import TestCase

from modules.sales_forecast import _history_item
from tools.refresh_sales_forecast_data import reorder_summary_sheets
from core.forecast_summary import D12_SHEET, SMALL_DAILY_SHEET, SMALL_CURVE_SHEET
from openpyxl import Workbook


class ReferenceParameterTests(TestCase):
    def test_history_ratios_preserve_missing_and_explicit_zero(self):
        fields = {'小订转化率':'conversion', '直接大定占比':'direct_share',
                  '大定到锁单率':'lock_rate', '留存大定率':'net_rate', '退订率':'cancel_rate'}
        missing = _history_item({'传播名':'测试车'}, True)
        zero = _history_item({'传播名':'测试车', **dict.fromkeys(fields, 0)}, True)
        for field in fields.values():
            self.assertIsNone(missing[field], field)
            self.assertEqual(zero[field], 0, field)

    def test_final_generator_sort_puts_cumulative_cancellation_rate_last(self):
        book = Workbook()
        book.remove(book.active)
        names = ('未知额外Sheet', '兼容源Sheet', '累计退订率', D12_SHEET, SMALL_CURVE_SHEET, SMALL_DAILY_SHEET)
        for index, name in enumerate(names):
            sheet = book.create_sheet(name)
            sheet['A1'] = f'={index}+1'
            sheet['B1'] = index
            sheet['B1'].number_format = '0.00%'
        before = {
            name: (
                book[name],
                book[name]['A1'].value,
                book[name]['B1'].value,
                book[name]['B1'].number_format,
                tuple(book[name]['B1']._style),
            )
            for name in names
        }
        reorder_summary_sheets(book)
        self.assertEqual(
            book.sheetnames,
            [SMALL_DAILY_SHEET, SMALL_CURVE_SHEET, D12_SHEET, '未知额外Sheet', '兼容源Sheet', '累计退订率'],
        )
        for name, (sheet, formula, value, number_format, style) in before.items():
            self.assertIs(book[name], sheet)
            self.assertEqual(book[name]['A1'].value, formula)
            self.assertEqual(book[name]['B1'].value, value)
            self.assertEqual(book[name]['B1'].number_format, number_format)
            self.assertEqual(tuple(book[name]['B1']._style), style)
        first_order = book.sheetnames
        reorder_summary_sheets(book)
        self.assertEqual(book.sheetnames, first_order)
        book.close()

        book = Workbook()
        book.remove(book.active)
        for name in (D12_SHEET, SMALL_CURVE_SHEET, SMALL_DAILY_SHEET, '其他表'):
            book.create_sheet(name)
        reorder_summary_sheets(book)
        self.assertNotIn('累计退订率', book.sheetnames)
        self.assertEqual(len(book.sheetnames), 4)
        book.close()
