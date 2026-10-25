"""An observed zero amount has the same source priority as any other number."""
from __future__ import annotations

from pathlib import Path
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook

from core.excel import WorkbookItem, WorkbookStore
from core.models import Subject
from modules.option_fee import OptionFeeModule, _period_amount_value


MODEL = "\u95ee\u754c M9 2026\u6b3e"
FREE = "\u514d\u8d39"
TOTAL = "\u603b\u8ba1"


class FeeZeroSourcePriorityTests(TestCase):
    def amount(self, values, labels=None):
        book = Workbook()
        self.addCleanup(book.close)
        sheet = book.active
        sheet.cell(3, 3, "26-08")
        labels = labels or [FREE, "0-5000", TOTAL]
        for column, label in enumerate(labels, 3):
            sheet.cell(4, column, label)
            sheet.cell(5, column, 10)
        sheet.cell(5, 2, "\u8ba2\u5355\u6570\u91cf")
        for column, value in enumerate(values, 3):
            sheet.cell(6, column, value)
        return _period_amount_value(sheet, 6, 3, 2 + len(labels))

    def test_labelled_zero_total_beats_nonzero_amount_bands(self):
        for total_label in (TOTAL, "\u5408\u8ba1", "\u6c47\u603b"):
            for zero in (0, "0", "0\u5143"):
                with self.subTest(total_label=total_label, zero=zero):
                    self.assertEqual(self.amount([1000, 3000, zero], [FREE, "0-5000", total_label]), 0)

    def test_without_total_first_valid_zero_is_not_skipped(self):
        labels = [FREE, "0-5000", "5000-10000"]
        for values in ([0, 3000, 7000], [None, 0, 7000], ["invalid", "0\u5143", 7000]):
            with self.subTest(values=values):
                self.assertEqual(self.amount(values, labels), 0)

    def test_nonzero_total_still_beats_earlier_amount_band_values(self):
        self.assertEqual(self.amount([1000, 3000, 6500]), 6500)
        self.assertEqual(self.amount([0, 3000, "6,500\u5143"]), 6500)

    def test_blank_or_invalid_total_falls_back_to_first_valid_numeric_band(self):
        for total in (None, "", "\u2014", "#DIV/0!", "invalid", float("nan"), float("inf")):
            with self.subTest(total=total):
                self.assertEqual(self.amount([4200, 5000, total]), 4200)
                self.assertEqual(self.amount([0, 5000, total]), 0)
        self.assertEqual(self.amount(["invalid", "3,000\u5143", None]), 3000)

    def test_multiple_labelled_totals_use_first_valid_number_including_zero(self):
        labels = [FREE, "0-5000", TOTAL, "\u5408\u8ba1"]
        self.assertEqual(self.amount([1000, 3000, 0, 9000], labels), 0)
        self.assertEqual(self.amount([1000, 3000, "#DIV/0!", 0], labels), 0)
        self.assertEqual(self.amount([1000, 3000, 6500, 9000], labels), 6500)

    def test_boolean_flags_are_not_amounts_and_do_not_hide_later_numbers(self):
        for flag in (False, True):
            with self.subTest(flag=flag):
                self.assertEqual(self.amount([4200, 5000, flag]), 4200)
                self.assertEqual(self.amount([0, 5000, flag]), 0)
                self.assertEqual(self.amount([flag, 3000, None]), 3000)
                self.assertEqual(self.amount([flag, 0, None]), 0)

    def test_no_valid_amount_keeps_existing_zero_fallback(self):
        self.assertEqual(self.amount([None, "invalid", "#DIV/0!"]), 0)
        self.assertEqual(self.amount([float("nan"), float("inf"), "\u2014"]), 0)

    def page(self, totals):
        book = Workbook()
        self.addCleanup(book.close)
        sheet = book.active
        sheet.title = MODEL + "_\u9009\u914d\u6e17\u900f\u7387"
        periods = ["26-07", "26-08"][-len(totals):]
        for index, (period, (option_total, actual_total)) in enumerate(zip(periods, totals)):
            start = 3 + index * 3
            sheet.cell(3, start, period)
            for offset, label in enumerate((FREE, "0-5000", TOTAL)):
                sheet.cell(4, start + offset, label)
            sheet.cell(5, 2, "\u8ba2\u5355\u6570\u91cf")
            for offset, value in enumerate((10, 0, 10)):
                sheet.cell(5, start + offset, value)
            for row, values in enumerate((
                ["Max", "\u9009\u914d\u91d1\u989d"],
                [None, "\u5b9e\u9645\u91d1\u989d"],
                ["\u5408\u8ba1", "\u9009\u914d\u91d1\u989d\u603b\u8ba1"],
                ["\u5408\u8ba1", "\u5b9e\u9645\u91d1\u989d\u603b\u8ba1"],
            ), 6):
                for column, value in enumerate(values, 1):
                    sheet.cell(row, column, value)
            for row in (6, 8):
                for offset, value in enumerate((0, 5000, option_total)):
                    sheet.cell(row, start + offset, value)
            for row in (7, 9):
                for offset, value in enumerate((0, 3000, actual_total)):
                    sheet.cell(row, start + offset, value)
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("\u9009\u914d\u91d1\u7edf\u8ba1.xlsx"), book)]
        with patch("modules.option_fee.LOGGER.warning"):
            board = OptionFeeModule().build(store, Subject("car", MODEL, "generation", "\u95ee\u754c"))
        self.assertIsNotNone(board)
        return board.views["month"]["pages"][periods[-1]]

    def test_generation_build_preserves_explicit_zero_in_kpis_and_version_table(self):
        page = self.page([(0, 0)])
        kpis = {card["label"]: card["value"] for card in page["kpis"]}
        self.assertEqual(kpis["\u5e73\u5747\u9009\u914d\u91d1\u989d"], 0)
        self.assertEqual(kpis["\u5e73\u5747\u5b9e\u9645\u91d1\u989d"], 0)
        rows = next(part["data"]["rows"] for part in page["sections"] if part["kind"] == "table")
        self.assertEqual(rows, [["Max", 0, 0, 0]])
        data = next(part["data"] for part in page["sections"] if part["kind"] == "fee_period_bands")
        self.assertEqual(data["actual_total"], 0)
        self.assertEqual(data["actual_totals"], [0])

    def test_generation_build_binds_zero_and_positive_amounts_to_their_months(self):
        page = self.page([(0, 0), (12000, 6500)])
        data = next(part["data"] for part in page["sections"] if part["kind"] == "fee_period_bands")
        self.assertEqual(data["periods"], ["26-07", "26-08"])
        self.assertEqual(data["actual_totals"], [0, 6500])
        self.assertEqual(data["actual_total"], 6500)
        kpis = {card["label"]: card["value"] for card in page["kpis"]}
        self.assertEqual(kpis["\u5e73\u5747\u9009\u914d\u91d1\u989d"], 12000)
        self.assertEqual(kpis["\u5e73\u5747\u5b9e\u9645\u91d1\u989d"], 6500)
        rows = next(part["data"]["rows"] for part in page["sections"] if part["kind"] == "table")
        self.assertEqual(rows, [["Max", 12000, 6500, (12000 - 6500) / 12000]])
