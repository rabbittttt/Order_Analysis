"""Invalid maintained totals stay unknown until real quantities can replace them."""
from pathlib import Path
from unittest import TestCase

from openpyxl import Workbook, load_workbook

from core.forecast_summary import D12_SHEET, MASTER_SHEET, table_records
from tests import test_forecast_source_backfill_e2e as fixtures
from tools.refresh_sales_forecast_data import (
    build_secondary, normalize_reference_sheet, optional_reference_quantity,
    optional_reference_rate,
)


class ForecastInvalidQuantitiesTests(TestCase):
    quantity_fields = ('总小订', '总小转大', '总大定', '总直接大定', '总退订',
                       '首销期留存大定', '首销期锁单')
    rate_fields = ('小订转化率', '直接大定占比', '退订率', '留存大定率',
                   '留存大定到锁单率', '大定到锁单率')

    def build(self, maintained=None, *, launch_quantities=None, add_mix=False,
              source_quantities=None):
        fixture = fixtures.SourceBackfillEndToEndTests()
        self.addCleanup(fixture.doCleanups)

        def edit(path):
            book = load_workbook(path)
            sheet = book['车型汇总']
            headers = [cell.value for cell in sheet[1]]
            for field, value in (maintained or {}).items():
                if field not in headers:
                    headers.append(field)
                    sheet.cell(1, len(headers), field)
                sheet.cell(2, headers.index(field)+1, value)
            book.save(path)
            book.close()

        def mix(folder):
            for kind, metrics in (
                ('大定', [('大定', [50, 50, 30]), ('留存大定', [45, 45, 27])]),
                ('锁单', [('交车锁单', [25, 25, 15])]),
            ):
                book = Workbook()
                sheet = book.active
                sheet.title = fixture.model+'by天'
                sheet.append(['指标', '统计类型', '分类', '2026-09-01', '2026-09-02', '2026-09-03'])
                for metric, values in metrics:
                    sheet.append([metric, '数量', '数量', *values])
                book.save(folder/f'鸿蒙智行{kind}选配比例分析.xlsx')
                book.close()

        return fixture.build(maintained=maintained, source_edit=edit,
                             launch_quantities=launch_quantities,
                             source_quantities=source_quantities,
                             add_sources=mix if add_mix else None)

    def row(self, book, name=MASTER_SHEET):
        return table_records(book, name)[0]

    def test_count_parser_preserves_zero_and_supported_numeric_strings(self):
        for value, expected in ((0, 0), ('0', 0), (42, 42), ('42', 42),
                                (' 1,234 ', 1234), ('1，234', 1234), ('+12', 12)):
            with self.subTest(value=value):
                self.assertEqual(optional_reference_quantity(value), expected)

    def test_invalid_count_types_are_unknown_instead_of_zero(self):
        for value in (None, '', ' ', '待更新', '—', '#DIV/0!', '#N/A', True, False,
                      float('nan'), float('inf'), -float('inf'), -1, '-1',
                      'NaN', 'Infinity', '1,23', '5%', '12单'):
            with self.subTest(value=value):
                self.assertIsNone(optional_reference_quantity(value))

    def test_temporary_roster_normalization_supports_old_and_modern_identity_headers(self):
        for identity in ('代际名', '历史传播名'):
            with self.subTest(identity=identity):
                book = Workbook()
                self.addCleanup(book.close)
                sheet = book.active
                sheet.append(['车型资料'])
                sheet.append([identity, '总小转大', '总小订', '小订转化率', '品牌'])
                sheet.append(['测试车', '#DIV/0!', '0', '.5', '待更新'])
                sheet['C3'].number_format = '0.00'
                normalize_reference_sheet(sheet)
                self.assertEqual([cell.value for cell in sheet[3]], ['测试车', None, 0, .5, '待更新'])
                self.assertEqual(sheet['C3'].number_format, '0.00')

    def test_invalid_maintained_total_is_filled_from_real_launch_days(self):
        for value in (None, '待更新', '#DIV/0!', True, False, -1, 'NaN', 'Infinity'):
            with self.subTest(value=value):
                book = self.build({'总小转大': value})
                master = self.row(book)
                self.assertEqual(master['总小转大'], 40)
                self.assertEqual(master['小订转化率'], .4)
                curve = self.row(book, '小转大累计完成度')
                self.assertEqual([curve[f'D{i}'] for i in (1, 2, 3)], [.25, .75, 1])
                self.assertEqual(self.row(book, D12_SHEET)['D1小转大/总小转大'], .25)

    def test_all_invalid_totals_allow_available_real_quantities_to_fill(self):
        for value in ('待更新', True):
            with self.subTest(value=value):
                book = self.build(dict.fromkeys(self.quantity_fields, value), add_mix=True)
                row = self.row(book)
                for field, expected in zip(self.quantity_fields, (100, 40, 130, 90, 6, 117, 65)):
                    self.assertEqual(row[field], expected, field)
                self.assertAlmostEqual(row['留存大定率'], .9)
                self.assertAlmostEqual(row['大定到锁单率'], .5)

    def test_true_zero_total_keeps_priority_over_nonzero_daily_sum(self):
        book = self.build({'总小转大': 0})
        self.assertEqual(self.row(book)['总小转大'], 0)
        self.assertEqual(self.row(book)['小订转化率'], 0)
        daily = self.row(book, '小转大当日数量')
        self.assertEqual([daily[f'D{i}'] for i in (1, 2, 3)], [10, 20, 10])
        self.assertTrue(all(self.row(book, '小转大累计完成度')[f'D{i}'] is None for i in (1, 2, 3)))

    def test_maintained_numeric_strings_keep_valid_quantity_priority(self):
        book = self.build({'总小订': '1,000', '总小转大': '+40',
                           '总大定': '130', '总直接大定': ' 90 '})
        master = self.row(book)
        self.assertEqual([master[field] for field in ('总小订', '总小转大', '总大定', '总直接大定')],
                         [1000, 40, 130, 90])

    def test_internal_daily_gap_without_a_real_anchor_keeps_terminal_unknown(self):
        book = self.build({'总小转大': '#DIV/0!'}, launch_quantities={
            '当日小订转大定数量': [10, None, 10], '累计小订转大定数量': [None]*3,
            '当日大定数量': [50, None, 30], '当日直接大定数量': [40, None, 20],
            '累计直接大定数量': [None]*3,
        })
        self.assertIsNone(self.row(book)['总小转大'])
        self.assertIsNone(self.row(book)['小订转化率'])
        daily = self.row(book, '小转大当日数量')
        self.assertEqual([daily[f'D{i}'] for i in (1, 2, 3)], [10, None, 10])
        self.assertTrue(all(self.row(book, '小转大累计完成度')[f'D{i}'] is None for i in (1, 2, 3)))

    def test_initial_unknown_totals_do_not_fabricate_ratios_or_d12_totals(self):
        for value in (None, '#DIV/0!', True, .6, 0):
            with self.subTest(value=value):
                maintained = dict.fromkeys(self.quantity_fields)
                maintained.update(dict.fromkeys(self.rate_fields, value))
                full = self.build(maintained, source_quantities={'小转大进度': [10, None, 5]})
                root = Path(full._test_path).parent
                output = root/'initial.xlsx'
                source = root/'小订及首销数据整理.xlsx'
                build_secondary(source, output, source)
                book = load_workbook(output, data_only=True)
                self.addCleanup(book.close)
                master = self.row(book, '预测基准总表')
                expected = value if value is not None and type(value) in (int, float) else None
                for field in self.rate_fields:
                    self.assertEqual(master[field], expected, field)
                for field in ('总小订', '总小转大', '总直接大定', '总大定'):
                    self.assertIsNone(self.row(book, D12_SHEET)[field], field)
                self.assertIsNone(self.row(book, '小转大累计完成度')['D1'])

    def test_rate_validation_does_not_turn_invalid_independent_rate_into_zero(self):
        for value in (None, '', '#DIV/0!', '待更新', True, False, -1, 1.2,
                      float('nan'), float('inf')):
            with self.subTest(value=value):
                self.assertIsNone(optional_reference_rate(value))
        self.assertEqual(optional_reference_rate('0.6'), .6)
        self.assertEqual(optional_reference_rate(0), 0)

    def test_invalid_input_warning_identifies_file_sheet_model_and_field(self):
        with self.assertLogs('sales_forecast_refresh', level='WARNING') as captured:
            self.build({'总小转大': '待更新'})
        warning = next(line for line in captured.output if '[历史维护值无效]' in line)
        for text in ('小订及首销数据整理.xlsx', '车型汇总', '享界 G9 2027款', '小订转大定量', '待更新'):
            self.assertIn(text, warning)
