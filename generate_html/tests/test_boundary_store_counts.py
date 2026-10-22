from __future__ import annotations

from datetime import date
from pathlib import Path
import logging
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from core.excel import WorkbookItem, WorkbookStore
from core.models import Subject
from modules.channel_analysis import ChannelAnalysisModule


CHANNEL = "华为旗舰店"
NEW = ("门店类型", "贡献门店数")
OLD = ("贡献订单门店数", "数量")
MISSING = object()
RIGHT = (date(2026, 1, 1), date(2026, 10, 22))
LEFT = (date(2022, 1, 14), date(2025, 12, 31))
FILENAMES = ("鸿蒙智行锁单选配比例分析.xlsx", "鸿蒙智行锁单选配比例分析_2021-2025.xlsx")


class BoundaryStoreCountTests(unittest.TestCase):
    def setUp(self):
        # A test must not leave process-wide conflict suppression enabled for
        # later tests which independently assert physical-cell diagnostics.
        self.conflicts = patch("core.excel._CONFLICT_CHECKS", None)
        self.conflicts.start()
        self.addCleanup(self.conflicts.stop)

    def source(self, index, active, *, labels=NEW, span=MISSING, period="26WK01",
               filename=None, active_row=3, column=4, additional_rows=()):
        book = Workbook()
        sheet = book.active
        sheet.title = "鸿蒙智行 by周"
        sheet.append(["指标", "统计类型", "分类"])
        sheet.cell(1, column, period)
        for col, value in enumerate(["门店类型", "数量", CHANNEL], 1):
            sheet.cell(2, col, value)
        sheet.cell(2, column, 10 + index)
        if active is not MISSING:
            for col, value in enumerate([*labels, CHANNEL], 1):
                sheet.cell(active_row, col, value)
            sheet.cell(active_row, column, active)
        for metric, stat, value in additional_rows:
            row = sheet.max_row + 1
            for col, entry in enumerate([metric, stat, CHANNEL], 1):
                sheet.cell(row, col, entry)
            sheet.cell(row, column, value)
        item = WorkbookItem(Path(filename or FILENAMES[index]), book)
        item._boundary_date_range = ([RIGHT, LEFT][index] if span is MISSING else span)
        return item

    def merge(self, items):
        store = WorkbookStore(Path("."))
        store.items = items
        self.addCleanup(store.close)
        keyword = "大定选配比例" if "大定" in items[0].path.name else "锁单选配比例"
        item, sheet = store.find_subject_sheet(keyword, "鸿蒙智行", "week")
        return store, item, sheet

    def cells(self, sheet, labels):
        return [sheet.cell(row, 4).value for row in range(2, sheet.max_row + 1)
                if tuple(sheet.cell(row, col).value for col in (1, 2, 3)) == (*labels, CHANNEL)]

    def active(self, sheet, labels=NEW):
        cells = self.cells(sheet, labels)
        self.assertEqual(len(cells), 1)
        return cells[0]

    def test_user_cross_year_ranges_zero_and_nonzero_are_exact_and_located(self):
        items = [self.source(0, 0, active_row=64), self.source(1, 3, active_row=54, column=207)]
        with self.assertLogs("core.excel", level="INFO") as logs:
            _, _, sheet = self.merge(items)
        self.assertEqual(self.active(sheet), 3)
        self.assertEqual(self.cells(sheet, ("门店类型", "数量")), [21])
        self.assertTrue(all(record.levelno == logging.INFO for record in logs.records))
        output = "\n".join(logs.output)
        self.assertIn("边界周门店数合并", output)
        self.assertIn("真实0", output)
        self.assertIn("2026-01-01~2026-10-22", output)
        self.assertIn("2022-01-14~2025-12-31", output)
        self.assertIn("单元格=D64", output)
        self.assertIn("单元格=GY54", output)
        for filename in FILENAMES:
            self.assertIn(filename, output)
        self.assertNotIn("多文件数值冲突", output)
        self.assertNotIn("未在边界日衔接", output)
        self.assertEqual(items[0].workbook.active["D64"].value, 0)
        self.assertEqual(items[1].workbook.active["GY54"].value, 3)

    def test_zero_union_is_independent_of_source_order_and_accepted_row_format(self):
        for active in [(0, 3), (3, 0), (0, 0)]:
            for formats in [(NEW, NEW), (OLD, OLD), (NEW, OLD), (OLD, NEW)]:
                for reverse in (False, True):
                    with self.subTest(active=active, formats=formats, reverse=reverse):
                        items = [self.source(index, value, labels=formats[index]) for index, value in enumerate(active)]
                        with self.assertLogs("core.excel", level="INFO") as logs:
                            _, _, sheet = self.merge(items[::-1] if reverse else items)
                        for labels in set(formats):
                            self.assertEqual(self.active(sheet, labels), max(active))
                        self.assertFalse(any(record.levelno >= logging.WARNING for record in logs.records))

    def test_two_positive_counts_remain_unknown_without_false_value_conflicts(self):
        for active in [(4, 3), (4, 4)]:
            with self.subTest(active=active), self.assertLogs("core.excel", level="WARNING") as logs:
                _, _, sheet = self.merge([self.source(index, value) for index, value in enumerate(active)])
            self.assertIsNone(self.active(sheet))
            self.assertEqual(len(logs.records), 1)
            self.assertIn("缺少门店明细无法去重", logs.output[0])
            self.assertNotIn("多文件数值冲突", logs.output[0])
            self.assertNotIn("未在边界日衔接", logs.output[0])

    def test_blank_missing_and_invalid_counts_never_become_zero(self):
        for unknown in [None, "", MISSING, -1, 1.5, "not a count", "3元", "3e0", "0%",
                        "3％", True, False, float("nan"), float("inf")]:
            for reverse in (False, True):
                with self.subTest(unknown=unknown, reverse=reverse):
                    items = [self.source(0, 0), self.source(1, unknown)]
                    with self.assertLogs("core.excel", level="WARNING") as logs:
                        _, _, sheet = self.merge(items[::-1] if reverse else items)
                    self.assertIsNone(self.active(sheet))
                    self.assertIn("缺失、无效或存在重复键差异", "\n".join(logs.output))
                    self.assertNotIn("多文件数值冲突", "\n".join(logs.output))

    def test_new_format_priority_is_resolved_in_each_source_before_union(self):
        for new_value, expected in [(0, 3), (None, None), (2, None)]:
            with self.subTest(new_value=new_value):
                current = self.source(0, 9, labels=OLD, additional_rows=[(*NEW, new_value)])
                historical = self.source(1, 3, labels=OLD)
                with self.assertLogs("core.excel", level="INFO"):
                    _, _, sheet = self.merge([current, historical])
                self.assertEqual(self.active(sheet, OLD), expected)
                self.assertEqual(self.active(sheet, NEW), expected)

    def test_half_year_boundary_uses_the_same_empty_set_rule(self):
        ranges = [(date(2026, 6, 29), date(2026, 6, 30)), (date(2026, 7, 1), date(2026, 7, 5))]
        with self.assertLogs("core.excel", level="INFO"):
            _, _, sheet = self.merge([self.source(index, value, span=ranges[index], period="26WK27")
                                      for index, value in enumerate([0, 5])])
        self.assertEqual(self.active(sheet), 5)

    def test_extra_source_cannot_make_a_zero_union_look_complete(self):
        for extra in [0, 3, 7, None]:
            with self.subTest(extra=extra):
                items = [self.source(0, 0), self.source(1, 3),
                         self.source(0, extra, filename="鸿蒙智行锁单选配比例分析_补充.xlsx")]
                with self.assertLogs("core.excel", level="WARNING") as logs:
                    _, _, sheet = self.merge(items)
                self.assertIsNone(self.active(sheet))
                output = "\n".join(logs.output)
                self.assertIn("存在额外来源", output)
                self.assertIn("多文件数值冲突", output)

    def test_conflicting_duplicate_inside_a_shard_stays_unknown_and_warns(self):
        current = self.source(0, 0, additional_rows=[(*NEW, 2)])
        historical = self.source(1, 3)
        with self.assertLogs("core.excel", level="WARNING") as logs:
            _, _, sheet = self.merge([current, historical])
        self.assertIsNone(self.active(sheet))
        output = "\n".join(logs.output)
        self.assertIn("同文件重复键冲突", output)
        self.assertIn("同一物理文件存在重复业务键", output)
        self.assertIn("全周结果见专用门店分段合并或去重说明", output)
        self.assertNotIn("多文件数值冲突", output)
        self.assertNotIn("未在边界日衔接", output)

    def test_identical_duplicate_inside_a_shard_does_not_add_a_store_twice(self):
        current = self.source(0, 0, additional_rows=[(*NEW, 0)])
        with self.assertLogs("core.excel", level="INFO") as logs:
            _, _, sheet = self.merge([current, self.source(1, 3)])
        self.assertEqual(self.active(sheet), 3)
        self.assertFalse(any(record.levelno >= logging.WARNING for record in logs.records))

    def test_valid_numeric_text_keeps_the_channel_input_contract(self):
        for empty, count in [("0", "3"), ("+0.0", "3.0"), (0, "1,000"), ("0.0", "+3")]:
            with self.subTest(empty=empty, count=count), self.assertLogs("core.excel", level="INFO"):
                _, _, sheet = self.merge([self.source(0, empty), self.source(1, count)])
            self.assertEqual(self.active(sheet), float(count.replace(",", "")))

    def test_conflicting_duplicate_week_column_cannot_resolve_an_empty_set(self):
        current = self.source(0, 0)
        original = current.workbook.active
        original.cell(1, 5, "2026WK01")
        original.cell(2, 5, 10)
        original.cell(3, 5, 2)
        with self.assertLogs("core.excel", level="WARNING") as logs:
            _, _, sheet = self.merge([current, self.source(1, 3)])
        self.assertIsNone(self.active(sheet))
        self.assertIn("同文件重复键冲突", "\n".join(logs.output))

    def test_unknown_and_overlapping_coverage_keep_real_conflict_diagnostics(self):
        for span in [None, (date(2025, 12, 29), date(2026, 1, 4)),
                     (date(2026, 1, 1), date(2026, 10, 22))]:
            with self.subTest(span=span), self.assertLogs("core.excel", level="WARNING") as logs:
                _, _, sheet = self.merge([self.source(0, 0), self.source(1, 3, span=span)])
            self.assertEqual(self.active(sheet), 0)
            self.assertIn("多文件数值冲突", "\n".join(logs.output))
            self.assertNotIn("边界周门店数合并", "\n".join(logs.output))

    def test_ordinary_week_keeps_first_source_and_real_conflict(self):
        with self.assertLogs("core.excel", level="WARNING") as logs:
            _, _, sheet = self.merge([self.source(0, 0, period="26WK02"), self.source(1, 3, period="26WK02")])
        self.assertEqual(self.active(sheet), 0)
        self.assertEqual(self.cells(sheet, ("门店类型", "数量")), [10])
        self.assertIn("多文件数值冲突", "\n".join(logs.output))

    def test_channel_dashboard_receives_exact_count_and_uses_its_configured_denominator(self):
        items = [self.source(index, value, filename=filename.replace("锁单", "大定"))
                 for index, (value, filename) in enumerate(zip([0, 3], FILENAMES))]
        store = WorkbookStore(Path("."))
        store.items = items
        self.addCleanup(store.close)
        with self.assertLogs("core.excel", level="INFO"):
            dashboard = ChannelAnalysisModule().build(store, Subject("鸿蒙智行", "鸿蒙智行", "group"))
        row = dashboard.views["week"]["pages"]["26WK01"]["sections"][0]["data"]["rows"][0]
        self.assertEqual(row["active_count"], 3)
        self.assertEqual(row["count"], 21)
        self.assertEqual(row["active_rate"], 3 / row["store_count"])


if __name__ == "__main__":
    unittest.main()
