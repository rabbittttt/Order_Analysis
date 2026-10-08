from datetime import date, datetime, timedelta
import logging
from pathlib import Path
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook

from core.excel import WorkbookItem, WorkbookStore, _merge_records
from core.forecast_summary import MASTER_SHEET, DAILY_SHEET
from modules.sales_forecast import _complete_steady_weeks, _stage_window, _read_steady_history
from tools.refresh_sales_forecast_data import audit_final_forecast, curve_for


MODEL = '问界 M9 2026款'


class ForecastDiagnosticsTests(TestCase):
    def test_all_real_problems_and_repeated_occurrences_are_warning(self):
        book, _ = self.final_book(None, None, None, has_small=False)
        sheet = book[MASTER_SHEET]
        values = [cell.value for cell in sheet[2]]
        sheet.delete_rows(2)
        for index in range(9):
            sheet.append([f'测试车{index}', *values[1:]])
        with self.assertLogs('sales_forecast_refresh', level='WARNING') as logs:
            first = audit_final_forecast(book, date(2026, 9, 3))
            second = audit_final_forecast(book, date(2026, 9, 3))
        self.assertEqual(len(first), 9)
        self.assertEqual(first, second)
        self.assertEqual([r.getMessage() for r in logs.records], first + second)
        self.assertTrue(all(r.levelno == logging.WARNING for r in logs.records))
        self.assertIn('测试车8', logs.records[8].getMessage())
        book.close()

    def test_unmatched_preliminary_curve_is_debug_not_final_failure(self):
        with self.assertLogs('sales_forecast_refresh', level='DEBUG') as logs:
            self.assertEqual(curve_for('未来车', {}, {}, {}, '退订'), [])
        self.assertEqual([r.levelno for r in logs.records], [logging.DEBUG])

    def test_all_alias_curve_conflicts_print_warning_and_keep_values(self):
        curves = {'标准车': [35]*9, '别名': [2]*9}
        with self.assertLogs('sales_forecast_refresh', level='WARNING') as logs:
            result = curve_for('标准车', curves, {}, {'标准车': {'标准车', '别名'}}, '小转大')
        self.assertEqual(result, [35]*9)
        self.assertEqual(len(logs.records), 9)
        self.assertIn('D9 | 标准车=35；别名=2', logs.output[-1])

    def cancellation_sheet(self, book, title, value):
        sheet = book.create_sheet(title)
        sheet.append(['退订分析'])
        sheet.append(['取消日期', '当日小订退'])
        sheet.append([datetime(2026, 7, 27), value])
        return sheet

    def test_same_file_conflict_locates_both_original_sheets_and_cells(self):
        book = Workbook()
        first = self.cancellation_sheet(book, 'V800  2026款_日度退订', 35)
        second = self.cancellation_sheet(book, 'V800 2026款_日度退订', 2)
        item = WorkbookItem(Path('鸿蒙智行小订退订分析.xlsx'), book)
        with self.assertLogs('core.excel', level='DEBUG') as logs:
            merged = _merge_records([(item, first), (item, second)], ({'取消日期'},))
        self.assertEqual(merged.cell(3, 2).value, 35)
        self.assertEqual(merged._merge_conflict_kinds['同文件重复键'], 1)
        detail = '\n'.join(logs.output)
        self.assertIn('Sheet=V800  2026款_日度退订', detail)
        self.assertIn('Sheet=V800 2026款_日度退订', detail)
        self.assertIn('单元格=B3: 35', detail)
        self.assertIn('单元格=B3: 2', detail)
        book.close()

    def test_same_sheet_duplicate_dates_are_located_not_added(self):
        book = Workbook()
        first = self.cancellation_sheet(book, 'V800_日度退订', 35)
        first.append([datetime(2026, 7, 27), 2])
        other = Workbook(); second = self.cancellation_sheet(other, first.title, 35)
        with self.assertLogs('core.excel', level='DEBUG') as logs:
            merged = _merge_records([(WorkbookItem(Path('小订退订分析.xlsx'), book), first),
                                     (WorkbookItem(Path('小订退订分析_历史.xlsx'), other), second)], ({'取消日期'},))
        self.assertEqual(merged.cell(3, 2).value, 35)
        self.assertIn('单元格=B4: 2', '\n'.join(logs.output))
        book.close(); other.close()

    def test_cross_file_conflict_keeps_priority_and_reports_locations(self):
        books = [Workbook(), Workbook()]
        for book, count in zip(books, (35, 2)):
            book.remove(book.active)
            self.cancellation_sheet(book, 'V800_日度退订', count)
        store = WorkbookStore(Path('.'))
        store.items = [WorkbookItem(Path(f'小订退订分析{i}.xlsx'), b) for i, b in enumerate(books)]
        with self.assertLogs('core.excel', level='INFO') as logs:
            item = store.find('小订退订分析')
        self.assertEqual(item.workbook.worksheets[0].cell(3, 2).value, 35)
        self.assertIn('冲突1项', '\n'.join(logs.output))
        self.assertIn('Sheet=V800_日度退订 / 单元格=B3', logs.output[0])
        store.close()

    def test_conflicting_cells_are_grouped_with_full_range_not_three_examples(self):
        books = [Workbook(), Workbook()]
        for book in books:
            book.remove(book.active)
        sheets = [self.cancellation_sheet(b, 'V800_日度退订', value) for b, value in zip(books, (35, 2))]
        for index in range(1, 9):
            for sheet, value in zip(sheets, (35, 2)):
                sheet.append([datetime(2026, 7, 27)+timedelta(days=index), value])
        store = WorkbookStore(Path('.'))
        store.items = [WorkbookItem(Path(f'小订退订分析{i}.xlsx'), b) for i, b in enumerate(books)]
        with self.assertLogs('core.excel', level='WARNING') as logs:
            merged = store.find('小订退订分析').workbook.worksheets[0]
        self.assertEqual(len(logs.records), 1)
        self.assertIn('冲突9项', logs.output[0])
        self.assertIn('2026-07-27～2026-08-04', logs.output[0])
        self.assertIn('单元格=B3:B11', logs.output[0])
        self.assertEqual(merged.cell(11, 2).value, 35)
        store.close()

    def test_stage_end_is_shared_and_input_record_is_not_modified(self):
        window = {'generation': MODEL, 'launch_date': '2026-06-01', 'days': 30}
        parsed = _stage_window({'m9': window}, MODEL)
        self.assertEqual(parsed['end_date'], '2026-06-30')
        self.assertNotIn('end_date', window)
        invalid = {**window, 'end_date': '2026-05-30'}
        self.assertEqual(_stage_window({'m9': invalid}, MODEL)['end_date'], '2026-05-30')

    def test_steady_reader_uses_derived_launch_end(self):
        book = Workbook(); sheet = book.active; sheet.title = MODEL+'by周'
        sheet.append(['指标', '统计类型', '分类', '26WK28'])
        sheet.append(['交车锁单', '数量', '数量', 700])
        item = WorkbookItem(Path('锁单选配比例.xlsx'), book)
        with patch('modules.sales_forecast._read_model_mapping', return_value={}), patch('modules.sales_forecast._read_model_master', return_value={}):
            rows, _ = _read_steady_history(SimpleNamespace(find=lambda _: item),
                {'m9': {'generation': MODEL, 'launch_date': '2026-06-01', 'days': 30}}, date(2026, 7, 20))
        self.assertEqual(rows[0]['steady_start_date'], '2026-07-01')
        self.assertEqual(rows[0]['weeks'][0]['lock'], 700)
        book.close()

    def test_zero_week_requires_full_coverage_and_blank_week_stays_unknown(self):
        start = date(2026, 7, 6); today = date(2026, 7, 28)
        weeks = [{'period': '26WK28', 'start_date': '2026-07-06', 'end_date': '2026-07-12', 'lock': 7}]
        coverage = [(start, date(2026, 7, 26))]
        result = _complete_steady_weeks(list(weeks), {'26WK28': 7, '26WK30': None}, [], coverage, start, today)
        self.assertEqual([(r['period'], r['lock']) for r in result], [('26WK28', 7), ('26WK29', 0)])
        short = _complete_steady_weeks(list(weeks), {'26WK28': 7}, [], [(start, date(2026, 7, 18))], start, today)
        self.assertEqual(len(short), 1)
        unknown_day = [{'date': '2026-07-13', 'lock': None, 'complete': True}]
        blocked = _complete_steady_weeks(list(weeks), {'26WK28': 7}, unknown_day, coverage, start, date(2026, 7, 21))
        self.assertEqual(len(blocked), 1)

    def final_book(self, small_end, launch_start, launch_end, total=None, has_small=True):
        book = Workbook(); book.remove(book.active)
        master = book.create_sheet(MASTER_SHEET)
        master.append(['代际名', '二级代际名', '有小订', '小订开始', '小订结束', '首销开始', '首销结束', '首销天数', '总小订', '首销期留存大定', '首销期锁单'])
        master.append([MODEL, None, has_small, date(2026, 9, 1) if has_small else None, small_end, launch_start, launch_end, 2, total, 0, 0])
        daily = book.create_sheet(DAILY_SHEET)
        daily.append(['代际名', '二级代际名', '日期', '小订数量', '大定', '小转大', '直接大定'])
        return book, daily

    def test_future_terminal_fields_are_not_required_and_complete_actuals_resolve_gaps(self):
        book, daily = self.final_book(date(2026, 9, 4), date(2026, 9, 5), date(2026, 9, 6))
        book[MASTER_SHEET]['J2'] = None
        book[MASTER_SHEET]['K2'] = None
        for day in (1, 2): daily.append([MODEL, None, date(2026, 9, day), 5, None, None, None])
        self.assertEqual(audit_final_forecast(book, date(2026, 9, 3)), [])
        book.close()

    def test_final_missing_small_days_and_total_disagreement_remain_visible(self):
        book, daily = self.final_book(date(2026, 9, 2), date(2026, 9, 5), date(2026, 9, 6), total=20)
        daily.append([MODEL, None, date(2026, 9, 1), 5])
        with self.assertLogs('sales_forecast_refresh', level='WARNING'):
            issues = audit_final_forecast(book, date(2026, 9, 3))
        self.assertIn('D2', '\n'.join(issues))
        daily.append([MODEL, None, date(2026, 9, 2), 5])
        with self.assertLogs('sales_forecast_refresh', level='WARNING'):
            issues = audit_final_forecast(book, date(2026, 9, 3))
        self.assertIn('多来源逐日合计=10', '\n'.join(issues))
        self.assertNotIn('数量缺失', '\n'.join(issues))
        book.close()

    def test_no_small_vehicle_and_true_zero_terminal_are_valid(self):
        book, daily = self.final_book(None, date(2026, 9, 1), date(2026, 9, 2), has_small=False)
        for day in (1, 2): daily.append([MODEL, None, date(2026, 9, day), None, 0, None, 0])
        self.assertEqual(audit_final_forecast(book, date(2026, 9, 3)), [])
        book.close()

    def test_ended_terminal_gap_limits_only_corresponding_reference(self):
        book, daily = self.final_book(None, date(2026, 9, 1), date(2026, 9, 2), has_small=False)
        book[MASTER_SHEET]['J2'] = None
        book[MASTER_SHEET]['K2'] = None
        for day in (1, 2): daily.append([MODEL, None, date(2026, 9, day), None, 0, None, 0])
        with self.assertLogs('sales_forecast_refresh', level='WARNING'):
            issues = audit_final_forecast(book, date(2026, 9, 3))
        self.assertEqual(len(issues), 1)
        self.assertIn('不停止其他有效预测', issues[0])
        self.assertIn('首销期留存大定、首销期锁单', issues[0])
        book.close()
