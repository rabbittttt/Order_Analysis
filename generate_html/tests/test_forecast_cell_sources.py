"""Selected physical-cell provenance, independent of equal-valued candidates."""
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest import TestCase
from unittest.mock import patch
from openpyxl import Workbook
from core.excel import (WorkbookItem, WorkbookStore, cell_source_text, cell_origin_points,
                        covered_date_source_text, join_source_texts, source_files_text)
from core.model_identity import model_key
from modules.sales_forecast import _read_actual_profiles, _read_steady_history, _resolve_stage_candidate


class ForecastCellSourceTests(TestCase):
    model = "问界 M9 2026款"
    day = date(2026, 9, 1)

    def item(self, filename, rows, title=None, coverage=None):
        book = Workbook(); self.addCleanup(book.close)
        book.active.title = title or self.model+"by天"
        for row in rows: book.active.append(row)
        item = WorkbookItem(Path(filename), book)
        if coverage: item._boundary_date_range = coverage
        return item

    def store(self, items):
        store = WorkbookStore(Path(".")); store.items = items
        return store

    def read(self, items, aliases=(), windows=None):
        mapping = {model_key(name): self.model for name in (self.model, *aliases)}
        with patch("modules.sales_forecast._read_model_mapping", return_value=mapping), \
             patch("modules.sales_forecast.generation_records", return_value=[]):
            profiles, _ = _read_actual_profiles(self.store(items), windows, today=self.day+timedelta(days=3))
        return next(p for p in profiles if p["model"] == self.model)

    def mix(self, suffix, values, title=None, days=None, coverage=None):
        return self.item("小订选配比例_"+suffix+".xlsx", [
            ["指标", "统计类型", "分类", *(days or [self.day, self.day+timedelta(days=1)])],
            ["小订", "数量", "数量", *values]], title, coverage)

    def test_zero_and_blank_fill_keep_their_own_file(self):
        first, second = self.mix("A", [0, None]), self.mix("B", [9, 30])
        p = self.read([first, second]); title=self.model+"by天"
        self.assertEqual([r["orders"] for r in p["small_daily_days"]], [0, 30])
        self.assertEqual({day: join_source_texts(origin) for day, origin in p["small_daily_sources"].items()}, {"2026-09-01": f"{first.path.name}｜{title}",
                                                  "2026-09-02": f"{second.path.name}｜{title}"})

    def test_equal_values_keep_only_first_selected_cell(self):
        first, second = self.mix("A", [0, 30]), self.mix("B", [0, 30])
        sheet=self.store([first, second]).find("小订选配比例").workbook.worksheets[0]
        self.assertEqual(cell_origin_points(sheet, 2, 5), ((first.path.name, self.model+"by天", "E2"),))

    def test_alias_keeps_original_sheet_for_second_file_fill(self):
        alias="M9旗舰2026款"
        first, second=self.mix("A", [0, None]), self.mix("B", [9, 30], alias+"by天")
        self.assertEqual(join_source_texts(self.read([first, second], (alias,))["small_daily_sources"]["2026-09-02"]),
                         f"{second.path.name}｜{alias}by天")

    def test_physical_omitted_date_zero_keeps_coverage_source(self):
        end=self.day+timedelta(days=2)
        first=self.mix("A", [1, None], days=[self.day, end], coverage=(self.day, end))
        second=self.mix("B", [9, 7], days=[self.day+timedelta(days=1), end], coverage=(self.day, end))
        p=self.read([first, second])
        self.assertEqual([r["orders"] for r in p["small_daily_days"]], [1, 0, 7])
        self.assertEqual(join_source_texts(p["small_daily_sources"]["2026-09-02"]), f"{first.path.name}｜{self.model}by天")
        self.assertEqual(covered_date_source_text(first, first.workbook.active, "2026-09-02"),
                         f"{first.path.name}｜{self.model}by天")

    def test_cancel_daily_zero_and_cumulative_have_independent_owners(self):
        title=self.model+"_日度退订"; header=["取消日期", "当日小订退", "累计小订退"]
        first=self.item("小订退订分析_A.xlsx", [header, [self.day, 0, None]], title)
        second=self.item("小订退订分析_B.xlsx", [header, [self.day, 9, 10]], title)
        row=self.read([first, second])["cancel_days"][0]
        self.assertEqual((row["daily_cancel"], row["cancel"]), (0, 10))
        self.assertEqual(row["_field_origins"], {"daily_cancel":f"{first.path.name}｜{title}", "cancel":f"{second.path.name}｜{title}"})

    def test_small_hour_sum_lists_every_actual_contributor(self):
        title=self.model+"_小订分时退订"; header=["小订日期", "小订小时", "小订数"]
        first=self.item("小订退订分析_A.xlsx", [header, [self.day, 1, 2]], title)
        second=self.item("小订退订分析_B.xlsx", [header, [self.day, 2, 3]], title)
        p=self.read([first, second]); bucket=p["small_hourly_days"][0]
        self.assertEqual(bucket["orders"], 5)
        self.assertEqual(p["small_daily_sources"]["2026-09-01"], f"{first.path.name}｜{title}、{second.path.name}｜{title}")
        self.assertEqual(bucket["hours"][1]["_field_origins"]["orders"], f"{second.path.name}｜{title}")

    def test_launch_day_and_cumulative_follow_selected_cells(self):
        def item(suffix,gross,direct,cumulative):
            return self.item("首销期订单节奏_"+suffix+".xlsx", [["周期", "D1"], ["时间", self.day],
                ["当日大定数量", gross], ["当日直接大定数量", direct], ["累计直接大定数量", cumulative]])
        first, second=item("A",0,None,None),item("B",9,3,3)
        p=self.read([first,second]); row=p["days"][0]; title=self.model+"by天"
        self.assertEqual((row["gross"],row["direct"]),(0,3))
        self.assertEqual(row["_field_origins"], {"gross":f"{first.path.name}｜{title}","direct":f"{second.path.name}｜{title}"})
        self.assertEqual(p["cumulative_days"][0]["_field_origins"]["direct"], f"{second.path.name}｜{title}")

    def test_complete_hour_day_sum_keeps_both_selected_file_sources(self):
        title=self.model+"by时"; times=[datetime(2026,9,1,h) for h in range(24)]
        def item(suffix,values):
            return self.item("首销期订单节奏_"+suffix+".xlsx", [["周期",*range(24)], ["时间",*times], ["当日大定数量",*values]],title)
        first,second=item("A",[1]*12+[None]*12),item("B",[None]*12+[2]*12)
        p=self.read([first,second]); row=p["days"][0]
        self.assertEqual(row["gross"],36)
        self.assertEqual(row["_field_origins"]["gross"],f"{first.path.name}｜{title}、{second.path.name}｜{title}")
        self.assertEqual(p["hourly_days"][0]["hours"][12]["_field_origins"]["gross"],f"{second.path.name}｜{title}")

    def test_boundary_week_sum_has_both_contributing_cells(self):
        title=self.model+"by周"
        def item(suffix,value,first,last):
            return self.item("大定选配比例_"+suffix+".xlsx", [["指标","统计类型","分类","26WK27"],
                ["大定","数量","数量",value]],title,(first,last))
        first=item("A",20,date(2026,6,29),date(2026,6,30))
        second=item("B",30,date(2026,7,1),date(2026,7,5))
        sheet=self.store([first,second]).find("大定选配比例").workbook.worksheets[0]
        self.assertEqual(sheet.cell(2,4).value,50)
        self.assertEqual(cell_source_text(sheet,2,4),f"{first.path.name}｜{title}、{second.path.name}｜{title}")
        self.assertEqual(source_files_text(cell_source_text(sheet,2,4)),first.path.name+"、"+second.path.name)

    def test_stage_merge_moves_origin_with_selected_value(self):
        first={"days":[{"date":"2026-09-01","gross":0,"net":None,"_field_origins":{"gross":"A.xlsx｜A","net":"stale.xlsx｜stale"}}]}
        second={"days":[{"date":"2026-09-01","gross":9,"net":3,"_field_origins":{"gross":"B.xlsx｜B","net":"B.xlsx｜B"}}]}
        row=_resolve_stage_candidate({"history":first,"launch":second},"active")["days"][0]
        self.assertEqual((row["gross"],row["net"]),(0,3))
        self.assertEqual(row["_field_origins"],{"gross":"A.xlsx｜A","net":"B.xlsx｜B"})
        self.assertEqual(join_source_texts("A.xlsx｜A","A.xlsx｜A","B.xlsx｜B"),"A.xlsx｜A、B.xlsx｜B")

    def test_no_small_direct_orders_use_the_selected_gross_and_net_sources(self):
        title=self.model+"by天"; header=["指标","统计类型","分类",self.day]
        first=self.item("大定选配比例_A.xlsx", [header,["大定","数量","数量",None],["净大定","数量","数量",0]],title)
        second=self.item("大定选配比例_B.xlsx", [header,["大定","数量","数量",30],["净大定","数量","数量",20]],title)
        windows={model_key(self.model):{"generation":self.model,"has_small":False,
            "launch_date":"2026-09-01","end_date":"2026-09-01"}}
        row=self.read([first,second],windows=windows)["days"][0]
        self.assertEqual((row["gross"],row["net"],row["direct"],row["small_to_big"]),(30,0,30,0))
        self.assertEqual(row["_field_origins"]["gross"],f"{second.path.name}｜{title}")
        self.assertEqual(row["_field_origins"]["direct"],f"{second.path.name}｜{title}")
        self.assertEqual(row["_field_origins"]["net"],f"{first.path.name}｜{title}")

    def test_steady_day_and_week_do_not_revert_to_the_first_file_label(self):
        title=self.model+"by天"; header=["指标","统计类型","分类",date(2026,8,31),self.day]
        first=self.item("锁单选配比例_A.xlsx",[header,["交车锁单","数量","数量",0,None]],title)
        second=self.item("锁单选配比例_B.xlsx",[header,["交车锁单","数量","数量",9,5]],title)
        for item,value in ((first,None),(second,10)):
            sheet=item.workbook.create_sheet(self.model+"by周")
            sheet.append(["指标","统计类型","分类","26WK36"])
            sheet.append(["交车锁单","数量","数量",value])
        mapping={model_key(self.model):self.model}
        windows={model_key(self.model):{"generation":self.model,"launch_date":"2026-08-29","end_date":"2026-08-30"}}
        with patch("modules.sales_forecast._read_model_mapping",return_value=mapping), \
             patch("modules.sales_forecast._read_model_master",return_value={}), \
             patch("modules.sales_forecast.generation_records",return_value=[]):
            result,_=_read_steady_history(self.store([first,second]),windows,date(2026,9,7))
        row=result[0]
        self.assertEqual(row["daily_sources"]["2026-08-31"],f"{first.path.name}｜{title}")
        self.assertEqual(row["daily_sources"]["2026-09-01"],f"{second.path.name}｜{title}")
        self.assertEqual(row["weeks"][0]["_field_origins"]["lock"],f"{second.path.name}｜{self.model}by周")
