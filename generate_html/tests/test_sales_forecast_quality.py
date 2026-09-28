from __future__ import annotations

import unittest
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from openpyxl import Workbook, load_workbook

from tools.refresh_sales_forecast_data import cumulative, curve_for, day_structure_quality, mapping_record_for, summary_quality, extract_quantity_table, workbook_header_index, combine_daily, build_secondary


class SalesForecastQualityTests(unittest.TestCase):
    def test_partial_source_survives_generated_workbook_and_derived_ratios(self):
        with TemporaryDirectory() as temp:
            source, output = Path(temp) / "raw.xlsx", Path(temp) / "summary.xlsx"
            book = Workbook()
            book.active.title = "车型汇总"
            book.active.append(["车型", "总小订", "小订转大定量", "大定量", "直接大定量", "开始大定日期", "小转大结束日期", "首销期天数"])
            book.active.append(["进行中车型", 1000, None, None, None, datetime(2026, 9, 1), datetime(2026, 9, 4), 4])
            for name, values in [("小转大进度", [10, 20, None, None]), ("直接大定进度", [5, 0, None, None]), ("大定进度", [15, None, None, None]), ("退订进度", [0, 1, None, None])]:
                sheet = book.create_sheet(name)
                sheet.append(["车型", "D1", "D2", "D3", "D4"])
                sheet.append(["进行中车型", *values])
                sheet.append(["车型", *[f"D{i}" for i in range(1, 8)]])
                sheet.append(["进行中车型", *[None] * 7])
            book.save(source)
            book.close()
            with patch("tools.refresh_sales_forecast_data.load_mapping", return_value=({}, {})):
                build_secondary(source, output, Path(temp) / "mapping.xlsx")
            result = load_workbook(output, data_only=True)
            try:
                def values(name):
                    return list(next(result[name].iter_rows(min_row=2, max_row=2, values_only=True)))[1:]
                self.assertEqual(values("小转大当日数量"), [10, 20, None, None])
                self.assertEqual(values("直接大定当日数量"), [5, 0, None, None])
                self.assertEqual(values("总大定当日数量"), [15, 20, None, None])
                self.assertEqual(values("小转大累计完成度"), [None] * 4)
                self.assertEqual(values("累计小订转化率"), [.01, .03, None, None])
                for actual, expected in zip(values("累计直接大定占比")[:2], [1 / 3, 1 / 7]):
                    self.assertAlmostEqual(actual, expected)
            finally:
                result.close()

    def test_partial_quantity_block_is_not_replaced_by_longer_percentage_block(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "小转大进度"
        sheet.append(["车型", "D1", "D2", "D3", "D4"])
        sheet.append(["进行中车型", 10, 0, None, None])
        sheet.append([])
        sheet.append([])
        sheet.append(["第二车型", 30, 20, None, None])
        sheet.append(["车型", *[f"D{i}" for i in range(1, 20)]])
        sheet.append(["进行中车型", *[None] * 19])
        rows, count, _, header = extract_quantity_table(book, workbook_header_index(book), "small")
        self.assertEqual(header, 1)
        self.assertEqual(count, 4)
        self.assertEqual(rows["进行中车型"], [10, 0, None, None])
        self.assertEqual(rows["第二车型"], [30, 20, None, None])
        self.assertEqual(cumulative(rows["进行中车型"], 100), [.1, .1, None, None])
        book.close()

    def test_explicit_aliases_fill_known_days_without_erasing_partial_curve(self):
        aliases = {"传播名": {"传播名", "简称", "旧简称"}}
        result = curve_for("传播名", {"传播名": [10, None, None], "简称": [None, 0, None], "旧简称": [10, None, 20]}, {}, aliases, "小转大")
        self.assertEqual(result, [10, 0, 20])
        self.assertEqual(combine_daily([10, None, 0], [5, 20, 0]), [15, None, 0])

    def test_cumulative_keeps_future_dates_blank(self):
        self.assertEqual(cumulative([100, 50, None, None], 1000), [.1, .15, None, None])

    def test_cumulative_keeps_explicit_zero_observed(self):
        self.assertEqual(cumulative([100, 0, 50], 1000), [.1, .1, .15])

    def test_cumulative_does_not_claim_complete_total_after_gap(self):
        self.assertEqual(cumulative([100, None, 50], 1000), [.1, None, None])
        self.assertEqual(cumulative([None, 100], 1000), [None, None])

    def test_refresh_queries_accept_propagation_generation_and_alias_names(self):
        record = {"历史传播名": "尊界 G9", "订单分析代际名": "尊界 G9 2027款"}
        mapping = {"尊界g9": record}
        aliases = {
            "尊界g9": {"尊界g9", "尊界g92027", "g9旗舰"},
            "尊界g92027": {"尊界g9", "尊界g92027", "g9旗舰"},
        }
        self.assertIs(mapping_record_for("尊界 G9 2027款", mapping, aliases), record)
        self.assertIs(mapping_record_for("G9旗舰", mapping, aliases), record)
        self.assertEqual(curve_for("尊界 G9 2027款", {"尊界 G9": [1, 2]}, mapping, aliases, "测试"), [1, 2])

    def test_day_structure_rejects_impossible_components(self):
        valid, issue = day_structure_quality(80, 20, 100, "D1")
        self.assertTrue(valid)
        self.assertEqual(issue, "")
        valid, issue = day_structure_quality(120, 10, 100, "D1")
        self.assertFalse(valid)
        self.assertIn("高于总大定", issue)

    def test_summary_quality_separates_completeness_and_consistency(self):
        completeness, consistency, issues = summary_quality(
            1000, 300, .30, 400, 100, .25, 50, .05,
            350, .875, 280, .70,
        )
        self.assertEqual(completeness, 1)
        self.assertEqual(consistency, 1)
        self.assertEqual(issues, [])

        completeness, consistency, issues = summary_quality(
            1000, 450, .45, 400, 100, .25, 50, .05,
            350, .875, 280, .70,
        )
        self.assertEqual(completeness, 1)
        self.assertLess(consistency, 1)
        self.assertTrue(any("总小转大+总直接大定" in issue for issue in issues))

    def test_summary_quality_can_treat_an_explicit_zero_as_present(self):
        completeness, _, _ = summary_quality(
            1000, 300, .30, 500, 200, .40, 0, 0,
            450, .90, 400, .80,
            (True,) * 12,
        )
        self.assertEqual(completeness, 1)


if __name__ == "__main__":
    unittest.main()
