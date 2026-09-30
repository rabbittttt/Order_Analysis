"""Regression coverage for campaign applicability and scoped forecast conditions."""
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import Workbook
from core.excel import WorkbookItem, WorkbookStore, parse_metric_sheet
from core.forecast_summary import MASTER_SHEET, read_public_forecast
from modules.sales_forecast import _apply_small_order_applicability, _profile_hard_errors, _target_options
from tools.refresh_sales_forecast_data import forecast_weekly_orders


class ForecastConditionTests(unittest.TestCase):
    def target(self, **extra):
        return dict(name="直接大定车", stage="active", launch_date="2026-09-01",
                    end_date="2026-09-03", days=3, launch_days_maintained=True, **extra)

    def test_membership_not_missing_total_decides_reservation_applicability(self):
        profiles = [{"model":"有小订车"}, {"model":"直接大定车"}]
        history = [{"model":"有小订车","launch_date":"2026-09-01","days":3,"small":0}]
        history[0].update(tier="SUV",energy="纯电",node="标准发布",launch_weekday="周二",launch_period="下午")
        result = _target_options(history, profiles, {}, today=date(2026,9,2))
        self.assertTrue(result[0]["has_small"])
        self.assertFalse(result[1]["has_small"])

    def test_explicit_summary_flag_overrides_generated_reference_membership(self):
        history = [{"model":"直接大定车","launch_date":"2026-09-01","days":3,"small":0}]
        window = {"generation":"直接大定车","has_small":False,"launch_date":"2026-09-01","days":3}
        history[0].update(tier="SUV",energy="纯电",node="标准发布",launch_weekday="周二",launch_period="下午")
        result = _target_options(history, [{"model":"直接大定车"}], {"direct":window}, today=date(2026,9,2))
        self.assertFalse(result[0]["has_small"])

    def test_no_reservation_is_a_business_identity_not_a_ratio_estimate(self):
        target = self.target(has_small=False)
        profile = {"days":[{"date":"2026-09-01","gross":100,"small_to_big":None,"direct":None}],
                   "missing_fields":["总小订"], "stage_profiles":{"active":{"days":[{"gross":0}]}}}
        _apply_small_order_applicability(target, profile)
        self.assertEqual(profile["days"][0]["small_to_big"], 0)
        self.assertEqual(profile["days"][0]["direct"], 100)
        self.assertEqual(profile["stage_profiles"]["active"]["days"][0]["direct"], 0)
        self.assertEqual(profile["missing_fields"], [])
        self.assertEqual(target["small_stage"], "not_applicable")
        self.assertEqual(_profile_hard_errors(target,profile,today=date(2026,9,2)), [])

    def test_missing_small_total_does_not_block_progress_method(self):
        profile={"total_small":0,"days":[{"date":"2026-09-01","gross":100,"small_to_big":80,"direct":20}]}
        self.assertEqual(_profile_hard_errors(self.target(has_small=True),profile,today=date(2026,9,2)), [])

    def test_ended_gross_survives_unknown_components_and_small_total(self):
        target=self.target(has_small=True)
        target.update(stage="ended",days=1,end_date="2026-09-01")
        profile={"total_small":0,"days":[{"date":"2026-09-01","gross":100,"small_to_big":None,"direct":None}]}
        self.assertEqual(_profile_hard_errors(target,profile,today=date(2026,9,2)), [])

    def test_missing_gross_is_still_a_blocking_real_day_issue(self):
        profile={"days":[{"date":"2026-09-01","gross":None,"small_to_big":None,"direct":None}]}
        self.assertTrue(_profile_hard_errors(self.target(has_small=False),profile,today=date(2026,9,2)))

    def test_aggregate_columns_neither_participate_nor_produce_merge_conflicts(self):
        store=WorkbookStore(Path("."))
        for year,total in [(2025,1000),(2026,200)]:
            book=Workbook();sheet=book.active;sheet.title="直接大定车by天"
            sheet.append(["指标","统计类型","分类","2026-09-01","总计","近28天"])
            sheet.append(["大定","数量","数量",10,total,total])
            store.items.append(WorkbookItem(Path(f"大定选配比例分析{year}.xlsx"),book))
        try:
            with patch("core.excel.LOGGER.warning") as warning:
                item=store.find("大定选配比例")
            warning.assert_not_called()
            self.assertEqual(list(parse_metric_sheet(item.workbook.worksheets[0])),["2026-09-01"])
            self.assertEqual(store.items[0].workbook.active.cell(1,5).value,"总计")
        finally:
            store.close()

    def test_current_week_uses_updated_day_without_false_stage_split_warning(self):
        book=Workbook();sheet=book.active;sheet.title="直接大定车by天"
        sheet.append(["指标","统计类型","分类",date(2026,9,28)])
        for label,value in [("大定",7),("留存大定",6),("交车锁单",5)]:
            sheet.append([label,"数量","数量",value])
        week=book.create_sheet("直接大定车by周")
        week.append(["指标","统计类型","分类","26WK40"])
        for label,value in [("大定",7),("留存大定",6),("交车锁单",5)]:
            week.append([label,"数量","数量",value])
        item=WorkbookItem(Path("大定选配比例分析.xlsx"),book)
        store=SimpleNamespace(find=lambda key:item if key=="大定选配比例" else None)
        data={"targets":[{"name":"直接大定车","launch_date":"2026-09-01","end_date":"2026-09-20"}],"actuals":[]}
        dashboard=SimpleNamespace(views={"week":{"pages":{"预测方案":{"workspace":{"data":data}}}}})
        with patch("modules.sales_forecast._read_model_mapping",return_value={}),patch("tools.refresh_sales_forecast_data.LOGGER.warning") as warning:
            rows=forecast_weekly_orders(store,dashboard,date(2026,9,30))
        current=next(row for row in rows if row[1]=="26WK40")
        self.assertEqual(current[5:8],[7,6,5])
        self.assertEqual(current[4].date(),date(2026,9,28))
        warning.assert_not_called()
        book.close()
