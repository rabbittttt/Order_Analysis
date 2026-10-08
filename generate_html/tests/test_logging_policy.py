from datetime import date
import logging
from pathlib import Path
from unittest import TestCase
from unittest.mock import MagicMock, patch

from openpyxl import Workbook

from core.excel import (WorkbookItem, WorkbookStore, _merge_matrix, _merge_partitioned, _merge_records,
                        compact_diagnostic_ranges, _coordinate_ranges, reset_conflict_diagnostics)
import main


class LoggingPolicyTests(TestCase):
    def matrix(self, periods, values, *, title='车型_SKU', filename='SKU分析.xlsx', span=None):
        book = Workbook()
        self.addCleanup(book.close)
        sheet = book.active
        sheet.title = title
        sheet.append(['累计占比', *([None] * len(periods))])
        sheet.append(['排名', *periods])
        for index, row in enumerate(values, 1):
            sheet.append([f'TOP{index}', *row])
        item = WorkbookItem(Path(filename), book)
        item._boundary_date_range = span
        return item, sheet

    def test_cross_file_aggregate_is_not_comparable_and_values_stay_unchanged(self):
        periods = ['总计', '近28天', '总计占比']
        matches = [self.matrix(periods, [[.4, .5, .6]], filename='SKU2026.xlsx'),
                   self.matrix(periods, [[.8, .9, 1]], filename='SKU2025.xlsx')]
        with self.assertLogs('core.excel', level='INFO') as logs:
            result = _merge_matrix(matches, 2, 1)
        self.assertTrue(all(r.levelno == logging.INFO for r in logs.records))
        self.assertIn('没有共同统计范围', logs.output[0])
        self.assertEqual({result.cell(2, c).value: result.cell(3, c).value for c in range(2, 5)},
                         dict(zip(periods, [.4, .5, .6])))

    def test_duplicate_aggregate_inside_same_file_still_warns(self):
        matches = [self.matrix(['总计'], [[.4]], title='车_SKU', filename='SKU.xlsx'),
                   self.matrix(['总计'], [[.8]], title='车 _SKU', filename='SKU.xlsx')]
        with self.assertLogs('core.excel', level='WARNING') as logs:
            _merge_matrix(matches, 2, 1)
        self.assertEqual(len(logs.records), 1)
        self.assertIn('同文件重复键冲突', logs.output[0])

    def test_proven_split_ratio_has_info_but_no_false_conflict(self):
        for period, left, right in [('26WK01', (date(2025, 12, 29), date(2025, 12, 31)), (date(2026, 1, 1), date(2026, 1, 4))),
                                    ('26WK27', (date(2026, 6, 29), date(2026, 6, 30)), (date(2026, 7, 1), date(2026, 7, 5)))]:
            with self.subTest(period=period):
                matches = [self.matrix([period], [[.4]], filename='SKU_A.xlsx', span=left),
                           self.matrix([period], [[.8]], filename='SKU_B.xlsx', span=right)]
                with self.assertLogs('core.excel', level='INFO') as logs:
                    result = _merge_matrix(matches, 2, 1)
                self.assertTrue(all(r.levelno == logging.INFO for r in logs.records))
                self.assertIn('不代表合并全周指标', logs.output[0])
                self.assertEqual(result.cell(3, 2).value, .4)

    def test_unknown_boundary_scope_lists_all_cells_without_claiming_value_conflict(self):
        matches = [self.matrix(['26WK01'], [[.4]]*80, filename='SKU_2021-2025.xlsx'),
                   self.matrix(['26WK01'], [[.8]]*80, filename='SKU_2026.xlsx')]
        with self.assertLogs('core.excel', level='WARNING') as logs:
            _merge_matrix(matches, 2, 1)
        self.assertEqual(len(logs.records), 1)
        self.assertIn('统计范围待核对', logs.output[0])
        self.assertIn('B3:B82', logs.output[0])
        self.assertIn('涉及80处差异', logs.output[0])
        self.assertNotIn('多文件数值冲突', logs.output[0])

    def test_full_overlapping_boundary_groups_all_real_differences(self):
        span = (date(2025, 12, 29), date(2026, 1, 4))
        matches = [self.matrix(['26WK01'], [[.4]]*8, filename='SKU_A.xlsx', span=span),
                   self.matrix(['26WK01'], [[.8]]*8, filename='SKU_B.xlsx', span=span)]
        with self.assertLogs('core.excel', level='WARNING') as logs:
            result = _merge_matrix(matches, 2, 1)
        self.assertEqual(len(logs.records), 1)
        self.assertIn('冲突8项', logs.output[0])
        self.assertIn('TOP1～TOP8', logs.output[0])
        self.assertIn('B3:B10', logs.output[0])
        self.assertEqual(result.cell(10, 2).value, .4)

    def test_float_tail_is_ignored_but_real_difference_and_zero_are_not(self):
        matches = [self.matrix(['26WK02'], [[.1 + .2], [.3], [0]], filename='SKU_A.xlsx'),
                   self.matrix(['26WK02'], [[.3], [.300001], [.2]], filename='SKU_B.xlsx')]
        with self.assertLogs('core.excel', level='WARNING') as logs:
            result = _merge_matrix(matches, 2, 1)
        self.assertEqual(len(logs.records), 1)
        self.assertIn('冲突2项', logs.output[0])
        self.assertTrue(all('TOP1' not in line for line in logs.output))
        self.assertEqual(result.cell(3, 2).value, .1 + .2)
        self.assertEqual(result.cell(5, 2).value, 0)

    def test_partitioned_conflict_reports_original_not_temporary_coordinate(self):
        matches = []
        for filename, column, value in [('SKU_A.xlsx', 5, .4), ('SKU_B.xlsx', 9, .8)]:
            book = Workbook(); self.addCleanup(book.close)
            sheet = book.active; sheet.title = '车型_SKU'
            sheet.cell(1, 1, '订单数量'); sheet.cell(2, 1, 'SKU'); sheet.cell(2, 2, '26WK02'); sheet.cell(3, 1, 'A'); sheet.cell(3, 2, 10)
            sheet.cell(1, column, '总计占比'); sheet.cell(2, column, '排名'); sheet.cell(2, column+1, '26WK02')
            sheet.cell(3, column, 'TOP1'); sheet.cell(3, column+1, value)
            matches.append((WorkbookItem(Path(filename), book), sheet))
        with self.assertLogs('core.excel', level='WARNING') as logs:
            _merge_partitioned(matches, horizontal=True)
        self.assertEqual(len(logs.records), 1)
        self.assertIn('单元格=F3: 0.4', logs.output[0])
        self.assertIn('单元格=J3: 0.8', logs.output[0])

    def test_chart_conflicts_include_every_period_and_original_cell(self):
        store = WorkbookStore(Path('.')); self.addCleanup(store.close)
        for i in range(2):
            book = Workbook(); sheet = book.active; sheet.title = '问界by天图表'
            sheet.append(['问界'])
            sheet.append([None, '留存大定', *[f'2026-01-{d:02d}' for d in range(1, 6)]])
            sheet.append([None, '版本A', *([.4 + i*.2]*5)])
            sheet.append([None, '总计', *([100 + i]*5)])
            store.items.append(WorkbookItem(Path(f'大定选配比例{i}.xlsx'), book))
        with self.assertLogs('core.excel', level='WARNING') as logs:
            result = store.find_chart_data('大定选配比例', '问界', 'day')[2]
        self.assertEqual(len(logs.records), 2)
        self.assertTrue(all('冲突5项' in line and '2026-01-01～2026-01-05' in line for line in logs.output))
        self.assertTrue(any('单元格=C4:G4' in line for line in logs.output))
        self.assertTrue(any('单元格=C3:G3' in line for line in logs.output))
        self.assertEqual(result['totals'], [100]*5)
        self.assertFalse(any(key.startswith('_log_') for key in result))

    def test_many_duplicate_records_group_by_metric_without_losing_dates(self):
        matches = []
        for title, amount in [('车型 _日度退订', 35), ('车型_日度退订', 2)]:
            book = Workbook(); self.addCleanup(book.close)
            sheet = book.active; sheet.title = title
            sheet.append(['取消日期', '当日小订退', '累计小订退'])
            for day in range(1, 29):
                sheet.append([date(2026, 7, day), amount, amount*day])
            matches.append((WorkbookItem(Path('小订退订分析.xlsx'), book), sheet))
        with self.assertLogs('core.excel', level='WARNING') as logs:
            result = _merge_records(matches, ({'取消日期'},))
        self.assertEqual(len(logs.records), 2)
        self.assertTrue(all('冲突28项' in line and '2026-07-01～2026-07-28' in line for line in logs.output))
        self.assertTrue(any('单元格=B2:B29' in line for line in logs.output))
        self.assertEqual(result.cell(2, 2).value, 35)

    def test_same_key_equal_or_complementary_fields_do_not_warn(self):
        matches = []
        for value, extra in [(0, None), (0, 7)]:
            book = Workbook(); self.addCleanup(book.close)
            sheet = book.active; sheet.title = '车_日度退订'
            sheet.append(['日期', '数量', '补充数量'])
            sheet.append([date(2026, 7, 1), value, extra])
            matches.append((WorkbookItem(Path('小订退订分析.xlsx'), book), sheet))
        with patch('core.excel.LOGGER.warning') as warning:
            result = _merge_records(matches, ({'日期'},))
        warning.assert_not_called()
        self.assertEqual([result.cell(2, c).value for c in (2, 3)], [0, 7])

    def test_large_daily_conflicts_are_bounded_by_metrics_not_cells(self):
        from datetime import timedelta
        matches = []
        for title, value in [('车 _日度退订', 35), ('车_日度退订', 2)]:
            book = Workbook(); self.addCleanup(book.close)
            sheet = book.active; sheet.title = title
            sheet.append(['日期', *[f'指标{i}' for i in range(13)]])
            for index in range(1420):
                sheet.append([date(2022, 1, 1)+timedelta(days=index), *([value]*13)])
            matches.append((WorkbookItem(Path('小订退订分析.xlsx'), book), sheet))
        with self.assertLogs('core.excel', level='WARNING') as logs:
            result = _merge_records(matches, ({'日期'},))
        self.assertEqual(len(logs.records), 13)  # 18,460 cell differences, all 13 metrics visible.
        self.assertTrue(all('冲突1420项' in line for line in logs.output))
        self.assertEqual(result._merge_conflict_count, 18460)
        self.assertEqual(result.cell(1421, 14).value, 35)
        self.assertLess(sum(len(line) for line in logs.output), 10000)

    def test_run_dedup_preserves_new_values_and_resets_between_runs(self):
        matches = [self.matrix(['26WK02'], [[.4]], filename='SKU_A.xlsx'),
                   self.matrix(['26WK02'], [[.8]], filename='SKU_B.xlsx')]
        reset_conflict_diagnostics()
        self.addCleanup(reset_conflict_diagnostics, False)
        with self.assertLogs('core.excel', level='WARNING') as logs:
            _merge_matrix(matches, 2, 1)
            _merge_matrix(matches, 2, 1)
            matches[1][1].cell(3, 2).value = .9
            _merge_matrix(matches, 2, 1)
            reset_conflict_diagnostics()
            _merge_matrix(matches, 2, 1)
        self.assertEqual(len(logs.records), 3)

    def test_ranges_preserve_gaps_and_rectangles_have_no_missing_cells(self):
        self.assertEqual(compact_diagnostic_ranges(['D1', 'D2', 'D4', 'D5', 'D8']), 'D1～D2、D4～D5、D8')
        self.assertEqual(compact_diagnostic_ranges(f'D{i}' for i in range(1, 187)), 'D1～D186')
        self.assertEqual(compact_diagnostic_ranges(['26/07/27', '26/07/28', '26/07/30']), '2026-07-27～2026-07-28、2026-07-30')
        self.assertEqual(compact_diagnostic_ranges(['26WK52', '26WK53', '27WK01']), '26WK52～27WK01')
        self.assertEqual(_coordinate_ranges({'B2', 'C2', 'D2'}), 'B2:D2')
        self.assertEqual(_coordinate_ranges({'B2', 'B3', 'C2', 'C3'}), 'B2:C3')
        self.assertEqual(_coordinate_ranges({'B2', 'B3', 'C2'}), 'C2、B2:B3')

    def test_filter_dedups_refresh_and_parent_but_not_other_issues_or_errors(self):
        diagnostic_filter = main.DuplicateDiagnosticFilter()
        def record(message, level=logging.WARNING):
            return logging.LogRecord('test', level, '', 0, message, (), None)
        first = record('[销量预测刷新] [冲突] A文件 B2=3 / C文件 B2=5')
        # Shared filter on three handlers must accept the same first record for all.
        self.assertEqual([diagnostic_filter.filter(first) for _ in range(3)], [True]*3)
        self.assertFalse(diagnostic_filter.filter(record('[冲突] A文件 B2=3 / C文件 B2=5')))
        self.assertTrue(diagnostic_filter.filter(record('[冲突] A文件 B2=3 / C文件 B2=6')))
        self.assertTrue(diagnostic_filter.filter(record('[冲突] D文件 B2=3 / C文件 B2=5')))
        self.assertTrue(diagnostic_filter.filter(record('读取失败', logging.ERROR)))
        self.assertTrue(diagnostic_filter.filter(record('读取失败', logging.ERROR)))

    def test_known_extension_warnings_are_info_and_context_line_is_not_printed(self):
        with self.assertLogs('order_analysis', level='INFO') as logs:
            level = main.relay_refresh_line('C:/lib/openpyxl/worksheet/_reader.py:329: UserWarning: Unknown extension is not supported and will be removed')
            level = main.relay_refresh_line('  warn(msg)', level)
            level = main.relay_refresh_line('C:/lib/openpyxl/worksheet/_reader.py:329: UserWarning: Conditional Formatting extension is not supported and will be removed', level)
            level = main.relay_refresh_line('  warn(msg)', level)
            main.relay_refresh_line('C:/lib/other.py:1: UserWarning: 数据类型未知', level)
        self.assertEqual([r.levelno for r in logs.records], [20, 20, 30])
        self.assertTrue(all('warn(msg)' not in line for line in logs.output))

    def test_relay_preserves_severity_duplicates_and_one_traceback_issue(self):
        collector = main.DiagnosticCollector()
        before = list(main.BUILD_DIAGNOSTICS); counts = dict(main.DIAGNOSTIC_COUNTS)
        self.addCleanup(lambda: main.BUILD_DIAGNOSTICS.__setitem__(slice(None), before))
        self.addCleanup(lambda: main.DIAGNOSTIC_COUNTS.update(counts))
        main.BUILD_DIAGNOSTICS.clear(); main.DIAGNOSTIC_COUNTS.update(WARNING=0, ERROR=0)
        with self.assertLogs('order_analysis', level='DEBUG') as logs:
            main.LOGGER.addHandler(collector)
            try:
                level = logging.INFO
                for line in ['[WARNING] 缺D3', '[WARNING] 缺D3', '[ERROR] 刷新失败',
                             'Traceback (most recent call last):', '  File "source.py", line 1',
                             'ValueError: 输入格式错误', '[INFO] 收尾', '[DEBUG] 读取细节']:
                    level = main.relay_refresh_line(line, level)
            finally:
                main.LOGGER.removeHandler(collector)
        self.assertEqual([r.levelno for r in logs.records], [30,30,40,40,40,40,20,10])
        self.assertTrue(all('[WARNING]' not in r.getMessage() and '[ERROR]' not in r.getMessage() for r in logs.records))
        self.assertEqual(main.DIAGNOSTIC_COUNTS, {'WARNING': 2, 'ERROR': 1})
        self.assertEqual(len(main.BUILD_DIAGNOSTICS), 2)

    def test_refresh_relays_before_wait_and_propagates_failure(self):
        process = MagicMock()
        process.__enter__.return_value = process
        observed = []
        def output():
            yield '[WARNING] 第一个问题\n'
            self.assertTrue(observed)  # Already relayed while the child is still producing output.
            yield '[ERROR] 第二个问题\n'
        process.stdout = output(); process.wait.return_value = 1
        with patch('main.subprocess.Popen', return_value=process) as popen, patch('main.relay_refresh_line', side_effect=lambda line, level: observed.append(line) or logging.WARNING):
            with self.assertRaisesRegex(RuntimeError, '退出码 1'):
                main.refresh_sales_forecast_data({'refresh_sales_forecast_data_before_build': True})
        self.assertEqual(len(observed), 2)
        self.assertIn('-u', popen.call_args.args[0])
        self.assertEqual(popen.call_args.kwargs['stderr'], main.subprocess.STDOUT)
