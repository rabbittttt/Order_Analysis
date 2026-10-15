"""Exercise physical source workbooks through the production summary builder."""
from datetime import date, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from core.forecast_summary import (MASTER_SHEET, D12_SHEET, DAILY_SHEET, SMALL_DAILY_SHEET,
                                  SMALL_CURVE_SHEET, LAUNCH_HOURLY_SHEET, WEEKLY_SHEET,
                                  table_records, summary_scope, read_public_forecast)
from modules.sales_forecast import _read_history, _optional_count
from tools.refresh_sales_forecast_data import build_secondary, audit_source_retention


class SourceBackfillEndToEndTests(TestCase):
    model = '享界 G9 2027款'

    def build(self, *, end='2026-09-03', daily=True, hourly=True, dates=None,
              cancel_headers=None, maintained=None, text_counts=False, as_of='2026-09-05',
              source_quantities=None, launch_quantities=None, add_sources=None, source_edit=None):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        root = Path(temporary.name)
        source, output, orders = root/'小订及首销数据整理.xlsx', root/'汇总.xlsx', root/'orders'
        orders.mkdir()
        book = Workbook()
        sheet = book.active
        sheet.title = '车型汇总'
        headers = ['代际名', '二级代际名', '产品档位', '能源类型', '发布类型',
                   '小订发布时段', '首销发布时段', '小订开始日期', '小订结束日期',
                   '首销开始日期', '首销结束日期', '首销天数', '总小订',
                   '总小转大', '小订转化率', '总大定', '总直接大定', '直接大定占比',
                   '总退订', '退订率']
        record = dict(zip(headers[:13], [self.model, None, '豪华轿车', '纯电', '标准发布',
                      '下午', '晚上', datetime(2026, 8, 30), datetime(2026, 8, 31),
                      datetime(2026, 9, 1), datetime.fromisoformat(end), 3, 100]))
        record.update(maintained or {})
        sheet.append(headers)
        sheet.append([record.get(k) for k in headers])
        for name in ('小转大进度', '直接大定进度', '大定进度', '退订进度'):
            sheet = book.create_sheet(name)
            sheet.append(['代际名', 'D1', 'D2', 'D3'])
            sheet.append([self.model, *(source_quantities or {}).get(name, [None, None, None])])
        sheet = book.create_sheet('小订by天')
        sheet.append(['代际名', datetime(2026, 8, 30), datetime(2026, 8, 31), '合计'])
        sheet.append([self.model, 40, 60, 100])
        book.save(source)
        book.close()
        if source_edit:
            source_edit(source)

        book = Workbook()
        book.remove(book.active)
        for grain, enabled in (('天', daily), ('时', hourly)):
            if not enabled:
                continue
            sheet = book.create_sheet(self.model+'by'+grain)
            prefix = 'D' if grain == '天' else 'H'
            sheet.append(['周期', *[f'{prefix}{i}' for i in range(1, 4)]])
            sheet.append(['时间', *(dates or [datetime(2026, 9, i, 20 if grain == '时' else 0) for i in range(1, 4)])])
            for name, values in (
                ('当日大定数量', [50, 50, 30]), ('当日小订转大定数量', [10, 20, 10]),
                ('累计小订转大定数量', [10, 30, 40]), ('当日直接大定数量', [40, 30, 20]),
                ('累计直接大定数量', [40, 70, 90]), ('当日交车锁单数量', [25, 25, 15])):
                values = (launch_quantities or {}).get(name, values)
                sheet.append([name, *[f' {value:,} ' if text_counts and value is not None else value for value in values]])
        book.save(orders/'鸿蒙智行首销期订单节奏.xlsx')
        book.close()

        book = Workbook()
        sheet = book.active
        sheet.title = self.model+'_日度退订'
        sheet.append(['日度退订'])
        sheet.append(cancel_headers or ['取消日期', '当日小订退', '累计小订退'])
        for i, cumulative in enumerate((1, 3, 6), 1):
            sheet.append([datetime(2026, 9, i), *[str(n) if text_counts else n for n in (i, cumulative)]])
        book.save(orders/'鸿蒙智行小订退订分析.xlsx')
        book.close()

        if add_sources:
            add_sources(orders)
        build_secondary(source, output, source, orders, as_of)
        result = load_workbook(output, data_only=True)
        result._test_path = output
        self.addCleanup(result.close)
        return result

    def assert_complete(self, book):
        rows = table_records(book, MASTER_SHEET)
        self.assertEqual(len(rows), 1, rows)
        master = rows[0]
        for field, value in {'总小订':100, '总小转大':40, '小订转化率':.4,
                             '总直接大定':90, '总大定':130, '直接大定占比':90/130,
                             '总退订':6, '退订率':.06}.items():
            self.assertAlmostEqual(master[field], value, msg=field)
        metrics = table_records(book, D12_SHEET)[0]
        for field, value in {'D1小转大':10, 'D2小转大':20, 'D1直接大':40,
                             'D2直接大':30, 'D1小转大/总小转大':.25,
                             'D1退订率':.01, 'D2退订率':.02}.items():
            self.assertAlmostEqual(metrics[field], value, msg=field)
        for sheet, values in (('小转大累计完成度', [.25, .75, 1]),
                              ('直接大定累计完成度', [40/90, 70/90, 1]),
                              ('累计大定完成度', [50/130, 100/130, 1]),
                              ('累计小订转化率', [.1, .3, .4]),
                              ('累计直接大定占比', [.8, .7, 90/130]),
                              ('每日直接大定占比', [.8, .6, 20/30]),
                              ('累计退订率', [.01, .03, .06])):
            row = table_records(book, sheet)[0]
            for i, expected in enumerate(values, 1):
                self.assertAlmostEqual(row[f'D{i}'], expected, msg=f'{sheet}/D{i}')

    def test_only_small_total_with_both_launch_sheets_fills_saved_workbook(self):
        self.assert_complete(self.build())

    def test_text_quantities_in_both_launch_and_cancel_sheets_fill_saved_workbook(self):
        self.assert_complete(self.build(text_counts=True))

    def web_inputs(self, book, today):
        with summary_scope(book._test_path, workbook=book):
            _, history = _read_history()
            profiles, windows, small, steady = read_public_forecast(book, history, today=today)
        return history, profiles, windows, small, steady

    def test_saved_facts_and_web_eligibility_across_every_stage(self):
        for stamp, phase, completed in [
            ('2026-08-29', '未开始', 0), ('2026-08-30', '未开始', 0),
            ('2026-08-31', '未开始', 0), ('2026-09-01', '进行中', 0),
            ('2026-09-02', '进行中', 1), ('2026-09-03', '进行中', 2),
            ('2026-09-04', '已结束', 3), ('2026-09-20', '已结束', 3)]:
            with self.subTest(stamp=stamp):
                book = self.build(as_of=stamp)
                self.assert_complete(book)  # Saving never drops an explicit source observation.
                master = table_records(book, MASTER_SHEET)[0]
                self.assertEqual(master['首销阶段'], phase)
                self.assertEqual(master['首销参考终值有效'], phase == '已结束')
                rows = table_records(book, DAILY_SHEET)
                self.assertEqual(len(rows), 5)
                self.assertEqual(len({(r['订单分析代际名'], r['日期']) for r in rows}), len(rows))
                self.assertEqual(len(table_records(book, LAUNCH_HOURLY_SHEET)), 3)
                history, profiles, _, small, _ = self.web_inputs(book, date.fromisoformat(stamp))
                item = history[0]
                self.assertEqual(item['daily_small'], [10, 20, 10][:completed])
                if completed < 3:
                    for key in ('conversion', 'direct_share', 'cancel_rate', 'gross'):
                        self.assertIsNone(item[key], key)
                    self.assertEqual(item['small_progress'], [])
                else:
                    self.assertEqual(item['conversion'], .4)
                    self.assertEqual(item['small_progress'], [.25, .75, 1])
                self.assertEqual(item['d1_valid'], completed >= 1)
                self.assertEqual(item['d2_valid'], completed >= 2)
                self.assertTrue(all(r['date'] <= stamp for p in profiles for r in p['days']))
                self.assertEqual(small[0]['total_complete'], stamp > '2026-08-31')
                self.assertTrue(all(d < stamp for d in small[0]['dates']))
                self.assertTrue(all(s.sheet_state == 'visible' for s in book))

    def test_high_priority_partial_and_zero_cells_fill_per_field(self):
        book = self.build(source_quantities={
            '小转大进度': [0, None, None], '直接大定进度': [50, None, None],
            '大定进度': [50, None, None]})
        master = table_records(book, MASTER_SHEET)[0]
        self.assertEqual((master['总小转大'], master['总直接大定'], master['总大定']), (30, 100, 130))
        small = table_records(book, '小转大当日数量')[0]
        self.assertEqual([small[f'D{i}'] for i in (1, 2, 3)], [0, 20, 10])
        self.assertEqual(table_records(book, D12_SHEET)[0]['D2小转大'], 20)
        self.assertEqual(table_records(book, '小转大累计完成度')[0]['D3'], 1)

    def test_daily_source_alone_is_sufficient(self):
        self.assert_complete(self.build(hourly=False))

    def test_chart_sheets_cannot_overwrite_actual_day_and_hour_sheets(self):
        def add_sources(folder):
            path = folder/'鸿蒙智行首销期订单节奏.xlsx'
            book = load_workbook(path)
            for grain in ('天', '时'):
                sheet = book.create_sheet(self.model+'by'+grain+'图表')
                sheet.append(['图表展示'])
            book.save(path)
            book.close()
        book = self.build(add_sources=add_sources)
        self.assert_complete(book)
        self.assertEqual(len(table_records(book, LAUNCH_HOURLY_SHEET)), 3)

    def test_matching_hourly_snapshot_fills_missing_day_components(self):
        def add_sources(folder):
            path = folder/'鸿蒙智行首销期订单节奏.xlsx'
            book = load_workbook(path)
            sheet = book[self.model+'by天']
            for row in (4, 5, 6, 7):
                for col in range(2, 5):
                    sheet.cell(row, col).value = None
            book.save(path)
            book.close()
        self.assert_complete(self.build(add_sources=add_sources))

    def test_hourly_components_work_without_explicit_gross_and_provenance_is_saved(self):
        def add_sources(folder):
            path = folder/'鸿蒙智行首销期订单节奏.xlsx'
            book = load_workbook(path)
            for col in range(2, 5):
                for row in (4, 5, 6, 7):
                    book[self.model+'by天'].cell(row, col).value = None
                book[self.model+'by时'].cell(3, col).value = None
            book.save(path)
            book.close()
        book = self.build(add_sources=add_sources)
        self.assert_complete(book)
        row = next(r for r in table_records(book, DAILY_SHEET) if r.get('小转大') == 10)
        self.assertIn('分时补缺', row['字段来源'])
        self.assertIn('by时', row['字段来源'])

    def test_partial_hour_snapshot_cannot_fill_full_day_or_replace_zero(self):
        def add_sources(folder):
            path = folder/'鸿蒙智行首销期订单节奏.xlsx'
            book = load_workbook(path)
            day, hour = book[self.model+'by天'], book[self.model+'by时']
            for col in range(2, 5):
                for row in (4, 5, 6, 7):
                    day.cell(row, col).value = None
            day['C4'] = 0
            hour['B3'] = 20  # only part of D1's independent 50 orders
            hour['B4'] = 2
            hour['B6'] = 18
            hour['D4'] = None  # known gross alone must not fabricate a component
            book.save(path)
            book.close()
        book = self.build(add_sources=add_sources)
        rows = table_records(book, '小转大当日数量')[0]
        self.assertEqual([rows[f'D{i}'] for i in (1, 2, 3)], [None, 0, None])
        self.assertIsNone(table_records(book, MASTER_SHEET)[0]['总小转大'])

    def test_retention_audit_detects_dropped_source_and_corrupted_related_cells(self):
        book = self.build()
        raw = [{'model': self.model, 'days': [{'date': '2026-09-01', 'small_to_big': 10}]}]
        self.assertEqual(audit_source_retention(book, raw), [])
        def clear(sheet, field, row=2):
            headers = [c.value for c in book[sheet][1]]
            book[sheet].cell(row, headers.index(field)+1).value = None
        clear('小转大当日数量', 'D1')
        clear(D12_SHEET, 'D2直接大')
        clear('累计退订率', 'D2')
        clear(MASTER_SHEET, '总直接大定')
        details = '\n'.join(audit_source_retention(book, raw))
        for label in ('小转大当日数量', 'D2', '累计退订率', '总直接大定'):
            self.assertIn(label, details)
        headers = [c.value for c in book[DAILY_SHEET][1]]
        for cells in book[DAILY_SHEET].iter_rows(min_row=2):
            if cells[headers.index('日期')].value == datetime(2026, 9, 1):
                cells[headers.index('小转大')].value = None
        self.assertIn(DAILY_SHEET, '\n'.join(audit_source_retention(book, raw)))

    def test_retention_failure_aborts_refresh_before_publishing(self):
        saved = []
        def seed_previous_output(path):
            output = path.parent/'汇总.xlsx'
            book = Workbook()
            book.active['A1'] = '上次可用汇总'
            book.save(output)
            book.close()
            saved.append(output)
        with patch('tools.refresh_sales_forecast_data.audit_source_retention', return_value=['模拟检测到落值遗漏']):
            with self.assertRaisesRegex(ValueError, '已阻止覆盖正式文件'):
                self.build(source_edit=seed_previous_output)
        book = load_workbook(saved[0], data_only=True)
        try:
            self.assertEqual(book.active['A1'].value, '上次可用汇总')
        finally:
            book.close()

    def test_cancellation_header_is_located_by_fields_not_fixed_row_or_column(self):
        def add_sources(folder):
            path = folder/'鸿蒙智行小订退订分析.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.title = self.model+'_日度退订'
            sheet.append(['当日小订退', '累计小订退', '取消日期'])
            for i, cumulative in enumerate((1, 3, 6), 1):
                sheet.append([i, cumulative, datetime(2026, 9, i)])
            book.save(path)
            book.close()
        self.assert_complete(self.build(add_sources=add_sources))

    def test_missing_end_and_duration_does_not_invent_a_35_day_window(self):
        book = self.build(maintained={'首销结束日期': None, '首销天数': None})
        master = table_records(book, MASTER_SHEET)[0]
        self.assertIsNone(master['首销结束'])
        self.assertFalse(master['首销参考终值有效'])
        self.assertEqual(table_records(book, '小转大当日数量')[0]['D3'], 10)
        history, *_ = self.web_inputs(book, date(2026, 9, 5))
        self.assertIsNone(history[0]['conversion'])

    def test_missing_gross_does_not_erase_independent_components(self):
        book = self.build(launch_quantities={'当日大定数量': [None, None, None]})
        self.assert_complete(book)

    def test_count_parser_distinguishes_zero_text_and_invalid_values(self):
        for value, expected in [('0', 0), (' 1,234 ', 1234), ('1，234.0', 1234),
                                ('+12', 12), (0, 0), ('', None), (None, None),
                                ('12%', None), ('1,2', None), ('12单', None),
                                ('-1', None), (True, None), (float('inf'), None)]:
            with self.subTest(value=value):
                self.assertEqual(_optional_count([value], 0), expected)

    def test_multiple_files_fill_blank_cells_and_keep_existing_values(self):
        def sources(folder):
            book = load_workbook(folder/'鸿蒙智行首销期订单节奏.xlsx')
            for sheet in book:
                sheet['B4'] = 10  # fill the original D1 blank from the second physical file
                sheet['C4'] = 99  # conflicting D2 small: original 20 must win
            book.save(folder/'鸿蒙智行首销期订单节奏_补充.xlsx')
            book.close()
        self.assert_complete(self.build(add_sources=sources,
            launch_quantities={'当日小订转大定数量': [None, 20, 10]}))

    def test_mix_fallback_and_steady_week_keep_text_counts_and_future_observations(self):
        def sources(folder):
            for filename, metrics in [('大定', [('大定', 80), ('留存大定', 60)]),
                                      ('锁单', [('交车锁单', 25)])]:
                book = Workbook()
                sheet = book.active
                sheet.title = self.model+'by天'
                sheet.append(['指标', '统计类型', '分类', *[datetime(2026, 9, i) for i in range(1, 15)]])
                for label, value in metrics:
                    sheet.append([label, '数量', '数量', *[str(value)]*14])
                sheet = book.create_sheet(self.model+'by周')
                sheet.append(['指标', '统计类型', '分类', '26WK37', '26WK38', '26WK39'])
                for label, value in metrics:
                    sheet.append([label, '数量', '数量', str(value*7), str(value*7+1), str(value*7+2)])
                book.save(folder/f'鸿蒙智行{filename}选配比例分析.xlsx')
                book.close()
        book = self.build(as_of='2026-09-14', add_sources=sources,
                          launch_quantities={'当日交车锁单数量': [None, None, None]})
        self.assert_complete(book)  # lower-priority gross 80 must not replace launch values
        master = table_records(book, MASTER_SHEET)[0]
        self.assertEqual((master['首销期锁单'], master['首销期留存大定']), (75, 180))
        weeks = {r['周期']: r for r in table_records(book, WEEKLY_SHEET)}
        self.assertEqual((weeks['26WK37']['大定'], weeks['26WK37']['交车锁单']), (560, 175))
        self.assertEqual(weeks['26WK38']['交车锁单'], 176)
        self.assertEqual(weeks['26WK39']['交车锁单'], 177)
        steady = self.web_inputs(book, date(2026, 9, 14))[4]
        self.assertEqual([r['period'] for r in steady[0]['weeks']], ['26WK37'])
        self.assertEqual(len({r['日期'] for r in table_records(book, DAILY_SHEET)}), len(table_records(book, DAILY_SHEET)))

    def test_quantity_only_model_is_saved_with_unknown_dates_and_not_a_final_reference(self):
        name = '独立历史车2024款'
        def source_edit(path):
            book = load_workbook(path)
            book['小转大进度'].append([name, '1,000', '20', None])
            book['直接大定进度'].append([name, 30, None, 0])
            book.save(path)
            book.close()
        book = self.build(source_edit=source_edit)
        master = next(r for r in table_records(book, MASTER_SHEET) if r['历史传播名'] == name)
        self.assertFalse(master['首销参考终值有效'])
        self.assertIsNone(master['首销开始'])
        row = next(r for r in table_records(book, '小转大当日数量') if r['传播名'] == name)
        self.assertEqual((row['D1'], row['D2']), (1000, 20))
        item = next(r for r in self.web_inputs(book, date(2026, 9, 5))[0] if r['model'] == name)
        self.assertFalse(item['d1_valid'])
        self.assertIsNone(item['conversion'])

    def test_no_small_model_preserves_direct_launch_actuals(self):
        def source_edit(path):
            book = load_workbook(path)
            headers = [c.value for c in book['车型汇总'][1]]
            for key in ('小订开始日期', '小订结束日期', '总小订'):
                book['车型汇总'].cell(2, headers.index(key)+1).value = None
            del book['小订by天']
            book.save(path)
            book.close()
        book = self.build(source_edit=source_edit, launch_quantities={
            '当日小订转大定数量': [None]*3, '累计小订转大定数量': [None]*3,
            '当日直接大定数量': [None]*3, '累计直接大定数量': [None]*3})
        master = table_records(book, MASTER_SHEET)[0]
        self.assertFalse(master['有小订'])
        self.assertEqual(master['总直接大定'], 130)
        self.assertEqual(master['直接大定占比'], 1)
        self.assertIsNone(master['小订转化率'])

    def test_missing_denominator_is_filled_before_computing_rates_not_defaulted_to_zero(self):
        self.assert_complete(self.build(maintained={'总小订': None, '总小转大': 40, '总直接大定': 90}))

    def test_secondary_generation_keeps_its_own_source_facts(self):
        self.model = '问界 M7 2024款 Ultra版'
        book = self.build(maintained={'代际名': '问界 M7 2024款', '二级代际名': self.model})
        self.assert_complete(book)
        master = table_records(book, MASTER_SHEET)[0]
        self.assertEqual(master['primary_generation'], '问界 M7 2024款')
        self.assertEqual(master['secondary_generation'], self.model)
        self.assertTrue(all(r['订单分析代际名'] == self.model for r in table_records(book, DAILY_SHEET)))
        self.assertEqual(self.web_inputs(book, date(2026, 9, 5))[0][0]['conversion'], .4)

    def test_small_mix_text_counts_fill_only_missing_history_day(self):
        def source_edit(path):
            book = load_workbook(path)
            book['小订by天']['C2'] = None
            book['小订by天']['D2'] = 1100
            book.save(path)
            book.close()
        def add_sources(folder):
            book = Workbook()
            sheet = book.active
            sheet.title = self.model+'by天'
            sheet.append(['指标', '统计类型', '分类', datetime(2026, 8, 30), datetime(2026, 8, 31)])
            sheet.append(['小订', '数量', '数量', '9,999', '1,060'])
            book.save(folder/'鸿蒙智行小订选配比例分析.xlsx')
            book.close()
        book = self.build(maintained={'总小订':1100}, source_edit=source_edit, add_sources=add_sources)
        row = table_records(book, SMALL_DAILY_SHEET)[0]
        self.assertEqual((row['D1'], row['D2']), (40, 1060))
        self.assertEqual(table_records(book, SMALL_CURVE_SHEET)[0]['D2'], 1)
        self.assertNotIn('.refreshing.xlsx', table_records(book, MASTER_SHEET)[0]['订单来源文件'])
