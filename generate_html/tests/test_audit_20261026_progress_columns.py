"""Explicit Dn headers keep their calendar positions in forecast references."""
from datetime import datetime
from pathlib import Path
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook

from core.forecast_summary import reservation_reference_shape
from modules import sales_forecast as forecast


class ProgressColumnCalendarTests(TestCase):
    def book(self, headers, values):
        book = Workbook()
        sheet = book.active
        sheet.title = "progress"
        sheet.append(["model", *headers])
        sheet.append(["test", *values])
        self.addCleanup(book.close)
        return book, sheet

    def read(self, headers, values):
        book, sheet = self.book(headers, values)
        return forecast._progress_rows(book, sheet.title)["test"]

    def test_missing_physical_columns_keep_unknown_day_slots(self):
        self.assertEqual(self.read(["D1", "D3", "D5"], [.2, .6, 1]),
                         [.2, None, .6, None, 1])

    def test_leading_missing_days_do_not_move_first_observation_to_d1(self):
        self.assertEqual(self.read(["D3", "D5"], [.6, 1]),
                         [None, None, .6, None, 1])

    def test_unordered_columns_follow_explicit_day_numbers(self):
        self.assertEqual(self.read(["D3", "D1", "D2"], [1, .2, .6]),
                         [.2, .6, 1])

    def test_sparse_unordered_columns_keep_all_calendar_gaps(self):
        self.assertEqual(self.read(["D5", "D1", "D3"], [1, .2, .6]),
                         [.2, None, .6, None, 1])

    def test_dense_curves_and_real_zero_are_unchanged(self):
        self.assertEqual(self.read(["D1", "D2", "D3"], [0, .6, 1]),
                         [0, .6, 1])
        self.assertEqual(self.read(["D1", "D2", "D3"], [.2, None, 1]),
                         [.2, None, 1])

    def test_sparse_curve_cannot_synthesize_missing_daily_orders(self):
        curve = self.read(["D1", "D3", "D5"], [.2, .6, 1])
        shape = reservation_reference_shape(curve, 100)
        self.assertEqual(shape["daily_orders"], [])
        self.assertEqual(shape["standard_progress"], [.2, None, .6, None, 1])

    def small_history(self, headers, values):
        book = Workbook()
        self.addCleanup(book.close)
        book.active.title = "\u5c0f\u8ba2by\u5929"
        book.active.append(["\u5c0f\u8ba2"])
        summary = book.create_sheet("\u8f66\u578b\u6c47\u603b")
        summary.append(["\u8f66\u578b", "\u5c0f\u8ba2\u5f00\u59cb\u65e5\u671f",
                        "\u5c0f\u8ba2\u7ed3\u675f\u65e5\u671f", "\u603b\u5c0f\u8ba2", "\u5c0f\u8ba2\u5929\u6570"])
        summary.append(["test", datetime(2026, 9, 1), datetime(2026, 9, 5), 100, 5])
        progress = book.create_sheet("\u5c0f\u8ba2\u8fdb\u5ea6")
        progress.append(["\u8f66\u578b", *headers])
        progress.append(["test", *values])
        with patch.object(Path, "exists", return_value=True), \
             patch.object(forecast, "open_input", return_value=book), \
             patch.object(forecast, "_read_model_mapping", return_value={"test": "test"}), \
             patch.object(forecast, "_read_model_master", return_value={}):
            _, history = forecast._read_small_order_history(Path("memory.xlsx"))
        self.assertEqual(len(history), 1)
        return history[0]

    def test_full_small_history_reader_retains_sparse_reference_positions(self):
        item = self.small_history(["D1", "D3", "D5"], [.2, .6, 1])
        self.assertEqual(item["standard_progress"], [.2, None, .6, None, 1])
        self.assertEqual(item["daily_orders"], [])
        self.assertEqual(item["dates"], [])
        self.assertEqual(item["total"], 100)
        self.assertEqual(item["days"], 5)
        self.assertFalse(item["daily_actual"])

    def test_duplicate_day_columns_leave_that_day_unknown(self):
        self.assertEqual(self.read(["D1", "D1", "D2"], [.2, .8, 1]),
                         [None, 1])

    def test_alias_duplicate_day_columns_leave_that_day_unknown(self):
        self.assertEqual(self.read(["D1", "D01", "D2"], [.2, .8, 1]),
                         [None, 1])

    def test_duplicate_days_cannot_create_daily_orders_in_full_reader(self):
        for headers in (["D1", "D1", "D2"], ["D1", "D01", "D2"]):
            with self.subTest(headers=headers):
                item = self.small_history(headers, [.2, .8, 1])
                self.assertEqual(item["standard_progress"], [None, 1])
                self.assertEqual(item["daily_orders"], [])
                self.assertEqual(item["dates"], [])
                self.assertEqual(item["total"], 100)
                self.assertFalse(item["daily_actual"])
