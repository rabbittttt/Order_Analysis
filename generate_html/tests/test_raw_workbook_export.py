"""Whole-workbook downloads retain the source, independently of display blocks."""
import base64
import json
from pathlib import Path
import shutil
import subprocess
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest import TestCase, skipUnless

from openpyxl import Workbook, load_workbook
from main import raw_tables

ROOT = Path(__file__).resolve().parents[1]


class RawWorkbookSnapshotTests(TestCase):
    def test_snapshot_preserves_all_sheets_and_original_bytes(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / '中文原表.xlsx'
            book = Workbook()
            self.addCleanup(book.close)
            sheet = book.active
            sheet.title = '并排区块'
            sheet.append(['销量', None, None, '累计占比', None])
            sheet.append(['配置', '数量', None, '配置', '占比'])
            sheet.append(['黑色', 10, None, '黑色', .25])
            sheet['E3'].number_format = '0.0%'
            detail = book.create_sheet('明细')
            detail.append(['公式', '=SUM(1,2)'])
            detail.merge_cells('A3:B3')
            hidden = book.create_sheet('隐藏辅助表')
            hidden.append(['保留隐藏内容', 123])
            hidden.sheet_state = 'hidden'
            book.save(path)

            files, blocks = raw_tables(SimpleNamespace(items=[SimpleNamespace(path=path, workbook=book)]))
            self.assertEqual([s['source_sheet'] for s in files[0]['sheets']], ['并排区块', '并排区块', '明细'])
            exported = base64.b64decode(blocks[files[0]['workbook_key']])
            self.assertEqual(exported, path.read_bytes())
            copy = Path(folder) / '下载.xlsx'
            copy.write_bytes(exported)
            loaded = load_workbook(copy, data_only=False)
            self.addCleanup(loaded.close)
            self.assertEqual(loaded.sheetnames, ['并排区块', '明细', '隐藏辅助表'])
            self.assertEqual(loaded['明细']['B1'].value, '=SUM(1,2)')
            self.assertEqual(loaded['并排区块']['E3'].number_format, '0.0%')
            self.assertEqual(loaded['隐藏辅助表'].sheet_state, 'hidden')

    def test_unreadable_source_does_not_break_the_raw_table(self):
        with TemporaryDirectory() as folder:
            book = Workbook()
            self.addCleanup(book.close)
            book.active.append(['名称', '数量'])
            book.active.append(['测试', 0])
            with self.assertLogs('order_analysis', level='WARNING') as logs:
                files, blocks = raw_tables(SimpleNamespace(items=[SimpleNamespace(
                    path=Path(folder) / 'missing.xlsx', workbook=book)]))
            self.assertNotIn('workbook_key', files[0])
            self.assertIn('0|0', blocks)
            self.assertIn('整表导出不可用', logs.output[0])


@skipUnless(shutil.which('node'), 'Node.js required')
class RawWorkbookDownloadTests(TestCase):
    def run_download(self, invalid=False):
        source = (ROOT / 'templates/dashboard.js').read_text(encoding='utf-8')
        begin = source.index('  function exportRaw(')
        end = source.index('  async function jumpToSource', begin)
        script = r'''
const exportsSource=JSON.parse(process.argv[1]),invalid=process.argv[2]==='true';
const original=Buffer.from(Array.from({length:170003},(_,i)=>i%251)),downloads=[],messages=[],blobs=new Map();
const files={original:{name:'中文原表.xlsx',workbook_key:'0|xlsx'},other:{name:'其他.xlsx',workbook_key:'1|xlsx'}};
const blocks={'0|xlsx':invalid?'?bad-base64?':original.toString('base64'),'1|xlsx':Buffer.from('wrong file').toString('base64')};
const state={rawQuery:'匹配',rawFile:0},button={disabled:false,textContent:'导出整个 Excel'};
const doc={body:{appendChild(){}},createElement:()=>({click(){downloads.push({name:this.download,blob:blobs.get(this.href)})},remove(){}})};
const urls={createObjectURL:blob=>{const url='blob:'+blobs.size;blobs.set(url,blob);return url},revokeObjectURL(){}};
const api=new Function('RAW_BLOCKS','document','URL','showToast','state','rawRowsCache','visibleRawRows',exportsSource+
 ';return {exportRawWorkbook,exportRaw};')(blocks,doc,urls,text=>messages.push(text),state,
 [['标题','数量'],['匹配',0],['不显示',9]],(rows,query)=>[rows[0],...rows.slice(1).filter(r=>r.includes(query))]);
(async()=>{
 const pending=api.exportRawWorkbook(files.original,button);
 const loading=button.disabled&&button.textContent==='正在导出…';
 state.rawFile=1;state.rawQuery='无匹配';await pending;
 const workbook=downloads[0];
 const same=workbook?Buffer.from(await workbook.blob.arrayBuffer()).equals(original):false;
 state.rawQuery='匹配';api.exportRaw(files.other,{name:'当前页'});
 const csv=await downloads.at(-1).blob.text();
 console.log(JSON.stringify({same,loading,button,messages,workbookName:workbook?.name,
 workbookType:workbook?.blob.type,csv}));
})().catch(error=>{console.error(error);process.exitCode=1});
'''
        result = subprocess.run(['node', '-e', script, json.dumps(source[begin:end]), json.dumps(invalid)],
                                capture_output=True, text=True, encoding='utf-8', cwd=ROOT)
        self.assertEqual(result.returncode, 0, result.stderr)
        return json.loads(result.stdout)

    def test_chunked_download_keeps_source_after_search_and_file_switch(self):
        result = self.run_download()
        self.assertTrue(result['same'])
        self.assertTrue(result['loading'])
        self.assertEqual(result['workbookName'], '中文原表.xlsx')
        self.assertEqual(result['workbookType'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertEqual(result['button'], {'disabled': False, 'textContent': '导出整个 Excel'})
        self.assertIn('匹配,0', result['csv'])
        self.assertNotIn('不显示', result['csv'])

    def test_invalid_snapshot_restores_button_and_explains_failure(self):
        result = self.run_download(invalid=True)
        self.assertFalse(result['same'])
        self.assertFalse(result['button']['disabled'])
        self.assertIn('整表导出失败', result['messages'][0])
