from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook

from core.excel import WorkbookItem
from core.forecast_summary import (MASTER_SHEET, DAILY_SHEET, D12_SHEET, public_forecast_tables,
                                   table_records, restrict_launch_reference)
from modules.sales_forecast import _read_actual_profiles, _history_item
from tools.refresh_sales_forecast_data import write_rows


class ForecastReferenceBackfillTests(TestCase):
    def build(self, *, daily=None, cumulative=None, small_total=20, as_of=date(2026, 1, 6),
              history=None, cancel=None, cumulative_start='2026-01-03', maintained=None, raw_fields=None):
        book = Workbook()
        self.addCleanup(book.close)
        book.active.title = MASTER_SHEET
        fields = ['历史传播名', '订单分析代际名', '首销开始', '首销结束', '首销天数',
                  '小订开始', '小订结束', '总小订', '总小转大', '小订转化率', '总直接大定',
                  '直接大定占比', '总大定', '总退订', '退订率', '首销期留存大定', '首销期锁单',
                  '留存大定率', '留存大定到锁单率', '大定到锁单率']
        row = dict(zip(fields[:8], ['测试车', '测试车', '2026-01-03', '2026-01-05', 3,
                                   '2026-01-01', '2026-01-02', small_total]))
        row.update(maintained or {})
        book.active.append(fields)
        book.active.append([row.get(k) for k in fields])
        raw = {'model': '测试车', 'days': daily or [], 'cumulative_days': cumulative or [],
               'cumulative_launch_date': cumulative_start,
               'day_source': {'file': '首销期订单节奏.xlsx', 'sheet': '测试车by天'},
               'small_daily_days': [{'date': '2026-01-01', 'orders': 8},
                                    {'date': '2026-01-02', 'orders': 12}],
               'small_daily_sources': {d: {'file': '小订选配比例分析.xlsx', 'sheet': '测试车by天'}
                                       for d in ('2026-01-01', '2026-01-02')},
               'cancel_days': cancel or [],
               'cancel_source': {'file': '小订退订分析.xlsx', 'sheet': '测试车日度退订'}}
        raw.update(raw_fields or {})
        target = dict(name='测试车', history_model='测试车', launch_date='2026-01-03',
                      end_date='2026-01-05', days=3, small_start_date='2026-01-01',
                      small_end_date='2026-01-02')
        data = {'targets': [target], 'actuals': [raw],
                'history': [{'model': '测试车', 'generation': '测试车', 'launch_date': '2026-01-03',
                             'days': 3, 'daily_orders': []}],
                'reference_daily': history or {}}

        def emit(name, headers, rows, percent_headers=None):
            if name in book.sheetnames:
                del book[name]
            write_rows(book, name, headers, rows, 'BackfillTest', percent_headers)

        public_forecast_tables(book, data, emit, as_of_date=as_of, raw_profiles=[raw])
        return book

    def record(self, book, name=MASTER_SHEET):
        return table_records(book, name)[0]

    def quantities(self):
        return [dict(date=f'2026-01-0{i+3}', day=f'D{i+1}', small_to_big=s, direct=d, gross=s+d)
                for i, (s, d) in enumerate(((2, 6), (3, 9), (5, 15)))]

    def anchors(self):
        return [dict(date=f'2026-01-0{i+3}', small_to_big=s, direct=d, gross=s+d)
                for i, (s, d) in enumerate(((2, 6), (5, 15), (10, 30)))]

    def test_resolved_small_denominator_precedes_all_dependent_features(self):
        book = self.build(daily=self.quantities(), small_total=None,
                          cancel=[{'date': f'2026-01-0{i+3}', 'daily_cancel': n, 'cancel': c}
                                  for i, (n, c) in enumerate(((1, 1), (1, 2), (3, 5)))])
        row = self.record(book)
        self.assertEqual((row['总小订'], row['总小转大'], row['总直接大定'], row['总退订']), (20, 10, 30, 5))
        self.assertEqual((row['小订转化率'], row['直接大定占比'], row['退订率']), (.5, .75, .25))
        self.assertEqual(self.record(book, '累计小订转化率')['D1'], .1)
        self.assertEqual(self.record(book, '累计退订率')['D3'], .25)
        self.assertEqual(self.record(book, D12_SHEET)['D1退订率'], .05)

    def test_g9_available_totals_recalculate_rates_despite_missing_final_breakdown(self):
        book = self.build(daily=[
            dict(date='2026-01-03', small_to_big=2412, direct=140, gross=2552, net=2318, lock=2263),
            dict(date='2026-01-04', small_to_big=4886, direct=3497, gross=8383, net=7569, lock=7394),
            dict(date='2026-01-05', gross=378, net=313, lock=211),
        ], maintained={'留存大定率': 0, '留存大定到锁单率': '#DIV/0!', '大定到锁单率': 0})
        row = self.record(book)
        self.assertEqual((row['总大定'], row['首销期留存大定'], row['首销期锁单']), (11313, 10200, 9868))
        for field, expected in (('留存大定率', 10200/11313),
                                ('留存大定到锁单率', 9868/10200), ('大定到锁单率', 9868/11313)):
            self.assertAlmostEqual(row[field], expected)
        self.assertIsNone(row['总小转大'])
        self.assertIsNone(row['总直接大定'])
        item = _history_item(row, True)
        self.assertAlmostEqual(item['net_rate'], 10200/11313)
        self.assertAlmostEqual(item['lock_rate'], 9868/11313)

    def test_all_derived_rates_follow_final_quantities_not_stale_maintained_ratios(self):
        expected = {'小订转化率': .5, '直接大定占比': .75, '退订率': .25,
                    '留存大定率': .9, '留存大定到锁单率': 32/36, '大定到锁单率': .8}
        daily = [dict(r, net=r['gross']*.9, lock=r['gross']*.8) for r in self.quantities()]
        cancel = [dict(date=f'2026-01-0{i+3}', daily_cancel=n, cancel=c)
                  for i, (n, c) in enumerate(((1, 1), (1, 2), (3, 5)))]
        for old in (0, '#DIV/0!', .123):
            with self.subTest(old=old):
                book = self.build(daily=daily, cancel=cancel, maintained=dict.fromkeys(expected, old))
                row = self.record(book)
                for field, value in expected.items():
                    self.assertAlmostEqual(row[field], value, msg=field)
                self.assertEqual(row['总大定'], 40)
                self.assertEqual(row['总小转大'], 10)
                self.assertEqual(row['总直接大定'], 30)

    def test_missing_quantities_preserve_only_valid_maintained_ratios(self):
        fields = ('小订转化率', '直接大定占比', '退订率', '留存大定率',
                  '留存大定到锁单率', '大定到锁单率')
        for value in (0, .6, 1, None, '', '#DIV/0!', '#N/A', -1, 1.2, True, False,
                      float('nan'), float('inf')):
            with self.subTest(value=value):
                book = self.build(maintained=dict.fromkeys(fields, value))
                row = self.record(book)
                expected = value if type(value) in (int, float) and 0 <= value <= 1 else None
                for field in fields:
                    self.assertEqual(row[field], expected, field)

    def test_real_zero_numerators_recalculate_zero_rates(self):
        book = self.build(daily=[dict(r, net=0, lock=0) for r in self.quantities()],
                          maintained={'留存大定率': .9, '大定到锁单率': .8,
                                      '留存大定到锁单率': .6})
        row = self.record(book)
        self.assertEqual(row['留存大定率'], 0)
        self.assertEqual(row['大定到锁单率'], 0)
        self.assertEqual(row['留存大定到锁单率'], .6)  # 0/0 cannot replace an independent maintained rate.

    def test_zero_denominator_does_not_invent_a_rate(self):
        book = self.build(maintained={'总大定': 0, '首销期留存大定': 0, '首销期锁单': 0,
                                      '留存大定率': '#DIV/0!', '大定到锁单率': None,
                                      '留存大定到锁单率': '#DIV/0!'})
        row = self.record(book)
        for field in ('留存大定率', '大定到锁单率', '留存大定到锁单率'):
            self.assertIsNone(row[field], field)

    def test_invalid_quantity_cannot_be_coerced_to_zero_to_overwrite_valid_rate(self):
        for value in ('#DIV/0!', -1, True, float('nan'), float('inf')):
            with self.subTest(value=value):
                book = self.build(maintained={'总大定': 40, '首销期留存大定': value, '留存大定率': .9})
                self.assertEqual(self.record(book)['留存大定率'], .9)

    def test_recalculation_uses_selected_maintained_totals_without_changing_priority(self):
        book = self.build(daily=[dict(r, net=r['gross']*.9, lock=r['gross']*.8)
                                 for r in self.quantities()], maintained={
            '总大定': 100, '首销期留存大定': 80, '首销期锁单': 60,
            '留存大定率': 0, '留存大定到锁单率': 0, '大定到锁单率': 0})
        row = self.record(book)
        self.assertEqual((row['总大定'], row['首销期留存大定'], row['首销期锁单']), (100, 80, 60))
        self.assertEqual((row['留存大定率'], row['留存大定到锁单率'], row['大定到锁单率']), (.8, .75, .6))

    def test_partial_history_fills_each_field_from_real_launch_days(self):
        book = self.build(daily=self.quantities(), history={
            '小转大当日数量': {'测试车': [0, None, None]},
            '直接大定当日数量': {'测试车': [8, None, None]},
            '总大定当日数量': {'测试车': [8, None, None]}})
        row = self.record(book)
        self.assertEqual((row['总小转大'], row['总直接大定'], row['总大定']), (8, 32, 40))
        self.assertEqual(self.record(book, D12_SHEET)['D2小转大'], 3)
        self.assertEqual(self.record(book, '小转大累计完成度')['D3'], 1)
        self.assertEqual(self.record(book, '直接大定累计完成度')['D3'], 1)

    def test_real_cumulative_quantities_backfill_only_adjacent_daily_gaps(self):
        book = self.build(cumulative=self.anchors())
        row = self.record(book)
        self.assertEqual((row['总小转大'], row['总直接大定'], row['总大定']), (10, 30, 40))
        self.assertEqual(self.record(book, D12_SHEET)['D2小转大'], 3)
        self.assertEqual(self.record(book, '小转大当日数量')['D3'], 5)
        self.assertEqual(self.record(book, '小转大累计完成度')['D2'], .5)
        self.assertTrue(any('累计差分' in str(r.get('字段来源')) for r in table_records(book, DAILY_SHEET)))

    def test_cumulative_across_a_gap_fills_progress_not_invented_daily_values(self):
        book = self.build(cumulative=[self.anchors()[0], self.anchors()[2]])
        self.assertEqual(self.record(book)['总小转大'], 10)
        curve = self.record(book, '小转大累计完成度')
        self.assertEqual((curve['D1'], curve['D2'], curve['D3']), (.2, None, 1))
        daily = self.record(book, '小转大当日数量')
        self.assertEqual((daily['D1'], daily['D2'], daily['D3']), (2, None, None))
        self.assertIsNone(self.record(book, D12_SHEET)['D1+D2小转大'])

    def test_ongoing_stage_saves_observed_snapshot_but_excludes_terminal_reference(self):
        book = self.build(cumulative=self.anchors(), as_of=date(2026, 1, 5))
        self.assertEqual(self.record(book)['总小转大'], 10)
        self.assertEqual(self.record(book)['小订转化率'], .5)
        self.assertEqual(self.record(book, D12_SHEET)['D2小转大'], 3)
        self.assertEqual(self.record(book, D12_SHEET)['D1小转大/总小转大'], .2)
        self.assertEqual(self.record(book, '累计小订转化率')['D2'], .25)
        self.assertEqual(self.record(book, '累计小订转化率')['D3'], .5)
        item = _history_item(self.record(book, D12_SHEET), True)
        restrict_launch_reference(item, self.record(book), date(2026, 1, 5))
        self.assertIsNone(item['small_to_big'])
        self.assertIsNone(item['d1_small_completion'])
        self.assertEqual(item['small_d2_d1'], 1.5)

    def test_wrong_campaign_cumulative_is_not_borrowed(self):
        book = self.build(cumulative=self.anchors(), cumulative_start='2025-01-03')
        self.assertIsNone(self.record(book)['总小转大'])
        self.assertIsNone(self.record(book, '小转大当日数量')['D1'])

    def test_missing_unrelated_component_does_not_blank_independent_d12_ratios(self):
        book = self.build(cumulative=[{k: v for k, v in r.items() if k in ('date', 'small_to_big')}
                                     for r in self.anchors()])
        metric = self.record(book, D12_SHEET)
        self.assertEqual(metric['D1小转大/总小转大'], .2)
        self.assertEqual(metric['D2小转大/D1小转大'], 1.5)
        self.assertIsNone(metric['D1直接大/D1大定'])
        self.assertEqual(metric['D1口径状态'], '部分字段有效')
        item = _history_item(metric, True)
        self.assertTrue(item['d1_valid'])
        self.assertEqual(item['d1_small_completion'], .2)
        self.assertEqual(item['small_d2_d1'], 1.5)
        for field in ('d1_small_share', 'd1_direct_share', 'd12_small_share', 'd12_direct_share'):
            self.assertIsNone(item[field])

    def test_partial_structure_conflict_is_not_mistaken_for_partial_eligibility(self):
        book = self.build(daily=[{'date': '2026-01-03', 'small_to_big': 9, 'gross': 4}],
                          maintained={'总小转大': 10})
        metric = self.record(book, D12_SHEET)
        self.assertEqual(metric['D1小转大/总小转大'], .9)
        self.assertEqual(metric['D1口径状态'], '数据不足或异常')
        self.assertFalse(_history_item(metric, True)['d1_valid'])
        self.assertEqual(_history_item(metric, True)['d1_small_completion'], 0)

    def test_real_zero_cumulative_is_preserved_and_never_divided_by_zero(self):
        book = self.build(cumulative=[dict(date=r['date'], small_to_big=0) for r in self.anchors()])
        self.assertEqual(self.record(book)['总小转大'], 0)
        self.assertEqual(self.record(book)['小订转化率'], 0)
        self.assertEqual(self.record(book, D12_SHEET)['D2小转大'], 0)
        self.assertIsNone(self.record(book, '小转大累计完成度')['D2'])

    def test_negative_nonfinite_boolean_and_decreasing_cumulatives_do_not_fill(self):
        for values in ((-1, -2, -3), (float('nan'), float('inf'), True), (2, 1, 5)):
            with self.subTest(values=values):
                book = self.build(cumulative=[dict(date=r['date'], small_to_big=v)
                                             for r, v in zip(self.anchors(), values)])
                self.assertIsNone(self.record(book)['总小转大'])
                self.assertIsNone(self.record(book, '小转大当日数量')['D1'])

    def test_mismatched_cancel_dates_do_not_fill_historical_terminal(self):
        book = self.build(daily=self.quantities(), cancel=[{'date': '2026-02-03', 'cancel': 6, 'daily_cancel': 6}])
        self.assertIsNone(self.record(book)['总退订'])
        self.assertIsNone(self.record(book)['退订率'])

    def test_no_borrowing_cumulative_after_stage_end(self):
        book = self.build(cumulative=[self.anchors()[0], dict(date='2026-01-06', small_to_big=10)])
        self.assertIsNone(self.record(book)['总小转大'])

    def test_history_conflict_does_not_let_lower_priority_cumulative_set_terminal(self):
        book = self.build(cumulative=[self.anchors()[0], self.anchors()[2]],
                          history={'小转大当日数量': {'测试车': [0, None, None]}})
        self.assertEqual(self.record(book, '小转大当日数量')['D1'], 0)
        self.assertIsNone(self.record(book)['总小转大'])

    def test_maintained_terminal_is_preserved_if_daily_incomplete(self):
        book = self.build(cumulative=[self.anchors()[2]], maintained={'总小转大': 12})
        self.assertEqual(self.record(book)['总小转大'], 12)

    def test_cancel_cumulative_keeps_prelaunch_baseline_and_recovers_after_gap(self):
        book = self.build(daily=self.quantities(), cancel=[
            {'date': '2026-01-02', 'cancel': 2},
            {'date': '2026-01-03', 'cancel': 3, 'daily_cancel': 1},
            {'date': '2026-01-04', 'cancel': None, 'daily_cancel': 1},
            {'date': '2026-01-05', 'cancel': None, 'daily_cancel': 1}])
        self.assertEqual(self.record(book)['总退订'], 5)
        self.assertEqual(self.record(book, '累计退订率')['D2'], .2)
        self.assertEqual(self.record(book, '累计退订率')['D3'], .25)

    def test_parser_preserves_real_cumulative_with_missing_daily_rows(self):
        book = Workbook()
        self.addCleanup(book.close)
        sheet = book.active
        sheet.title = '测试车2026款by天'
        sheet.append(['周期', 'D1', 'D2', 'D3'])
        sheet.append(['时间', datetime(2026, 1, 3), datetime(2026, 1, 4), datetime(2026, 1, 5)])
        sheet.append(['累计小订转大定数量', 2, 5, 10])
        sheet.append(['累计直接大定数量', 6, 15, 30])
        sheet.append(['累计交车锁单数量', 0, 3, 6])
        sheet.append(['当日大定数量', 8, 12, 20])
        item = WorkbookItem(Path('首销期订单节奏.xlsx'), book)
        store = SimpleNamespace(find=lambda keyword: item if keyword == '首销期订单节奏' else None)
        with patch('modules.sales_forecast._read_model_mapping', return_value={}):
            profiles, _ = _read_actual_profiles(store)
        self.assertEqual(profiles[0]['cumulative_launch_date'], '2026-01-03')
        self.assertEqual([r['small_to_big'] for r in profiles[0]['cumulative_days']], [2, 5, 10])
        self.assertEqual([r['lock'] for r in profiles[0]['cumulative_days']], [0, 3, 6])
        summary = self.build(raw_fields={**profiles[0], 'model': '测试车'})
        master = self.record(summary)
        self.assertEqual((master['总小转大'], master['总直接大定'], master['直接大定占比']), (10, 30, .75))
        self.assertEqual(master['首销期锁单'], 6)
        self.assertEqual(self.record(summary, D12_SHEET)['D1小转大/总小转大'], .2)
        self.assertEqual(self.record(summary, '直接大定累计完成度')['D3'], 1)
