from __future__ import annotations

import gzip
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from core.excel import load_data_workbook


class OptionalWorkbookCacheTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.path, self.cache = self.root / 'source.xlsx', self.root / 'cache'
        book = Workbook()
        book.active.append([123, 0, .3, 'source'])
        book.active['C1'].number_format = '0.0%'
        book.save(self.path)
        book.close()

    def assert_source(self):
        book = load_data_workbook(self.path, self.cache)
        try:
            self.assertEqual([c.value for c in book.active[1]], [123, 0, .3, 'source'])
            self.assertEqual(book.active['C1'].number_format, '0.0%')
        finally:
            book.close()

    def writer(self, fail_write=None, fail_close=False):
        original = gzip.open

        class FaultyWriter:
            def __init__(self, *args, **kwargs):
                self.stream = original(*args, **kwargs)
                self.writes = 0

            def write(self, text):
                self.writes += 1
                if self.writes == fail_write:
                    raise OSError('cache write failed')
                return self.stream.write(text)

            def close(self):
                self.stream.close()
                if fail_close:
                    raise OSError('cache flush failed')

        return FaultyWriter

    def test_flush_failure_keeps_source_and_never_commits_cache(self):
        with patch('core.excel.gzip.open', side_effect=self.writer(fail_close=True)):
            self.assert_source()
        self.assertFalse(list(self.cache.glob('*.jsonl.gz')))
        self.assert_source()

    def test_header_write_failure_never_commits_truncated_cache(self):
        with patch('core.excel.gzip.open', side_effect=self.writer(fail_write=1)):
            self.assert_source()
        self.assertFalse(list(self.cache.glob('*.jsonl.gz')))

    def test_record_write_and_close_failures_keep_source(self):
        with patch('core.excel.gzip.open', side_effect=self.writer(fail_write=3, fail_close=True)):
            self.assert_source()
        self.assertFalse(list(self.cache.glob('*.jsonl.gz')))

    def test_commit_and_cleanup_failures_keep_source(self):
        with patch('core.excel.os.replace', side_effect=PermissionError('cache locked')), \
             patch.object(Path, 'unlink', side_effect=PermissionError('temp cache locked')):
            self.assert_source()
        self.assertFalse(list(self.cache.glob('*.jsonl.gz')))

    def test_unavailable_post_read_stat_only_disables_cache_commit(self):
        original = Path.stat
        calls = 0

        def unavailable(path, *args, **kwargs):
            nonlocal calls
            if path == self.path:
                calls += 1
                if calls == 2:
                    raise PermissionError('source metadata temporarily locked')
            return original(path, *args, **kwargs)

        with patch.object(Path, 'stat', unavailable):
            self.assert_source()
        self.assertFalse(list(self.cache.glob('*.jsonl.gz')))

    def test_unreadable_cache_does_not_hide_readable_source(self):
        original = Path.exists

        def unavailable(path):
            if path.name.endswith('.jsonl.gz'):
                raise PermissionError('cache cannot be inspected')
            return original(path)

        with patch.object(Path, 'exists', unavailable):
            self.assert_source()

    def test_cleanup_failure_does_not_mask_real_source_error(self):
        class BrokenCells(dict):
            def values(self):
                raise ValueError('source cells corrupt')

        book = Workbook()
        book.active._cells = BrokenCells()
        with patch('core.excel.load_workbook', return_value=book), \
             patch('core.excel.gzip.open', side_effect=self.writer(fail_close=True)), \
             patch.object(Path, 'unlink', side_effect=PermissionError('cache locked')):
            with self.assertRaisesRegex(ValueError, 'source cells corrupt'):
                load_data_workbook(self.path, self.cache)

    def test_real_source_open_errors_still_propagate(self):
        with patch('core.excel.load_workbook', side_effect=ValueError('source corrupt')):
            with self.assertRaisesRegex(ValueError, 'source corrupt'):
                load_data_workbook(self.path, self.cache)


if __name__ == '__main__':
    unittest.main()
