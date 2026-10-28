"""SKU missing observations must not become zero concentration or zero activity."""
from pathlib import Path
from unittest import TestCase

from openpyxl import Workbook

from core.excel import WorkbookItem, WorkbookStore
from core.models import Subject
from modules.sku import SkuModule


class SkuUnknownObservationTests(TestCase):
    def build(self, current, cumulative=None, number_format="General"):
        book = Workbook()
        self.addCleanup(book.close)
        sheet = book.active
        sheet.title = "测试车_SKU"
        sheet.append(["SKU", "版本", "外观", "内饰", "轮毂", "26WK01", "26WK02", "总计占比",
                      "累计锁单", "26WK01", "26WK02", "总计"])
        if cumulative is None:
            cumulative = [("TOP3", current[0]), ("TOP5", current[0])]
        for index in range(max(len(current), len(cumulative))):
            values = ([f"SKU{index + 1}", "Max", "黑", "灰", "20", .4, current[index], .4]
                      if index < len(current) else [None] * 8)
            rank, share = cumulative[index] if index < len(cumulative) else (None, None)
            sheet.append([*values, rank, .6 if rank else None, share, .6 if rank else None])
            sheet.cell(index + 2, 7).number_format = number_format
            sheet.cell(index + 2, 11).number_format = number_format
        combo = book.create_sheet("测试车_组合")
        combo.append(["累计锁单", "配置", "26WK01", "26WK02", "总计"])
        combo.append(["TOP1", "Max", .6, current[0], .6])
        combo.cell(2, 4).number_format = number_format
        store = WorkbookStore(Path("."))
        store.items = [WorkbookItem(Path("鸿蒙智行SKU收敛度.xlsx"), book)]
        view = SkuModule().build(store, Subject("car", "测试车", "generation")).to_dict()["views"]["week"]
        page = view["pages"][view["default_period"]]
        kpis = {item["label"]: item for item in page["kpis"]}
        sections = {item["kind"]: item["data"] for item in page["sections"]}
        return view, kpis, sections

    def test_invalid_or_missing_rate_cells_stay_unknown_in_every_sku_block(self):
        for value in (None, "", "待更新", "#DIV/0!", True, False, "NaN", "Infinity"):
            with self.subTest(value=value):
                view, kpis, sections = self.build([value])
                self.assertEqual(view["default_period"], "26WK02")
                self.assertEqual(sections["sku_heatmap"]["rows"][0]["values"], [.6, None])
                self.assertEqual(sections["combo_matrix"]["groups"][0]["rows"][0]["values"], [.6, None])
                self.assertEqual(sections["table"]["rows"], [])
                for name in ("有效SKU", "TOP3累计", "TOP5累计"):
                    self.assertIsNone(kpis[name]["value"])
                    self.assertEqual(kpis[name]["comparison"]["status"], "unavailable")
                self.assertIsNone(sections["tail_distribution"][-1]["value"])

    def test_partial_observations_do_not_claim_complete_counts(self):
        _, kpis, sections = self.build([.03, None, 0], [("TOP3", .6), ("TOP5", .8)])
        self.assertIsNone(kpis["有效SKU"]["value"])
        self.assertIsNone(sections["tail_distribution"][-1]["value"])
        self.assertEqual([row[-1] for row in sections["table"]["rows"]], [.03, 0])
        self.assertEqual(kpis["TOP3累计"]["value"], 60)
        self.assertEqual(kpis["TOP5累计"]["value"], 80)

    def test_real_zero_is_still_a_known_zero(self):
        for value in (0, "0", "0%"):
            with self.subTest(value=value):
                _, kpis, sections = self.build([value])
                self.assertEqual(kpis["有效SKU"]["value"], 0)
                self.assertEqual(kpis["TOP3累计"]["value"], 0)
                self.assertEqual(kpis["TOP5累计"]["value"], 0)
                self.assertEqual(sections["table"]["rows"][0][-1], 0)
                self.assertEqual(sections["combo_matrix"]["groups"][0]["rows"][0]["values"][-1], 0)
                self.assertEqual(sections["tail_distribution"][-1]["value"], 0)

    def test_percentage_formats_and_known_counting_are_preserved(self):
        for value, fmt in ((.6, "0.0%"), (60, "General"), ("60%", "General"), ("60", "General")):
            with self.subTest(value=value, fmt=fmt):
                _, kpis, sections = self.build([value, .03], number_format=fmt)
                self.assertEqual(kpis["有效SKU"]["value"], 2)
                self.assertEqual(kpis["TOP3累计"]["value"], 60)
                self.assertEqual(sections["table"]["rows"][0][-1], .6)
                self.assertEqual(sections["combo_matrix"]["groups"][0]["rows"][0]["values"][-1], .6)
                self.assertEqual(sections["tail_distribution"][-1]["value"], 1)

    def test_missing_prefix_cannot_claim_a_known_80_percent_threshold(self):
        _, kpis, _ = self.build([.2], [("TOP1", None), ("TOP3", .9), ("TOP5", 1)])
        self.assertEqual(kpis["覆盖80%的SKU"]["value"], "-")

    def test_threshold_already_reached_before_missing_suffix_remains_known(self):
        _, kpis, _ = self.build([.9], [("TOP1", .9), ("TOP3", None), ("TOP5", None)])
        self.assertEqual(kpis["覆盖80%的SKU"]["value"], 1)
