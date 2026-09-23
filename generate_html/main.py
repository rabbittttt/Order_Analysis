from __future__ import annotations

import argparse
import json
import logging
import os
import subprocess
import sys
import webbrowser
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from time import perf_counter
from typing import Any
from openpyxl import load_workbook

from core.components import add_kpi_comparisons
from core.china_calendar import PUBLISHED_YEARS
from core.config import load_config
from core.discovery import discover_subjects, set_capabilities
from core.excel import WorkbookItem, WorkbookStore, clean_text, format_excel_cell
from core.forecast_summary import SUMMARY_NAME
from core.renderer import b64gzip, render_dashboard
from core.validation import validate_manifest
from modules.registry import MODULES


ROOT = Path(__file__).resolve().parent
CODE_ROOT = ROOT.parent
PROJECT_ROOT = CODE_ROOT.parent if CODE_ROOT.name.lower() == "scripts" else CODE_ROOT
AUTO_OPEN_HTML = True
LOGGER = logging.getLogger("order_analysis")
BUILD_DIAGNOSTICS: list[str] = []


class DiagnosticCollector(logging.Handler):
    """Keep every emitted warning/error for the user-facing build report."""

    def emit(self, record: logging.LogRecord) -> None:
        if record.levelno >= logging.WARNING:
            BUILD_DIAGNOSTICS.append(f"{record.levelname} | {record.getMessage()}")


def refresh_sales_forecast_data(config: dict[str, Any], input_dir: Path | None = None) -> None:
    """Optionally rebuild the secondary forecast workbook with this Python runtime."""
    if not config.get("refresh_sales_forecast_data_before_build", False):
        LOGGER.info("销量预测二次处理刷新：已按 config.json 跳过")
        return
    script = ROOT / "tools" / "refresh_sales_forecast_data.py"
    if not script.exists():
        raise FileNotFoundError(f"已开启销量预测二次处理刷新，但找不到脚本：{script}")
    LOGGER.info("销量预测二次处理刷新：开始运行 %s", script)
    child_env = os.environ.copy()
    child_env["PYTHONIOENCODING"] = "utf-8"
    child_env["PYTHONUTF8"] = "1"
    command = [sys.executable, str(script)]
    if input_dir is not None:
        command.extend(["--orders", str(input_dir.resolve())])
        command.extend(["--output", str((input_dir / SUMMARY_NAME).resolve())])
    if config.get("forecast_as_of_date"):
        command.extend(["--as-of-date", str(config["forecast_as_of_date"])])
    result = subprocess.run(
        command,
        cwd=script.parent,
        env=child_env,
        text=True,
        capture_output=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    for line in result.stdout.splitlines():
        if line.startswith("[ERROR]"):
            LOGGER.error("[销量预测刷新] %s", line)
        elif line.startswith("[WARNING]"):
            LOGGER.warning("[销量预测刷新] %s", line)
        else:
            LOGGER.info("[销量预测刷新] %s", line)
    for line in result.stderr.splitlines():
        LOGGER.warning("[销量预测刷新] %s", line)
    if result.returncode:
        raise RuntimeError(f"销量预测二次处理刷新失败，退出码 {result.returncode}")
    LOGGER.info("销量预测二次处理刷新：完成")


class LevelFormatter(logging.Formatter):
    """Console formatter that colors problem levels; files keep plain text."""

    COLORS = {"WARNING": "\033[33m", "ERROR": "\033[31m", "CRITICAL": "\033[35m"}

    def format(self, record: logging.LogRecord) -> str:
        message = super().format(record)
        color = self.COLORS.get(record.levelname)
        return f"{color}{message}\033[0m" if color else message


def configure_runtime(output_dir: Path, debug: bool = False) -> Path:
    """Configure UTF-8 console/file logging before any workbook is opened."""
    output_dir.mkdir(parents=True, exist_ok=True)
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure:
            reconfigure(encoding="utf-8", errors="replace")
    log_dir = output_dir / "log"
    log_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_file = log_dir / f"generate_html_{timestamp}.log"
    if sys.platform == "win32":
        os.system("")  # enable ANSI colors in the Windows console
    fmt = "%(asctime)s | %(levelname)s | %(message)s"
    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(LevelFormatter(fmt, datefmt="%H:%M:%S"))
    file_handler = logging.FileHandler(log_file, mode="w", encoding="utf-8")
    file_handler.setFormatter(logging.Formatter(fmt, datefmt="%H:%M:%S"))
    BUILD_DIAGNOSTICS.clear()
    diagnostic_handler = DiagnosticCollector()
    logging.basicConfig(
        level=logging.DEBUG if debug else logging.INFO,
        handlers=[console, file_handler, diagnostic_handler],
        force=True,
    )
    return log_file


def serializable(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def log_payload_sizes(manifest: dict[str, Any]) -> None:
    """记录单HTML的主要体积来源，超阈值时提醒分页/延迟渲染。"""
    dashboard_sizes = {
        key: len(json.dumps(value, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))
        for key, value in manifest.get("dashboards", {}).items()
    }
    raw_sizes = {
        key: len(value.encode("ascii"))
        for key, value in manifest.get("raw_blocks", {}).items()
    }
    dashboard_total = sum(dashboard_sizes.values())
    raw_total = sum(raw_sizes.values())
    largest_dashboard = max(dashboard_sizes.items(), key=lambda item: item[1], default=("-", 0))
    largest_raw = max(raw_sizes.items(), key=lambda item: item[1], default=("-", 0))
    LOGGER.info(
        "HTML载荷: 看板JSON %.1fMB，底表压缩块 %.1fMB；最大看板=%s %.1fMB，最大底表块=%s %.1fMB",
        dashboard_total / 1024 / 1024,
        raw_total / 1024 / 1024,
        largest_dashboard[0],
        largest_dashboard[1] / 1024 / 1024,
        largest_raw[0],
        largest_raw[1] / 1024 / 1024,
    )
    if largest_raw[1] > 5 * 1024 * 1024:
        LOGGER.warning(
            "[性能校验] 单个底表压缩块超过5MB: %s（%.1fMB），建议对该Sheet分页或虚拟滚动",
            largest_raw[0],
            largest_raw[1] / 1024 / 1024,
        )


def split_parallel_blocks(name: str, rows: list[list]) -> list[dict[str, Any]]:
    """Split side-by-side blocks (like SKU's 表格A/B/C) into separate tables."""
    if len(rows) < 2:
        return [{"name": name, "rows": rows}]
    title_cols = [col for col, value in enumerate(rows[0]) if value not in (None, "")]
    # 多区块:第一行有多个标题，且标题列之间至少隔一个空列，第二行在各标题列有表头
    if len(title_cols) < 2:
        return [{"name": name, "rows": rows}]
    if any(right - left < 2 for left, right in zip(title_cols, title_cols[1:])):
        return [{"name": name, "rows": rows}]
    second = rows[1]
    if not all(col < len(second) and second[col] not in (None, "") for col in title_cols):
        return [{"name": name, "rows": rows}]
    width = max(len(row) for row in rows)
    blocks = []
    for index, start in enumerate(title_cols):
        end = title_cols[index + 1] if index + 1 < len(title_cols) else width
        block_rows = []
        for row in rows:
            sliced = row[start:end]
            while sliced and sliced[-1] in (None, ""):
                sliced.pop()
            block_rows.append(sliced)
        if block_rows:
            block_width = max(len(row) for row in block_rows)
            keep = [col for col in range(block_width) if any(col < len(row) and row[col] not in (None, "") for row in block_rows)]
            block_rows = [[row[col] if col < len(row) else None for col in keep] for row in block_rows]
        blocks.append({"name": f"{name} · {rows[0][start]}", "rows": block_rows})
    return blocks


def raw_tables(store: WorkbookStore) -> tuple[list[dict[str, Any]], dict[str, str]]:
    """底表数据:元信息放主数据,行数据按表分块压缩(base64+gzip),前端按需解压。"""
    files = []
    blocks: dict[str, str] = {}
    for file_index, item in enumerate(store.items):
        sheets = []
        for sheet in item.workbook.worksheets:
            if getattr(sheet, "sheet_state", "visible") != "visible":
                continue
            rows = []
            for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=sheet.max_column):
                values = [format_excel_cell(cell) for cell in row]
                while values and values[-1] in (None, ""):
                    values.pop()
                if any(value not in (None, "") for value in values):
                    rows.append(values)
            # 去掉整列全空的列，让并排区块（如 SKU 的 A/B/C 表）紧凑连排
            if rows:
                width = max(len(row) for row in rows)
                keep = [
                    col for col in range(width)
                    if any(col < len(row) and row[col] not in (None, "") for row in rows)
                ]
                rows = [[row[col] if col < len(row) else None for col in keep] for row in rows]
            for block in split_parallel_blocks(sheet.title, rows):
                block_rows = block["rows"]
                # 首行是区块大标题（单值）、次行才是列名时，去掉冗余标题行让列名成为表头
                if (
                    len(block_rows) > 1
                    and sum(1 for value in block_rows[0] if value not in (None, "")) <= 1
                    and sum(1 for value in block_rows[1] if value not in (None, "")) > 1
                ):
                    block_rows = block_rows[1:]
                key = f"{file_index}|{len(sheets)}"
                blocks[key] = b64gzip(json.dumps({"rows": block_rows}, ensure_ascii=False, separators=(",", ":")))
                sheets.append({
                    "name": block["name"],
                    "total_rows": len(block_rows),
                    "total_cols": max((len(row) for row in block_rows), default=0),
                    "source_sheet": sheet.title,
                })
        files.append({"name": item.path.name, "sheets": sheets})
    return files, blocks


def build(input_dir: Path, output_file: Path) -> tuple[dict[str, Any], list[str]]:
    build_started = perf_counter()
    LOGGER.info("开始生成订单分析看板")
    LOGGER.info("输入目录: %s", input_dir)
    LOGGER.info("输出文件: %s", output_file)
    config = load_config(ROOT / "config.json")
    if config.get("chart_render_mode") not in {"original", "generated"}:
        raise ValueError("config.json 的 chart_render_mode 必须是 original 或 generated")
    if date.today().year not in PUBLISHED_YEARS:
        LOGGER.warning(
            "[日历校验] %d年中国法定节假日数据尚未维护；预测日期将仅按自然周判断工作日/周末",
            date.today().year,
        )
    manifest_mode = config.get("build_manifest_mode", "compact")
    if manifest_mode not in {"compact", "full", "off"}:
        raise ValueError("config.json 的 build_manifest_mode 必须是 compact、full 或 off")
    stage_started = perf_counter()
    refresh_sales_forecast_data(config, input_dir)
    refresh_seconds = perf_counter() - stage_started
    store = WorkbookStore(input_dir)
    stage_started = perf_counter()
    store.load(exclude_names={SUMMARY_NAME})
    load_seconds = perf_counter() - stage_started
    if not store.items:
        details = "；".join(store.load_errors) if store.load_errors else "目录中没有xlsx文件"
        raise RuntimeError(f"没有成功读取任何Excel，无法生成看板：{details}")
    try:
        stage_started = perf_counter()
        subjects = discover_subjects(store)
        discovery_seconds = perf_counter() - stage_started
        LOGGER.info("已识别分析主体 %d 个: %s", len(subjects), "、".join(subject.name for subject in subjects))
        dashboards: dict[str, Any] = {}
        capabilities: dict[str, set[str]] = {}
        runtime_warnings = list(store.load_errors)
        module_seconds: dict[str, float] = defaultdict(float)
        module_builds: dict[str, int] = defaultdict(int)
        for module in MODULES:
            if module.id == "sales_forecast":
                module.summary_path = input_dir / SUMMARY_NAME
            set_config = getattr(module, "set_config", None)
            if set_config:
                set_config(config)
        for subject in subjects:
            matched_modules = []
            for module in MODULES:
                if hasattr(module, "set_dashboards"):
                    module.set_dashboards(dashboards)
                module_started = perf_counter()
                try:
                    dashboard = module.build(store, subject)
                except Exception as exc:
                    message = f"模块生成失败: 主体={subject.name}, 模块={module.label}, 错误={exc}"
                    runtime_warnings.append(message)
                    LOGGER.exception(message)
                    continue
                finally:
                    module_seconds[module.label] += perf_counter() - module_started
                    module_builds[module.label] += 1
                if not dashboard:
                    continue
                dashboards[f"{subject.id}|{module.id}"] = dashboard.to_dict()
                capabilities.setdefault("".join(subject.name.split()), set()).add(module.id)
                matched_modules.append(module.label)
            LOGGER.debug("主体映射完成: %s -> %s", subject.name, "、".join(matched_modules) if matched_modules else "仅底表")
        configured_order = config.get("module_order") or []
        known_modules = {module.id for module in MODULES} | {"raw"}
        unknown_modules = [module for module in configured_order if module not in known_modules]
        if unknown_modules:
            LOGGER.warning("[\u914d\u7f6e\u6821\u9a8c] module_order \u5305\u542b\u672a\u77e5\u6a21\u5757: %s", "\u3001".join(unknown_modules))
        missing_modules = [module.id for module in MODULES if module.id not in configured_order]
        if missing_modules:
            LOGGER.warning("[\u914d\u7f6e\u6821\u9a8c] module_order \u672a\u914d\u7f6e\u5df2\u6ce8\u518c\u6a21\u5757: %s", "\u3001".join(missing_modules))
        set_capabilities(subjects, capabilities, configured_order)
        add_kpi_comparisons(dashboards)
        stage_started = perf_counter()
        # Add the consolidated forecast workbook only after subject/module builds,
        # so imported source tabs never create extra analysis subjects.
        summary_path = input_dir / SUMMARY_NAME
        if summary_path.exists() and all(item.path.resolve() != summary_path.resolve() for item in store.items):
            store.items.append(WorkbookItem(summary_path, load_workbook(summary_path, data_only=True)))
        raw_files, raw_blocks = raw_tables(store)
        raw_seconds = perf_counter() - stage_started
        manifest = {
            "meta": {
                "title": config["title"],
                "generated_at": datetime.now().isoformat(timespec="seconds"),
                "input_dir": str(input_dir),
            },
            "config": config,
            "subjects": [subject.to_dict() for subject in subjects],
            "dashboards": dashboards,
            "chart_assets": store.chart_assets,
            "raw_files": raw_files,
            "raw_blocks": raw_blocks,
        }
        validation_warnings = validate_manifest(manifest)
        for warning in validation_warnings:
            LOGGER.warning("[生成校验] %s", warning)
        LOGGER.info("开始写入 HTML（主体 %d 个，看板 %d 个）", len(subjects), len(dashboards))
        stage_started = perf_counter()
        log_payload_sizes(manifest)
        render_dashboard(manifest, ROOT / "templates", output_file)
        render_seconds = perf_counter() - stage_started
        output_size = output_file.stat().st_size
        LOGGER.info("生成HTML体积: %.1fMB", output_size / 1024 / 1024)
        if output_size > 35 * 1024 * 1024:
            LOGGER.warning("[性能校验] 单HTML超过35MB，建议优先拆分大底表或图表资源")
        warnings = list(dict.fromkeys([
            *BUILD_DIAGNOSTICS,
            *(f"WARNING | {message}" for message in runtime_warnings if message not in "\n".join(BUILD_DIAGNOSTICS)),
        ]))
        stage_started = perf_counter()
        manifest_path = output_file.parent / "build_manifest.json"
        if manifest_mode != "off":
            with manifest_path.open("w", encoding="utf-8") as handle:
                if manifest_mode == "full":
                    json.dump(manifest, handle, ensure_ascii=False, indent=2)
                else:
                    json.dump(manifest, handle, ensure_ascii=False, separators=(",", ":"))
            LOGGER.info(
                "构建清单: %s 模式，%.1fMB",
                manifest_mode,
                manifest_path.stat().st_size / 1024 / 1024,
            )
        else:
            if manifest_path.exists():
                manifest_path.unlink()
            LOGGER.info("构建清单: 已按 config.json 跳过（旧清单已清理，避免误读过期数据）")
        legacy_warning_file = output_file.parent / "build_warnings.log"
        if legacy_warning_file.exists():
            legacy_warning_file.unlink()
        manifest_seconds = perf_counter() - stage_started
        if warnings:
            LOGGER.info("诊断完成：发现 %d 条映射、数据或生成异常，已逐条写入本次运行日志", len(warnings))
        else:
            LOGGER.info("校验完成：没有发现警告")
        module_summary = "；".join(
            f"{label} {seconds:.2f}s/{module_builds[label]}次"
            for label, seconds in sorted(module_seconds.items(), key=lambda item: item[1], reverse=True)
        )
        LOGGER.info("模块耗时: %s", module_summary)
        LOGGER.info(
            "性能汇总: 预测刷新 %.2fs | Excel加载 %.2fs | 主体识别 %.2fs | 模块生成 %.2fs | "
            "底表整理 %.2fs | HTML输出 %.2fs | 清单/诊断 %.2fs | 总计 %.2fs",
            refresh_seconds,
            load_seconds,
            discovery_seconds,
            sum(module_seconds.values()),
            raw_seconds,
            render_seconds,
            manifest_seconds,
            perf_counter() - build_started,
        )
        return manifest, warnings
    finally:
        store.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate the HarmonyOS Auto order dashboard.")
    parser.add_argument("--input", type=Path, default=PROJECT_ROOT / "output_file")
    parser.add_argument(
        "--output",
        type=Path,
        default=PROJECT_ROOT / "output_file" / "鸿蒙智行订单分析汇总.html",
    )
    parser.add_argument("--check", action="store_true", help="Return a non-zero exit code when validation warnings exist.")
    parser.add_argument("--no-open", action="store_true", help="Generate the HTML without opening the default browser.")
    parser.add_argument("--debug", action="store_true", help="输出调试日志，用于排查数据解析问题。")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    log_file = configure_runtime(args.output.parent, args.debug)
    try:
        manifest, warnings = build(args.input, args.output)
    except Exception:
        LOGGER.exception("生成失败，请根据上方错误和日志排查")
        LOGGER.error("运行日志: %s", log_file)
        return 1
    LOGGER.info("生成成功: %s", args.output)
    LOGGER.info("汇总: 主体=%d，看板=%d, 警告=%d", len(manifest["subjects"]), len(manifest["dashboards"]), len(warnings))
    LOGGER.info("运行日志: %s", log_file)
    if AUTO_OPEN_HTML and not args.no_open:
        try:
            if sys.platform == "win32":
                os.startfile(args.output)  # 用系统默认应用打开
                opened = True
            else:
                opened = webbrowser.open(args.output.resolve().as_uri())
            if opened:
                LOGGER.info("已调用系统默认应用打开看板")
            else:
                LOGGER.warning("系统没有确认看板已打开，请手动打开: %s", args.output)
        except Exception:
            LOGGER.exception("自动打开看板失败，请手动打开HTML")
    return 1 if args.check and warnings else 0


if __name__ == "__main__":
    raise SystemExit(main())
