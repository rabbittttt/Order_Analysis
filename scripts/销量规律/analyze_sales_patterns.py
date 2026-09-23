"""Single-workbook descriptive sales analysis. Python 3.10+ and openpyxl."""
import argparse, calendar, hashlib, importlib.util, json, math, re, time
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from statistics import mean, median
import openpyxl
SCRIPT_PATH=Path(__file__).resolve()
ROOT=next((parent for parent in SCRIPT_PATH.parents if (parent/'scripts'/'generate_html'/'core'/'china_calendar.py').exists()), SCRIPT_PATH.parents[2])

def log(message):
    print(f'[{datetime.now():%Y-%m-%d %H:%M:%S}] {message}',flush=True)
ORDER_METRIC_COLUMNS={
    '大定':(('大定','大定来源'),),
    '留存大定':(('留存大定','留存大定来源'),('净大定','净大定来源')),
    '小转大':(('小转大','小转大来源'),),
    '直接大定':(('直接大定','直接大定来源'),),
    '交车锁单':(('交车锁单','锁单来源'),),
}

def dt(v):
    if isinstance(v,datetime): return v.date()
    if isinstance(v,date): return v
    try: return date.fromisoformat(str(v)[:10])
    except ValueError: return None

def span(a,b):
    while a<=b:
        yield a
        a+=timedelta(days=1)

def group(rows,fields):
    g=defaultdict(list)
    for r in rows: g[tuple(r[f] for f in fields)].append(r)
    return g

def load_calendar(extra=None):
    spec=importlib.util.spec_from_file_location('calendar_rules',ROOT/'scripts/generate_html/core/china_calendar.py')
    m=importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
    periods=dict(m.HOLIDAY_PERIODS); work=dict(m.ADJUSTED_WORKDAYS); sources=dict(m.SOURCE_URLS)
    periods[2021]=(('元旦','2021-01-01','2021-01-03'),('春节','2021-02-11','2021-02-17'),('清明节','2021-04-03','2021-04-05'),('劳动节','2021-05-01','2021-05-05'),('端午节','2021-06-12','2021-06-14'),('中秋节','2021-09-19','2021-09-21'),('国庆节','2021-10-01','2021-10-07'))
    work[2021]=('2021-02-07','2021-02-20','2021-04-25','2021-05-08','2021-09-18','2021-09-26','2021-10-09')
    sources[2021]='https://app.www.gov.cn/govdata/gov/202011/26/465359/article.html'
    periods[2024]=tuple((n,'2023-12-30' if n=='元旦' else '2024-06-08' if n=='端午节' else a,b) for n,a,b in periods[2024])
    if extra:
        for y,c in json.loads(Path(extra).read_text(encoding='utf-8-sig')).items():
            y=int(y); periods[y]=c['periods']; work[y]=c['workdays']; sources[y]=c['source']
    events=[(y,n,dt(a),dt(b)) for y,items in sorted(periods.items()) for n,a,b in items]
    if any(not a or not b or b<a for _,_,a,b in events): raise ValueError('Invalid calendar')
    return dict(years=set(periods),events=events,holidays={d:n for _,n,a,b in events for d in span(a,b)},work={dt(d) for ds in work.values() for d in ds},sources=sources)

def ordinary(d,c):
    return d.year in c['years'] and d not in c['holidays'] and d not in c['work']

def read_panel(path):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    panel=[]; audit=Counter(); seen={}
    try:
        for name,small in [('小订及退订逐日',True),('当前订单逐日',False)]:
            it=wb[name].iter_rows(values_only=True); head=next(it)
            required={'日期','订单分析代际名'}|({'小订数量','真实逐日'} if small else {'订单阶段'})
            if required-set(head): raise ValueError(name+'缺少字段'+str(required-set(head)))
            metrics={'小订数量':('小订数量','小订来源')} if small else {}
            if not small:
                for metric,candidates in ORDER_METRIC_COLUMNS.items():
                    present=[(value_col,source_col) for value_col,source_col in candidates if value_col in head]
                    if not present:
                        continue
                    value_col,source_col=present[0]
                    if source_col not in head:
                        raise ValueError(name+'缺少字段'+source_col)
                    metrics[metric]=(value_col,source_col)
                    if len(present)>1:
                        audit[metric+'新旧字段同时存在，优先新字段']+=1
            for rowno,vs in enumerate(it,2):
                if all(v is None for v in vs): continue
                r=dict(zip(head,vs)); audit[name+'原始行']+=1
                d=dt(r.get('日期')); model=r.get('订单分析代际名')
                if not d or not model: audit['无效日期或车型']+=1; continue
                match=re.fullmatch(r'D(\d+)',str(r.get('生命周期') or ''),re.I)
                life=int(match[1]) if match else None
                stage='小订阶段' if small else str(r.get('订单阶段') or '未知阶段')
                anchor=dt(r.get('小订开始')) if small else None
                if not anchor and life and life>0: anchor=d-timedelta(days=life-1)
                cycle=str(anchor) if anchor and stage!='平销' else stage
                for metric,(value_col,scol) in metrics.items():
                    v=r.get(value_col)
                    if v is None: audit[metric+'缺失']+=1; continue
                    if isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v): audit[metric+'非数值']+=1; continue
                    if small and str(r.get('真实逐日')).lower() not in ('true','1','是'): audit['非真实小订排除']+=1; continue
                    src=str(r.get(scol) or '')
                    if not src or any(s in src for s in ('预测','合成','模拟','缺失')): audit[metric+'来源排除']+=1; continue
                    key=(str(model),d,metric); signature=(v,stage,src,cycle)
                    if key in seen:
                        if seen[key]!=signature: raise ValueError('重复键冲突'+str(key))
                        audit['相同重复去重']+=1; continue
                    seen[key]=signature
                    if v<0: audit[metric+'负值不参与倍率']+=1
                    panel.append(dict(model=str(model),date=d,metric=metric,stage=stage,source=src,cycle=cycle,value=float(v),life=life,sheet=name,row=rowno))
        attach_model_classes(wb, panel, audit)
    finally: wb.close()
    return panel,dict(audit)

def analyze(panel,c):
    out={k:[] for k in ('coverage','weekly','monthly','annual','holidays','lifecycle')}
    for (metric,stage),rs in group(panel,('metric','stage')).items():
        ds={r['date'] for r in rs}
        out['coverage'].append(dict(metric=metric,stage=stage,models=len({r['model'] for r in rs}),observations=len(rs),dates=len(ds),start=min(ds),end=max(ds)))
    fields=('model','metric','stage','cycle','source')
    for key,rs in group(panel,fields).items():
        base=dict(zip(fields,key)); days={r['date']:r for r in rs if r['value']>=0}
        if not days: continue
        life={r['life']:r['value'] for r in rs if r['life'] and r['value']>=0}
        if key[2]!='平销' and all(i in life for i in range(1,8)) and sum(life[i] for i in range(3,8))>0:
            out['lifecycle'].append(dict(base,d1_ratio=life[1]/mean(life[i] for i in range(3,8)),first2_share=(life[1]+life[2])/sum(life[i] for i in range(1,8))))
        wg=defaultdict(list); mg=defaultdict(list)
        for d,r in days.items():
            wg[d-timedelta(days=d.weekday())].append(r)
            mg[(d.year,d.month)].append(r)
        for monday,rows in wg.items():
            if len(rows)!=7 or not all(ordinary(r['date'],c) for r in rows): continue
            ys=[days[monday+timedelta(days=i)]['value'] for i in range(7)]
            if mean(ys)<=0 or mean(ys[:5])<=0: continue
            out['weekly'].append(dict(base,week=monday,indexes=[100*y/mean(ys) for y in ys],ratio=mean(ys[5:])/mean(ys[:5])))
        cm={}
        for (y,m),rows in mg.items():
            n=calendar.monthrange(y,m)[1]
            if len(rows)!=n: continue
            ys=[days[date(y,m,i)]['value'] for i in range(1,n+1)]
            cm[(y,m)]=mean(ys)
            denominator=mean(ys[:-7])
            out['monthly'].append(dict(base,year=y,month=m,
                ratio=mean(ys[-7:])/denominator if denominator>0 else None,
                status='有效' if denominator>0 else '月内前期日均为零'))
        for y in {y for y,m in cm}:
            if all((y,m) in cm for m in range(1,13)):
                ys=[cm[(y,m)] for m in range(1,13)]
                annual_mean=mean(ys)
                out['annual'].append(dict(base,year=y,
                    indexes=[100*v/annual_mean for v in ys] if annual_mean>0 else [None]*12,
                    status='有效' if annual_mean>0 else '全年日均为零'))
        if key[2]!='平销': continue
        for y,name,a,b in c['events']:
            if a>max(days) or b<min(days): continue
            target=list(span(a-timedelta(days=7),b+timedelta(days=7)))
            if not all(d in days for d in target): continue
            controls=[d for d in days if ordinary(d,c) and (a-timedelta(days=35)<=d<=a-timedelta(days=8) or b+timedelta(days=8)<=d<=b+timedelta(days=35)) and not any(x-timedelta(days=7)<=d<=z+timedelta(days=7) for _,_,x,z in c['events'])]
            ratios={}; used=set()
            for d in target:
                cs=[x for x in controls if x.weekday()==d.weekday()]
                if len(cs)<4 or not any(x<a for x in cs) or not any(x>b for x in cs): break
                expected=mean(days[x]['value'] for x in cs)
                if expected<=0: break
                ratios[d]=days[d]['value']/expected; used.update(cs)
            else:
                out['holidays'].append(dict(base,year=y,holiday=name,before=mean(v for d,v in ratios.items() if d<a),during=mean(v for d,v in ratios.items() if a<=d<=b),after=mean(v for d,v in ratios.items() if d>b),control_days=len(used)))
    return out

# Comparisons always retain each order metric and sales stage separately.
def period_bounds(d, kind):
    if kind == '日':
        return d, d
    if kind == '周':
        a = d - timedelta(days=d.weekday())
        return a, a + timedelta(days=6)
    month = d.month if kind == '月' else ((d.month-1)//3)*3+1 if kind == '季' else 1
    a = date(d.year, month, 1)
    width = 1 if kind == '月' else 3 if kind == '季' else 12
    endmonth = month + width - 1
    return a, date(d.year, endmonth, calendar.monthrange(d.year, endmonth)[1])

def quantile(values, p):
    vs = sorted(values)
    if not vs:
        return None
    pos = (len(vs)-1)*p
    low = int(pos); high = min(low+1, len(vs)-1)
    return vs[low] + (vs[high]-vs[low])*(pos-low)

def order_role(metric):
    return '锁单' if metric == '交车锁单' else '小订' if metric == '小订数量' else '大定'

def attach_model_classes(wb, panel, audit):
    definitions = ('brand', 'segment', 'energy')
    fields = ('品牌', '产品档位', '能源类型')
    mapping = defaultdict(lambda: defaultdict(set))
    if '车型基本信息' in getattr(wb, 'sheetnames', ()):
        it = wb['车型基本信息'].iter_rows(values_only=True)
        headers = next(it)
        for values in it:
            r = dict(zip(headers, values)); model = r.get('订单分析代际名')
            if not model:
                continue
            for key, field in zip(definitions, fields):
                if r.get(field):
                    mapping[str(model)][key].add(str(r[field]).strip())
    for r in panel:
        r['role'] = order_role(r['metric'])
        for key in definitions:
            choices = mapping[r['model']][key]
            r[key] = next(iter(choices)) if len(choices) == 1 else '映射冲突' if choices else '未提供'
    audit['属性缺失车型数'] = len({r['model'] for r in panel if any(r[k]=='未提供' for k in definitions)})
    audit['属性冲突车型数'] = len({r['model'] for r in panel if any(r[k]=='映射冲突' for k in definitions)})

def period_comparisons(panel, cal):
    comparisons = []
    interval_calendar_cache = {}
    for key, rs in group(panel, ('model','metric','stage','cycle','source')).items():
        base = dict(zip(('model','metric','stage','cycle','source'),key))
        base.update({k:rs[0].get(k,'未提供') for k in ('brand','segment','energy')})
        base['role'] = order_role(base['metric'])
        for kind in ('日','周','月','季','年'):
            buckets = defaultdict(list)
            for r in rs:
                a,b = period_bounds(r['date'],kind)
                buckets[a].append(r)
            for a, current in sorted(buckets.items()):
                _, b = period_bounds(a,kind)
                pa,pb = period_bounds(a-timedelta(days=1),kind)
                previous = buckets.get(pa,[])
                cn,pn = (b-a).days+1,(pb-pa).days+1
                cv,pv = sum(r['value'] for r in current),sum(r['value'] for r in previous)
                current_negative_days=sum(r['value']<0 for r in current)
                previous_negative_days=sum(r['value']<0 for r in previous)
                status = '有效'
                if current_negative_days:
                    status='本期含负值'
                elif previous_negative_days:
                    status='前期含负值'
                elif len(current)!=cn:
                    status='本期不完整'
                elif len(previous)!=pn:
                    status='前期缺失或不完整'
                elif pv==0:
                    status='前期为零'
                elif pv<0:
                    status='前期为负'
                transition = (str(pa.month)+'→'+str(a.month)+'月') if kind=='月' else (
                    '周'+str(pa.weekday()+1)+'→周'+str(a.weekday()+1) if kind=='日' else
                    'Q'+str((pa.month-1)//3+1)+'→Q'+str((a.month-1)//3+1) if kind=='季' else '前期→本期')
                interval=(pa,b)
                if interval not in interval_calendar_cache:
                    known=all(d.year in cal['years'] for d in span(pa,b))
                    special=sum(d in cal['holidays'] or d in cal['work'] for d in span(pa,b)) if known else None
                    interval_calendar_cache[interval]=(known,special)
                known,special=interval_calendar_cache[interval]
                comparisons.append(dict(base,kind=kind,start=a,end=b,previous_start=pa,previous_end=pb,
                    current=cv,previous=pv if previous else None,current_observed=len(current),
                    previous_observed=len(previous),current_days=cn,previous_days=pn,
                    current_negative_days=current_negative_days,previous_negative_days=previous_negative_days,
                    ratio=cv/pv if status=='有效' else None,
                    daily_ratio=(cv/cn)/(pv/pn) if status=='有效' else None,
                    transition=transition,status=status,holiday_adjusted_days=special))
    return comparisons

def summarize_comparisons(rows):
    buckets=defaultdict(list)
    for r in rows:
        if r['status']!='有效':
            continue
        # Each dimension is a separate grouping, never an additional volume to sum.
        categories=[('整体','全部车型')]+[(label,r.get(key,'未提供')) for label,key in
                     [('品牌','brand'),('产品档位','segment'),('能源类型','energy')]]
        for dim,value in categories:
            key=(r['role'],r['metric'],r['stage'],r['kind'],r['transition'],dim,value)
            buckets[key].append(r)
    summaries=[]
    for key,rs in sorted(buckets.items()):
        base=dict(zip(('role','metric','stage','kind','transition','dimension','category'),key))
        bymodel=group(rs,('model',))
        byperiod=group(rs,('start',))
        timeline=[median(x['daily_ratio'] for x in g) for g in byperiod.values()]
        years=defaultdict(list)
        for r in rs:
            years[r['start'].year].append(r)
        year_values=[]
        for y,yr in sorted(years.items()):
            model_values=[median(x['daily_ratio'] for x in g) for g in group(yr,('model',)).values()]
            year_values.append((y,median(model_values)))
        summaries.append(dict(base,pairs=len(rs),models=len(bymodel),periods=len(byperiod),
            years=len(years),ratio=median(median(x['ratio'] for x in g) for g in bymodel.values()),
            daily_ratio=median(median(x['daily_ratio'] for x in g) for g in bymodel.values()),
            p25=quantile(timeline,.25),p75=quantile(timeline,.75),
            above100=mean(mean(x['daily_ratio']>1 for x in g) for g in bymodel.values()),
            yearly='；'.join(str(y)+'年 '+format(v,'.1%') for y,v in year_values),
            evidence='不足跨年验证' if len(years)<2 else '跨年可比较，未控活动与趋势',
            status='少于8个不同周期' if len(byperiod)<8 else '已有8个以上周期，仍需留出验证'))
    return summaries

def correlation(xs, ys):
    if len(xs)<3:
        return None
    ax,ay=mean(xs),mean(ys)
    xx=sum((x-ax)**2 for x in xs); yy=sum((y-ay)**2 for y in ys)
    if xx==0 or yy==0:
        return None
    return sum((x-ax)*(y-ay) for x,y in zip(xs,ys))/math.sqrt(xx*yy)

def lock_lag_analysis(panel):
    results=[]
    for key, rs in group(panel, ('model','stage','cycle')).items():
        locks=group([r for r in rs if r['metric']=='交车锁单'],('source',))
        for metric in ('大定','留存大定'):
            deposits=group([r for r in rs if r['metric']==metric],('source',))
            for (ds,), dr in deposits.items():
                for (ls,), lr in locks.items():
                    x={r['date']:r['value'] for r in dr}
                    y={r['date']:r['value'] for r in lr}
                    # Use one non-negative date set for every lag and one for all daily changes.
                    matched=sorted(d for d in x if x[d]>=0 and
                                   all(d+timedelta(days=k) in y and y[d+timedelta(days=k)]>=0
                                       for k in range(4)))
                    changes=sorted(d for d in x if d-timedelta(days=1) in x and x[d]>=0 and x[d-timedelta(days=1)]>=0 and
                                   all(d+timedelta(days=k) in y and y[d+timedelta(days=k)]>=0
                                       for k in range(-1,4)))
                    for lag in range(4):
                        xs=[x[d] for d in matched]
                        ys=[y[d+timedelta(days=lag)] for d in matched]
                        dx=[x[d]-x[d-timedelta(days=1)] for d in changes]
                        dy=[y[d+timedelta(days=lag)]-y[d+timedelta(days=lag-1)] for d in changes]
                        enough=len(matched)>=14
                        results.append(dict(model=key[0],stage=key[1],cycle=key[2],metric=metric,
                            brand=rs[0].get('brand','未提供'),segment=rs[0].get('segment','未提供'),
                            energy=rs[0].get('energy','未提供'),lag=lag,pairs=len(matched),
                            start=min(matched) if matched else None,end=max(matched) if matched else None,
                            level_corr=correlation(xs,ys) if enough else None,
                            change_corr=correlation(dx,dy) if len(changes)>=14 else None,
                            change_pairs=len(changes),deposit_source=ds,lock_source=ls,
                            status='仅相关性，非转化率或逐单锁单时长' if enough else '不足14个共同日期'))
    return results

def extended_analysis(panel,cal):
    comparisons=period_comparisons(panel,cal)
    model_rows=[]
    for (model,),rs in group(panel,('model',)).items():
        model_rows.append(dict(model=model,**{k:rs[0].get(k,'未提供') for k in ('brand','segment','energy')},
                               stages='、'.join(sorted({r['stage'] for r in rs})),
                               metrics='、'.join(sorted({r['metric'] for r in rs}))))
    return dict(period_comparisons=comparisons,period_summary=summarize_comparisons(comparisons),
                lock_lags=lock_lag_analysis(panel),model_classes=model_rows)


def sha256_file(path):
    digest=hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda:stream.read(1024*1024),b''):
            digest.update(block)
    return digest.hexdigest()

def load_component(path, name):
    spec=importlib.util.spec_from_file_location(name,path)
    if spec is None or spec.loader is None:
        raise RuntimeError('无法加载分析组件：'+str(path))
    module=importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def main():
    started=time.perf_counter()
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--input',type=Path,default=ROOT/'output_file/鸿蒙智行销量数据汇总.xlsx')
    ap.add_argument('--out',type=Path,default=None,help='已弃用；为兼容旧命令保留，不再使用该路径')
    ap.add_argument('--calendar',type=Path)
    ap.add_argument('--sample-data',action='store_true',help='将输出标记为样例数据')
    ap.add_argument('--real-data',action='store_true',help=argparse.SUPPRESS)
    a=ap.parse_args()
    src=a.input.resolve()
    final=(ROOT/'output_file'/'鸿蒙智行销量规律分析.xlsx').resolve()
    exporter=SCRIPT_PATH.with_name('export_excel_openpyxl.py')
    calendar_path=ROOT/'scripts'/'generate_html'/'core'/'china_calendar.py'
    config_path=SCRIPT_PATH.with_name('report_config.json')
    rules_path=SCRIPT_PATH.with_name('report_rules.py')
    forecast_path=SCRIPT_PATH.with_name('forecast_insights.py')
    web_exporter=SCRIPT_PATH.with_name('export_html.py')
    web_final=final.with_suffix('.html')
    templates=SCRIPT_PATH.with_name('templates')

    log('销量规律分析启动')
    if a.out is not None:
        log('--out参数已弃用，忽略该路径；结果仍写入固定输出文件')
    log('输入文件：'+str(src))
    if src==final:
        raise ValueError('输入文件与固定输出文件相同，已停止以保护源数据：'+str(src))
    log('分析目的：辅助销量预测；净大定（留存大定）与交车锁单分别研究')
    required=(('输入工作簿',src),('节假日日历',calendar_path),
              ('Excel导出器',exporter),('报告规则配置',config_path),
              ('规律汇总模块',rules_path),('预测辅助模块',forecast_path),('网页导出器',web_exporter),
              *((f'网页模板 {name}',templates/name) for name in ('patterns.html','patterns.css','patterns.js')))
    missing=[f'{label}不存在：{path}' for label,path in required if not path.is_file()]
    if a.calendar is not None and not a.calendar.is_file():
        missing.append(f'额外节假日日历不存在：{a.calendar.resolve()}')
    if missing:
        raise FileNotFoundError('\n'.join(missing))

    log('正在读取并校验源数据...')
    stage=time.perf_counter()
    digest=sha256_file(src)
    panel,audit=read_panel(src)
    if not panel:
        raise ValueError('源工作簿中没有有效观测，请检查日期、车型、指标值和来源字段')
    log(f'源数据读取及校验完成：有效观测 {len(panel):,} 条，车型 {len({r["model"] for r in panel}):,} 个，用时 {time.perf_counter()-stage:.1f} 秒')
    log(f'数据日期范围：{min(r["date"] for r in panel)} 至 {max(r["date"] for r in panel)}')

    stage=time.perf_counter()
    c=load_calendar(a.calendar)
    data_years=sorted({r['date'].year for r in panel})
    unknown=sorted(set(data_years)-c['years'])
    log('节假日日历已加载：'+('、'.join(map(str,sorted(c['years']))) or '无')+f'，用时 {time.perf_counter()-stage:.1f} 秒')
    if unknown:
        log('提示：数据覆盖年份缺少日历规则：'+'、'.join(map(str,unknown))+'；这些年份不参与普通周/节假日识别')

    log('正在计算周、月、年、节假日和生命周期规律...')
    stage=time.perf_counter()
    result=analyze(panel,c)
    log(f'周、月、年、节假日和生命周期计算完成，用时 {time.perf_counter()-stage:.1f} 秒')
    log('正在计算周期环比、车型分类和锁单滞后...')
    stage=time.perf_counter()
    result.update(extended_analysis(panel,c))
    log(f'周期比较、车型分类和锁单滞后计算完成，用时 {time.perf_counter()-stage:.1f} 秒')

    log('正在建立净大定、锁单预测基准并做滚动回测...')
    stage=time.perf_counter()
    forecast=load_component(forecast_path,'sales_forecast_insights').build_forecast_insights(panel,c,log=log)
    for table in ('profiles','backtests'):
        for row in forecast.get(table,[]):
            row['grain']={'day':'日','week':'周','month':'月'}.get(row.get('grain'),row.get('grain'))
    log(f'预测辅助及回测计算完成：{len(forecast.get("profiles", [])):,} 组，耗时 {time.perf_counter()-stage:.1f} 秒')

    if sha256_file(src)!=digest:
        raise RuntimeError('分析期间源文件发生变化，已停止导出')
    payload=dict(sample=a.sample_data and not a.real_data,source=str(src),sha256=digest,
                 audit=audit,calendar_years=sorted(c['years']),calendar_sources=c['sources'],
                 analysis_end=max(r['date'] for r in panel),results=result,forecast=forecast,
                 calendar_events=[dict(year=y,holiday=n,start=a,end=b) for y,n,a,b in c['events']])
    log('正在准备离线交互网页...')
    web=load_component(web_exporter,'sales_html_exporter')
    html=web.render_html(web.build_web_data(payload,panel))
    final.parent.mkdir(parents=True,exist_ok=True)
    log('分析计算完成，正在生成Excel...')
    stage=time.perf_counter()
    module=load_component(exporter,'sales_excel_exporter')
    module.export_workbook(payload,final,log=log)
    if not final.is_file():
        raise RuntimeError('Excel导出命令已结束，但未找到最终文件：'+str(final))
    log(f'Excel生成完成，用时 {time.perf_counter()-stage:.1f} 秒')
    log('正在保存离线交互网页...')
    web.export_html(html,web_final,log=log)
    log(f'分析完成，总耗时 {time.perf_counter()-started:.1f} 秒')
    print('Excel：'+str(final),flush=True)
    print('网页：'+str(web_final),flush=True)
    print('源文件未修改；结果数量：'+str({k:len(v) for k,v in result.items()}),flush=True)
if __name__=='__main__': main()


