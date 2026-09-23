from __future__ import annotations

import json
import os
import re
import tempfile
import time
import zipfile
import importlib.util
import math
import unicodedata
from functools import lru_cache
from collections import defaultdict
from datetime import date
from pathlib import Path
from statistics import mean, median

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT = "Microsoft YaHei"
DARK = "24455F"
TEXT = "243447"
TITLE = "18344B"
SECTION = "E7EEF4"
ALT = "F2F5F8"


@lru_cache(maxsize=8192)
def _text_units(text):
    return max((sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in line)
                for line in text.splitlines()), default=0)


def _row_height(text, width):
    lines = max(len(text.splitlines()), math.ceil(_text_units(text) / max(width - 2, 1)))
    return min(409, max(27, lines * 13 + 6))



def _groups(rows, keys):
    result = defaultdict(list)
    for row in rows:
        result[tuple(row.get(k) for k in keys)].append(row)
    return list(result.items())


def _within(rows, field):
    values = []
    for _, model_rows in _groups(rows, ["model"]):
        nums = [r.get(field) for r in model_rows if r.get(field) is not None]
        if nums:
            values.append(median(nums))
    return median(values) if values else None


def _choose_grouping(models, config):
    dimensions = [("segment", "产品档位"), ("brand", "品牌"), ("energy", "能源类型")]
    group_by = config.get("groupBy", "auto")
    min_groups = config.get("minGroups", 3)
    max_groups = config.get("maxGroups", 8)
    if not isinstance(min_groups, int) or not isinstance(max_groups, int) or min_groups < 1 or max_groups < min_groups:
        raise ValueError("分组数量范围无效")
    if group_by != "auto" and group_by not in {x[0] for x in dimensions}:
        raise ValueError("groupBy须为auto、segment、brand或energy")
    candidates = []
    for priority, (key, label) in enumerate(dimensions):
        valid = [str(r.get(key) or "").strip() for r in models]
        valid = [x for x in valid if x and x not in {"未提供", "映射冲突"}]
        count = len(set(valid))
        coverage = len(valid) / len(models) if models else 0
        candidates.append(dict(key=key, label=label, priority=priority, count=count,
                               coverage=coverage, fit=min_groups <= count <= max_groups))
    complete = [r for r in candidates if r["coverage"] >= .8 and r["count"] > 1]
    pool = complete or candidates
    if group_by == "auto":
        selected = sorted(pool, key=lambda x: (-int(x["fit"]), -x["coverage"], x["priority"]))[0]
    else:
        selected = next(r for r in candidates if r["key"] == group_by)
    entries = []
    for row in models:
        item = dict(row)
        value = str(row.get(selected["key"]) or "").strip() or "未提供"
        item["category"] = f'{selected["label"]}：{value}'
        entries.append(item)
    categories = sorted({r["category"] for r in entries})
    reason = (
        ("自动选择" if group_by == "auto" else "配置指定") + f' {selected["label"]}；'
        f'有效属性覆盖 {selected["coverage"] * 100:.1f}%；原始类别 {selected["count"]} 个；缺失/冲突独立展示'
    )
    return dict(entries=entries, categories=categories, candidates=candidates, reason=reason)


def _selected_metrics(comparisons, config=None):
    available = list(dict.fromkeys(r["metric"] for r in comparisons))
    available.sort(key=lambda metric: ({"留存大定": 0, "交车锁单": 1}.get(metric, 2), metric))
    requested = (config or {}).get("mainMetrics", "auto")
    if requested == "auto":
        return available, []
    if not isinstance(requested, list) or not requested or any(not isinstance(x, str) for x in requested):
        raise ValueError("mainMetrics须为auto或非空指标名称列表")
    unknown = [x for x in requested if x not in available]
    if unknown:
        raise ValueError("mainMetrics指定的指标在本次数据中不存在：" + "、".join(unknown))
    return list(dict.fromkeys(requested)), [x for x in available if x not in requested]


def _report_series(comparisons, categories, model_map, selected_metrics=None):
    allowed = set(selected_metrics) if selected_metrics is not None else {r["metric"] for r in comparisons}
    buckets = defaultdict(dict)
    for row in comparisons:
        if row["metric"] not in allowed:
            continue
        key = (row["metric"], row["stage"])
        category = model_map.get(row["model"], {}).get("category")
        buckets["全部车型"][key] = None
        if category is not None:
            buckets[category][key] = None
    return [dict(category=category, metric=metric, stage=stage)
            for category in ["全部车型", *categories] for metric, stage in buckets[category]]


def _as_date(value):
    if not value:
        return None
    return date.fromisoformat(str(value)[:10])


class ExcelBuilder:
    """Reuse workbook styles and report sheet progress without retaining preview files."""
    def __init__(self, sample=False, log=None):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.sample = sample
        self.table_no = 0
        self.log = log or (lambda message: None)
        self.styles = {}
        self.font = Font(name=FONT, size=10, color=TEXT)
        self.header_font = Font(name=FONT, size=10, bold=True, color="FFFFFF")

    def _style(self, number_format="General", stripe=False, numeric=False, header=False):
        key = (number_format, stripe, numeric, header)
        if key not in self.styles:
            style = NamedStyle(name=f"SalesStyle{len(self.styles)}")
            style.font = self.header_font if header else self.font
            style.alignment = Alignment(horizontal="center" if header else "right" if numeric else "left",
                                        vertical="center", wrap_text=header or not numeric)
            style.number_format = number_format
            if header or stripe:
                style.fill = PatternFill("solid", fgColor=DARK if header else ALT)
            self.wb.add_named_style(style)
            self.styles[key] = style
        return self.styles[key]

    def sheet(self, name, widths):
        self.log(f"正在生成工作表：{name}")
        ws = self.wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws["A2"] = name
        ws["A2"].font = Font(name=FONT, size=16, bold=True, color=TITLE)
        ws.row_dimensions[2].height = 30
        self.note(ws, 3, "样例结果，仅验证分析流程，不代表真实经营规律。" if self.sample
                  else "本次计算结果快照；更新数据后需重新运行分析。")
        ws["A3"].font = Font(name=FONT, size=10, color="945B10")
        return ws

    def note(self, ws, row, text):
        cell = ws.cell(row, 1, text)
        cell.font = self.font
        cell.alignment = Alignment(vertical="center", wrap_text=False)
        ws.row_dimensions[row].height = 25

    def section(self, ws, row, title, last):
        for col in range(1, last + 1):
            cell = ws.cell(row, col)
            cell.fill = PatternFill("solid", fgColor=SECTION)
        cell = ws.cell(row, 1, title)
        cell.font = Font(name=FONT, size=10, bold=True, color=TITLE)
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 25

    def table(self, ws, start, headers, rows, formats=None, filter_table=False):
        started = time.perf_counter()
        formats = formats or {}
        if start + len(rows) > 1_048_576:
            raise ValueError(f"工作表“{ws.title}”有 {len(rows):,} 行，超过Excel单表容量，请缩小输入日期范围。")
        header_style = self._style(header=True)
        for col, header in enumerate(headers, 1):
            cell = ws.cell(start, col, header)
            cell.style = header_style.name
        ws.row_dimensions[start].height = 34
        styles = {
            (col, stripe, numeric): self._style(formats.get(col - 1, "General"), stripe, numeric)
            for col in range(1, len(headers) + 1) for stripe in (False, True) for numeric in (False, True)
        }
        for ridx, row in enumerate(rows, start + 1):
            stripe = (ridx - start) % 2 == 0
            height = 27
            for cidx, value in enumerate(row, 1):
                cell = ws.cell(ridx, cidx, value)
                if isinstance(value, str):
                    cell.data_type = "s"  # Source labels remain text, even when they begin with '='.
                    height = max(height, _row_height(value, ws.column_dimensions[get_column_letter(cidx)].width))
                cell.style = styles[cidx, stripe, isinstance(value, (int, float, date))].name
            ws.row_dimensions[ridx].height = height
            written = ridx - start
            if written % 10000 == 0:
                self.log(f"工作表 {ws.title}：已写入 {written:,}/{len(rows):,} 行")
        if filter_table and rows:
            self.table_no += 1
            ref = f"A{start}:{get_column_letter(len(headers))}{start + len(rows)}"
            tab = Table(displayName=f"Data{self.table_no}", ref=ref)
            tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
            ws.add_table(tab)
        self.log(f"工作表 {ws.title}：写入 {len(rows):,} 行，耗时 {time.perf_counter()-started:.2f} 秒")
        return start + len(rows) + 1

    def feature(self, name, headers, rows, widths, formats, method):
        ws = self.sheet(name, widths)
        self.note(ws, 5, method)
        self.table(ws, 7, headers, rows, formats, True)
        if not rows:
            self.note(ws, 8, "当前没有满足条件的可比样本，详见“数据覆盖与方法”。")
        ws.freeze_panes = "C8"
        ws.print_title_rows = "7:7"
        return ws


def _atomic_save(workbook, output_path, log):
    """Only replace the previous report after the new workbook passes basic validation."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=f".{output_path.stem}-", suffix=".xlsx", dir=output_path.parent)
    os.close(fd)
    temp_path = Path(temp_name)
    started = time.perf_counter()
    try:
        log(f"正在保存Excel：{output_path}")
        workbook.save(temp_path)
        with zipfile.ZipFile(temp_path) as archive:
            bad = archive.testzip()
            if bad:
                raise ValueError(f"Excel保存校验失败：{bad}")
        check = load_workbook(temp_path, read_only=True, data_only=False)
        try:
            if check.sheetnames != workbook.sheetnames:
                raise ValueError("Excel保存校验失败：工作表不完整")
        finally:
            check.close()
        os.replace(temp_path, output_path)
        log(f"Excel保存完成：{len(workbook.sheetnames)} 张工作表，{output_path.stat().st_size / 1024 / 1024:.2f} MB，耗时 {time.perf_counter()-started:.2f} 秒")
    except PermissionError as exc:
        raise PermissionError(f"无法保存Excel，请关闭已打开的“{output_path.name}”并检查目录写入权限。上次成功的结果已保留。") from exc
    finally:
        temp_path.unlink(missing_ok=True)


def export_workbook(payload, output_path, log=None):
    log = log or (lambda message: None)
    d = payload["results"]
    script_dir = Path(__file__).resolve().parent
    config = json.loads((script_dir / "report_config.json").read_text(encoding="utf-8-sig"))
    grouping = _choose_grouping(d["model_classes"], config)
    categories = grouping["categories"]
    model_map = {r["model"]: r for r in grouping["entries"]}
    selected_metrics, omitted_metrics = _selected_metrics(d["period_comparisons"], config)
    series = _report_series(d["period_comparisons"], categories, model_map, selected_metrics)

    rule_spec = importlib.util.spec_from_file_location("sales_report_rules", script_dir / "report_rules.py")
    rule_module = importlib.util.module_from_spec(rule_spec)
    rule_spec.loader.exec_module(rule_module)
    rule_rows = rule_module.build_rule_summary(d, model_map, categories, selected_metrics)

    def pick(rows, category, metric, stage):
        return [r for r in rows if (category == "全部车型" or model_map.get(r.get("model"), {}).get("category") == category)
                and r.get("metric") == metric and r.get("stage") == stage]

    month_rows = [r for r in d["period_comparisons"] if r.get("kind") == "月"]
    monthly_profiles = []
    for item in series:
        scoped = pick(month_rows, item["category"], item["metric"], item["stage"])
        for month in range(1, 13):
            rows = [r for r in scoped if int(str(r["start"])[5:7]) == month]
            complete = [dict(r, daymean=r["current"] / r["current_days"]) for r in rows
                        if r.get("current_observed") == r.get("current_days") and r.get("current_negative_days", 0) == 0 and r.get("current", -1) >= 0]
            paired = [r for r in rows if r.get("status") == "有效"]
            years = sorted({str(r["start"])[:4] for r in paired})
            annual = [(year, _within([r for r in paired if str(r["start"]).startswith(year)], "daily_ratio")) for year in years]
            if not paired:
                conclusion = "缺少完整相邻月"
            elif len(years) < 2:
                conclusion = "单年样本，不能认定季节性"
            elif all(v > 1 for _, v in annual):
                conclusion = "各年中位数均上升，待控制活动"
            elif all(v < 1 for _, v in annual):
                conclusion = "各年中位数均下降，待控制活动"
            else:
                conclusion = "跨年方向不一致"
            monthly_profiles.append(dict(
                **item, month=month, daymean=_within(complete, "daymean"), ratio=_within(paired, "ratio"),
                daily=_within(paired, "daily_ratio"), models=len({r["model"] for r in complete}),
                complete=len(complete), pairs=len(paired), pairmodels=len({r["model"] for r in paired}),
                years=len(years), periods=len({r["start"] for r in paired}),
                above=mean([mean([1 if r["daily_ratio"] > 1 else 0 for r in rs]) for _, rs in _groups(paired, ["model"])]) if paired else None,
                yearly="；".join(f"{year}：{value * 100:.1f}%" for year, value in annual),
                conclusion=conclusion
            ))

    b = ExcelBuilder(payload.get("sample", False), log=log)
    overview = b.sheet("分析总览", [22, 12, 24, 24, 24, 24, 24, 15])
    overview.sheet_properties.tabColor = TITLE
    b.section(overview, 5, "车型分组｜" + grouping["reason"], 8)
    class_rows = [[c, sum(r["category"] == c for r in model_map.values()),
                   "、".join(r["model"] for r in model_map.values() if r["category"] == c)] for c in categories]
    b.table(overview, 6, ["车型组", "车型数", "车型成员（代际）"], class_rows)
    overview.merge_cells("C6:H6")
    for i in range(len(categories)):
        overview.merge_cells(start_row=7 + i, start_column=3, end_row=7 + i, end_column=8)
        member_width = sum(overview.column_dimensions[get_column_letter(c)].width for c in range(3, 9))
        overview.row_dimensions[7 + i].height = max(44, _row_height(class_rows[i][2], member_width))
    note_row = 8 + len(categories)
    summary_start = note_row + 3
    b.note(overview, note_row, "类别及车型自动读取；不合并原始类别、不凑固定组数。分组维度可在report_config.json中指定。")
    b.section(overview, summary_start - 1, "预测目标优先｜净大定（留存大定）、锁单分别分析；100%＝持平", 8)
    overview_rows = []
    for item in series:
        w = pick(d["weekly"], item["category"], item["metric"], item["stage"])
        m = [r for r in monthly_profiles if all(r[k] == item[k] for k in ("category", "metric", "stage"))]
        pm = [r for r in pick(month_rows, item["category"], item["metric"], item["stage"]) if r.get("status") == "有效"]
        overview_rows.append([item["category"], item["metric"], item["stage"], _within(w, "ratio"),
                              len({r["week"] for r in w}), sum(r["pairs"] > 0 for r in m), len(pm),
                              "见全年月份" if pm else "暂无相邻月"])
    b.table(overview, summary_start, ["车型组", "指标", "阶段", "周末/工作日日均", "完整普通周数", "可比月份/12", "相邻月比较对数", "月份结论"],
            overview_rows, {3: "0.0%"})
    end = summary_start + len(overview_rows) + 1
    b.note(overview, end + 1, "全年月份页：逐月看日均比、月度日均量及覆盖；月份证据页：看总量比、年份、跨年方向。")
    b.note(overview, end + 2, "周末比例先车型内取中位数，再车型等权；空白＝无合格样本，0%才表示真实为零。")
    b.note(overview, end + 3, "首销及小订受上市节奏影响；常年季节性应优先检验平销，并控制上市、促销和节假日。")
    overview.freeze_panes = "A7"

    selected_note = ("主视图按配置展示：" + "、".join(selected_metrics) + "；未展示：" + "、".join(omitted_metrics)
                     if omitted_metrics else "主视图展示本次数据中全部可用指标，指标和销售阶段分别统计。")
    b.note(overview, end + 4, selected_note)
    summary_row_start = end + 7
    b.section(overview, summary_row_start, "规律摘要：各组、指标、阶段选取周期数最多的一项，全部结果见“规律汇总”", 8)
    featured = {}
    for rule in rule_rows:
        if rule.get("typical") is None:
            continue
        key = (rule["category"], rule["metric"], rule["stage"])
        prior = featured.get(key)
        if prior is None or (rule["periods"], rule["years"]) > (prior["periods"], prior["years"]):
            featured[key] = rule
    b.table(overview, summary_row_start + 1,
            ["车型组", "指标", "阶段", "规律", "典型比例", "年份数", "不同周期数", "验证情况"],
            [[r["category"], r["metric"], r["stage"], r["pattern"], r["typical"],
              r["years"], r["periods"], r["validation_status"]] for r in featured.values()], {4: "0.0%"})
    for ridx in range(summary_row_start + 2, summary_row_start + 2 + len(featured)):
        overview.row_dimensions[ridx].height = max(44, overview.row_dimensions[ridx].height or 0)
    rule_ws = b.feature("规律汇总",
        ["车型组", "指标", "阶段", "规律", "典型比例", "周期P25", "周期P75", "车型数", "样本数",
         "不同周期数", "年份数", "描述", "历史比例", "验证年度", "验证年度比例", "验证情况", "明细工作表"],
        [[r["category"], r["metric"], r["stage"], r["pattern"], r["typical"], r["p25"], r["p75"], r["models"],
          r["samples"], r["periods"], r["years"], r["conclusion"], r["train_ratio"], r["validation_year"],
          r["validation_ratio"], r["validation_status"], r["detail_sheet"]] for r in rule_rows],
        [24, 16, 14, 32, 16, 16, 16, 10, 12, 14, 10, 40, 16, 14, 17, 48, 22],
        {4: "0.0%", 5: "0.0%", 6: "0.0%", 12: "0.0%", 14: "0.0%"},
        "典型比例先车型内再车型等权取中位数；P25/P75按日历周期计算；方向验证不等于统计显著或因果关系。")
    rule_ws.sheet_properties.tabColor = "197A83"

    months = b.sheet("全年月份", [22, 16, 14] + [12] * 12)
    months.sheet_properties.tabColor = "197A83"
    b.note(months, 5, "1月对比上年12月；每格按车型等权中位数。空白表示缺样本，不能用零替代。")
    row = 7
    for field, title, fmt in [
        ("daily", "① 本月日均 / 上月日均｜排除自然月天数不同的影响", "0.0%"),
        ("daymean", "② 完整月日均量｜单车型日均的中位数，单位：单/天", "#,##0.0"),
        ("pairs", "③ 完整相邻月比较对数｜与①逐格对应", "0")]:
        b.section(months, row, title, 15)
        rows = []
        for item in series:
            p = [r for r in monthly_profiles if all(r[k] == item[k] for k in ("category", "metric", "stage"))]
            rows.append([item["category"], item["metric"], item["stage"], *[r[field] for r in p]])
        b.table(months, row + 1, ["车型组", "指标", "阶段", *[f"{i}月" for i in range(1, 13)]],
                rows, {i: fmt for i in range(3, 15)})
        if field == "daily":
            for ridx, values in enumerate(rows, row + 2):
                for cidx, value in enumerate(values[3:], 4):
                    if value is not None:
                        months.cell(ridx, cidx).fill = PatternFill("solid", fgColor="DCEFE8" if value > 1.02 else "FBE7DD" if value < .98 else "EDF0F3")
        row += len(rows) + 4
    b.note(months, row, "不同月份的车型构成可能不同，完整月日均量不可直接用于淡旺季排名；应以同车型相邻月和分年复核为主。")
    months.freeze_panes = "D9"

    b.feature("月份证据",
        ["车型组", "指标", "阶段", "月份", "单车型日均中位数", "完整月车型数", "完整车型月数", "相邻月车型数", "相邻月比较对数",
         "不同年月数", "年份数", "总量占上月", "日均占上月", "上升占比", "分年日均比", "证据判断"],
        [[r["category"], r["metric"], r["stage"], r["month"], r["daymean"], r["models"], r["complete"], r["pairmodels"], r["pairs"],
          r["periods"], r["years"], r["ratio"], r["daily"], r["above"], r["yearly"], r["conclusion"]] for r in monthly_profiles],
        [22, 16, 14, 9, 20, 16, 16, 16, 18, 16, 12, 17, 17, 15, 48, 38],
        {4: "#,##0.0", 11: "0.0%", 12: "0.0%", 13: "0.0%"},
        "每个组×指标×阶段均保留1—12月。先车型内中位数，再车型等权；比较对限定同车型、阶段、批次和来源。")

    b.feature("周内明细",
        ["车型", "指标", "阶段", "批次", "周一日期", "周一", "周二", "周三", "周四", "周五", "周六", "周日", "周末倍率", "数据来源"],
        [[r["model"], r["metric"], r["stage"], r["cycle"], _as_date(r["week"]), *r["indexes"], r["ratio"], r["source"]] for r in d["weekly"]],
        [29, 15, 14, 15, 15, 10, 10, 10, 10, 10, 10, 10, 14, 55],
        {4: "yyyy-mm-dd", **{i: "0.0" for i in range(5, 12)}, 12: '0.000"倍"'},
        "同车型、阶段、批次和来源的完整普通周；节假日与调休所在周排除。")
    b.feature("月内明细",
        ["车型", "指标", "阶段", "批次", "年份", "月份", "月末7天倍率", "状态", "数据来源"],
        [[r["model"], r["metric"], r["stage"], r["cycle"], r["year"], r["month"], r["ratio"], r.get("status", "有效"), r["source"]] for r in d["monthly"]],
        [29, 15, 14, 15, 10, 10, 17, 24, 55], {6: '0.000"倍"'},
        "完整自然月；最后7天日均/其余日期日均。不补齐缺失日期。")
    b.feature("生命周期明细",
        ["车型", "指标", "阶段", "批次", "首日相对D3—D7", "前两天占首周", "数据来源"],
        [[r["model"], r["metric"], r["stage"], r["cycle"], r["d1_ratio"], r["first2_share"], r["source"]] for r in d["lifecycle"]],
        [29, 15, 14, 15, 18, 19, 55], {4: '0.000"倍"', 5: "0.0%"},
        "要求D1—D7完整；小转大与直接大定分别分析，避免混淆集中转化和新增需求。")

    methods = b.sheet("数据覆盖与方法", [30, 55, 14, 16, 16, 17, 17, 18])
    b.table(methods, 6, ["指标", "阶段", "车型数", "车型日数", "不同日期数", "开始日期", "结束日期"],
            [[r["metric"], r["stage"], r["models"], r["observations"], r["dates"], _as_date(r["start"]), _as_date(r["end"])] for r in d["coverage"]],
            {5: "yyyy-mm-dd", 6: "yyyy-mm-dd"})
    source_row = 8 + len(d["coverage"])
    b.section(methods, source_row, "数据来源与检查", 8)
    b.note(methods, source_row + 1, "业务来源：鸿蒙智行销量数据汇总.xlsx / 小订及退订逐日、当前订单逐日")
    b.note(methods, source_row + 2, "结果模式：" + ("本地样例" if payload.get("sample") else "实际数据") + "；所有数值为本次分析快照，原业务工作簿未修改。")
    b.note(methods, source_row + 3, "SHA-256：" + payload["sha256"])
    audit_row = source_row + 5
    b.table(methods, audit_row, ["检查项", "记录数"], list(payload["audit"].items()))
    method_row = audit_row + len(payload["audit"]) + 3
    b.section(methods, method_row, "分析方法与适用边界", 8)
    notes = [
        ("周", "完整普通周；指数每日/该周日均×100；先车型内平均再车型等权。"),
        ("月", "完整自然月；月末最后7天与其余日期日均比较，未调整星期与节日。"),
        ("年", "完整年度逐年显示；通用季节性需多年度验证。零月保留，全年为零时指数留空。"),
        ("节假日", "平销节前7天、节中、节后7天均须完整；同星期匹配前后对照。"),
        ("缺失与来源", "缺失不补零；预测、合成、模拟和缺失来源排除。负值保留在观测合计，含负值周期不算倍率。"),
        ("统计边界", "规律汇总提供完整年度的留出方向验证；未做回归、置信区间或因果识别。"),
        ("日历", "已配置年份：" + "、".join(map(str, payload.get("calendar_years", []))) + "；未配置年份不判断普通周与节日。"),
        ("首页分类", grouping["reason"]),
        ("更新", "重新运行分析脚本；Excel中的快照不会自行读取源文件。"),
    ]
    for i, (label, text) in enumerate(notes, method_row + 2):
        methods.cell(i, 1, label).font = b.font
        methods.merge_cells(start_row=i, start_column=2, end_row=i, end_column=8)
        methods.cell(i, 2, text).font = b.font
        methods.cell(i, 2).alignment = Alignment(vertical="center", wrap_text=True)
        methods.row_dimensions[i].height = 36
    methods.freeze_panes = "A7"

    metric_rank = ["小订数量", "大定", "交车锁单", "留存大定", "小转大", "直接大定"]
    stage_rank = ["平销", "首销", "小订阶段"]
    kind_rank = ["周", "月", "季", "年", "日"]
    dim_rank = ["整体", "品牌", "产品档位", "能源类型"]
    rank = lambda value, seq: seq.index(value) if value in seq else 99
    summaries = sorted(d["period_summary"], key=lambda r: (
        rank(r["metric"], metric_rank), rank(r["stage"], stage_rank), rank(r["kind"], kind_rank),
        rank(r["dimension"], dim_rank), int(re.search(r"\d+", str(r.get("transition") or "")).group()) if re.search(r"\d+", str(r.get("transition") or "")) else 0, r.get("category") or ""))
    summary_headers = ["环节", "指标", "销售阶段", "周期", "周期转换", "分类维度", "类别", "车型数", "比较对数", "不同周期", "年份数",
                       "总量比中位数", "日均比中位数", "周期P25", "周期P75", "高于100%占比", "分年日均比", "证据说明"]
    summary_row = lambda r: [r["role"], r["metric"], r["stage"], r["kind"], r["transition"], r["dimension"], r["category"], r["models"],
                             r["pairs"], r["periods"], r["years"], r["ratio"], r["daily_ratio"], r["p25"], r["p75"], r["above100"],
                             r["yearly"], r["evidence"] + "；" + r["status"]]
    widths = [10, 15, 13, 9, 16, 14, 25, 10, 12, 12, 10, 17, 17, 14, 14, 19, 45, 48]
    formats = {i: "0.0%" for i in range(11, 16)}
    b.feature("下订周期规律", summary_headers, [summary_row(r) for r in summaries if r["role"] != "锁单"], widths, formats,
              "小订、大定各指标分别统计；100%＝持平，120%＝增长20%。车型内取中位数，再车型等权取中位数。")
    b.feature("锁单周期规律", summary_headers, [summary_row(r) for r in summaries if r["role"] == "锁单"], widths, formats,
              "按实际锁单日期独立分析；不与小订、大定加总，不将锁单统一提前3天。滞后另页核验。")

    b.feature("周期比例明细",
        ["车型", "品牌", "产品档位", "能源类型", "环节", "指标", "销售阶段", "周期", "周期转换", "本期开始", "本期结束", "前期开始", "前期结束",
         "本期观测合计", "前期观测合计", "本期天数", "本期有效日", "前期天数", "前期有效日", "本期占前期", "日均占前期", "假期调休日", "状态", "数据来源", "本期负值天数", "前期负值天数"],
        [[r["model"], r["brand"], r["segment"], r["energy"], r["role"], r["metric"], r["stage"], r["kind"], r["transition"],
          _as_date(r["start"]), _as_date(r["end"]), _as_date(r["previous_start"]), _as_date(r["previous_end"]), r["current"], r["previous"],
          r["current_days"], r["current_observed"], r["previous_days"], r["previous_observed"], r["ratio"], r["daily_ratio"],
          r["holiday_adjusted_days"], r["status"], r["source"], r.get("current_negative_days", 0), r.get("previous_negative_days", 0)] for r in d["period_comparisons"]],
        [30, 12, 22, 18, 10, 15, 13, 9, 16, 16, 16, 16, 16, 19, 19, 12, 13, 12, 13, 17, 17, 14, 28, 55, 16, 16],
        {**{i: "yyyy-mm-dd" for i in range(9, 13)}, 13: "#,##0", 14: "#,##0", 19: "0.0%", 20: "0.0%"},
        "分母仅取紧邻前一自然周期；缺日、缺前期或前期为零时比例留空。")

    b.feature("锁单滞后",
        ["车型", "销售阶段", "周期批次", "大定指标", "品牌", "产品档位", "能源类型", "滞后天数", "共同日期数", "对齐开始", "对齐结束",
         "数量相关系数", "日变动相关系数", "变动对数", "解释", "大定来源", "锁单来源"],
        [[r["model"], r["stage"], r["cycle"], r["metric"], r["brand"], r["segment"], r["energy"], r["lag"], r["pairs"],
          _as_date(r.get("start")), _as_date(r.get("end")), r["level_corr"], r["change_corr"], r["change_pairs"], r["status"],
          r["deposit_source"], r["lock_source"]] for r in d["lock_lags"]],
        [30, 13, 16, 15, 12, 22, 18, 13, 15, 16, 16, 18, 20, 13, 48, 35, 55],
        {9: "yyyy-mm-dd", 10: "yyyy-mm-dd", 11: "0.000", 12: "0.000"},
        "比较大定(t)与锁单(t+0/1/2/3)。四个滞后共用同一批日期；至少14对才计算。")

    b.feature("车型分组",
        ["车型代际", "首页车型组", "品牌", "产品档位", "能源类型", "销售阶段", "可用指标"],
        [[r["model"], model_map[r["model"]]["category"], r["brand"], r["segment"], r["energy"], r["stages"], r["metrics"]] for r in d["model_classes"]],
        [32, 22, 14, 26, 22, 28, 60], {},
        "仅使用同一汇总文件的车型基本信息。增程、纯电作为混合组；映射冲突保留提示，不强行归类。")

    annual_rows = []
    for r in sorted(d.get("annual", []), key=lambda x: (x["metric"], x["stage"], x["model"], x["year"])):
        annual_rows.append([r["model"], model_map.get(r["model"], {}).get("category", "未提供"), r["metric"],
                            r["stage"], r["cycle"], r["year"],
                            *[v / 100 if v is not None else None for v in r["indexes"]],
                            r.get("status", "有效"), r["source"]])
    b.feature("年度季节性",
        ["车型", "车型组", "指标", "阶段", "批次", "年份", *[f"{i}月" for i in range(1, 13)], "状态", "数据来源"],
        annual_rows, [30, 24, 16, 14, 16, 10] + [12] * 12 + [24, 55], {i: "0.0%" for i in range(6, 18)},
        "完整自然年：各月日均/12个月日均的均值。真实零月保留；全年为零时比例留空。单年结果不能认定通用季节性。")
    holiday_rows = [[r["model"], model_map.get(r["model"], {}).get("category", "未提供"),
                     r["metric"], r["stage"], r["cycle"], r["year"], r["holiday"],
                     r["before"], r["during"], r["after"], r["control_days"], r["source"]]
                    for r in sorted(d.get("holidays", []), key=lambda x: (x["year"], x["holiday"], x["metric"], x["model"]))]
    b.feature("节假日明细",
        ["车型", "车型组", "指标", "阶段", "批次", "年份", "节假日", "节前7天/对照", "节中/对照",
         "节后7天/对照", "对照日期数", "数据来源"],
        holiday_rows, [30, 24, 16, 14, 16, 10, 16, 19, 19, 19, 15, 55], {7: "0.0%", 8: "0.0%", 9: "0.0%"},
        "仅平销；节前7天、节中、节后7天须完整。按同星期、同车型、阶段、批次和来源匹配节前后对照。")
    sources_row = methods.max_row + 3
    b.section(methods, sources_row, "日历来源", 8)
    for year, url in sorted(payload.get("calendar_sources", {}).items(), key=lambda x: str(x[0])):
        sources_row += 1
        b.note(methods, sources_row, f"{year}年：{url}")
    coverage_rows = []
    for keys, records in _groups(d["period_comparisons"], ["role", "metric", "stage", "kind"]):
        coverage_rows.append([*keys, len(records), sum(r["status"] == "有效" for r in records),
                              sum(r["status"] == "本期不完整" for r in records),
                              sum(r["status"] == "前期缺失或不完整" for r in records),
                              sum("负" in r["status"] for r in records),
                              sum(r["status"] == "前期为零" for r in records)])
    coverage_start = sources_row + 3
    b.section(methods, coverage_start, "周期覆盖与不可比原因", 10)
    b.table(methods, coverage_start + 1,
            ["环节", "指标", "阶段", "周期", "观测周期数", "有效比较", "本期不完整", "前期不完整", "含负值", "前期为零"],
            coverage_rows)
    for column in ("I", "J"):
        methods.column_dimensions[column].width = 18
    b.note(methods, methods.max_row + 2, selected_note)
    b.note(methods, methods.max_row + 1, "未形成年度或节日结果时仍保留空结果页，空白不表示销量为零。")
    b.note(methods, methods.max_row + 1, "留出验证以数据截至日确定最近完整年度，并要求相同车型、阶段、批次、来源的完整年度可比记录。")
    b.note(methods, methods.max_row + 1, "历史/验证年度同方向仅表示描述性重复，未控制价格权益、改款、门店、供给与长期趋势。")

    if "forecast" in payload:
        forecast = payload["forecast"]
        targets = {"留存大定": "净大定（留存大定）", "交车锁单": "锁单"}
        forecast_status = {"ok": "可查看回测参考", "insufficient_history": "历史训练样本不足",
                           "insufficient_backtest": "独立回测样本不足", "no_complete_target_period": "当前无完整可预测周期"}
        forecast_rows = []
        for r in forecast.get("profiles", []):
            forecast_rows.append([targets.get(r["metric"], r["metric"]), r["model"],
                model_map.get(r["model"], {}).get("category", "未提供"), r["stage"], r["grain"],
                _as_date(r.get("cutoff")), _as_date(r.get("target_start")), _as_date(r.get("target_end")),
                r.get("baseline"), r.get("estimate"), r.get("factor"), r.get("p25"), r.get("p75"),
                r.get("train_samples"), r.get("test_samples"), r.get("wape"), r.get("baseline_wape"),
                r.get("bias"), "；".join(str(v) for v in (forecast_status.get(r.get("status"), r.get("status", "")), r.get("tail_note"), {"zero_actual_denominator": "回测实际绝对值合计为0，WAPE不可计算", "zero_baseline_wape": "基准WAPE为0，相对改善不可定义"}.get(r.get("evaluation_status"))) if v), r.get("method"), r.get("cycle"), r.get("source")])
        b.feature("预测辅助",
            ["预测口径", "车型", "车型组", "阶段", "周期", "评估截至", "目标开始", "目标结束",
             "基准预测（量）", "规律预测（量）", "历史日均倍率", "倍率P25", "倍率P75", "训练样本", "共同回测样本",
             "规律WAPE", "基准WAPE", "预测偏差", "适用情况", "规则方法", "批次", "来源"],
            forecast_rows, [26, 30, 24, 12, 10, 16, 16, 16, 19, 19, 18, 15, 15, 14, 16, 16, 16, 16, 52, 58, 20, 55],
            {**{i: "yyyy-mm-dd" for i in (5, 6, 7)}, **{i: "#,##0.0" for i in (8, 9)},
             **{i: "0.0%" for i in (10, 11, 12, 15, 16, 17)}},
            "辅助销量预测：净大定与锁单分开；所有模型仅用当时可见的历史数据。P25/P75为历史倍率波动，不是预测区间。")
        tests = forecast.get("backtests", [])
        b.feature("预测回测明细",
            ["预测口径", "车型", "阶段", "周期", "评估截至", "目标开始", "目标结束", "实际量", "基准预测", "规律预测", "训练样本", "批次", "来源"],
            [[targets.get(r["metric"], r["metric"]), r["model"], r["stage"], r["grain"],
              _as_date(r.get("origin")), _as_date(r.get("target_start")), _as_date(r.get("target_end")),
              r.get("actual"), r.get("baseline"), r.get("predicted"), r.get("train_samples"), r.get("cycle"), r.get("source")] for r in tests],
            [26, 30, 12, 10, 16, 16, 16, 18, 18, 18, 14, 20, 55],
            {**{i: "yyyy-mm-dd" for i in (4, 5, 6)}, **{i: "#,##0.0" for i in (7, 8, 9)}},
            "滚动回测逐期先拟合、再核对实际。WAPE＝绝对误差合计/实际绝对值合计；相同有效窗口对照基准。")
        b.wb.move_sheet("预测辅助", offset=1-b.wb.sheetnames.index("预测辅助"))
        b.note(methods, methods.max_row + 2, "研究宗旨：各种规律均用于辅助销量预测；净大定（源字段留存大定）与交车锁单分别校准。")
        for note in forecast.get("notes", []):
            b.note(methods, methods.max_row + 1, str(note))
        b.note(methods, methods.max_row + 1, "交互网页与Excel由同一次分析生成，支持筛选、图形对比、预测试算与明细导出。")

    _atomic_save(b.wb, output_path, log)
    return output_path


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("evidence", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    payload = json.loads(args.evidence.read_text(encoding="utf-8-sig"))
    print(json.dumps({"output": str(export_workbook(payload, args.output))}, ensure_ascii=False))