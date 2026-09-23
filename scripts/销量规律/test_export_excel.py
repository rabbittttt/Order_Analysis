"""Portable Excel export regression tests; every workbook stays in a temporary directory."""
import importlib.util
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

spec = importlib.util.spec_from_file_location("sales_export_test", Path(__file__).with_name("export_excel_openpyxl.py"))
e = importlib.util.module_from_spec(spec)
spec.loader.exec_module(e)


def fixture():
    base = dict(model="测试车型", metric="留存大定", stage="平销", cycle="平销", source="原始逐日")
    period = dict(base, brand="测试品牌", segment="测试档位", energy="纯电", role="大定", kind="月",
                  transition="8→9月", start=date(2025, 9, 1), end=date(2025, 9, 30),
                  previous_start=date(2025, 8, 1), previous_end=date(2025, 8, 31),
                  current=300, previous=310, current_days=30, previous_days=31,
                  current_observed=30, previous_observed=31, current_negative_days=0, previous_negative_days=0,
                  ratio=300 / 310, daily_ratio=1, status="有效", holiday_adjusted_days=0)
    return dict(sample=True, sha256="test-only", calendar_years=[2025], calendar_sources={},
                audit={"测试记录": 1}, results=dict(
                    coverage=[dict(metric="留存大定", stage="平销", models=1, observations=365,
                                   dates=365, start=date(2025, 1, 1), end=date(2025, 12, 31))],
                    model_classes=[dict(model="测试车型", brand="测试品牌", segment="测试档位", energy="纯电",
                                        stages="平销", metrics="留存大定、小订数量")],
                    weekly=[dict(base, week=date(2025, 8, 4), indexes=[100] * 7, ratio=1)],
                    monthly=[dict(base, year=2025, month=2, ratio=None, status="月初基线为零")],
                    annual=[dict(base, year=2025, indexes=[100, 0] + [100] * 10, status="有效")],
                    holidays=[dict(base, year=2025, holiday="测试假期", before=1.1, during=1.4, after=.9, control_days=20)],
                    lifecycle=[], period_comparisons=[period, dict(period, metric="小订数量", stage="小订阶段", role="小订")],
                    period_summary=[], lock_lags=[]))


class ExportTests(unittest.TestCase):
    def test_auto_includes_retained_deposits_and_manual_filter_is_explicit(self):
        rows = fixture()["results"]["period_comparisons"]
        selected, omitted = e._selected_metrics(rows, {})
        self.assertEqual(set(selected), {"留存大定", "小订数量"})
        grouping = e._choose_grouping(fixture()["results"]["model_classes"], {})
        model_map = {r["model"]: r for r in grouping["entries"]}
        self.assertEqual({r["metric"] for r in e._report_series(rows, grouping["categories"], model_map)}, set(selected))
        self.assertEqual(e._selected_metrics(rows, {"mainMetrics": ["留存大定"]}), (["留存大定"], ["小订数量"]))
        with self.assertRaisesRegex(ValueError, "不存在"):
            e._selected_metrics(rows, {"mainMetrics": ["不存在的指标"]})

    def test_export_keeps_holiday_annual_zero_and_retained_metric(self):
        messages = []
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "report.xlsx"
            e.export_workbook(fixture(), target, log=messages.append)
            wb = load_workbook(target)
            try:
                self.assertIn("规律汇总", wb.sheetnames)
                self.assertEqual(wb["年度季节性"]["H8"].value, 0)  # February, not a missing month.
                self.assertEqual(wb["节假日明细"]["I8"].value, 1.4)
                self.assertEqual(wb["节假日明细"]["I8"].number_format, "0.0%")
                self.assertIsNone(wb["月内明细"]["G8"].value)
                self.assertEqual(wb["月内明细"]["H8"].value, "月初基线为零")
                main_values = {cell.value for row in wb["分析总览"] for cell in row if cell.value is not None}
                self.assertIn("留存大定", main_values)
                self.assertEqual(wb["规律汇总"].freeze_panes, "C8")
                self.assertTrue(wb["节假日明细"].tables)
                self.assertLess(len(wb._cell_styles), 80)
            finally:
                wb.close()
            self.assertEqual([x.name for x in Path(folder).iterdir()], ["report.xlsx"])
        self.assertTrue(any("工作表" in message and "行" in message for message in messages))
        self.assertTrue(any("Excel保存完成" in message for message in messages))

    def test_negative_observation_is_not_in_complete_month_profile(self):
        payload = fixture()
        r = payload["results"]["period_comparisons"][0]
        r.update(current_negative_days=1, status="本期含负值", ratio=None, daily_ratio=None)
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "report.xlsx"
            e.export_workbook(payload, target)
            wb = load_workbook(target)
            try:
                row = next(row for row in wb["月份证据"].iter_rows(min_row=8, values_only=True)
                           if row[0] == "全部车型" and row[1] == "留存大定" and row[3] == 9)
                self.assertIsNone(row[4])
                self.assertEqual(row[6], 0)
            finally:
                wb.close()

    def test_forecast_targets_are_separate_and_backtests_are_exported(self):
        payload=fixture()
        base=dict(model="测试车型",stage="平销",cycle="平销",source="真实逐日",grain="日",
                  cutoff="2025-09-30",target_start="2025-10-01",target_end="2025-10-01",
                  baseline=100,estimate=120,factor=1.2,p25=1.1,p75=1.3,train_samples=12,
                  test_samples=20,wape=.12,baseline_wape=.2,bias=.01,status="ok",method="历史日均倍率")
        payload["forecast"]={"profiles":[dict(base,metric="留存大定"),dict(base,metric="交车锁单",estimate=140)],
                             "backtests":[dict(base,metric="留存大定",origin="2025-09-30",actual=115,predicted=120)],"notes":[]}
        with tempfile.TemporaryDirectory() as folder:
            target=Path(folder)/"report.xlsx"
            e.export_workbook(payload,target)
            wb=load_workbook(target)
            try:
                self.assertEqual(len(wb.sheetnames),17)
                self.assertEqual(wb.sheetnames[1],"预测辅助")
                self.assertEqual(wb["预测辅助"]["A8"].value,"净大定（留存大定）")
                self.assertEqual(wb["预测辅助"]["A9"].value,"锁单")
                self.assertEqual(wb["预测辅助"]["J8"].value,120)
                self.assertEqual(wb["预测辅助"]["J9"].value,140)
                self.assertEqual(wb["预测回测明细"]["H8"].value,115)
            finally:wb.close()

    def test_failed_save_preserves_previous_file_and_only_removes_own_temp(self):
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "report.xlsx"
            unrelated = Path(folder) / "notes.txt"
            target.write_bytes(b"previous-success")
            unrelated.write_text("keep", encoding="utf-8")
            def broken_save(path):
                Path(path).write_bytes(b"partial")
                raise OSError("disk full")
            wb = Workbook()
            with patch.object(wb, "save", side_effect=broken_save):
                with self.assertRaisesRegex(OSError, "disk full"):
                    e._atomic_save(wb, target, lambda message: None)
            self.assertEqual(target.read_bytes(), b"previous-success")
            self.assertEqual(unrelated.read_text(encoding="utf-8"), "keep")
            self.assertEqual(sorted(p.name for p in Path(folder).iterdir()), ["notes.txt", "report.xlsx"])

    def test_locked_target_keeps_previous_report(self):
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "report.xlsx"
            target.write_bytes(b"previous-success")
            wb = Workbook()
            wb.active["A1"] = "test"
            with patch.object(e.os, "replace", side_effect=PermissionError("locked")):
                with self.assertRaisesRegex(PermissionError, "关闭已打开"):
                    e._atomic_save(wb, target, lambda message: None)
            self.assertEqual(target.read_bytes(), b"previous-success")
            self.assertEqual([x.name for x in Path(folder).iterdir()], ["report.xlsx"])


if __name__ == "__main__":
    unittest.main()
