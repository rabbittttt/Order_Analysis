"""Read forecast inputs from one workbook while retaining original source identities."""
from __future__ import annotations

import base64
import gzip
import json
from contextlib import contextmanager
from contextvars import ContextVar
from pathlib import Path

from openpyxl import load_workbook

from core.excel import WorkbookItem, WorkbookStore

SUMMARY_NAME = "鸿蒙智行销量数据汇总.xlsx"
INDEX_SHEET = "数据来源目录"
SNAPSHOT_SHEET = "_预测缓存"
ACTIVE_SUMMARY = ContextVar("forecast_summary", default=None)


def write_summary_snapshot(workbook, payload):
    """Store the resolved forecast payload, not copied source worksheets."""
    if SNAPSHOT_SHEET in workbook.sheetnames:
        del workbook[SNAPSHOT_SHEET]
    packed = base64.b64encode(
        gzip.compress(json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))
    ).decode("ascii")
    sheet = workbook.create_sheet(SNAPSHOT_SHEET)
    sheet.append(["格式", "版本"])
    sheet.append(["gzip+base64", 1])
    for offset in range(0, len(packed), 30000):
        sheet.cell(sheet.max_row + 1, 1, packed[offset:offset + 30000])
    sheet.sheet_state = "veryHidden"


def read_summary_snapshot(workbook):
    if SNAPSHOT_SHEET not in workbook.sheetnames:
        return None
    sheet = workbook[SNAPSHOT_SHEET]
    if sheet["A2"].value != "gzip+base64":
        raise ValueError("销量预测缓存格式无法识别，请运行main重新生成")
    packed = "".join(str(sheet.cell(row, 1).value or "") for row in range(3, sheet.max_row + 1))
    try:
        return json.loads(gzip.decompress(base64.b64decode(packed)).decode("utf-8"))
    except Exception as exc:
        raise ValueError("销量预测缓存损坏，请运行main重新生成") from exc


class SheetView:
    def __init__(self, sheet, title):
        self._sheet, self.title = sheet, title

    def __getattr__(self, name):
        return getattr(self._sheet, name)

    def __getitem__(self, key):
        return self._sheet[key]


class WorkbookView:
    def __init__(self, sheets):
        self.worksheets = list(sheets)
        self.sheetnames = [sheet.title for sheet in self.worksheets]

    def __getitem__(self, name):
        return next(sheet for sheet in self.worksheets if sheet.title == name)

    def close(self):
        pass  # The enclosing summary_scope owns the workbook.


def input_path(default):
    active = ACTIVE_SUMMARY.get()
    return active["path"] if active else default


def open_input(path, **kwargs):
    active = ACTIVE_SUMMARY.get()
    if active and Path(path) == active["path"]:
        return active["inputs"]
    return load_workbook(path, **kwargs)


@contextmanager
def summary_scope(path, workbook=None):
    path = Path(path)
    owns_workbook = workbook is None
    if owns_workbook:
        workbook = load_workbook(path, read_only=False, data_only=True)
    token = None
    try:
        snapshot = read_summary_snapshot(workbook)
        if snapshot is not None:
            inputs = [sheet for sheet in workbook.worksheets if sheet.title != SNAPSHOT_SHEET]
            store = WorkbookStore(path.parent)
            store.items = []
            active = {
                "path": path,
                "inputs": WorkbookView(inputs),
                "store": store,
                "snapshot": snapshot,
            }
            token = ACTIVE_SUMMARY.set(active)
            yield active
            return
        if INDEX_SHEET not in workbook.sheetnames:
            raise ValueError(f"{path.name}缺少数据来源目录，请运行main重新生成")
        groups, inputs = {}, list(workbook.worksheets)
        seen = set()
        for kind, filename, original, stored, *_ in workbook[INDEX_SHEET].iter_rows(min_row=2, values_only=True):
            if not stored:
                continue
            if stored not in workbook.sheetnames or (filename, original) in seen:
                raise ValueError(f"汇总目录重复或缺少Sheet：{filename}/{original}")
            seen.add((filename, original))
            view = SheetView(workbook[stored], original)
            if kind == "订单":
                groups.setdefault(filename, []).append(view)
            elif kind in {"历史", "映射"}:
                inputs = [sheet for sheet in inputs if sheet.title != original]
                inputs.append(view)
        store = WorkbookStore(path.parent)
        store.items = [WorkbookItem(Path(name), WorkbookView(sheets)) for name, sheets in groups.items()]
        active = {"path": path, "inputs": WorkbookView(inputs), "store": store}
        token = ACTIVE_SUMMARY.set(active)
        yield active
    finally:
        if token is not None:
            ACTIVE_SUMMARY.reset(token)
        if owns_workbook:
            workbook.close()
