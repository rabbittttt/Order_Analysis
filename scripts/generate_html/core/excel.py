from __future__ import annotations

import base64
import logging
import re
import weakref
from statistics import median
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.styles.numbers import is_date_format


GRAIN_MAP = {"时": "hour", "天": "day", "周": "week", "月": "month"}
GRAIN_LABELS = {value: key for key, value in GRAIN_MAP.items()}
LOGGER = logging.getLogger(__name__)


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\u2011", "-")).strip()


def compact_text(value: Any) -> str:
    return re.sub(r"\s+", "", clean_text(value))


def compact_identifier(value: Any) -> str:
    """Normalize separators used inconsistently in workbook and Sheet names."""
    return re.sub(r"[\s_\-—–:：/\\（）()【】\[\]]+", "", clean_text(value)).lower()


def display_period(value: Any, number_format: str = "") -> str:
    if isinstance(value, (datetime, date)):
        return format_excel_value(value, number_format or "yyyy-mm-dd")
    return clean_text(value)


def safe_number(value: Any, default: float = 0.0) -> float:
    if value in (None, ""):
        return default
    if isinstance(value, str):
        text = clean_text(value)
        negative = text.startswith("(") and text.endswith(")")
        percent = "%" in text or "％" in text
        text = text.replace(",", "").replace("，", "").replace("%", "").replace("％", "")
        text = re.sub(r"(?:人民币|RMB|CNY|元|万元|亿元|￥|¥|\$)", "", text, flags=re.IGNORECASE).strip(" ()")
        try:
            number = float(text)
            if negative:
                number = -number
            return number / 100 if percent else number
        except ValueError:
            return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def cell_number(cell: Any, rate: bool = False, default: float = 0.0) -> float:
    """Read a number while respecting percent-like strings and cell formats."""
    value = safe_number(cell.value, default)
    if not rate or cell.value in (None, ""):
        return value
    if isinstance(cell.value, str):
        text = clean_text(cell.value)
        return value / 100 if "%" not in text and "％" not in text and abs(value) > 1.5 else value
    number_format = clean_text(getattr(cell, "number_format", ""))
    if "%" not in number_format and abs(value) > 1.5:
        return value / 100
    return value


def _format_precision(number_format: str, marker: str = "") -> int:
    section = number_format.split(";", 1)[0]
    if marker and marker in section:
        section = section.split(marker, 1)[0]
    match = re.search(r"[0#](?:\.([0#]+))?[^0#]*$", section)
    return len(match.group(1) or "") if match else 0


def _format_date(value: date | datetime, number_format: str) -> str:
    fmt = number_format.lower().replace("\\", "")
    has_time = any(token in fmt for token in ("h", "s", "am/pm"))
    has_day = "d" in fmt
    has_month = "m" in fmt
    has_year = "y" in fmt
    separator = "/" if "/" in fmt else "." if "." in fmt else "-"
    if has_year and has_month and has_day:
        result = value.strftime(f"%Y{separator}%m{separator}%d")
    elif has_year and has_month:
        result = value.strftime(f"%Y{separator}%m")
    elif has_month and has_day:
        result = value.strftime(f"%m{separator}%d")
    else:
        result = value.isoformat()
    if has_time and isinstance(value, datetime):
        time_text = value.strftime("%H:%M:%S" if "s" in fmt else "%H:%M")
        result = f"{result} {time_text}" if has_year or has_month or has_day else time_text
    return result


def format_excel_value(value: Any, number_format: str = "General") -> str:
    """Render common Excel formats without assuming fixed precision or date syntax."""
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return _format_date(value, number_format)
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if not isinstance(value, (int, float)):
        return clean_text(value)
    fmt = clean_text(number_format or "General")
    if fmt.lower() == "general" or fmt == "@":
        return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
    if is_date_format(fmt):
        return str(value)
    if "%" in fmt:
        precision = _format_precision(fmt, "%")
        return f"{value * 100:,.{precision}f}%"
    precision = _format_precision(fmt)
    grouping = "," in fmt.split(";", 1)[0]
    rendered = f"{value:,.{precision}f}" if grouping else f"{value:.{precision}f}"
    currency = next((symbol for symbol in ("￥", "¥", "$", "€", "£") if symbol in fmt), "")
    return f"{currency}{rendered}"


def format_excel_cell(cell: Any) -> str:
    return format_excel_value(cell.value, getattr(cell, "number_format", "General"))


def safe_rate(numerator: Any, denominator: Any) -> float:
    den = safe_number(denominator)
    return safe_number(numerator) / den if den else 0.0


def sheet_subject(name: str) -> str:
    value = clean_text(name)
    patterns = [
        r"净(?:大定|小订)by[时天周月](?:第\s*\d+\s*期)?图表$",
        r"by[时天周月](?:第\s*\d+\s*期)?图表$",
        r"by[时天周月](?:第\s*\d+\s*期)?$",
        r"_(SKU|组合|日度退订|选配退订|小订分时退订|选配渗透率排名|付费选配占比|选配渗透率)$",
    ]
    for pattern in patterns:
        value = re.sub(pattern, "", value, flags=re.IGNORECASE)
    # Generic analysis sheets may append a free-form suffix. Keep the suffix as
    # the analysis name while resolving the sheet back to its known group,
    # brand, or generation subject.
    subject_prefixes = [
        # “总计/汇总”可以是代际名称本身的一部分，例如“问界 M9 2026款总计”。
        # 只清理它后面的分析后缀，不要把它误当作汇总周期截掉。
        r"^(.+?20\d{2}\s*款(?:\s*(?:总计|汇总))?)\s*[_\-—]\s*.+$",
        # 代号型代际（无款号），例如“问界 F2N”。
        r"^([\u4e00-\u9fff]{1,3}界\s*[A-Za-z0-9]+(?:\s*20\d{2}\s*款)?(?:\s*(?:总计|汇总))?)\s*[_\-—]\s*.+$",
        r"^(鸿蒙智行)\s*[_\-—]\s*.+$",
        r"^([\u4e00-\u9fff]{1,3}界)\s*[_\-—]\s*.+$",
    ]
    for pattern in subject_prefixes:
        match = re.match(pattern, value, flags=re.IGNORECASE)
        if match:
            value = match.group(1)
            break
    return clean_text(value.rstrip("_"))


def is_aggregate_generation(name: str) -> bool:
    """以“总计/汇总”结尾的代际是多个代际的聚合视图，品牌级统计时须排除以免重复。"""
    return clean_text(name).endswith(("总计", "汇总"))


def subject_type(name: str) -> str:
    if compact_text(name) == "鸿蒙智行":
        return "group"
    if re.search(r"20\d{2}\s*款", name):
        return "generation"
    if re.fullmatch(r"[\u4e00-\u9fff]{1,3}界\s*[A-Za-z0-9]+", compact_text(name)):
        return "generation"
    return "brand"


def subject_parent(name: str, kind: str) -> str | None:
    if kind == "brand":
        return "鸿蒙智行"
    if kind != "generation":
        return None
    match = re.match(r"^([\u4e00-\u9fff]{1,3}界)", compact_text(name))
    return match.group(1) if match else None


def grain_from_sheet(name: str) -> str | None:
    match = re.search(r"by([时天周月])", name)
    return GRAIN_MAP.get(match.group(1)) if match else None


@dataclass
class WorkbookItem:
    path: Path
    workbook: Any


class WorkbookStore:
    def __init__(self, input_dir: Path):
        self.input_dir = input_dir
        self.items: list[WorkbookItem] = []
        self.chart_assets: dict[str, str] = {}
        self.load_errors: list[str] = []
        self._find_all_cache: dict[str, list[WorkbookItem]] = {}
        self._subject_sheets_cache: dict[tuple[Any, ...], list[tuple[WorkbookItem, Any]]] = {}

    def load(self, exclude_names: Iterable[str] = ()) -> None:
        if not self.input_dir.exists():
            raise FileNotFoundError(f"输入目录不存在: {self.input_dir}")
        self._find_all_cache.clear()
        self._subject_sheets_cache.clear()
        excluded = set(exclude_names)
        paths = sorted(path for path in self.input_dir.glob("*.xlsx") if not path.name.startswith("~$") and path.name not in excluded)
        if not paths:
            LOGGER.warning("输入目录中没有找到 .xlsx 文件: %s", self.input_dir)
        for path in paths:
            try:
                workbook = load_workbook(path, data_only=True, read_only=False)
                self.items.append(WorkbookItem(path=path, workbook=workbook))
                LOGGER.debug("已读取 Excel: %s（%d 个 Sheet）", path.name, len(workbook.sheetnames))
            except Exception as exc:
                message = f"Excel读取失败: {path.name}: {exc}"
                self.load_errors.append(message)
                LOGGER.exception(message)

    def close(self) -> None:
        for item in self.items:
            item.workbook.close()

    def find(self, keyword: str) -> WorkbookItem | None:
        return next((item for item in self.items if keyword in item.path.name), None)

    def find_all(self, keyword: str) -> list[WorkbookItem]:
        cached = self._find_all_cache.get(keyword)
        if cached is None:
            cached = [item for item in self.items if keyword in item.path.name]
            self._find_all_cache[keyword] = cached
        return cached

    def find_subject_sheet(
        self,
        keyword: str,
        subject: str,
        grain: str | None = None,
        suffix: str | tuple[str, ...] | list[str] | None = None,
        include_charts: bool = False,
    ) -> tuple[WorkbookItem, Any] | None:
        matches = self.find_subject_sheets(keyword, subject, grain, suffix, include_charts)
        return matches[0] if matches else None

    def find_subject_sheets(
        self,
        keyword: str,
        subject: str,
        grain: str | None = None,
        suffix: str | tuple[str, ...] | list[str] | None = None,
        include_charts: bool = False,
    ) -> list[tuple[WorkbookItem, Any]]:
        """Return every matching sheet; callers can retain multiple launch phases."""
        suffixes = () if suffix is None else ((suffix,) if isinstance(suffix, str) else tuple(suffix))
        cache_key = (keyword, compact_text(subject), grain, suffixes, include_charts)
        cached = self._subject_sheets_cache.get(cache_key)
        if cached is not None:
            return cached
        items = self.find_all(keyword)
        if not items:
            return []
        if len(items) > 1:
            LOGGER.warning("关键词“%s”匹配到多个Excel，将逐个查找: %s", keyword, [item.path.name for item in items])
        result: list[tuple[WorkbookItem, Any]] = []
        for item in items:
            for sheet in item.workbook.worksheets:
                if not include_charts and "图表" in sheet.title:
                    continue
                if compact_text(sheet_subject(sheet.title)) != compact_text(subject):
                    continue
                if grain and grain_from_sheet(sheet.title) != grain:
                    continue
                if suffixes:
                    normalized_title = compact_identifier(sheet.title)
                    if not any(compact_identifier(candidate) in normalized_title for candidate in suffixes):
                        continue
                result.append((item, sheet))
                LOGGER.debug("Sheet映射: 关键词=%s, 主体=%s -> %s / %s", keyword, subject, item.path.name, sheet.title)
        self._subject_sheets_cache[cache_key] = result
        return result

    def matching_sheet_names(self, keyword: str, subject: str) -> list[str]:
        """Return candidate Sheet names for diagnostics when a module mapping fails."""
        items = self.find_all(keyword)
        target = compact_text(subject)
        return [
            f"{item.path.name} / {sheet.title}"
            for item in items
            for sheet in item.workbook.worksheets
            if compact_text(sheet_subject(sheet.title)) == target
        ]

    def all_sheet_names(self) -> Iterable[tuple[str, str]]:
        for item in self.items:
            for name in item.workbook.sheetnames:
                yield item.path.name, name

    def find_chart_data(
        self, keyword: str, subject: str, grain: str
    ) -> tuple[WorkbookItem, Any, dict[str, Any]] | None:
        """Read one subject block, preferring its own/parent chart sheet before broader sheets."""
        item = self.find(keyword)
        if not item:
            return None
        target = compact_text(subject)
        kind = subject_type(subject)
        preferred_owner = subject_parent(subject, kind) if kind == "generation" else subject
        sheets = [
            sheet for sheet in item.workbook.worksheets
            if "图表" in sheet.title and grain_from_sheet(sheet.title) == grain
        ]
        sheets.sort(key=lambda sheet: compact_text(sheet_subject(sheet.title)) != compact_text(preferred_owner))
        for sheet in sheets:
            for row in range(1, sheet.max_row + 1):
                if compact_text(sheet.cell(row, 1).value) != target:
                    continue
                data = self._parse_chart_block(item, sheet, row, target)
                if data:
                    return item, sheet, data
        return None

    def find_chart_blocks(
        self, keyword: str, subject: str, grain: str
    ) -> tuple[WorkbookItem, Any, list[dict[str, Any]]] | None:
        """Return all blocks from the subject's chart sheet; generations return only their own block."""
        item = self.find(keyword)
        if not item:
            return None
        kind = subject_type(subject)
        owner = subject_parent(subject, kind) if kind == "generation" else subject
        target_owner = compact_text(owner)
        for sheet in item.workbook.worksheets:
            if "图表" not in sheet.title or grain_from_sheet(sheet.title) != grain:
                continue
            if compact_text(sheet_subject(sheet.title)) != target_owner:
                continue
            blocks = []
            for row in range(1, sheet.max_row + 1):
                block_subject = clean_text(sheet.cell(row, 1).value)
                if not block_subject:
                    continue
                if kind == "generation" and compact_text(block_subject) != compact_text(subject):
                    continue
                data = self._parse_chart_block(item, sheet, row, compact_text(block_subject))
                if data:
                    blocks.append(data)
            if blocks:
                return item, sheet, blocks
        exact = self.find_chart_data(keyword, subject, grain)
        return (exact[0], exact[1], [exact[2]]) if exact else None

    def _parse_chart_block(
        self, item: WorkbookItem, sheet: Any, row: int, target: str
    ) -> dict[str, Any] | None:
        header_row = row + 1
        columns: list[int] = []
        periods: list[str] = []
        for col in range(3, sheet.max_column + 1):
            cell = sheet.cell(header_row, col)
            value = display_period(cell.value, cell.number_format)
            if not value or value == "总计":
                continue
            columns.append(col)
            periods.append(value)
        if not columns:
            return None
        series = []
        totals = [0.0] * len(columns)
        for data_row in range(header_row + 1, sheet.max_row + 1):
            label = clean_text(sheet.cell(data_row, 2).value)
            next_subject = clean_text(sheet.cell(data_row, 1).value)
            if next_subject:
                break
            if not label:
                if series:
                    break
                continue
            if label == "总计":
                totals = [cell_number(sheet.cell(data_row, col)) for col in columns]
                break
            values = [cell_number(sheet.cell(data_row, col), rate=True) for col in columns]
            series.append({"name": label, "values": values})
        if not series:
            return None
        block_subject = clean_text(sheet.cell(row, 1).value)
        image_key = self._extract_chart_image(item, sheet, row, compact_text(target))
        return {
            "subject": block_subject,
            "headline": clean_text(sheet.cell(header_row, 2).value),
            "periods": periods,
            "series": series,
            "totals": totals,
            "image_key": image_key,
        }

    def _extract_chart_image(self, item: WorkbookItem, sheet: Any, subject_row: int, target: str) -> str | None:
        """Extract the image embedded in a subject block once and keep it as a shared HTML asset."""
        asset_key = f"{item.path.name}|{sheet.title}|{target}"
        if asset_key in self.chart_assets:
            return asset_key
        next_subject_row = None
        for row in range(subject_row + 1, sheet.max_row + 1):
            if clean_text(sheet.cell(row, 1).value):
                next_subject_row = row
                break

        if next_subject_row is None:
            # The final subject has no real lower boundary. Infer one from the
            # preceding block spacing instead of relying on sheet.max_row: Excel
            # images may be anchored one or more rows below the last data cell.
            subject_rows = [
                row for row in range(1, subject_row + 1)
                if clean_text(sheet.cell(row, 1).value)
            ]
            spacings = [
                later - earlier
                for earlier, later in zip(subject_rows, subject_rows[1:])
                if later > earlier
            ]
            if spacings:
                block_span = max(1, int(round(median(spacings))))
            else:
                # A single-block sheet still gets a finite search window. Cover
                # its existing cells plus the immediately following anchor row.
                block_span = max(30, sheet.max_row - subject_row + 2)
            next_subject_row = subject_row + block_span

        images = sorted(
            getattr(sheet, "_images", []),
            key=lambda image: getattr(getattr(getattr(image, "anchor", None), "_from", None), "row", -1),
        )
        for image in images:
            anchor = getattr(image, "anchor", None)
            marker = getattr(anchor, "_from", None)
            image_row = getattr(marker, "row", -1) + 1
            if not subject_row <= image_row < next_subject_row:
                continue
            try:
                image_format = clean_text(getattr(image, "format", "png")).lower() or "png"
                mime = "jpeg" if image_format in {"jpg", "jpeg"} else image_format
                encoded = base64.b64encode(image._data()).decode("ascii")
                self.chart_assets[asset_key] = f"data:image/{mime};base64,{encoded}"
                return asset_key
            except (OSError, ValueError, AttributeError):
                return None
        return None


_METRIC_SHEET_CACHE: "weakref.WeakKeyDictionary[Any, dict[str, dict[str, Any]]]" = weakref.WeakKeyDictionary()


def parse_metric_sheet(sheet: Any) -> dict[str, dict[str, Any]]:
    """Parse one metric sheet once per build.

    Dashboard modules only read the returned structure.  Caching by Worksheet
    identity avoids reparsing the same order/lock sheet for every subject and
    every period while allowing closed workbooks to be garbage-collected.
    """
    cached = _METRIC_SHEET_CACHE.get(sheet)
    if cached is not None:
        return cached
    headers = [display_period(sheet.cell(1, col).value, sheet.cell(1, col).number_format) for col in range(4, sheet.max_column + 1)]
    # 原表周期列完整保留，包括“总计/汇总”；KPI Card 直接使用原表对应列的值。
    periods = [header for header in headers if header]
    result = {period: {"metrics": {}, "structures": {}} for period in periods}
    current_metric = ""
    current_stat = ""
    for row in range(2, sheet.max_row + 1):
        metric = clean_text(sheet.cell(row, 1).value)
        stat = clean_text(sheet.cell(row, 2).value)
        category = clean_text(sheet.cell(row, 3).value)
        if metric:
            current_metric = metric
        if stat:
            current_stat = stat
        if not current_metric or not category:
            continue
        for offset, period in enumerate(headers, start=4):
            if period not in result:
                continue
            cell = sheet.cell(row, offset)
            if current_stat == "数量" and category == "数量":
                result[period]["metrics"][current_metric] = cell_number(cell)
            elif current_stat == "占比":
                result[period]["structures"].setdefault(current_metric, []).append(
                    {"label": category, "share": cell_number(cell, rate=True), "count": None}
                )
            elif current_stat == "数量":
                rows = result[period]["structures"].setdefault(current_metric, [])
                target = next((entry for entry in rows if entry["label"] == category), None)
                if target:
                    target["count"] = cell_number(cell)
    _METRIC_SHEET_CACHE[sheet] = result
    return result
