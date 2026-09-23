from __future__ import annotations

import re
from datetime import date, datetime

from openpyxl.utils.datetime import from_excel

from core.components import kpi, section, table
from core.excel import WorkbookStore, cell_number, clean_text, display_period, safe_number, safe_rate
from core.models import Dashboard, SourceRef, Subject


def _phase_number(sheet_name: str) -> int | None:
    match = re.search(r"第\s*(\d+)\s*期", sheet_name)
    return int(match.group(1)) if match else None


def _date_label(cell) -> str:
    value = cell.value
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)):
        try:
            return from_excel(value, epoch=cell.parent.parent.epoch).strftime("%Y-%m-%d")
        except (TypeError, ValueError, OverflowError):
            pass
    return display_period(value, cell.number_format).split(" ", 1)[0]


def _read_hourly(sheet) -> tuple[list[dict], str]:
    rows: list[dict] = []
    start_label = ""
    metric_rows = {
        clean_text(sheet.cell(row, 1).value): row
        for row in range(3, sheet.max_row + 1)
        if clean_text(sheet.cell(row, 1).value)
    }
    for col in range(2, sheet.max_column + 1):
        time_cell = sheet.cell(2, col)
        time_value = time_cell.value
        if time_value in (None, ""):
            continue
        if isinstance(time_value, (datetime, date)):
            period = time_value.strftime("%Y-%m-%d")
            hour = time_value.hour if isinstance(time_value, datetime) else col - 2
            timestamp_label = time_value.strftime("%Y-%m-%d %H:%M") if isinstance(time_value, datetime) else f"{period} {int(hour):02d}:00"
        elif isinstance(time_value, (int, float)):
            timestamp = from_excel(time_value, epoch=sheet.parent.epoch)
            period, hour = timestamp.strftime("%Y-%m-%d"), timestamp.hour
            timestamp_label = timestamp.strftime("%Y-%m-%d %H:%M")
        else:
            period = display_period(time_value, time_cell.number_format).split(" ", 1)[0]
            header = clean_text(sheet.cell(1, col).value)
            digits = "".join(char for char in header if char.isdigit())
            hour = max(safe_number(digits, col - 1) - 1, 0)
            timestamp_label = f"{period} {int(hour):02d}:00"
        start_label = start_label or timestamp_label

        def metric(*names: str) -> float:
            row = next((metric_rows[name] for name in names if name in metric_rows), None)
            return cell_number(sheet.cell(row, col)) if row else 0.0

        rows.append({
            "period": period, "hour": hour,
            "orders": metric("当日大定数量"), "net": metric("当日留存大定数量", "当日净大定数量"),
            "small": metric("当日小订转大定数量"), "direct": metric("当日直接大定数量"),
            "lock": metric("当日交车锁单数量"),
        })
    return rows, start_label


class LaunchRhythmModule:
    id = "launch_rhythm"
    label = "首销节奏"

    def build(self, store: WorkbookStore, subject: Subject) -> Dashboard | None:
        views = {}
        sources: list[SourceRef] = []
        hourly_matches = store.find_subject_sheets("首销期订单节奏", subject.name, "hour")
        for grain in ("day", "week"):
            matches = store.find_subject_sheets("首销期订单节奏", subject.name, grain)
            if not matches:
                continue
            matches.sort(key=lambda found: (_phase_number(found[1].title) is None, _phase_number(found[1].title) or 0))
            phases = []
            for index, (item, sheet) in enumerate(matches, start=1):
                phase_number = _phase_number(sheet.title)
                hourly = self._matching_hourly(hourly_matches, phase_number, len(matches))
                phase = self._build_phase(item, sheet, phase_number or index, hourly, len(matches) > 1 or phase_number is not None)
                if phase:
                    phases.append(phase)
                    sources.extend(phase["sources"])
            if not phases:
                continue
            if len(phases) == 1:
                phase = phases[0]
                views[grain] = {"periods": [phase["period_label"]], "default_period": phase["period_label"], "pages": {phase["period_label"]: phase["page"]}}
                continue
            pages = {"全部阶段": self._combined_page(phases)}
            pages.update({phase["selector_label"]: phase["page"] for phase in phases})
            selectors = ["全部阶段", *[phase["selector_label"] for phase in phases]]
            views[grain] = {"periods": selectors, "default_period": phases[-1]["selector_label"], "pages": pages}

        unique_sources: list[SourceRef] = []
        seen = set()
        for source in sources:
            key = (source.file, source.sheet, source.note)
            if key not in seen:
                seen.add(key)
                unique_sources.append(source)
        return Dashboard(self.id, subject.id, views, unique_sources) if views else None

    @staticmethod
    def _matching_hourly(hourly_matches, phase_number: int | None, phase_count: int):
        if phase_number is not None:
            return next((found for found in hourly_matches if _phase_number(found[1].title) == phase_number), None)
        return hourly_matches[0] if phase_count == 1 and hourly_matches else None

    def _build_phase(self, item, sheet, phase_number: int, hourly, show_phase: bool) -> dict | None:
        source = SourceRef(item.path.name, sheet.title, "首销期订单节奏")
        period_columns = [col for col in range(2, sheet.max_column + 1) if clean_text(sheet.cell(1, col).value) and clean_text(sheet.cell(1, col).value) != "总计"]
        periods = [clean_text(sheet.cell(1, col).value) for col in period_columns]
        if not periods:
            return None
        metrics = {}
        for row in range(3, sheet.max_row + 1):
            label = clean_text(sheet.cell(row, 1).value)
            if label:
                is_rate = "率" in label or "进度" in label
                display_label = label.replace("净大定", "留存大定")
                metrics[display_label] = [cell_number(sheet.cell(row, col), rate=is_rate) for col in period_columns]
        date_labels = [_date_label(sheet.cell(2, col)) for col in period_columns]
        date_labels = [label for label in date_labels if label]
        start_date = date_labels[0] if date_labels else ""
        end_date = date_labels[-1] if date_labels else periods[-1]

        def latest(name: str) -> float:
            values = metrics.get(name, [0])
            return values[-1] if values else 0

        cumulative_order = sum(metrics.get("当日大定数量", []))
        cumulative_net = sum(metrics.get("当日留存大定数量", []))
        cumulative_small = latest("累计小订转大定数量") or latest("累计小订转大数量")
        cumulative_small_rate = latest("累计小订转化率") or latest("累计小转大率")
        cumulative_small_base = safe_rate(cumulative_small, cumulative_small_rate) if cumulative_small_rate > 0 else 0
        cumulative_direct = latest("累计直接大定数量")
        matrix_rows = [[name, *[f"{value * 100:.1f}%" if ("率" in name or "进度" in name) else value for value in values]] for name, values in metrics.items()]
        page = {
            "kpis": [
                kpi("累计大定", cumulative_order, "单", periods[-1], "blue"),
                kpi("累计留存大定", cumulative_net, "单", periods[-1], "green"),
                kpi("累计小转大率", cumulative_small_rate * 100, "%", "累计小转大 ÷ 总小订", "purple"),
                kpi("累计直接大定占比", safe_rate(cumulative_direct, cumulative_order) * 100, "%", "累计大定来源", "blue"),
            ],
            "sections": [
                section("launch_composite", f"第{phase_number}期 · 首销订单来源" if show_phase else "首销订单来源", {"periods": periods, "metrics": metrics}, f"{start_date or '日期缺失'} 至 {end_date or '日期缺失'} · 各期独立计算" if show_phase else "累计直接大定 / 累计小订转大 / 来源进度", source=source),
                section("matrix", f"第{phase_number}期 · 完整指标" if show_phase else "首销期完整指标", table(["首销指标", *periods], matrix_rows), f"全部 {len(metrics)} 项指标 × {len(periods)} 个周期", source=source),
            ],
        }
        phase_sources = [source]
        if hourly:
            hourly_item, hourly_sheet = hourly
            hourly_rows, hourly_start = _read_hourly(hourly_sheet)
            if hourly_rows:
                hourly_source = SourceRef(hourly_item.path.name, hourly_sheet.title, "分时首销节奏")
                hourly_meta = (f"首销开启 {hourly_start} · 仅对应第{phase_number}期" if hourly_start else f"仅对应第{phase_number}期") if show_phase else (f"首销开启 {hourly_start} · 识别高峰时段与大定波动" if hourly_start else "识别高峰时段与大定波动")
                page["sections"].insert(1, section("launch_hourly", f"第{phase_number}期 · 分时首销节奏" if show_phase else "分时首销节奏", hourly_rows, hourly_meta, source=hourly_source))
                phase_sources.append(hourly_source)
        range_label = f"{start_date}—{end_date}" if start_date and end_date else end_date
        return {
            "phase": phase_number, "start_date": start_date, "end_date": end_date,
            "period_label": end_date or periods[-1],
            "selector_label": f"第{phase_number}期 · {range_label}" if range_label else f"第{phase_number}期",
            "days": len(periods), "orders": cumulative_order, "net": cumulative_net,
            "small": cumulative_small, "small_rate": cumulative_small_rate,
            "small_base": cumulative_small_base, "direct": cumulative_direct,
            "page": page, "metrics": metrics, "periods": periods,
            "source": source, "sources": phase_sources,
        }

    @staticmethod
    def _combined_page(phases: list[dict]) -> dict:
        total_orders = sum(phase["orders"] for phase in phases)
        total_net = sum(phase["net"] for phase in phases)
        total_small = sum(phase["small"] for phase in phases)
        total_direct = sum(phase["direct"] for phase in phases)
        comparable_phases = [phase for phase in phases if phase["small_base"] > 0]
        combined_small_rate = safe_rate(
            sum(phase["small"] for phase in comparable_phases),
            sum(phase["small_base"] for phase in comparable_phases),
        )
        summary_rows = []
        for index, phase in enumerate(phases):
            gap: int | str = "—"
            if index and phases[index - 1]["end_date"] and phase["start_date"]:
                try:
                    previous = datetime.strptime(phases[index - 1]["end_date"], "%Y-%m-%d").date()
                    current = datetime.strptime(phase["start_date"], "%Y-%m-%d").date()
                    gap = max((current - previous).days - 1, 0)
                except ValueError:
                    pass
            summary_rows.append([f"第{phase['phase']}期", phase["start_date"] or "日期缺失", phase["end_date"] or "日期缺失", phase["days"], gap, phase["orders"], phase["net"]])
        sections = [section(
            "table", "阶段概览",
            table(["阶段", "开始日期", "结束日期", "周期数", "距上期空档(天)", "累计大定", "累计留存大定"], summary_rows,
                  formats=["text", "text", "text", "number", "number", "number", "number"]),
            "按绝对日期展示阶段边界；空档不计为0，也不参与趋势连线", source=phases[0]["source"],
        )]
        for phase in phases:
            sections.append(section(
                "launch_composite", f"第{phase['phase']}期 · 首销订单来源",
                {"periods": phase["periods"], "metrics": phase["metrics"]},
                f"{phase['start_date'] or '日期缺失'} 至 {phase['end_date'] or '日期缺失'} · 与其他期次断开显示",
                source=phase["source"],
            ))
        return {
            "kpis": [
                kpi("多期累计大定", total_orders, "单", f"共{len(phases)}期", "blue"),
                kpi("多期累计留存大定", total_net, "单", f"共{len(phases)}期", "green"),
                kpi("多期小转大率", combined_small_rate * 100, "%", "按各期总小订加权计算", "purple"),
                kpi("多期直接大定占比", safe_rate(total_direct, total_orders) * 100, "%", "分期累计后计算", "blue"),
            ],
            "sections": sections,
        }
