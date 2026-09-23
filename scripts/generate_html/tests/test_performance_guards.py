from __future__ import annotations

import os
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook

from core.excel import parse_metric_sheet
from modules.option_fee import _option_layout
from tools.refresh_sales_forecast_data import (
    REFRESH_CODE_DEPENDENCIES,
    find_header_row,
    output_is_current,
    scan_header_rows,
)


class PerformanceGuardTests(unittest.TestCase):
    def test_forecast_refresh_skips_only_when_every_dependency_is_older(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            mapping = root / "mapping.xlsx"
            output = root / "output.xlsx"
            source.write_bytes(b"source")
            mapping.write_bytes(b"mapping")
            output.write_bytes(b"PK\x03\x04generated")

            newest_code = max(path.stat().st_mtime_ns for path in REFRESH_CODE_DEPENDENCIES)
            dependency_time = newest_code - 2_000_000_000
            os.utime(source, ns=(dependency_time, dependency_time))
            os.utime(mapping, ns=(dependency_time, dependency_time))
            output_time = newest_code + 2_000_000_000
            os.utime(output, ns=(output_time, output_time))
            self.assertTrue(output_is_current(source, output, mapping))

            source_time = output_time + 2_000_000_000
            os.utime(source, ns=(source_time, source_time))
            self.assertFalse(output_is_current(source, output, mapping))

    def test_header_scan_can_be_reused_without_changing_detection(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["说明"])
        sheet.append(["车型", "D1", "D2"])
        scanned = scan_header_rows(sheet)

        row, headers, days = find_header_row(sheet, require_days=True, scanned_rows=scanned)

        self.assertEqual(row, 2)
        self.assertEqual(headers[:3], ["车型", "D1", "D2"])
        self.assertEqual(days, [(1, 1), (2, 2)])

    def test_metric_and_option_layout_results_are_cached_per_sheet(self):
        workbook = Workbook()
        metric = workbook.active
        metric.cell(1, 4, "26WK35")
        metric.cell(2, 1, "大定")
        metric.cell(2, 2, "数量")
        metric.cell(2, 3, "数量")
        metric.cell(2, 4, 100)
        self.assertIs(parse_metric_sheet(metric), parse_metric_sheet(metric))

        option = workbook.create_sheet("选配金")
        option.cell(3, 3, "26-08")
        option.cell(4, 3, "免费")
        option.cell(4, 4, "0-5000")
        option.cell(5, 2, "订单数量")
        option.cell(5, 3, 10)
        option.cell(5, 4, 20)
        first = _option_layout(option)
        self.assertEqual(first, (3, 4, 5))
        self.assertIs(first, _option_layout(option))


if __name__ == "__main__":
    unittest.main()
