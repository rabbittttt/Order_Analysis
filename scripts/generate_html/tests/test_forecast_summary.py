from copy import copy
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side

from core.forecast_summary import (
    ACTIVE_SUMMARY,
    INDEX_SHEET,
    SNAPSHOT_SHEET,
    SUMMARY_NAME,
    summary_scope,
    write_summary_snapshot,
)
from core.excel import WorkbookStore
from modules.sales_forecast import _read_actual_profiles, _read_model_mapping, _read_model_master, _read_stage_windows
from tools.refresh_sales_forecast_data import append_forecast_inputs, forecast_signature


class ForecastSummaryTests(unittest.TestCase):
    def test_fast_import_keeps_values_and_number_formats_without_blank_styles(self):
        with TemporaryDirectory() as temp:
            source, mapping, orders = self.make_inputs(Path(temp))
            raw = load_workbook(source)
            raw.active['Z500'].font = Font(bold=True)
            raw.save(source)
            raw.close()
            full, fast = Workbook(), Workbook()
            append_forecast_inputs(full, source, mapping, orders, '2026-01-05')
            append_forecast_inputs(fast, source, mapping, orders, '2026-01-05', preserve_layout=False)
            for name in ('车型汇总', '车型基本信息'):
                facts = lambda sheet: {cell.coordinate:(cell.value, cell.number_format) for row in sheet for cell in row if cell.value is not None}
                self.assertEqual(facts(full[name]), facts(fast[name]))
            self.assertEqual(full['车型汇总'].max_row, 500)
            self.assertEqual(fast['车型汇总'].max_row, 2)
            full.close()
            fast.close()

    def test_imported_row_and_column_styles_survive_cross_workbook_round_trip(self):
        with TemporaryDirectory() as temp:
            root = Path(temp)
            source, mapping, orders = self.make_inputs(root)
            attributes = ("font", "fill", "border", "alignment", "protection",
                          "number_format", "quotePrefix", "pivotButton")
            for path, color in ((source, "FF123456"), (mapping, "FF654321")):
                book = load_workbook(path)
                sheet = book.active
                for dimension in (sheet.column_dimensions["G"], sheet.row_dimensions[8]):
                    dimension.font = Font(name="Arial", bold=True, color=color)
                    dimension.fill = PatternFill("solid", fgColor=color)
                    dimension.border = Border(left=Side(style="thin", color=color))
                    dimension.alignment = Alignment(horizontal="right", wrap_text=True)
                    dimension.protection = Protection(locked=False, hidden=True)
                    dimension.number_format = '0.0000" custom"'
                    dimension.quotePrefix = True
                    dimension.pivotButton = True
                    dimension.hidden = True
                    dimension.outlineLevel = 2
                    dimension.collapsed = True
                sheet.column_dimensions["G"].width = 27.5
                sheet.row_dimensions[8].height = 36
                sheet.column_dimensions["H"].width = 19
                book.save(path)
                book.close()

            summary = root / "summary.xlsx"
            book = Workbook()
            # Populate a different destination style registry, including a custom format.
            book.active["A1"].number_format = '0.00" target"'
            append_forecast_inputs(book, source, mapping, orders, "2026-01-05")
            for name in ("车型汇总", "车型基本信息"):
                self.assertIs(book[name].column_dimensions["G"].parent, book[name])
                self.assertIs(book[name].row_dimensions[8].parent, book[name])
            book.save(summary)
            book.close()

            # This non-streaming reopen is the operation that failed on intranet files.
            with summary_scope(summary) as active:
                for path, name in ((source, "车型汇总"), (mapping, "车型基本信息")):
                    original = load_workbook(path)
                    try:
                        actual = active["inputs"][name]
                        expected = original.active
                        for left, right in (
                            (actual.column_dimensions["G"], expected.column_dimensions["G"]),
                            (actual.row_dimensions[8], expected.row_dimensions[8]),
                        ):
                            for attribute in attributes + ("hidden", "outlineLevel", "collapsed"):
                                self.assertEqual(copy(getattr(left, attribute)), copy(getattr(right, attribute)), attribute)
                        self.assertEqual(actual.column_dimensions["G"].width, 27.5)
                        self.assertEqual(actual.column_dimensions["H"].width, 19)
                        self.assertEqual(actual.row_dimensions[8].height, 36)
                        self.assertEqual(actual["B2"].value, expected["B2"].value)
                        self.assertEqual(actual["B2"].number_format, expected["B2"].number_format)
                    finally:
                        original.close()

    def test_compact_snapshot_round_trip_needs_no_source_sheets(self):
        with TemporaryDirectory() as temp:
            path = Path(temp) / SUMMARY_NAME
            book = Workbook()
            book.active.title = "汇总说明"
            book.create_sheet(INDEX_SHEET)
            payload = {"views": {"week": {"value": 123}}, "sources": []}
            write_summary_snapshot(book, payload)
            book.save(path)
            book.close()

            with summary_scope(path) as active:
                self.assertEqual(active["snapshot"], payload)
                self.assertEqual(active["store"].items, [])
                self.assertNotIn(SNAPSHOT_SHEET, active["inputs"].sheetnames)
            check = load_workbook(path)
            self.assertEqual(check[SNAPSHOT_SHEET].sheet_state, "veryHidden")
            check.close()

    def test_summary_in_order_directory_does_not_become_a_raw_subject(self):
        with TemporaryDirectory() as temp:
            root = Path(temp)
            for filename, title in [(SUMMARY_NAME, "汇总内部Sheet"), ("锁单选配比例.xlsx", "问界 M9 2026款by天")]:
                book = Workbook()
                book.active.title = title
                book.save(root / filename)
                book.close()
            store = WorkbookStore(root)
            try:
                store.load(exclude_names={SUMMARY_NAME})
                self.assertEqual([item.path.name for item in store.items], ["锁单选配比例.xlsx"])
                self.assertEqual(list(store.all_sheet_names()), [("锁单选配比例.xlsx", "问界 M9 2026款by天")])
            finally:
                store.close()

    def make_inputs(self, root):
        source, mapping = root / "source.xlsx", root / "mapping.xlsx"
        orders = root / "orders"
        orders.mkdir()
        book = Workbook()
        sheet = book.active
        sheet.title = "车型汇总"
        sheet.append(["车型", "开始大定日期", "小转大结束日期", "小订开始日期", "小订结束日期"])
        sheet.append(["历史车", datetime(2026, 1, 3), datetime(2026, 1, 4), datetime(2026, 1, 1), datetime(2026, 1, 2)])
        book.save(source)
        book.close()
        book = Workbook()
        sheet = book.active
        sheet.title = "车型基本信息"
        sheet.append(["历史传播名", "订单分析代际名", "原始表简称/别名"])
        sheet.append(["历史车", "问界 M9 2026款", "简称"])
        book.save(mapping)
        book.close()
        book = Workbook()
        sheet = book.active
        sheet.title = "问界 M9 2026款by天"
        sheet.append([None, "D1", "D2"])
        sheet.append(["时间", datetime(2026, 1, 3), datetime(2026, 1, 4)])
        sheet.append(["当日大定数量", 0, None])
        sheet["B2"].number_format = "yyyy/mm/dd"
        book.save(orders / "鸿蒙智行首销期订单节奏.xlsx")
        book.close()
        return source, mapping, orders

    def test_summary_readers_work_without_original_files_and_keep_zero_blank_date(self):
        with TemporaryDirectory() as temp:
            root = Path(temp)
            source, mapping, orders = self.make_inputs(root)
            summary = root / "summary.xlsx"
            book = Workbook()
            append_forecast_inputs(book, source, mapping, orders, "2026-01-05")
            book.save(summary)
            book.close()
            # Originals made inaccessible: all forecast readers must use summary.
            source.rename(root / "source.saved")
            mapping.rename(root / "mapping.saved")
            (orders / "鸿蒙智行首销期订单节奏.xlsx").rename(orders / "orders.saved")
            with summary_scope(summary) as active:
                with patch("modules.sales_forecast.load_workbook", side_effect=AssertionError("external read")):
                    self.assertEqual(_read_model_mapping()["简称"], "问界 M9 2026款")
                    self.assertIn("历史车", _read_model_master())
                    _, windows = _read_stage_windows()
                    self.assertEqual(next(iter(windows.values()))["launch_date"], "2026-01-03")
                    profiles, sources = _read_actual_profiles(active["store"])
                sheet = active["store"].find("首销期订单节奏").workbook.worksheets[0]
                self.assertEqual(sheet.title, "问界 M9 2026款by天")
                self.assertEqual(sheet["B3"].value, 0)
                self.assertIsNone(sheet["C3"].value)
                self.assertEqual(sheet["B2"].number_format, "yyyy/mm/dd")
                self.assertEqual(profiles[0]["days"][0]["gross"], 0)
                self.assertTrue(sources)
            self.assertIsNone(ACTIVE_SUMMARY.get())

    def test_refresh_signature_tracks_added_removed_changed_files_and_date(self):
        with TemporaryDirectory() as temp:
            source, mapping, orders = self.make_inputs(Path(temp))
            signature = lambda day="2026-01-05": forecast_signature(source, mapping, orders, day)
            original = signature()
            self.assertNotEqual(original, signature("2026-01-06"))
            new_file = orders / "新增锁单选配比例.xlsx"
            new_file.write_bytes(b"test")
            self.assertNotEqual(original, signature())
            new_file.unlink()
            self.assertEqual(original, signature())
            with source.open("ab") as handle:
                handle.write(b"changed")
            self.assertNotEqual(original, signature())

    def test_broken_source_directory_is_not_silently_ignored(self):
        with TemporaryDirectory() as temp:
            path = Path(temp) / "summary.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = INDEX_SHEET
            sheet.append(["类别", "原始文件", "原始Sheet", "汇总Sheet"])
            sheet.append(["订单", "锁单选配比例.xlsx", "车型by天", "不存在"])
            book.save(path)
            book.close()
            with self.assertRaisesRegex(ValueError, "缺少Sheet"):
                with summary_scope(path):
                    pass
            self.assertIsNone(ACTIVE_SUMMARY.get())
