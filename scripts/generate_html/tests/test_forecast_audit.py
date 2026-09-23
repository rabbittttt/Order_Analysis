from datetime import date, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from core.forecast_summary import INDEX_SHEET, summary_scope
from modules.sales_forecast import _attach_actual_shapes, _fill_sparse_dates, _read_actual_profiles, _read_history
from tools.refresh_sales_forecast_data import append_forecast_views


class ForecastAuditTests(unittest.TestCase):
    def test_history_curves_prefer_exact_event_name_over_shared_generation(self):
        with TemporaryDirectory() as temp:
            path = Path(temp) / 'history.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.title = '预测基准总表'
            sheet.append(['传播名', '总大定'])
            sheet.append(['历史车', 100])
            sheet.append(['历史车Pro', 30])
            curve = book.create_sheet('小转大当日数量')
            curve.append(['传播名', 'D1', 'D2'])
            curve.append(['历史车', 80, 20])
            curve.append(['历史车Pro', 20, 10])
            book.save(path)
            book.close()
            with patch('modules.sales_forecast.input_path', return_value=path), patch('modules.sales_forecast._read_model_mapping', return_value={'历史车':'同一代际','历史车pro':'同一代际'}), patch('modules.sales_forecast._read_model_master', return_value={}):
                _, history = _read_history()
            self.assertEqual(history[0]['daily_small'], [80,20])
            self.assertEqual(history[1]['daily_small'], [20,10])

    def test_hourly_reference_keeps_actual_clock_hours_and_excludes_unfinished_d1(self):
        history = [{'model':'test', 'launch_date':'2026-09-01', 'daily_orders':[100]}]
        profiles = [{'model':'test','hourly_days':[{'date':'2026-09-01','hours':[{'hour':20,'gross':60},{'hour':18,'gross':40}]}]}]
        with patch('modules.sales_forecast._read_model_mapping', return_value={}):
            _attach_actual_shapes(history, profiles, today=date(2026,9,2))
            curve = history[0]['hourly_curve']
            self.assertEqual(len(curve),24)
            self.assertEqual(curve[:18], [0]*18)
            self.assertEqual(curve[18:21], [.4,.4,1])
            _attach_actual_shapes(history, profiles, today=date(2026,9,1))
            self.assertEqual(history[0]['hourly_curve'], [])

    def test_sparse_dates_only_fill_omitted_columns_inside_observed_interval(self):
        rows = [{'date':'2026-09-01','orders':8}, {'date':'2026-09-03','orders':None}, {'date':'2026-09-04','orders':3}]
        result = _fill_sparse_dates(rows, ('orders',))
        self.assertEqual([row['orders'] for row in result], [8, 0, None, 3])
        self.assertTrue(result[1]['absent_as_zero'])
        self.assertEqual(len(_fill_sparse_dates(rows[:1], ('orders',))), 1)
        self.assertEqual(_fill_sparse_dates([rows[0], rows[0]], ('orders',)), [rows[0], rows[0]])

    def test_hourly_missing_components_remain_missing_not_zero(self):
        book = Workbook()
        sheet = book.active
        sheet.title = '问界 M9 2026款by时'
        sheet.append(['时间', datetime(2026,9,1,8), datetime(2026,9,1,9)])
        sheet.append(['当日大定数量', 50, 50])
        sheet.append(['当日小订转大定数量', 30, None])
        item = SimpleNamespace(path=Path('首销期订单节奏.xlsx'), workbook=book)
        store = SimpleNamespace(find=lambda key: item if key == '首销期订单节奏' else None)
        with patch('modules.sales_forecast._read_model_mapping', return_value={}):
            profiles, _ = _read_actual_profiles(store)
        bucket = profiles[0]['hourly_days'][0]
        self.assertEqual(bucket['gross'], 100)
        self.assertIsNone(bucket['small_to_big'])
        self.assertIsNone(bucket['direct'])
        book.close()

    def test_borrowed_summary_scope_does_not_reload_or_close_workbook(self):
        book = Workbook()
        book.active.title = INDEX_SHEET
        book.active.append(['类别', '文件', 'Sheet', '内部Sheet'])
        with patch('core.forecast_summary.load_workbook', side_effect=AssertionError('unnecessary reload')), patch.object(book, 'close') as close:
            with summary_scope(Path('memory.xlsx'), workbook=book) as active:
                self.assertEqual(active['inputs'].sheetnames, [INDEX_SHEET])
            close.assert_not_called()
        book.close()

    def test_materialized_summary_matches_resolved_priority_and_does_not_duplicate_cancel(self):
        with TemporaryDirectory() as temp:
            path = Path(temp) / 'summary.xlsx'
            book = Workbook()
            book.active.title = INDEX_SHEET
            book.active.append(['类别', '文件', 'Sheet', '内部Sheet'])
            master = book.create_sheet('车型基本信息')
            master.append(['订单分析代际名', '历史传播名'])
            master.append(['测试车', '历史车'])
            book.create_sheet('字段说明')
            book.save(path)
            data = {'targets':[{'name':'测试车', 'end_date':'2026-09-02', 'small_start_date':'2026-08-01'}],
                    'small_order_history':[], 'steady_history':[],
                    'actuals':[{'model':'测试车', 'total_small':100,
                                'cancel_days':[{'date':'2026-09-01', 'cancel':9, 'cancel_rate':.09}],
                                'days':[{'date':'2026-09-01','gross':20,'cancel':5,'_field_sources':{'cancel':'history'}},
                                        {'date':'2026-09-02','gross':30}]}]}
            payload = {'views':{'week':{'pages':{'预测方案':{'workspace':{'data':data}}}}}}
            dashboard = SimpleNamespace(views=payload['views'], to_dict=lambda:payload)
            with patch('modules.sales_forecast.SalesForecastModule._build_from_sources', return_value=dashboard), patch('tools.refresh_sales_forecast_data.load_workbook', side_effect=AssertionError('unnecessary reload')):
                append_forecast_views(path, '2026-09-02', workbook=book)
            records = list(book['小订及退订逐日'].values)
            row = dict(zip(records[0], records[1]))
            self.assertEqual(row['累计退订'], 5)
            self.assertEqual(row['累计退订率'], .05)
            self.assertEqual(row['退订来源'], '小订及首销数据整理')
            orders = list(book['当前订单逐日'].values)
            self.assertNotIn('退订', orders[0])
            self.assertEqual(orders[-1][-1], '当日快照（未结束）')
            self.assertFalse(any(name.startswith('源_') for name in book.sheetnames))
            book.close()
