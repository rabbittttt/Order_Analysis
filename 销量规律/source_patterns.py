"""Extract explainable patterns directly from the consolidated sales workbook.

The workbook is deliberately treated as the only source of truth here.  This
module does not follow the ``来源文件`` columns to open another workbook and it
does not rebuild retained orders or locks from component counts.  The three
families returned by :func:`build_source_insights` are source-oriented:

``early_launch``
    Long-form records from ``首销D1D2预测指标``.  Each record has a stable
    ``metric``/``feature``/``value`` shape and retains the original field and
    source row number. The source already contains D1/D2 proportions and status
    columns, so they are exposed as-is; a few clearly named aliases are added
    without replacing the source fields.

``curves``
    Long-form D-day snapshots from the five requested ratio curves and the
    related quantity/completion curves.  Cumulative sheets are never summed.
    Each row carries the model's dynamic anchor date, denominator (when it is
    available in ``车型基本信息``), and the original source sheet.

``hourly``
    Long-form hourly records from ``小订by时`` and ``首销by时``.  Shares are
    calculated only for a complete, unique, non-negative 0--23 sequence with a
    positive daily total.  Missing, duplicate, negative and all-zero input is
    retained with an explicit status and no fabricated share.

The module intentionally returns ordinary JSON-friendly values (dates become
ISO strings).  It has no dependency on the report renderer and is safe to use
when the source workbook is missing or partially populated.
"""

from __future__ import annotations

import math
import re
from collections import defaultdict
from datetime import date, datetime, timedelta
from numbers import Real
from pathlib import Path
from typing import Any, Iterable, Mapping

try:  # Imported lazily by callers in environments that do not need Excel.
    import openpyxl
except Exception:  # pragma: no cover - the empty-source path is still useful.
    openpyxl = None  # type: ignore[assignment]


MODEL_SHEET = "车型基本信息"
EARLY_SHEET = "首销D1D2预测指标"
HOURLY_SHEETS = ("小订by时", "首销by时")

# Curves used by the current pattern pages, plus their count/completion inputs.
# Keeping this list at sheet-name level avoids hard-coding any model names.
CURVE_SHEETS = (
    "每日直接大定占比",
    "累计直接大定占比",
    "累计小订转化率",
    "小订期累计退订率",
    "首销期累计退订率",
    "小订当日数量",
    "小订期退订当日数量",
    "小订累计完成度",
    "小转大当日数量",
    "直接大定当日数量",
    "大定当日数量",
    "首销期退订当日数量",
    "小转大累计完成度",
    "直接大定累计完成度",
    "大定累计完成度",
)

# These four sheets describe the small-order window itself.  Other curves
# whose names mention 小订 (for example 累计小订转化率) describe conversion
# after launch and must use the 首销 anchor.
SMALL_ORDER_CURVE_SHEETS = frozenset({
    "小订当日数量",
    "小订期退订当日数量",
    "小订累计完成度",
    "小订期累计退订率",
})

_D_HEADER = re.compile(r"^\s*D\s*(\d+)\s*$", re.IGNORECASE)
_DATE_WORDS = ("日期", "日期", "时间", "时间")
_MODEL_WORDS = ("代际名", "车型", "模型", "车型名称", "车型代际")


def _json_value(value: Any) -> Any:
    """Convert openpyxl scalar values to values accepted by ``json.dumps``."""

    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, (str, int, bool)) or value is None:
        return value
    # Decimal and other scalar-like values occur in some Excel exports.
    if isinstance(value, Real):
        number = float(value)
        return number if math.isfinite(number) else None
    return str(value)


def _date_value(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None or isinstance(value, bool):
        return None
    text = str(value).strip()
    if not text:
        return None
    # Excel exports sometimes use a timestamp with a space or ``T``.
    text = text.replace("/", "-")
    for candidate in (text, text[:10]):
        try:
            return date.fromisoformat(candidate)
        except ValueError:
            continue
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _number(value: Any) -> float | int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Real):
        number = float(value)
    else:
        text = str(value).strip().replace(",", "")
        if not text:
            return None
        percent = text.endswith("%")
        if percent:
            text = text[:-1].strip()
        try:
            number = float(text)
        except (TypeError, ValueError):
            return None
        if percent:
            number /= 100.0
    if not math.isfinite(number):
        return None
    if number.is_integer() and abs(number) < 10**15:
        return int(number)
    return number


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _normalise_header(value: Any, index: int) -> str:
    text = _text(value)
    return text or f"列{index + 1}"


def _headers_and_rows(ws: Any) -> tuple[list[str], list[dict[str, Any]]]:
    iterator = ws.iter_rows(values_only=True)
    try:
        first = next(iterator)
    except StopIteration:
        return [], []
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(first):
        header = _normalise_header(value, index)
        count = seen.get(header, 0)
        seen[header] = count + 1
        if count:
            header = f"{header}__{count + 1}"
        headers.append(header)
    rows: list[dict[str, Any]] = []
    for values in iterator:
        if not any(value is not None and _text(value) != "" for value in values):
            continue
        row = {
            header: _json_value(values[index] if index < len(values) else None)
            for index, header in enumerate(headers)
        }
        rows.append(row)
    return headers, rows


def _find_header(headers: Iterable[str], *candidates: str,
                 contains: bool = False) -> str | None:
    values = list(headers)
    for candidate in candidates:
        if candidate in values:
            return candidate
    if contains:
        for header in values:
            if any(candidate and candidate in header for candidate in candidates):
                return header
    return None


def _model_field(headers: Iterable[str]) -> str | None:
    headers = list(headers)
    exact = _find_header(headers, "代际名", "车型", "车型名称", "模型", "车型代际")
    if exact:
        return exact
    return next((header for header in headers if any(word in header for word in _MODEL_WORDS)), None)


def _secondary_field(headers: Iterable[str]) -> str | None:
    return _find_header(headers, "二级代际名", "二级车型", "二级代际", contains=True)


def _first_present(row: Mapping[str, Any], *names: str) -> Any:
    for name in names:
        if name in row and row.get(name) not in (None, ""):
            return row.get(name)
    return None


def _model_name(row: Mapping[str, Any], headers: Iterable[str]) -> str:
    # Keep the same display key as the main daily panel: a non-empty second-
    # level generation is the model identity; the primary generation is only
    # a fallback when no second-level value exists.
    model = _first_present(
        row, "二级代际名", "订单分析代际名", "代际名", "车型", "车型名称", "模型", "车型代际"
    )
    if model in (None, ""):
        field = _model_field(headers)
        model = row.get(field) if field else None
    return _text(model)


def _secondary_name(row: Mapping[str, Any], headers: Iterable[str]) -> str:
    field = _secondary_field(headers)
    return _text(row.get(field)) if field else ""


def _metadata_maps(rows: list[dict[str, Any]], headers: list[str]) -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
    """Build exact and conservative alias maps from ``车型基本信息``."""

    # Collect all exact aliases before deciding whether an alias is safe.  A
    # primary generation may legitimately contain several second-level models;
    # silently keeping the first row would mix dates, denominators and model
    # attributes.
    aliases: dict[str, list[dict[str, Any]]] = defaultdict(list)
    canonical: dict[str, Any] = {}
    for row in rows:
        primary = _text(_first_present(row, "代际名", "订单分析代际名"))
        secondary = _secondary_name(row, headers)
        display = secondary or primary
        if not display:
            continue
        item = dict(row)
        item["model"] = display
        item["model_level"] = display
        item["primary_model"] = primary
        item["secondary_model"] = secondary
        for key in (secondary, primary):
            key = _text(key)
            if key:
                aliases[key].append(item)
                compact = _compact(key)
                if compact != key:
                    aliases[compact].append(item)
        canonical[display] = item
    maps: dict[str, dict[str, Any]] = {}
    for alias, candidates in aliases.items():
        distinct: dict[tuple[Any, ...], dict[str, Any]] = {}
        for item in candidates:
            identity = (
                item.get("primary_model"), item.get("secondary_model"),
                item.get("小订开始"), item.get("小订结束"),
                item.get("首销开始"), item.get("首销结束"),
            )
            distinct.setdefault(identity, item)
        if len(distinct) == 1:
            maps[alias] = next(iter(distinct.values()))
        else:
            maps[alias] = {"_mapping_status": "车型映射冲突/未明确版本"}
    return maps, canonical


def _compact(value: Any) -> str:
    return re.sub(r"\s+", "", _text(value)).replace("（", "(").replace("）", ")")


def _lookup_metadata(model: str, maps: Mapping[str, dict[str, Any]]) -> dict[str, Any]:
    model = _text(model)
    if not model:
        return {"_mapping_status": "车型映射未匹配"}
    # Exact (or whitespace-normalized exact) aliases only.  There is no
    # containment fallback: a short generation name must not guess a version.
    for key in (model, _compact(model)):
        if key in maps:
            return maps[key]
    return {"_mapping_status": "车型映射未匹配"}


def _metadata_value(meta: Mapping[str, Any], *names: str) -> Any:
    return _first_present(meta, *names)


def _meta_dimensions(meta: Mapping[str, Any]) -> dict[str, Any]:
    return {
        "category": _metadata_value(meta, "产品档位", "车型分类", "档位") or "",
        "brand": _metadata_value(meta, "品牌") or "",
        "energy": _metadata_value(meta, "能源类型", "能源") or "",
        "publish_type": _metadata_value(meta, "发布类型", "上市类型") or "",
    }


def _anchor(meta: Mapping[str, Any], stage: str) -> tuple[str | None, str]:
    if stage == "小订":
        value = _metadata_value(meta, "小订开始", "小订开始日期", "小订起始日期")
        return (_date_value(value).isoformat() if _date_value(value) else None, "小订开始")
    value = _metadata_value(meta, "首销开始", "首销开始日期", "首销起始日期")
    return (_date_value(value).isoformat() if _date_value(value) else None, "首销开始")


def _model_status(meta: Mapping[str, Any]) -> str:
    value = _metadata_value(meta, "销售阶段", "小订阶段", "首销阶段", "阶段")
    return _text(value)


def _is_open_status(value: Any) -> bool:
    text = _text(value)
    return any(token in text for token in ("进行中", "未结束", "当日", "未来"))


def _unit_for_feature(feature: str, sheet: str = "") -> str:
    text = f"{feature}{sheet}"
    if "/" in feature or any(token in text for token in ("率", "占比", "完成度", "%")):
        return "比例"
    if any(token in text for token in ("数量", "大定", "小订", "退订", "锁单")):
        return "数量"
    if any(token in text for token in ("状态", "说明", "来源", "类型", "档位", "品牌", "能源")):
        return "文本"
    return "未知"


def _status_for_value(value: Any, *, source_cutoff: date | None = None,
                      actual_date: date | None = None, meta: Mapping[str, Any] | None = None) -> str:
    if meta and meta.get("_mapping_status"):
        return str(meta["_mapping_status"])
    if source_cutoff and actual_date and actual_date >= source_cutoff:
        return "未来数据"
    number = _number(value)
    if number is not None and number < 0:
        return "需核对（负值）"
    if meta and _is_open_status(_model_status(meta)):
        return "未结束/进行中"
    if value is None or _text(value) == "":
        return "原始空白"
    return "有效"


def _early_metric(feature: str) -> str:
    """Map source features to a human-readable, non-mixing metric family."""

    text = _text(feature)
    if "小转大" in text:
        return "小转大"
    if "直接大" in text:
        return "直接大定"
    if "大定" in text:
        return "大定"
    if "退订" in text:
        return "退订"
    if "状态" in text or "说明" in text:
        return "源表核对"
    return "首销节奏"


def _parse_cutoff(rows: list[dict[str, Any]]) -> date | None:
    for row in rows:
        text = " ".join(_text(value) for value in row.values())
        if "数据判定日期" in text or "判定日期" in text:
            for value in row.values():
                parsed = _date_value(value)
                if parsed:
                    return parsed
    return None


def _note_lines(rows: list[dict[str, Any]]) -> list[str]:
    notes: list[str] = []
    for row in rows:
        values = [_text(value) for value in row.values() if _text(value)]
        if not values:
            continue
        notes.append("｜".join(values))
    return notes


def _raw_row(row: Mapping[str, Any]) -> dict[str, Any]:
    return {str(key): _json_value(value) for key, value in row.items()}


def _ratio_denominator(meta: Mapping[str, Any], curve: str) -> tuple[Any, str | None]:
    """Return the existing total used by a curve, never a recomputed total."""

    if "小订转化" in curve:
        names = ("总小订", "总小订数", "小订总量")
    elif "退订率" in curve:
        names = ("总小订", "总小订数", "小订总量")
    elif "直接大定" in curve:
        names = ("总直接大定", "直接大定", "总直接大定数")
    elif "小转大" in curve:
        names = ("总小转大", "小转大", "总小转大数")
    elif "大定" in curve:
        names = ("总大定", "大定", "总大定数")
    else:
        names = ()
    for name in names:
        value = _metadata_value(meta, name)
        if value not in (None, ""):
            return value, name
    return None, None


def _is_ratio_curve(sheet: str) -> bool:
    return any(token in sheet for token in ("占比", "率", "完成度"))


def _early_features(rows: list[dict[str, Any]], headers: list[str],
                    meta_maps: Mapping[str, dict[str, Any]],
                    source_cutoff: date | None) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    model_headers = {"代际名", "二级代际名", "车型", "车型名称", "模型", "车型代际"}
    static_headers = {"产品档位", "发布类型", "品牌", "能源类型"}
    for source_row, row in enumerate(rows, start=2):
        model = _model_name(row, headers)
        if not model:
            continue
        meta = _lookup_metadata(model, meta_maps)
        stage = "首销"
        anchor_date, anchor_type = _anchor(meta, stage)
        dimensions = _meta_dimensions(meta)
        common = {
            "model": model,
            "model_level": _text(row.get("二级代际名")) or model,
            **dimensions,
            "stage": stage,
            "source": EARLY_SHEET,
            "source_sheet": EARLY_SHEET,
            "source_row": source_row,
            "anchor_date": anchor_date,
            "anchor_type": anchor_type,
        }
        # Keep every original field, including status and quality explanation.
        for header in headers:
            if header in model_headers or header in static_headers:
                continue
            value = row.get(header)
            number = _number(value)
            feature_value: Any = number if number is not None else value
            status = _status_for_value(value, source_cutoff=source_cutoff,
                                       meta=meta)
            denominator = None
            denominator_field = None
            numerator = None
            numerator_field = None
            if "/" in header:
                numerator_field = header.split("/", 1)[0]
                denominator_field = header.rsplit("/", 1)[-1]
                numerator = row.get(numerator_field)
                denominator = row.get(denominator_field)
            result.append({
                **common,
                "metric": _early_metric(header),
                "feature": header,
                "value": feature_value,
                "unit": _unit_for_feature(header, EARLY_SHEET),
                "status": status,
                "numerator": _number(numerator) if _number(numerator) is not None else numerator,
                "numerator_field": numerator_field,
                "denominator": _number(denominator) if _number(denominator) is not None else denominator,
                "denominator_field": denominator_field,
                "note": "首销D1/D2原始字段；比例沿用源表口径。",
            })

        # Stable, human-readable aliases make the UI independent of slight
        # source-header changes while leaving all original columns untouched.
        # These ratios are derived from the source counts only; no net-order or
        # lock-order value is reconstructed.
        aliases = (
            ("D1占最终大定比例", ("D1大定",), ("总大定", "大定")),
            ("D1+D2占最终大定比例", ("D1+D2大定",), ("总大定", "大定")),
            ("D2/D1大定比", ("D2大定",), ("D1大定",)),
            ("D1小转大构成", ("D1小转大",), ("D1大定",)),
            ("D1直接大定构成", ("D1直接大",), ("D1大定",)),
            ("D1+D2小转大构成", ("D1+D2小转大",), ("D1+D2大定",)),
            ("D1+D2直接大定构成", ("D1+D2直接大",), ("D1+D2大定",)),
            ("小转大/最终大定构成", ("总小转大",), ("总大定", "大定")),
            ("直接大定/最终大定构成", ("总直接大定",), ("总大定", "大定")),
            ("D1+D2退订/总小订构成", ("D1+D2退订",), ("总小订",)),
        )
        for alias, numerator_candidates, denominator_candidates in aliases:
            numerator_field = next((name for name in numerator_candidates if name in row), None)
            denominator_field = next((name for name in denominator_candidates if name in row), None)
            numerator = _number(row.get(numerator_field)) if numerator_field else None
            denominator = _number(row.get(denominator_field)) if denominator_field else None
            if numerator is None or denominator is None:
                continue
            value = numerator / denominator if denominator != 0 else None
            base_status = _status_for_value(
                numerator,
                source_cutoff=source_cutoff,
                actual_date=_date_value(anchor_date),
                meta=meta,
            )
            status = base_status if base_status != "有效" else ("分母为零" if denominator == 0 else "有效")
            result.append({
                **common,
                "metric": _early_metric(alias),
                "feature": alias,
                "value": value,
                "unit": "比例",
                "status": status,
                "source_feature": f"{numerator_field}/{denominator_field}",
                "numerator": numerator,
                "numerator_field": numerator_field,
                "denominator": denominator,
                "denominator_field": denominator_field,
                "note": "由源表D1/D2数量构成计算；不重算留存大定或锁单。",
            })
    return result


def _curve_records(sheet: str, rows: list[dict[str, Any]], headers: list[str],
                   meta_maps: Mapping[str, dict[str, Any]], source_cutoff: date | None) -> list[dict[str, Any]]:
    d_headers: list[tuple[str, int]] = []
    for header in headers:
        match = _D_HEADER.match(header)
        if match:
            d_headers.append((header, int(match.group(1))))
    if not d_headers:
        return []
    model_headers = headers
    result: list[dict[str, Any]] = []
    for row in rows:
        model = _model_name(row, model_headers)
        if not model:
            continue
        meta = _lookup_metadata(model, meta_maps)
        stage = "小订" if sheet in SMALL_ORDER_CURVE_SHEETS else "首销"
        # A populated source stage field is more specific than a sheet-name
        # fallback.  Only accept the two known stage labels; do not treat an
        # unrelated text value as evidence for a phase.
        explicit_stage = _text(row.get("小订阶段")) if "小订阶段" in row else ""
        if explicit_stage:
            if "小订" in explicit_stage:
                stage = "小订"
            elif "首销" in explicit_stage or "平销" in explicit_stage:
                stage = "首销"
        anchor_date, anchor_type = _anchor(meta, stage)
        dimensions = _meta_dimensions(meta)
        denominator, denominator_field = _ratio_denominator(meta, sheet)
        # Do not emit trailing empty columns, but keep internal blanks.
        observed_days = [day for header, day in d_headers if row.get(header) not in (None, "")]
        max_day = max(observed_days, default=0)
        cumulative = "累计" in sheet
        for header, day in d_headers:
            if day > max_day:
                continue
            value = _number(row.get(header))
            actual = None
            if anchor_date:
                parsed_anchor = _date_value(anchor_date)
                actual = parsed_anchor + timedelta(days=day - 1) if parsed_anchor else None
            status = _status_for_value(row.get(header), source_cutoff=source_cutoff,
                                       actual_date=actual, meta=meta)
            if value is not None and status == "有效" and cumulative:
                status = "有效（累计快照）"
            result.append({
                "model": model,
                "model_level": _text(row.get("二级代际名")) or model,
                **dimensions,
                "metric": sheet,
                "curve": sheet,
                "stage": stage,
                "day": day,
                "date": actual.isoformat() if actual else None,
                "value": value,
                "unit": "比例" if _is_ratio_curve(sheet) else "数量",
                "is_cumulative": cumulative,
                "snapshot_note": "原始累计快照，不求和" if cumulative else "原始D日值",
                "base": _number(denominator) if _number(denominator) is not None else denominator,
                "denominator": _number(denominator) if _number(denominator) is not None else denominator,
                "denominator_field": denominator_field,
                "denominator_note": (f"车型基本信息.{denominator_field}（原始总量）"
                                      if denominator_field else "源表未提供可用分母"),
                "source": sheet,
                "source_sheet": sheet,
                "anchor_date": anchor_date,
                "anchor_type": anchor_type,
                "status": status,
                "note": ("原始累计快照，不求和" if cumulative else "原始D日值；结构曲线独立于留存大定/锁单"),
                "raw": {key: _json_value(value) for key, value in row.items()
                        if key in ("代际名", "二级代际名", "小订阶段", header)},
            })
    return result


def _hour_field(headers: list[str]) -> str | None:
    return _find_header(headers, "小时", "小时数", "小时序号", contains=True)


def _date_field(headers: list[str]) -> str | None:
    return _find_header(headers, "日期", "时间", "订单日期", "销售日期", contains=True)


def _hour_metric(source_metric: str) -> str:
    """Normalize hourly headers without losing the original field name."""

    text = _text(source_metric)
    if "小订" in text:
        return "小订数量"
    if "小转大" in text:
        return "小转大"
    if "直接大定" in text:
        return "直接大定"
    if "交车锁单" in text or "锁单" in text:
        return "交车锁单"
    if "大定" in text or "总量" in text:
        return "大定"
    return text


def _hourly_records(sheet: str, rows: list[dict[str, Any]], headers: list[str],
                    meta_maps: Mapping[str, dict[str, Any]], source_cutoff: date | None) -> list[dict[str, Any]]:
    model_field = _model_field(headers)
    date_field = _date_field(headers)
    hour_field = _hour_field(headers)
    if not model_field or not date_field or not hour_field:
        return []
    excluded = {model_field, date_field, hour_field, "来源文件", "数据来源"}
    excluded.update(header for header in headers if "来源" in header)
    metric_fields = [header for header in headers if header not in excluded]
    # Ignore text-only metadata fields accidentally placed beside measurements.
    metric_fields = [header for header in metric_fields
                     if any(token in header for token in ("小时", "数量", "大定", "小订", "锁单", "退订"))]
    if not metric_fields:
        return []
    groups: dict[tuple[str, date, str], list[tuple[int | None, Any, dict[str, Any]]]] = defaultdict(list)
    for row in rows:
        model = _model_name(row, headers)
        day = _date_value(row.get(date_field))
        if not model or not day:
            continue
        hour_value = _number(row.get(hour_field))
        hour = int(hour_value) if hour_value is not None and float(hour_value).is_integer() else None
        for metric in metric_fields:
            groups[(model, day, metric)].append((hour, row.get(metric), row))

    result: list[dict[str, Any]] = []
    for (model, day, source_metric), observations in groups.items():
        meta = _lookup_metadata(model, meta_maps)
        dimensions = _meta_dimensions(meta)
        stage = "小订" if sheet == "小订by时" else "首销"
        anchor_date, anchor_type = _anchor(meta, stage)
        hour_values: dict[int, list[Any]] = defaultdict(list)
        for hour, value, _row in observations:
            if hour is not None:
                hour_values[hour].append(value)
        missing = sorted(set(range(24)) - set(hour_values))
        duplicate = sorted(hour for hour, values in hour_values.items() if len(values) > 1)
        flat_values = [values[0] for hour, values in hour_values.items() if len(values) == 1]
        numeric = [_number(value) for value in flat_values]
        has_missing_value = any(_number(value) is None for value in flat_values)
        has_negative = any(value is not None and value < 0 for value in numeric)
        complete = not missing and not duplicate and not has_missing_value and len(numeric) == 24
        total = sum(value for value in numeric if value is not None) if complete else None
        if complete and total == 0:
            group_status = "总量为零"
        elif has_negative:
            group_status = "含负值"
        elif duplicate:
            group_status = "小时重复"
        elif missing:
            group_status = "小时缺失"
        elif has_missing_value:
            group_status = "小时空白"
        elif total is None or total <= 0:
            group_status = "总量无效"
        else:
            group_status = "有效"
        mapping_status = _text(meta.get("_mapping_status"))
        if mapping_status:
            group_status = mapping_status
        elif source_cutoff and day >= source_cutoff:
            group_status = "未来数据"
        # An open model may still have complete historical hours.  Unlike
        # final-scale D1/D2 and cumulative curves, a historical hourly day is
        # usable when its own 24-hour quality checks pass.
        metric = _hour_metric(source_metric)
        for hour in range(24):
            values = hour_values.get(hour, [])
            value = _number(values[0]) if len(values) == 1 else (None if not values else _number(values[0]))
            share = value / total if group_status == "有效" and total and value is not None else None
            if len(values) > 1:
                row_status = "小时重复"
            elif not values:
                row_status = "小时缺失"
            elif _number(values[0]) is None:
                row_status = "小时空白"
            elif value is not None and value < 0:
                row_status = "含负值"
            else:
                row_status = group_status
            result.append({
                "model": model,
                "model_level": _text(_lookup_metadata(model, meta_maps).get("二级代际名")) or model,
                **dimensions,
                "metric": metric,
                "source_metric": source_metric,
                "stage": stage,
                "date": day.isoformat(),
                "hour": hour,
                "value": value,
                "share": share,
                "total": total,
                "observed_hours": len(hour_values),
                "expected_hours": 24,
                "duplicate_count": len(values) if len(values) > 1 else 0,
                "duplicate_values": ([
                    _number(item) if _number(item) is not None else _json_value(item)
                    for item in values
                ] if len(values) > 1 else []),
                "status": row_status,
                "group_status": group_status,
                "note": "源表无留存大定分时字段；按原始小时指标独立计算，不互相替代。",
                "source": sheet,
                "source_sheet": sheet,
                "anchor_date": anchor_date,
                "anchor_type": anchor_type,
                "raw": {"hour": hour, source_metric: value},
            })
    return result


def _inventory(kind_sheet: str, headers: list[str], rows: list[dict[str, Any]],
               note: str = "") -> dict[str, Any]:
    if kind_sheet == MODEL_SHEET:
        kind = "metadata"
    elif kind_sheet == EARLY_SHEET:
        kind = "early_launch"
    elif kind_sheet in HOURLY_SHEETS:
        kind = "hourly"
    elif kind_sheet in CURVE_SHEETS:
        kind = "curve"
    elif kind_sheet == "说明与来源":
        kind = "notes"
    else:
        kind = "other"
    return {"sheet": kind_sheet, "rows": len(rows), "fields": headers,
            "kind": kind, "note": note}


def build_source_insights(path: str | Path) -> dict[str, Any]:
    """Read one consolidated workbook and return source-backed insight data.

    Missing files, missing optional sheets and malformed rows intentionally
    produce empty result families plus notes; callers can render a clear
    source warning rather than failing the complete report.
    """

    empty = {"inventory": [], "early_launch": [], "curves": [], "hourly": [],
             "notes": [], "source_cutoff": None}
    if openpyxl is None:
        empty["notes"].append("未安装 openpyxl，无法读取源表。")
        return empty
    source = Path(path)
    if not source.exists() or not source.is_file():
        empty["notes"].append(f"源表不存在：{source}")
        return empty
    try:
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        empty["notes"].append(f"源表读取失败：{exc}")
        return empty

    cache: dict[str, tuple[list[str], list[dict[str, Any]]]] = {}
    for sheet in workbook.sheetnames:
        try:
            cache[sheet] = _headers_and_rows(workbook[sheet])
        except Exception as exc:
            empty["notes"].append(f"工作表 {sheet} 读取失败：{exc}")
            cache[sheet] = ([], [])
    note_rows = cache.get("说明与来源", ([], []))[1]
    source_cutoff = _parse_cutoff(note_rows)
    empty["source_cutoff"] = source_cutoff.isoformat() if source_cutoff else None
    empty["notes"].extend(_note_lines(note_rows))
    empty["notes"].append("源表外部来源列仅作为追溯文本，不打开其指向的文件。")
    empty["notes"].append("累计曲线按D日原始快照输出，不求和、不重算留存大定或锁单。")
    empty["notes"].append("小订曲线仅小订当日数量、小订期退订当日数量、小订累计完成度和小订期累计退订率；累计小订转化率按首销锚点处理。")
    empty["notes"].append("by时无留存大定分时字段；小时指标按原字段分别展示，不用大定替代留存大定，也不用锁单替代大定。")
    if source_cutoff:
        empty["notes"].append(f"数据判定日期：{source_cutoff.isoformat()}；之后日期标记为未来数据，不参与完整训练样本。")

    meta_headers, meta_rows = cache.get(MODEL_SHEET, ([], []))
    meta_maps, _canonical = _metadata_maps(meta_rows, meta_headers)
    for sheet in workbook.sheetnames:
        headers, rows = cache.get(sheet, ([], []))
        note = ""
        if sheet == MODEL_SHEET:
            note = "动态车型分类、品牌、档位、能源、发布类型和小订/首销锚点。"
        elif sheet in CURVE_SHEETS:
            note = "D日原始值；累计字段为快照，不求和。"
        elif sheet in HOURLY_SHEETS:
            note = "小时原始数量；完整24小时且总量为正才计算分时占比。"
        empty["inventory"].append(_inventory(sheet, headers, rows, note))

    if EARLY_SHEET in cache:
        headers, rows = cache[EARLY_SHEET]
        empty["early_launch"] = _early_features(rows, headers, meta_maps, source_cutoff)
    else:
        empty["notes"].append(f"缺少工作表：{EARLY_SHEET}")
    for sheet in CURVE_SHEETS:
        if sheet not in cache:
            empty["notes"].append(f"缺少曲线工作表：{sheet}")
            continue
        headers, rows = cache[sheet]
        empty["curves"].extend(_curve_records(sheet, rows, headers, meta_maps, source_cutoff))
    for sheet in HOURLY_SHEETS:
        if sheet not in cache:
            empty["notes"].append(f"缺少小时工作表：{sheet}")
            continue
        headers, rows = cache[sheet]
        empty["hourly"].extend(_hourly_records(sheet, rows, headers, meta_maps, source_cutoff))
    # ``read_only=True`` keeps an archive handle open; close it explicitly so
    # Windows callers/tests can replace or remove a temporary source file.
    try:
        workbook.close()
    except Exception:
        pass
    return empty


__all__ = ["build_source_insights"]
