"""Legacy daily exports must not contaminate historical forecast evidence."""
import importlib.util
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def load(name, filename):
    spec = importlib.util.spec_from_file_location(name, ROOT / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


analysis = load('legacy_audit_analysis', 'analyze_sales_patterns.py')
forecast = load('legacy_audit_forecast', 'forecast_insights.py')


class LegacyEvidenceTests(unittest.TestCase):
    def book(self, orders=(), small=(), cutoff=None, statuses=False):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        schemas = (
            ('小订及退订逐日', ['订单分析代际名', '日期', '小订数量', '真实逐日', '小订来源'], small),
            ('当前订单逐日', ['订单分析代际名', '日期', '订单阶段', '留存大定', '留存大定来源'], orders),
        )
        for name, header, records in schemas:
            ws = wb.create_sheet(name)
            if statuses:
                header += ['日期状态']
            ws.append(header)
            for row in records:
                ws.append(row)
        if cutoff is not None:
            ws = wb.create_sheet('说明与来源')
            ws.append(['项目', '内容'])
            ws.append(['数据判定日期', cutoff])
        return wb

    def read(self, wb):
        with patch.object(analysis.openpyxl, 'load_workbook', return_value=wb):
            return analysis.read_panel(Path('in-memory.xlsx'))

    def test_cutoff_controls_both_metrics_and_forecast_origin(self):
        start = date(2025, 1, 1)
        cutoff = start + timedelta(days=70)
        orders = [['车型', start + timedelta(days=i), '平销', 10, '原始逐日'] for i in range(71)]
        orders += [['车型', cutoff + timedelta(days=1), '平销', 99999, '原始逐日'],
                   ['车型', date(2099, 1, 1), '平销', 88888, '原始逐日']]
        small = [['车型', cutoff, 0, True, '原始逐日'],
                 ['车型', date(2099, 1, 1), 77777, True, '原始逐日']]
        panel, audit = self.read(self.book(orders, small, cutoff))
        self.assertEqual(max(r['date'] for r in panel), cutoff)
        self.assertEqual(sum(r['value'] for r in panel), 710)
        self.assertEqual(audit['未来或超过数据判定日期排除'], 3)
        result = forecast.build_forecast_insights(panel)
        self.assertEqual(result['as_of'], cutoff)
        self.assertEqual({r['cutoff'] for r in result['profiles']}, {cutoff})

    def test_today_caps_missing_or_future_workbook_cutoff(self):
        today = date.today()
        for cutoff in (None, today + timedelta(days=20)):
            with self.subTest(cutoff=cutoff):
                panel, _ = self.read(self.book([
                    ['车型', today, '平销', 10, '原始逐日'],
                    ['车型', today + timedelta(days=1), '平销', 90, '原始逐日'],
                ], cutoff=cutoff))
                # No status column: preserve the existing inclusive old-format rule.
                self.assertEqual([(r['date'], r['value']) for r in panel], [(today, 10)])

    def test_optional_date_status_requires_finished_before_cutoff(self):
        cutoff = date(2025, 2, 10)
        orders = [['车型', cutoff - timedelta(days=1), '平销', 0, '原始逐日', '已结束日'],
                  ['车型', cutoff, '平销', 50, '原始逐日', '已结束日']]
        small = [['车型', cutoff - timedelta(days=1), 0, True, '原始逐日', '已结束日']]
        for index, status in enumerate(('当前日', '未结束日', '未知', None), 2):
            orders.append(['车型', cutoff - timedelta(days=index), '平销', 80, '原始逐日', status])
            small.append(['车型', cutoff - timedelta(days=index), 90, True, '原始逐日', status])
        panel, audit = self.read(self.book(orders, small, cutoff, statuses=True))
        self.assertEqual(len(panel), 2)
        self.assertTrue(all(r['value'] == 0 for r in panel))
        self.assertEqual(audit['未结束或未知日期状态排除'], 8)
        self.assertEqual(audit['数据判定当日未完整排除'], 1)

    def test_unknown_order_stages_and_snapshots_never_train(self):
        start = date(2025, 1, 1)
        orders = [['车型', start, '首销', 0, '原始逐日'],
                  ['车型', start + timedelta(days=1), '平销', -2, '原始逐日']]
        for index, stage in enumerate((None, '未知阶段', '阶段未知', '首销前', '其他'), 2):
            orders.append(['车型', start + timedelta(days=index), stage, 100, '原始逐日'])
        small = []
        for index, source in enumerate(('分时累计', '累计快照', '参考曲线', '未更新'), 10):
            orders.append(['车型', start + timedelta(days=index), '平销', 100, source])
            small.append(['车型', start + timedelta(days=index), 100, True, source])
        panel, audit = self.read(self.book(orders, small, date(2025, 2, 1)))
        self.assertEqual([(r['stage'], r['value']) for r in panel], [('首销', 0), ('平销', -2)])
        self.assertEqual(audit['阶段不明确排除'], 5)
        self.assertEqual(audit['留存大定来源排除'], 4)
        self.assertEqual(audit['小订数量来源排除'], 4)

    def test_rejected_future_duplicates_do_not_raise_conflict(self):
        panel, audit = self.read(self.book([
            ['车型', date(2025, 1, 1), '平销', 10, '原始逐日'],
            ['车型', date(2099, 1, 1), '平销', 20, '原始逐日'],
            ['车型', date(2099, 1, 1), '平销', 30, '原始逐日'],
        ], cutoff=date(2025, 1, 1)))
        self.assertEqual([r['value'] for r in panel], [10])
        self.assertEqual(audit['未来或超过数据判定日期排除'], 2)

    def test_unified_schema_delegation_is_unchanged(self):
        wb = self.book()
        wb.create_sheet(analysis.UNIFIED_DAILY_SHEET)
        expected = ([{'value': 17, 'source': 'resolved'}], {'unified': 1})
        with patch.object(analysis, 'read_unified_panel', return_value=expected) as reader:
            self.assertEqual(self.read(wb), expected)
        self.assertIs(reader.call_args.args[0], wb)


if __name__ == '__main__':
    unittest.main()
