"""Single-workbook descriptive sales analysis. Python 3.10+ and openpyxl."""
import argparse, calendar, hashlib, importlib.util, json, math, os, re, sys, time, webbrowser
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from statistics import mean, median
import openpyxl
SCRIPT_PATH=Path(__file__).resolve()
CODE_ROOT=SCRIPT_PATH.parent.parent
ROOT=CODE_ROOT.parent if CODE_ROOT.name.lower() == 'scripts' else CODE_ROOT

def log(message):
    print(f'[{datetime.now():%Y-%m-%d %H:%M:%S}] {message}',flush=True)
ORDER_METRIC_COLUMNS={
    '大定':(('大定','大定来源'),),
    '留存大定':(('留存大定','留存大定来源'),('净大定','净大定来源')),
    '小转大':(('小转大','小转大来源'),),
    '直接大定':(('直接大定','直接大定来源'),),
    '交车锁单':(('交车锁单','锁单来源'),),
}

def dt(v):
    if isinstance(v,datetime): return v.date()
    if isinstance(v,date): return v
    try: return date.fromisoformat(str(v)[:10])
    except ValueError: return None

def span(a,b):
    while a<=b:
        yield a
        a+=timedelta(days=1)

def group(rows,fields):
    g=defaultdict(list)
    for r in rows: g[tuple(r[f] for f in fields)].append(r)
    return g

def load_calendar(extra=None):
    spec=importlib.util.spec_from_file_location('calendar_rules',SCRIPT_PATH.parent.parent/'generate_html/core/china_calendar.py')
    m=importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
    periods=dict(m.HOLIDAY_PERIODS); work=dict(m.ADJUSTED_WORKDAYS); sources=dict(m.SOURCE_URLS)
    periods[2021]=(('元旦','2021-01-01','2021-01-03'),('春节','2021-02-11','2021-02-17'),('清明节','2021-04-03','2021-04-05'),('劳动节','2021-05-01','2021-05-05'),('端午节','2021-06-12','2021-06-14'),('中秋节','2021-09-19','2021-09-21'),('国庆节','2021-10-01','2021-10-07'))
    work[2021]=('2021-02-07','2021-02-20','2021-04-25','2021-05-08','2021-09-18','2021-09-26','2021-10-09')
    sources[2021]='https://app.www.gov.cn/govdata/gov/202011/26/465359/article.html'
    periods[2024]=tuple((n,'2023-12-30' if n=='元旦' else '2024-06-08' if n=='端午节' else a,b) for n,a,b in periods[2024])
    if extra:
        for y,c in json.loads(Path(extra).read_text(encoding='utf-8-sig')).items():
            y=int(y); periods[y]=c['periods']; work[y]=c['workdays']; sources[y]=c['source']
    events=[(y,n,dt(a),dt(b)) for y,items in sorted(periods.items()) for n,a,b in items]
    if any(not a or not b or b<a for _,_,a,b in events): raise ValueError('Invalid calendar')
    return dict(years=set(periods),events=events,holidays={d:n for _,n,a,b in events for d in span(a,b)},work={dt(d) for ds in work.values() for d in ds},sources=sources)

def ordinary(d,c):
    return d.year in c['years'] and d not in c['holidays'] and d not in c['work']

UNIFIED_DAILY_SHEET = '小订首销平销by天'
# Canonical names remain stable for analyses/configuration across workbook versions.
RESOLVED_METRIC_COLUMNS = {
    '小订数量': ('小订', '小订数量'),
    **{metric: tuple(column for column, _ in aliases)
       for metric, aliases in ORDER_METRIC_COLUMNS.items()},
}

def _workbook_records(wb, name):
    rows = iter(wb[name].iter_rows(values_only=True))
    header = next(rows, ())
    return header, ((n, dict(zip(header, values))) for n, values in enumerate(rows, 2)
                    if any(value is not None for value in values))


def _daily_source_rank(row, metric, model_info, cutoff):
    kind, filename = row['来源类型'], row['来源文件']
    if metric == '小订数量':
        if '小订选配' in filename: return 0
        if '小订退订' in filename: return 1
        if kind == '历史小订': return 2
        return 10
    if metric == '交车锁单' and row['订单阶段'] == '平销' and kind == '平销锁单':
        return 0
    start, end = dt(model_info.get('首销开始')), dt(model_info.get('首销结束'))
    if start and cutoff < start:
        order = ('小订退订', '历史首销', '首销订单')
    elif (end and cutoff >= end) or row['订单阶段'] == '平销':
        order = ('历史首销', '首销订单', '小订退订')
    else:
        order = ('首销订单', '历史首销', '小订退订')
    for rank, source in enumerate(order, 1):
        if source == kind or (source == '小订退订' and source in filename): return rank
    return 10


def _match_model_phase(records, row, day):
    """Match repeated model entries by original name and period, not row order."""
    records = list(records)
    # Keep an undated model row available.  It is useful evidence for the
    # explicit "阶段未知" rule: a missing boundary must not be treated as a
    # valid launch period merely because the row has an order value.
    if records and not any(dt(r.get('首销开始')) or dt(r.get('小订开始')) for r in records):
        if len(records) > 1:
            raise ValueError('车型阶段日期冲突：'+str(row.get('订单分析代际名') or row.get('代际名')))
        return records[0]
    dated = [r for r in records if dt(r.get('首销开始')) or dt(r.get('小订开始'))]
    same_name = [r for r in dated if row.get('历史传播名') and
                 r.get('历史传播名') == row['历史传播名']]
    choices = same_name or dated
    prefix = '小订' if row.get('订单阶段') in ('小订', '小订阶段') else '首销'
    covering = [r for r in choices if dt(r.get(prefix+'开始')) and dt(r.get(prefix+'结束'))
                and dt(r[prefix+'开始']) <= day <= dt(r[prefix+'结束'])]
    choices = covering or choices
    starts = [(dt(r.get(prefix+'开始')), r) for r in choices if dt(r.get(prefix+'开始'))]
    if not starts: return {}
    past = [(start, r) for start, r in starts if start <= day]
    selected_start = max(start for start, _ in past) if past else min(start for start, _ in starts)
    matches = [r for start, r in starts if start == selected_start]
    boundaries = {tuple(dt(r.get(k)) for k in ('小订开始','小订结束','首销开始','首销结束')) for r in matches}
    if len(boundaries) > 1:
        raise ValueError('车型阶段日期冲突：'+str(row.get('订单分析代际名')))
    return matches[0]


_ORDER_STAGE_FIELDS = frozenset({
    '大定', '留存大定', '净大定', '小转大', '直接大定', '交车锁单',
    '小订退订', '退订数量', '首销期退订', '首销期退订数量',
})


def _derived_stage_tags(row, day, small_info, order_info, columns):
    """Derive independent small-order and order stages from row data.

    The workbook's stage label is retained as an audit hint only.  A small
    order tag is based on the presence of its own value and its small-order
    window; the order tag is based on any order/return field and the launch
    boundaries.  The two tags are then concatenated by callers, so a day can
    legitimately be ``小订/首销``.
    """
    small_column = columns.get('小订数量')
    small_value_present = bool(small_column and row.get(small_column) is not None)
    small_info = small_info or {}
    order_info = order_info or {}
    if not small_info and not order_info:
        # No model metadata means both order boundaries are missing.  Keep a
        # small-order value (the user explicitly allows missing small bounds),
        # but mark order data unknown so it cannot enter launch pacing.
        tags = ['小订'] if small_value_present else []
        if any(field in row and row.get(field) is not None
               for field in (_ORDER_STAGE_FIELDS | {metric for metric in columns
                                                     if metric != '小订数量'})):
            tags.append('阶段未知')
        return tags
    small_start, small_end = dt(small_info.get('小订开始')), dt(small_info.get('小订结束'))
    tags = []
    if small_value_present and (not small_start or not small_end or
                                small_start <= day <= small_end):
        tags.append('小订')

    order_fields = set(_ORDER_STAGE_FIELDS)
    order_fields.update(metric for metric in columns if metric != '小订数量')
    # In the legacy unified layout the canonical key may resolve to an alias
    # such as 净大定.  Check both the canonical metric and the actual header so
    # an order value is still enough to derive 首销前/首销/平销.
    order_fields.update(column for metric, column in columns.items()
                        if metric != '小订数量' and column)
    has_order_value = any(field in row and row.get(field) is not None for field in order_fields)
    if has_order_value:
        launch_start, launch_end = dt(order_info.get('首销开始')), dt(order_info.get('首销结束'))
        if not launch_start or not launch_end:
            tags.append('阶段未知')
        elif day < launch_start:
            tags.append('首销前')
        elif day <= launch_end:
            tags.append('首销')
        else:
            tags.append('平销')
    return tags


def _model_name(row):
    return row.get('二级代际名') or row.get('订单分析代际名') or row.get('代际名')


def _field_sources(value):
    """Expand metric/source assignments; never borrow another metric's source."""
    result = {}
    for part in re.split(r'[；;\n]', str(value or '')):
        if '=' not in part: continue
        fields, source = part.split('=', 1)
        source = source.strip()
        for field in fields.split('/'):
            field = field.strip()
            if field in result and result[field] != source:
                raise ValueError('字段来源冲突：'+field)
            result[field] = source
    return result


def _resolved_field_source(sources, aliases):
    # The new value header can coexist with the old name in field provenance.
    declared = {sources[name] for name in aliases if name in sources}
    if len(declared) > 1:
        raise ValueError('字段来源冲突：'+'/'.join(aliases))
    return next(iter(declared), '')


def read_resolved_panel(wb, audit):
    """The current workbook already resolves each metric; respect its provenance."""
    cutoff = date.today()
    if '说明与来源' in wb.sheetnames:
        _, notes = _workbook_records(wb, '说明与来源')
        for _, row in notes:
            if row.get('项目') == '数据判定日期' and dt(row.get('内容')):
                cutoff = min(cutoff, dt(row['内容']))
    info = defaultdict(list)
    if '车型基本信息' in wb.sheetnames:
        _, model_rows = _workbook_records(wb, '车型基本信息')
        for _, row in model_rows:
            names = {_model_name(row), row.get('代际名'), row.get('订单分析代际名')}
            for name in names:
                if name: info[str(name)].append(row)
    head, rows = _workbook_records(wb, UNIFIED_DAILY_SHEET)
    required = {'日期', '订单阶段', '字段来源'}
    if required-set(head) or not ({'代际名','订单分析代际名'} & set(head)):
        raise ValueError(UNIFIED_DAILY_SHEET+'缺少字段：日期、代际名、订单阶段或字段来源')
    columns = {}
    for metric, aliases in RESOLVED_METRIC_COLUMNS.items():
        present = [column for column in aliases if column in head]
        if present:
            columns[metric] = present[0]
            if len(present)>1: audit[metric+'新旧字段同时存在，优先新字段'] += 1
    if not columns: raise ValueError(UNIFIED_DAILY_SHEET+'没有可识别的订单指标列')
    panel, seen = [], {}
    for rowno, row in rows:
        audit[UNIFIED_DAILY_SHEET+'原始行'] += 1
        d, model = dt(row.get('日期')), _model_name(row)
        if not d or not model:
            audit['无效日期或车型'] += 1
            continue
        if d > cutoff:
            audit['未来或超过数据判定日期排除'] += 1
            continue
        if '日期状态' in head:
            # New-schema rows explicitly describe completion. Unknown/blank status
            # and the judgement day itself must not train complete-day patterns.
            if str(row.get('日期状态') or '').strip() != '已结束日':
                audit['未结束或未知日期状态排除'] += 1
                continue
            if d == cutoff:
                audit['数据判定当日未完整排除'] += 1
                continue
        sources = _field_sources(row.get('字段来源'))
        raw_stage = str(row.get('订单阶段') or '').strip()
        records = info.get(str(model), [])
        # A repeated model can have different small-order and launch batches.
        # Resolve each stage against its own date window before deriving the
        # composite tags; matching once from the raw composite label can pick
        # the launch batch and incorrectly discard a valid small-order value.
        small_info = _match_model_phase(records, dict(row, 订单阶段='小订'), d)
        order_info = _match_model_phase(records, dict(row, 订单阶段='首销'), d)
        # The source stage is an audit hint.  Rebuild it from the row's own
        # fields and model dates so a same-day small-order + launch row keeps
        # both tags and launch boundaries remain inclusive.
        stages = _derived_stage_tags(row, d, small_info, order_info, columns)
        raw_tags = {('小订' if part == '小订阶段' else part)
                    for part in raw_stage.split('/') if part}
        if raw_tags != set(stages):
            audit['阶段标签与日期字段推导不一致'] += 1
        if not stages:
            audit['阶段不明确排除'] += 1
            continue
        if row.get('首销分时核对'):
            # Retain an audit count only: hourly checks do not override daily values.
            audit['原表首销分时核对标记行'] += 1
        for metric, column in columns.items():
            value = row.get(column)
            if value is None:
                audit[metric+'缺失'] += 1
                continue
            if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value):
                audit[metric+'非数值'] += 1
                continue
            source = _resolved_field_source(sources, RESOLVED_METRIC_COLUMNS[metric])
            if not source or any(word in source for word in ('预测','合成','模拟','缺失','累计','快照','参考曲线','未更新')):
                audit[metric+'字段来源排除'] += 1
                continue
            if metric == '小订数量':
                if '小订' not in stages:
                    audit[metric+'阶段不明确排除'] += 1
                    continue
                stage = '小订阶段'
            else:
                order_stages = [part for part in stages if part != '小订']
                if len(order_stages) != 1 or order_stages[0] not in ('首销', '平销'):
                    audit[metric+'阶段不明确排除'] += 1
                    continue
                stage = order_stages[0]
            if len(stages)>1: audit['复合阶段按指标拆分'] += 1
            metric_info = _match_model_phase(
                records,
                dict(row, 订单阶段=('小订' if stage == '小订阶段' else stage)), d)
            prefix = '小订' if stage == '小订阶段' else '首销'
            start, end = dt(metric_info.get(prefix+'开始')), dt(metric_info.get(prefix+'结束'))
            match = re.fullmatch(r'D(\d+)',str(row.get('生命周期') or ''),re.I)
            same_stage = ((raw_stage in ('小订', '小订阶段')) if stage == '小订阶段'
                          else raw_stage == stage)
            life = int(match[1]) if match and len(stages) == 1 and same_stage else None
            anchor = d-timedelta(days=life-1) if life and life>0 else None
            if stage in ('小订阶段','首销') and start and end and start<=d<=end:
                anchor, life = start,(d-start).days+1
            if stage == '平销': life, anchor = None,None
            cycle = str(anchor) if anchor else stage
            launch_start = dt(order_info.get('首销开始'))
            launch_end = dt(order_info.get('首销结束'))
            if stage != '小订阶段' and (launch_start is None or launch_end is None):
                launch_start = dt(metric_info.get('首销开始'))
                launch_end = dt(metric_info.get('首销结束'))
            key = (str(model),d,metric)
            signature = (value,stage,source,cycle,launch_start,launch_end)
            if key in seen:
                if seen[key] != signature: raise ValueError('已合并逐日表重复键冲突：'+str(key))
                audit['相同重复去重'] += 1
                continue
            seen[key] = signature
            if value<0: audit[metric+'负值不参与倍率'] += 1
            panel.append(dict(model=str(model),date=d,metric=metric,stage=stage,source=source,
                              cycle=cycle,value=float(value),life=life,
                              small_start=dt(metric_info.get('小订开始')),
                              small_end=dt(metric_info.get('小订结束')),
                              launch_start=launch_start,launch_end=launch_end,
                              sheet=UNIFIED_DAILY_SHEET,row=rowno))
    audit['字段级来源表已识别'] = 1
    attach_model_classes(wb,panel,audit)
    return panel,dict(audit)


def read_unified_panel(wb, audit):
    """Resolve daily source candidates without adding snapshots or weekly totals."""
    header = next(wb[UNIFIED_DAILY_SHEET].iter_rows(values_only=True), ())
    if '字段来源' in header:
        return read_resolved_panel(wb, audit)
    cutoff = date.today()
    if '说明与来源' in wb.sheetnames:
        _, notes = _workbook_records(wb, '说明与来源')
        for _, row in notes:
            if row.get('项目') == '数据判定日期' and dt(row.get('内容')):
                cutoff = min(cutoff, dt(row['内容']))
    info = defaultdict(list)
    if '车型基本信息' in wb.sheetnames:
        _, model_rows = _workbook_records(wb, '车型基本信息')
        for _, row in model_rows:
            model = row.get('订单分析代际名')
            if model:
                info[str(model)].append(row)
    head, rows = _workbook_records(wb, UNIFIED_DAILY_SHEET)
    required = {'日期', '订单分析代际名', '订单阶段', '来源类型', '来源文件', '来源Sheet'}
    if required - set(head):
        raise ValueError(UNIFIED_DAILY_SHEET + '缺少字段' + str(sorted(required-set(head))))
    columns = {'小订数量': '小订数量'} if '小订数量' in head else {}
    for metric, aliases in ORDER_METRIC_COLUMNS.items():
        present = [col for col, _ in aliases if col in head]
        if present:
            columns[metric] = present[0]
            if len(present) > 1: audit[metric+'新旧字段同时存在，优先新字段'] += 1
    if not columns:
        raise ValueError(UNIFIED_DAILY_SHEET + '没有可识别的订单指标列')
    candidates = defaultdict(list)
    seen = {}
    for rowno, row in rows:
        audit[UNIFIED_DAILY_SHEET+'原始行'] += 1
        d, model = dt(row.get('日期')), row.get('订单分析代际名')
        if not d or not model:
            audit['无效日期或车型'] += 1
            continue
        if d > cutoff:
            audit['未来或超过数据判定日期排除'] += 1
            continue
        kind = str(row.get('来源类型') or '').strip()
        filename = str(row.get('来源文件') or '').strip()
        source_sheet = str(row.get('来源Sheet') or '').strip()
        if any(word in kind or word in source_sheet for word in ('累计', '快照', '参考曲线')):
            audit['累计快照或参考曲线排除'] += 1
            continue
        if not kind or not filename or not source_sheet or any(
                word in kind for word in ('预测', '合成', '模拟', '缺失')):
            audit['逐日来源排除'] += 1
            continue
        if '真实逐日' in head and str(row.get('真实逐日')).lower() not in ('true', '1', '是'):
            audit['非真实逐日排除'] += 1
            continue
        raw_stage = str(row.get('订单阶段') or '').strip()
        if not raw_stage:
            audit['订单阶段缺失排除'] += 1
            continue
        records = info.get(str(model), [])
        # A composite source label such as 小订/首销 contains two independent
        # facts.  Derive them from the row's own values and each metric's date
        # window before assigning a stage to the metric.  In particular, a
        # small-order value can never leak into 首销/平销 rows.
        small_info = _match_model_phase(records, dict(row, 订单阶段='小订'), d)
        order_info = _match_model_phase(records, dict(row, 订单阶段='首销'), d)
        stage_tags = _derived_stage_tags(row, d, small_info, order_info, columns)
        # Older unified exports did not carry 车型基本信息.  Preserve their
        # explicit 首销/平销 stage as a compatibility hint when no boundaries
        # exist at all; the field-level format still uses 阶段未知 until its
        # missing boundaries are resolved.
        if not small_info and not order_info and raw_stage in ('首销前', '首销', '平销'):
            has_order_value = any(
                row.get(field) is not None
                for field in (_ORDER_STAGE_FIELDS | {
                    column for metric, column in columns.items()
                    if metric != '小订数量' and column
                })
            )
            if has_order_value:
                # Without phase metadata the raw label is the only safe
                # small-order hint.  A small value on a plain 首销/平销 row
                # must not be promoted to 小订期 merely because it is filled.
                small_tag = (['小订'] if '小订' in raw_stage and
                             row.get(columns.get('小订数量')) is not None else [])
                stage_tags = small_tag + [raw_stage]
        raw_tags = {('小订' if part == '小订阶段' else part)
                    for part in raw_stage.split('/') if part}
        if raw_tags != set(stage_tags):
            audit['阶段标签与日期字段推导不一致'] += 1
        if not stage_tags:
            audit['阶段不明确排除'] += 1
            continue
        source = ' | '.join((kind, filename, source_sheet))
        normalized_row = dict(row, 来源类型=kind, 来源文件=filename, 订单阶段=raw_stage)
        for metric, column in columns.items():
            value = row.get(column)
            if value is None:
                audit[metric+'缺失'] += 1
                continue
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                audit[metric+'非数值'] += 1
                continue
            if metric == '小订数量':
                if '小订' not in stage_tags:
                    audit[metric+'阶段不明确排除'] += 1
                    continue
                metric_stage = '小订阶段'
            else:
                order_stages = [part for part in stage_tags if part != '小订']
                if len(order_stages) != 1 or order_stages[0] not in ('首销', '平销'):
                    audit[metric+'阶段不明确排除'] += 1
                    continue
                metric_stage = order_stages[0]
            metric_info = _match_model_phase(
                records,
                dict(row, 订单阶段=('小订' if metric_stage == '小订阶段' else metric_stage)), d)
            match = re.fullmatch(r'D(\d+)', str(row.get('生命周期') or ''), re.I)
            life = int(match[1]) if match else None
            prefix = '小订' if metric_stage == '小订阶段' else '首销'
            start, end = dt(metric_info.get(prefix+'开始')), dt(metric_info.get(prefix+'结束'))
            anchor = d-timedelta(days=life-1) if life and life > 0 else None
            if metric_stage in ('小订阶段', '首销') and start and end and start <= d <= end:
                anchor, life = start, (d-start).days+1
            if metric_stage == '平销':
                life, anchor = None, None
            cycle = str(anchor) if anchor and metric_stage != '平销' else metric_stage
            launch_start = dt(order_info.get('首销开始')) or dt(row.get('首销开始'))
            launch_end = dt(order_info.get('首销结束')) or dt(row.get('首销结束'))
            key = (str(model), d, metric)
            source_key = key + (source,)
            signature = (value, metric_stage, cycle)
            if source_key in seen:
                if seen[source_key] != signature:
                    raise ValueError('重复键冲突（同来源）：'+str(source_key))
                audit['相同重复去重'] += 1
                continue
            seen[source_key] = signature
            item = dict(model=str(model), date=d, metric=metric, stage=metric_stage,
                        source=source, cycle=cycle, value=float(value), life=life,
                        small_start=dt(metric_info.get('小订开始')),
                        small_end=dt(metric_info.get('小订结束')),
                        launch_start=launch_start, launch_end=launch_end,
                        sheet=UNIFIED_DAILY_SHEET, row=rowno)
            rank_info = small_info if metric == '小订数量' else order_info
            rank = _daily_source_rank(normalized_row, metric, rank_info, cutoff)
            candidates[key].append((rank, item))
    panel = []
    for key, choices in sorted(candidates.items()):
        choices.sort(key=lambda choice: (choice[0], choice[1]['source']))
        rank, selected = choices[0]
        tied = [r for priority, r in choices if priority == rank]
        if any((r['value'], r['stage'], r['cycle']) !=
               (selected['value'], selected['stage'], selected['cycle']) for r in tied):
            raise ValueError('同优先级来源冲突，无法自动选取：'+str(key))
        audit['跨来源候选未累加'] += len(choices)-1
        if selected['value'] < 0: audit[selected['metric']+'负值不参与倍率'] += 1
        panel.append(selected)
    audit['新版逐日表已识别'] = 1
    attach_model_classes(wb, panel, audit)
    return panel, dict(audit)


def read_panel(path):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    panel=[]; audit=Counter(); seen={}
    try:
        if UNIFIED_DAILY_SHEET in getattr(wb, 'sheetnames', ()):
            return read_unified_panel(wb, audit)
        # Legacy exports are historical evidence too: use the same date fence
        # as the unified reader before duplicate checks or forecast training.
        cutoff = date.today()
        if '说明与来源' in getattr(wb, 'sheetnames', ()):
            _, notes = _workbook_records(wb, '说明与来源')
            for _, note in notes:
                if note.get('项目') == '数据判定日期' and dt(note.get('内容')):
                    cutoff = min(cutoff, dt(note['内容']))
        for name,small in [('小订及退订逐日',True),('当前订单逐日',False)]:
            it=wb[name].iter_rows(values_only=True); head=next(it)
            required={'日期','订单分析代际名'}|({'小订数量','真实逐日'} if small else {'订单阶段'})
            if required-set(head): raise ValueError(name+'缺少字段'+str(required-set(head)))
            metrics={'小订数量':('小订数量','小订来源')} if small else {}
            if not small:
                for metric,candidates in ORDER_METRIC_COLUMNS.items():
                    present=[(value_col,source_col) for value_col,source_col in candidates if value_col in head]
                    if not present:
                        continue
                    value_col,source_col=present[0]
                    if source_col not in head:
                        raise ValueError(name+'缺少字段'+source_col)
                    metrics[metric]=(value_col,source_col)
                    if len(present)>1:
                        audit[metric+'新旧字段同时存在，优先新字段']+=1
            for rowno,vs in enumerate(it,2):
                if all(v is None for v in vs): continue
                r=dict(zip(head,vs)); audit[name+'原始行']+=1
                d=dt(r.get('日期')); model=r.get('订单分析代际名')
                if not d or not model: audit['无效日期或车型']+=1; continue
                if d > cutoff:
                    audit['未来或超过数据判定日期排除'] += 1
                    continue
                if '日期状态' in head:
                    if str(r.get('日期状态') or '').strip() != '已结束日':
                        audit['未结束或未知日期状态排除'] += 1
                        continue
                    if d == cutoff:
                        audit['数据判定当日未完整排除'] += 1
                        continue
                stage='小订阶段' if small else str(r.get('订单阶段') or '').strip()
                if not small and stage not in ('首销', '平销'):
                    audit['阶段不明确排除'] += 1
                    continue
                match=re.fullmatch(r'D(\d+)',str(r.get('生命周期') or ''),re.I)
                life=int(match[1]) if match else None
                anchor=dt(r.get('小订开始')) if small else None
                if not anchor and life and life>0: anchor=d-timedelta(days=life-1)
                cycle=str(anchor) if anchor and stage!='平销' else stage
                for metric,(value_col,scol) in metrics.items():
                    v=r.get(value_col)
                    if v is None: audit[metric+'缺失']+=1; continue
                    if isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v): audit[metric+'非数值']+=1; continue
                    if small and str(r.get('真实逐日')).lower() not in ('true','1','是'): audit['非真实小订排除']+=1; continue
                    src=str(r.get(scol) or '')
                    if not src or any(s in src for s in ('预测','合成','模拟','缺失','累计','快照','参考曲线','未更新')): audit[metric+'来源排除']+=1; continue
                    key=(str(model),d,metric); signature=(v,stage,src,cycle)
                    if key in seen:
                        if seen[key]!=signature: raise ValueError('重复键冲突'+str(key))
                        audit['相同重复去重']+=1; continue
                    seen[key]=signature
                    if v<0: audit[metric+'负值不参与倍率']+=1
                    panel.append(dict(model=str(model),date=d,metric=metric,stage=stage,source=src,
                                      cycle=cycle,value=float(v),life=life,
                                      small_start=dt(r.get('小订开始')),
                                      small_end=dt(r.get('小订结束')),
                                      launch_start=dt(r.get('首销开始')),
                                      launch_end=dt(r.get('首销结束')),
                                      sheet=name,row=rowno))
        attach_model_classes(wb, panel, audit)
    finally: wb.close()
    return panel,dict(audit)

def analyze(panel,c):
    out={k:[] for k in ('coverage','weekly','monthly','annual','holidays','lifecycle')}
    for (metric,stage),rs in group(panel,('metric','stage')).items():
        ds={r['date'] for r in rs}
        out['coverage'].append(dict(metric=metric,stage=stage,models=len({r['model'] for r in rs}),observations=len(rs),dates=len(ds),start=min(ds),end=max(ds)))
    fields=('model','metric','stage','cycle','source')
    for key,rs in group(panel,fields).items():
        base=dict(zip(fields,key)); days={r['date']:r for r in rs if r['value']>=0}
        if not days: continue
        life={r['life']:r['value'] for r in rs if r['life'] and r['value']>=0}
        if key[2]!='平销' and all(i in life for i in range(1,8)) and sum(life[i] for i in range(3,8))>0:
            out['lifecycle'].append(dict(base,d1_ratio=life[1]/mean(life[i] for i in range(3,8)),first2_share=(life[1]+life[2])/sum(life[i] for i in range(1,8))))
        wg=defaultdict(list); mg=defaultdict(list)
        for d,r in days.items():
            wg[d-timedelta(days=d.weekday())].append(r)
            mg[(d.year,d.month)].append(r)
        for monday,rows in wg.items():
            if len(rows)!=7 or not all(ordinary(r['date'],c) for r in rows): continue
            ys=[days[monday+timedelta(days=i)]['value'] for i in range(7)]
            if mean(ys)<=0 or mean(ys[:5])<=0: continue
            out['weekly'].append(dict(base,week=monday,indexes=[100*y/mean(ys) for y in ys],ratio=mean(ys[5:])/mean(ys[:5])))
        cm={}
        for (y,m),rows in mg.items():
            n=calendar.monthrange(y,m)[1]
            if len(rows)!=n: continue
            ys=[days[date(y,m,i)]['value'] for i in range(1,n+1)]
            cm[(y,m)]=mean(ys)
            denominator=mean(ys[:-7])
            out['monthly'].append(dict(base,year=y,month=m,
                ratio=mean(ys[-7:])/denominator if denominator>0 else None,
                status='有效' if denominator>0 else '月内前期日均为零'))
        for y in {y for y,m in cm}:
            if all((y,m) in cm for m in range(1,13)):
                ys=[cm[(y,m)] for m in range(1,13)]
                annual_mean=mean(ys)
                out['annual'].append(dict(base,year=y,
                    indexes=[100*v/annual_mean for v in ys] if annual_mean>0 else [None]*12,
                    status='有效' if annual_mean>0 else '全年日均为零'))
        if key[2]!='平销': continue
        for y,name,a,b in c['events']:
            if a>max(days) or b<min(days): continue
            target=list(span(a-timedelta(days=7),b+timedelta(days=7)))
            if not all(d in days for d in target): continue
            controls=[d for d in days if ordinary(d,c) and (a-timedelta(days=35)<=d<=a-timedelta(days=8) or b+timedelta(days=8)<=d<=b+timedelta(days=35)) and not any(x-timedelta(days=7)<=d<=z+timedelta(days=7) for _,_,x,z in c['events'])]
            ratios={}; expected_by_day={}; used=set()
            for d in target:
                cs=[x for x in controls if x.weekday()==d.weekday()]
                if len(cs)<4 or not any(x<a for x in cs) or not any(x>b for x in cs): break
                expected=mean(days[x]['value'] for x in cs)
                if expected<=0: break
                ratios[d]=days[d]['value']/expected; expected_by_day[d]=expected; used.update(cs)
            else:
                out['holidays'].append(dict(base,year=y,holiday=name,before=mean(v for d,v in ratios.items() if d<a),during=mean(v for d,v in ratios.items() if a<=d<=b),after=mean(v for d,v in ratios.items() if d>b),control_days=len(used),
                    daily=[dict(date=d,actual=days[d]['value'],expected=expected_by_day[d]) for d in target]))
    return out

# Comparisons always retain each order metric and sales stage separately.
def period_bounds(d, kind):
    if kind == '日':
        return d, d
    if kind == '周':
        a = d - timedelta(days=d.weekday())
        return a, a + timedelta(days=6)
    month = d.month if kind == '月' else ((d.month-1)//3)*3+1 if kind == '季' else 1
    a = date(d.year, month, 1)
    width = 1 if kind == '月' else 3 if kind == '季' else 12
    endmonth = month + width - 1
    return a, date(d.year, endmonth, calendar.monthrange(d.year, endmonth)[1])

def quantile(values, p):
    vs = sorted(values)
    if not vs:
        return None
    pos = (len(vs)-1)*p
    low = int(pos); high = min(low+1, len(vs)-1)
    return vs[low] + (vs[high]-vs[low])*(pos-low)

def order_role(metric):
    return '锁单' if metric == '交车锁单' else '小订' if metric == '小订数量' else '大定'

def attach_model_classes(wb, panel, audit):
    definitions = ('brand', 'segment', 'energy')
    fields = ('品牌', '产品档位', '能源类型')
    mapping = defaultdict(lambda: defaultdict(set))
    if '车型基本信息' in getattr(wb, 'sheetnames', ()):
        it = wb['车型基本信息'].iter_rows(values_only=True)
        headers = next(it)
        for values in it:
            r = dict(zip(headers, values)); model = _model_name(r)
            if not model:
                continue
            for key, field in zip(definitions, fields):
                if r.get(field):
                    for name in {model, r.get('代际名'), r.get('订单分析代际名')} - {None, ''}:
                        mapping[str(name)][key].add(str(r[field]).strip())
    for r in panel:
        r['role'] = order_role(r['metric'])
        for key in definitions:
            choices = mapping[r['model']][key]
            r[key] = next(iter(choices)) if len(choices) == 1 else '映射冲突' if choices else '未提供'
    audit['属性缺失车型数'] = len({r['model'] for r in panel if any(r[k]=='未提供' for k in definitions)})
    audit['属性冲突车型数'] = len({r['model'] for r in panel if any(r[k]=='映射冲突' for k in definitions)})

def period_comparisons(panel, cal):
    comparisons = []
    interval_calendar_cache = {}
    for key, rs in group(panel, ('model','metric','stage','cycle','source')).items():
        base = dict(zip(('model','metric','stage','cycle','source'),key))
        base.update({k:rs[0].get(k,'未提供') for k in ('brand','segment','energy')})
        base['role'] = order_role(base['metric'])
        for kind in ('日','周','月','季','年'):
            buckets = defaultdict(list)
            for r in rs:
                a,b = period_bounds(r['date'],kind)
                buckets[a].append(r)
            for a, current in sorted(buckets.items()):
                _, b = period_bounds(a,kind)
                pa,pb = period_bounds(a-timedelta(days=1),kind)
                previous = buckets.get(pa,[])
                cn,pn = (b-a).days+1,(pb-pa).days+1
                cv,pv = sum(r['value'] for r in current),sum(r['value'] for r in previous)
                current_negative_days=sum(r['value']<0 for r in current)
                previous_negative_days=sum(r['value']<0 for r in previous)
                status = '有效'
                if current_negative_days:
                    status='本期含负值'
                elif previous_negative_days:
                    status='前期含负值'
                elif len(current)!=cn:
                    status='本期不完整'
                elif len(previous)!=pn:
                    status='前期缺失或不完整'
                elif pv==0:
                    status='前期为零'
                elif pv<0:
                    status='前期为负'
                transition = (str(pa.month)+'→'+str(a.month)+'月') if kind=='月' else (
                    '周'+str(pa.weekday()+1)+'→周'+str(a.weekday()+1) if kind=='日' else
                    'Q'+str((pa.month-1)//3+1)+'→Q'+str((a.month-1)//3+1) if kind=='季' else '前期→本期')
                interval=(pa,b)
                if interval not in interval_calendar_cache:
                    known=all(d.year in cal['years'] for d in span(pa,b))
                    special=sum(d in cal['holidays'] or d in cal['work'] for d in span(pa,b)) if known else None
                    interval_calendar_cache[interval]=(known,special)
                known,special=interval_calendar_cache[interval]
                comparisons.append(dict(base,kind=kind,start=a,end=b,previous_start=pa,previous_end=pb,
                    current=cv,previous=pv if previous else None,current_observed=len(current),
                    previous_observed=len(previous),current_days=cn,previous_days=pn,
                    current_negative_days=current_negative_days,previous_negative_days=previous_negative_days,
                    ratio=cv/pv if status=='有效' else None,
                    daily_ratio=(cv/cn)/(pv/pn) if status=='有效' else None,
                    transition=transition,status=status,holiday_adjusted_days=special))
    return comparisons

def summarize_comparisons(rows):
    buckets=defaultdict(list)
    for r in rows:
        if r['status']!='有效':
            continue
        # Each dimension is a separate grouping, never an additional volume to sum.
        categories=[('整体','全部车型')]+[(label,r.get(key,'未提供')) for label,key in
                     [('品牌','brand'),('产品档位','segment'),('能源类型','energy')]]
        for dim,value in categories:
            key=(r['role'],r['metric'],r['stage'],r['kind'],r['transition'],dim,value)
            buckets[key].append(r)
    summaries=[]
    for key,rs in sorted(buckets.items()):
        base=dict(zip(('role','metric','stage','kind','transition','dimension','category'),key))
        bymodel=group(rs,('model',))
        byperiod=group(rs,('start',))
        timeline=[median(x['daily_ratio'] for x in g) for g in byperiod.values()]
        years=defaultdict(list)
        for r in rs:
            years[r['start'].year].append(r)
        year_values=[]
        for y,yr in sorted(years.items()):
            model_values=[median(x['daily_ratio'] for x in g) for g in group(yr,('model',)).values()]
            year_values.append((y,median(model_values)))
        summaries.append(dict(base,pairs=len(rs),models=len(bymodel),periods=len(byperiod),
            years=len(years),ratio=median(median(x['ratio'] for x in g) for g in bymodel.values()),
            daily_ratio=median(median(x['daily_ratio'] for x in g) for g in bymodel.values()),
            p25=quantile(timeline,.25),p75=quantile(timeline,.75),
            above100=mean(mean(x['daily_ratio']>1 for x in g) for g in bymodel.values()),
            yearly='；'.join(str(y)+'年 '+format(v,'.1%') for y,v in year_values),
            evidence='不足跨年验证' if len(years)<2 else '跨年可比较，未控活动与趋势',
            status='少于8个不同周期' if len(byperiod)<8 else '已有8个以上周期，仍需留出验证'))
    return summaries

def correlation(xs, ys):
    if len(xs)<3:
        return None
    ax,ay=mean(xs),mean(ys)
    xx=sum((x-ax)**2 for x in xs); yy=sum((y-ay)**2 for y in ys)
    if xx==0 or yy==0:
        return None
    return sum((x-ax)*(y-ay) for x,y in zip(xs,ys))/math.sqrt(xx*yy)

def lock_lag_analysis(panel):
    """Compare retained deposits with locks at 0--3 days.

    Ordinary ``大定`` is intentionally excluded.  A missing retained-deposit
    series remains unavailable instead of silently changing the denominator.
    """
    results=[]
    for key, rs in group(panel, ('model','stage','cycle')).items():
        locks=group([r for r in rs if r['metric']=='交车锁单'],('source',))
        for metric in ('留存大定',):
            deposits=group([r for r in rs if r['metric']==metric],('source',))
            for (ds,), dr in deposits.items():
                for (ls,), lr in locks.items():
                    x={r['date']:r['value'] for r in dr}
                    y={r['date']:r['value'] for r in lr}
                    # Use one non-negative date set for every lag and one for all daily changes.
                    matched=sorted(d for d in x if x[d]>=0 and
                                   all(d+timedelta(days=k) in y and y[d+timedelta(days=k)]>=0
                                       for k in range(4)))
                    changes=sorted(d for d in x if d-timedelta(days=1) in x and x[d]>=0 and x[d-timedelta(days=1)]>=0 and
                                   all(d+timedelta(days=k) in y and y[d+timedelta(days=k)]>=0
                                       for k in range(-1,4)))
                    for lag in range(4):
                        xs=[x[d] for d in matched]
                        ys=[y[d+timedelta(days=lag)] for d in matched]
                        dx=[x[d]-x[d-timedelta(days=1)] for d in changes]
                        dy=[y[d+timedelta(days=lag)]-y[d+timedelta(days=lag-1)] for d in changes]
                        enough=len(matched)>=14
                        results.append(dict(model=key[0],stage=key[1],cycle=key[2],metric=metric,
                            brand=rs[0].get('brand','未提供'),segment=rs[0].get('segment','未提供'),
                            energy=rs[0].get('energy','未提供'),lag=lag,pairs=len(matched),
                            start=min(matched) if matched else None,end=max(matched) if matched else None,
                            level_corr=correlation(xs,ys) if enough else None,
                            change_corr=correlation(dx,dy) if len(changes)>=14 else None,
                            change_pairs=len(changes),deposit_source=ds,lock_source=ls,
                            status='仅相关性，非转化率或逐单锁单时长' if enough else '不足14个共同日期'))
    return results

def month_progress_analysis(panel):
    """Describe within-month completion without treating missing days as zero."""
    output=[]
    fields=('model','metric','stage','cycle','source')
    for key,records in group(panel,fields).items():
        months=defaultdict(list)
        for row in records:
            months[(row['date'].year,row['date'].month)].append(row)
        for (year,month),items in sorted(months.items()):
            first=date(year,month,1); count=calendar.monthrange(year,month)[1]
            last=date(year,month,count); by_day={r['date']:r['value'] for r in items}
            duplicate=len(by_day)!=len(items)
            complete=not duplicate and all(d in by_day for d in span(first,last))
            negative=any(v<0 for v in by_day.values())
            total=sum(by_day.values()) if not duplicate else None
            status=('同日期重复，需核对' if duplicate else '月份不完整' if not complete
                    else '月份含负值' if negative else '整月为零' if total==0 else '有效')
            for checkpoint in (7,14,21,count):
                cutoff=date(year,month,checkpoint)
                cumulative=sum(v for d,v in by_day.items() if first<=d<=cutoff) if complete and not negative else None
                share=cumulative/total if status=='有效' else None
                output.append(dict(zip(fields,key),year=year,month=month,start=first,end=last,
                    checkpoint=checkpoint,cutoff=cutoff,current=total,cumulative=cumulative,
                    observed=len(by_day),expected_days=count,share=share,
                    calendar_share=checkpoint/count,pace=share/(checkpoint/count) if share is not None else None,
                    status=status))
    return output


def extended_analysis(panel,cal):
    comparisons=period_comparisons(panel,cal)
    model_rows=[]
    for (model,),rs in group(panel,('model',)).items():
        model_rows.append(dict(model=model,**{k:rs[0].get(k,'未提供') for k in ('brand','segment','energy')},
                               stages='、'.join(sorted({r['stage'] for r in rs})),
                               metrics='、'.join(sorted({r['metric'] for r in rs}))))
    return dict(period_comparisons=comparisons,period_summary=summarize_comparisons(comparisons),
                lock_lags=lock_lag_analysis(panel),model_classes=model_rows,
                month_progress=month_progress_analysis(panel))


def sha256_file(path):
    digest=hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda:stream.read(1024*1024),b''):
            digest.update(block)
    return digest.hexdigest()

def load_component(path, name):
    spec=importlib.util.spec_from_file_location(name,path)
    if spec is None or spec.loader is None:
        raise RuntimeError('无法加载分析组件：'+str(path))
    module=importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module



def open_generated_webpage(path):
    """Use the OS default opener without choosing a specific browser."""
    path=Path(path).resolve()
    try:
        if sys.platform == 'win32':
            # A native path uses the same association as double-clicking the HTML;
            # do not send file:// through URL dispatch or choose a fallback browser.
            os.startfile(str(path))
            opened=True
        else:
            opened=webbrowser.open_new_tab(path.as_uri())
    except (OSError, webbrowser.Error) as exc:
        log(f'网页已生成，自动打开失败（{exc}）；请手动打开：{path}')
        return False
    if opened:
        log('已请求系统默认方式打开网页：'+str(path))
    else:
        log('网页已生成，未能启动默认浏览器；请手动打开：'+str(path))
    return opened


def main():
    started=time.perf_counter()
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--input',type=Path,default=ROOT/'output_file/鸿蒙智行销量数据汇总.xlsx')
    ap.add_argument('--out',type=Path,default=None,help='已弃用；为兼容旧命令保留，不再使用该路径')
    ap.add_argument('--calendar',type=Path)
    ap.add_argument('--sample-data',action='store_true',help='将输出标记为样例数据')
    ap.add_argument('--real-data',action='store_true',help=argparse.SUPPRESS)
    a=ap.parse_args()
    src=a.input.resolve()
    # The interactive HTML is the sole generated deliverable.  Keep the
    # historical workbook path out of this run so existing Excel files are
    # never overwritten or mistaken for a fresh result.
    web_final=(ROOT/'output_file'/'鸿蒙智行销量规律分析.html').resolve()
    exporter=SCRIPT_PATH.with_name('export_excel_openpyxl.py')
    calendar_path=SCRIPT_PATH.parent.parent/'generate_html'/'core'/'china_calendar.py'
    config_path=SCRIPT_PATH.with_name('report_config.json')
    rules_path=SCRIPT_PATH.with_name('report_rules.py')
    forecast_path=SCRIPT_PATH.with_name('forecast_insights.py')
    launch_path=SCRIPT_PATH.with_name('launch_patterns.py')
    source_patterns_path=SCRIPT_PATH.with_name('source_patterns.py')
    web_exporter=SCRIPT_PATH.with_name('export_html.py')
    templates=SCRIPT_PATH.with_name('templates')

    log('销量规律分析启动')
    if a.out is not None:
        log('--out参数已弃用，忽略该路径；网页仍写入固定输出文件')
    log('输入文件：'+str(src))
    if src==web_final:
        raise ValueError('输入文件与固定网页输出文件相同，已停止以保护源数据：'+str(src))
    log('分析目的：辅助销量预测；小订、留存大定与交车锁单分别研究')
    required=(('输入工作簿',src),('节假日日历',calendar_path),
              ('网页分组辅助模块',exporter),('报告规则配置',config_path),
              ('规律汇总模块',rules_path),('上市后销量节奏模块',launch_path),('源表扩展规律模块',source_patterns_path),('预测辅助模块',forecast_path),('网页导出器',web_exporter),
              *((f'网页模板 {name}',templates/name) for name in ('patterns.html','patterns.css','patterns.js')))
    missing=[f'{label}不存在：{path}' for label,path in required if not path.is_file()]
    if a.calendar is not None and not a.calendar.is_file():
        missing.append(f'额外节假日日历不存在：{a.calendar.resolve()}')
    if missing:
        raise FileNotFoundError('\n'.join(missing))

    log('正在读取并校验源数据...')
    stage=time.perf_counter()
    digest=sha256_file(src)
    panel,audit=read_panel(src)
    if not panel:
        raise ValueError('源工作簿中没有有效观测，请检查日期、车型、指标值和来源字段')
    log(f'源数据读取及校验完成：有效观测 {len(panel):,} 条，车型 {len({r["model"] for r in panel}):,} 个，用时 {time.perf_counter()-stage:.1f} 秒')
    log('读取工作表：'+'、'.join(sorted({r['sheet'] for r in panel})))
    for key in ('跨来源候选未累加', '累计快照或参考曲线排除', '未来或超过数据判定日期排除',
                '未结束或未知日期状态排除', '数据判定当日未完整排除', '阶段不明确排除',
                '阶段标签与日期字段推导不一致'):
        if audit.get(key): log(f'{key}：{audit[key]:,} 条')
    log(f'数据日期范围：{min(r["date"] for r in panel)} 至 {max(r["date"] for r in panel)}')

    stage=time.perf_counter()
    c=load_calendar(a.calendar)
    data_years=sorted({r['date'].year for r in panel})
    unknown=sorted(set(data_years)-c['years'])
    log('节假日日历已加载：'+('、'.join(map(str,sorted(c['years']))) or '无')+f'，用时 {time.perf_counter()-stage:.1f} 秒')
    if unknown:
        log('提示：数据覆盖年份缺少日历规则：'+'、'.join(map(str,unknown))+'；这些年份不参与普通周/节假日识别')

    log('正在计算周、月、年、节假日和生命周期规律...')
    stage=time.perf_counter()
    result=analyze(panel,c)
    log(f'周、月、年、节假日和生命周期计算完成，用时 {time.perf_counter()-stage:.1f} 秒')
    log('正在计算周期环比、车型分类和锁单滞后...')
    stage=time.perf_counter()
    result.update(extended_analysis(panel,c))
    launch_analysis=load_component(launch_path,'sales_launch_patterns').build_launch_analysis(panel)
    result['launch_periods']=launch_analysis['periods']
    result['launch_pacing']=launch_analysis['comparisons']
    log(f"上市后销量节奏：{len(result['launch_periods'])} 个周期、{len(result['launch_pacing'])} 条相邻比较，留存大定与锁单分别计算；小订仅保留小订期")
    log(f'周期比较、车型分类和锁单滞后计算完成，用时 {time.perf_counter()-stage:.1f} 秒')

    log('正在读取首销D1/D2、转化退订曲线和分时数据...')
    stage=time.perf_counter()
    source_insights=load_component(source_patterns_path,'sales_source_patterns').build_source_insights(src)
    log('源表扩展规律完成：'+str({k:len(source_insights.get(k,[])) for k in ('inventory','early_launch','curves','hourly')})+f'，用时 {time.perf_counter()-stage:.1f} 秒')

    log('正在建立三口径（日/周/月）预测辅助并做滚动回测：小订（仅小订期）、留存大定、锁单...')
    stage=time.perf_counter()
    forecast=load_component(forecast_path,'sales_forecast_insights').build_forecast_insights(panel,c,log=log)
    for table in ('profiles','backtests'):
        for row in forecast.get(table,[]):
            row['grain']={'day':'日','week':'周','month':'月'}.get(row.get('grain'),row.get('grain'))
    log(f'预测辅助及回测计算完成：{len(forecast.get("profiles", [])):,} 组，耗时 {time.perf_counter()-stage:.1f} 秒')

    if sha256_file(src)!=digest:
        raise RuntimeError('分析期间源文件发生变化，已停止导出')
    payload=dict(sample=a.sample_data and not a.real_data,source=str(src),sha256=digest,
                 source_sheets=sorted({r['sheet'] for r in panel}),source_mode='字段级来源' if audit.get('字段级来源表已识别') else '来源候选',audit=audit,calendar_years=sorted(c['years']),calendar_sources=c['sources'],
                 analysis_end=max(r['date'] for r in panel),results=result,forecast=forecast,source_insights=source_insights,
                 calendar_events=[dict(year=y,holiday=n,start=a,end=b) for y,n,a,b in c['events']])
    log('正在准备离线交互网页...')
    web=load_component(web_exporter,'sales_html_exporter')
    html=web.render_html(web.build_web_data(payload,panel))
    web_final.parent.mkdir(parents=True,exist_ok=True)
    log('正在保存离线交互网页...')
    web.export_html(html,web_final,log=log)
    log(f'分析完成，总耗时 {time.perf_counter()-started:.1f} 秒')
    print('网页：'+str(web_final),flush=True)
    print('源文件未修改；结果数量：'+str({k:len(v) for k,v in result.items()}),flush=True)
    open_generated_webpage(web_final)
if __name__=='__main__': main()


